"""Values-only Excel registers using the same calculated snapshot as the PDF."""

import textwrap

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from .calculator_report import project_calculator_report, _source_text
from .pricing_workbook import _serialize_exact


_RED = '941321'
_PALE = 'FFF1CE'
_INK = '332B30'
_NUMBER = '#,##0.00'
_SNAPSHOT = 'Current calculated snapshot. Recalculate in the app and download again to update this register.'


def _cell(sheet, row, column, value, *, heading=False, number_format=_NUMBER):
    cell = sheet.cell(row, column, value)
    # User references and source status/error text must never become executable
    # formulas, hyperlinks, or native Excel error cells.
    if isinstance(value, str):
        cell.data_type = 's'
    cell.font = Font(name='Arial', size=10, color='FFFFFF' if heading else _INK, bold=heading)
    cell.alignment = Alignment(horizontal='center',
                               vertical='center', wrap_text=True)
    cell.number_format = number_format
    if heading:
        cell.fill = PatternFill('solid', fgColor=_RED)
    return cell


def _height(sheet, row, values, widths, *, minimum=27):
    lines = max((sum(max(1, len(textwrap.wrap(line, max(8, int(width - 2)),
                                              replace_whitespace=False))) for line in str(value or '').split('\n'))
                 for value, width in zip(values, widths)), default=1)
    sheet.row_dimensions[row].height = min(409, max(minimum, 15 * lines + 10))


def _band(sheet, row, text, last, *, title=False):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
    cell = _cell(sheet, row, 1, text, heading=title)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    if title:
        cell.font = Font(name='Arial', size=14, color='FFFFFF', bold=True)
    else:
        cell.fill = PatternFill('solid', fgColor=_PALE)
    width = sum(sheet.column_dimensions[get_column_letter(col)].width for col in range(1, last + 1))
    _height(sheet, row, [text], [width], minimum=36 if title else 32)


def _sheet(workbook, name, title, widths):
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.outlinePr.summaryRight = False
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    _band(sheet, 1, title, len(widths), title=True)
    _band(sheet, 2, _SNAPSHOT, len(widths))
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=.25, right=.25, top=.35, bottom=.35, header=.15, footer=.15)
    sheet.oddFooter.right.text = 'Page &P of &N'
    return sheet


def _table(sheet, row, labels, records, widths, *, formats=None, filtered=False):
    for index, label in enumerate(labels, 1):
        _cell(sheet, row, index, label, heading=True).border = Border(right=Side(style='thin', color='FFFFFF'))
    _height(sheet, row, labels, widths, minimum=42)
    first = row
    for values in records:
        row += 1
        for index, value in enumerate(values, 1):
            cell = _cell(sheet, row, index, value, number_format=(formats or {}).get(index, _NUMBER))
            cell.border = Border(right=Side(style='thin', color='E4DBDD'))
            if row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F8F4F4')
        _height(sheet, row, values, widths)
    if filtered:
        sheet.auto_filter.ref = f'A{first}:{get_column_letter(len(labels))}{row}'
        sheet.freeze_panes = f'C{first + 1}'
        sheet.print_title_rows = f'1:{first}'
    return row + 2


def _schedule(data, workbook):
    identity = data['id']
    if identity == 'ductwork':
        labels = ['Line', 'Product', 'Duct dimensions (mm)', 'Length (m)', 'FRL', 'Thickness (mm)',
                  'Thickness basis', 'Duct surface (m²)', 'Net spray bags', 'Wrap material (m²)', 'Roll equivalents', 'Status']
        widths = [9, 26, 23, 16, 18, 17, 19, 19, 18, 20, 18, 65]
        records = [[item['line'], v['C'], v['B'], v['D'], v['E'],
                    item['wrap_layer_mm'] if item['wrap'] else v['L'], 'Per layer' if item['wrap'] else 'Coating',
                    v['K'], 'N/A' if item['wrap'] else v['M'], v['N'] if item['wrap'] else 'N/A',
                    v['O'] if item['wrap'] else 'N/A', v['J'] or 'No calculated status returned']
                   for item in data['rows'] for v in [item['values']]]
    elif identity == 'steel_vermiculite':
        labels = ['Line', 'Location', 'Mark', 'Product', 'Section', 'Quantity', 'Length (m)', 'Published thickness (mm)',
                  'Estimating thickness (mm)', 'Spray surface (m²)', 'Net bags', 'Whole bags per line', 'Status']
        widths = [9, 28, 25, 26, 24, 14, 16, 21, 21, 19, 18, 18, 65]
        records = [[item['line'], v['AA'], v['A'], v['B'], v['F'], v['I'], v['J'], v['O'], v['P'], v['R'], v['T'], v['U'],
                    '\n'.join(str(value) for value in (v['V'], v['W']) if value not in (None, '')) or 'No calculated status returned']
                   for item in data['rows'] for v in [item['values']]]
    else:
        labels = ['Line', 'Mark', 'Location', 'Product', 'Section', 'Design period (min)', 'Critical temperature (°C)', 'Board stack (mm)',
                  'Total thickness (mm)', 'Box reference area (m²)', 'Net board area (m²)', 'Area with waste (m²)',
                  'Sheets per line', 'Status']
        widths = [9, 25, 28, 26, 24, 19, 21, 20, 20, 21, 20, 20, 17, 65]
        records = [[item['line'], v['A'], v['B'], v['C'], v['D'], v['AN'], v['AO'], v['Z'], v['AB'], v['AD'], v['AE'], v['AF'], v['AG'],
                    '\n'.join(str(value) for value in (v['AR'], v['AS']) if value not in (None, '')) or 'No calculated status returned']
                   for item in data['rows'] for v in [item['values']]]
    sheet = _sheet(workbook, 'Schedule', data['title'] + ' — Schedule', widths)
    _band(sheet, 3, f"{len(records)} used schedule items. {data['incomplete_rows']} item(s) have incomplete or unavailable primary quantities.", len(widths))
    _table(sheet, 5, labels, records, widths, formats={1: '0'}, filtered=True)
    if not records:
        _band(sheet, 6, 'No schedule inputs are entered.', len(widths))


def _summary(data, workbook):
    widths = [36] + [19] * 10
    sheet = _sheet(workbook, 'Summary', data['title'] + ' — Excel register', widths)
    _band(sheet, 3, data['basis'], len(widths))
    row = 5
    # Keep the overview compact while accommodating descriptive measure labels.
    for label, value in data['totals']:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        _cell(sheet, row, 1, label).font = Font(name='Arial', size=10, color=_INK, bold=True)
        _cell(sheet, row, 5, value)
        sheet.row_dimensions[row].height = 30
        row += 1
    row += 1
    _band(sheet, row, f"{data['incomplete_rows']} schedule item(s) have incomplete or unavailable primary quantities. Totals retain the source workbook's exclusions; review item statuses before ordering.", len(widths))
    row += 2
    for note in data.get('summary_notes', []):
        if note:
            _band(sheet, row, _source_text(data['id'], 'BOARD SUMMARY', 'A8', note), len(widths))
            row += 1
    for summary in data['summaries']:
        columns = summary['columns']
        _band(sheet, row, summary['title'], len(widths), title=True)
        row += 1
        if summary['note']:
            _band(sheet, row, summary['note'], len(widths))
            row += 1
        records = []
        for item in summary['rows']:
            values = []
            for column in columns:
                value = item['values'][column]
                if data['id'] == 'ductwork' and summary['title'] == 'Product totals' and ((item['row'] < 11 and column in 'EFG') or (item['row'] == 11 and column == 'D')):
                    value = 'N/A'
                values.append(_source_text(data['id'], summary['sheet'], column + str(item['row']), value))
            records.append(values)
        row = _table(sheet, row, [summary['labels'][column] for column in columns], records, widths,
                     formats={6: '0.00%'} if data['id'] == 'steel_vermiculite' else None)
        for product, basis, interpretation in summary.get('qualifications', []):
            _band(sheet, row, '\n'.join(str(value) for value in (product, basis, interpretation) if value not in (None, '')), len(widths))
            row += 1
    _band(sheet, row, 'Source workbook: ' + data['source']['filename'], len(widths))
    _band(sheet, row + 1, 'Source SHA-256: ' + data['source']['sha256'], len(widths))
    sheet.freeze_panes = 'B5'
    sheet.print_title_rows = '1:2'


def _extras(data, workbook):
    # Match the PDF: the unused internal column L is excluded; evidence stays.
    columns = list('ABCDEFGHIJK') + ['M', 'N']
    widths = [9, 25, 26, 18, 15, 18, 18, 18, 16, 28, 19, 20, 55, 90]
    sheet = _sheet(workbook, 'Extra boards', 'EXTRA BOARDS', widths)
    _band(sheet, 3, 'Valid allowances are included in final board stock totals. Incomplete entries remain listed and are excluded by the workbook rules.', len(widths))
    if data['extra_rows']:
        labels = ['Item'] + [data['extra_rows'][0]['labels'][column] for column in columns]
        records = [[item['line']] + [item['values'][column] for column in columns] for item in data['extra_rows']]
        _table(sheet, 5, labels, records, widths, formats={1: '0', 9: '0.00%'}, filtered=True)
    else:
        _band(sheet, 5, 'No extra-board inputs are entered.', len(widths))


def build_calculator_register(calculator_id, inputs=None):
    """Export the current draft's report results without saving or Excel logic."""
    data = project_calculator_report(calculator_id, inputs)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = 'Ceasefire'
    workbook.properties.title = data['title'] + ' — Excel register'
    workbook.properties.description = _SNAPSHOT
    _summary(data, workbook)
    _schedule(data, workbook)
    if calculator_id == 'steel_board':
        _extras(data, workbook)
    for sheet in workbook:
        sheet.print_options.horizontalCentered = True
        sheet.print_area = sheet.dimensions
    return _serialize_exact(workbook)
