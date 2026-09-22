"""Penetration exports project one calculated snapshot, never recalculate totals."""

from io import BytesIO

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.platypus import SimpleDocTemplate, Spacer

from .calculator_register import _sheet, _table
from .pricing_workbook import _serialize_exact
from .report import ROOT, _Report, _company_header, _register_fonts, _number, _numeric, _LINE, _MUTED


SUMMARY = (
    ('Labour', 'labour'), ('Materials', 'materials'),
    ('Grand total', 'grand_total'),
    ('Total days', 'total_days'),
)
SCHEDULE = (
    ('Service', 'K'), ('Quantity', 'O'), ('Substrate', 'P'),
    ('Labour', 'F'), ('Materials', 'G'), ('Item total', 'H'),
)
LABOUR_INPUTS = {'pipe_labour_hours': 'pipe_hours'}


def _visible_outputs(definition):
    return [field for field in definition['output_fields']
            if field['column'] not in ('B', 'C', 'D', 'E', 'BI', 'BJ', 'BK')]


def _value(row, column):
    return row['outputs'].get(column, row['inputs'].get(column))


def _input_value(row, column):
    """Display the evaluated app allowance without filling its raw input."""
    raw = row['inputs'].get(column)
    if column not in LABOUR_INPUTS or raw not in (None, ''):
        return raw
    value = row.get('labour_policy', {}).get(LABOUR_INPUTS[column], row.get('input_defaults', {}).get(column))
    if value is None:
        if column == 'pipe_labour_hours' and row['inputs'].get('Y') in (None, ''):
            return 'Not applicable'
        return 'Unavailable'
    return value


def _input_basis(row, column):
    if column not in LABOUR_INPUTS:
        return ''
    manual = row['inputs'].get(column) not in (None, '')
    if column == 'pipe_labour_hours' and row['inputs'].get('Y') in (None, ''):
        return ('Manual' if manual else 'Automatic') + '; no collar selected'
    if not manual and _input_value(row, column) == 'Unavailable':
        return 'Manual value required'
    return 'Manual' if manual else 'Automatic'


def _input_records(row, fields):
    """Present paired dimensions once while preserving their source columns."""
    fields_by_column = {field['column']: field for field in fields}
    for field in fields:
        if field.get('paired_into'):
            continue
        column = field['column']
        paired_column = field.get('paired_column')
        if paired_column:
            paired = fields_by_column.get(paired_column, {})
            first, second = _input_value(row, column), _input_value(row, paired_column)
            if first in (None, '') and second in (None, ''):
                value = None
            else:
                value = f'{_display(first, field)} x {_display(second, paired)}'
            yield field, value, ''
        else:
            yield field, _input_value(row, column), _input_basis(row, column)


def _display(value, field=None):
    if value is None or value == '':
        return ''
    if _numeric(value):
        field = field or {}
        return _number(value, percent=field.get('format') == 'percent',
                       money=field.get('format') == 'currency')
    return str(value)


def _excel_format(field):
    return {'percent': '0.00%', 'currency': '$#,##0.00'}.get(field.get('format'), '#,##0.00')


def render_penetration_pdf(result, definition, project_details):
    """Show project details, totals and the compact calculated schedule."""
    _register_fonts()
    report = _Report({})
    # Let long schedule tables split on the current page. Keeping an entire
    # table with its heading can otherwise strand the schedule heading alone.
    report.styles['subheading'].keepWithNext = False
    report.styles['section'].keepWithNext = True
    width, height = landscape(A4)
    margin = 32
    content = width - 2 * margin
    logo = ImageReader(str(ROOT / 'static' / 'ceasefire-logo.png'))
    report.story.append(report.p('Firestopping estimate', 'title'))
    details = [('Project No.', project_details.get('project_no', '')),
               ('Client', project_details.get('client', '')),
               ('Site Address', project_details.get('site_address', ''))]
    report.story.append(report.table(['Project details', 'Value'],
        [[report.p(label, 'cell'), report.p(value, 'cell')] for label, value in details],
        [content * .25, content * .75]))
    report.story.append(Spacer(1, 12))
    report.story.append(report.table(['Estimate', 'Amount / quantity'], [
        [report.p(label, 'cell'), report.p(_display(result['summary'].get(key),
            {'format': 'currency' if key != 'total_days' else 'number'}), 'numeric')]
        for label, key in SUMMARY], [content * .65, content * .35]))
    if result.get('errors'):
        report.story.append(report.p('Some workbook results are unavailable. Review the calculation errors below.', 'alert'))
    report.story.append(report.p('Schedule', 'subheading'))
    fields_by_column = {field['column']: field for field in definition['row_fields'] + definition['output_fields']}
    widths = [content * part for part in (.045, .245, .09, .2, .14, .14, .14)]
    report.story.append(report.table(['Line'] + [label for label, _ in SCHEDULE], [
        [report.p(str(index), 'cell')] + [report.p(_display(_value(row, col), fields_by_column[col]), 'cell') for _, col in SCHEDULE]
        for index, row in enumerate(result['rows'], 1)], widths, compact=True))
    for index, row in enumerate(result['rows'], 1):
        for error in row.get('errors', []):
            report.story.append(report.p(f"Line {index} · {error['cell']}: {error['message']}", 'alert'))
    for error in result.get('errors', []):
        if not error.get('row_id'):
            report.story.append(report.p(f"{error['cell']}: {error['message']}", 'alert'))

    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=margin, rightMargin=margin,
        topMargin=91, bottomMargin=43, pageCompression=1,
        title='Ceasefire - Firestopping estimate', author='Ceasefire')

    def decorate(canvas, doc):
        canvas.saveState()
        _company_header(canvas, logo, width, height, margin, 'FIRESTOPPING ESTIMATE', label_below_logo=True)
        canvas.setStrokeColor(_LINE)
        canvas.line(margin, 32, width - margin, 32)
        canvas.setFont('CeasefireVera', 7)
        canvas.setFillColor(_MUTED)
        canvas.drawString(margin, 20, 'Ceasefire ESTIMATOR | Firestopping Estimator')
        canvas.drawRightString(width - margin, 20, f'Page {doc.page}')
        canvas.restoreState()

    document.build(report.story, onFirstPage=decorate, onLaterPages=decorate)
    return output.getvalue()


def build_penetration_register(result, definition, project_details):
    """Values-only Excel export; user text cannot become formulas or links."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = 'Ceasefire'
    workbook.properties.title = 'Firestopping estimate'
    summary = _sheet(workbook, 'Summary', 'FIRESTOPPING ESTIMATE', [36, 64])
    row = _table(summary, 4, ['Project details', 'Value'],
                 [[label, project_details.get(key, '')] for label, key in
                  [('Project No.', 'project_no'), ('Client', 'client'), ('Site Address', 'site_address')]], [36, 64])
    first = row + 1
    row = _table(summary, row, ['Estimate', 'Amount / quantity'],
                 [[label, result['summary'].get(key)] for label, key in SUMMARY], [36, 64])
    for offset, (_, key) in enumerate(SUMMARY):
        summary.cell(first + offset, 2).number_format = _excel_format({'format': 'number' if key == 'total_days' else 'currency'})
    widths = [9, 35, 14, 28, 20, 20, 22]
    schedule = _sheet(workbook, 'Schedule', 'PENETRATION SCHEDULE', widths)
    fields_by_column = {field['column']: field for field in definition['row_fields'] + definition['output_fields']}
    _table(schedule, 4, ['Line'] + [label for label, _ in SCHEDULE],
           [[index] + [_value(row, column) for _, column in SCHEDULE]
            for index, row in enumerate(result['rows'], 1)], widths,
           formats={index: _excel_format(fields_by_column[column]) for index, (_, column) in enumerate(SCHEDULE, 2)}, filtered=True)
    settings_fields, settings_schema = definition.get('global_fields', []), definition.get('settings', {})
    if settings_fields or settings_schema:
        settings = _sheet(workbook, 'Settings', 'FIRESTOPPING SETTINGS', [36, 70, 18, 18])
        row = _table(settings, 4, ['Setting', 'Value', 'Units'],
                     [[field['label'], result['draft']['globals'].get(field['column']), field.get('units', '')]
                      for field in settings_fields], [36, 70, 18])
        for offset, field in enumerate(settings_fields, 5):
            settings.cell(offset, 2).number_format = _excel_format(field)
        routes = result['draft']['globals'].get('service_routes', {})
        if settings_schema.get('service_routes'):
            row = _table(settings, row + 1, ['Service-tab routing', 'Service types separated by semicolons'],
                         [[route['label'], '; '.join(routes.get(route['key'], []))]
                          for route in settings_schema['service_routes']], [36, 106])
        bands = result['draft']['globals'].get('labour_bands', {})
        band_rows = [[band['label'], item['maximum'], band['units'], item['hours']]
                     for band in settings_schema.get('labour_bands', [])
                     for item in bands.get(band['key'], [])]
        if band_rows:
            first_band = row + 2
            _table(settings, row + 1, ['Task-hour band', 'Up to', 'Unit', 'Hours'],
                   band_rows, [36, 28, 18, 18])
            for offset in range(first_band, first_band + len(band_rows)):
                settings.cell(offset, 2).number_format = '#,##0.###############'
                settings.cell(offset, 4).number_format = '#,##0.###############'
    for title, fields, key in [('Inputs', definition['row_fields'], 'inputs'),
                               ('Calculated detail', _visible_outputs(definition), 'outputs')]:
        widths = [9, 42, 85, 20] + ([36] if key == 'inputs' else [])
        sheet = _sheet(workbook, title, title.upper(), widths)
        records = [
            (index, field, value, basis)
            for index, row in enumerate(result['rows'], 1)
            for field, value, basis in (
                _input_records(row, fields) if key == 'inputs' else
                ((item, row[key].get(item['column']), '') for item in fields))
        ]
        _table(sheet, 4, ['Line', 'Parameter', 'Value', 'Units'] + (['Basis'] if key == 'inputs' else []),
               [[index, field['label'], value, field.get('units', '')] + ([basis] if key == 'inputs' else [])
                for index, field, value, basis in records], widths, filtered=True)
        for offset, (_, field, _, _) in enumerate(records, 5):
            sheet.cell(offset, 3).number_format = _excel_format(field)
    if result.get('errors'):
        widths = [24, 22, 85]
        sheet = _sheet(workbook, 'Calculation errors', 'CALCULATION ERRORS', widths)
        _table(sheet, 4, ['Line ID', 'Cell', 'Error'],
               [[error.get('row_id', ''), error['cell'], error['message']] for error in result['errors']], widths)
    for sheet in workbook:
        sheet.print_area = sheet.dimensions
    return _serialize_exact(workbook)
