"""Values-only pricing workbooks: editable lists, with no executable spreadsheet logic.

An import is a proposed complete catalogue replacement. Persistence belongs to
the caller; parsing never writes configuration or saved quotes.
"""

from copy import deepcopy
import csv
import hashlib
from io import BytesIO, StringIO
import math
import posixpath
import re
from uuid import uuid4
import xml.etree.ElementTree as ET
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import column_index_from_string, coordinate_to_tuple, get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.views import Selection

from .catalog import ValidationError, effective_catalog, validate_configuration, yield_unit


MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_ROWS = 5000
MAX_PARTS = 200
XML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_TEXT_ESCAPE = re.compile(r"_x([0-9A-Fa-f]{4})_")
PROPERTY_HEADERS = {
    "Weight": "weight", "Width (mm)": "width_mm", "Length (mm)": "length_mm",
    "Thickness (mm)": "thickness_mm", "Area (m²)": "sqm", "Diameter (mm)": "diameter_mm",
    "Metres per tube": "metres_per_tube", "Area per box (m²)": "sqm_per_box",
    "Area per drum (m²)": "sqm_per_drum", "Volume (L)": "volume_litres",
    "Volume (mL)": "volume_ml",
}
INVENTORY_HEADERS = (
    "Inventory ID", "Item code", "Product name", "Sales description", "Pricing mode",
    "Supplier price", "Markup", "Sell price", *PROPERTY_HEADERS,
)
RATE_HEADERS = (
    "Rate ID", "Group", "Inventory ID", "Rate name", "Price source",
    "Unit sell rate", "Yield type", "Yield",
)
COMBINED_SHEET = "Inventory & Rates"
COMBINED_HEADERS = (
    "Row type", "Name", "Item code", "Group", "Supplier price", "Markup",
    "Sell price / rate", "Price source", "Yield type", "Yield", "Pricing mode",
    "Sales description", "Inventory ID", "Rate ID", "Use order", *PROPERTY_HEADERS,
)
LEGACY_COMPACT_HEADERS = (
    "Item code", "Product name", "Supplier price", "Markup", "Sell price",
    "Group", "Selection name", "Price source", "Sell rate", "Yield type", "Yield",
    "Pricing mode", "Sales description", "Inventory ID", "Rate ID", "Use order", *PROPERTY_HEADERS,
)
COMPACT_HEADERS = (*LEGACY_COMPACT_HEADERS[:11], "Yield unit", *LEGACY_COMPACT_HEADERS[11:])
LEGACY_USE_VECTOR_HEADERS = ("Group", "Selection name", "Price source", "Sell rate", "Yield type", "Yield", "Rate ID", "Use order")
USE_VECTOR_HEADERS = (*LEGACY_USE_VECTOR_HEADERS, "Yield unit")
# All layouts feed the same existing pricing parser. Product values have one
# owner; a use links to that owner explicitly, never by its physical position.
_INVENTORY_COLUMNS = {name: {"Product name": "Name", "Sell price": "Sell price / rate"}.get(name, name)
                      for name in INVENTORY_HEADERS}
_USE_COLUMNS = {name: {"Rate name": "Name", "Unit sell rate": "Sell price / rate"}.get(name, name)
                for name in RATE_HEADERS}
_COMPACT_USE_COLUMNS = {name: {"Rate name": "Selection name", "Unit sell rate": "Sell rate"}.get(name, name)
                        for name in RATE_HEADERS if name != "Inventory ID"}


def _pack_uses(values):
    """One native scalar, or a positional semicolon CSV record without rounding."""
    if not values:
        return None
    if len(values) == 1 and not isinstance(values[0], str):
        return values[0]
    buffer = StringIO(newline="")
    csv.writer(buffer, delimiter=";", lineterminator="\r\n").writerow(values)
    value = buffer.getvalue()[:-2]
    if len(value) > 32767:
        raise ValidationError("A product's use list exceeds Excel's 32,767-character cell limit.")
    return value


def _unpack_uses(value, label):
    if value in (None, ""):
        return [None]
    if not isinstance(value, str):
        return [value]
    if len(value) > 32767:
        raise ValidationError(f"{label}: a use list must have at most 32,767 characters.")
    # csv.reader accepts quotes embedded in bare fields. Reject those as well
    # as unterminated quotes so malformed lists cannot shift use associations.
    field = r'(?:"(?:[^"]|"")*"|[^";\r\n]*)'
    if not re.fullmatch(field + r'(?:;' + field + r')*', value):
        raise ValidationError(f"{label}: use valid semicolon-separated CSV with doubled quotes inside quoted entries.")
    try:
        rows = list(csv.reader(StringIO(value, newline=""), delimiter=";", strict=True))
    except csv.Error as error:
        raise ValidationError(f"{label}: invalid semicolon-separated CSV.") from error
    if len(rows) != 1:
        raise ValidationError(f"{label}: a use list must contain one CSV record.")
    return rows[0]


def _use_number(value, label):
    if value in (None, "") or isinstance(value, str) and not value.strip():
        return None
    if isinstance(value, str):
        if not re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?', value.strip()):
            raise ValidationError(f"{label}: each entry must be a number or an empty slot.")
        try:
            return float(value)
        except ValueError as error:
            raise ValidationError(f"{label}: invalid numeric entry.") from error
    return value


def _text(value, label, *, optional=False, limit=1000):
    if optional and value in (None, ""):
        return ""
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise ValidationError(f"{label} must be nonempty text of at most {limit} characters.")
    if any(ord(character) < 32 and character not in "\n\r\t" for character in value):
        raise ValidationError(f"{label} contains unsupported control characters.")
    return value


def _excel_precision(value, previous):
    """Excel saves 15 significant digits; preserve unchanged stored precision."""
    if (isinstance(value, (int, float)) and not isinstance(value, bool)
            and isinstance(previous, (int, float)) and not isinstance(previous, bool)
            and format(value, ".15g") == format(previous, ".15g")):
        return previous
    return value


def _property(value, label):
    if isinstance(value, str):
        if len(value) > 1000:
            raise ValidationError(f"{label} must be at most 1000 characters.")
        return value
    return _number(value, label, optional=True)


def _number(value, label, *, optional=False, minimum=0):
    if optional and value in (None, ""):
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValidationError(f"{label} must contain a number, not text or a formula.")
    if not math.isfinite(value) or abs(value) > 1e12 or value < minimum:
        raise ValidationError(f"{label} must be finite, at least {minimum}, and at most 1 trillion.")
    return value


def _identifier(value, label, prefix, seen):
    if value in (None, ""):
        value = f"{prefix}_{uuid4().hex}"
    if isinstance(value, int) and not isinstance(value, bool):
        value = str(value)
    value = _text(value, label, limit=120)
    if not re.fullmatch(r"[A-Za-z0-9_.:-]+", value):
        raise ValidationError(f"{label} may use letters, digits, underscores, dots, colons and hyphens only.")
    if value in seen:
        raise ValidationError(f"Duplicate {label}: {value}.")
    seen.add(value)
    return value


def _rate_mode(rate, configuration):
    if "price" in configuration.get("rates", {}).get(rate["id"], {}):
        return "override"
    return rate.get("price_mode", "inventory" if rate.get("inventory_id") else "override")


def _yield_type(value, uses_yield):
    if not uses_yield:
        return "Not used"
    return "Blank" if value is None else "Empty text" if value == "" else "Number"


def _format_sheet(sheet, headers, widths, numeric_columns=(), percent_columns=(), *, freeze_panes="E2"):
    # Reassigning a two-axis freeze in openpyxl retains the former pane
    # selections. Start clean and place each selection inside its actual pane.
    view = sheet.sheet_view
    view.pane = None
    view.selection = [Selection()]
    sheet.freeze_panes = freeze_panes
    row, column = coordinate_to_tuple(freeze_panes)
    if row > 1 and column > 1:
        positions = (("topRight", f"{get_column_letter(column)}1"),
                     ("bottomLeft", f"A{row}"), ("bottomRight", freeze_panes))
    elif row > 1:
        positions = (("bottomLeft", freeze_panes),)
    elif column > 1:
        positions = (("topRight", freeze_panes),)
    else:
        positions = ((None, "A1"),)
    view.selection = [Selection(pane=pane, activeCell=cell, sqref=cell) for pane, cell in positions]
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 36
    # Styles are immutable once assigned. Reuse them across cells instead of
    # constructing tens of thousands of identical validated style objects.
    header_fill = PatternFill("solid", fgColor="8D1726")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=11, color="174D8D")
    centred = Alignment(horizontal="center", vertical="center", wrap_text=True)
    band_fills = {0: PatternFill("solid", fgColor="F0F5FA"), 1: PatternFill("solid", fgColor="FFFFFF")}
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = centred
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 31
        for cell in row:
            cell.font = body_font
            cell.alignment = centred
            cell.fill = band_fills[cell.row % 2]
            if cell.column in numeric_columns:
                cell.number_format = '#,##0.00;[Red](#,##0.00);0.00'
            if cell.column in percent_columns:
                cell.number_format = '0.00%;[Red](0.00%);0.00%'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"


def _validation(sheet, column, choices, *, max_rows=MAX_ROWS, allow_blank=True):
    validation = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=allow_blank)
    validation.errorTitle = "Choose a listed value"
    validation.error = "Use one of the values shown in the dropdown."
    validation.showErrorMessage = True
    validation.errorStyle = "stop"
    sheet.add_data_validation(validation)
    validation.add(f"{column}2:{column}{max_rows + 1}")


def _serialize_exact(workbook, *, escape_text=False):
    """Retain binary-float round trips instead of openpyxl's 16-digit formatting.

    Cells remain ordinary numeric OOXML cells. Excel can display fewer digits;
    importing the generated file directly must preserve the estimator's values.
    Stored strings also retain their exact line breaks with either XML writer.
    """
    stream = BytesIO()
    workbook.save(stream)
    numbers = {
        f"xl/worksheets/sheet{index}.xml": {
            cell.coordinate: repr(cell.value)
            for row in sheet for cell in row
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
        }
        for index, sheet in enumerate(workbook, 1)
    }
    strings = {
        f"xl/worksheets/sheet{index}.xml": {
            cell.coordinate: cell.value for row in sheet for cell in row
            if cell.data_type == "s" and isinstance(cell.value, str)
        }
        for index, sheet in enumerate(workbook, 1)
    }
    output = BytesIO()
    with ZipFile(stream) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename in numbers:
                root = ET.fromstring(content)
                for cell in root.iter(f"{{{XML_NS}}}c"):
                    if cell.attrib.get("r") in numbers[item.filename]:
                        cell.find(f"{{{XML_NS}}}v").text = numbers[item.filename][cell.attrib["r"]]
                    text = cell.find(f"{{{XML_NS}}}is/{{{XML_NS}}}t")
                    if text is not None and cell.attrib.get("r") in strings[item.filename]:
                        # Without lxml, openpyxl emits literal CR characters.
                        # Restore text from the workbook before escaping it:
                        # the XML parse above has already normalized raw CR.
                        value = strings[item.filename][cell.attrib["r"]]
                        # Pricing reads Excel's ST_Xstring escapes. Protect
                        # literal escape-looking text before Excel reads it.
                        # Other users of this shared serializer keep their
                        # existing text contract unless they opt in.
                        text.text = _TEXT_ESCAPE.sub(lambda match: "_x005F_" + match[0][1:], value) if escape_text else value
                # XML normalizes literal CR to LF on read. Preserve carriage
                # returns in quoted use names as character references.
                content = ET.tostring(root, encoding="utf-8", xml_declaration=True).replace(b"\r", b"&#13;")
            target.writestr(item.filename, content)
    return output.getvalue()


def export_pricing_workbook(configuration):
    """Export one row per product with parallel use lists, including draft edits."""
    data = effective_catalog(configuration)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = COMBINED_SHEET
    sheet.append(COMPACT_HEADERS)
    uses = {}
    for group, rates in data["rate_groups"].items():
        for order, rate in enumerate(rates, 1):
            uses.setdefault(rate.get("inventory_id"), []).append((group, order, rate))

    def append(product, records):
        vectors = {header: [] for header in USE_VECTOR_HEADERS}
        for group, order, rate in records:
            value = rate.get("yield")
            use = {"Selection name": rate["name"], "Group": group,
                   "Sell rate": rate["price"], "Price source": _rate_mode(rate, configuration).title(),
                   "Yield type": _yield_type(value, bool(data["rate_group_rules"][group]["yield_column"])),
                   "Yield": value if isinstance(value, (int, float)) else None,
                   "Yield unit": yield_unit(group, rate),
                   "Rate ID": rate["id"], "Use order": order}
            for header in USE_VECTOR_HEADERS:
                vectors[header].append(use[header])
        values = {**product, **{header: _pack_uses(values) for header, values in vectors.items()}}
        # This derived display is not one of the editable positional vectors.
        # A product with two area-based uses only needs one unit label.
        units = dict.fromkeys(unit for unit in vectors["Yield unit"] if unit)
        values["Yield unit"] = "; ".join(units) or None
        sheet.append([values.get(header) for header in COMPACT_HEADERS])

    for item in data["inventory"]:
        append({"Product name": item["name"], "Item code": item.get("item_code", ""),
                "Supplier price": item.get("supplier_price"), "Markup": item["markup"],
                "Sell price": item["sales_price"],
                "Pricing mode": "Supplier markup" if item["pricing_mode"] == "supplier_markup" else "Manual",
                "Sales description": item["sales_description"], "Inventory ID": item["id"],
                **{header: item.get("properties", {}).get(key) for header, key in PROPERTY_HEADERS.items()}},
               uses.get(item["id"], []))
    for group, order, rate in uses.get(None, []):
        append({}, [(group, order, rate)])

    widths = dict(zip("ABCDEFGHIJKLMNOPQ", (15, 43, 17, 14, 19, 26, 43, 20, 20, 21, 20, 24, 21, 48, 21, 28, 18)))
    widths.update({get_column_letter(column): 19 for column in range(18, 29)})
    _format_sheet(sheet, COMPACT_HEADERS, widths, numeric_columns=(3, 4, 5, 9, 11, *range(18, 29)), percent_columns=(4,), freeze_panes="C2")
    sheet.sheet_properties.outlinePr.summaryRight = False
    use_fill = PatternFill("solid", fgColor="F0F5FA")
    inventory_fill = PatternFill("solid", fgColor="FFF0DE")
    unused_fill = PatternFill("solid", fgColor="ECECEC")
    product_fonts = {bold: Font(name="Calibri", size=11, color="174D8D", bold=bold) for bold in (True, False)}
    editable = Protection(locked=False)
    readonly = Protection(locked=True)
    unit_font = Font(name="Calibri", size=11, color="536271")
    # Column defaults also unlock cells in new product rows. Existing cells
    # have explicit styles, so apply their protection individually below.
    for column, header in enumerate(COMPACT_HEADERS, 1):
        sheet.column_dimensions[get_column_letter(column)].protection = readonly if header == "Yield unit" else editable
    # Optional dimensions share one editable column style, including new rows.
    sheet.column_dimensions.group("R", "AB", outline_level=1, hidden=True)
    sheet.column_dimensions["Q"].collapsed = True
    sheet.protection.sheet = True
    sheet.protection.autoFilter = False
    sheet.protection.sort = False
    sheet.protection.insertRows = False
    sheet.protection.deleteRows = False
    sheet.protection.formatRows = False
    sheet.protection.formatColumns = False
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = False
    for row in sheet.iter_rows(min_row=2):
        inventory = row[14].value is not None
        has_uses = row[5].value is not None
        for header, cell in zip(COMPACT_HEADERS, row):
            cell.protection = readonly if header == "Yield unit" else editable
            use_cell = header in USE_VECTOR_HEADERS
            cell.fill = use_fill if use_cell else inventory_fill
            if use_cell and not has_uses or not use_cell and not inventory:
                cell.fill = unused_fill
            if header == "Product name":
                cell.font = product_fonts[inventory]
            if header == "Yield unit":
                cell.font = unit_font
        # Multi-use CSV remains visible at normal zoom; account for explicit
        # newlines and conservative text width without changing stored values.
        lines = max(sum(max(1, math.ceil(len(line) / max(1, widths[cell.column_letter] - 3)))
                        for line in str(cell.value or "").splitlines() or [""])
                    for cell in row[:17])
        sheet.row_dimensions[row[0].row].height = min(409, max(31, 16 * lines + 8))
        sheet.cell(row[0].row, 17).number_format = "0"
    comments = {
        "F": "The eight editable use columns are parallel semicolon-separated lists. First entries belong together, then second entries, and so on. Keep every editable list the same length, including empty slots. Yield unit is read-only.",
        "G": 'Quote entries containing semicolons, quotes or newlines. Double quotes inside quoted entries: "Name; with ""quotes""". Spaces are significant; do not add separator padding.',
        "I": "Editing a sell rate creates an Override. Set its matching Price source to Inventory to restore the linked product price.",
        "J": "Number requires its matching Yield. Blank and Empty text preserve distinct empty-value behavior. Not used applies to groups without yields. Two empty yield slots are a single semicolon.",
        "L": "Read-only calculation unit: m² / unit for area coverage, m / unit for mastic, and blank for groups without yields. Repeated units appear once. This is not a positional use list; units follow Group on import.",
        "O": "Keep existing Inventory IDs. A new product with a blank ID receives one shared by its uses in this row. Standalone rates leave all product fields, including Inventory ID, blank.",
        "P": "Keep existing Rate IDs in their matching positions. An empty slot creates a new use identity; row order does not identify rates.",
        "Q": "Use unique positive integers within each Group. Empty slots append new choices; clear or update the matching slot when moving a use to another Group.",
    }
    for column, text in comments.items():
        sheet[f"{column}1"].comment = Comment(text, "Ceasefire")
    _validation(sheet, "M", ("Supplier markup", "Manual"), max_rows=2 * MAX_ROWS, allow_blank=True)

    instructions = workbook.create_sheet("Instructions")
    for row in [
        ["CEASEFIRE INVENTORY & RATES", "How to update the combined pricing library"],
        ["Save changes", "Import previews this entire workbook. Review additions, removals and updates, then Save pricing in ESTIMATOR. Existing saved quotes retain their own products and prices."],
        ["One row per product", "Product fields appear once. Group, Selection name, Price source, Sell rate, Yield type, Yield, Rate ID and Use order contain eight editable parallel semicolon-separated use lists. Keep their entry counts aligned. Yield unit is read-only and follows Group on import. Optional product properties are in columns R:AB."],
        ["Editing use lists", 'First entries belong to the first use, second entries to the second use. Preserve empty slots: two blank yields are ; . Quote entries containing semicolons, quotes or newlines, and double quotes inside quoted entries. Example: "Name; with ""quotes""";Second name. Spaces are preserved; do not add padding.'],
        ["Complete replacement", "Keep the sheet and exact headers. To remove a use, remove its entry from all eight editable lists. To remove a product, clear its editable cells in A:K and M:AB; leave the locked unit cell. Filtered or hidden rows still import. Keep existing IDs; Use order preserves dropdown order."],
        ["Add products and uses", "Add a product row with its pricing and aligned use lists. Blank Inventory ID allocates one product identity linked to all uses on that row; blank Rate ID slots allocate new uses. For standalone rates leave every product field blank and use Override pricing. Selection name supplies the standalone label."],
        ["Supplier markup", "On product rows enter supplier price and markup (30% means 0.30). Changing either recalculates sell price as supplier × (1 + markup). Unchanged inputs retain the existing stored sell price exactly."],
        ["Manual pricing", "Choose Manual to enter the product Sell price directly. Supplier price may be blank; markup remains information. Each use's Price source is Inventory for its linked price or Override for an independent price."],
        ["Rates and yields", "Price source is Inventory or Override per entry. Editing a Sell rate creates an Override; choose Inventory to restore its link. Yield type Number requires a nonnegative Yield; Blank and Empty text retain distinct empty values; Not used applies to groups without yields. Yield unit is locked, shows each distinct calculation unit once and follows Group on import."],
        ["Names and groups", "Selection name is the Estimator dropdown choice. Group places it in the existing calculation category. Names must be unique within a Group. These categories reproduce the original selections; they do not establish technical suitability."],
        ["Values only", "Use values, not formulas, macros or external links. Single numeric entries are numbers; multiple entries use exact numeric text separated by semicolons. Prices are a snapshot; ESTIMATOR evaluates edits on import. Earlier Inventory/Use-row and separate Inventory/Rates templates remain supported."],
        ["Limits", f"Maximum {MAX_ROWS:,} products and {MAX_ROWS:,} uses after expansion, 5 MB file, 32,767 characters per cell and numeric magnitude up to 1 trillion. Prices/yields cannot be negative; markup cannot be below -100%. All use cells blank means no uses."],
        ["Group keys", "Use these exact Group keys. New group definitions are not introduced by importing a pricing file."],
        *[[group, "Yield required" if rule["yield_column"] else "Yield not used"]
          for group, rule in data["rate_group_rules"].items()],
    ]:
        instructions.append(row)
    _format_sheet(instructions, (), {"A": 29, "B": 115}, freeze_panes="A2")
    for row in range(2, 13):
        instructions.row_dimensions[row].height = 61
    instructions.auto_filter.ref = None
    # Never allow a product label beginning '=' to become an Excel formula.
    for page in workbook:
        for row in page:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    return _serialize_exact(workbook, escape_text=True)


def _preflight(payload, filename):
    if not isinstance(filename, str) or not filename.lower().endswith(".xlsx"):
        raise ValidationError("Import an .xlsx pricing workbook exported by ESTIMATOR.")
    if not isinstance(payload, bytes) or not payload or len(payload) > MAX_FILE_BYTES:
        raise ValidationError("Pricing workbooks must be nonempty and no larger than 5 MB.")
    try:
        with ZipFile(BytesIO(payload)) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if len(entries) > MAX_PARTS or len(names) != len(set(names)):
                raise ValidationError("Workbook has too many or duplicate ZIP parts.")
            if sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                raise ValidationError("Workbook expands beyond the 20 MB limit.")
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                raise ValidationError("This is not an Excel .xlsx workbook.")
            workbook_xml = archive.read("xl/workbook.xml")
            if b"\x00" in workbook_xml or b"<!DOCTYPE" in workbook_xml.upper() or b"<!ENTITY" in workbook_xml.upper():
                raise ValidationError("XML declarations with entities are not supported.")
            sheet_names = {element.attrib.get("name") for element in ET.fromstring(workbook_xml).iter()
                           if element.tag.rsplit("}", 1)[-1] == "sheet"}
            max_rows = 2 * MAX_ROWS if COMBINED_SHEET in sheet_names else MAX_ROWS
            for entry in entries:
                name = entry.filename.lower()
                if entry.flag_bits & 1 or ".." in name.split("/") or name.startswith(("/", "\\")):
                    raise ValidationError("Encrypted or unsafe workbook ZIP parts are not supported.")
                if any(part in name for part in ("vbaproject", "macrosheet", "externallink", "activex", "embeddings/")):
                    raise ValidationError("Macros, embedded objects and external links are not allowed in pricing workbooks.")
                if not name.endswith((".xml", ".rels")):
                    continue
                content = archive.read(entry)
                if b"\x00" in content or b"<!DOCTYPE" in content.upper() or b"<!ENTITY" in content.upper():
                    raise ValidationError("XML declarations with entities are not supported.")
                root = ET.fromstring(content)
                worksheet_rows = 0
                for element in root.iter():
                    local = element.tag.rsplit("}", 1)[-1]
                    if local == "Relationship":
                        target = element.attrib.get("Target", "")
                        relation = element.attrib.get("Type", "").lower()
                        if element.attrib.get("TargetMode", "").lower() == "external" or "://" in target or target.startswith("//"):
                            raise ValidationError("External links are not allowed in pricing workbooks.")
                        if any(token in relation for token in ("vbaproject", "macrosheet", "oleobject", "activex")):
                            raise ValidationError("Macros and embedded objects are not supported.")
                    if local in {"Override", "Default"} and any(token in element.attrib.get("ContentType", "").lower() for token in ("macroenabled", "vba", "macrosheet")):
                        raise ValidationError("Macro-enabled workbooks are not supported.")
                    if local == "f":
                        raise ValidationError("Excel formulas are not supported. Replace formulas with values before importing.")
                    if name.startswith("xl/worksheets/") and local == "row":
                        worksheet_rows += 1
                        row_index = element.attrib.get("r", str(worksheet_rows))
                        if (worksheet_rows > max_rows + 1 or not row_index.isdigit()
                                or not 1 <= int(row_index) <= max_rows + 1):
                            raise ValidationError(f"A pricing sheet must have at most {max_rows:,} rows.")
                    if name.startswith("xl/worksheets/") and local == "c":
                        address = element.attrib.get("r", "")
                        match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]*)", address)
                        # Schedule imports share this ZIP/XML check and can use
                        # columns through Z. Each importer validates its own
                        # exact headers after this broad resource bound.
                        max_columns = len(COMPACT_HEADERS) if COMBINED_SHEET in sheet_names else 26
                        if not match or int(match.group(2)) > max_rows + 1 or column_index_from_string(match.group(1)) > max_columns:
                            raise ValidationError(f"A pricing sheet must have at most {max_rows:,} rows and only the provided columns.")
    except ValidationError:
        raise
    except (BadZipFile, ET.ParseError, KeyError, RuntimeError, ValueError, OSError, NotImplementedError) as error:
        raise ValidationError("The workbook is damaged or is not a supported .xlsx file.") from error


def _decode_excel_text(value):
    # Decode once: _x005F_x000D_ means literal "_x000D_", not CR.
    value = _TEXT_ESCAPE.sub(lambda match: chr(int(match[1], 16)), value)
    if any(0xD800 <= ord(character) <= 0xDFFF for character in value):
        try:
            value = value.encode("utf-16-le", errors="surrogatepass").decode("utf-16-le")
        except UnicodeError as error:
            raise ValidationError("Workbook text contains an invalid Unicode escape.") from error
    if any((ord(character) < 32 and character not in "\n\r\t") or ord(character) in (0xFFFE, 0xFFFF)
           for character in value):
        raise ValidationError("Workbook text contains unsupported control characters.")
    return value


def _pricing_text_cells(payload):
    """Read raw text before openpyxl removes some shared-string escapes.

    This runs only after ZIP/XML preflight. Numbers, formulas and worksheet
    structure remain the responsibility of openpyxl and the existing parser.
    """
    namespace = f"{{{XML_NS}}}"
    relationship_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

    def text_value(container):
        # Rich-text runs are displayed text; phonetic annotations are not.
        pieces = []
        for child in container:
            if child.tag == namespace + "t":
                pieces.append(_decode_excel_text(child.text or ""))
            elif child.tag == namespace + "r":
                pieces.extend(_decode_excel_text(text.text or "") for text in child if text.tag == namespace + "t")
        return "".join(pieces)

    with ZipFile(BytesIO(payload)) as archive:
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        parts = {relation.attrib["Id"]: (
            relation.attrib.get("Type", ""),
            posixpath.normpath(posixpath.join("xl", relation.attrib["Target"])).lstrip("/"))
            for relation in relationships}
        shared = []
        for kind, path in parts.values():
            if kind.endswith("/sharedStrings"):
                shared = [text_value(item) for item in ET.fromstring(archive.read(path))
                          if item.tag == namespace + "si"]
        result = {}
        document = ET.fromstring(archive.read("xl/workbook.xml"))
        for sheet in document.findall(namespace + "sheets/" + namespace + "sheet"):
            kind, path = parts[sheet.attrib[relationship_id]]
            if not kind.endswith("/worksheet"):
                continue
            values, seen = {}, set()
            for cell in ET.fromstring(archive.read(path)).iter(namespace + "c"):
                address = cell.attrib["r"]
                if address in seen:
                    raise ValidationError("A pricing sheet must not contain duplicate cell addresses.")
                seen.add(address)
                if cell.attrib.get("t") == "s":
                    index = cell.find(namespace + "v")
                    if index is not None:
                        if not re.fullmatch(r"[0-9]+", index.text or "") or int(index.text) >= len(shared):
                            raise ValidationError("Workbook contains an invalid shared text reference.")
                        values[address] = shared[int(index.text)]
                elif cell.attrib.get("t") == "inlineStr":
                    inline = cell.find(namespace + "is")
                    if inline is not None:
                        values[address] = text_value(inline)
                elif cell.attrib.get("t") == "str":
                    text = cell.find(namespace + "v")
                    if text is not None:
                        values[address] = _decode_excel_text(text.text or "")
            result[sheet.attrib["name"]] = values
        return result


def _text_rows(sheet, text_cells):
    for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        yield tuple(text_cells.get(f"{get_column_letter(column)}{number}", value)
                    for column, value in enumerate(row, 1))


def _rows(sheet, headers, *, max_rows=MAX_ROWS, text_cells=None):
    # The optional OOXML dimension is absent in some valid editors' output,
    # and may be stale after a list edit. Derive it from actual bounded cells
    # so a missing dimension cannot fail or a small one silently truncate data.
    sheet.reset_dimensions()
    try:
        sheet.calculate_dimension(force=True)
    except (UnboundLocalError, IndexError) as error:
        raise ValidationError(f"{sheet.title} must contain the exported template headers.") from error
    if sheet.max_row > max_rows + 1 or sheet.max_column > len(headers):
        raise ValidationError(f"{sheet.title} has too many rows or unexpected columns.")
    rows = _text_rows(sheet, text_cells or {})
    actual = tuple(next(rows, ()))
    if actual != tuple(headers):
        raise ValidationError(f"{sheet.title} headers must match the exported template exactly.")
    for row_number, values in enumerate(rows, 2):
        if all(value in (None, "") for value in values):
            continue
        yield row_number, dict(zip(headers, values))


def _pricing_rows(workbook, text_cells):
    """Adapt all three file layouts to the existing business parser."""
    names = set(workbook.sheetnames)
    if {"Inventory", "Rates"}.issubset(names) and not names - {"Inventory", "Rates", "Instructions"}:
        return (list(_rows(workbook["Inventory"], INVENTORY_HEADERS, text_cells=text_cells.get("Inventory"))),
                list(_rows(workbook["Rates"], RATE_HEADERS, text_cells=text_cells.get("Rates"))))
    if COMBINED_SHEET not in names or names - {COMBINED_SHEET, "Instructions"}:
        raise ValidationError("Use the Inventory & Rates sheet, or the earlier Inventory and Rates sheets, with only an optional Instructions sheet.")
    sheet = workbook[COMBINED_SHEET]
    sheet.reset_dimensions()
    sheet_text = text_cells.get(COMBINED_SHEET, {})
    headers = tuple(next(_text_rows(sheet, sheet_text), ()))
    if headers in (COMPACT_HEADERS, LEGACY_COMPACT_HEADERS):
        return _compact_rows(sheet, sheet_text, headers)
    inventory, uses, orders = [], [], set()
    for number, row in _rows(workbook[COMBINED_SHEET], COMBINED_HEADERS, max_rows=2 * MAX_ROWS, text_cells=sheet_text):
        kind = row["Row type"]
        if kind not in ("Inventory", "Use"):
            raise ValidationError(f"{COMBINED_SHEET} row {number}: Row type must be Inventory or Use.")
        mapping = _INVENTORY_COLUMNS if kind == "Inventory" else _USE_COLUMNS
        allowed = {"Row type", *mapping.values()} | ({"Use order"} if kind == "Use" else set())
        unexpected = [header for header, value in row.items() if header not in allowed and value not in (None, "")]
        if unexpected:
            raise ValidationError(f"{COMBINED_SHEET} row {number}: {', '.join(unexpected)} must be blank on an {kind} row.")
        record = {header: row[column] for header, column in mapping.items()}
        if kind == "Inventory":
            inventory.append((number, record))
        else:
            order = row["Use order"]
            if order not in (None, ""):
                order = _number(order, f"{COMBINED_SHEET} row {number} Use order", minimum=1)
                if order != int(order) or order > MAX_ROWS:
                    raise ValidationError(f"{COMBINED_SHEET} row {number}: Use order must be a whole number from 1 to {MAX_ROWS}.")
                key = (row["Group"], order)
                if key in orders:
                    raise ValidationError(f"{COMBINED_SHEET} row {number}: Duplicate Use order in {row['Group']}.")
                orders.add(key)
            else:
                order = None
            uses.append((order, number, record))
        if len(inventory) > MAX_ROWS or len(uses) > MAX_ROWS:
            raise ValidationError(f"The library must have at most {MAX_ROWS:,} Inventory rows and {MAX_ROWS:,} Use rows.")
    # Each group keeps its own sequence. Blank orders append in file order;
    # this permits new uses without renumbering existing choices.
    uses.sort(key=lambda item: (item[0] is None, item[0] or 0, item[1]))
    return inventory, [(number, row) for _, number, row in uses]


def _compact_rows(sheet, text_cells, headers=COMPACT_HEADERS):
    inventory, uses, orders = [], [], set()
    # Yield unit is an output column, including in older editable-unit files.
    # Ignore it rather than requiring edits to locked cells when uses change.
    vector_headers = LEGACY_USE_VECTOR_HEADERS
    for number, row in _rows(sheet, headers, max_rows=2 * MAX_ROWS, text_cells=text_cells):
        label = f"{COMBINED_SHEET} row {number}"
        has_product = any(row[header] not in (None, "") for header in INVENTORY_HEADERS)
        has_uses = any(row[header] not in (None, "") for header in vector_headers)
        record = {header: row[header] for header in INVENTORY_HEADERS}
        if has_product:
            # A compact row declares this product and its own uses together;
            # generate one shared identity, never infer links to another row.
            if has_uses and record["Inventory ID"] in (None, ""):
                record["Inventory ID"] = f"inv_{uuid4().hex}"
            inventory.append((number, record))
        if len(inventory) > MAX_ROWS:
            raise ValidationError(f"The library must have at most {MAX_ROWS:,} products and {MAX_ROWS:,} uses after expansion.")
        if not has_uses:
            continue
        vectors = {header: _unpack_uses(row[header], f"{label} {header}") for header in vector_headers}
        counts = {header: len(values) for header, values in vectors.items()}
        if len(set(counts.values())) != 1:
            details = ", ".join(f"{header}={count}" for header, count in counts.items())
            raise ValidationError(f"{label}: all use lists must have the same number of entries, including empty slots ({details}).")
        count = counts["Group"]
        if len(uses) + count > MAX_ROWS:
            raise ValidationError(f"The library must have at most {MAX_ROWS:,} products and {MAX_ROWS:,} uses after expansion.")
        for index in range(count):
            use_label = f"{label} use {index + 1}"
            use = {header: vectors[column][index] for header, column in _COMPACT_USE_COLUMNS.items()}
            for header in ("Group", "Price source", "Yield type", "Rate ID"):
                if isinstance(use[header], str):
                    use[header] = use[header].strip()
            use["Inventory ID"] = record["Inventory ID"] if has_product else None
            for header in ("Unit sell rate", "Yield"):
                use[header] = _use_number(use[header], f"{use_label} {header}")
            order = _use_number(vectors["Use order"][index], f"{use_label} Use order")
            if order is not None:
                order = _number(order, f"{use_label} Use order", minimum=1)
                if order != int(order) or order > MAX_ROWS:
                    raise ValidationError(f"{use_label}: Use order must be a whole number from 1 to {MAX_ROWS}.")
                key = (use["Group"], order)
                if key in orders:
                    raise ValidationError(f"{use_label}: Duplicate Use order in {use['Group']}.")
                orders.add(key)
            uses.append((order, number, index, use))
    uses.sort(key=lambda item: (item[0] is None, item[0] or 0, item[1], item[2]))
    return inventory, [(number, row) for _, number, _, row in uses]


def _inventory_comparison(item):
    return {key: item.get(key) for key in (
        "id", "item_code", "name", "sales_description", "pricing_mode", "supplier_price",
        "markup", "sales_price", "properties",
    )}


def _rate_comparison(group, item, configuration):
    return {"group": group, **{key: item.get(key) for key in ("name", "inventory_id", "price", "yield")},
            "price_mode": _rate_mode(item, configuration), "yield_unit": yield_unit(group, item)}


def _changes(before, after):
    return {"added": len(after.keys() - before.keys()), "removed": len(before.keys() - after.keys()),
            "updated": sum(before[key] != after[key] for key in before.keys() & after.keys())}


def import_pricing_workbook(payload, filename, current_configuration):
    """Validate a workbook and return a replacement configuration without saving."""
    _preflight(payload, filename)
    current = effective_catalog(current_configuration)
    previous_inventory = {item["id"]: item for item in current["inventory"]}
    previous_rates = {rate["id"]: (group, rate) for group, rows in current["rate_groups"].items() for rate in rows}
    try:
        text_cells = _pricing_text_cells(payload)
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False, keep_links=False)
    except ValidationError:
        raise
    except Exception as error:
        raise ValidationError("The workbook could not be read. Export a fresh .xlsx template and try again.") from error
    try:
        inventory_rows, rate_rows = _pricing_rows(workbook, text_cells)
        proposed = deepcopy(current)
        proposed["sources"]["pricing_import"] = {
            "filename": filename.replace("\\", "/").rsplit("/", 1)[-1],
            "sha256": hashlib.sha256(payload).hexdigest(),
        }
        proposed["inventory"] = []
        proposed["rate_groups"] = {group: [] for group in current["rate_groups"]}
        inventory_ids = set()
        for row_number, row in inventory_rows:
            label = f"Inventory row {row_number}"
            identity = _identifier(row["Inventory ID"], "Inventory ID", "inv", inventory_ids)
            old = previous_inventory.get(identity)
            mode = {"Supplier markup": "supplier_markup", "Manual": "manual"}.get(row["Pricing mode"])
            if mode is None:
                raise ValidationError(f"{label}: Pricing mode must be Supplier markup or Manual.")
            item = deepcopy(old) if old else {"source": {}, "status": "Active", "inventory_type": "Untracked"}
            item_code = row["Item code"]
            if item_code is None:
                item_code = ""
            if item_code is not None and (isinstance(item_code, bool) or not isinstance(item_code, (str, int, float))):
                raise ValidationError(f"{label}: Item code must be text or a number.")
            if isinstance(item_code, str) and len(item_code) > 120:
                raise ValidationError(f"{label}: Item code is too long.")
            if isinstance(item_code, (int, float)):
                _number(item_code, f"{label} Item code")
            description = row["Sales description"]
            if description is None:
                description = ""
            if not isinstance(description, str):
                raise ValidationError(f"{label}: Sales description must contain text.")
            item.update({"id": identity, "item_code": item_code,
                         "name": _text(row["Product name"], f"{label} Product name"),
                         "sales_description": description,
                         "pricing_mode": mode,
                         "supplier_price": _number(row["Supplier price"], f"{label} Supplier price", optional=mode == "manual"),
                         "markup": _number(row["Markup"], f"{label} Markup", minimum=-1),
                         "properties": {key: _property(row[header], f"{label} {header}")
                                        for header, key in PROPERTY_HEADERS.items()}})
            entered_sell = _number(row["Sell price"], f"{label} Sell price", optional=mode == "supplier_markup")
            if old:
                for field in ("supplier_price", "markup"):
                    item[field] = _excel_precision(item[field], old.get(field))
                entered_sell = _excel_precision(entered_sell, old["sales_price"])
                for key, value in item["properties"].items():
                    item["properties"][key] = _excel_precision(value, old.get("properties", {}).get(key))
            if mode == "supplier_markup":
                unchanged = old and all(item[key] == old.get(key) for key in ("pricing_mode", "supplier_price", "markup"))
                if unchanged:
                    if entered_sell is not None and entered_sell != old["sales_price"]:
                        raise ValidationError(f"{label}: Sell price is calculated. Edit Supplier price or Markup, or choose Manual pricing.")
                    item["sales_price"] = old["sales_price"]
                else:
                    item["sales_price"] = _number(item["supplier_price"] * (1 + item["markup"]), f"{label} calculated Sell price")
                item["calculated_sell_price"] = item["sales_price"] if not unchanged else old.get("calculated_sell_price", item["sales_price"])
            else:
                item["sales_price"] = entered_sell
            proposed["inventory"].append(item)

        inventory = {item["id"]: item for item in proposed["inventory"]}
        rate_ids = set()
        group_names = {group: set() for group in proposed["rate_groups"]}
        for row_number, row in rate_rows:
            label = f"Rates row {row_number}"
            identity = _identifier(row["Rate ID"], "Rate ID", "rate", rate_ids)
            group = row["Group"]
            if group not in proposed["rate_groups"]:
                raise ValidationError(f"{label}: Unknown Group; use an exact group key from Instructions.")
            name = _text(row["Rate name"], f"{label} Rate name")
            if name.casefold() in group_names[group]:
                raise ValidationError(f"{label}: Duplicate rate name {name!r} in {group}.")
            group_names[group].add(name.casefold())
            linked_id = row["Inventory ID"]
            if isinstance(linked_id, int) and not isinstance(linked_id, bool):
                linked_id = str(linked_id)
            if linked_id in (None, ""):
                linked_id = None
            if linked_id is not None and linked_id not in inventory:
                raise ValidationError(f"{label}: Inventory ID {linked_id!r} is missing. Restore the product, change the link, or remove this rate.")
            price_mode = {"Inventory": "inventory", "Override": "override"}.get(row["Price source"])
            if price_mode is None or (price_mode == "inventory" and linked_id is None):
                raise ValidationError(f"{label}: Choose Inventory with a valid Inventory ID, or Override for an independent rate.")
            entered_price = _number(row["Unit sell rate"], f"{label} Unit sell rate", optional=price_mode == "inventory")
            old_pair = previous_rates.get(identity)
            old = old_pair[1] if old_pair else None
            if old:
                entered_price = _excel_precision(entered_price, old["price"])
            if price_mode == "inventory":
                price = inventory[linked_id]["sales_price"]
                if old and _rate_mode(old, current_configuration) == "inventory":
                    if entered_price is not None and entered_price != old["price"]:
                        price_mode, price = "override", entered_price
                    elif linked_id == old.get("inventory_id") and inventory[linked_id]["sales_price"] == previous_inventory[linked_id]["sales_price"]:
                        price = old["price"]
                elif not old and entered_price is not None and entered_price != price:
                    price_mode, price = "override", entered_price
            else:
                price = entered_price
            uses_yield = bool(current["rate_group_rules"][group]["yield_column"])
            yield_type = row["Yield type"]
            if not uses_yield and yield_type != "Not used":
                raise ValidationError(f"{label}: {group} uses Yield type Not used.")
            if uses_yield and yield_type not in {"Number", "Blank", "Empty text"}:
                raise ValidationError(f"{label}: Choose Yield type Number, Blank or Empty text.")
            if yield_type == "Number":
                yield_value = _number(row["Yield"], f"{label} Yield")
                if old:
                    yield_value = _excel_precision(yield_value, old.get("yield"))
            else:
                if row["Yield"] not in (None, ""):
                    raise ValidationError(f"{label}: Clear Yield when Yield type is {yield_type}.")
                yield_value = "" if yield_type == "Empty text" else None
            rate = deepcopy(old) if old else {"source": {}}
            rate.update({"id": identity, "name": name, "inventory_id": linked_id,
                         "price": price, "price_mode": price_mode,
                         "yield": yield_value, "uses_yield": uses_yield})
            if not uses_yield:
                # A use can move out of a yield category. Older descriptive
                # metadata is then inapplicable, while numeric rules stay put.
                rate.pop("yield_unit", None)
            if old and old["name"] != name:
                rate.pop("display_name", None)
            proposed["rate_groups"][group].append(rate)

        configuration = validate_configuration({"catalog": proposed, "inventory": {}, "rates": {}})
        summary = {
            "inventory": _changes({key: _inventory_comparison(item) for key, item in previous_inventory.items()},
                                  {key: _inventory_comparison(item) for key, item in inventory.items()}),
            "rates": _changes({key: _rate_comparison(group, item, current_configuration)
                               for key, (group, item) in previous_rates.items()},
                              {item["id"]: _rate_comparison(group, item, {})
                               for group, rows in proposed["rate_groups"].items() for item in rows}),
        }
        return {"configuration": configuration, "summary": summary}
    except ValidationError:
        raise
    except (ValueError, TypeError, KeyError, IndexError, OverflowError, ET.ParseError) as error:
        # Read-only workbooks decode cells lazily during iteration. Malformed
        # numeric cells and XML must fail as a rejected import, not a server 500.
        raise ValidationError("The workbook contains malformed cells. Export a fresh template and copy valid values into it.") from error
    finally:
        workbook.close()
