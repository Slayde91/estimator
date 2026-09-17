"""Values-only schedule exchange for the packaged estimating calculators.

Uploads replace schedule inputs as a draft; they cannot replace settings,
reference databases or calculation formulas. Original example rows are cleared.
"""

from copy import deepcopy
import hashlib
from io import BytesIO
import math
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formula.translate import Translator
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from .catalog import ValidationError
from .pricing_workbook import _preflight, _serialize_exact
from .workbook_catalog import column_name, load_workbook_catalog, range_addresses
from .workbook_runtime import application_editable_cells, load_application_catalog


def _schedule(calculator_id):
    catalog = load_application_catalog(calculator_id)
    schedule = catalog["schedule"]
    columns = [column for column in schedule["columns"] if column.get("editable")]
    location = "AA" if calculator_id == "steel_vermiculite" else None
    if location:
        columns.sort(key=lambda field: field["column"] != location)
    return catalog, schedule, columns


def _template_formula(formula, schedule, columns):
    """Move source row references and remap input columns into the template."""
    translated = Translator("=" + formula, origin=f"A{schedule['first_row']}").translate_formula("A2")
    mapping = {field["column"]: column_name(index) for index, field in enumerate(columns, 1)
               if not field.get("generated")}
    tokens = Tokenizer(translated).items
    for token in tokens:
        if token.type == "OPERAND" and token.subtype == "RANGE" and "!" not in token.value:
            token.value = re.sub(r"(?<![A-Z0-9_])(\$?)([A-Z]+)(\$?)([1-9]\d*)(?![A-Z0-9_])",
                                 lambda match: match[1] + mapping.get(match[2], match[2]) + match[3] + match[4], token.value)
    return "".join(token.value for token in tokens)


def _list_values(catalog, reference):
    reference = catalog["defined_names"].get(reference, reference).lstrip("=")
    match = re.fullmatch(r"(?:'((?:[^']|'')+)'|([^!]+))!([A-Z$0-9:]+)", reference)
    if not match:
        return None
    sheet_name = (match[1] or match[2]).replace("''", "'")
    sheet = next((sheet for sheet in catalog["sheets"] if sheet["name"] == sheet_name), None)
    if sheet is None:
        return None
    values = []
    for address in range_addresses(match[3]):
        cell = sheet["cells"].get(address, {})
        # Reference lists must be literal data, never formula-cache authority.
        if "formula" in cell:
            return None
        value = cell.get("value")
        if value is not None:
            values.append(value)
    return values


def export_schedule_template(calculator_id):
    """Create a blank template containing the exact editable input headings."""
    catalog, schedule, columns = _schedule(calculator_id)
    capacity = schedule["last_row"] - schedule["first_row"] + 1
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = schedule["sheet"]
    instructions = workbook.create_sheet("Instructions")
    instructions.append(["CEASEFIRE SCHEDULE TEMPLATE", catalog["title"]])
    instructions.append(["Import", "Import replaces the complete schedule. Review it before saving the estimate. Settings and other calculator inputs are retained."])
    instructions.append(["Rows", f"Enter up to {capacity} rows below the headings. Schedule line numbers are generated from physical row order on import, including hidden or filtered rows. Blank rows keep their positions."])
    instructions.append(["Values only", "Enter values, not Excel formulas. Do not add worksheets or calculated columns. Keep the headings and worksheet name unchanged."])
    instructions.append(["Numbers", "Keep numbers numeric. Display uses at most two decimal places; stored precision is retained. Enter percentage values as percentages, for example 10%."])
    instructions.append(["Dropdowns", "Lists reproduce the source workbook choices. Source warning dropdowns allow specified intermediate values; the calculator reports unsupported cases."])
    instructions.append(["Reference lists", "The following lists support the template dropdowns. Import ignores this sheet and never changes the application reference data."])
    instructions.column_dimensions["A"].width = 32
    instructions.column_dimensions["B"].width = 110
    instructions.freeze_panes = "A8"
    for row in instructions:
        instructions.row_dimensions[row[0].row].height = 44
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    local_lists = {}

    def register(reference):
        if reference in local_lists:
            return local_lists[reference]
        values = _list_values(catalog, reference)
        if values is None:
            return None
        name = f"ScheduleChoices{len(local_lists) + 1}"
        column = column_name(len(local_lists) + 1)
        instructions[f"{column}9"] = name
        for index, value in enumerate(values or [None], 10):
            instructions[f"{column}{index}"] = value
        workbook.defined_names.add(DefinedName(name, attr_text=f"'Instructions'!${column}$10:${column}${9 + max(1, len(values))}"))
        local_lists[reference] = name
        return name

    for index, field in enumerate(columns, 1):
        column = column_name(index)
        sheet.cell(1, index, field["label"])
        sheet.column_dimensions[column].width = 9 if field.get("generated") else max(18, min(38, len(field["label"]) + 2))
        validation = field.get("validation", {})
        source_formula = validation.get("formula1", "")
        translated = source_formula
        if validation.get("type") == "list":
            if source_formula.startswith('"') and source_formula.endswith('"'):
                pass
            elif source_formula in catalog["defined_names"]:
                translated = register(source_formula)
            else:
                # Replace quoted source ranges with local named lists. Retain
                # conditional logic for the board FRL/temperature dropdowns.
                unsupported = []
                def replace_reference(match):
                    value = match[1].replace('""', '"')
                    if "!" not in value:
                        return match[0]
                    name = register(value)
                    if name is None:
                        unsupported.append(value)
                    return '"' + (name or value) + '"'
                translated = re.sub(r'"((?:[^"]|"")*)"', replace_reference, source_formula)
                if unsupported:
                    translated = None
                elif translated:
                    translated = _template_formula(translated, schedule, columns)
        if validation and translated is not None:
            data_validation = DataValidation(
                type=validation["type"], formula1=translated,
                formula2=validation.get("formula2"), operator=validation.get("operator"),
                allow_blank=True, showErrorMessage=validation.get("showErrorMessage") == "1",
                errorStyle=validation.get("errorStyle", "stop"), showDropDown=False,
            )
            data_validation.errorTitle = validation.get("errorTitle", "Check input")
            data_validation.error = validation.get("error", "Use the source workbook input choices.")
            sheet.add_data_validation(data_validation)
            data_validation.add(f"{column}2:{column}{capacity + 1}")
        note = f"{field['label']}. " + ("Optional advanced input. " if field.get("advanced") else "")
        note += ("Generated reference number; import follows row position." if field.get("generated") else
                 "Enter a number or leave blank." if field["type"] == "number" else "Enter a value or leave blank.")
        sheet.cell(1, index).comment = Comment(note, "Ceasefire")
        for row in range(2, capacity + 2):
            cell = sheet.cell(row, index)
            if field.get("generated"):
                cell.value = row - 1
            cell.font = Font(name="Calibri", size=11, color="174D8D")
            cell.fill = PatternFill("solid", fgColor="E7E6E6" if field.get("generated") else "F0F5FA" if row % 2 == 0 else "FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.number_format = "0" if field.get("generated") else "0.00%" if "%" in field["label"] else "0.##" if field["type"] == "number" else "General"
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:{column_name(len(columns))}{capacity + 1}"
    sheet.row_dimensions[1].height = 62
    for target in (sheet, instructions):
        target.sheet_view.showGridLines = False
        for cell in target[1]:
            cell.fill = PatternFill("solid", fgColor="8D1726")
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in target:
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    return _serialize_exact(workbook)


def _value(cell, field, row_number):
    value = cell.value
    if value is None:
        return "" if cell.data_type == "inlineStr" else None
    label = f"Schedule row {row_number}, {field['label']}"
    if isinstance(value, bool) or cell.data_type in {"e", "f", "d"}:
        raise ValidationError(f"{label}: enter text or a number, not a formula, date, boolean or Excel error.")
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or abs(value) > 1e12:
            raise ValidationError(f"{label}: numbers must be finite and no larger than 1 trillion in magnitude.")
        return value
    if not isinstance(value, str) or len(value) > 1000 or any(ord(char) < 32 and char not in "\n\r\t" for char in value):
        raise ValidationError(f"{label}: enter text of at most 1000 characters or a number.")
    if field["type"] == "number" and value != "":
        raise ValidationError(f"{label}: enter a numeric Excel value, not text.")
    return value


def _check_order(payload):
    """Reject ambiguous duplicate or unordered rows/cells before openpyxl.

    Excel parsers may silently overwrite or ignore these malformed records.
    The shared preflight has already bounded and safely parsed every ZIP part.
    """
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(BytesIO(payload)) as archive:
        for name in archive.namelist():
            if not name.startswith("xl/worksheets/") or not name.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(name))
            last_row = 0
            for row in root.findall(f"{namespace}sheetData/{namespace}row"):
                index = int(row.attrib.get("r", last_row + 1))
                if index <= last_row:
                    raise ValidationError("Schedule workbook has duplicate or unordered rows.")
                last_row = index
                previous_column = 0
                for cell in row.findall(f"{namespace}c"):
                    match = re.fullmatch(r"([A-Z])([1-9]\d*)", cell.attrib["r"])
                    if match is None:
                        raise ValidationError("Use the selected calculator's schedule template; schedule columns must be within A:Z.")
                    column = ord(match[1]) - 64
                    if int(match[2]) != index or column <= previous_column:
                        raise ValidationError("Schedule workbook has duplicate, unordered or misplaced cells.")
                    previous_column = column


def import_schedule_workbook(calculator_id, payload, filename, current_inputs=None):
    """Parse a bounded untrusted upload into a complete schedule input overlay."""
    catalog, schedule, columns = _schedule(calculator_id)
    try:
        _preflight(payload, filename)
    except ValidationError as error:
        raise ValidationError(str(error).replace("pricing workbook", "schedule workbook").replace("Pricing workbooks", "Schedule workbooks")) from error
    _check_order(payload)
    if current_inputs is not None and (not isinstance(current_inputs, dict) or any(not isinstance(value, dict) for value in current_inputs.values())):
        raise ValidationError("Current calculator inputs must be grouped by worksheet.")
    proposed = deepcopy(current_inputs or {})
    # Accept only application-authorized inputs even in the caller's overlay.
    for name, cells in proposed.items():
        if set(cells) - application_editable_cells(calculator_id, name):
            raise ValidationError("Current calculator inputs contain a calculated or reference cell.")
    target = proposed.setdefault(schedule["sheet"], {})
    capacity = schedule["last_row"] - schedule["first_row"] + 1
    for row in range(schedule["first_row"], schedule["last_row"] + 1):
        for field in columns:
            if not field.get("generated"):
                target[f"{field['column']}{row}"] = None
    workbook = None
    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False, keep_links=False)
        if schedule["sheet"] not in workbook.sheetnames or set(workbook.sheetnames) - {schedule["sheet"], "Instructions"}:
            raise ValidationError(f"Use the {catalog['title']} template: the workbook must contain {schedule['sheet']} and only an optional Instructions worksheet.")
        sheet = workbook[schedule["sheet"]]
        sheet.reset_dimensions()
        sheet.calculate_dimension(force=True)
        if sheet.max_row > capacity + 1 or sheet.max_column > len(columns) + 1:
            raise ValidationError(f"The schedule must contain at most {capacity} rows and only the exported input columns.")
        legacy_columns = [field for field in load_workbook_catalog(calculator_id)["schedule"]["columns"] if field.get("editable")]
        rows = sheet.iter_rows(max_col=sheet.max_column)
        actual = tuple(cell.value for cell in next(rows, ()))
        numbered_columns = [{"label": "Line", "type": "number", "generated": True}] + columns
        if actual == tuple(field["label"] for field in numbered_columns):
            columns = numbered_columns
        elif actual == tuple(field["label"] for field in legacy_columns):
            columns = legacy_columns
        elif actual != tuple(field["label"] for field in columns):
            raise ValidationError("Schedule headers must match the selected calculator's exported template exactly.")
        imported_rows = 0
        last_populated_row = schedule['first_row']
        for row_index, row in enumerate(rows, 2):
            values = [_value(cell, field, row_index) for cell, field in zip(row, columns)]
            if any(value not in (None, "") for field, value in zip(columns, values) if not field.get("generated")):
                imported_rows += 1
                last_populated_row = schedule['first_row'] + row_index - 2
            source_row = schedule["first_row"] + row_index - 2
            for field, value in zip(columns, values):
                if field.get("generated"):
                    if value not in (None, "") and (type(value) not in (int, float) or int(value) != value):
                        raise ValidationError(f"Schedule row {row_index}, Line: use a whole reference number or leave blank.")
                    if value not in (None, "") and not 1 <= value <= capacity:
                        raise ValidationError(f"Schedule row {row_index}, Line: use a reference number from 1 to {capacity}.")
                    continue
                target[f"{field['column']}{source_row}"] = value
    except ValidationError:
        raise
    except (UnboundLocalError, IndexError, StopIteration) as error:
        raise ValidationError("The schedule must contain the exported template headers.") from error
    except Exception as error:
        raise ValidationError("The schedule workbook contains malformed cells or is not a readable .xlsx file.") from error
    finally:
        if workbook is not None:
            workbook.close()
    return {"inputs": proposed, "imported_rows": imported_rows,
            "schedule_rows": list(range(schedule['first_row'], last_populated_row + 1)),
            "source_sha256": hashlib.sha256(payload).hexdigest()}
