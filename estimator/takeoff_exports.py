"""Literal-value takeoff registers retaining IDs, sources and transfer lineage."""

import csv
from copy import deepcopy
from io import StringIO
import json
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .catalog import ValidationError
from .takeoff_area import measured_area
from .takeoff_model import (area_layers, base_length, has_explicit_area_layers, item_digest, length_additions, measured_length, is_area_item, is_standalone_count,
                           validate_measurement_scope, digest, item_result, object_fields)
from .pricing_workbook import _serialize_exact


HEADERS = ('Item ID', 'Item version', 'Mode', 'Mark / run', 'Level', 'Zone', 'Group',
           'Member type', 'Section', 'Shape', 'Width mm', 'Height mm', 'Diameter mm',
           'Quantity', 'Length per item m', 'Total length m', 'FRL', 'Fire period min',
           'Exposure', 'Orientation', 'System', 'Product', 'Source document ID',
           'Source document', 'Source SHA-256', 'Source page', 'Geometry PDF points',
           'Measurement basis', 'Evidence references', 'Confirmation ID',
           'Confirmation digest', 'Confirmed at', 'Transfer references', 'Notes', 'Physical member IDs',
           'Confirmation checks', 'Confirmed by', 'Gross area m2', 'Excluded area m2', 'Net area m2',
           'Treatment', 'Substrate', 'Surface basis', 'Surface citation', 'Area exclusions',
           'Base length per item m', 'Riser/drop additions per item m', 'Riser/drop source dimensions', 'Count ID', 'Purpose')


def register_rows(snapshot, items, *, include_layers=False):
    documents = {d['id']: d for d in snapshot['documents']}
    rows = []
    for item in sorted(items, key=is_standalone_count):
        if item['state'] != 'confirmed' or not item['confirmation'] or item['confirmation']['digest'] != item_digest(item, snapshot):
            raise ValidationError('Only unchanged confirmed takeoff records can be exported.')
        fields, geometry, receipt = item['fields'], item['geometry'], item['confirmation']
        if not geometry or type(item['quantity']) is not int or item['quantity'] <= 0:
            raise ValidationError('Export requires source geometry and explicit physical quantities.')
        doc = documents[geometry['document_id']]
        validate_measurement_scope(item, snapshot)
        area = measured_area(item, snapshot) if is_area_item(item) else None
        length = None if area is not None or is_standalone_count(item) else measured_length(item, snapshot)
        basis = dict(item['measurement']) if item['measurement'] else {'method': 'count', 'quantity': item['quantity']}
        if basis['method'] == 'calibrated':
            basis['calibration'] = next(c for c in snapshot['calibrations'] if c['id'] == basis['calibration_id'])
        refs = [{**ref, 'document_sha256': documents[ref['document_id']]['sha256'],
                 'document_name': documents[ref['document_id']]['name']} for ref in item['evidence']]
        bindings = [{**{key: binding[key] for key in ('id', 'calculator_id', 'sheet', 'row', 'item_version', 'status')},
                     'destination_verification': 'Last transfer receipt; calculator draft not rechecked by register export'}
                    for binding in snapshot['transfers'] if binding['item_id'] == item['id']]
        row = [item['id'], item['version'], item['mode']]
        row.extend(fields.get(k) for k in ('mark', 'level', 'zone', 'group', 'member_type', 'section', 'shape', 'width_mm', 'height_mm', 'diameter_mm'))
        row.extend([item['quantity'], length, length * item['quantity'] if length is not None else None])
        row.extend(fields.get(k) for k in ('frl', 'fire_period_min', 'exposure', 'orientation', 'system', 'product'))
        row.extend([doc['id'], doc['name'], doc['sha256'], geometry['page'],
                    json.dumps(geometry['points'], separators=(',', ':')),
                    json.dumps(basis, ensure_ascii=False, sort_keys=True),
                    json.dumps(refs, ensure_ascii=False, sort_keys=True), receipt['id'], receipt['digest'], receipt['at'],
                    json.dumps(bindings, ensure_ascii=False, sort_keys=True), fields.get('notes'), json.dumps(item['member_ids']),
                    json.dumps(receipt.get('checks'), ensure_ascii=False, sort_keys=True),
                    json.dumps(receipt.get('actor'), ensure_ascii=False, sort_keys=True)])
        row.extend([*(area[key] if area else None for key in ('gross_area_m2', 'excluded_area_m2', 'net_area_m2')),
                    *(fields.get(key) if area else None for key in ('treatment', 'substrate', 'surface_basis', 'surface_citation')),
                    json.dumps(geometry['exclusions'], ensure_ascii=False, sort_keys=True) if area else None])
        additions = [{**addition, 'document_sha256': documents[addition['document_id']]['sha256'],
                      'document_name': documents[addition['document_id']]['name']}
                     for addition in item.get('length_additions', [])]
        row.extend([base_length(item, snapshot) if area is None and not is_standalone_count(item) else None,
                    length_additions(item, snapshot) if area is None and not is_standalone_count(item) else None,
                    json.dumps(additions, ensure_ascii=False, sort_keys=True) if area is None else None,
                    item.get('count_id'), item.get('purpose')])
        if include_layers:
            row.extend([area_layers(item) if area else None,
                        item_result(item, snapshot)['total_area_m2'] if area else None])
        rows.append(row)
    return rows


def export_register(snapshot, items, format):
    if format not in ('csv', 'xlsx'):
        raise ValidationError('Choose CSV or XLSX register export.')
    include_layers = any(has_explicit_area_layers(item) for item in items)
    headers = HEADERS + (('Number of layers', 'Total area m2') if include_layers else ())
    rows = register_rows(snapshot, items, include_layers=include_layers)
    if format == 'csv':
        output = StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow(headers)
        for row in rows:
            # Spreadsheet importers interpret leading formula characters even in
            # quoted CSV fields. A leading apostrophe intentionally marks text.
            writer.writerow(["'"+v if isinstance(v, str) and v.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')) else v for v in row])
        return output.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8', 'CEASEFIRE-Takeoffs.csv'
    workbook = Workbook()
    sheet = workbook.active; sheet.title = 'Confirmed Takeoffs'
    details = []
    for row_number, row in enumerate([headers, *rows], 1):
        for column, value in enumerate(row, 1):
            if row_number > 1 and isinstance(value, str) and len(value.encode('utf-16-le')) > 60000:
                # Excel cells are limited to 32,767 characters. openpyxl would
                # silently truncate geometry/member/evidence JSON otherwise.
                parts = [value[start:start+15000] for start in range(0, len(value), 15000)]
                details.extend([[row[0], headers[column-1], index+1, len(parts), part] for index, part in enumerate(parts)])
                value = f'See Provenance Detail: {row[0]} / {headers[column-1]} ({len(parts)} ordered parts)'
            cell = sheet.cell(row_number, column, value)
            if isinstance(value, str):
                cell.data_type = 's'  # Exact literal content, including =, +, - and @.
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if row_number == 1:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='1F2937')
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 22
    if details:
        detail_sheet = workbook.create_sheet('Provenance Detail')
        detail_sheet.append(['Item ID', 'Field', 'Part', 'Total Parts', 'Exact Text (concatenate in Part order)'])
        for row in details:
            detail_sheet.append(row)
            for cell in detail_sheet[detail_sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        detail_sheet.freeze_panes = 'C2'
        detail_sheet.auto_filter.ref = detail_sheet.dimensions
        for column, width in (('A', 40), ('B', 30), ('C', 10), ('D', 12), ('E', 100)):
            detail_sheet.column_dimensions[column].width = width
    info = workbook.create_sheet('Provenance')
    for row in [('Project ID', snapshot['project_id']), ('Takeoff revision', snapshot['revision']),
                ('Audit head SHA-256', snapshot['audit_head']),
                ('Scope', 'Confirmed takeoff records. Technical suitability requires separate assessment.'),
                ('Coordinates', 'Unrotated source PDF coordinates. Traces use the retained calibration or cited dimension; Steel counts retain independent member markers and an explicitly entered manual length, without requiring a scale.'),
                ('Quantity', 'Explicit physical quantity; never inferred from photographs. Steel count quantity is derived from its retained member markers, with ordered point/member-ID correspondence. Linear total length is quantity times the sum of base length and each explicitly entered riser/drop addition per member. A wall/slab polygon remains one physical surface; net area is gross area less explicit exclusions. Number of layers is an explicit positive whole number, defaulting to one when absent; total area is net area times layers. No additional face or physical member is inferred.'),
                ('Surface basis', 'Polygon areas use the traced source projection. Historical surface basis, treatment and citation properties remain as entered. Calibration scale is squared for area. Printed presets include PDF UserUnit once; manual known-distance calibration is direct. Rendering zoom and rotation do not change quantities.'),
                ('Calculator links', 'Transfer statuses describe retained receipts. This register export does not recheck the current calculator draft; stale/conflicting links do not become current by exporting.'),
                ('Source links', 'Document IDs, page numbers and hashes identify the retained source originals.')]:
        info.append(row)
    info.column_dimensions['A'].width = 28; info.column_dimensions['B'].width = 100
    for row in info:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if isinstance(cell.value, str):
                cell.data_type = 's'
    payload = _serialize_exact(workbook); workbook.close()
    return payload, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'CEASEFIRE-Takeoffs.xlsx'


def _local_authority(snapshot, store):
    """Receipts copied into project JSON cannot grant export approval."""
    with store.connect() as db:
        approvals = {row[0]: json.loads(row[1]) for row in db.execute(
            'SELECT receipt_id,receipt FROM takeoff_approvals WHERE project_id=? AND kind=?',
            (snapshot['project_id'], 'confirmation'))}
        bindings = {(row[0], row[1]) for row in db.execute(
            'SELECT binding_id,digest FROM takeoff_transfer_receipts WHERE project_id=?',
            (snapshot['project_id'],))}
    return approvals, bindings


def _confirmed(item, snapshot, result, approvals):
    receipt = item.get('confirmation')
    return bool(item['state'] == 'confirmed' and not result['issues'] and receipt
                and receipt['digest'] == item_digest(item, snapshot)
                and approvals.get(receipt['id']) == receipt)


def _calculator_drafts(value):
    from .schedule_rows import blank_schedule_defaults, normalize_schedule_rows
    from .takeoff_transfer import DESTINATIONS
    from .workbook_calculators import normalize_calculator_inputs
    if not isinstance(value, dict) or set(value) - set(DESTINATIONS):
        raise ValidationError('Export accepts the current inputs of the three existing calculators only.')
    drafts = {}
    for calculator_id, supplied in value.items():
        object_fields(supplied, {'inputs', 'schedule_rows'}, 'Current calculator draft', {'inputs', 'schedule_rows'})
        if not isinstance(supplied['inputs'], dict) or any(not isinstance(cells, dict) for cells in supplied['inputs'].values()):
            raise ValidationError('Current calculator inputs must be a map of worksheet input maps.')
        if not isinstance(supplied['schedule_rows'], list):
            raise ValidationError('Current calculator schedule rows must be an explicit list.')
        # This is a read-only calculation of an existing draft, including saved
        # legacy overrides. The public calculator input validator still rejects
        # formula outputs, unknown cells, nonfinite numbers and excess fields.
        inputs = normalize_calculator_inputs(calculator_id, blank_schedule_defaults(calculator_id, supplied['inputs']))
        rows = normalize_schedule_rows(calculator_id, inputs, supplied['schedule_rows'])
        drafts[calculator_id] = {'inputs': inputs, 'schedule_rows': rows,
                                'declared_rows': supplied['schedule_rows'],
                                'fingerprint': digest({'inputs': supplied['inputs'], 'schedule_rows': supplied['schedule_rows']})}
    return drafts


def _linked_results(snapshot, selected, confirmations, registered_bindings, drafts):
    """Project original result cells only after verifying each current link."""
    from .takeoff_transfer import row_values
    from .excel_engine import FormulaError
    from .workbook_calculators import calculator_session, source_model
    output = {}; sessions = {}
    number = lambda value: not isinstance(value, bool) and isinstance(value, (int, float)) and math.isfinite(value)
    for item in selected:
        results = []
        for binding in snapshot['transfers']:
            if binding['item_id'] != item['id']:
                continue
            calculator_id = binding['calculator_id']; draft = drafts.get(calculator_id)
            result = {'calculator_id': calculator_id, 'binding_id': binding['id'],
                      'sheet': binding['sheet'], 'row': binding['row'], 'status': 'Unavailable',
                      'reason': 'Current confirmed source and matching calculator draft are required.',
                      'published_thickness_mm': None, 'estimating_thickness_mm': None,
                      'board_stack_mm': None, 'board_layers': None, 'total_thickness_mm': None,
                      'duct_thickness_mm': None, 'duct_thickness_basis': None,
                      'wrap_layers': None, 'wrap_layer_thickness_mm': None}
            results.append(result)
            if (not confirmations[item['id']] or binding['status'] != 'current'
                    or binding['item_version'] != item['version'] or binding['item_digest'] != item_digest(item, snapshot)
                    or (binding['id'], digest({key: value for key, value in binding.items() if key != 'status'})) not in registered_bindings
                    or draft is None):
                continue
            model = source_model(calculator_id); schedule = model['schedule']
            if (binding['source_sha256'] != model['source']['sha256'] or binding['sheet'] != schedule['sheet']
                    or binding['row'] not in draft['declared_rows']
                    or binding['input_hash'] != digest(row_values(draft['inputs'], schedule, binding['row']))):
                result['reason'] = 'The linked row was edited, removed or uses a different calculator source.'
                continue
            try:
                if calculator_id not in sessions:
                    sessions[calculator_id] = calculator_session(calculator_id, draft['inputs'])
                _, engine, lock = sessions[calculator_id]
                with lock:
                    value = lambda column: engine.value(schedule['sheet'], f"{column}{binding['row']}")
                    if calculator_id == 'steel_vermiculite':
                        status, published, estimating = value('W'), value('O'), value('P')
                        if not isinstance(status, str) or not status.startswith('QUANTIFIED') or not number(estimating) or estimating <= 0:
                            result['reason'] = str(status or 'The calculator has no usable coating thickness.')
                            continue
                        result.update(published_thickness_mm=published if number(published) else None,
                                      estimating_thickness_mm=estimating)
                    elif calculator_id == 'steel_board':
                        status, stack, total, layers = value('AR'), value('Z'), value('AB'), value('AA')
                        if status != 'CLADDING ESTIMATE' or not number(total) or total <= 0:
                            result['reason'] = str(status or 'The calculator has no usable board selection.')
                            continue
                        result.update(board_stack_mm=stack, total_thickness_mm=total,
                                      board_layers=layers if number(layers) else None)
                    else:
                        spray, layers = value('L'), value('R')
                        if value('CB') != 1:
                            result['reason'] = str(value('J') or 'The Ductwork calculator has no usable body result.')
                            continue
                        if number(spray) and spray > 0:
                            result.update(duct_thickness_mm=spray, duct_thickness_basis='Spray DFT')
                        else:
                            layer_m = engine.value('PRODUCT SETTINGS', 'B96')
                            if not number(layers) or layers <= 0 or not number(layer_m) or layer_m <= 0:
                                result['reason'] = 'The Ductwork calculator has no usable spray or continuous-wrap thickness.'
                                continue
                            # Unit projection of the native continuous-layer count
                            # and selected layer thickness, not a new fire rule.
                            result.update(duct_thickness_mm=layers * layer_m * 1000,
                                          duct_thickness_basis='Continuous wrap; local penetration layers excluded',
                                          wrap_layers=layers, wrap_layer_thickness_mm=layer_m * 1000)
                    result.update(status='Current', reason='', calculator_source_sha256=model['source']['sha256'],
                                  calculator_draft_sha256=draft['fingerprint'])
            except (ValueError, ArithmeticError, TypeError, KeyError, FormulaError):
                result['reason'] = 'The existing calculator could not return a usable result.'
        output[item['id']] = results
    return output


def linked_register_results(snapshot, request, *, store):
    """Read original calculator outputs through the same authority gates as exports.

    Results are a projection of the supplied current draft, never persisted as
    Takeoff properties or promoted into calculator inputs.
    """
    object_fields(request, {'expected_revision', 'calculator_drafts'}, 'Linked register results',
                  {'expected_revision'})
    if type(request['expected_revision']) is not int or request['expected_revision'] != snapshot['revision']:
        raise ValidationError('The Takeoffs draft changed. Read its current linked results again.')
    drafts = _calculator_drafts(request.get('calculator_drafts', {}))
    selected = [item for item in snapshot['items'] if item['mode'] in ('steel', 'duct')]
    approvals, binding_registry = _local_authority(snapshot, store)
    confirmations = {item['id']: _confirmed(item, snapshot, item_result(item, snapshot), approvals)
                     for item in selected}
    return {'project_id': snapshot['project_id'], 'revision': snapshot['revision'],
            'linked_results': _linked_results(snapshot, selected, confirmations, binding_registry, drafts)}


def linked_result_text(results):
    """Human-readable result values; unknowns never become zero thickness."""
    sections = []
    for result in results:
        label = {'steel_vermiculite': 'Spray', 'steel_board': 'Board', 'ductwork': 'Ductwork'}[result['calculator_id']]
        if result['status'] != 'Current':
            sections.append(f"{label}: Unavailable ({result['reason']})")
        elif result['calculator_id'] == 'steel_vermiculite':
            published = result['published_thickness_mm']
            sections.append(f"Spray: {result['estimating_thickness_mm']:.2f} mm estimating"
                            + (f"; {published:.2f} mm published" if published is not None else '; published thickness unavailable'))
        elif result['calculator_id'] == 'ductwork':
            sections.append(f"Ductwork: {result['duct_thickness_mm']:.2f} mm; {result['duct_thickness_basis']}")
        else:
            layers = result['board_layers']
            sections.append(f"Board: {result['board_stack_mm']} mm stack; "
                            + (f"{layers:g} layers; " if layers is not None else '')
                            + f"{result['total_thickness_mm']:.2f} mm total")
    return ' | '.join(sections) or 'Unavailable (no current linked calculator result)'


def _schedule_xlsx(snapshot, selected, results, confirmations, linked, *, mode, confirmation_filter):
    """Current register category, with locally verified status and literal provenance."""
    documents = {doc['id']: doc for doc in snapshot['documents']}
    include_layers = any(has_explicit_area_layers(item) for item in selected)
    headers = ['Item ID', 'Item version', 'Mode', 'Confirmation', 'Mark / Run ID', 'Member type', 'Section',
               'Shape', 'Width mm', 'Height mm', 'Diameter mm', 'Quantity', 'Length per item m', 'Total length m',
               *(['Gross area m2', 'Excluded area m2'] if mode == 'walls_floors' else []), 'Net area m2', 'Product', 'FRL', 'Fire period min', 'Exposed sides', 'Critical temperature C',
               'Exposure', 'Orientation', 'Level', 'Zone', 'Group', 'Notes', 'Source document', 'Source page',
               'Source SHA-256', 'Linked calculator result', 'Issues', 'Confirmation ID', 'Confirmation digest',
               'Geometry', 'Measurement basis', 'All item properties', 'Supporting evidence', 'Riser/drop additions',
               'Calculator result provenance', 'Appearance', 'Physical member IDs', 'Source page metadata', 'Confirmation receipt', 'Count ID', 'Purpose']
    if include_layers:
        headers.extend(['Number of layers', 'Total area m2'])
    workbook = Workbook(); sheet = workbook.active; sheet.title = 'Current Takeoffs'
    sheet.append(headers); details = []
    for item in selected:
        field, geometry = item['fields'], item['geometry']; result = results[item['id']]
        doc = documents[geometry['document_id']] if geometry else None
        confirmation = item['confirmation'] if confirmations[item['id']] else None
        def bound_reference(reference):
            source = documents[reference['document_id']]
            return {**reference, 'document_sha256': source['sha256'], 'document_name': source['name']}
        basis = deepcopy(item['measurement'])
        if basis and basis['method'] == 'calibrated':
            calibration = next(value for value in snapshot['calibrations'] if value['id'] == basis['calibration_id'])
            basis['calibration'] = bound_reference(calibration)
        evidence = [bound_reference(reference) for reference in item['evidence']]
        additions = [bound_reference(reference) for reference in item.get('length_additions', [])]
        row = [item['id'], item['version'], item['mode'], 'Confirmed' if confirmation else 'Unconfirmed',
               *(field.get(key) for key in ('mark', 'member_type', 'section', 'shape', 'width_mm', 'height_mm', 'diameter_mm')),
               item['quantity'], result.get('length_m'), result.get('total_length_m'),
               *([result.get('gross_area_m2'), result.get('excluded_area_m2')] if mode == 'walls_floors' else []), result.get('net_area_m2'),
               *(field.get(key) for key in ('product', 'frl', 'fire_period_min', 'sides', 'critical_temperature', 'exposure',
                                          'orientation', 'level', 'zone', 'group', 'notes')),
               doc['name'] if doc else None, geometry['page'] if geometry else None, doc['sha256'] if doc else None,
               linked_result_text(linked[item['id']]), '\n'.join(issue['message'] for issue in result['issues']),
               confirmation['id'] if confirmation else None, confirmation['digest'] if confirmation else None]
        row.extend(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':')) if value is not None else None
                   for value in (geometry, basis, field, evidence, additions,
                                 linked[item['id']], item.get('appearance'), item['member_ids'],
                                 doc['pages'][geometry['page']-1] if doc else None, confirmation))
        row.extend([item.get('count_id'), item.get('purpose')])
        if include_layers:
            row.extend([result.get('layers'), result.get('total_area_m2')])
        for index, value in enumerate(row):
            if isinstance(value, str) and len(value.encode('utf-16-le')) > 60000:
                parts = [value[start:start+15000] for start in range(0, len(value), 15000)]
                details.extend([[item['id'], headers[index], i+1, len(parts), part] for i, part in enumerate(parts)])
                row[index] = f'See Provenance Detail: {item["id"]} / {headers[index]} ({len(parts)} ordered parts)'
        sheet.append(row)
    if details:
        detail = workbook.create_sheet('Provenance Detail')
        detail.append(['Item ID', 'Field', 'Part', 'Total parts', 'Exact text in part order'])
        for row in details: detail.append(row)
    info = workbook.create_sheet('Export scope')
    for row in [('Project ID', snapshot['project_id']), ('Takeoff revision', snapshot['revision']),
                ('Audit SHA-256', snapshot['audit_head']), ('Register', mode), ('Confirmation category', confirmation_filter),
                ('Scope', 'All items in the selected register and confirmation category, including hidden rows. Wall and slab records retain their original measurement types.'),
                ('Approval', 'Unconfirmed items are drafts. Unknown quantities remain blank. Export does not confirm an item or authorize technical suitability.'),
                ('Surface layers', 'Physical surface quantity remains one. Total area equals net measured area times the explicitly entered number of layers, defaulting to one when absent. Gross, excluded and net areas are unchanged; extra faces are not inferred.'),
                ('Linked results', 'Only locally registered current confirmations and exact current calculator row inputs permit a linked steel result. Missing or stale results are unavailable.'),
                ('Precision', 'Stored numeric values retain full precision. Cell display rounding does not change quantities.')]:
        info.append(row)
    for table in workbook:
        table.freeze_panes = 'A2'; table.auto_filter.ref = table.dimensions
        for row in table:
            for cell in row:
                if isinstance(cell.value, str): cell.data_type = 's'
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                if cell.row == 1:
                    cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F2937')
                elif isinstance(cell.value, (float, int)):
                    measurement_headers = {'Width mm', 'Height mm', 'Diameter mm', 'Length per item m', 'Total length m', 'Gross area m2', 'Excluded area m2', 'Net area m2', 'Total area m2'}
                    cell.number_format = '0.00' if table.title == 'Current Takeoffs' and headers[cell.column-1] in measurement_headers else '0'
        for column in table.columns: table.column_dimensions[column[0].column_letter].width = 24
    payload = _serialize_exact(workbook); workbook.close()
    label = 'Walls-Floors' if mode == 'walls_floors' else mode.title()
    suffix = '' if confirmation_filter == 'all' else f'-{confirmation_filter.title()}'
    return payload, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', f'CEASEFIRE-{label}{suffix}-Takeoffs.xlsx'


def export_workspace(snapshot, request, format, *, documents, store):
    """Read-only current-register/drawing exports after service evidence gates."""
    if format not in ('schedule-xlsx', 'marked-pdf'):
        raise ValidationError('Choose the current schedule XLSX or marked drawing PDF export.')
    if request.get('mode') == 'penetrations' and format == 'marked-pdf':
        if type(request.get('expected_revision')) is not int or request['expected_revision'] != snapshot['revision']:
            raise ValidationError('The Takeoffs draft changed before export. Retry from its current state.')
        from .takeoff_physical_markers import export_physical_pdf
        return export_physical_pdf(snapshot, request, documents)
    allowed = {'expected_revision', 'mode', 'item_ids', 'calculator_drafts'} | ({'document_id', 'annotation_ids'} if format == 'marked-pdf' else {'confirmation'})
    required = {'expected_revision', 'mode'} | ({'document_id', 'item_ids'} if format == 'marked-pdf' else set())
    object_fields(request, allowed, 'Current Takeoffs export', required)
    if type(request['expected_revision']) is not int or request['expected_revision'] != snapshot['revision']:
        raise ValidationError('The Takeoffs draft changed before export. Retry from its current state.')
    if request['mode'] not in ('steel', 'duct', 'wall', 'slab', 'walls_floors'):
        raise ValidationError('Choose the active Steel, Duct or Walls/Floors register.')
    modes = ('wall', 'slab') if request['mode'] == 'walls_floors' else (request['mode'],)
    active = sorted((item for item in snapshot['items'] if item['mode'] in modes), key=is_standalone_count)
    confirmation_filter = request.get('confirmation', 'all')
    if not isinstance(confirmation_filter, str) or confirmation_filter not in ('confirmed', 'unconfirmed', 'all'):
        raise ValidationError('Choose confirmed, unconfirmed or all items for the XLSX download.')
    ids = request.get('item_ids', [item['id'] for item in active])
    if (not isinstance(ids, list) or len(ids) > 10000 or any(not isinstance(value, str) for value in ids)
            or len(ids) != len(set(ids)) or set(ids) - {item['id'] for item in active}):
        raise ValidationError('Export item IDs must be distinct existing items in the selected Takeoffs mode.')
    if format == 'schedule-xlsx':
        if set(ids) != {item['id'] for item in active}:
            raise ValidationError('Download XLSX includes every item in the current mode, including hidden rows.')
        if not active:
            raise ValidationError('There are no items in this Takeoffs schedule to export.')
        selected = active
    else:
        document = next((doc for doc in snapshot['documents'] if doc['id'] == request['document_id']), None)
        if document is None:
            raise ValidationError('Choose a source document retained in this Takeoffs project.')
        by_id = {item['id']: item for item in active}; selected = [by_id[identifier] for identifier in ids]
        if any(not item['geometry'] or item['geometry']['document_id'] != document['id'] for item in selected):
            raise ValidationError('Every exported markup must belong to the selected original PDF.')
        active_annotations = {value['id']: value for value in snapshot.get('annotations', {}).get('callouts', [])
                              if value['mode'] in modes and value['document_id'] == document['id']}
        annotation_ids = request.get('annotation_ids', list(active_annotations))
        if (not isinstance(annotation_ids, list) or len(annotation_ids) > 1000
                or any(not isinstance(value, str) for value in annotation_ids)
                or len(annotation_ids) != len(set(annotation_ids)) or set(annotation_ids) - active_annotations.keys()):
            raise ValidationError('Export call-out IDs must identify distinct free call-outs in the selected mode and original PDF.')
        annotations = [active_annotations[identifier] for identifier in annotation_ids]
    approvals, binding_registry = _local_authority(snapshot, store)
    results = {item['id']: item_result(item, snapshot) for item in selected}
    confirmations = {item['id']: _confirmed(item, snapshot, results[item['id']], approvals) for item in selected}
    if format == 'schedule-xlsx' and confirmation_filter != 'all':
        selected = [item for item in selected if confirmations[item['id']] == (confirmation_filter == 'confirmed')]
        if not selected:
            raise ValidationError(f'There are no {confirmation_filter} items in this Takeoffs register to export.')
    drafts = _calculator_drafts(request.get('calculator_drafts', {}))
    linked = _linked_results(snapshot, selected, confirmations, binding_registry, drafts)
    if format == 'schedule-xlsx':
        return _schedule_xlsx(snapshot, selected, results, confirmations, linked, mode=request['mode'], confirmation_filter=confirmation_filter)
    from .takeoff_markup_pdf import export_marked_pdf
    return export_marked_pdf(document, selected, results, confirmations, linked, documents,
                             project_id=snapshot['project_id'], revision=snapshot['revision'], mode=request['mode'], snapshot=snapshot,
                             annotations=annotations)
