"""Literal, hierarchy-preserving UNAPPROVED DRAFT diagnostic exports.

This formatter neither verifies source bytes nor creates an approval or lock.
Locators, display names and uncertainty remain recorded assertions. Every
retained entity is exported, including tombstones. Version two uses numbered
Defect -> Barrier -> Service IDs; legacy exports retain their original schema.
An empty parent never invents a service or quantity.
"""

import csv
from datetime import datetime
from hashlib import sha256
from io import BytesIO, StringIO
import json
from uuid import UUID
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .catalog import ValidationError
from .pricing_workbook import _serialize_exact
from .takeoff_physical import entity_references, graph_collections, graph_parents, graph_digest, validate_graph

STATUS = 'UNAPPROVED DRAFT'
SOURCE_STATUS = 'UNVERIFIED ASSERTIONS'
# The marker is also in the header so an empty CSV is still explicitly a draft.
STATUS_COLUMN = 'export_status [UNAPPROVED DRAFT]'
SOURCE_STATUS_COLUMN = 'source_status [UNVERIFIED ASSERTIONS]'
MAX_SOURCE_NAMES = 20000
MAX_SOURCE_NAME_LENGTH = 2000
MAX_SOURCE_NAME_TEXT = 1024 * 1024
COMMON_HEADERS = (STATUS_COLUMN, SOURCE_STATUS_COLUMN, 'record_scope', 'project_id',
    'graph_id', 'graph_revision', 'graph_sha256', 'entity_type', 'entity_id',
    'entity_revision', 'deleted', 'deleted_at_revision', 'parent_type', 'parent_id',
    'barrier_id', 'defect_id', 'opening_id', 'service_id')
FIELD_HEADERS = ('label', 'location', 'frl', 'barrier_type', 'substrate', 'orientation',
    'opening_type', 'opening_size', 'shape', 'service', 'service_type', 'service_size',
    'quantity', 'thickness_mm', 'width_mm', 'height_mm', 'diameter_mm', 'depth_mm',
    'insulation_mm', 'notes')
DETAIL_HEADERS = ('fields_json', 'evidence_json', 'source_names_json',
                  'uncertainty_state', 'uncertainty_note', 'uncertainty_json')
CSV_HEADERS = COMMON_HEADERS + FIELD_HEADERS + DETAIL_HEADERS
V2_COMMON_HEADERS = COMMON_HEADERS[:-4] + ('display_id', 'defect_id', 'barrier_id',
    'service_id', 'defect_uuid', 'barrier_uuid', 'service_uuid')
V3_COMMON_HEADERS = tuple(key for key in V2_COMMON_HEADERS if key not in ('defect_id', 'defect_uuid'))
V2_FIELD_HEADERS = tuple(field for field in FIELD_HEADERS
                        if field not in ('opening_type', 'opening_size', 'shape', 'depth_mm'))
KIND_FIELDS = {
    'barrier': ('label', 'location', 'barrier_type', 'substrate', 'orientation', 'thickness_mm', 'notes'),
    'defect': ('label', 'location', 'frl', 'notes'),
    'opening': ('label', 'opening_type', 'opening_size', 'shape', 'width_mm', 'height_mm',
                'diameter_mm', 'depth_mm', 'notes'),
    'service': ('label', 'service', 'service_type', 'service_size', 'quantity', 'width_mm',
                'height_mm', 'diameter_mm', 'insulation_mm', 'notes'),
}
EVIDENCE_HEADERS = COMMON_HEADERS + ('association_key', 'association_index', 'document_id',
    'document_name', 'document_sha256', 'page', 'image_id', 'image_sha256', 'occurrence_id',
    'region_json', 'supported_fields_json', 'evidence_note', 'association_json')
CHUNK_HEADERS = ('sheet', 'record_id', 'column', 'part', 'total_parts', 'sha256',
                 'exact_text [concatenate in part order]')
MATRIX_HEADERS = ('Defect ID', 'Barrier ID', 'Service ID', 'Location', 'FRL',
                  'Substrate', 'Orientation', 'Category', 'Service type',
                  'Service quantity', 'Service Size (mm)')
_EXCEL_DATE = (2000, 1, 1, 0, 0, 0)


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)


def _names(value):
    if value is None:
        return {}
    if not isinstance(value, dict) or len(value) > MAX_SOURCE_NAMES:
        raise ValidationError('Source display names must be a bounded document-ID lookup.')
    total = 0
    for identifier, name in value.items():
        try:
            valid_id = isinstance(identifier, str) and str(UUID(identifier)) == identifier
        except ValueError:
            valid_id = False
        if not valid_id:
            raise ValidationError('Source display name keys must be canonical document UUIDs.')
        if (not isinstance(name, str) or len(name) > MAX_SOURCE_NAME_LENGTH
                or any(ord(char) < 32 and char not in '\n\r\t' for char in name)
                or any(0xD800 <= ord(char) <= 0xDFFF for char in name)):
            raise ValidationError('Source display names must contain bounded plain text.')
        total += len(name)
        if total > MAX_SOURCE_NAME_TEXT:
            raise ValidationError('Source display names exceed the total text limit.')
    return dict(value)


def _rows(graph, names):
    collections, parents = graph_collections(graph), graph_parents(graph)
    current = graph['version'] in (2, 3)
    common_headers = {1: COMMON_HEADERS, 2: V2_COMMON_HEADERS, 3: V3_COMMON_HEADERS}[graph['version']]
    index = {entity['id']: (kind, entity) for kind, collection in collections.items()
             for entity in graph[collection]}
    fingerprint = graph_digest(graph)
    entities, associations = [], []
    for kind, collection in collections.items():
        for entity in graph[collection]:
            row = {STATUS_COLUMN: STATUS, SOURCE_STATUS_COLUMN: SOURCE_STATUS,
                'record_scope': 'historical' if entity['deleted'] else 'active',
                'project_id': graph['project_id'], 'graph_id': graph['id'],
                'graph_revision': graph['revision'], 'graph_sha256': fingerprint,
                'entity_type': kind, 'entity_id': entity['id'], 'entity_revision': entity['revision'],
                'deleted': entity['deleted'], 'deleted_at_revision': entity['deleted_at_revision'],
                'parent_type': parents[kind][0] if kind in parents else None,
                'parent_id': entity[parents[kind][1]] if kind in parents else None}
            if current:
                row['display_id'] = entity['display_id']
            ancestor_kind, ancestor = kind, entity
            while True:
                row[f'{ancestor_kind}_id'] = ancestor['display_id'] if current else ancestor['id']
                if current:
                    row[f'{ancestor_kind}_uuid'] = ancestor['id']
                if ancestor_kind not in parents:
                    break
                ancestor_kind, ancestor = index[ancestor[parents[ancestor_kind][1]]]
            # Exact own fields remain lossless. Service Plans additionally
            # displays the FRL recorded by its explicit current barrier parent.
            for key, value in entity['fields'].items():
                row[f'{kind}_size' if key == 'size' else key] = value
            if kind == 'service':
                row['quantity'] = entity['quantity']
                if graph['version'] == 3:
                    row['frl'] = index[entity['barrier_id']][1]['fields'].get('frl')
            row['marker_json'] = _json(entity.get('marker'))
            row['annotation_json'] = _json(entity.get('annotation'))
            row.update(fields_json=_json(entity['fields']), evidence_json=_json(entity['evidence']),
                source_names_json=_json({reference['document_id']: names[reference['document_id']]
                    for reference in entity_references(entity) if reference['document_id'] in names}),
                uncertainty_state=entity['uncertainty']['state'], uncertainty_note=entity['uncertainty']['note'],
                uncertainty_json=_json(entity['uncertainty']))
            entities.append(row)
            for ordinal, reference in enumerate(entity['evidence'], 1):
                association = {key: row.get(key) for key in common_headers}
                association.update(association_key=f'{entity["id"]}/{ordinal}', association_index=ordinal,
                    document_id=reference['document_id'], document_name=names.get(reference['document_id']),
                    document_sha256=reference['document_sha256'], page=reference['page'],
                    image_id=reference.get('image_id'), image_sha256=reference.get('image_sha256'),
                    occurrence_id=reference.get('occurrence_id'),
                    region_json=_json(reference['region']) if 'region' in reference else None,
                    supported_fields_json=_json(reference['fields']) if 'fields' in reference else None,
                    evidence_note=reference.get('note'), association_json=_json(reference))
                associations.append(association)
    return entities, associations, fingerprint


def _csv_text(value):
    # Quoting alone does not prevent spreadsheet formula interpretation. Prefix
    # dangerous leading text, including tab/CR-leading strings, with apostrophe.
    if isinstance(value, str) and (value.startswith(('\t', '\r'))
            or value.lstrip().startswith(('=', '+', '-', '@'))):
        return "'" + value
    return value


def _put(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'
    cell.alignment = Alignment(vertical='top', wrap_text=True)


def _sheet(workbook, title, headers, rows, details, *, record_key='entity_id'):
    sheet = workbook.create_sheet(title)
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(1, column)
        _put(cell, label)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='8B2F13')
        sheet.column_dimensions[cell.column_letter].width = 25
    for row_number, row in enumerate(rows, 2):
        for column, label in enumerate(headers, 1):
            value = row.get(label)
            if isinstance(value, str) and len(value.encode('utf-16-le')) > 60000:
                # At most 30,000 UTF-16 units even if every character uses a
                # surrogate pair. No source characters are truncated or added.
                parts = [value[start:start+15000] for start in range(0, len(value), 15000)]
                identity = row[record_key]
                fingerprint = sha256(value.encode()).hexdigest()
                details.extend({'sheet': title, 'record_id': identity, 'column': label,
                    'part': index, 'total_parts': len(parts), 'sha256': fingerprint,
                    CHUNK_HEADERS[-1]: part} for index, part in enumerate(parts, 1))
                value = f'See Provenance Detail: {title} / {identity} / {label} ({len(parts)} ordered parts)'
            _put(sheet.cell(row_number, column), value)
    sheet.freeze_panes = 'J2'
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def _stable_archive(payload):
    """Remove serialization timestamps; identical draft input has identical bytes."""
    result = BytesIO()
    with ZipFile(BytesIO(payload)) as source, ZipFile(result, 'w', ZIP_DEFLATED) as target:
        for filename in sorted(source.namelist()):
            data = source.read(filename)
            if filename == 'docProps/core.xml':
                root = ET.fromstring(data)
                for tag in ('created', 'modified'):
                    node = root.find(f'{{http://purl.org/dc/terms/}}{tag}')
                    if node is not None:
                        node.text = '2000-01-01T00:00:00Z'
                data = ET.tostring(root, encoding='utf-8', xml_declaration=True)
            entry = ZipInfo(filename, _EXCEL_DATE)
            entry.compress_type = ZIP_DEFLATED
            target.writestr(entry, data)
    return result.getvalue()


def matrix_rows(graph):
    """Active hierarchy leaves, retaining parents with no children as empty rows."""
    graph = validate_graph(graph)
    if graph['version'] not in (2, 3):
        raise ValidationError('The matrix PDF requires the current Defect Reports or Service Plans hierarchy.')
    barriers = [entry for entry in graph['barriers'] if not entry['deleted']]
    services = [entry for entry in graph['services'] if not entry['deleted']]
    services_by_barrier, barriers_by_defect = {}, {}
    for entry in services:
        services_by_barrier.setdefault(entry['barrier_id'], []).append(entry)
    for entry in barriers:
        barriers_by_defect.setdefault(entry.get('defect_id'), []).append(entry)
    rows = []

    def append(defect=None, barrier=None, service=None):
        df, bf, sf = (entry['fields'] if entry else {} for entry in (defect, barrier, service))
        size = sf.get('size')
        if not size and sf.get('width_mm') is not None and sf.get('height_mm') is not None:
            size = f"{sf['width_mm']} x {sf['height_mm']}"
        if not size and sf.get('diameter_mm') is not None:
            size = sf['diameter_mm']
        rows.append([defect['display_id'] if defect else '', barrier['display_id'] if barrier else '',
            service['display_id'] if service else '', bf.get('location') or df.get('location', ''),
            (bf if graph['version'] == 3 else df).get('frl', ''), bf.get('substrate', ''),
            bf.get('orientation', ''), sf.get('service', ''), sf.get('service_type', ''),
            service['quantity'] if service else '', size or ''])

    def append_barrier(defect, barrier):
        children = services_by_barrier.get(barrier['id'], [])
        if children:
            for service in children:
                append(defect, barrier, service)
        else:
            append(defect, barrier)

    if graph['version'] == 3:
        for barrier in barriers:
            append_barrier(None, barrier)
    else:
        for defect in graph['defects']:
            if defect['deleted']:
                continue
            children = barriers_by_defect.get(defect['id'], [])
            if children:
                for barrier in children:
                    append_barrier(defect, barrier)
            else:
                append(defect)
    if len(rows) > 20000:
        raise ValidationError('The matrix PDF exceeds its 20,000-row limit. No rows were omitted; use the CSV/XLSX register.')
    return rows


def export_matrix_pdf(graph):
    from html import escape
    import re
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, LongTable, TableStyle, Spacer

    rows = matrix_rows(graph)
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A3), leftMargin=30,
        rightMargin=30, topMargin=30, bottomMargin=36, title='Passive_Fire_Matrix',
        author='CEASEFIRE ESTIMATOR', invariant=True)
    cell_style = ParagraphStyle('MatrixCell', fontName='Helvetica', fontSize=8, leading=10)
    header_style = ParagraphStyle('MatrixHeader', parent=cell_style, fontName='Helvetica-Bold', textColor=colors.white)
    def cell(value, heading=False):
        text = re.sub(r'\s+', ' ', str(value if value is not None else '')).strip()
        style = header_style if heading else cell_style
        if len(text) > 1000 and not heading:
            style = ParagraphStyle('MatrixLongCell', parent=cell_style, fontSize=6, leading=8)
        return Paragraph(escape(text), style)
    table = LongTable([[cell(value, True) for value in MATRIX_HEADERS]] +
                      [[cell(value) for value in row] for row in rows],
                      colWidths=[65,65,65,125,65,125,85,120,155,90,100], repeatRows=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#841824')),
        ('GRID',(0,0),(-1,-1),.35,colors.HexColor('#D4C8CA')),
        ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),
        ('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),7),
        ('BOTTOMPADDING',(0,0),(-1,-1),7),
        ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F8F5F5')])]))
    title = Paragraph('Passive_Fire_Matrix', ParagraphStyle('MatrixTitle',fontName='Helvetica-Bold',fontSize=18,leading=22))
    def footer(canvas, doc):
        canvas.saveState(); canvas.setFont('Helvetica',7)
        canvas.setFillColor(colors.HexColor('#63575B'))
        canvas.drawString(30,18,f"UNAPPROVED DRAFT | {('Service Plans' if graph['version'] == 3 else 'Defect Reports')} | Physical revision {graph['revision']}")
        canvas.drawRightString(doc.pagesize[0]-30,18,f'Page {doc.page}')
        canvas.restoreState()
    document.build([title,Spacer(1,12),table],onFirstPage=footer,onLaterPages=footer)
    return output.getvalue(), 'application/pdf', 'Passive_Fire_Matrix.pdf'


def export_physical_graph(graph, format, source_names=None):
    """Return bytes, MIME and filename for a draft-only diagnostic register.

    CSV includes all retained entities. XLSX places active entities in their
    typed sheets and tombstones in Historical Entities; Evidence includes both.
    Source names are optional unverified display labels, never file paths to
    read. There is deliberately no approved/export-with-lock switch.
    """
    if not isinstance(format, str) or format not in ('csv', 'xlsx', 'pdf'):
        raise ValidationError('Choose CSV, XLSX or the matrix PDF draft physical export.')
    graph = validate_graph(graph)
    if format == 'pdf':
        return export_matrix_pdf(graph)
    current = graph['version'] in (2, 3)
    common_headers = {1: COMMON_HEADERS, 2: V2_COMMON_HEADERS, 3: V3_COMMON_HEADERS}[graph['version']]
    details_headers = DETAIL_HEADERS + (('marker_json',) if graph['version'] == 3 or any('marker' in entry for entry in graph['barriers']) else ())
    if graph['version'] == 2 and any('annotation' in entry for entry in graph['defects']):
        details_headers += ('annotation_json',)
    csv_headers = common_headers + (V2_FIELD_HEADERS if current else FIELD_HEADERS) + details_headers
    evidence_headers = common_headers + EVIDENCE_HEADERS[len(COMMON_HEADERS):]
    names = _names(source_names)
    rows, associations, fingerprint = _rows(graph, names)
    label = 'Service-Plans' if graph['version'] == 3 else 'Physical'
    filename = f'CEASEFIRE-{label}-UNAPPROVED-DRAFT.{format}'
    if format == 'csv':
        output = StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow(csv_headers)
        writer.writerows([_csv_text(row.get(column)) for column in csv_headers] for row in rows)
        return output.getvalue().encode('utf-8-sig'), 'text/csv; charset=utf-8', filename
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = 'CEASEFIRE physical graph - UNAPPROVED DRAFT'
    workbook.properties.subject = 'Diagnostic assertions only; no approval or Physical Model Lock'
    workbook.properties.creator = 'CEASEFIRE'
    workbook.properties.created = datetime(*_EXCEL_DATE)
    workbook.properties.modified = datetime(*_EXCEL_DATE)
    details = []
    try:
        for kind, collection in graph_collections(graph).items():
            own_fields = KIND_FIELDS[kind] + (('frl',) if graph['version'] == 3 else ())
            _sheet(workbook, collection.title(), common_headers + own_fields + details_headers,
                   [row for row in rows if row['entity_type'] == kind and not row['deleted']], details)
        _sheet(workbook, 'Evidence', evidence_headers, associations, details, record_key='association_key')
        _sheet(workbook, 'Historical Entities', csv_headers, [row for row in rows if row['deleted']], details)
        if details:
            _sheet(workbook, 'Provenance Detail', CHUNK_HEADERS, details, [], record_key='record_id')
        info = workbook.create_sheet('Provenance')
        for row_number, row in enumerate((
            ('Export status', STATUS),
            ('Authority', 'No approval, technical suitability, human review acceptance, or Physical Model Lock is asserted.'),
            ('Source status', 'Evidence locators, source names and uncertainty are unverified assertions. This formatter does not read or verify source bytes.'),
            ('Project ID', graph['project_id']), ('Physical graph ID', graph['id']),
            ('Physical graph revision', graph['revision']), ('Physical graph SHA-256', fingerprint),
            ('State', 'draft'),
            ('Hierarchy', ('Barrier -> Service. FRL belongs to the Barrier; service FRL is displayed from its current parent.' if graph['version'] == 3 else 'Defect -> Barrier -> Service. Numbered IDs are project-local display IDs; entity_id, parent_id and *_uuid retain exact UUID links.' if current else
                           'Barrier -> Defect -> Opening -> Service. Typed parent and ancestor IDs are explicit links.')),
            ('Facts', 'Exact entity facts remain in fields_json. Service Plans displays service FRL from its current Barrier; this does not copy FRL into service fields.' if graph['version'] == 3 else 'Field columns contain only the entity\'s own recorded facts. Parent facts are not copied into child fields. Exact typed fields remain in fields_json.'),
            ('Quantity', ('Only services carry explicit positive quantities. Empty barriers have no service row or quantity. Images and repeated views never create counts.' if current else
                          'Only services carry explicit positive quantities. Empty openings have no service row or quantity. Images and repeated views never create counts.')),
            ('History', 'All retained entities are included. XLSX typed sheets contain active entities; Historical Entities contains tombstones. Evidence includes associations of both.'),
            ('Evidence associations', 'association_index identifies the retained list position, not physical quantity. Image UUID, SHA-256 and occurrence UUID remain distinct.'),
            ('Long text', 'Provenance Detail stores ordered exact text chunks identified by sheet, record ID and column. Concatenate in part order and check the recorded SHA-256.'),
            ('File metadata time', 'Fixed serialization timestamp 2000-01-01; not a source, review or export event time.'),
        ), 1):
            for column, value in enumerate(row, 1):
                _put(info.cell(row_number, column), value)
        info.column_dimensions['A'].width = 30
        info.column_dimensions['B'].width = 110
        payload = _stable_archive(_serialize_exact(workbook))
    finally:
        workbook.close()
    return payload, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', filename
