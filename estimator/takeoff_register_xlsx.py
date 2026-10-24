"""Presentation copies of visible register columns; never approval or input data."""
from math import ceil

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .catalog import ValidationError
from .pricing_workbook import _serialize_exact

MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
STEEL = ('Confirmation', 'Member mark', 'Level', 'Member type', 'Steel section', 'Product',
         'Fire period (min)', 'Crit. Temp (°C)', 'Exposure')
DUCT = ('Confirmation', 'Item', 'Level', 'WxH (mm)', 'FRL', 'Orientation', 'Product')
SURFACE = ('Confirmation', 'Surface ID', 'Surface type', 'Level', 'Surface basis', 'Substrate',
           'Treatment', 'Protection system', 'Protection product', 'FRL / fire rating', 'Number of layers')
LENGTH = ('Qty', 'Length each (m)', 'Total (m)', 'Thickness (mm)', 'Evidence / issues')
DUCT_LENGTH = LENGTH[:3] + ('Duct surface (m²)',) + LENGTH[3:]
AREA = ('Gross (m²)', 'Excluded (m²)', 'Net (m²)', 'Total (m²)', 'Evidence / issues')
STANDALONE = ('Confirmation', 'Item', 'Level', 'WxH (mm)', 'FRL', 'Orientation')


def _width(header):
    if header in ('Evidence / issues', 'Source evidence', 'Legacy hierarchy / label'):
        return 58
    if header in ('Protection system', 'Protection product', 'Substrate', 'Surface basis', 'Service type'):
        return 30
    return max(14, min(26, len(header) + 3))


def register_workbook(tables, title='Current Takeoffs', *, subject='Current register copy', number_formats=None):
    """One presentation sheet with the exact current register columns.

    Values stay typed and exact; number formats only control their display.
    Literal text cannot become a formula. Oversized cells fail instead of losing
    text or adding an unrequested metadata sheet.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    workbook.properties.creator = 'CEASEFIRE'
    workbook.properties.subject = subject
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.tabColor = 'B90A15'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.freeze_panes = 'C2'
    sheet.print_title_rows = '1:1'
    numeric_headers = {'Qty', 'Service quantity', 'Total Count/QTY', 'Number of layers',
                       *LENGTH[1:4], 'Duct surface (m²)', *AREA[:4], 'Length (m)'}
    whole = {'Qty', 'Number of layers', 'Total Count/QTY'}
    try:
        for table_index, (headers, rows) in enumerate(tables):
            if table_index:
                sheet.append([None])
            start = sheet.max_row + 1 if table_index else 1
            sheet.append(list(headers))
            for row in rows:
                if len(row) != len(headers):
                    raise ValidationError('The register column count changed during export.')
                if any(isinstance(value, str) and len(value.encode('utf-16-le')) // 2 > 32767 for value in row):
                    raise ValidationError('Register text exceeds an Excel cell limit. Shorten the affected displayed field before downloading.')
                sheet.append(row)
            for cells in sheet.iter_rows(min_row=start, max_row=sheet.max_row, max_col=len(headers)):
                is_header = cells[0].row == start
                height = 32 if is_header else 26
                for cell, header in zip(cells, headers):
                    value = cell.value
                    if isinstance(value, str):
                        if len(value.encode('utf-16-le')) // 2 > 32767:
                            raise ValidationError('Register text exceeds an Excel cell limit. Shorten the affected displayed field before downloading.')
                        cell.data_type = 's'
                    cell.font = Font(name='Arial', size=10, bold=is_header, color='FFFFFF' if is_header else '2B252A')
                    cell.fill = PatternFill('solid', fgColor='B90A15' if is_header else ('F8F3F3' if (cell.row-start) % 2 == 0 else 'FFFFFF'))
                    cell.alignment = Alignment(horizontal='center' if is_header else ('right' if isinstance(value, (int, float)) else 'left'),
                                               vertical='center' if is_header else 'top', wrap_text=True)
                    cell.border = Border(bottom=Side(style='thin', color='E8E2E3'),
                                         right=Side(style='thin', color='FFFFFF') if is_header else Side())
                    if not is_header and header in numeric_headers:
                        cell.number_format = (number_formats or {}).get(header, '#,##0' if header in whole else '#,##0.00')
                    width = _width(header)
                    letter = cell.column_letter
                    sheet.column_dimensions[letter].width = max(sheet.column_dimensions[letter].width or 0, width)
                    if not is_header and isinstance(value, str):
                        lines = sum(max(1, ceil(len(line) / max(10, width - 3))) for line in value.split('\n'))
                        height = max(height, min(409, lines * 14 + 8))
                sheet.row_dimensions[cells[0].row].height = height
            if table_index == 0:
                sheet.auto_filter.ref = f'A{start}:{get_column_letter(len(headers))}{sheet.max_row}'
        sheet.print_options.horizontalCentered = True
        sheet.print_area = sheet.dimensions
        return _serialize_exact(workbook)
    finally:
        workbook.close()


def _size(fields):
    width, height = fields.get('width_mm'), fields.get('height_mm')
    def label(value):
        return '' if value is None else str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
    return '' if width is None and height is None else f'{label(width)} x {label(height)}'


def export_current_register(snapshot, selected, results, confirmations, linked, *, mode, target, confirmation):
    area = mode in ('wall', 'slab', 'walls_floors')
    main = [item for item in selected if not (area and item.get('purpose') == 'length-only')]
    lengths = [item for item in selected if area and item.get('purpose') == 'length-only']
    if mode == 'steel':
        headers = STEEL + (('Exposed sides',) if target == 'steel_board' else ()) + LENGTH
    else:
        headers = SURFACE + AREA if area else DUCT + DUCT_LENGTH
    documents = {doc['id']: doc for doc in snapshot['documents']}
    rows = []
    for item in main:
        fields, result = item['fields'], results[item['id']]
        review = 'Confirmed' if confirmations[item['id']] else 'Unconfirmed'
        if area:
            row = [review, fields.get('mark'), 'Wall' if item['mode'] == 'wall' else 'Floor',
                   *[fields.get(key) for key in ('level', 'surface_basis', 'substrate', 'treatment', 'system', 'product', 'frl')],
                   fields.get('layers', 1), *[result.get(key) for key in ('gross_area_m2', 'excluded_area_m2', 'net_area_m2', 'total_area_m2')]]
            row[4] = {'wall-face': 'Wall face (true elevation)', 'slab-soffit': 'Slab soffit', 'slab-top': 'Slab top'}.get(fields.get('surface_basis'), fields.get('surface_basis'))
        else:
            row = [review, *[fields.get(key) for key in ('mark', 'level')]]
            row += [fields.get(key) for key in ('member_type', 'section', 'product', 'fire_period_min', 'critical_temperature', 'exposure')] if mode == 'steel' else [_size(fields), fields.get('frl'), fields.get('orientation'), fields.get('product')]
            if mode == 'steel' and target == 'steel_board':
                row.append(fields.get('sides'))
            row += [item['quantity'], result.get('length_m'), result.get('total_length_m')]
            record = next((value for value in linked[item['id']] if value['calculator_id'] == target), None)
            if mode == 'duct':
                row.append(record.get('duct_surface_m2') if record and record.get('duct_surface_status') == 'Current' else None)
            key = {'steel_vermiculite': 'estimating_thickness_mm', 'steel_intumescent': 'intumescent_thickness_mm', 'steel_board': 'total_thickness_mm', 'ductwork': 'duct_thickness_mm'}[target]
            row.append(record.get(key) if record and record['status'] == 'Current' else None)
        geometry = item['geometry']
        evidence = f'{documents[geometry["document_id"]]["name"]} · p{geometry["page"]}' if geometry else 'Source markup missing'
        if item['mode'] == 'duct' and item.get('purpose') != 'count-only' and fields.get('shape') != 'rectangular':
            evidence += '\nRetained non-rectangular duct; calculator transfer unavailable.'
        row.append('\n'.join([evidence, *[issue['message'] for issue in result['issues']]]))
        rows.append(row)
    tables = [(headers, rows)]
    if lengths:
        tables.append((STANDALONE + (('Length (m)',) if area else ('Total Count/QTY',)), [
            ['Confirmed' if confirmations[item['id']] else 'Unconfirmed', item['fields'].get('mark'), item['fields'].get('level'),
             _size(item['fields']), item['fields'].get('frl'), item['fields'].get('orientation'),
             item['quantity'] if item['purpose'] == 'count-only' else results[item['id']].get('length_m')]
            for item in lengths]))
    formats = {'Thickness (mm)': '0.000############'} if mode == 'steel' and target == 'steel_intumescent' else None
    payload = register_workbook(tables, number_formats=formats)
    label = 'Walls-Floors' if area else mode.title()
    suffix = '' if confirmation == 'all' else f'-{confirmation.title()}'
    return payload, MIME, f'CEASEFIRE-{label}{suffix}-Takeoffs.xlsx'


def export_physical_register(graph, *, confirmation='all'):
    from .takeoff_physical import effective_confirmation, graph_collections, graph_parents, validate_graph
    from .takeoff_physical_exports import _stable_archive
    if not isinstance(confirmation, str) or confirmation not in ('all', 'confirmed', 'unconfirmed'):
        raise ValidationError('Choose all, confirmed or unconfirmed items for the XLSX register.')
    graph = validate_graph(graph)
    parents = graph_parents(graph)
    index = {entity['id']: (kind, entity) for kind, collection in graph_collections(graph).items() for entity in graph[collection]}
    legacy, plans = graph['version'] == 1, graph['version'] == 3
    kinds = ['barrier', 'defect', 'opening', 'service'] if legacy else ['barrier', 'service'] if plans else ['defect', 'barrier', 'service']
    headers = (*(['Legacy hierarchy / label'] if legacy else []), *[f'{kind.title()} ID' for kind in kinds],
               'Confirmation', 'Location', 'FRL', *(['Opening type', 'Opening size'] if legacy else []),
               'Substrate', 'Orientation', 'Category', 'Service type', 'Service quantity', 'Service Size (mm)', 'Source evidence')
    children = {}
    for kind, entry in index.values():
        parent = entry[parents[kind][1]] if kind in parents else None
        children.setdefault(parent, []).append((kind, entry))
    rows = []
    def visit(kind, entity, ancestors):
        if entity['deleted']:
            return
        relation = {**ancestors, kind: entity}
        state = effective_confirmation(graph, entity['id'], _index=index)
        if confirmation == 'all' or state == confirmation:
            owner = relation.get('barrier' if plans else 'defect', {})
            if legacy and not owner:
                owner = relation.get('barrier', {})
            fields, barrier = entity['fields'], relation.get('barrier', {}).get('fields', {})
            ids = [relation.get(key, {}).get('display_id', relation.get(key, {}).get('id', '')) for key in kinds]
            service = fields if kind == 'service' else {}
            size = service.get('size')
            row = ([f'{kind.title()}: {fields.get("label") or entity["id"]}'] if legacy else []) + ids
            row += [state.title(), owner.get('fields', {}).get('location'), owner.get('fields', {}).get('frl')]
            if legacy:
                opening = relation.get('opening', {}).get('fields', {})
                row += [opening.get('opening_type'), opening.get('size')]
            row += [barrier.get('substrate'), barrier.get('orientation'), service.get('service'), service.get('service_type'),
                    (entity['quantity'] if entity['quantity'] is not None else 'Unknown') if kind == 'service' else None,
                    size, f'{len(entity["evidence"])} source associations']
            rows.append(row)
        for child_kind, child in children.get(entity['id'], []):
            visit(child_kind, child, relation)
    for kind, entity in children.get(None, []):
        visit(kind, entity, {})
    title = 'Service Plans' if plans else 'Defect Reports'
    payload = _stable_archive(register_workbook([(headers, rows)], title, subject='UNAPPROVED DRAFT register copy'))
    return payload, MIME, f'CEASEFIRE-{title.replace(" ", "-")}-{confirmation.title()}-UNAPPROVED-DRAFT.xlsx'
