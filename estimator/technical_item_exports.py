"""Read-only Technical Library item exports from canonical, refetched evidence.

Only a library ID is accepted. Images use the library's registered, hash-checked
asset reader; report links use the shared whitelist. No caller text, paths,
URLs, calculator inputs or library mutations enter these downloads.
"""

from copy import deepcopy
from hashlib import sha256
from html import escape
from io import BytesIO
from .company_settings import company_logo
import json
import re

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font
from PIL import Image as PillowImage
from reportlab.lib.pagesizes import A4
from reportlab.platypus import Image, KeepTogether, SimpleDocTemplate, Spacer

from .calculator_register import _sheet, _table
from .catalog import ValidationError
from .firestopping_library import MAX_DIAGRAM_BYTES, _render_diagram
from .library_report_links import validate_report_links
from .pricing_workbook import _serialize_exact
from .reference_library import identifier
from .report import _Report, _company_header, _escaped, _register_fonts, _LINE, _MUTED
from .technical_fields import is_hidden_technical_label
from .technical_display import compact_technical_text, display_technical_title

_NOTE = 'Technical Library item details.'
_HIDDEN = {'diagram captions (visually checked)', 'workbook report revision', 'reference review',
           'reference qualification', 'workbook / source differences', 'workbook / source difference',
           'source issues', 'assessment basis', 'basis / review note', 'thickness source',
           'source information', 'source references', 'technical basis'}
_HIDDEN_COLUMNS = {'id', 'configuration', 'source configuration'}
MAX_IMAGES = 64
MAX_IMAGE_BYTES = 60 * 1_048_576
MAX_IMAGE_PIXELS = 80_000_000
MAX_CELLS = 20_000
MAX_TEXT_BYTES = 2 * 1_048_576
_SELECTOR_ID = re.compile(r'^(?:\d+|sys_var_\d+_au)(?:\s*\(\+\d+ variants listed below\))?$', re.I)


def _text(value):
    if value is None:
        return ''
    return json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else str(value)


def _hidden_field(record, field):
    label = field['label'].strip().casefold()
    return (is_hidden_technical_label(label) or label in _HIDDEN
            or (label == 'id' and record.get('id', '').lower().startswith(('promat-', 'trafalgar-selector-'))
                and _SELECTOR_ID.fullmatch(_text(field.get('value')).strip())))


def prepare_technical_item(body, *, library):
    if not isinstance(body, dict) or set(body) - {'id', 'download'} or 'id' not in body:
        raise ValidationError('Include only one Technical Library item ID and an optional download destination.')
    key = identifier(body['id'])
    record = library.detail('technical', key)
    fields, tables, image_refs = [], [], {}
    cells = text_bytes = 0

    def account(values):
        nonlocal cells, text_bytes
        cells += len(values)
        text_bytes += sum(len(_text(value).encode('utf-8')) for value in values)
        if cells > MAX_CELLS or text_bytes > MAX_TEXT_BYTES:
            raise ValidationError('This technical item exceeds the bounded item download size.')

    account([record.get(name, '') for name in ('id', 'title', 'subtitle', 'notice', 'diagram_status')])
    for field in record.get('fields', []):
        label = field['label']
        if _hidden_field(record, field):
            continue
        if not _text(field.get('value')).strip() and not any(field.get(name) for name in ('images', 'table', 'tables', 'table_links')):
            continue
        try:
            validate_report_links(field, user_authored=record.get('user_authored') is True)
        except (ValueError, TypeError, KeyError) as exc:
            raise ValidationError('The technical item contains an invalid report reference.') from exc
        captured = {'label': label, 'value': _text(field.get('value')), 'report_links': deepcopy(field.get('report_links', []))}
        fields.append(captured)
        account([label, captured['value']])
        for ordinal, table in enumerate(([field['table']] if field.get('table') else []) + field.get('tables', [])):
            visible = [i for i, name in enumerate(table['columns']) if ' '.join(name.strip().casefold().split()) not in _HIDDEN_COLUMNS]
            if not visible:
                continue
            columns = [table['columns'][i] for i in visible]
            rows = [[_text(row[i]) for i in visible] for row in table['rows']]
            account(columns)
            for row in rows:
                account(row)
            captions = field.get('table_captions', [])
            caption = _text(captions[ordinal]) if ordinal < len(captions) else ''
            account([caption])
            tables.append({'label': label, 'caption': caption, 'columns': columns, 'rows': rows})
        for image in field.get('images', []):
            image_refs.setdefault(identifier(image['id']), image)
    for image in record.get('images', []):
        image_refs.setdefault(identifier(image['id']), image)
    if len(image_refs) > MAX_IMAGES:
        raise ValidationError('This technical item has too many source diagrams for one download.')
    images, image_bytes, pixels = [], 0, 0
    for image_id, ref in image_refs.items():
        payload, mime, filename = library.asset(image_id, pdf=False)
        image_bytes += len(payload)
        if mime not in {'image/png', 'image/jpeg'} or not 0 < len(payload) <= MAX_DIAGRAM_BYTES or image_bytes > MAX_IMAGE_BYTES:
            raise ValidationError('Technical source diagrams must be bounded registered PNG or JPEG images.')
        try:
            with PillowImage.open(BytesIO(payload)) as header:
                pixels += header.width * header.height
        except Exception as exc:
            raise ValidationError('The technical source diagram cannot be read.') from exc
        if pixels > MAX_IMAGE_PIXELS:
            raise ValidationError('This technical item exceeds the source diagram pixel budget.')
        rendered = _render_diagram(payload)
        caption = _text(ref.get('caption') or 'Source diagram')
        account([caption, filename])
        images.append({'id': image_id, 'filename': filename, 'caption': caption,
                       'source_sha256': sha256(payload).hexdigest(), **rendered})
    sources = []
    for source in record.get('sources', []):
        # Public provenance only; never serialize storage paths or private audit metadata.
        row = [_text(source.get(name)) for name in ('label', 'filename', 'document_id', 'page', 'sheet', 'row')]
        account(row)
        sources.append(row)
    related = [[_text(link.get(name)) for name in ('title', 'id', 'relationship', 'notice')] for link in record.get('links', [])]
    for row in related:
        account(row)
    return {'record': {name: _text(record.get(name)) for name in ('id', 'title', 'subtitle', 'notice', 'diagram_status')},
            'fields': fields, 'tables': tables, 'images': images, 'sources': sources, 'related': related}


def _field_rows(snapshot):
    for field in snapshot['fields']:
        if _hidden_field(snapshot['record'], field):
            continue
        links = {link['label']: link['url'] for link in field['report_links']}
        if links:
            for label in field['value'].splitlines():
                if label.strip():
                    yield compact_technical_text(field['label']), compact_technical_text(label.strip()), links.get(label.strip())
        else:
            yield compact_technical_text(field['label']), compact_technical_text(field['value']), None


def _diagram_caption(image):
    return 'Diagram'


def _new_sheet(workbook, name, title, widths):
    _xlsx_cells([title, _NOTE])
    sheet = _sheet(workbook, name, title, widths)
    sheet.cell(2, 1).value = _NOTE
    return sheet


def _xlsx_cells(values):
    for value in values:
        if isinstance(value, str) and len(value.encode('utf-16-le')) // 2 > 32767:
            raise ValidationError('A technical item cell exceeds Excel\'s 32,767-character limit. Download the PDF or consult the original source instead; no source text was shortened.')


def _xlsx_table(sheet, row, columns, records, widths):
    _xlsx_cells(columns)
    for values in records:
        _xlsx_cells(values)
    return _table(sheet, row, columns, records, widths)


def build_technical_item_xlsx(snapshot):
    record = snapshot['record']
    intro = []
    rows = list(_field_rows(snapshot))
    # Preflight the actual emitted cell text before constructing any workbook.
    # Report-number fields may split into several cells; the PDF snapshot keeps
    # its existing capacity and is never restricted by Excel's per-cell bound.
    for values in intro + [[label, value] for label, value, _ in rows] + snapshot['sources'] + snapshot['related']:
        _xlsx_cells(values)
    for table in snapshot['tables']:
        _xlsx_cells([table['label'], table['caption'], *table['columns']])
        for values in table['rows']:
            _xlsx_cells(values)
    for image in snapshot['images']:
        _xlsx_cells([image['caption'], image['filename'], image['source_sha256']])
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = 'CLASSIFIRE'
    workbook.properties.title = 'Technical Library Item Details'
    sheet = _new_sheet(workbook, 'Item Details', 'TECHNICAL LIBRARY ITEM DETAILS', [32, 82])
    _xlsx_table(sheet, 4, ['Field', 'Value'], intro + [[label, value] for label, value, _ in rows], [32, 82])
    for index, (_, _, url) in enumerate(rows, 5 + len(intro)):
        if url:
            cell = sheet.cell(index, 2)
            cell.hyperlink = url
            cell.font = Font(name='Arial', size=10, color='0563C1', underline='single')
    sheet.freeze_panes = 'A5'
    for ordinal, table in enumerate(snapshot['tables'], 1):
        widths = [max(16, 114 / len(table['columns']))] * len(table['columns'])
        target = _new_sheet(workbook, f'Installation Table {ordinal}', compact_technical_text(table['label']), widths)
        _xlsx_table(target, 4,
                    [compact_technical_text(value) for value in table['columns']],
                    [[compact_technical_text(value) for value in row] for row in table['rows']], widths)
    diagrams = _new_sheet(workbook, 'Source Diagrams', 'SOURCE DIAGRAMS', [32, 82])
    row = 4
    if not snapshot['images']:
        _xlsx_table(diagrams, row, ['Diagram'], [['No diagram is recorded.']], [114])
    for image in snapshot['images']:
        row = _xlsx_table(diagrams, row, ['Source diagram'], [[_diagram_caption(image)]], [114])
        drawing = ExcelImage(BytesIO(image['image_data']))
        ratio = min(760 / drawing.width, 320 / drawing.height, 1)
        drawing.width *= ratio
        drawing.height *= ratio
        diagrams.add_image(drawing, f'A{row}')
        # Each reserved row stays below Excel's height limit, keeping images clear of text.
        height_rows = max(1, int(drawing.height / 20) + 1)
        for offset in range(height_rows):
            diagrams.row_dimensions[row + offset].height = 15
        row += height_rows + 2
    for target in workbook:
        for cells in target:
            for cell in cells:
                if cell.row > 2:
                    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    return _serialize_exact(workbook)


def render_technical_item_pdf(snapshot, *, company=None):
    _register_fonts()
    report = _Report({})
    width, height = A4
    margin, content = 32, width - 64
    record = snapshot['record']
    report.story.append(report.p(compact_technical_text(display_technical_title(record) or 'Technical Library item'), 'subheading'))
    rows = []
    for label, value, url in _field_rows(snapshot):
        cell = report.rich(f'<link href="{escape(url, quote=True)}" color="#BD0018"><u>{_escaped(value)}</u></link>') if url else report.p(value, 'cell')
        rows.append([report.p(label, 'cell'), cell])
    if rows:
        report.story.append(report.table(['Field', 'Value'], rows, [content * .3, content * .7]))
    for table in snapshot['tables']:
        report.story.append(report.p(compact_technical_text(table['label']), 'subheading'))
        if len(table['columns']) <= 4:
            report.story.append(report.table([compact_technical_text(value) for value in table['columns']], [[report.p(compact_technical_text(value), 'cell') for value in row] for row in table['rows']], [content / len(table['columns'])] * len(table['columns']), compact=True))
        else:
            # Keep source columns readable rather than shrinking 20 columns into
            # an A4 page. The explicit source row links each horizontal part.
            for start in range(0, len(table['columns']), 3):
                columns = [compact_technical_text(value) for value in table['columns'][start:start + 3]]
                report.story.append(report.p(f'Columns {start + 1}-{start + len(columns)} of {len(table["columns"])}; row numbers align these parts.', 'small'))
                values = [[report.p(str(index), 'cell')] + [report.p(compact_technical_text(value), 'cell') for value in row[start:start + 3]] for index, row in enumerate(table['rows'], 1)]
                report.story.append(report.table(['Row', *columns], values, [55] + [(content - 55) / len(columns)] * len(columns), compact=True))
    if not snapshot['images']:
        report.story.extend([report.p('Diagram', 'subheading'), report.p('No diagram is recorded.')])
    for index, image in enumerate(snapshot['images']):
        scale = min(content / image['width'], 330 / image['height'])
        diagram = ([report.p('Source diagram', 'subheading')] if index == 0 else []) + [
            Image(BytesIO(image['image_data']), width=image['width'] * scale, height=image['height'] * scale), Spacer(1, 10)]
        report.story.append(KeepTogether(diagram))
    logo = company_logo(company)
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=91, bottomMargin=43,
                                 pageCompression=1, title='CLASSIFIRE - Technical Library Item Details', author='CLASSIFIRE')

    def decorate(canvas, doc):
        canvas.saveState()
        _company_header(canvas, logo, width, height, margin, 'TECHNICAL ITEM', label_below_logo=True, company=company)
        canvas.setStrokeColor(_LINE)
        canvas.line(margin, 32, width - margin, 32)
        canvas.setFont('CeasefireVera', 7)
        canvas.setFillColor(_MUTED)
        canvas.drawString(margin, 20, 'CLASSIFIRE ESTIMATOR | Technical Library record snapshot')
        canvas.drawRightString(width - margin, 20, f'Page {doc.page}')
        canvas.restoreState()

    document.build(report.story, onFirstPage=decorate, onLaterPages=decorate)
    return output.getvalue()
