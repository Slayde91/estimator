"""Prepare a private BOSS matrix index; never change the workbook or live data.

Each system remains one source row. Only substrate marks equal to 1 are used.
Existing broad pipe groups are retained; unmatched fixture systems are held.
Report URLs must come from the captured BOSS resources table, never guesses.
"""
import argparse
from collections import Counter
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import sys
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from estimator.reference_library import ReferenceLibrary, MAX_INDEX
from estimator.library_report_links import valid_report_url
from estimator.technical_fields import normalize_technical_item
from estimator.technical_field_reviewed import review_fingerprint

IMPORT_KEY = 'boss_matrix_import'
SOURCE_URL = 'https://bossfire.com/resources/'
MAX_WORKBOOK_BYTES = 32 * 1024 * 1024
MAX_UNPACKED_BYTES = 128 * 1024 * 1024
EXPECTED_HEADERS = {
    'A':'System','B':'Service (Tier 1)','C':'Service (Tier 2)',
    'D':'Type - Blank Seal / Linear Joint / Service Penetration / Cavity Barrier',
    'E':'Trade','F':'Service Diameter Range (mm)','G':'Min Dia (mm)','H':'Max Dia (mm)',
    'I':'Element','AR':'Substrate Thickness (mm)','AS':'Localised build-up required?',
    'AT':'Build-up requirement',
    **{column:'Fire Stopping Information #'+str(n) for n,column in enumerate(('AU','AV','AW','AX','AY','AZ'),1)},
    'BA':'Annular gap (mm)','BB':'Aperture (mm)','BC':'Sealant depth (mm)',
    'BD':'Sealant Fillet (mm)','BE':'Wrap Length (mm)','BF':'Product Code',
    'BG':'Collar Size / Wrap Product Code','BH':'FRL / FRR',
    'BJ':'Specimen Size & Details','BK':'Report Ref.','BR':'Additional information',
}
EMPTY = {'', '-', 'NA', 'N/A', 'None'}
SUBSTRATES = ('K','L','M','N','O','P','Q','R','S','T','U','V','W','X','Y','Z',
              'AA','AB','AC','AD','AE','AF','AG','AH','AI','AJ','AK','AL','AM','AN','AO','AP','AQ')
SOURCE_COLUMNS = tuple(chr(n) for n in range(65,91)) + tuple('A'+chr(n) for n in range(65,91)) + tuple('B'+chr(n) for n in range(65,83))
PIPE_TYPES = {
    'UPVC': 'uPVC pipe', 'HDPE': 'HDPE pipe', 'PEX': 'PEX pipe',
    'PPR': 'Plastic Pipes', 'PP': 'Plastic Pipes', 'PE': 'Plastic Pipes',
    'CPVC': 'Plastic Pipes', 'ABS': 'Plastic Pipes', 'PEX-AL': 'Plastic Pipes',
    'Kelox / KeKilit': 'Plastic Pipes', 'Concentric 60/100 plastic pipe': 'Plastic Pipes',
    'Concentric c80/120 plastic pipe': 'Plastic Pipes',
    'Copper Pipes': 'Copper Pipes', 'Steel Pipes': 'Steel Pipes',
    'Iron Pipes': 'Unlagged Pipes', 'Sprinkler pipe': 'Unlagged Pipes',
    'Insulated Copper Pipes': 'Lagged Copper Pipes',
    'Insulated Steel Pipes': 'Lagged Steel Pipes', 'Insulated Iron Pipes': 'Lagged Pipes',
}
ELECTRICAL_TYPES = {
    'Fibre Optic / NBN': 'Fibre Optic', 'Fire Alarm': 'TPS & Fire Alarm Cables',
    'TPS': 'TPS & Fire Alarm Cables', 'CAT': 'Data Cables',
    'Telecommunications': 'Communications Cables', 'Mixed Comms Bundle': 'Cable Bundles',
    'Appendix D2 Config': 'D2 Comms Cables', 'Appendix D1 Config': 'D1 Power Cables',
    'Multi Core / Orange Circular / Mains': 'Power Cables', 'Aluminium Core': 'Power Cables',
    'COAX': 'Coaxial Cables', 'Cable tray': 'Cable Trays', 'Busbar': 'Busbar Trunking',
}


def compact(value):
    return (json.dumps(value, ensure_ascii=False, separators=(',', ':'), allow_nan=False, default=str)+'\n').encode('utf-8')


def text(value):
    return '' if value is None else str(value).strip()


def meaningful(value):
    return text(value) not in EMPTY


def read_matrix(path):
    import openpyxl
    if not 0 < Path(path).stat().st_size <= MAX_WORKBOOK_BYTES:
        raise ValueError('Matrix exceeds the bounded workbook size.')
    with ZipFile(path) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED_BYTES:
            raise ValueError('Matrix exceeds the bounded expanded workbook size.')
    raw = Path(path).read_bytes()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError('Review the matrix sheet selection before importing.')
        sheet = workbook.worksheets[0]
        if not 9 <= sheet.max_row <= 10_000 or not 70 <= sheet.max_column <= 256:
            raise ValueError('Review the matrix dimensions before importing.')
        headers = {c.column_letter:text(c.value) for c in next(sheet.iter_rows(min_row=8,max_row=8)) if c.value is not None}
        for column, expected in EXPECTED_HEADERS.items():
            if headers.get(column) != expected:
                raise ValueError('Matrix column changed: '+column)
        if any(not headers.get(column) for column in SUBSTRATES):
            raise ValueError('Missing approved-substrate column.')
        rows, seen = [], set()
        for number, cells in enumerate(sheet.iter_rows(min_row=9,max_col=70),9):
            values = {c.column_letter:c.value for c in cells if c.value is not None and c.column_letter in SOURCE_COLUMNS}
            identity = text(values.get('A'))
            if not identity:
                if any(meaningful(v) for k,v in values.items() if k!='BQ'):
                    raise ValueError('Matrix row has data but no System ID: '+str(number))
                continue
            if not re.fullmatch(r'BF-\d+', identity) or identity in seen:
                raise ValueError('Invalid or duplicate System ID: '+identity)
            seen.add(identity)
            if text(values.get('I')) not in {'Wall','Floor','Ceiling'}:
                raise ValueError('Review substrate orientation for '+identity)
            for column in SUBSTRATES:
                mark=values.get(column)
                if not (type(mark) is int and mark==1 or mark in (None,'-')):
                    raise ValueError('Unexpected substrate mark in '+column+str(number))
            if not meaningful(values.get('BH')) or not meaningful(values.get('BK')):
                raise ValueError('Missing FRL or report for '+identity)
            rows.append({'row':number,'cells':values})
        if not rows:
            raise ValueError('Matrix has no system rows.')
        return {'filename':Path(path).name,'sha256':hashlib.sha256(raw).hexdigest(),
                'sheet':sheet.title,'headers':headers,'rows':rows}
    finally:
        workbook.close()


def held_reason(values):
    if text(values.get('B'))=='Water_Valves_and_Tap Sets':
        return 'Valve/tap service type held by user instruction.'
    if text(values.get('C'))=='GPO' and re.search(r'floor.?box',text(values.get('BJ')),re.I):
        return 'Electrical floorbox service type held by user instruction.'
    return None


def map_services(values):
    family, subtype, kind = (text(values.get(c)) for c in ('B','C','D'))
    if kind=='Blank Seal':return ['Blank Seal']
    if kind=='Linear Joint':return ['Linear Joints']
    if family=='Conduit':return ['Conduits']
    if family.startswith('HVAC_Bundles'):return ['Pair Coils']
    if family.startswith('Mixed_Trades'):return ['Mixed Services']
    if subtype in PIPE_TYPES:return [PIPE_TYPES[subtype]]
    if subtype in ELECTRICAL_TYPES:return [ELECTRICAL_TYPES[subtype]]
    if family=='Other' and subtype in {'GPO','Damper','Steel Purlin','Timber'}:
        return [{'GPO':'Wall Sockets','Damper':'Fire Dampers','Steel Purlin':'Structural Steel','Timber':'Structural Timber'}[subtype]]
    raise ValueError('Unmapped service requires review: '+text(values.get('A'))+' / '+family+' / '+subtype)


def report_key(value):
    value=text(value).upper()
    for prefix, length in [('EWFA',8),('FAS',6),('FRT',6),('FSP',4),('FSV',4),('FAR',4),('PF',5)]:
        match=re.search(r'\b'+prefix+r'[\s_-]*([0-9]{'+str(length)+r'})(?![0-9])',value)
        if match:return prefix+match[1]
    return None


def report_catalog(capture):
    if capture.get('source_url') != SOURCE_URL or not isinstance(capture.get('pdf_links'),list):
        raise ValueError('Use a capture of the BOSS resources report table.')
    result={}
    for link in capture['pdf_links']:
        label,url=link.get('label'),link.get('url')
        if not isinstance(label,str) or not valid_report_url(url) or not url.startswith('https://bossfire.com/documents/test-reports/'):
            raise ValueError('Invalid captured BOSS report link.')
        key=report_key(label) or report_key(url.rsplit('/',1)[-1])
        if key:result.setdefault(key,[]).append({'label':label,'url':url})
    return result


def specimen_parts(value):
    """Move whole, purely descriptive service sentences; keep mixed clauses intact."""
    service,installation=[],[]
    for part in re.split(r'(?<=[.!?])\s+(?=[A-Z0-9•])|\n+',text(value)):
        if not part.strip():continue
        is_service=bool(re.search(r'\b(?:pipes?|cables?|conduits?|pair.?coils?|busbar|damper|purlin|timber|radiata|CAT[5-7])\b',part,re.I))
        is_installation=bool(re.search(r'\b(?:sealant|seal|fillet|collar|wrap|aperture|penetrat\w*|install\w*|appl\w*|fixed|fitted|coating|BOSS|tested|approved|sided)\b',part,re.I))
        (service if is_service and not is_installation else installation).append(part.strip())
    return '\n'.join(service),'\n'.join(installation)


def bind_review(item):
    # Use the established source/projection fingerprint contract, as for Promat.
    from unittest.mock import patch
    fields=deepcopy(item['fields']);captured=[]
    def capture(value, review=None, source_fields=None):
        captured.append(deepcopy(value));return value
    with patch('estimator.technical_fields.normalize_reviewed_content',capture):
        normalize_technical_item(item)
    item['technical_field_review']={'source_fields_sha256':review_fingerprint(fields),
        'pre_review_fields_sha256':review_fingerprint(captured[0]),'fields':fields}


def make_entry(source,row,catalog):
    v=row['cells'];identity=text(v['A']);services=map_services(v)
    fields=[]
    def add(label,value,columns):
        if meaningful(value):
            fields.append({'label':label,'value':text(value)})
    add('ID',identity,['A']);add('Manufacturer','BOSS',[])
    products=list(dict.fromkeys(text(v[c]) for c in ('AU','AV','AW','AX','AY','AZ') if meaningful(v.get(c))))
    product_text='\n'.join('• '+p for p in products)
    for column,label in [('BF','Product code'),('BG','Collar / wrap component code')]:
        if meaningful(v.get(column)):product_text+='\n'+label+': '+text(v[column])
    add('System Products',product_text,['AU','AV','AW','AX','AY','AZ','BF','BG'])
    add('Installation Type',v.get('D'),['D'])
    add('Barrier Type',v.get('I'),['I'])
    approved=[c for c in SUBSTRATES if v.get(c)==1]
    thickness='Minimum substrate thickness: '+text(v.get('AR'))+' mm' if meaningful(v.get('AR')) else 'Minimum substrate thickness not stated.'
    bullets=['• '+re.sub(r'\s*BF-XXXX-\d+\s*$','',source['headers'][c]).strip().replace('\n',' ')+' — '+thickness for c in approved]
    add('Barrier Construction','\n'.join(bullets) or 'No approved substrate is marked in matrix columns K–AQ.',approved+['AR'])
    subtype=text(v.get('C'))
    service_name='; '.join(services)
    if meaningful(subtype) and subtype not in services:service_name+=' — '+subtype
    add('Service',service_name,['B','C','D','E'])
    description,installation=specimen_parts(v.get('BJ'))
    sizes=[label+' (mm): '+text(v[c]) for c,label in [('F','Diameter range'),('G','Minimum diameter'),('H','Maximum diameter')] if meaningful(v.get(c))]
    if description:sizes.append(description)
    add('Service Size / Configuration','\n'.join(sizes),['F','G','H','BJ'])
    add('FRL',v.get('BH'),['BH'])
    add('Maximum Opening Size',('Aperture (mm): '+text(v['BB'])) if meaningful(v.get('BB')) else '',['BB'])
    add('Annular Gap',('Annular gap (mm): '+text(v['BA'])) if meaningful(v.get('BA')) else '',['BA'])
    add('Service Wrap',('Wrap length (mm): '+text(v['BE'])) if meaningful(v.get('BE')) else '',['BE'])
    seal=[label+' (mm): '+text(v[c]) for c,label in [('BC','Sealant depth'),('BD','Sealant fillet')] if meaningful(v.get(c))]
    add('Seal Depth / Fillet Size','\n'.join(seal),['BC','BD'])
    requirements=[installation] if installation else []
    if meaningful(v.get('AS')):requirements.append('Localised build-up: '+text(v['AS']))
    if meaningful(v.get('AT')):requirements.append('Build-up requirements: '+text(v['AT']))
    if meaningful(v.get('BR')):requirements.append('Additional requirements: '+text(v['BR']))
    add('Installation Details','\n\n'.join(requirements),['BJ','AS','AT','BR'])
    report=text(v['BK']);add('Report Number',report,['BK'])
    candidates=catalog.get(report_key(report),[])
    issues=[];link=None
    if not approved:
        issues.append('No approved substrate is marked in matrix columns K–AQ. The originally tested construction has not been substituted as an approved alternative.')
    if len({r['url'] for r in candidates})==1:
        link=candidates[0];fields[-1]['report_links']=[{'label':report,'url':link['url']}]
        if re.search(r'\bR\d|Rev|Issue|V\d',report,re.I) and re.sub(r'\W','',report).casefold() not in re.sub(r'\W','',link['label']).casefold():
            issues.append('The matrix cites '+report+'. BOSS currently lists '+link['label']+'. Revision equivalence has not been verified.')
    else:
        issues.append('The matrix cites '+report+', which is not uniquely listed in section 6 of the captured BOSS resources page. No report PDF link has been assigned.')
    add('Source Issues','\n'.join(issues),['BK'])
    item={'id':'boss-'+identity.lower(),'title':identity,'subtitle':report,'summary':'BOSS matrix system '+identity,
          'source_label':'BOSS Fire System Matrix · '+source['sheet']+' row '+str(row['row']),
          'sources':[{'label':'BOSS Fire System Matrix · '+identity,'filename':source['filename'],'sheet':source['sheet'],'row':row['row']}],
          'fields':fields,'filter_values':{'manufacturer':['BOSS'],'document':[report],'services':services,
              'orientation':['Vertical' if v['I']=='Wall' else 'Horizontal']},
          'boss_matrix_source':{'schema_version':1,'workbook_sha256':source['sha256'],'row':row['row'],
              'source_row_sha256':hashlib.sha256(compact(v)).hexdigest(),
              'approved_substrate_columns':approved,'report_catalog_key':report_key(report)}}
    bind_review(item)
    return item


def prepare(matrix,base,report_links,output):
    base,output=Path(base).resolve(),Path(output).resolve()
    if output.exists() or output==base or output.is_relative_to(base) or base.is_relative_to(output):
        raise ValueError('Choose a new output directory separate from the installed library.')
    if output.is_relative_to(ROOT) or Path(matrix).resolve().is_relative_to(ROOT):
        raise ValueError('Supplier workbook and prepared index must remain outside the source repository.')
    source=read_matrix(matrix);capture=json.loads(Path(report_links).read_text(encoding='utf-8'))
    catalog=report_catalog(capture)
    original=(base/'library.json').read_bytes();bundle=json.loads(original)
    if IMPORT_KEY in bundle:raise ValueError('BOSS already imported; review existing identities before an update.')
    existing={item['id'] for item in bundle['libraries']['technical']['items']}
    entries=[];held=[]
    for row in source['rows']:
        reason=held_reason(row['cells'])
        if reason:held.append({'system':row['cells']['A'],'row':row['row'],'reason':reason});continue
        item=make_entry(source,row,catalog)
        if item['id'] in existing:raise ValueError('Existing library identity collision: '+item['id'])
        entries.append(item)
    filters=bundle['libraries']['technical'].setdefault('filters',[])
    for key,label in [('manufacturer','Manufacturer'),('document','Report'),('orientation','Orientation'),('services','Services')]:
        if not any(f['key']==key for f in filters):filters.append({'key':key,'label':label})
    bundle['libraries']['technical']['items'].extend(entries)
    bundle[IMPORT_KEY]={'schema_version':1,'source_filename':source['filename'],'workbook_sha256':source['sha256'],
        'sheet':source['sheet'],'source_rows':len(source['rows']),'imported':len(entries),'held':held,
        'mapping_policy':'Existing broad pipe groups; unmatched valve/tap and floorbox systems held by user.',
        'report_catalog_sha256':hashlib.sha256(Path(report_links).read_bytes()).hexdigest(),'report_catalog_source':SOURCE_URL,
        'report_catalog':capture['pdf_links'],
        'field_columns':{'ID':['A'],'Service':['B','C','D','E'],'Service Size / Configuration':['F','G','H','BJ'],
            'Barrier Type':['I'],'Barrier Construction':list(SUBSTRATES)+['AR'],
            'System Products':['AU','AV','AW','AX','AY','AZ','BF','BG'],'Annular Gap':['BA'],
            'Maximum Opening Size':['BB'],'Seal Depth / Fillet Size':['BC','BD'],'Service Wrap':['BE'],
            'FRL':['BH'],'Installation Details':['BJ','AS','AT','BR'],'Report Number':['BK']}}
    payload=compact(bundle)
    if len(payload)>=MAX_INDEX:raise ValueError('Candidate needs '+str(len(payload))+' bytes; the current bounded index limit is '+str(MAX_INDEX)+'.')
    output.mkdir(parents=True);(output/'library.json').write_bytes(payload)
    library=ReferenceLibrary(output);summary=library.overview()
    for item in entries:
        actual={f['label']:{k:v for k,v in f.items() if k!='source_labels'} for f in library.detail('technical',item['id'])['fields']}
        expected={f['label']:f for f in item['fields']}
        for label,value in expected.items():
            if actual.get(label)!=value:raise ValueError('Projected BOSS field changed: '+item['id']+' / '+label)
    linked=sum(any(f.get('report_links') for f in item['fields']) for item in entries)
    audit={'source':{k:v for k,v in source.items() if k!='rows'},'base_index_sha256':hashlib.sha256(original).hexdigest(),
        'candidate_sha256':hashlib.sha256(payload).hexdigest(),'candidate_bytes':len(payload),
        'source_systems':len(source['rows']),'imported':len(entries),'held':held,'linked_records':linked,
        'unlinked_reports':dict(Counter(next(f['value'] for f in item['fields'] if f['label']=='Report Number') for item in entries if not any(f.get('report_links') for f in item['fields']))),
        'libraries':summary['libraries'],'original_records_and_assets_preserved':True}
    (output/'import-audit.json').write_bytes(compact(audit))
    return audit


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('matrix',type=Path);parser.add_argument('--base',required=True,type=Path)
    parser.add_argument('--report-links',required=True,type=Path);parser.add_argument('--output',required=True,type=Path)
    args=parser.parse_args()
    result=prepare(args.matrix,args.base,args.report_links,args.output)
    print(json.dumps({k:result[k] for k in ['source_systems','imported','held','linked_records','unlinked_reports','candidate_bytes']}))


if __name__=='__main__':main()
