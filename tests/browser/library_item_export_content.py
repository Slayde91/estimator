"""Check documents from the disposable library editor fixture only."""
import argparse
import base64
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import sys
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from openpyxl import load_workbook
from PIL import Image, ImageChops, ImageStat
from pypdf import PdfReader
from estimator.penetration_calculator import calculate, definition
from estimator.firestopping_library import diagram_upload


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--capture', required=True, type=Path)
    parser.add_argument('--file', required=True, type=Path)
    args = parser.parse_args()
    captured = json.loads(args.capture.read_text(encoding='utf-8'))
    expected = calculate({'globals': captured['globals'], 'rows': [captured['item']]}, captured['configuration'])
    row = expected['rows'][0]
    diagram = captured.get('diagram')
    diagram_hash = sha256(base64.b64decode(diagram['content_base64'])).hexdigest() if diagram else None
    if args.file.suffix == '.pdf':
        reader = PdfReader(args.file)
        text = '\n'.join(page.extract_text() for page in reader.pages)
        assert captured['item']['inputs']['T'] in text
        assert 'Unscheduled composer retained' not in text
        images = [image.image for page in reader.pages for image in page.images]
        if diagram:
            # The existing upload path normalizes to JPEG before PDF embedding.
            # Compare those retained display pixels, not pre-normalization bytes.
            rendered = diagram_upload(diagram)
            with Image.open(BytesIO(rendered['image_data'])) as source:
                matching = [image for image in images if image.size == source.size]
                assert len(matching) == 1
                # PDF image extraction can re-encode JPEG pixels; retain a tight
                # mean channel tolerance across the entire source diagram.
                difference = ImageChops.difference(matching[0].convert('RGB'), source.convert('RGB'))
                assert max(ImageStat.Stat(difference).mean) < 2
            public_text = text + str(reader.metadata)
            assert diagram_hash not in public_text and rendered['sha256'] not in public_text
            assert diagram['filename'] not in public_text
        else:
            assert all(image.size != (640, 360) for image in images)
            assert 'No source diagram is attached to this current item.' in text
    else:
        workbook = load_workbook(args.file, data_only=False)
        assert 'Schedule' not in workbook.sheetnames
        inputs = list(workbook['Inputs'].iter_rows(min_row=5, values_only=True))
        metadata = definition(captured['configuration'])
        assert len(inputs) == len(metadata['row_fields'])
        assert [cell.value for cell in workbook['Inputs'][4]] == ['Section', 'Field', 'Raw input', 'Effective value', 'Units']
        assert 'Source' not in workbook.sheetnames
        for field, record in zip(metadata['row_fields'], inputs):
            value = row['inputs'].get(field['column'])
            assert record[:2] == (field['group'], field['label']), field['column']
            assert record[2] == (None if value == '' else value), field['column']
        assert workbook['Item']['B5'].value == captured['item']['id']
        assert workbook['Item']['B8'].value == expected['summary']['grand_total']
        values = '\n'.join(str(cell.value) for sheet in workbook for record in sheet for cell in record)
        assert 'Unscheduled composer retained' not in values
        assert all(cell.data_type != 'f' for sheet in workbook for record in sheet for cell in record)
        assert all(cell.hyperlink is None or sheet.title == 'Reports' for sheet in workbook for record in sheet for cell in record)
        with ZipFile(args.file) as archive:
            xml = '\n'.join(archive.read(name).decode('utf-8') for name in archive.namelist()
                            if name.endswith(('.xml', '.rels')))
        if diagram_hash:
            assert diagram_hash not in xml and diagram['filename'] not in xml
        workbook.close()
    print(json.dumps({'passed': True, 'file': args.file.name, 'diagram_sha256': diagram_hash, 'fields_exact': True}))


if __name__ == '__main__':
    main()
