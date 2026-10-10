"""Verify only the disposable browser's current-item documents."""
import argparse
import base64
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from openpyxl import load_workbook
from pypdf import PdfReader

from estimator.penetration_calculator import calculate, definition


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--directory', required=True)
    args = parser.parse_args()
    folder = Path(args.directory)
    captured = json.loads((folder / 'captured-item.json').read_text(encoding='utf-8'))
    item = captured['item']
    expected = calculate({'globals': captured['globals'], 'rows': [item]}, captured['configuration'])
    assert isinstance(expected['summary']['grand_total'], (int, float)) and expected['summary']['grand_total'] > 0
    reader = PdfReader(folder / 'CEASEFIRE-Firestopping-Item.pdf')
    text = '\n'.join(page.extract_text() for page in reader.pages)
    assert item['inputs']['T'] in text and item['inputs']['U'] in text
    assert item['id'] in text
    assert 'SCHEDULE ONLY - RETAIN THIS ROW' not in text
    diagram_sha = sha256(base64.b64decode(captured['diagram']['content_base64'])).hexdigest()
    assert diagram_sha in text and captured['diagram']['filename'] in text
    images = [image.image for page in reader.pages for image in page.images]
    assert sum(image.size == (320, 160) for image in images) == 1
    source_image = next(image for image in images if image.size == (320, 160)).convert('RGB')
    assert source_image.getpixel((80, 60))[0] > source_image.getpixel((80, 60))[1] * 3
    assert min(source_image.getpixel((250, 130))) > 240
    workbook = load_workbook(folder / 'CEASEFIRE-Firestopping-Item.xlsx', data_only=False)
    assert 'Schedule' not in workbook.sheetnames
    inputs = {row[2]: row for row in workbook['Inputs'].iter_rows(min_row=5, values_only=True)}
    assert len(inputs) == len(definition(captured['configuration'])['row_fields'])
    for field in definition(captured['configuration'])['row_fields']:
        assert inputs[field['column']][3] == expected['rows'][0]['inputs'].get(field['column'])
    assert workbook['Item']['B5'].value == item['id']
    # OOXML empty text reads as None; numeric/error values retain their type/value.
    assert workbook['Item']['B8'].value == (None if expected['summary']['grand_total'] == '' else expected['summary']['grand_total'])
    assert all(cell.data_type != 'f' and cell.hyperlink is None for sheet in workbook for row in sheet for cell in row)
    all_values = '\n'.join(str(cell.value) for sheet in workbook for row in sheet for cell in row)
    assert 'SCHEDULE ONLY - RETAIN THIS ROW' not in all_values
    assert diagram_sha in all_values
    result = {'passed': True, 'pdf_pages': len(reader.pages), 'diagram_sha256': diagram_sha,
              'source_image_embedded_once': True, 'source_pixels_visible': True,
              'all_raw_fields_exact': True, 'calculated_total_exact': expected['summary']['grand_total'],
              'no_formulas_or_hyperlinks': True, 'no_schedule_item_in_either_file': True}
    (folder / 'content-result.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    print(json.dumps(result))


if __name__ == '__main__':
    main()
