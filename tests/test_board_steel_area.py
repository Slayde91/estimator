"""Steel surface grouping, with independent retained girth/length expectations."""

from copy import deepcopy
from io import BytesIO
import unittest

from openpyxl import load_workbook
from pypdf import PdfReader

from estimator.board_steel_area import board_net_steel_areas
from estimator.calculator_report import _material_rows, build_calculator_summary_report, project_calculator_report
from estimator.calculator_register import build_calculator_register
from estimator.excel_engine import WorkbookEngine
from estimator.workbook_calculators import calculate_worksheet, source_model
from tests.test_board_product_totals import example_inputs


def areas(inputs):
    return board_net_steel_areas(WorkbookEngine(source_model('steel_board'), inputs))


def stock(result, product='PROMATECT-XS', thickness=15):
    return next(row for row in result['rows'] if row['product'] == product and row['thickness_mm'] == thickness)


class BoardSteelAreaTests(unittest.TestCase):
    def test_pdf_retains_steel_demand_or_unavailability_when_board_demand_is_zero(self):
        rows = [{'row': number, 'values': dict.fromkeys('EFGHIJ', 0)} for number in (12, 13, 14)]
        summary = {'rows': rows, 'net_steel_areas': {
            12: {'net_steel_sqm': 0, 'issues': []},
            13: {'net_steel_sqm': 0.00001, 'issues': []},
            14: {'net_steel_sqm': None, 'issues': [{'line': 1, 'message': 'Missing geometry'}]},
        }}
        self.assertEqual([row['row'] for row in _material_rows('steel_board', summary)], [13, 14])

    def test_same_thickness_double_layer_counts_steel_once_without_waste_or_extras(self):
        inputs = example_inputs(9, length=10)
        # 200UC46's retained 4-side nominal profile girth is 1.1838283271233896 m.
        # It differs from the 0.812 m boxed perimeter and 1.864 m two-layer board girth.
        expected = 11.838283271233896
        first = areas(inputs)
        self.assertAlmostEqual(stock(first)['net_steel_sqm'], expected, places=12)
        inputs['CALCULATOR']['L9'] = 1
        inputs['EXTRA BOARDS'].update({'B6': 'PROMATECT-XS', 'C6': 15, 'G6': 999})
        self.assertEqual(stock(areas(inputs))['net_steel_sqm'], stock(first)['net_steel_sqm'])

    def test_three_sides_use_exposed_girth_and_total_lineal_metres(self):
        inputs = example_inputs(9, length=0.123456789)
        inputs['CALCULATOR']['G9'] = 3
        result = areas(inputs)
        used = [row for row in result['rows'] if row['net_steel_sqm']]
        self.assertTrue(used)
        for row in used:
            self.assertAlmostEqual(row['net_steel_sqm'], 0.9808283271233895 * 0.123456789, places=14)

    def test_different_layer_thicknesses_each_receive_the_member_once(self):
        inputs = example_inputs(9, length=2)
        inputs['CALCULATOR'].update({'C9': 'TRAFALGAR COREX', 'H9': 180})
        engine = WorkbookEngine(source_model('steel_board'), inputs)
        layer1, layer2 = engine.value('CALCULATOR', 'AJ9'), engine.value('CALCULATOR', 'AK9')
        self.assertGreater(layer1, 0)
        self.assertGreater(layer2, 0)
        self.assertNotEqual(layer1, layer2)
        result = board_net_steel_areas(engine)
        for thickness in (layer1, layer2):
            self.assertAlmostEqual(stock(result, 'TRAFALGAR COREX', thickness)['net_steel_sqm'], 2.367656654246779)

    def test_full_schedule_including_last_row_and_typographic_section_identity(self):
        inputs = example_inputs(9, 1008, length=0.123456789)
        result = areas(inputs)
        self.assertAlmostEqual(stock(result)['net_steel_sqm'], 1.1838283271233896 * 0.123456789 * 2, places=14)
        inputs = example_inputs(1008, length=2)
        inputs['CALCULATOR']['D1008'] = '100x100x10 SHS'
        self.assertAlmostEqual(stock(areas(inputs))['net_steel_sqm'], 0.8)

    def test_missing_or_unsupported_profile_withholds_whole_group_and_explains_lines(self):
        inputs = example_inputs(9, 10, length=1)
        inputs['CALCULATOR']['M10'] = 'Rotate'
        result = areas(inputs)
        group = stock(result)
        self.assertIsNone(group['net_steel_sqm'])
        self.assertEqual(group['issues'][0]['line'], 2)
        self.assertIn('Standard exposure', group['issues'][0]['message'])
        self.assertTrue(any('lines 2' in note for note in result['notes']))
        inputs['CALCULATOR'].update({'M10': 'Standard', 'F10': 0})
        self.assertIsNone(stock(areas(inputs))['net_steel_sqm'])
        self.assertIn('positive total lineal metres', stock(areas(inputs))['issues'][0]['message'])

    def test_custom_geometry_and_missing_design_never_fall_back_to_box_area(self):
        inputs = example_inputs(9)
        inputs['CALCULATOR'].update({'R9': 203, 'S9': 203})
        result = areas(inputs)
        assigned = [row for row in result['rows'] if row['issues']]
        self.assertTrue(assigned)
        self.assertTrue(all(row['net_steel_sqm'] is None for row in assigned))
        self.assertIn('Custom steel geometry', assigned[0]['issues'][0]['message'])
        inputs = example_inputs(9)
        inputs['CALCULATOR']['D9'] = 'Unknown steel ID'
        result = areas(inputs)
        self.assertEqual(result['unassigned'][0]['line'], 1)
        self.assertTrue(all(row['net_steel_sqm'] == 0 for row in result['rows']))
        self.assertIn('Unassigned steel area', result['notes'][0])

    def test_blank_is_zero_and_calculation_does_not_mutate_model_inputs_or_source(self):
        inputs = example_inputs(1008)
        model, before_inputs = source_model('steel_board'), deepcopy(inputs)
        before_model = deepcopy(model)
        engine = WorkbookEngine(model, inputs)
        one = board_net_steel_areas(engine)
        self.assertEqual(board_net_steel_areas(engine), one)
        self.assertEqual(model, before_model)
        self.assertEqual(inputs, before_inputs)
        blank = areas(example_inputs())
        self.assertTrue(all(row['net_steel_sqm'] == 0 for row in blank['rows']))
        self.assertEqual(blank['unassigned'], [])

    def test_worksheet_pdf_and_register_share_value_and_keep_native_purchasing_cells(self):
        inputs = example_inputs(9, length=10)
        api = calculate_worksheet('steel_board', inputs, 'BOARD SUMMARY')
        data = project_calculator_report('steel_board', inputs)
        self.assertEqual(api['board_net_steel_areas'], data['board_net_steel_areas'])
        stock_row = stock(api['board_net_steel_areas'])['row']
        purchase = next(row for row in data['summaries'][0]['rows'] if row['row'] == stock_row)['values']['J']
        workbook = load_workbook(BytesIO(build_calculator_register('steel_board', inputs)))
        summary = workbook['Summary']
        header = next(row for row in summary if any(cell.value == 'Net Steel sqm' for cell in row))
        labels = {cell.value: cell.column for cell in header if cell.value}
        self.assertEqual(labels['Net Steel sqm'] + 1, labels['Purchase sqm'])
        target = header[0].row + (stock_row - 12) + 1
        self.assertAlmostEqual(summary.cell(target, labels['Net Steel sqm']).value, 11.838283271233896)
        self.assertEqual(summary.cell(target, labels['Purchase sqm']).value, purchase)
        text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(build_calculator_summary_report('steel_board', inputs))).pages)
        self.assertIn('Net Steel sqm', text)
        self.assertIn('11.84', text)
        self.assertIn('Different thickness rows can include the same member', text)


if __name__ == '__main__':
    unittest.main()
