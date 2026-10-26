"""Synthetic matrix rows only; supplier workbook and records stay outside Git."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest

from scripts.import_boss_matrix import (make_entry, map_services, held_reason, read_matrix,
    report_catalog, report_key, prepare, SUBSTRATES, SOURCE_URL, EXPECTED_HEADERS)
from estimator.reference_library import ReferenceLibrary
from estimator.technical_fields import normalize_technical_item
from estimator.library_facets import classify_facets, SERVICES, CATEGORIES
from estimator.library_substrates import classify_substrates
from estimator.library_report_links import valid_report_url
from estimator.technical_item_exports import prepare_technical_item, build_technical_item_xlsx, render_technical_item_pdf

REPORT='https://bossfire.com/documents/test-reports/Synthetic-FAS990001.pdf'
HEADERS={**{c:'Synthetic substrate '+c for c in SUBSTRATES},**EXPECTED_HEADERS}
HEADERS.update(K='Single layer plasterboard system\nBF-XXXX-1',N='Concrete wall\nBF-XXXX-4')
SOURCE={'filename':'Synthetic-matrix.xlsx','sha256':'a'*64,'sheet':'Synthetic systems','headers':HEADERS}
CATALOG={'FAS990001':[{'label':'Synthetic FAS 990001','url':REPORT}]}


def sample(**changes):
    v={'A':'BF-9001','B':'Plastic_Pipes','C':'PPR','D':'Service Penetration','E':'Plumbing / Hydraulic / Active Fire',
       'F':'Up to 32','G':0,'H':32,'I':'Wall','K':1,'N':1,'AR':90,'AS':'As required',
       'AT':'Build up locally using a synthetic board.','AU':'Synthetic sealant','BA':5,'BB':60,'BC':0,'BD':15,
       'BE':'Exposed 300 Unexposed 600','BH':'-/90/60 - RISF 24 min',
       'BJ':'32 mm PPR pipe. Apply sealant to both sides. Preserve the exposed-side condition.',
       'BK':'FAS990001 R2.0','BR':'Maintain 200 mm separation.'}
    v.update(changes);return {'row':9,'cells':v}


class BossMatrixTests(unittest.TestCase):
    def test_owner_verified_edition_pair_does_not_create_revision_issue(self):
        catalog = {'FAS190100': [{'label': 'WF FAS 190100',
                    'url': 'https://bossfire.com/documents/test-reports/WF-FAS-190100.pdf'}]}
        item = make_entry(SOURCE, sample(BK='FAS190100 R2.1'), catalog)
        self.assertFalse(any('Revision equivalence' in field.get('value', '') for field in item['fields']))
        other = make_entry(SOURCE, sample(BK='FAS190100 R2.2'), catalog)
        self.assertTrue(any('Revision equivalence' in field.get('value', '') for field in other['fields']))

    def test_mapping_reuses_existing_broad_groups_and_holds_fixture_types(self):
        before=(SERVICES,CATEGORIES)
        for subtype,expected in [('PPR','Plastic Pipes'),('PP','Plastic Pipes'),('PEX-AL','Plastic Pipes'),
                                 ('Kelox / KeKilit','Plastic Pipes'),('Iron Pipes','Unlagged Pipes'),
                                 ('Insulated Iron Pipes','Lagged Pipes'),('Copper Pipes','Copper Pipes')]:
            self.assertEqual(map_services(sample(C=subtype)['cells']),[expected])
        self.assertEqual(map_services(sample(B='Conduit',C='UPVC')['cells']),['Conduits'])
        self.assertIn('Valve/tap',held_reason(sample(B='Water_Valves_and_Tap Sets')['cells']))
        self.assertIn('floorbox',held_reason(sample(B='Other',C='GPO',BJ='Recessed electrical floorbox.')['cells']))
        self.assertIsNone(held_reason(sample(B='Other',C='GPO',BJ='Wall socket outlet.')['cells']))
        self.assertEqual((SERVICES,CATEGORIES),before)
        with self.assertRaisesRegex(ValueError,'Unmapped service'):map_services(sample(B='Unknown',C='Unknown')['cells'])

    def test_every_condition_uses_its_own_field_and_exact_source_rating(self):
        row=sample();item=make_entry(SOURCE,row,CATALOG);projected=normalize_technical_item(item)
        fields={f['label']:f for f in projected['fields']}
        self.assertEqual(fields['FRL']['value'],row['cells']['BH'])
        self.assertEqual(fields['Barrier Construction']['value'].count('• '),2)
        self.assertEqual(fields['Barrier Construction']['value'].count('90 mm'),2)
        self.assertIn('Minimum diameter (mm): 0',fields['Service Size / Configuration']['value'])
        self.assertIn('32 mm PPR pipe.',fields['Service Size / Configuration']['value'])
        self.assertIn('both sides',fields['Installation Details']['value'])
        self.assertIn('exposed-side condition',fields['Installation Details']['value'])
        self.assertIn('200 mm separation',fields['Installation Details']['value'])
        self.assertIn('Sealant depth (mm): 0',fields['Seal Depth / Fillet Size']['value'])
        self.assertEqual(fields['Service Wrap']['value'],'Wrap length (mm): Exposed 300 Unexposed 600')
        self.assertEqual(fields['Report Number']['report_links'],[{'label':row['cells']['BK'],'url':REPORT}])
        self.assertIn('Revision equivalence',fields['Source Issues']['value'])
        self.assertEqual(classify_substrates(item,projected),['Plasterboard wall','Concrete/masonry wall'])
        self.assertIn('Plumbing & Hydraulic',classify_facets(projected)['category'])
        changed=deepcopy(item);changed['fields'][0]['value']='BF-9999'
        with self.assertRaisesRegex(ValueError,'source fields changed'):normalize_technical_item(changed)

    def test_only_marked_substrates_are_searchable_and_zero_marks_stay_unresolved(self):
        row=sample(K='-',N=1)
        row['cells']['AT']='Local build-up with plasterboard and CLT.'
        item=make_entry(SOURCE,row,{})
        projected=normalize_technical_item(item)
        self.assertEqual(classify_substrates(item,projected),['Concrete/masonry wall'])
        row=sample(K='-',N='-');item=make_entry(SOURCE,row,{})
        self.assertEqual(classify_substrates(item,normalize_technical_item(item)),[])
        self.assertIn('No approved substrate',next(f['value'] for f in item['fields'] if f['label']=='Source Issues'))
        row=sample(I='Ceiling',K='-',N='-',AM=1);source=deepcopy(SOURCE);source['headers']['AM']='BOSS FlexiBatt (Horiz)'
        item=make_entry(source,row,{})
        self.assertEqual(item['filter_values']['orientation'],['Horizontal'])
        self.assertEqual(classify_substrates(item,normalize_technical_item(item)),['BOSS FlexiBatt ceiling'])

    def test_report_matches_are_exact_families_and_never_guessed(self):
        self.assertEqual(report_key('WF FAS 990001 R1.0'),'FAS990001')
        self.assertNotEqual(report_key('FAR 3921.3'),report_key('FAR 3821'))
        capture={'source_url':SOURCE_URL,'pdf_links':[{'label':'Synthetic FAS 990001','url':REPORT}]}
        self.assertEqual(report_catalog(capture),CATALOG)
        unknown=make_entry(SOURCE,sample(BK='FAS990002'),CATALOG)
        self.assertNotIn('report_links',next(f for f in unknown['fields'] if f['label']=='Report Number'))
        for url in [REPORT+'?redirect=evil',REPORT+'#',REPORT+'?',REPORT.replace('https:','http:'),
                    REPORT.replace('bossfire.com','bossfire.com.evil.test'),
                    REPORT.replace('bossfire.com','user@bossfire.com'),
                    'https://bossfire.com/documents/test-reports/../secret.pdf']:
            self.assertFalse(valid_report_url(url),url)
        self.assertTrue(valid_report_url(REPORT))

    def write_matrix(self,path,rows):
        # Test-only workbook fixtures are generated in a disposable folder.
        from openpyxl import Workbook
        wb=Workbook();ws=wb.active;ws.title=SOURCE['sheet']
        for col,value in HEADERS.items():ws[col+'8']=value
        for number,row in enumerate(rows,9):
            for col,value in row['cells'].items():ws[col+str(number)]=value
        wb.save(path);wb.close()

    def test_workbook_column_contract_duplicate_ids_and_bad_marks(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'Synthetic.xlsx';self.write_matrix(path,[sample()])
            source=read_matrix(path);self.assertEqual(len(source['rows']),1);self.assertEqual(source['rows'][0]['row'],9)
            self.write_matrix(path,[sample(),sample()])
            with self.assertRaisesRegex(ValueError,'duplicate System ID'):read_matrix(path)
            self.write_matrix(path,[sample(K=2)])
            with self.assertRaisesRegex(ValueError,'substrate mark'):read_matrix(path)
            import openpyxl
            self.write_matrix(path,[sample()])
            wb=openpyxl.load_workbook(path);wb.active['BG8']='Unexpected component column';wb.save(path);wb.close()
            with self.assertRaisesRegex(ValueError,'Matrix column changed: BG'):read_matrix(path)
            self.write_matrix(path,[sample()])
            wb=openpyxl.load_workbook(path);wb.active['BR10001']='Unexpected distant row';wb.save(path);wb.close()
            with self.assertRaisesRegex(ValueError,'matrix dimensions'):read_matrix(path)

    def test_candidate_preserves_original_records_and_refuses_an_unreviewed_rerun(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);base=root/'base';base.mkdir()
            original={'schema_version':1,'notice':'Synthetic test','documents':[],'images':[],'links':[],
                'libraries':{'technical':{'filters':[],'items':[{'id':'existing-system','title':'Existing','fields':[]}]},
                             'penetration':{'filters':[],'items':[{'id':'existing-item','title':'Existing item','fields':[]}]}}}
            (base/'library.json').write_text(json.dumps(original))
            matrix=root/'Synthetic.xlsx';self.write_matrix(matrix,[sample(),sample(A='BF-9002',B='Water_Valves_and_Tap Sets')])
            links=root/'links.json';links.write_text(json.dumps({'source_url':SOURCE_URL,'pdf_links':[{'label':'Synthetic FAS 990001','url':REPORT}]}))
            output=root/'candidate';result=prepare(matrix,base,links,output)
            self.assertEqual(result['imported'],1);self.assertEqual(len(result['held']),1)
            data=json.loads((output/'library.json').read_bytes())
            self.assertEqual(data['libraries']['technical']['items'][0],original['libraries']['technical']['items'][0])
            self.assertEqual(data['libraries']['penetration'],original['libraries']['penetration'])
            self.assertEqual(json.loads((base/'library.json').read_bytes()),original)
            self.assertEqual(ReferenceLibrary(output).listing('technical')['total'],2)
            with self.assertRaisesRegex(ValueError,'already imported'):prepare(matrix,output,links,root/'again')

    def test_pdf_and_workbook_exports_keep_conditions_and_report_link(self):
        from openpyxl import load_workbook
        from pypdf import PdfReader
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            item=make_entry(SOURCE,sample(),CATALOG)
            bundle={'schema_version':1,'notice':'Synthetic export test','documents':[],'images':[],'links':[],
                'libraries':{'technical':{'filters':[{'key':key,'label':key} for key in ('manufacturer','document','services','orientation')],'items':[item]},'penetration':{'filters':[],'items':[]}}}
            payload=json.dumps(bundle).encode()
            (root/'library.json').write_bytes(payload)
            snapshot=prepare_technical_item({'id':item['id']},library=ReferenceLibrary(root))
            before=deepcopy(snapshot)
            workbook=load_workbook(BytesIO(build_technical_item_xlsx(snapshot)),data_only=False)
            cells=[cell for sheet in workbook for row in sheet for cell in row]
            self.assertFalse(any(cell.data_type=='f' for cell in cells))
            self.assertEqual([(cell.value,cell.hyperlink.target) for cell in cells if cell.hyperlink],
                             [('FAS990001 R2.0',REPORT)])
            self.assertTrue(any('• Single layer plasterboard' in str(cell.value) and '90 mm' in str(cell.value) for cell in cells))
            self.assertTrue(any('Exposed 300 Unexposed 600' in str(cell.value) for cell in cells))
            reader=PdfReader(BytesIO(render_technical_item_pdf(snapshot)))
            pdf_text=' '.join(' '.join(page.extract_text() for page in reader.pages).split())
            for value in ['-/90/60 - RISF 24 min','exposed-side condition','Sealant depth (mm): 0',
                          'FAS990001 R2.0','32 mm PPR pipe.']:
                self.assertIn(value,pdf_text)
            self.assertNotIn('Revision equivalence', pdf_text)
            self.assertFalse(any('Revision equivalence' in str(cell.value) for cell in cells))
            links=[annotation.get_object()['/A']['/URI'] for page in reader.pages for annotation in page.get('/Annots',[])
                   if '/A' in annotation.get_object() and '/URI' in annotation.get_object()['/A']]
            self.assertEqual(links,[REPORT])
            self.assertEqual(snapshot,before)
            self.assertEqual((root/'library.json').read_bytes(),payload)


if __name__=='__main__':unittest.main()
