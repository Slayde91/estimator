"""Complete worksheet and calculator PDF HTTP boundaries on an isolated store."""

from copy import deepcopy
import hashlib
import http.client
from io import BytesIO
import json
from pathlib import Path
import tempfile
import threading
import unittest

from pypdf import PdfReader
from openpyxl import load_workbook

from estimator.catalog import ROOT
from estimator.schedule_rows import blank_schedule_defaults
from estimator.server import create_server
from estimator.storage import Store
from estimator.workbook_calculators import source_model


EXTENTS = {
    "steel_vermiculite": {"CALCULATOR": (41, 14), "SCHEDULE": (1009, 27), "BAGS": (29, 14), "SETTINGS": (568, 65)},
    "ductwork": {"CALCULATOR": (1010, 90), "SUMMARY": (45, 12), "PRODUCT SETTINGS": (164, 17)},
    "steel_board": {"START": (100, 12), "CALCULATOR": (1008, 35), "BOARD SUMMARY": (38, 12), "EXTRA BOARDS": (45, 14), "SETTINGS": (51, 17)},
}


class CalculatorPresentationApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.database = Path(cls.temp.name) / "presentation-api.sqlite3"
        cls.server = create_server(0, cls.database)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.source_hashes = {path.name: hashlib.sha256(path.read_bytes()).hexdigest()
                             for path in (ROOT / "data/calculators").glob("*.json.gz")}

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()
        cls.temp.cleanup()

    def setUp(self):
        self.store = Store(self.database)
        with self.store.connect() as db:
            db.execute("DELETE FROM calculator_states")

    def request(self, method, path, body=None, headers=None):
        connection = http.client.HTTPConnection("127.0.0.1", self.server.server_port, timeout=120)
        try:
            connection.request(method, path, body=json.dumps(body) if isinstance(body, dict) else body,
                               headers={"Content-Type": "application/json", **(headers or {})})
            response = connection.getresponse()
            return response.status, dict(response.getheaders()), response.read()
        finally:
            connection.close()

    def json_request(self, method, path, body=None, expected=200):
        status, _, payload = self.request(method, path, body)
        self.assertEqual(status, expected, payload.decode("utf-8", errors="replace")[:1000])
        return json.loads(payload)

    def worksheet(self, identity, sheet, **options):
        return self.json_request("POST", f"/api/calculators/{identity}/worksheet", {"sheet": sheet, **options})

    @staticmethod
    def cells(worksheet):
        return {cell["address"]: cell for row in worksheet["rows"] for cell in row["cells"]}

    def stored_rows(self):
        with self.store.connect() as db:
            return {table: db.execute(f"SELECT * FROM {table} ORDER BY 1").fetchall()
                    for table in ("calculator_states", "quotes", "settings")}

    def test_all_twelve_worksheets_return_every_source_row_without_saving(self):
        before = self.stored_rows()
        page_count = 0
        for identity, pages in EXTENTS.items():
            for name, (max_row, max_column) in pages.items():
                with self.subTest(identity=identity, page=name):
                    result = self.worksheet(identity, name)
                    self.assertEqual((result["sheet"], result["start_row"], result["end_row"]), (name, 1, max_row))
                    self.assertEqual((result["max_row"], result["max_column"]), (max_row, max_column))
                    self.assertEqual([row["row"] for row in result["rows"]], list(range(1, max_row + 1)))
                    self.assertTrue(all([cell["column"] for cell in row["cells"]] == result["visible_columns"] for row in result["rows"]))
                    self.assertFalse(set(result["visible_columns"]) & set(result["hidden_columns"]))
                    self.assertTrue(all(isinstance(cell["presentation"]["bold"], bool) and isinstance(cell["presentation"]["role"], str)
                                        for row in result["rows"] for cell in row["cells"]))
                    for row in result["rows"]:
                        for cell in row["cells"]:
                            if "options_ref" in cell:
                                self.assertIn(cell["options_ref"], result["option_sets"])
                                self.assertNotIn("options", cell, "Shared options must not also be repeated in every cell")
                    page_count += 1
        self.assertEqual(page_count, 12)
        self.assertEqual(self.stored_rows(), before)
        self.assertEqual({path.name: hashlib.sha256(path.read_bytes()).hexdigest()
                          for path in (ROOT / "data/calculators").glob("*.json.gz")}, self.source_hashes)

    def test_final_prepared_rows_remain_editable_and_preserve_raw_precision(self):
        before = self.stored_rows()
        precision = 12.345678901234567
        cases = [("ductwork", "CALCULATOR", "D1010"), ("steel_board", "CALCULATOR", "F1008"),
                 ("steel_vermiculite", "SCHEDULE", "J1009"), ("steel_board", "EXTRA BOARDS", "G45")]
        for identity, sheet, address in cases:
            with self.subTest(identity=identity, sheet=sheet):
                result = self.worksheet(identity, sheet, inputs={sheet: {address: precision}})
                last = self.cells(result)[address]
                self.assertTrue(last["editable"])
                self.assertFalse(last["calculated"])
                self.assertEqual(last["value"], precision)
                self.assertEqual(result["inputs"][sheet][address], precision)
        self.assertEqual(self.stored_rows(), before)

    def test_thousandth_row_and_location_survive_save_reopen_without_draft_writes(self):
        precision = 0.12345678901234566
        cases = {
            'steel_vermiculite': {'SCHEDULE': {'AA1009': 'Level 9 / East wing - grid A.1000', 'J1009': precision}},
            'steel_board': {'CALCULATOR': {'B1008': 'Level 9 / West wing', 'F1008': precision}},
            'ductwork': {'CALCULATOR': {'B1010': '400x200', 'D1010': precision}},
        }
        for identity, inputs in cases.items():
            with self.subTest(calculator=identity):
                sheet = next(iter(inputs))
                expected = blank_schedule_defaults(identity, inputs)
                before = self.stored_rows()
                self.assertEqual(self.worksheet(identity, sheet, inputs=inputs)['inputs'], expected)
                self.assertEqual(self.stored_rows(), before)
                saved = self.json_request('PUT', f'/api/calculators/{identity}/state', {'inputs': inputs})
                self.assertEqual(saved['inputs'], expected)
                self.assertEqual(Store(self.database).calculator_state(identity)['inputs'], expected)
                self.assertEqual(self.json_request('GET', f'/api/calculators/{identity}')['inputs'], expected)
                after_save = self.stored_rows()
                self.worksheet(identity, sheet, inputs={})
                self.assertEqual(self.stored_rows(), after_save)
        before = self.stored_rows()
        for identity, sheet, address in [('steel_vermiculite', 'SCHEDULE', 'AA1010'),
                                          ('steel_vermiculite', 'SCHEDULE', 'Z1009'),
                                          ('steel_board', 'CALCULATOR', 'F1009'),
                                          ('ductwork', 'CALCULATOR', 'D1011')]:
            self.json_request('PUT', f'/api/calculators/{identity}/state',
                              {'inputs': {sheet: {address: 2}}}, expected=400)
        self.assertEqual(self.stored_rows(), before)

    def test_section_controls_and_text_aliases_leave_raw_worksheet_values_intact(self):
        before = self.stored_rows()
        calculator = self.worksheet('steel_vermiculite', 'CALCULATOR')
        cells = self.cells(calculator)
        self.assertEqual(calculator['navigation_mode'], 'hidden')
        self.assertEqual(calculator['display_text']['L6'], 'PUBLISHED VALUE')
        self.assertIn('mm', cells['L6']['value'])
        self.assertIsInstance(cells['H6']['value'], (int, float))
        self.assertEqual(calculator['display_cells']['H6']['suffix'], ' mm')
        bags = self.worksheet('steel_vermiculite', 'BAGS')
        self.assertEqual(bags['display_table_order'], [1, 0])
        self.assertEqual(bags['navigation_mode'], 'hidden')
        for row in range(20, 26):
            self.assertTrue(bags['display_cells'][f'A{row}']['bold'])
        for identity, sheet, count, field in (
                ('steel_vermiculite', 'SETTINGS', 12, 'D42'),
                ('ductwork', 'PRODUCT SETTINGS', 5, 'B96'),
                ('steel_board', 'SETTINGS', 3, 'B34')):
            page = self.worksheet(identity, sheet)
            self.assertEqual(page['navigation_mode'], 'select')
            self.assertEqual(len(page['settings_sections']), count)
            self.assertEqual(len(page['rows']), page['max_row'])
            self.assertIn(field, self.cells(page))
            if identity == 'steel_vermiculite':
                self.assertFalse(page['display_cells']['D42']['bold'])
                self.assertFalse(self.cells(page)['D42']['editable'])
            elif identity == 'ductwork':
                self.assertEqual(page['display_text']['J94'], 'FYREWRAP APPLICATION TABLE')
                self.assertIn('MANUAL', self.cells(page)['J94']['value'])
                for address in ('B35', 'B65', 'B73'):
                    self.assertEqual(page['display_cells'][address]['control'], 'select')
                application = next(table for table in page['presentation_tables']
                                   if table['title_address'] == 'J94')
                self.assertIn('directional requirements', application['note'])
                self.assertIn('external 120/120/60', application['note'])
                self.assertEqual(page['warnings'], [])
        self.assertEqual(self.stored_rows(), before)

    def test_display_tabs_keep_the_original_settings_api_and_saved_input_scope(self):
        definition = self.json_request('GET', '/api/calculators/steel_vermiculite')
        self.assertEqual(definition['pages'], ['CALCULATOR', 'SCHEDULE', 'BAGS', 'SETTINGS'])
        self.assertEqual([page['id'] for page in definition['display_pages']],
                         ['START', 'CALCULATOR', 'SCHEDULE', 'BAGS', 'SETTINGS', 'FACTOR CALCS'])
        self.assertEqual([page['label'] for page in definition['display_pages']],
                         ['START', 'CALCULATOR', 'SCHEDULE', 'SUMMARY', 'SETTINGS', 'FACTOR CALCS'])
        draft = {'SETTINGS': {'D346': 123.4567890123, 'D358': 219.8765432109}}
        saved = self.json_request('PUT', '/api/calculators/steel_vermiculite/state', {'inputs': draft})
        expected = blank_schedule_defaults('steel_vermiculite', draft)
        self.assertEqual(saved['inputs'], expected)
        before = self.stored_rows()
        page = self.worksheet('steel_vermiculite', 'SETTINGS')
        for address, value in draft['SETTINGS'].items():
            self.assertEqual(self.cells(page)[address]['value'], value)
        for browser_page in ('START', 'FACTOR CALCS'):
            self.json_request('POST', '/api/calculators/steel_vermiculite/worksheet',
                              {'sheet': browser_page}, expected=400)
            self.json_request('PUT', '/api/calculators/steel_vermiculite/state',
                              {'inputs': {browser_page: {'D346': 999}}}, expected=400)
        self.assertEqual(self.json_request('GET', '/api/calculators/steel_vermiculite')['inputs'], expected)
        self.assertEqual(self.stored_rows(), before)

    def test_dropdowns_share_full_lists_and_keep_dependent_numeric_choices(self):
        inputs = {"CALCULATOR": {"C9": "TRAFALGAR COREX", "C10": "PROMATECT 250", "I9": "Beam", "I10": "Column"}}
        result = self.worksheet("steel_board", "CALCULATOR", inputs=inputs)
        cells = self.cells(result)
        choices = lambda address: result["option_sets"][cells[address]["options_ref"]]
        self.assertEqual(choices("C9"), ["TRAFALGAR COREX", "PROMATECT 250", "PROMATECT 100", "PROMATECT-XS"])
        self.assertEqual(cells["C9"]["options_ref"], cells["C1008"]["options_ref"])
        self.assertEqual(cells["D9"]["options_ref"], cells["D1008"]["options_ref"])
        self.assertEqual(len(choices("D9")), 1342)
        self.assertEqual(choices("H9"), [30, 60, 90, 120, 180])
        self.assertEqual(choices("H10"), [60, 90, 120, 180])
        self.assertTrue(all(type(value) in (int, float) for value in choices("H9")))
        self.assertEqual(choices("J9"), [620])
        self.assertEqual(choices("J10"), [550])
        self.assertTrue(cells["H9"]["allow_other"])
        self.assertEqual(cells["H9"]["error_style"], "warning")
        self.assertFalse(cells["C9"]["allow_other"])
        self.assertEqual(cells["C9"]["error_style"], "stop")
        self.assertEqual(result["inputs"], blank_schedule_defaults("steel_board", inputs))

    def test_advanced_columns_preserve_hidden_inputs_and_readonly_helpers(self):
        normal = self.worksheet("steel_board", "CALCULATOR")
        advanced = self.worksheet("steel_board", "CALCULATOR", include_advanced=True)
        self.assertNotIn("M9", self.cells(normal))
        self.assertNotIn("X1008", self.cells(normal))
        self.assertEqual(advanced["visible_columns"], list(range(1, 36)))
        self.assertTrue(self.cells(advanced)["X1008"]["editable"])
        self.assertFalse(self.cells(advanced)["AI1008"]["editable"])
        duct = self.worksheet("ductwork", "CALCULATOR", include_advanced=True)
        self.assertEqual(duct["visible_columns"], list(range(1, 91)))
        self.assertFalse(self.cells(duct)["CL1010"]["editable"])
        self.assertTrue(self.cells(duct)["CL1010"]["calculated"])

    def test_current_draft_and_saved_inputs_are_separate(self):
        saved = self.store.save_calculator_state("ductwork", {"CALCULATOR": {"D11": 3}})
        before = self.stored_rows()
        draft = {"CALCULATOR": {"D11": 12.345678901234567}, "PRODUCT SETTINGS": {"B46": 20}}
        changed = self.worksheet("ductwork", "CALCULATOR", inputs=draft)
        self.assertEqual(changed["inputs"], blank_schedule_defaults("ductwork", draft))
        self.assertEqual(self.cells(changed)["D11"]["value"], 12.345678901234567)
        self.assertEqual(self.cells(changed)["K11"]["value"], "")
        current = self.worksheet("ductwork", "CALCULATOR")
        self.assertEqual(current["inputs"], saved["inputs"])
        self.assertEqual(self.cells(current)["D11"]["value"], 3)
        self.assertEqual(self.cells(current)["K11"]["value"], "")
        self.assertEqual(self.stored_rows(), before)

    def test_formula_backed_settings_remain_formulas_until_explicitly_overridden(self):
        initial_state = self.store.calculator_state("steel_vermiculite")
        initial_rows = self.stored_rows()
        # An explicit empty input snapshot uses the retained source formulas;
        # a new unsaved state separately carries reviewed application defaults.
        page = self.worksheet("steel_vermiculite", "SETTINGS", inputs={})
        cells = self.cells(page)
        for address in ("D70", "D102", "D179", "D235"):
            self.assertTrue(cells[address]["editable"])
            self.assertTrue(cells[address]["calculated"])
            self.assertIsInstance(cells[address]["value"], (int, float))
        self.assertEqual(page["inputs"], blank_schedule_defaults("steel_vermiculite", {}))
        changed = self.worksheet("steel_vermiculite", "SETTINGS", inputs={"SETTINGS": {"D70": 0.012345678901234567}})
        self.assertEqual(self.cells(changed)["D70"]["value"], 0.012345678901234567)
        self.assertEqual(self.store.calculator_state("steel_vermiculite"), initial_state)
        self.assertEqual(self.stored_rows(), initial_rows)

    def test_source_titles_sections_and_secondary_table_headers_are_distinguished(self):
        cases = [
            ("steel_vermiculite", "CALCULATOR", {"A1": "title", "A5": "section", "H5": "section", "B28": "column_header", "K16": "output"}),
            ("steel_vermiculite", "BAGS", {"A17": "section", "A19": "column_header", "D10": "output", "H6": "note", "H10": "body", "H11": "note"}),
            ("ductwork", "CALCULATOR", {"A1": "title", "A3": "note"}),
            ("ductwork", "SUMMARY", {"A1": "compact_summary_title", "A8": "column_header", "A18": "column_header", "A30": "column_header", "A39": "column_header"}),
            ("steel_board", "START", {"A1": "compact_title", "A8": "section"}),
            ("steel_board", "EXTRA BOARDS", {"A5": "column_header", "M5": "column_header"}),
        ]
        for identity, sheet, roles in cases:
            with self.subTest(identity=identity, page=sheet):
                cells = self.cells(self.worksheet(identity, sheet))
                for address, role in roles.items():
                    self.assertEqual(cells[address]["presentation"]["role"], role, address)

    def test_worksheet_rejects_invalid_shapes_hidden_pages_and_caller_extents(self):
        before = self.stored_rows()
        invalid = [{"include_advanced": value} for value in (None, 0, 1, "true", [], {})]
        invalid += [{"sheet": value} for value in ([], {}, False, 0, "STEEL LIBRARY", "missing")]
        invalid += [{"start_row": 1}, {"row_count": 100000}, {"max_row": 1}, {"formula_overrides": {}},
                    {"inputs": []}, {"inputs": True}, {"inputs": {"CALCULATOR": {"K11": 10}}}]
        for body in invalid:
            with self.subTest(body=body):
                self.json_request("POST", "/api/calculators/ductwork/worksheet", body, expected=400)
        self.json_request("PUT", "/api/calculators/ductwork/worksheet", {}, expected=405)
        self.json_request("POST", "/api/calculators/missing/worksheet", {}, expected=400)
        self.assertEqual(self.stored_rows(), before)

    @staticmethod
    def single_row_draft(identity, marker):
        model = source_model(identity)
        schedule = model["schedule"]
        sheet = next(sheet for sheet in model["sheets"] if sheet["name"] == schedule["sheet"])
        columns = [field["column"] for field in schedule["columns"] if field["editable"]]
        cells = {f"{column}{row}": None for row in range(schedule["first_row"], schedule["last_row"] + 1) for column in columns}
        for column in columns:
            address = column + str(schedule["first_row"])
            cells[address] = sheet["cells"].get(address, {}).get("value")
        row = schedule["first_row"]
        if identity == "ductwork":
            cells[f"B{row}"] = "333x222"
            cells[f"D{row}"] = 12.345678901234567
        else:
            cells[f"A{row}"] = marker
            cells[f"F{row}" if identity == "steel_board" else f"J{row}"] = 12.345678901234567
        return {schedule["sheet"]: cells}

    def test_both_pdf_downloads_for_all_three_calculators_isolate_content_and_use_the_draft(self):
        for identity in EXTENTS:
            with self.subTest(identity=identity):
                draft = self.single_row_draft(identity, f"CURRENT DRAFT {identity}")
                extra_marker = "MATERIALS ONLY EXTRA ALLOWANCE"
                if identity == "steel_vermiculite":
                    draft['SCHEDULE']['L10'] = 987.654321
                if identity == "steel_board":
                    draft['EXTRA BOARDS'] = {'A45': extra_marker, 'B45': 'TRAFALGAR COREX',
                                             'C45': 12.5, 'G45': 1.25}
                pristine = deepcopy(draft)
                model = source_model(identity)
                sheet = model["schedule"]["sheet"]
                marker_cell = "B11" if identity == "ductwork" else "A9" if identity == "steel_board" else "A10"
                self.store.save_calculator_state(identity, {sheet: {marker_cell: "SAVED VALUE EXCLUDED FROM DRAFT"}})
                before = self.stored_rows()
                marker = "333x222" if identity == "ductwork" else f"CURRENT DRAFT {identity}"
                project_details = {"project_no": "CF-908", "client": "Draft client", "site_address": "8 Current Street"}
                for route, filename in (("report.pdf", "APPENDIX A.pdf"),
                                        ("summary.pdf", f"ceasefire-{identity}-materials-summary.pdf")):
                    with self.subTest(route=route):
                        status, headers, payload = self.request("POST", f"/api/calculators/{identity}/{route}",
                                                                {"inputs": draft, "project_details": project_details})
                        self.assertEqual(status, 200, payload[:300])
                        self.assertEqual(headers["Content-Type"], "application/pdf")
                        self.assertEqual(headers["Content-Disposition"], f'attachment; filename="{filename}"')
                        self.assertEqual(headers["Cache-Control"], "no-store")
                        self.assertEqual(int(headers["Content-Length"]), len(payload))
                        reader = PdfReader(BytesIO(payload))
                        text = "\n".join(page.extract_text() for page in reader.pages)
                        compact_text = "".join(text.split())
                        self.assertIn("Current calculator snapshot", text)
                        self.assertNotIn(model['source']['filename'], compact_text)
                        self.assertNotIn(model['source']['sha256'], compact_text)
                        self.assertNotIn('Authoritative workbook', text)
                        self.assertNotIn('SHA-256', text)
                        for label in ('Project No.', 'Client', 'Site Address'):
                            self.assertIn(label, text)
                        for value in project_details.values():
                            self.assertIn(value, text)
                        for page in reader.pages:
                            page_text = page.extract_text()
                            for contact in ('ABN: 50 612 231 562', 'Phone: 1300 92 62 88', 'sales@ceasefire.com.au'):
                                self.assertIn(contact, page_text)
                        for heading in ("Schedule inputs, calculations and complete notes",
                                        "Single-member calculator - separate from the schedule",
                                        "Manual bag calculation - separate from the schedule",
                                        "Settings used for this report"):
                            self.assertNotIn(heading, text)
                        if route == "report.pdf":
                            self.assertIn("Full schedule", text)
                            self.assertIn("".join(marker.split()), compact_text)
                            for heading in ("Final product and material summary", "Overall schedule totals",
                                            "Product totals", "Product order totals", "Board stock totals by product and thickness",
                                            "Extra-board item"):
                                self.assertNotIn(heading, text)
                            self.assertNotIn("".join(extra_marker.split()), compact_text)
                            if identity != "steel_board":
                                self.assertIn("12.35", text)
                            else:
                                self.assertIn("Net board", text)
                        else:
                            self.assertNotIn("Full schedule", text)
                            self.assertNotIn("".join(marker.split()), compact_text)
                            self.assertIn("Material quantities and summary", text)
                            if identity == "steel_vermiculite":
                                self.assertNotIn("Final product and material summary", text)
                                self.assertNotIn("Product order totals", text)
                            else:
                                self.assertIn("Final product and material summary", text)
                            self.assertNotIn("Overall schedule totals", text)
                            if identity == "steel_board":
                                self.assertIn("".join(extra_marker.split()), compact_text)
                                self.assertIn("Extra-board item 40", text)
                            elif identity == "ductwork":
                                # Surface is perimeter (1.11 m) × draft length.
                                self.assertIn("13.70", text)
                            else:
                                self.assertIn("987.65", text)
                        self.assertNotIn("SAVED VALUE EXCLUDED FROM DRAFT", text)
                        self.assertGreaterEqual(len(reader.pages), 1)
                        self.assertTrue(any(obj.get_object().get("/Subtype") == "/Image"
                                            for obj in reader.pages[0]["/Resources"]["/XObject"].values()))
                        self.assertEqual(self.stored_rows(), before)
                self.assertEqual(self.stored_rows(), before)
                self.assertEqual(draft, pristine)
        self.assertEqual({path.name: hashlib.sha256(path.read_bytes()).hexdigest()
                          for path in (ROOT / "data/calculators").glob("*.json.gz")}, self.source_hashes)

    def test_both_pdf_routes_use_saved_fallback_but_explicit_empty_inputs_use_source_values(self):
        saved = self.single_row_draft('ductwork', 'unused')
        saved['CALCULATOR'].update({'B11': '777x444', 'D11': 987.654321})
        self.store.save_calculator_state('ductwork', saved)
        before = self.stored_rows()
        for route in ('report.pdf', 'summary.pdf'):
            with self.subTest(route=route):
                status, headers, payload = self.request('POST', f'/api/calculators/ductwork/{route}', {})
                self.assertEqual(status, 200, payload[:300])
                filename = 'APPENDIX A.pdf' if route == 'report.pdf' else 'ceasefire-ductwork-materials-summary.pdf'
                self.assertEqual(headers['Content-Disposition'], f'attachment; filename="{filename}"')
                saved_text = ' '.join(page.extract_text() for page in PdfReader(BytesIO(payload)).pages)
                # The schedule shows length; its summary shows surface area.
                expected_value = '987.65' if route == 'report.pdf' else '2,411.85'
                self.assertIn(expected_value, saved_text)
                status, headers, payload = self.request('POST', f'/api/calculators/ductwork/{route}', {'inputs': {}})
                self.assertEqual(status, 200, payload[:300])
                self.assertEqual(headers['Content-Disposition'], f'attachment; filename="{filename}"')
                source_text = ' '.join(page.extract_text() for page in PdfReader(BytesIO(payload)).pages)
                self.assertNotIn(expected_value, source_text)
                self.assertNotIn('777x444', source_text)
                self.assertNotEqual(source_text, saved_text)
                self.assertEqual(self.stored_rows(), before)
        self.assertEqual(self.store.calculator_state('ductwork')['inputs'], saved)

    def test_pdf_request_boundaries_reject_invalid_inputs_without_writes(self):
        self.store.save_calculator_state("ductwork", {"CALCULATOR": {"D11": 5}})
        before = self.stored_rows()
        invalid = [{"sheet": "SUMMARY"}, {"include_advanced": True}, {"title": "Unsupported"},
                   {"inputs": []}, {"inputs": True}, {"inputs": {"CALCULATOR": {"M11": 123}}},
                   {"inputs": {"CALCULATOR": {"D11": float("inf")}}}]
        for route in ('report.pdf', 'summary.pdf'):
            for body in invalid:
                with self.subTest(route=route, body=body):
                    self.json_request("POST", f"/api/calculators/ductwork/{route}", body, expected=400)
            self.json_request("PUT", f"/api/calculators/ductwork/{route}", {}, expected=405)
            self.json_request("POST", f"/api/calculators/missing/{route}", {}, expected=400)
            for headers in ({'Origin': 'https://attacker.example'}, {'Host': 'attacker.example'}):
                with self.subTest(route=route, headers=headers):
                    self.assertEqual(self.request('POST', f'/api/calculators/ductwork/{route}', {}, headers)[0], 403)
        self.assertEqual(self.stored_rows(), before)

    def test_all_excel_registers_use_the_current_draft_without_saving(self):
        for identity in EXTENTS:
            with self.subTest(identity=identity):
                draft = self.single_row_draft(identity, f'CURRENT REGISTER {identity}')
                sheet = source_model(identity)['schedule']['sheet']
                marker_cell = 'B11' if identity == 'ductwork' else 'A9' if identity == 'steel_board' else 'A10'
                self.store.save_calculator_state(identity, {sheet: {marker_cell: 'SAVED VALUE EXCLUDED FROM REGISTER'}})
                before = self.stored_rows()
                pristine = deepcopy(draft)
                status, headers, payload = self.request('POST', f'/api/calculators/{identity}/register.xlsx', {'inputs': draft})
                self.assertEqual(status, 200, payload[:300])
                self.assertEqual(headers['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                self.assertEqual(headers['Content-Disposition'], 'attachment; filename="APPENDIX A.xlsx"')
                self.assertEqual(headers['Cache-Control'], 'no-store')
                self.assertEqual(int(headers['Content-Length']), len(payload))
                workbook = load_workbook(BytesIO(payload))
                values = [cell.value for worksheet in workbook for row in worksheet for cell in row]
                marker = '333x222' if identity == 'ductwork' else f'CURRENT REGISTER {identity}'
                self.assertIn(marker, values)
                self.assertNotIn('SAVED VALUE EXCLUDED FROM REGISTER', values)
                if identity != 'steel_board':
                    self.assertIn(12.345678901234567, values)
                self.assertEqual(workbook['Schedule'].max_row, 6)
                self.assertEqual(self.stored_rows(), before)
                self.assertEqual(draft, pristine)

    def test_excel_register_saved_fallback_and_request_boundaries(self):
        self.store.save_calculator_state('ductwork', {'CALCULATOR': {'B11': '700x500', 'D11': 5}})
        before = self.stored_rows()
        status, headers, payload = self.request('POST', '/api/calculators/ductwork/register.xlsx', {})
        self.assertEqual(status, 200, payload[:300])
        self.assertEqual(headers['Content-Disposition'], 'attachment; filename="APPENDIX A.xlsx"')
        workbook = load_workbook(BytesIO(payload))
        self.assertEqual(workbook['Schedule']['C6'].value, '700x500')
        self.assertEqual(workbook['Schedule']['D6'].value, 5)
        invalid = [{'sheet': 'SUMMARY'}, {'include_advanced': True}, {'title': 'Unsupported'},
                   {'inputs': []}, {'inputs': True}, {'inputs': {'CALCULATOR': {'M11': 123}}},
                   {'inputs': {'CALCULATOR': {'D11': float('inf')}}}]
        for body in invalid:
            with self.subTest(body=body):
                self.json_request('POST', '/api/calculators/ductwork/register.xlsx', body, expected=400)
        self.json_request('PUT', '/api/calculators/ductwork/register.xlsx', {}, expected=405)
        self.json_request('POST', '/api/calculators/missing/register.xlsx', {}, expected=400)
        self.assertEqual(self.stored_rows(), before)

    def test_pdf_retains_zero_and_incomplete_items_at_the_last_schedule_row(self):
        inputs = self.single_row_draft("steel_board", "ZERO LENGTH REVIEW")
        inputs["CALCULATOR"]["F9"] = 0
        inputs["CALCULATOR"]["A208"] = "LAST ITEM REVIEW"
        inputs["CALCULATOR"]["C208"] = "PROMATECT 250"
        inputs["CALCULATOR"]["F208"] = 0
        before = self.stored_rows()
        status, _, payload = self.request("POST", "/api/calculators/steel_board/report.pdf", {"inputs": inputs})
        self.assertEqual(status, 200, payload[:300])
        text = " ".join(" ".join(page.extract_text() for page in PdfReader(BytesIO(payload)).pages).split())
        self.assertIn("ZERO LENGTH REVIEW", text)
        self.assertIn("LAST ITEM REVIEW", text)
        self.assertIn("2 used schedule items", text)
        self.assertIn("2 item(s) have incomplete or unavailable primary quantities", text)
        self.assertIn("Not available", text)
        self.assertEqual(self.stored_rows(), before)

    def test_worksheet_and_report_keep_a_stale_saved_source_untouched(self):
        stale = {"inputs": {"CALCULATOR": {"D11": 12}}, "source_sha256": "prior-version", "updated_at": "original-time"}
        with self.store.connect() as db:
            db.execute("INSERT INTO calculator_states VALUES(?,?,?)", ("ductwork", json.dumps(stale), "original-time"))
        before = self.stored_rows()
        for route in ("worksheet", "report.pdf", "summary.pdf", "register.xlsx"):
            with self.subTest(route=route):
                error = self.json_request("POST", f"/api/calculators/ductwork/{route}", {"inputs": {}}, expected=400)
                self.assertIn("version", error["error"])
                self.assertEqual(self.stored_rows(), before)


if __name__ == "__main__":
    unittest.main()
