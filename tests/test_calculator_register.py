"""Excel registers retain report calculations as safe, full-precision values."""

from copy import deepcopy
from functools import lru_cache
from io import BytesIO
import unittest
from unittest.mock import patch
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook

from estimator.calculator_register import build_calculator_register
from estimator.calculator_report import project_calculator_report, _source_text
from estimator.workbook_calculators import source_model
from tests.test_calculator_report import cleared_inputs, REMOVED_SECTIONS
from tests.test_workbook_parity import read_fixture, scenario_inputs, excel_equal


IDS = ('ductwork', 'steel_vermiculite', 'steel_board')
VARIATIONS = {
    'ductwork': 'calibration_reference_precision',
    'steel_vermiculite': 'missing_yield_and_zero_area',
    'steel_board': 'advanced_choices_boundaries_and_all_stock',
}
SCHEDULE_COLUMNS = {
    'ductwork': {'Product': 'C', 'Duct dimensions (mm)': 'B', 'Length (m)': 'D',
                 'FRL': 'E', 'Duct surface (m²)': 'K'},
    'steel_vermiculite': {'Location': 'AA', 'Mark': 'A', 'Product': 'B', 'Section': 'F', 'Quantity': 'I',
                         'Length (m)': 'J', 'Published thickness (mm)': 'O',
                         'Estimating thickness (mm)': 'P', 'Spray surface (m²)': 'R',
                         'Net bags': 'T', 'Whole bags per line': 'U'},
    'steel_board': {'Mark': 'A', 'Location': 'B', 'Product': 'C', 'Section': 'D', 'Design period (min)': 'AN',
                    'Critical temperature (°C)': 'AO', 'Board stack (mm)': 'Z',
                    'Total thickness (mm)': 'AB', 'Box reference area (m²)': 'AD',
                    'Net board area (m²)': 'AE', 'Area with waste (m²)': 'AF', 'Sheets per line': 'AG'},
}


@lru_cache(maxsize=6)
def snapshot(identity, kind='default'):
    fixture = read_fixture(identity, kind)
    scenario = fixture['scenarios'][0] if kind == 'default' else next(
        value for value in fixture['scenarios'] if value['id'] == VARIATIONS[identity])
    inputs = scenario_inputs(scenario)
    data = project_calculator_report(identity, inputs)
    payload = build_calculator_register(identity, inputs)
    return data, load_workbook(BytesIO(payload)), scenario['expected'], payload


def sheet_text(sheet):
    return '\n'.join(str(cell.value) for row in sheet for cell in row if cell.value is not None)


def values(workbook):
    return {sheet.title: tuple(tuple(cell.value for cell in row) for row in sheet) for sheet in workbook}


class CalculatorRegisterTests(unittest.TestCase):
    def assert_cell_value(self, cell, expected):
        if expected == '':
            self.assertIsNone(cell.value, cell.coordinate)
            self.assertEqual(cell.data_type, 'inlineStr', cell.coordinate)
        else:
            self.assertEqual(cell.value, expected, cell.coordinate)
            if isinstance(expected, str):
                self.assertEqual(cell.data_type, 's', cell.coordinate)
            elif type(expected) in (int, float):
                self.assertEqual(cell.data_type, 'n', cell.coordinate)

    def assert_schedule(self, identity, data, workbook, native):
        sheet = workbook['Schedule']
        headers = {cell.value: cell.column for cell in sheet[5]}
        self.assertEqual(sheet.max_row, 5 + len(data['rows']))
        comparisons = 0
        for target_row, item in enumerate(data['rows'], 6):
            self.assert_cell_value(sheet.cell(target_row, headers['Line']), item['line'])
            for label, column in SCHEDULE_COLUMNS[identity].items():
                cell = sheet.cell(target_row, headers[label])
                self.assert_cell_value(cell, item['values'][column])
                address = column + str(item['row'])
                if address in native.get(data['sheet'], {}):
                    expected = native[data['sheet']][address]
                    actual = '' if cell.value is None and cell.data_type == 'inlineStr' else cell.value
                    self.assertTrue(excel_equal(actual, expected), (identity, address, actual, expected))
                    comparisons += 1
            if identity == 'ductwork':
                output = item['values']
                for label, expected in [('Thickness (mm)', item['wrap_layer_mm'] if item['wrap'] else output['L']),
                                        ('Thickness basis', 'Per layer' if item['wrap'] else 'Coating'),
                                        ('Net spray bags', 'N/A' if item['wrap'] else output['M']),
                                        ('Wrap material (m²)', output['N'] if item['wrap'] else 'N/A'),
                                        ('Roll equivalents', output['O'] if item['wrap'] else 'N/A')]:
                    self.assert_cell_value(sheet.cell(target_row, headers[label]), expected)
                for label, column in ([('Wrap material (m²)', 'N'), ('Roll equivalents', 'O')] if item['wrap'] else
                                      [('Thickness (mm)', 'L'), ('Net spray bags', 'M')]):
                    address = column + str(item['row'])
                    if address in native.get(data['sheet'], {}):
                        cell = sheet.cell(target_row, headers[label])
                        actual = '' if cell.value is None and cell.data_type == 'inlineStr' else cell.value
                        self.assertTrue(excel_equal(actual, native[data['sheet']][address]), (identity, address))
                        comparisons += 1
                status = output['J'] or 'No calculated status returned'
            else:
                columns = ('V', 'W') if identity == 'steel_vermiculite' else ('AR', 'AS')
                status = '\n'.join(str(item['values'][column]) for column in columns
                                   if item['values'][column] not in (None, '')) or 'No calculated status returned'
            self.assert_cell_value(sheet.cell(target_row, headers['Status']), status)
        self.assertGreaterEqual(comparisons, 5)

    def assert_summaries(self, identity, data, workbook, native):
        sheet = workbook['Summary']
        for row, (label, value) in enumerate(data['totals'], 5):
            self.assertEqual(sheet.cell(row, 1).value, label)
            self.assert_cell_value(sheet.cell(row, 5), value)
        text = sheet_text(sheet)
        self.assertIn(data['basis'], text)
        self.assertIn(data['source']['filename'], text)
        self.assertIn(data['source']['sha256'], text)
        comparisons = 0
        for table in data['summaries']:
            title_row = next(cell.row for row in sheet for cell in row if cell.value == table['title'])
            labels = [table['labels'][column] for column in table['columns']]
            header_row = next(row for row in range(title_row + 1, sheet.max_row + 1)
                              if [sheet.cell(row, column).value for column in range(1, len(labels) + 1)] == labels)
            if table['note']:
                self.assertIn(table['note'], text)
            for target_row, item in enumerate(table['rows'], header_row + 1):
                for target_column, column in enumerate(table['columns'], 1):
                    unavailable = identity == 'ductwork' and table['title'] == 'Product totals' and (
                        (item['row'] < 11 and column in 'EFG') or (item['row'] == 11 and column == 'D'))
                    address = column + str(item['row'])
                    expected = 'N/A' if unavailable else _source_text(
                        identity, table['sheet'], address, item['values'][column])
                    cell = sheet.cell(target_row, target_column)
                    self.assert_cell_value(cell, expected)
                    if not unavailable and address in native.get(table['sheet'], {}):
                        expected_native = _source_text(identity, table['sheet'], address, native[table['sheet']][address])
                        actual = '' if cell.value is None and cell.data_type == 'inlineStr' else cell.value
                        self.assertTrue(excel_equal(actual, expected_native), (identity, table['sheet'], address))
                        comparisons += 1
            for qualification in table.get('qualifications', []):
                for note in qualification:
                    if note not in (None, ''):
                        self.assertIn(str(note), text)
        for note in data.get('summary_notes', []):
            if note:
                self.assertIn(_source_text(identity, 'BOARD SUMMARY', 'A8', note), text)
        self.assertGreater(comparisons, 20)

    def test_defaults_and_changed_native_scenarios_retain_calculated_values(self):
        for identity in IDS:
            for kind in ('default', 'variations'):
                with self.subTest(identity=identity, kind=kind):
                    data, workbook, native, _ = snapshot(identity, kind)
                    self.assertEqual(workbook.sheetnames, ['Summary', 'Schedule'] +
                                     (['Extra boards'] if identity == 'steel_board' else []))
                    self.assert_schedule(identity, data, workbook, native)
                    self.assert_summaries(identity, data, workbook, native)
                    if identity == 'steel_board':
                        self.assert_extras(data, workbook, native)

    def assert_extras(self, data, workbook, native=None):
        sheet = workbook['Extra boards']
        columns = list('ABCDEFGHIJK') + ['M', 'N']
        if not data['extra_rows']:
            self.assertIn('No extra-board inputs are entered.', sheet_text(sheet))
            return
        self.assertEqual(sheet.max_row, len(data['extra_rows']) + 5)
        for target_row, item in enumerate(data['extra_rows'], 6):
            self.assertEqual(sheet.cell(target_row, 1).value, item['line'])
            for target_column, column in enumerate(columns, 2):
                cell = sheet.cell(target_row, target_column)
                self.assert_cell_value(cell, item['values'][column])
                address = column + str(item['row'])
                if native and address in native.get('EXTRA BOARDS', {}):
                    actual = '' if cell.value is None and cell.data_type == 'inlineStr' else cell.value
                    self.assertTrue(excel_equal(actual, native['EXTRA BOARDS'][address]), address)

    def test_pooled_orders_keep_the_native_rounding_and_duct_fractions(self):
        data, workbook, _, _ = snapshot('steel_board', 'variations')
        pooled = dict(data['totals'])['Pooled whole sheets']
        self.assertNotEqual(pooled, sum(item['values']['AG'] for item in data['rows']
                                       if type(item['values']['AG']) in (int, float)))
        self.assertEqual(workbook['Summary']['E7'].value, pooled)
        _, workbook, _, _ = snapshot('steel_vermiculite')
        self.assertEqual(workbook['Summary']['E8'].value, 219)
        self.assertEqual(workbook['Summary']['E7'].value, 218.37775141101426)
        data, workbook, _, _ = snapshot('ductwork')
        self.assertEqual(workbook['Schedule']['I6'].value, data['rows'][0]['values']['M'])
        self.assertNotEqual(workbook['Schedule']['I6'].value, round(workbook['Schedule']['I6'].value))

    def test_last_rows_zero_and_hidden_input_only_items_remain_incomplete(self):
        for identity, address, value in [('ductwork', 'D1010', 0),
                                         ('steel_vermiculite', 'A1009', 'Last incomplete spray member'),
                                         ('steel_board', 'X1008', 'Hidden design reference')]:
            with self.subTest(identity=identity):
                inputs = cleared_inputs(identity)
                schedule = source_model(identity)['schedule']
                inputs[schedule['sheet']][address] = value
                workbook = load_workbook(BytesIO(build_calculator_register(identity, inputs)))
                sheet = workbook['Schedule']
                self.assertEqual(sheet.max_row, 6)
                self.assertEqual(sheet['A6'].value, schedule['last_row'] - schedule['first_row'] + 1)
                self.assertIn('1 used schedule items. 1 item(s)', sheet['A3'].value)
                self.assertTrue(sheet.cell(6, sheet.max_column).value)
                if identity == 'ductwork':
                    self.assert_cell_value(sheet['D6'], 0)
                    self.assertIsNone(sheet['B6'].value)
                elif identity == 'steel_vermiculite':
                    self.assertEqual(sheet['C6'].value, value)
                else:
                    self.assertNotIn(value, sheet_text(sheet))  # Match the compact PDF scope.

    def test_steel_location_line_and_mark_are_separate_and_keep_final_row(self):
        for identity, source_location, first in [('steel_vermiculite', 'AA', 10), ('steel_board', 'B', 9)]:
            with self.subTest(identity=identity):
                inputs = cleared_inputs(identity)
                schedule = source_model(identity)['schedule']
                cells = inputs[schedule['sheet']]
                for row, text in ((first, 'First location'), (first + 999, '=literal location <west>')):
                    cells[f'{source_location}{row}'] = text
                    cells[f'A{row}'] = f'MARK-{row}'
                before = deepcopy(inputs)
                book = load_workbook(BytesIO(build_calculator_register(identity, inputs)))
                sheet = book['Schedule']
                leading = ['Line', 'Location', 'Mark'] if identity == 'steel_vermiculite' else ['Line', 'Mark', 'Location']
                self.assertEqual([cell.value for cell in sheet[5]][:3], leading)
                headers = {cell.value: cell.column for cell in sheet[5]}
                self.assert_cell_value(sheet.cell(6, headers['Line']), 1)
                self.assert_cell_value(sheet.cell(7, headers['Line']), 1000)
                self.assert_cell_value(sheet.cell(7, headers['Location']), '=literal location <west>')
                self.assert_cell_value(sheet.cell(7, headers['Mark']), f'MARK-{first + 999}')
                self.assertEqual(sheet.max_row, 7)
                self.assertEqual(inputs, before)
                book.close()

    def test_literal_text_precision_and_no_executable_workbook_content(self):
        inputs = cleared_inputs('steel_board')
        literals = ['=HYPERLINK("https://example.test", "open")', '+SUM(1,2)', '-2+3',
                    '@SUM(1,2)', '#N/A', 'https://example.test/reference']
        for row, literal in enumerate(literals, 9):
            inputs['CALCULATOR']['A' + str(row)] = literal
        long_note = '=1+1 <literal evidence> ' + 'Reference 1.2345678901234567 ' * 60
        precise = -1.2345678901234567
        inputs['EXTRA BOARDS'].update({'A45': 'Final allowance', 'G45': precise,
                                      'H45': 0.14505123456789013, 'N45': long_note})
        before = deepcopy(inputs)
        payload = build_calculator_register('steel_board', inputs)
        workbook = load_workbook(BytesIO(payload), data_only=False)
        for row, literal in enumerate(literals, 6):
            self.assert_cell_value(workbook['Schedule'].cell(row, 2), literal)
        extras = workbook['Extra boards']
        self.assert_cell_value(extras['H6'], precise)
        self.assert_cell_value(extras['I6'], inputs['EXTRA BOARDS']['H45'])
        self.assertEqual(extras['I6'].number_format, '0.00%')
        self.assert_cell_value(extras['N6'], long_note)
        self.assertEqual(inputs, before)
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    self.assertNotIn(cell.data_type, ('f', 'e'))
                    self.assertIsNone(cell.hyperlink)
        with ZipFile(BytesIO(payload)) as archive:
            self.assertFalse(any('externallink' in name.lower() or 'vbaproject' in name.lower()
                                 for name in archive.namelist()))
            for name in archive.namelist():
                if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                    root = ElementTree.fromstring(archive.read(name))
                    self.assertFalse(any(node.tag.rsplit('}', 1)[-1] in ('f', 'hyperlink') for node in root.iter()))

    def test_evidence_only_extra_board_row_is_retained_without_a_false_quantity(self):
        inputs = cleared_inputs('steel_board')
        inputs['EXTRA BOARDS']['N45'] = 'Evidence only; quantity unresolved'
        data = project_calculator_report('steel_board', inputs)
        workbook = load_workbook(BytesIO(build_calculator_register('steel_board', inputs)))
        self.assert_extras(data, workbook)
        self.assertEqual(workbook['Extra boards']['A6'].value, 40)
        self.assertEqual(workbook['Extra boards']['N6'].value, inputs['EXTRA BOARDS']['N45'])
        self.assertIsNone(workbook['Extra boards']['K6'].value)
        self.assertEqual(workbook['Summary']['E6'].value, 0)

    def test_empty_registers_keep_summaries_and_clear_empty_messages(self):
        for identity in IDS:
            with self.subTest(identity=identity):
                workbook = load_workbook(BytesIO(build_calculator_register(identity, cleared_inputs(identity))))
                self.assertEqual(workbook['Schedule']['A6'].value, 'No schedule inputs are entered.')
                self.assertIn('0 used schedule items. 0 item(s)', workbook['Schedule']['A3'].value)
                self.assertIn('Source SHA-256:', sheet_text(workbook['Summary']))
                if identity == 'steel_board':
                    self.assertEqual(workbook['Extra boards']['A5'].value, 'No extra-board inputs are entered.')

    def test_one_snapshot_excludes_separate_helpers_and_does_not_mutate_inputs(self):
        inputs = {'CALCULATOR': {'D6': 'MANDOLITE CP2', 'D16': 456.789},
                  'BAGS': {'D6': 'MANDOLITE CP2', 'D7': 1234.56789, 'D8': 19.875}}
        before = deepcopy(inputs)
        with patch('estimator.calculator_register.project_calculator_report', wraps=project_calculator_report) as project:
            workbook = load_workbook(BytesIO(build_calculator_register('steel_vermiculite', inputs)))
        project.assert_called_once_with('steel_vermiculite', inputs)
        self.assertEqual(inputs, before)
        self.assertEqual(values(workbook), values(snapshot('steel_vermiculite')[1]))
        text = '\n'.join(sheet_text(sheet) for sheet in workbook)
        for heading in REMOVED_SECTIONS:
            self.assertNotIn(heading, text)


if __name__ == '__main__':
    unittest.main()
