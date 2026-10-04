"""Defect visual annotations preserve source evidence, topology and quantities."""
from copy import deepcopy
import csv
from io import BytesIO, StringIO
import json
import unittest
from unittest.mock import patch
from uuid import uuid4

from openpyxl import load_workbook
from pypdf import PdfReader

from estimator.catalog import ValidationError
from estimator.takeoff_model import new_snapshot
from estimator.takeoff_physical import entity_references, validate_graph
from estimator.takeoff_physical_exports import export_physical_graph
from estimator.takeoff_physical_markers import defect_annotation, defect_summary, export_physical_pdf
from estimator.takeoff_physical_operations import prepare_changes
from tests import test_takeoff_physical_v2 as v2
from tests import test_takeoff_workspace as fixtures
from tests import test_takeoff_project as project_fixtures


class DefectCalloutModelTests(unittest.TestCase):
    def setUp(self):
        self.snapshot = new_snapshot(); self.document = fixtures.document()
        self.snapshot['documents'] = [self.document]
        self.annotation = {'document_id': self.document['id'], 'document_sha256': self.document['sha256'],
                           'page': 1, 'point': [50.123456789, 50.987654321]}
        self.defect = v2.create('defect', label='Observed defect', location='L01', frl='-/60/60')
        self.evidence = {key: self.annotation[key] for key in ('document_id', 'document_sha256', 'page')} | {
            'region': [[45, 45], [55, 45], [55, 55], [45, 55]], 'note': 'Original source interpretation'}
        self.defect['entity'].update(annotation=deepcopy(self.annotation), evidence=[deepcopy(self.evidence)])
        self.barrier = v2.create('barrier', self.defect['entity']['id'], substrate='Concrete')
        self.service = v2.create('service', self.barrier['entity']['id'], service='Pipe', service_type='Copper pipe')
        self.service['entity']['quantity'] = 7

    def graph(self):
        return prepare_changes(self.snapshot, [self.defect, self.barrier, self.service])['graph']

    def test_annotation_is_additive_and_old_region_display_does_not_rewrite_graph(self):
        legacy = deepcopy(self.defect); del legacy['entity']['annotation']
        graph = prepare_changes(self.snapshot, [legacy])['graph']; before = deepcopy(graph)
        locator = defect_annotation(graph['defects'][0], self.document['id'])
        self.assertEqual(locator['point'], [50, 50]); self.assertEqual(graph, before)
        self.assertNotIn('annotation', graph['defects'][0]); self.assertNotIn('marker', graph['defects'][0])
        self.assertEqual(validate_graph(graph), before)

    def test_legacy_multi_page_anchor_uses_first_original_region_globally(self):
        legacy = deepcopy(self.defect['entity']); del legacy['annotation']
        legacy['evidence'] = [{**self.evidence, 'page': 2, 'region': [[40, 50], [60, 50], [60, 70], [40, 70]]}, self.evidence]
        before = deepcopy(legacy)
        locator = defect_annotation(legacy, self.document['id'])
        self.assertEqual(locator['page'], 2); self.assertEqual(locator['point'], [50, 60])
        self.assertEqual(legacy, before); self.assertIsNone(defect_annotation(legacy, str(uuid4())))

    def test_legacy_multi_page_pdf_exports_the_same_single_anchor_without_evidence_rewrite(self):
        command = deepcopy(self.defect); del command['entity']['annotation']
        command['entity']['evidence'] = [{**self.evidence, 'page': 2, 'region': [[40, 50], [60, 50], [60, 70], [40, 70]]}, self.evidence]
        self.document['pages'].append({**deepcopy(self.document['pages'][0]), 'page': 2})
        self.snapshot['physical'] = prepare_changes(self.snapshot, [command])['graph']; before = deepcopy(self.snapshot)
        with patch('estimator.takeoff_markup_pdf.export_marked_pdf', return_value=(b'pdf', 'application/pdf', 'draft.pdf')) as export:
            export_physical_pdf(self.snapshot, {'expected_revision': 0, 'mode': 'penetrations', 'physical_scope': 'defect_reports',
                'document_id': self.document['id'], 'item_ids': [command['entity']['id']]}, object())
        rows = export.call_args.kwargs['physical_rows']; self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['geometry']['page'], 2); self.assertEqual(rows[0]['geometry']['points'], [[50, 60]])
        self.assertEqual(rows[0]['mark'], 'D-0001'); self.assertEqual(self.snapshot, before)

    def test_explicit_null_removes_legacy_fallback_without_removing_evidence(self):
        command = deepcopy(self.defect); del command['entity']['annotation']
        graph = prepare_changes(self.snapshot, [command])['graph']
        self.assertIsNotNone(defect_annotation(graph['defects'][0], self.document['id']))
        graph = v2.apply(graph, {'op': 'update', 'entity_id': command['entity']['id'], 'changes': {'annotation': None}})
        self.assertIsNone(defect_annotation(graph['defects'][0], self.document['id']))
        self.assertEqual(graph['defects'][0]['evidence'], [self.evidence]); self.assertEqual(graph['defects'][0]['annotation'], None)

    def test_move_resize_remove_restore_preserve_evidence_identity_and_all_children(self):
        graph = self.graph(); before = deepcopy(graph)
        layout = {'offset': [23.123456789, -11.987654321], 'width': 130.25, 'height': 45.75}
        for annotation in ({**self.annotation, 'point': [60.125, 70.875], 'callout': layout}, None, self.annotation):
            graph = v2.apply(graph, {'op': 'update', 'entity_id': self.defect['entity']['id'], 'changes': {'annotation': annotation}})
            defect = graph['defects'][0]
            self.assertEqual(defect['annotation'], annotation); self.assertEqual(defect['evidence'], before['defects'][0]['evidence'])
            self.assertEqual(defect['fields'], before['defects'][0]['fields']); self.assertEqual(defect['id'], before['defects'][0]['id'])
            self.assertEqual(defect['display_id'], 'D-0001'); self.assertNotIn('quantity', defect); self.assertNotIn('marker', defect)
            self.assertEqual(graph['barriers'], before['barriers']); self.assertEqual(graph['services'], before['services'])
        graph = v2.apply(graph, {'op': 'delete', 'entity_id': self.defect['entity']['id'], 'cascade': True})
        graph = v2.apply(graph, {'op': 'restore', 'entity_id': self.defect['entity']['id'], 'mode': 'same_deletion'})
        self.assertEqual(graph['defects'][0]['annotation'], self.annotation)
        self.assertEqual(graph['defects'][0]['evidence'], [self.evidence]); self.assertEqual(graph['services'][0]['quantity'], 7)

    def test_annotation_rejects_wrong_source_bounds_extra_fields_and_nonfinite_layout(self):
        before = deepcopy(self.snapshot)
        patches = [{'point': [-1, 50]}, {'page': 999}, {'document_id': str(uuid4())}, {'document_sha256': 'f' * 64},
                   {'quantity': 1}, {'point': [True, 1]}, {'point': [float('nan'), 1]},
                   {'callout': {'offset': [0, 0], 'width': 0, 'height': 5}},
                   {'callout': {'offset': [10001, 0], 'width': 5, 'height': 5}},
                   {'callout': {'offset': [0, 0], 'width': 5, 'height': float('inf')}},
                   {'callout': {'offset': [0, 0], 'width': 5, 'height': 5, 'inferred_frl': '120'}}]
        for patch in patches:
            with self.subTest(patch=patch):
                command = deepcopy(self.defect); command['entity']['annotation'].update(patch)
                with self.assertRaises(ValidationError): prepare_changes(self.snapshot, [command])
                self.assertEqual(self.snapshot, before)
        for command in (self.barrier, self.service):
            wrong = deepcopy(command); wrong['entity']['annotation'] = self.annotation
            with self.assertRaises(ValidationError): prepare_changes(self.snapshot, [self.defect, wrong])

    def test_summary_uses_only_explicit_defect_facts_and_active_linked_values(self):
        graph = self.graph()
        self.assertEqual(defect_summary(graph, graph['defects'][0]), [
            'D-0001 · Observed defect · L01 · FRL -/60/60',
            'B-0001 · Concrete · FRL -/60/60', 'S-0001 · 7 × · Pipe · Copper pipe'])
        self.assertIn('7 ×', '\n'.join(defect_summary(graph, graph['defects'][0])))
        graph['defects'][0]['fields'].pop('frl'); graph['barriers'][0]['fields']['frl'] = 'custom child rating'
        self.assertNotIn('FRL', '\n'.join(defect_summary(graph, graph['defects'][0])))
        graph['services'][0]['deleted'] = True
        self.assertEqual(defect_summary(graph, graph['defects'][0])[-1], '0 services')
        self.assertEqual(entity_references(graph['defects'][0]), [self.evidence, self.annotation])

    def test_summary_retains_every_explicit_member_parent_and_collection_order(self):
        masonry = v2.create('barrier', self.defect['entity']['id'], substrate='Masonry')
        duplicate = v2.create('barrier', self.defect['entity']['id'], substrate='Concrete')
        unknown = v2.create('barrier', self.defect['entity']['id'])
        other = v2.create('defect', label='Other defect')
        unrelated = v2.create('barrier', other['entity']['id'], substrate='Unrelated substrate')
        cable = v2.create('service', masonry['entity']['id'], service_type='Cable bundle')
        pipe = v2.create('service', duplicate['entity']['id'], service_type='Copper pipe')
        untyped = v2.create('service', unknown['entity']['id'], service='Not a recorded service type')
        excluded = v2.create('service', unrelated['entity']['id'], service_type='Unrelated service type')
        historical = v2.create('service', self.barrier['entity']['id'], service_type='Deleted service type')
        graph = prepare_changes(self.snapshot, [self.defect, self.barrier, masonry, duplicate, unknown,
            other, unrelated, self.service, cable, pipe, untyped, excluded, historical])['graph']
        graph = v2.apply(graph, {'op': 'delete', 'entity_id': historical['entity']['id'], 'cascade': False})
        # Render order follows the retained collections, not a deduplicated value set.
        graph['barriers'][:2] = reversed(graph['barriers'][:2])
        graph['services'][:2] = reversed(graph['services'][:2])
        before = deepcopy(graph)
        lines = defect_summary(graph, graph['defects'][0])
        self.assertEqual([line.split(' · ')[0] for line in lines], ['D-0001', 'B-0002', 'S-0002', 'B-0001', 'S-0001', 'B-0003', 'S-0003', 'B-0004', 'S-0004'])
        self.assertIn('Masonry', lines[1]); self.assertIn('Cable bundle', lines[2]); self.assertIn('Copper pipe', lines[4]); self.assertIn('Service type not recorded', lines[-1])
        self.assertTrue(all('B-' not in line for line in lines if line.startswith('S-')))
        self.assertEqual(graph, before)
        graph['services'][0]['fields']['service_type'] = 'Updated\nservice\ttype'
        self.assertTrue(any(line.startswith('S-0002 · 1 ×') and 'Updated service type' in line for line in defect_summary(graph, graph['defects'][0])))

    def test_deleted_barrier_children_and_category_do_not_supply_defect_values(self):
        graph = self.graph()
        graph = v2.apply(graph, {'op': 'delete', 'entity_id': self.barrier['entity']['id'], 'cascade': True})
        self.assertEqual(defect_summary(graph, graph['defects'][0]), [
            'D-0001 · Observed defect · L01 · FRL -/60/60', '0 substrates · 0 services'])
        graph = self.graph(); graph['services'][0]['fields'].pop('service_type')
        self.assertEqual(defect_summary(graph, graph['defects'][0])[-1],
                         'S-0001 · 7 × · Pipe · Service type not recorded')

    def test_csv_xlsx_retain_separate_annotation_bytes_without_new_evidence_rows(self):
        graph = self.graph(); before = deepcopy(graph)
        rows = list(csv.DictReader(StringIO(export_physical_graph(graph, 'csv')[0].decode('utf-8-sig'))))
        defect = next(row for row in rows if row['entity_type'] == 'defect')
        self.assertEqual(json.loads(defect['annotation_json']), self.annotation)
        self.assertEqual(json.loads(defect['evidence_json']), [self.evidence]); self.assertEqual(defect['quantity'], '')
        workbook = load_workbook(BytesIO(export_physical_graph(graph, 'xlsx')[0])); self.addCleanup(workbook.close)
        rows = list(workbook['Defects'].values); row = dict(zip(rows[0], rows[1]))
        self.assertEqual(json.loads(row['annotation_json']), self.annotation)
        self.assertEqual(workbook['Evidence'].max_row, 2); self.assertEqual(graph, before)


class DefectCalloutProjectTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): project_fixtures.TakeoffProjectTests.setUpClass()

    def setUp(self):
        self.case = project_fixtures.TakeoffProjectTests(); self.case.setUp(); self.addCleanup(self.case.doCleanups)
        self.doc = self.case.session['snapshot']['documents'][0]
        self.annotation = {'document_id': self.doc['id'], 'document_sha256': self.doc['sha256'], 'page': 1,
                           'point': [120.123456789, 300.987654321], 'callout': {'offset': [60.125, -35.875], 'width': 180.25, 'height': 65.75}}
        self.defect = v2.create('defect', label='Defect source', location='Level 1', frl='-/60/60')
        self.evidence = {key: self.annotation[key] for key in ('document_id', 'document_sha256', 'page')} | {
            'region': [[115, 295], [125, 295], [125, 305], [115, 305]], 'note': 'Original source location'}
        self.defect['entity'].update(annotation=self.annotation, evidence=[self.evidence]); self.apply([self.defect])

    def apply(self, commands):
        case = self.case
        preview = case.service.preview_physical(case.session['session_id'], {'expected_revision': case.session['revision'], 'scope': 'defect_reports', 'commands': commands})
        case.session = case.service.apply_physical(case.session['session_id'], {'expected_revision': case.session['revision'], 'scope': 'defect_reports', 'request_id': str(uuid4()), 'preview_id': preview['preview_id']})

    def test_real_pdf_contains_all_linked_values_without_changing_source_or_graph(self):
        concrete = v2.create('barrier', self.defect['entity']['id'], substrate='Concrete wall')
        masonry = v2.create('barrier', self.defect['entity']['id'], substrate='Masonry wall')
        pipe = v2.create('service', concrete['entity']['id'], service_type='Copper pipe')
        cable = v2.create('service', masonry['entity']['id'], service_type='Cable bundle')
        other = v2.create('defect', label='Separate source annotation')
        unrelated = v2.create('barrier', other['entity']['id'], substrate='Unrelated substrate')
        excluded = v2.create('service', unrelated['entity']['id'], service_type='Unrelated service')
        self.apply([concrete, masonry, pipe, cable, other, unrelated, excluded])
        case = self.case; before = deepcopy(case.session['snapshot'])
        source = case.documents.document_path(self.doc).read_bytes()
        payload, mime, _ = case.service.export_workspace(case.session['session_id'], 'marked-pdf', {
            'expected_revision': case.session['revision'], 'mode': 'penetrations', 'physical_scope': 'defect_reports',
            'document_id': self.doc['id'], 'item_ids': [self.defect['entity']['id']]})
        self.assertEqual(mime, 'application/pdf')
        text = ' '.join(' '.join(page.extract_text().split()) for page in PdfReader(BytesIO(payload)).pages)
        for expected in ('B-0001', 'Concrete wall', 'B-0002', 'Masonry wall',
                         'S-0001 · 1 × · Copper pipe', 'S-0002 · 1 × · Cable bundle'):
            self.assertIn(expected, text)
        self.assertNotIn('Unrelated substrate', text); self.assertNotIn('Unrelated service', text)
        self.assertNotIn('more services', text)
        self.assertNotIn('more linked records', text)
        self.assertEqual(case.documents.document_path(self.doc).read_bytes(), source)
        self.assertEqual(case.service.get(case.session['session_id'])['snapshot'], before)

    def test_real_pdf_export_save_reopen_preserve_source_and_calculator_bytes(self):
        case = self.case; before = deepcopy(case.session['snapshot']); source = case.documents.document_path(self.doc).read_bytes()
        payload, mime, _ = case.service.export_workspace(case.session['session_id'], 'marked-pdf', {
            'expected_revision': case.session['revision'], 'mode': 'penetrations', 'physical_scope': 'defect_reports',
            'document_id': self.doc['id'], 'item_ids': [self.defect['entity']['id']]})
        self.assertEqual(mime, 'application/pdf'); text = ' '.join(' '.join(page.extract_text().split()) for page in PdfReader(BytesIO(payload)).pages)
        for expected in ('D-0001', 'Defect source', 'Level 1', 'FRL -/60/60', '0 substrates', '0 services'):
            self.assertIn(expected, text)
        self.assertEqual(case.documents.document_path(self.doc).read_bytes(), source)
        self.assertEqual(case.service.get(case.session['session_id'])['snapshot'], before)
        case.library.save_as({**deepcopy(case.base), 'takeoffs': before, 'takeoffs_session_id': case.session['session_id']})
        saved = json.loads(case.target.read_bytes()); case.dialogs.opened = str(case.target); reopened = case.library.open_file()
        for value in (saved, reopened):
            self.assertEqual(value['takeoffs']['physical'], before['physical'])
            self.assertEqual(value['takeoffs']['physical']['defects'][0]['evidence'], [self.evidence])
        # The public open response includes recalculated display metadata. The
        # saved calculator/estimate payload stays byte-equivalent to the fixture,
        # and reopening retains every calculator input and schedule row.
        for key in ('estimate', 'calculators'): self.assertEqual(saved[key], json.loads(case.legacy)[key])
        for name, calculator in saved['calculators'].items():
            for key in ('inputs', 'schedule_rows'): self.assertEqual(reopened['calculators'][name][key], calculator[key])
        self.assertEqual(reopened['takeoffs_issues'], []); case.documents.validate_audit(reopened['takeoffs'])


if __name__ == '__main__': unittest.main()
