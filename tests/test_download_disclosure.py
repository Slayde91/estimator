"""Generated downloads publish report IDs without internal assessment evidence."""
from copy import deepcopy
from io import BytesIO
import json
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader

from estimator.calculator_report import build_calculator_report, build_calculator_summary_report
from estimator.calculator_register import build_calculator_register
from estimator.download_disclosure import calculator_download_data, without_fingerprints
from estimator.excel_engine import column_name
from estimator.technical_item_exports import build_technical_item_xlsx, render_technical_item_pdf
from estimator.takeoff_exports import linked_result_text


PRIVATE = 'INTERNAL-ASSESSMENT-DETAIL-ONLY'
HASH = 'd9309321a9ca747cf8baf6256acd5516af41f67ba998980a6cee6e2ffba7d243'
REPORT = 'FC13107-01-05'
URL = 'https://example.com/report.pdf'


def snapshot(identity):
    values = {column_name(index): 1 for index in range(1, 95)}
    values.update(A='Member 1', AA='Level 1', B='Product', C='Case', F='Section 1')
    if identity == 'ductwork':
        values.update(B='400 x 300', C='CAFCO 300', D=5, E='-/120/120', J=PRIVATE, AP=PRIVATE, AQ=PRIVATE)
    elif identity == 'steel_board':
        values.update(B='Level 1', C='PROMATECT 100', D='100UC15', F=5,
                      AI=PRIVATE, AR=PRIVATE, AS=PRIVATE, AT=PRIVATE)
    else:
        values.update(V=PRIVATE, W=PRIVATE, X=PRIVATE, Y=PRIVATE, O=26, J=5)
    titles = {'ductwork': 'Product totals', 'steel_vermiculite': 'Product order totals',
              'steel_intumescent': 'Product orders', 'steel_board': 'Board stock totals by product and thickness'}
    columns = 'ABCDEFGHIJK' if identity == 'steel_board' else 'ABCDEFGHIJ' if identity == 'ductwork' else 'ABCDEFGHI'
    summary_values = {key: 1 for key in columns}
    summary_values['A'] = 'Product'
    if identity == 'steel_board':
        summary_values['K'] = PRIVATE
    elif identity != 'ductwork':
        summary_values['I'] = PRIVATE
    return {'id': identity, 'title': identity, 'source': {'filename': PRIVATE, 'sha256': HASH},
            'inputs': {}, 'sheet': 'SCHEDULE', 'basis': PRIVATE, 'incomplete_rows': 0,
            'assessment_scope': PRIVATE, 'assessment_receipt': HASH, 'assessment_notes': [PRIVATE],
            'application_notes': [PRIVATE], 'application_summary_note': PRIVATE, 'summary_notes': [PRIVATE],
            'totals': [('Area', 20)], 'extra_rows': [],
            'rows': [{'line': 1, 'row': 10, 'values': values, 'complete': True,
                      'product_code': 1, 'wrap': False, 'wrap_layer_mm': None,
                      'assessment_source': PRIVATE, 'assessment_reference': PRIVATE,
                      'report_references': [{'label': REPORT, 'url': URL}]}],
            'summaries': [{'title': titles[identity], 'sheet': 'SUMMARY', 'columns': list(columns),
                           'labels': {key: key for key in columns}, 'note': PRIVATE,
                           'qualifications': [('Product', 'Part A: 25.00 kg', 'Pack: 26.00 kg / 18.00 L.')]
                           if identity == 'steel_intumescent' else [('Product', PRIVATE, HASH)],
                           'rows': [{'row': 9, 'values': summary_values}]}]}


def pdf_text(payload):
    reader = PdfReader(BytesIO(payload))
    text = '\n'.join(page.extract_text() for page in reader.pages)
    return reader, text + str(reader.metadata)


def workbook_xml(payload):
    with ZipFile(BytesIO(payload)) as archive:
        return '\n'.join(archive.read(name).decode('utf-8') for name in archive.namelist()
                         if name.endswith(('.xml', '.rels')))


class DownloadDisclosureTests(unittest.TestCase):
    def test_all_calculator_formats_keep_report_and_quantities_without_private_evidence(self):
        for identity in ('ductwork', 'steel_vermiculite', 'steel_intumescent', 'steel_board'):
            data = snapshot(identity)
            original = deepcopy(data)
            with self.subTest(identity=identity), \
                    patch('estimator.calculator_report.project_calculator_report', return_value=data), \
                    patch('estimator.calculator_register.project_calculator_report', return_value=data):
                for builder in (build_calculator_report, build_calculator_summary_report):
                    reader, text = pdf_text(builder(identity))
                    self.assertIn(REPORT, text)
                    self.assertNotIn(PRIVATE, text)
                    self.assertNotIn(HASH, text)
                    links = [ref.get_object()['/A'].get('/URI') for page in reader.pages
                             for ref in page.get('/Annots', []) if '/A' in ref.get_object()]
                    self.assertIn(URL, links)
                payload = build_calculator_register(identity)
                xml = workbook_xml(payload)
                self.assertIn(REPORT, xml)
                self.assertNotIn(PRIVATE, xml)
                self.assertNotIn(HASH, xml)
                workbook = load_workbook(BytesIO(payload))
                self.addCleanup(workbook.close)
                self.assertTrue(all(sheet.sheet_state == 'visible' for sheet in workbook))
                self.assertFalse(any(cell.comment or cell.data_type == 'f' for sheet in workbook for row in sheet for cell in row))
                self.assertEqual(workbook['Summary']['E5'].value, 20)
                headers = [cell.value for cell in workbook['Schedule'][5]]
                self.assertNotIn('Status', headers)
                self.assertNotIn('Status / basis', headers)
                self.assertNotIn('Table A/V (m⁻¹)', headers)
                self.assertEqual(data, original)

    def test_projection_omits_extra_board_internal_status_and_evidence_without_mutation(self):
        data = snapshot('steel_board')
        data['extra_rows'] = [{'values': {'A': 'Allowance', 'G': 3.25, 'M': PRIVATE, 'N': 'Customer allowance note'}, 'labels': {}}]
        original = deepcopy(data)
        public = calculator_download_data(data)
        self.assertEqual(public['extra_rows'][0]['values'], {'A': 'Allowance', 'G': 3.25, 'N': 'Customer allowance note'})
        self.assertEqual(data, original)

    def test_technical_item_keeps_installation_fields_and_report_links_only_from_provenance(self):
        data = {'record': {'id': 'example', 'title': 'Pipe system', 'subtitle': '',
                           'notice': PRIVATE, 'diagram_status': PRIVATE},
                'fields': [{'label': 'Report Number', 'value': REPORT, 'report_links': [{'label': REPORT, 'url': URL}]},
                           {'label': 'Service', 'value': 'Copper pipe 25 mm', 'report_links': []},
                           {'label': 'Reference qualification', 'value': PRIVATE, 'report_links': []}],
                'tables': [{'label': 'Installation options', 'caption': PRIVATE,
                            'columns': ['Service', 'FRL'], 'rows': [['Copper', '-/120/120']]}],
                'images': [], 'sources': [[PRIVATE, 'private.pdf', 'private', 12, 'Table 7', 42]],
                'related': [['Candidate', 'private', PRIVATE, PRIVATE]]}
        original = deepcopy(data)
        _, text = pdf_text(render_technical_item_pdf(data))
        payload = build_technical_item_xlsx(data)
        for output in (text, workbook_xml(payload)):
            self.assertIn(REPORT, output)
            self.assertIn('Copper pipe 25 mm', output)
            self.assertNotIn(PRIVATE, output)
            self.assertNotIn('Source Information', output)
        workbook = load_workbook(BytesIO(payload))
        self.addCleanup(workbook.close)
        links = [cell.hyperlink.target for sheet in workbook for row in sheet for cell in row if cell.hyperlink]
        self.assertEqual(links, [URL])
        self.assertEqual(data, original)

    def test_drawing_geometry_retained_while_fingerprints_are_not_exported(self):
        source = {'document_id': 'document', 'document_sha256': HASH, 'region': [1.25, 2.5, 3, 4],
                  'nested': [{'source_sha256': HASH, 'image_sha256': HASH, 'page': 2,
                              'library_quantity': {'state': 'unknown', 'metadata_sha256': HASH},
                              'future_asset_sha256': HASH, 'future_source_sha256s': [HASH]}],
                  'notes': '=Literal user text'}
        original = deepcopy(source)
        public = without_fingerprints(source)
        self.assertEqual(public, {'document_id': 'document', 'region': [1.25, 2.5, 3, 4],
                                  'nested': [{'page': 2, 'library_quantity': {'state': 'unknown'}}],
                                  'notes': '=Literal user text'})
        self.assertNotIn(HASH, json.dumps(public))
        self.assertEqual(source, original)

    def test_unavailable_takeoff_result_does_not_disclose_lookup_reason(self):
        result = [{'calculator_id': 'steel_intumescent', 'status': 'Unavailable', 'reason': PRIVATE}]
        self.assertIn('Unavailable', linked_result_text(result))
        self.assertNotIn(PRIVATE, linked_result_text(result))
        self.assertEqual(result[0]['reason'], PRIVATE)


if __name__ == '__main__':
    unittest.main()
