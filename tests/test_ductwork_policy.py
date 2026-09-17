"""Approved choices preserve source evidence and unknown historical inputs."""

from copy import deepcopy
import hashlib
from io import BytesIO
import json
import unittest

from openpyxl import Workbook, load_workbook

from estimator.catalog import ValidationError
from estimator.ductwork_policy import (DUCTWORK_CHOICES, apply_ductwork_choices, application_frl,
    canonical_ductwork_value, canonical_exposure, canonical_frl, legacy_exhaust_frl)
from estimator.pricing_workbook import _serialize_exact
from estimator.schedule_rows import empty_schedule_inputs
from estimator.schedule_workbook import export_schedule_template, import_schedule_workbook
from estimator.workbook_catalog import DATA_DIRECTORY, load_workbook_catalog
from estimator.workbook_calculators import calculate_page, normalize_calculator_inputs
from estimator.workbook_runtime import load_application_catalog


class DuctworkPolicyTests(unittest.TestCase):
    def test_frl_aliases_are_exact_and_do_not_round_unsupported_ratings(self):
        for minutes in (60, 90, 120, 180):
            for value in (minutes, float(minutes), str(minutes), f' {minutes}.00 ', f' {minutes}/{minutes}/{minutes} '):
                with self.subTest(value=value):
                    self.assertEqual(canonical_frl(value), '/'.join([str(minutes)] * 3))
        for value in ('120/120/-', '120/120/60', ' 120 / 120 / 60 '):
            self.assertEqual(canonical_frl(value), '120/120/120')
        for value in (None, '', True, False, 120.0000000001, 240, '75/75/75', '120/120/119.99', 'Unknown rating', 'NaN', 'Infinity'):
            with self.subTest(value=value):
                self.assertEqual(canonical_frl(value), value)

    def test_known_product_and_orientation_aliases_preserve_unknown_values(self):
        for column, value, expected in [('C', 'CAFCO', 'CAFCO 300'), ('C', ' cafco 300 ', 'CAFCO 300'),
                ('C', 'monokote', 'MONOKOTE'), ('I', 'Mixed', 'Both'), ('I', ' mixed ', 'Both'),
                ('I', 'Horizontal', 'Horizontal')]:
            self.assertEqual(canonical_ductwork_value(column, value), expected)
        for column, value in [('C', 'Legacy coating'), ('I', 'Diagonal'), ('D', 1.23456789), ('H', 'Unknown use')]:
            self.assertEqual(canonical_ductwork_value(column, value), value)
        self.assertEqual(canonical_ductwork_value('H', 'Smoke exhaust'), 'Both')

    def test_exposure_alias_helper_keeps_pressurisation_and_unknown_identity(self):
        for name in ('Kitchen exhaust - inside', 'Kitchen exhaust - outside', 'Diesel pump ventilation', 'Other exhaust'):
            self.assertEqual(canonical_exposure(name), 'Internal')
        for name in ('Kitchen + smoke exhaust', 'Smoke exhaust', 'Stair pressure relief', 'Mixed'):
            self.assertEqual(canonical_exposure(name), 'Both')
        for name in ('Stair pressurisation', 'Other pressurisation', 'Internal', 'External', 'Both', 'Legacy application'):
            self.assertEqual(canonical_exposure(name), name)

    def test_legacy_exhaust_minimum_never_downgrades_or_changes_spray(self):
        for name in ('Kitchen exhaust - inside', 'Kitchen exhaust - outside', 'Diesel pump ventilation',
                     'Other exhaust', 'Kitchen + smoke exhaust', 'Smoke exhaust', 'Stair pressure relief'):
            for frl in (60, 90, '60/60/60', '90/90/90'):
                self.assertEqual(legacy_exhaust_frl('FyreWrap', name, frl), '120/120/120')
            for frl in ('120/120/120', '180/180/180', '240/240/180', '75/75/75', None):
                self.assertEqual(legacy_exhaust_frl('FyreWrap', name, frl), frl)
        for product, exposure in [('MONOKOTE', 'Smoke exhaust'), ('CAFCO', 'Smoke exhaust'),
                                  ('FyreWrap', 'Internal'), ('FyreWrap', 'Mixed'), ('FyreWrap', 'Stair pressurisation')]:
            self.assertEqual(legacy_exhaust_frl(product, exposure, 60), '60/60/60')

    def test_normalization_handles_every_runtime_row_without_mutating_inputs(self):
        data = {'CALCULATOR': {f'{column}{row}': value for row in (11, 310, 311, 1010)
                              for column, value in [('C', 'CAFCO'), ('E', ' 120/120/60 '), ('H', 'Smoke exhaust'), ('I', 'Mixed'), ('D', 1.23456789012345)]}}
        before = deepcopy(data)
        result = normalize_calculator_inputs('ductwork', data)
        self.assertEqual(data, before)
        for row in (11, 310, 311, 1010):
            self.assertEqual(result['CALCULATOR'][f'C{row}'], 'CAFCO 300')
            self.assertEqual(result['CALCULATOR'][f'E{row}'], '120/120/120')
            self.assertEqual(result['CALCULATOR'][f'I{row}'], 'Both')
            self.assertEqual(result['CALCULATOR'][f'H{row}'], 'Both')
            self.assertEqual(result['CALCULATOR'][f'D{row}'], before['CALCULATOR'][f'D{row}'])
        self.assertEqual(normalize_calculator_inputs('ductwork', result), result)
        unknown = {'CALCULATOR': {'C11': 'Legacy product', 'E11': '75/75/75', 'H11': 'Unknown use', 'I11': 'Diagonal'}}
        self.assertEqual(normalize_calculator_inputs('ductwork', unknown), {'CALCULATOR': {**unknown['CALCULATOR'], 'I13': 'Both'}})
        self.assertEqual(normalize_calculator_inputs('ductwork', {'CALCULATOR': {'E11': '', 'I11': None}}),
                         {'CALCULATOR': {'E11': None, 'I11': None, 'I13': 'Both'}})
        for value in (True, float('inf'), float('nan'), '120\n'):
            with self.assertRaises(ValidationError):
                normalize_calculator_inputs('ductwork', {'CALCULATOR': {'E11': value}})

    def test_application_rating_uses_highest_case_without_filling_or_downgrading(self):
        for exposure in ('Internal', 'External', 'Both', 'Stair pressurisation', 'Other pressurisation', 'Smoke exhaust'):
            for rating in (60, 90, '60/60/60', '90/90/90', '120/120/-', '120/120/60'):
                self.assertEqual(application_frl('FyreWrap', exposure, rating), '120/120/120')
            for rating in (None, '', '75/75/75', '180/180/180', '240/240/180', 120.1):
                self.assertEqual(application_frl('FyreWrap', exposure, rating), rating)
        for exposure in ('Internal', 'Both', 'Kitchen exhaust - inside'):
            self.assertEqual(application_frl('FyreWrap', exposure, ' - / 30 / 30 '), '120/120/120')
        self.assertEqual(application_frl('MONOKOTE', 'Internal', 60), '60/60/60')
        self.assertEqual(application_frl('FyreWrap', 'Unknown use', 60), '60/60/60')

    def test_effective_source_fallbacks_and_partial_edits_are_idempotent(self):
        self.assertEqual(normalize_calculator_inputs('ductwork', {}), {'CALCULATOR': {'I13': 'Both'}})
        # C13/H13 come from the existing FyreWrap Internal example.
        result = normalize_calculator_inputs('ductwork', {'CALCULATOR': {'E13': 60}})
        self.assertEqual(result, {'CALCULATOR': {'E13': '120/120/120', 'I13': 'Both'}})
        # A product-only edit must evaluate the effective source exposure.
        result = normalize_calculator_inputs('ductwork', {'CALCULATOR': {'C11': 'FyreWrap', 'E11': 60}})
        self.assertEqual(result['CALCULATOR']['E11'], '120/120/120')
        result = normalize_calculator_inputs('ductwork', {'CALCULATOR': {'C11': 'FyreWrap', 'E11': 60, 'H11': 'Smoke exhaust'}})
        self.assertEqual(result['CALCULATOR']['E11'], '120/120/120')
        self.assertEqual(result['CALCULATOR']['H11'], 'Both')
        self.assertEqual(normalize_calculator_inputs('ductwork', result), result)

    def test_explicit_blanks_and_empty_schedules_never_resurrect_source_examples(self):
        empty = empty_schedule_inputs('ductwork')
        self.assertEqual(normalize_calculator_inputs('ductwork', empty), empty)
        partial = {'CALCULATOR': {'C13': None, 'E13': None, 'H13': None, 'I13': None}}
        self.assertEqual(normalize_calculator_inputs('ductwork', partial), partial)
        unknown = {'CALCULATOR': {'C1010': 'FyreWrap', 'E1010': None, 'H1010': 'Smoke exhaust'}}
        result = normalize_calculator_inputs('ductwork', unknown)
        self.assertIsNone(result['CALCULATOR']['E1010'])
        self.assertEqual(result['CALCULATOR']['H1010'], 'Both')
        self.assertNotIn('E1009', result['CALCULATOR'])

    def test_direct_schedule_import_returns_canonical_ratings_and_names(self):
        model = load_application_catalog('ductwork')
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'CALCULATOR'
        sheet.append([field['label'] for field in model['schedule']['columns'] if field['editable']])
        sheet.append(['250x250', 'FyreWrap', 10, '-/30/30', 1, 0, 'Kitchen exhaust - inside', 'Mixed'])
        sheet.append(['250x250', 'FyreWrap', 10, 90, 0, 0, 'External', 'Horizontal'])
        sheet.append(['250x250', 'FyreWrap', 10, '180/180/180', 0, 0, 'Smoke exhaust', 'Horizontal'])
        payload = _serialize_exact(workbook)
        workbook.close()
        imported = import_schedule_workbook('ductwork', payload, 'schedule.xlsx')
        self.assertEqual(imported['imported_rows'], 3)
        self.assertEqual(imported['schedule_rows'], [11, 12, 13])
        cells = imported['inputs']['CALCULATOR']
        self.assertEqual((cells['E11'], cells['H11'], cells['I11']), ('120/120/120', 'Internal', 'Both'))
        self.assertEqual(cells['E12'], '120/120/120')
        self.assertEqual((cells['E13'], cells['H13']), ('180/180/180', 'Both'))
        self.assertTrue(all(cells[f'{column}14'] is None for column in 'BCDEFGHI'))

    def test_source_formula_cells_and_packages_are_unchanged_by_metadata_policy(self):
        package = DATA_DIRECTORY / 'ductwork.json.gz'
        fingerprint = hashlib.sha256(package.read_bytes()).hexdigest()
        model = load_workbook_catalog('ductwork')
        before = hashlib.sha256(json.dumps([sheet['cells'] for sheet in model['sheets']], sort_keys=True).encode()).hexdigest()
        original_choices = deepcopy(model['schedule']['columns'])
        apply_ductwork_choices(model)
        after = hashlib.sha256(json.dumps([sheet['cells'] for sheet in model['sheets']], sort_keys=True).encode()).hexdigest()
        self.assertEqual(before, after)
        self.assertEqual(hashlib.sha256(package.read_bytes()).hexdigest(), fingerprint)
        self.assertEqual(load_workbook_catalog('ductwork')['schedule']['columns'], original_choices)

    def test_runtime_sheet_schedule_and_xml_metadata_have_identical_strict_choices(self):
        model = load_application_catalog('ductwork')
        sheet = next(item for item in model['sheets'] if item['name'] == 'CALCULATOR')
        fields = {field['column']: field for field in model['schedule']['columns']}
        rules = {rule['sqref']: rule for rule in sheet['validations']}
        xml_rules = {node['attributes']['sqref']: node for parent in sheet['metadata'] if parent['tag'] == 'dataValidations' for node in parent['children']}
        for column, choices in DUCTWORK_CHOICES.items():
            reference = f'{column}11:{column}1010'
            rule = rules[reference]
            self.assertEqual(fields[column]['validation'], rule)
            self.assertEqual(rule['formula1'], '"' + ','.join(choices) + '"')
            self.assertEqual({key: rule[key] for key in ('allowBlank', 'errorStyle', 'showErrorMessage')},
                             {'allowBlank': '1', 'errorStyle': 'stop', 'showErrorMessage': '1'})
            self.assertEqual(xml_rules[reference]['attributes'], {key: value for key, value in rule.items() if key != 'formula1'})
            self.assertEqual(xml_rules[reference]['children'][0]['text'], rule['formula1'])

    def test_api_controls_keep_unknown_values_visible_without_custom_entry(self):
        inputs = {'CALCULATOR': {'C1010': 'Legacy product', 'E1010': '75/75/75', 'H1010': 'Unknown use', 'I1010': 'Diagonal'}}
        page = calculate_page('ductwork', inputs, 'CALCULATOR', 1010, 1)
        cells = {cell['address']: cell for row in page['rows'] for cell in row['cells']}
        for column, choices in DUCTWORK_CHOICES.items():
            cell = cells[f'{column}1010']
            self.assertEqual(cell['options'], list(choices))
            self.assertEqual(cell['value'], inputs['CALCULATOR'][f'{column}1010'])
            self.assertFalse(cell['allow_other'])
            self.assertEqual(cell['error_style'], 'stop')

    def test_excel_template_has_same_canonical_lists_and_stop_validation(self):
        workbook = load_workbook(BytesIO(export_schedule_template('ductwork')))
        try:
            rules = {str(rule.sqref): rule for rule in workbook['CALCULATOR'].data_validations.dataValidation}
            for source_column, template_column in [('C', 'B'), ('E', 'D'), ('H', 'G'), ('I', 'H')]:
                rule = rules[f'{template_column}2:{template_column}1001']
                self.assertEqual(rule.formula1, '"' + ','.join(DUCTWORK_CHOICES[source_column]) + '"')
                self.assertEqual(rule.errorStyle, 'stop')
                self.assertTrue(rule.showErrorMessage)
                self.assertTrue(rule.allowBlank)
        finally:
            workbook.close()


if __name__ == '__main__':
    unittest.main()
