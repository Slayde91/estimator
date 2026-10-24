"""Mixed-product exports retain source identity and reviewed quantity status."""
from io import BytesIO
from copy import deepcopy
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from pypdf import PdfReader

from estimator.calculator_report import project_calculator_report, build_calculator_report, build_calculator_summary_report
from estimator.calculator_register import build_calculator_register
from estimator.firetex_assessment import Assessment, PRODUCTS
from estimator.calculator_defaults import default_calculator_inputs


FX6010_DENSITY_ESTIMATE = (25.915 + .3) / 18


def mixed_inputs(*, reviewed_defaults=True):
    values = default_calculator_inputs('steel_intumescent') if reviewed_defaults else {'SETTINGS': {}}
    values['SCHEDULE'] = {}
    for row, name in enumerate(PRODUCTS, 10):
        values['SCHEDULE'].update({f'B{row}': name, f'C{row}': 'I/H beam - 3 sides',
            f'D{row}': 550, f'E{row}': 'A/V', f'G{row}': 100, f'H{row}': 60,
            f'L{row}': 10, f'AA{row}': 'Level ' + str(row)})
    return values


class FiretexExportTests(unittest.TestCase):
    def setUp(self):
        snapshots = tuple(Assessment(name, str(index) * 64, str(index + 2) * 64)
                          for index, name in enumerate(PRODUCTS, 1))
        self.patchers = [patch('estimator.firetex_assessment.load_assessments', return_value=snapshots),
                         patch('estimator.firetex_assessment.lookup', return_value={
                             'dft': 1000, 'factor': 100, 'status': 'PUBLISHED',
                             'source': 'Synthetic certificate p2 Table 1', 'note': 'Invented 1000 µm test cell.'})]
        for patcher in self.patchers:
            patcher.start()
            self.addCleanup(patcher.stop)

    def test_projection_keeps_separate_orders_and_source_receipts(self):
        data = project_calculator_report('steel_intumescent', mixed_inputs())
        self.assertEqual([row['values']['B'] for row in data['rows']], list(PRODUCTS))
        self.assertEqual([item['product'] for item in data['assessment_evidences']], list(PRODUCTS))
        pools = {row['values']['A']: row['values'] for row in data['summaries'][0]['rows']}
        self.assertEqual(pools[PRODUCTS[0]]['G'], 1)
        self.assertEqual(pools[PRODUCTS[1]]['G'], 1)
        self.assertEqual(pools[PRODUCTS[1]]['H'], 18)
        self.assertAlmostEqual(pools[PRODUCTS[1]]['C'], 10 / .92 * FX6010_DENSITY_ESTIMATE)
        self.assertAlmostEqual(pools[PRODUCTS[1]]['D'], 10 / .92)
        self.assertIn('READY', pools[PRODUCTS[1]]['I'])
        self.assertEqual(data['incomplete_rows'], 0)
        self.assertEqual(data['title'], 'Steel (intumescent)')
        self.assertRegex(data['basis'], r'(?i)(nominal|estimated)')

    def test_pdf_and_xlsx_keep_ready_default_orders_and_product_receipts(self):
        values = mixed_inputs()
        for builder in (build_calculator_report, build_calculator_summary_report):
            payload = builder('steel_intumescent', values)
            text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(payload)).pages)
            for name in PRODUCTS:
                self.assertIn(name, text)
            self.assertIn('READY' if builder == build_calculator_summary_report else 'DFT', text)
            self.assertNotIn('INCOMPLETE', text)
            self.assertIn('dataset SHA-256', text)
            if builder == build_calculator_summary_report:
                self.assertNotIn('Steel protection assessment basis', text)
                self.assertIn('Whole packs', text)
                self.assertNotIn('Whole kits', text)
        workbook = load_workbook(BytesIO(build_calculator_register('steel_intumescent', values)))
        text = '\n'.join(str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value is not None)
        for name in PRODUCTS:
            self.assertIn(name, text)
        self.assertIn('READY', text)
        self.assertNotIn('INCOMPLETE', text)
        self.assertIn('Compatibility SHA-256', text)
        fx6010 = next(row for row in workbook['Summary'] if row[0].value == PRODUCTS[1])
        self.assertAlmostEqual(fx6010[2].value, 10 / .92 * FX6010_DENSITY_ESTIMATE)
        self.assertAlmostEqual(fx6010[3].value, 10 / .92)
        self.assertEqual(fx6010[6].value, 1)
        self.assertEqual(fx6010[7].value, 18)

    def test_legacy_part_c_is_retained_and_held_in_projection_pdf_and_xlsx(self):
        values = mixed_inputs()
        values['SETTINGS']['H74'] = .212345678901234
        before = deepcopy(values)
        data = project_calculator_report('steel_intumescent', values)
        pools = {row['values']['A']: row['values'] for row in data['summaries'][0]['rows']}
        self.assertEqual(pools[PRODUCTS[0]]['G'], 1)
        self.assertEqual(pools[PRODUCTS[1]]['G'], '')
        self.assertEqual(pools[PRODUCTS[1]]['H'], '')
        self.assertIn('INCOMPLETE', pools[PRODUCTS[1]]['I'])
        self.assertEqual(data['incomplete_rows'], 1)
        self.assertEqual(data['inputs']['SETTINGS']['H74'], .212345678901234)
        held_row = next(row for row in data['rows'] if row['values']['B'] == PRODUCTS[1])
        self.assertEqual(held_row['values']['O'], 1000)
        self.assertIn('legacy Part C', held_row['values']['W'])
        for builder in (build_calculator_report, build_calculator_summary_report):
            text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(builder('steel_intumescent', values))).pages)
            for name in PRODUCTS:
                self.assertIn(name, text)
            self.assertIn('INCOMPLETE' if builder == build_calculator_summary_report else 'DFT', text)
            self.assertIn('dataset SHA-256', text)
        workbook = load_workbook(BytesIO(build_calculator_register('steel_intumescent', values)))
        record = next(row for row in workbook['Summary'] if row[0].value == PRODUCTS[1])
        self.assertIsNone(record[6].value)
        self.assertIsNone(record[7].value)
        self.assertIn('INCOMPLETE', record[8].value)
        text = '\n'.join(str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value is not None)
        self.assertIn('INCOMPLETE', text)
        self.assertIn('Compatibility SHA-256', text)
        self.assertEqual(values, before)

    def test_sparse_saved_export_uses_effective_kit_estimates_without_rewriting_raw_settings(self):
        addresses = ('H66', 'H67', 'H68', 'H71', 'H72', 'H73')
        for settings in ({}, dict.fromkeys(addresses), dict.fromkeys(addresses, '')):
            with self.subTest(blank_kind=settings.get('H67', 'omitted')):
                values = mixed_inputs(reviewed_defaults=False)
                values['SETTINGS'] = settings
                before = deepcopy(values)
                data = project_calculator_report('steel_intumescent', values)
                pools = {row['values']['A']: row['values'] for row in data['summaries'][0]['rows']}
                normalized = {key: None if value == '' else value for key, value in settings.items()}
                self.assertEqual(data['inputs']['SETTINGS'], normalized)
                self.assertEqual(values, before)
                self.assertEqual(pools[PRODUCTS[0]]['G'], 1)
                self.assertEqual(pools[PRODUCTS[1]]['G'], 1)
                self.assertEqual(pools[PRODUCTS[1]]['H'], 18)
                self.assertAlmostEqual(pools[PRODUCTS[1]]['D'], 10 / .92)
                self.assertAlmostEqual(pools[PRODUCTS[1]]['C'], 10 / .92 * FX6010_DENSITY_ESTIMATE)
                self.assertIn('READY', pools[PRODUCTS[1]]['I'])
                self.assertEqual(data['incomplete_rows'], 0)

    def test_sparse_saved_pdf_and_register_include_estimates_without_storing_them(self):
        values = mixed_inputs(reviewed_defaults=False)
        before = deepcopy(values)
        for builder in (build_calculator_report, build_calculator_summary_report):
            text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(builder('steel_intumescent', values))).pages)
            self.assertNotIn('INCOMPLETE', text)
            if builder == build_calculator_summary_report:
                self.assertRegex(text, r'(?i)(nominal|estimated)')
            else:
                self.assertIn('QUANTIFIED', text)
        workbook = load_workbook(BytesIO(build_calculator_register('steel_intumescent', values)))
        record = next(row for row in workbook['Summary'] if row[0].value == PRODUCTS[1])
        self.assertAlmostEqual(record[2].value, 10 / .92 * FX6010_DENSITY_ESTIMATE)
        self.assertAlmostEqual(record[3].value, 10 / .92)
        self.assertEqual(record[6].value, 1)
        self.assertEqual(record[7].value, 18)
        self.assertEqual(values, before)
        self.assertEqual(values['SETTINGS'], {})

    def test_explicit_invalid_fx6010_values_stay_held_while_other_product_orders_remain_ready(self):
        for settings in ({'H67': 0}, {'H71': 'Unrecognised pack basis'}):
            with self.subTest(settings=settings):
                values = mixed_inputs(reviewed_defaults=False)
                values['SETTINGS'] = settings
                before = deepcopy(values)
                data = project_calculator_report('steel_intumescent', values)
                pools = {row['values']['A']: row['values'] for row in data['summaries'][0]['rows']}
                self.assertEqual(pools[PRODUCTS[0]]['G'], 1)
                self.assertEqual(pools[PRODUCTS[1]]['G'], '')
                self.assertEqual(pools[PRODUCTS[1]]['H'], '')
                self.assertIn('INCOMPLETE', pools[PRODUCTS[1]]['I'])
                self.assertEqual(data['inputs']['SETTINGS'], settings)
                self.assertEqual(values, before)

    def test_fx5090_explicit_blank_is_not_filled_by_fx6010_automatic_defaults(self):
        values = mixed_inputs(reviewed_defaults=False)
        values['SETTINGS']['F37'] = None
        before = deepcopy(values)
        data = project_calculator_report('steel_intumescent', values)
        pools = {row['values']['A']: row['values'] for row in data['summaries'][0]['rows']}
        self.assertEqual(pools[PRODUCTS[0]]['G'], '')
        self.assertIn('INCOMPLETE', pools[PRODUCTS[0]]['I'])
        self.assertEqual(pools[PRODUCTS[1]]['G'], 1)
        self.assertEqual(pools[PRODUCTS[1]]['H'], 18)
        self.assertEqual(data['inputs']['SETTINGS'], {'F37': None})
        self.assertEqual(values, before)

    def test_custom_pack_volume_keeps_coating_orders_but_holds_inconsistent_component_orders(self):
        values = mixed_inputs(reviewed_defaults=False)
        values['SETTINGS']['H67'] = 9
        before = deepcopy(values)
        data = project_calculator_report('steel_intumescent', values)
        record = next(row['values'] for row in data['summaries'][0]['rows'] if row['values']['A'] == PRODUCTS[1])
        self.assertAlmostEqual(record['C'], 10 / .92 * FX6010_DENSITY_ESTIMATE)
        self.assertAlmostEqual(record['D'], 10 / .92)
        self.assertEqual(record['G'], 2)
        self.assertEqual(record['H'], 18)
        self.assertEqual(record['I'], 'READY - component orders held')
        self.assertEqual(data['incomplete_rows'], 0)
        fx_row = next(row['values'] for row in data['rows'] if row['values']['B'] == PRODUCTS[1])
        self.assertEqual(fx_row['O'], 1000)
        self.assertIn('Component order withheld', fx_row['Y'])
        qualification = next(item for item in data['summaries'][0]['qualifications'] if item[0] == PRODUCTS[1])
        self.assertNotIn('Part A:', qualification[1])
        self.assertNotIn('Part B:', qualification[1])
        self.assertIn('Unavailable kg', qualification[2])
        for builder in (build_calculator_report, build_calculator_summary_report):
            text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(builder('steel_intumescent', values))).pages)
            self.assertIn('component orders held', ' '.join(text.split()))
        workbook = load_workbook(BytesIO(build_calculator_register('steel_intumescent', values)))
        exported = next(row for row in workbook['Summary'] if row[0].value == PRODUCTS[1])
        self.assertAlmostEqual(exported[2].value, record['C'])
        self.assertAlmostEqual(exported[3].value, record['D'])
        self.assertEqual(exported[6].value, 2)
        self.assertEqual(exported[7].value, 18)
        self.assertEqual(exported[8].value, 'READY - component orders held')
        self.assertEqual(values, before)
        self.assertEqual(data['inputs']['SETTINGS'], {'H67': 9})


if __name__ == '__main__':
    unittest.main()
