"""Product quantity arithmetic and shared calculator boundaries.

Table selection has separate tests against the reviewed authority. A fixed
1 mm result here provides hand-checkable quantity examples independent of DFT.
"""

from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace
from threading import RLock
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from estimator.catalog import ValidationError
from estimator.calculator_defaults import default_calculator_inputs
from estimator.intumescent import (AUTOMATIC_DEFAULTS, CASE_CHOICES, CASES, FX5090, FX6010,
                                  FX6010_KIT_DEFAULTS, PERIODS, PRODUCTS, TEMPERATURES,
                                  REPORT_SHA256, SC902_SETTING_DEFAULTS,
                                  SETTING_DEFAULTS, IntumescentEngine, PRODUCT, effective_inputs, model)
from estimator.workbook_calculators import calculator_definition, calculate_page, calculate_worksheet, normalize_calculator_inputs
from estimator.schedule_workbook import export_schedule_template, import_schedule_workbook
from estimator.schedule_rows import empty_schedule_inputs


ASSESSMENT = SimpleNamespace(digest='test', report_sha256='report', error='')
FIXED_RESULT = {'dft': 1000, 'factor': 100, 'status': 'PUBLISHED',
                'source': 'FC18580-01-01 p12 Table 6', 'note': '1 mm = 1000 µm DFT.'}


def inputs(*, area=10, row=10, **settings):
    return {'SCHEDULE': {f'B{row}': PRODUCT, f'C{row}': 'I/H beam - 3 sides',
                         f'D{row}': 550, f'E{row}': 'A/V', f'G{row}': 100,
                         f'H{row}': 60, f'L{row}': area},
            'SETTINGS': settings}


class IntumescentTests(unittest.TestCase):
    def setUp(self):
        self.lookup = patch('estimator.sc902_assessment.lookup', return_value=FIXED_RESULT)
        self.lookup.start()
        self.addCleanup(self.lookup.stop)

    def engine(self, values=None):
        return IntumescentEngine(model(), values or inputs(), assessment=ASSESSMENT)

    def test_published_consumption_drives_kg_litres_and_kits(self):
        engine = self.engine()
        # 10 m² × 1 mm × 1.75 kg/m²/mm = 17.5 kg.
        self.assertEqual(engine.value('SCHEDULE', 'O10'), 1000)
        self.assertAlmostEqual(engine.value('SCHEDULE', 'P10'), 1000 / .85)
        self.assertAlmostEqual(engine.value('SCHEDULE', 'S10'), 17.5 / 1.46)
        self.assertEqual(engine.value('SCHEDULE', 'T10'), .7)
        self.assertEqual(engine.value('SUMMARY', 'G7'), 1)
        self.assertEqual(engine.value('SUMMARY', 'D10'), 22.3)
        self.assertEqual(engine.value('SUMMARY', 'D11'), 2.7)
        self.assertAlmostEqual(engine.value('SETTINGS', 'D22'), 25 / 1.46)
        self.assertAlmostEqual(engine.value('SETTINGS', 'D23'), 25 / 1.75)
        self.assertAlmostEqual(engine.value('SETTINGS', 'D18'), 22.3 / 1.55)
        self.assertAlmostEqual(engine.value('SETTINGS', 'D19'), 2.7 / .99)

    def test_product_totals_pool_before_waste_and_rounding(self):
        values = inputs(area=1, D16=.1)
        values['SCHEDULE'].update(inputs(area=1, row=11)['SCHEDULE'])
        values['SCHEDULE'].update(inputs(area=1, row=1009)['SCHEDULE'])
        engine = self.engine(values)
        self.assertEqual(sum(engine.value('SCHEDULE', f'U{row}') for row in (10, 11, 1009)), 3)
        self.assertAlmostEqual(engine.value('SUMMARY', 'E7'), .21)
        self.assertEqual(engine.value('SUMMARY', 'G7'), 1)
        self.assertAlmostEqual(engine.totals()['kg'], 5.25)

    def test_editable_settings_change_quantities_without_changing_dft(self):
        engine = self.engine(inputs(D9=2, D10=15, D11=5, D15=2.4, D14=.8, D16=.25))
        self.assertEqual(engine.value('SCHEDULE', 'O10'), 1000)
        self.assertEqual(engine.value('SCHEDULE', 'P10'), 1250)
        self.assertEqual(engine.value('SCHEDULE', 'S10'), 12)
        self.assertEqual(engine.value('SCHEDULE', 'T10'), 1.2)
        self.assertEqual(engine.value('SUMMARY', 'G7'), 2)
        self.assertEqual(engine.value('SUMMARY', 'D10'), 30)
        self.assertEqual(engine.value('SUMMARY', 'D11'), 10)

    def test_incomplete_or_blocked_item_withholds_complete_order(self):
        values = inputs()
        values['SCHEDULE']['A11'] = 'incomplete member'
        engine = self.engine(values)
        self.assertEqual(engine.value('SUMMARY', 'G7'), '')
        self.assertEqual(engine.value('SUMMARY', 'D10'), '')
        self.assertEqual(engine.value('SUMMARY', 'D15'), 1)
        self.assertAlmostEqual(engine.value('SUMMARY', 'C7'), 17.5)
        with patch('estimator.sc902_assessment.lookup', return_value={**FIXED_RESULT, 'dft': None, 'status': 'UNAVAILABLE'}):
            blocked = self.engine()
            for column in ('O', 'P', 'S', 'T', 'U'):
                self.assertEqual(blocked.value('SCHEDULE', column + '10'), '')
            self.assertEqual(blocked.value('SUMMARY', 'G7'), '')

    def test_invalid_quantity_settings_withhold_without_zero_substitution(self):
        for address, value in [('D9', 0), ('D10', None), ('D11', -1), ('D14', 1.1), ('D15', 0), ('D16', -.1)]:
            with self.subTest(address=address):
                engine = self.engine(inputs(**{address: value}))
                self.assertEqual(engine.value('SCHEDULE', 'O10'), 1000)
                self.assertEqual(engine.value('SCHEDULE', 'S10'), '')
                self.assertEqual(engine.value('SUMMARY', 'G7'), '')

    def test_geometry_area_and_esa_conversion_are_separate_from_dft(self):
        values = inputs(area=None, D17=8000)
        values['SCHEDULE'].update(E10='ESA/M', G10=12.5, I10=2, J10=5, K10=.8)
        engine = self.engine(values)
        self.assertEqual(engine.value('SCHEDULE', 'M10'), 100)
        self.assertEqual(engine.value('SCHEDULE', 'R10'), 8)
        values['SCHEDULE']['L10'] = 3
        self.assertEqual(self.engine(values).value('SCHEDULE', 'R10'), 3)

    def test_invalid_present_overrides_do_not_silently_fall_back(self):
        for field, value in [('K10', -1), ('K10', 0), ('L10', -1), ('L10', 0)]:
            values = inputs()
            values['SCHEDULE'].update(I10=2, J10=5, K10=.8)
            values['SCHEDULE'][field] = value
            with self.subTest(field=field, value=value):
                engine = self.engine(values)
                self.assertEqual(engine.value('SCHEDULE', 'O10'), 1000)
                self.assertEqual(engine.value('SCHEDULE', 'S10'), '')
                self.assertEqual(engine.value('SUMMARY', 'G7'), '')

    def test_named_geometry_holds_and_family_mismatch_withhold(self):
        engine = self.engine()
        self.assertIsNone(engine._geometry('100X100X4SHS', 'R3')[0])
        self.assertIsNone(engine._geometry('310UB40.4', 'CHS')[0])
        self.assertIsNone(engine._geometry('unknown section', 'R4')[0])
        from estimator.excel_engine import WorkbookEngine
        original = WorkbookEngine.value
        for field, value in [('BA', 1), ('AM', 'EXCLUDED - invalid section')]:
            def held(instance, sheet, address):
                return value if sheet == 'SECTIONS' and address.startswith(field) else original(instance, sheet, address)
            with self.subTest(field=field), patch.object(WorkbookEngine, 'value', held):
                self.assertIsNone(engine._geometry('310UB40.4', 'R3')[0])

    def test_extreme_finite_inputs_withhold_instead_of_returning_infinity(self):
        engine = self.engine(inputs(D9=5e-324, D14=5e-324))
        self.assertEqual(engine.value('SCHEDULE', 'O10'), 1000)
        self.assertEqual(engine.value('SCHEDULE', 'P10'), '')
        self.assertEqual(engine.value('SCHEDULE', 'S10'), '')
        self.assertEqual(engine.value('SUMMARY', 'G7'), '')

    def test_unrepresentable_positive_quantities_do_not_become_zero_orders(self):
        for values in (inputs(area=5e-324, D15=5e-324), inputs(area=5e-324, D9=1e100)):
            with self.subTest(values=values):
                engine = self.engine(values)
                self.assertEqual(engine.value('SCHEDULE', 'O10'), 1000)
                self.assertEqual(engine.value('SCHEDULE', 'S10'), '')
                self.assertEqual(engine.value('SUMMARY', 'G7'), '')

    def test_three_visible_pages_retain_internal_start_choices_and_read_only_outputs(self):
        definition = calculator_definition('steel_intumescent')
        self.assertEqual(definition['pages'], ['SCHEDULE', 'SUMMARY', 'SETTINGS'])
        self.assertIn('START', [sheet['name'] for sheet in model()['sheets']])
        self.assertTrue(all(value.startswith('START!') for value in model()['defined_names'].values()))
        self.assertEqual(definition['title'], 'Steel (intumescent)')
        values = inputs()
        values['SCHEDULE']['C10'] = 'I/H beam/column - 4 sides'
        with patch('estimator.sc902_assessment.lookup', return_value=FIXED_RESULT) as selected:
            self.assertEqual(self.engine(values).value('SCHEDULE', 'O10'), 1000)
            self.assertEqual(selected.call_args.kwargs['case'], 'R4')
        with self.assertRaises(ValidationError):
            normalize_calculator_inputs('steel_intumescent', {'SCHEDULE': {'O10': 1}})
        for invalid in (float('inf'), True, '=1+1'):
            with self.subTest(invalid=invalid), self.assertRaises(ValidationError):
                normalize_calculator_inputs('steel_intumescent', {'SETTINGS': {'D9': invalid}})

    def test_blank_schedule_is_blank_and_mutation_does_not_change_model(self):
        original = deepcopy(model())
        engine = self.engine({'SCHEDULE': {}})
        self.assertEqual(engine.value('SCHEDULE', 'O10'), '')
        self.assertEqual(engine.value('SUMMARY', 'G7'), 0)
        self.assertEqual(model(), original)

    def test_template_round_trip_retains_settings_and_location(self):
        payload = export_schedule_template('steel_intumescent')
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook['SCHEDULE']
        headings = [cell.value for cell in sheet[1]]
        sheet.cell(2, headings.index('Location') + 1, 'Level 1')
        sheet.cell(2, headings.index('Product') + 1, PRODUCT)
        stream = BytesIO()
        workbook.save(stream)
        result = import_schedule_workbook('steel_intumescent', stream.getvalue(), 'schedule.xlsx', inputs(D16=.15))
        self.assertEqual(result['inputs']['SCHEDULE']['AA10'], 'Level 1')
        self.assertEqual(result['inputs']['SETTINGS']['D16'], .15)


class MultiProductIntumescentTests(unittest.TestCase):
    """Invented 1 mm DFT isolates quantity policies from private report data."""

    def setUp(self):
        self.firetex = (SimpleNamespace(product=FX5090, digest='invented-5090'),
                        SimpleNamespace(product=FX6010, digest='invented-6010'))
        sc902 = patch('estimator.sc902_assessment.lookup', return_value=FIXED_RESULT)
        firetex = patch('estimator.firetex_assessment.lookup', return_value={
            **FIXED_RESULT, 'source': 'Synthetic certificate p2 Table 1'})
        self.sc_lookup, self.fx_lookup = sc902.start(), firetex.start()
        self.addCleanup(sc902.stop)
        self.addCleanup(firetex.stop)

    def engine(self, values=None, *, installed=True):
        return IntumescentEngine(model(), inputs() if values is None else values,
                                 assessment=ASSESSMENT,
                                 firetex_assessments=self.firetex if installed else None)

    def values(self, product_name=FX5090, *, row=10, area=10, **settings):
        values = inputs(area=area, row=row, **settings)
        values['SCHEDULE'][f'B{row}'] = product_name
        return values

    def new_values(self, product_name=FX6010, *, row=10, area=10, **settings):
        values = default_calculator_inputs('steel_intumescent')
        row_inputs = self.values(product_name, row=row, area=area, **settings)
        values['SCHEDULE'] = row_inputs['SCHEDULE']
        values['SETTINGS'].update(row_inputs['SETTINGS'])
        return values

    def test_legacy_settings_and_source_identity_are_retained(self):
        source = model()
        cells = next(sheet for sheet in source['sheets'] if sheet['name'] == 'SETTINGS')['cells']
        for address, value in SC902_SETTING_DEFAULTS.items():
            self.assertEqual(cells[address]['value'], value)
        self.assertEqual(source['source']['sha256'], REPORT_SHA256)
        self.assertEqual(source['schedule']['last_column'], 'AA')
        self.assertEqual(source['schedule']['line_id_column'], 'Z')
        self.assertEqual(source['schedule']['location_column'], 'AA')
        for address in ('H66', 'H67', 'H68', 'H72', 'H73', 'H74'):
            self.assertIsNone(cells[address]['value'])
        self.assertEqual(PERIODS, (15,20,30,45,60,75,90,105,120,150,180))
        start = next(sheet for sheet in source['sheets'] if sheet['name'] == 'START')['cells']
        self.assertIn('exact published temperature', start['A14']['value'])
        self.assertNotIn('lower temperature', start['A14']['value'])

    def test_each_product_routes_to_its_own_assessment(self):
        engine = self.engine()
        self.assertIs(engine.product_assessment(PRODUCT), ASSESSMENT)
        self.assertIs(engine.product_assessment(FX5090), self.firetex[0])
        self.assertIs(engine.product_assessment(FX6010), self.firetex[1])
        self.assertIsNone(engine.product_assessment('Unknown'))
        for product_name, selected in zip((FX5090, FX6010), self.firetex):
            with self.subTest(product=product_name):
                engine = self.engine(self.values(product_name))
                self.assertEqual(engine.value('SCHEDULE', 'O10'), 1000)
                self.assertIs(self.fx_lookup.call_args.args[0], selected)

    def test_fx5090_uses_volume_solids_and_keeps_sc902_totals_separate(self):
        values = self.values()
        values['SCHEDULE'].update(inputs(row=11)['SCHEDULE'])
        engine = self.engine(values)
        self.assertAlmostEqual(engine.value('SCHEDULE', 'S10'), 10 / .69)
        self.assertAlmostEqual(engine.value('SCHEDULE', 'T10'), 10 / .69 / 18)
        self.assertAlmostEqual(engine.value('SCHEDULE', 'P10'), 1000 / .69)
        fx, sc = engine.product_totals(FX5090), engine.totals()
        self.assertEqual(tuple(engine.product_totals()), PRODUCTS)
        self.assertEqual(fx['whole'], 1)
        self.assertEqual(fx['order_litres'], 18)
        self.assertAlmostEqual(fx['kg'], 10 / .69 * 1.39)
        self.assertEqual(fx['part_a'], '')
        self.assertAlmostEqual(sc['net'], .7)
        self.assertEqual(sc['kg'], 17.5)
        self.assertAlmostEqual(engine.value('SCHEDULE', 'G5'), 10/.69 + 17.5/1.46)
        self.assertEqual(engine.value('SUMMARY', 'A7'), PRODUCT)
        self.assertEqual(engine.value('SUMMARY', 'A22'), FX5090)

    def test_sparse_fx6010_uses_au_estimates_without_changing_other_product(self):
        values = self.values(FX6010)
        values['SCHEDULE'].update(self.values(FX5090, row=11)['SCHEDULE'])
        before = deepcopy(values)
        engine = self.engine(values)
        self.assertEqual(engine.value('SCHEDULE', 'O10'), 1000)
        self.assertAlmostEqual(engine.value('SCHEDULE', 'P10'), 1000/.92)
        self.assertAlmostEqual(engine.value('SCHEDULE', 'S10'), 10/.92)
        self.assertIn('QUANTIFIED', engine.value('SCHEDULE', 'W10'))
        self.assertEqual(engine.product_totals(FX6010)['pending'], 0)
        self.assertEqual(engine.product_totals(FX6010)['whole'], 1)
        self.assertEqual(engine.product_totals(FX5090)['whole'], 1)
        self.assertEqual(engine.value('SCHEDULE', 'S5'), 0)
        self.assertEqual(values, before)

    def test_new_and_reset_fx6010_defaults_use_explicit_nominal_au_kit_estimates(self):
        frozen_model = deepcopy(model())
        for factory in (default_calculator_inputs, empty_schedule_inputs):
            with self.subTest(factory=factory.__name__):
                values = factory('steel_intumescent')
                self.assertEqual(values['SETTINGS']['H67'], 18)
                self.assertEqual(values['SETTINGS']['H68'], .92)
                for address in ('H66', 'H72', 'H73'):
                    self.assertEqual(values['SETTINGS'][address], FX6010_KIT_DEFAULTS[address], address)
                self.assertIsNone(values['SETTINGS']['H74'])
                values.setdefault('SCHEDULE', {}).update(self.values(FX6010)['SCHEDULE'])
                before = deepcopy(values)
                engine = self.engine(values)
                self.assertEqual(engine.value('SCHEDULE', 'O10'), 1000)
                self.assertAlmostEqual(engine.value('SCHEDULE', 'P10'), 1000 / .92)
                self.assertAlmostEqual(engine.value('SCHEDULE', 'S10'), 10 / .92)
                self.assertAlmostEqual(engine.value('SCHEDULE', 'T10'), 10 / .92 / 18)
                total = engine.product_totals(FX6010)
                self.assertEqual(total['pending'], 0)
                self.assertEqual(total['whole'], 1)
                self.assertEqual(total['order_litres'], 18)
                self.assertEqual(total['dry_litres'], 10)
                self.assertAlmostEqual(total['kg'], 10/.92*FX6010_KIT_DEFAULTS['H66'])
                self.assertAlmostEqual(total['pack_kg'], 26.215)
                self.assertEqual(total['part_a'], FX6010_KIT_DEFAULTS['H72'])
                self.assertEqual(total['part_b'], .3)
                self.assertEqual(total['part_c'], '')
                self.assertEqual(values, before)
        self.assertEqual(model(), frozen_model)

    def test_saved_blank_fx6010_settings_use_effective_defaults_without_rewriting_raw_inputs(self):
        for settings in ({}, {'H67': None, 'H68': None},
                         {'H67': None, 'H68': .92}, {'H67': 18, 'H68': None},
                         {address: '' for address in AUTOMATIC_DEFAULTS['SETTINGS']}):
            with self.subTest(settings=settings):
                values = self.values(FX6010, **settings)
                before = deepcopy(values)
                normalized = normalize_calculator_inputs('steel_intumescent', values)
                # Existing normalization canonicalizes literal blanks to None.
                expected_raw = {address:None if value == '' else value
                                for address,value in settings.items()}
                self.assertEqual(normalized['SETTINGS'], expected_raw)
                engine = self.engine(normalized)
                self.assertEqual(engine.value('SCHEDULE', 'O10'), 1000)
                self.assertAlmostEqual(engine.value('SCHEDULE', 'S10'), 10/.92)
                self.assertEqual(engine.product_totals(FX6010)['whole'], 1)
                self.assertEqual(normalized['SETTINGS'], expected_raw)
                self.assertEqual(values, before)

    def test_automatic_defaults_are_scoped_copied_and_preserve_custom_and_legacy_inputs(self):
        values = self.values(FX6010, D16=.15, F37=20, H69=None, H70='', H74=.25,
                             H66=1.23456789012345, H67=20, H68=.812345678901234,
                             H71='Component mass (kg)', H72=8, H73=2)
        before = deepcopy(values)
        effective = effective_inputs(values)
        self.assertEqual(effective, before)
        self.assertIsNot(effective, values)
        self.assertIsNot(effective['SETTINGS'], values['SETTINGS'])
        self.assertIsNot(effective['SCHEDULE'], values['SCHEDULE'])
        effective['SCHEDULE']['A10'] = 'isolated copy'
        effective['SETTINGS']['H67'] = 1
        self.assertEqual(values, before)
        self.assertEqual(set(AUTOMATIC_DEFAULTS), {'SETTINGS'})
        self.assertEqual(set(AUTOMATIC_DEFAULTS['SETTINGS']), {'H66','H67','H68','H71','H72','H73'})
        self.assertEqual(model()['automatic_defaults'], AUTOMATIC_DEFAULTS)
        self.assertAlmostEqual(FX6010_KIT_DEFAULTS['H66'], (25.915+.3)/18)
        self.assertAlmostEqual(FX6010_KIT_DEFAULTS['H72'], 25.915)
        self.assertEqual(FX6010_KIT_DEFAULTS['H73'], .3)
        self.assertTrue(all('Nominal' in next(sheet for sheet in model()['sheets']
                            if sheet['name']=='SETTINGS')['cells'][f'A{address[1:]}']['value']
                            for address in ('H66','H72','H73')))

    def test_explicit_zero_invalid_and_custom_au_settings_are_never_defaulted(self):
        for address in AUTOMATIC_DEFAULTS['SETTINGS']:
            for value in (0, False, -1, 'invalid', ' ', float('inf')):
                with self.subTest(address=address, value=value):
                    raw = {'SETTINGS': {address: value, 'H74': .25}}
                    effective = effective_inputs(raw)
                    self.assertIs(type(effective['SETTINGS'][address]), type(value))
                    self.assertEqual(effective['SETTINGS'][address], value)
                    self.assertEqual(effective['SETTINGS']['H74'], .25)
        for address in ('H66','H67','H68','H72','H73'):
            for value in (0, -1):
                with self.subTest(address=address, value=value):
                    engine = self.engine(self.values(FX6010, **{address:value}))
                    self.assertEqual(engine.value('SCHEDULE','O10'), 1000)
                    self.assertEqual(engine.value('SCHEDULE','S10'), '')
                    self.assertEqual(engine.product_totals(FX6010)['whole'], '')

    def test_effective_defaults_do_not_replace_explicit_blank_other_products_or_allowances(self):
        raw = {'SETTINGS': {'D9':None,'F37':'','H69':None,'H70':'','H74':None}}
        effective = effective_inputs(raw)
        for address, value in raw['SETTINGS'].items():
            self.assertEqual(effective['SETTINGS'][address], value)
        self.assertNotIn('H74', effective_inputs()['SETTINGS'])
        self.assertEqual(SETTING_DEFAULTS['H67'], None)

    def test_legacy_part_c_rows_hide_only_for_blank_or_numeric_zero_without_losing_input(self):
        for part_c in (None, 0, 0.0, .25, -1):
            with self.subTest(part_c=part_c):
                values = self.new_values(H74=part_c)
                before = deepcopy(values)
                engine = self.engine(values)
                with patch('estimator.workbook_calculators.calculator_session',
                           return_value=(values, engine, RLock())):
                    settings = calculate_page('steel_intumescent', values, 'SETTINGS', 74, 1)
                    summary = calculate_page('steel_intumescent', values, 'SUMMARY', 47, 1)
                hidden = part_c is None or part_c == 0
                self.assertEqual(74 in settings['omitted_rows'], hidden)
                self.assertEqual(47 in summary['omitted_rows'], hidden)
                self.assertEqual(settings['inputs']['SETTINGS']['H74'], part_c)
                self.assertEqual(values, before)

    def test_fx6010_custom_pack_volume_keeps_litres_but_withholds_conflicting_component_orders(self):
        engine = self.engine(self.values(FX6010, H67=8, H68=.8))
        total = engine.product_totals(FX6010)
        self.assertEqual(engine.value('SCHEDULE', 'P10'), 1250)
        self.assertEqual(engine.value('SCHEDULE', 'S10'), 12.5)
        self.assertEqual(engine.value('SCHEDULE', 'T10'), 1.5625)
        self.assertEqual(total['whole'], 2)
        self.assertEqual(total['order_litres'], 16)
        self.assertAlmostEqual(total['kg'], 12.5*FX6010_KIT_DEFAULTS['H66'])
        for field in ('pack_kg','part_a','part_b'):
            self.assertEqual(total[field], '', field)
        self.assertEqual(total['part_c'], '')
        self.assertIn('component orders held',total['status'])
        self.assertIn('Component order withheld',engine.value('SCHEDULE','Y10'))
        self.assertAlmostEqual(engine.value('SETTINGS','H80'),26.215)
        self.assertEqual(total['dry_litres'], 10)

    def test_fx6010_half_volume_custom_pack_cannot_order_full_default_components(self):
        values = self.values(FX6010,H67=9)
        before = deepcopy(values)
        engine = self.engine(values)
        total = engine.product_totals(FX6010)
        self.assertEqual(engine.value('SCHEDULE','O10'),1000)
        self.assertAlmostEqual(engine.value('SCHEDULE','P10'),1000/.92)
        self.assertAlmostEqual(total['litres'],10/.92)
        self.assertAlmostEqual(total['net'],10/.92/9)
        self.assertAlmostEqual(total['kg'],10/.92*FX6010_KIT_DEFAULTS['H66'])
        self.assertEqual(total['whole'],2)
        self.assertEqual(total['order_litres'],18)
        self.assertEqual(total['pending'],0)
        for field in ('part_a','part_b','pack_kg'):
            self.assertEqual(total[field],'',field)
        self.assertEqual(engine.value('SCHEDULE','W10'),'QUANTIFIED - component orders held')
        self.assertIn('pack volume, density and Part A/B masses',engine.value('SCHEDULE','Y10'))
        self.assertEqual(total['status'],'READY - component orders held')
        self.assertEqual(values,before)

    def test_fx6010_component_consistency_uses_selected_basis_and_nominal_relative_rounding(self):
        cases = (
            # Matching custom volume/density, and independently rounded nominal
            # published estimates, can retain complete component orders.
            ({'H67':9,'H66':26.215/9},True),
            ({'H67':18,'H66':1.46,'H72':25.9,'H73':.3},True),
            ({'H67':10,'H66':1,'H72':9,'H73':.951},True),
            ({'H67':10,'H66':1,'H72':9,'H73':.949},False),
            ({'H67':18,'H66':2},False),
            ({'H72':8,'H73':2},False),
            # A fixed absolute tolerance must not approve tiny packs whose
            # component total is twice the mass implied by volume/density.
            ({'H67':1e-6,'H66':1,'H72':1e-6,'H73':1e-6},False),
            # The mass basis derives its volume; unused saved H67 is untouched.
            ({'H71':'Component mass (kg)','H67':9,'H66':1.25,'H72':8,'H73':2},True),
        )
        for settings,consistent in cases:
            with self.subTest(settings=settings):
                values=self.values(FX6010,**settings)
                before=deepcopy(values)
                engine=self.engine(values)
                total=engine.product_totals(FX6010)
                self.assertEqual(engine.value('SCHEDULE','O10'),1000)
                self.assertGreater(total['litres'],0)
                self.assertGreater(total['kg'],0)
                self.assertGreater(total['whole'],0)
                self.assertEqual(bool(total['component_order_note']),not consistent)
                if consistent:
                    self.assertGreater(total['part_a'],0)
                    self.assertGreater(total['part_b'],0)
                    self.assertGreater(total['pack_kg'],0)
                else:
                    self.assertEqual((total['part_a'],total['part_b'],total['pack_kg']),('','',''))
                self.assertEqual(values,before)

    def test_fx6010_volume_basis_accepts_zero_c_with_nominal_component_estimates(self):
        baseline = self.engine(self.new_values()).product_totals(FX6010)
        values = self.new_values(H74=0)
        self.assertEqual(values['SETTINGS']['H72'], FX6010_KIT_DEFAULTS['H72'])
        self.assertEqual(values['SETTINGS']['H73'], FX6010_KIT_DEFAULTS['H73'])
        total = self.engine(values).product_totals(FX6010)
        self.assertEqual(total, baseline)
        self.assertEqual(total['whole'], 1)
        self.assertEqual(total['part_c'], '')

    def test_fx6010_two_component_mass_basis_accepts_blank_or_numeric_zero_c(self):
        values = self.values(FX6010, H66=1.25, H68=.8, H71='Component mass (kg)',
                             H72=8, H73=2)
        baseline = None
        for part_c in (None, '', 0, 0.0, -0.0):
            with self.subTest(part_c=part_c):
                values['SETTINGS']['H74'] = part_c
                before = deepcopy(values)
                engine = self.engine(values)
                total = engine.product_totals(FX6010)
                self.assertEqual(total['kit_litres'], 8)
                self.assertEqual(total['pack_kg'], 10)
                self.assertEqual(total['litres'], 12.5)
                self.assertEqual(total['kg'], 15.625)
                self.assertEqual(total['whole'], 2)
                self.assertEqual((total['part_a'], total['part_b'], total['part_c']), (16, 4, ''))
                self.assertEqual(engine.value('SUMMARY', 'D47'), '')
                self.assertEqual(engine.value('SETTINGS', 'H80'), 10)
                self.assertEqual(engine.value('SETTINGS', 'H81'), 8)
                self.assertEqual(values, before)
                if baseline is None:
                    baseline = total
                else:
                    self.assertEqual(total, baseline)
        for address in ('H66', 'H72', 'H73'):
            with self.subTest(missing=address):
                missing = deepcopy(values)
                missing['SETTINGS'][address] = 0
                engine = self.engine(missing)
                self.assertEqual(engine.value('SCHEDULE', 'O10'), 1000)
                self.assertEqual(engine.value('SCHEDULE', 'S10'), '')
                self.assertEqual(engine.product_totals(FX6010)['whole'], '')

    def test_legacy_or_invalid_part_c_holds_both_quantity_bases_without_erasing_input(self):
        for basis in ('Pack volume (L)', 'Component mass (kg)'):
            for part_c in (.212345678901234, -1, True, False, '0', 'legacy', float('inf'), float('nan')):
                with self.subTest(basis=basis, part_c=part_c):
                    values = self.new_values(H66=1.25, H71=basis, H72=8, H73=2, H74=part_c)
                    before = deepcopy(values)
                    engine = self.engine(values)
                    self.assertEqual(engine.value('SCHEDULE', 'O10'), 1000)
                    self.assertAlmostEqual(engine.value('SCHEDULE', 'P10'), 1000 / .92)
                    for column in ('S', 'T', 'U'):
                        self.assertEqual(engine.value('SCHEDULE', column + '10'), '', column)
                    self.assertIn('legacy Part C', engine.value('SCHEDULE', 'W10'))
                    total = engine.product_totals(FX6010)
                    self.assertEqual(total['pending'], 1)
                    for field in ('whole', 'order_litres', 'part_a', 'part_b', 'part_c'):
                        self.assertEqual(total[field], '', field)
                    self.assertEqual(values, before)

    def test_part_c_zero_normalizes_exactly_but_nonnumeric_and_nonfinite_values_reject(self):
        for value in (0, 0.0, -0.0):
            with self.subTest(value=value):
                result = normalize_calculator_inputs('steel_intumescent', {'SETTINGS': {'H74': value}})
                self.assertEqual(result['SETTINGS']['H74'], value)
                self.assertIs(type(result['SETTINGS']['H74']), type(value))
        for value in (True, False, '0', float('inf'), float('nan')):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                normalize_calculator_inputs('steel_intumescent', {'SETTINGS': {'H74': value}})

    def test_fx6010_volume_orders_pool_then_round_at_pack_boundary_and_apply_waste_once(self):
        values = self.new_values(area=8.28)
        values['SCHEDULE'].update(self.values(FX6010, row=1009, area=8.28)['SCHEDULE'])
        values['SCHEDULE'].update(self.values(FX5090, row=11)['SCHEDULE'])
        values['SCHEDULE'].update(self.values(PRODUCT, row=12)['SCHEDULE'])
        engine = self.engine(values)
        self.assertEqual(sum(engine.value('SCHEDULE', f'U{row}') for row in (10, 1009)), 2)
        total = engine.product_totals(FX6010)
        self.assertEqual(total['used'], 2)
        self.assertAlmostEqual(total['litres'], 18)
        self.assertEqual(total['whole'], 1)
        self.assertEqual(total['order_litres'], 18)
        self.assertEqual(engine.product_totals(FX5090)['whole'], 1)
        self.assertEqual(engine.product_totals(PRODUCT)['whole'], 1)
        values['SETTINGS']['H69'] = .1
        with_waste = self.engine(values)
        self.assertEqual(with_waste.product_totals(FX6010)['whole'], 2)
        self.assertEqual(with_waste.product_totals(FX6010)['order_litres'], 36)
        self.assertEqual(with_waste.product_totals(FX5090)['whole'], 1)
        self.assertEqual(with_waste.product_totals(PRODUCT)['whole'], 1)
        just_over = self.engine(self.new_values(area=16.560001))
        self.assertEqual(just_over.product_totals(FX6010)['whole'], 2)

    def test_product_waste_is_pooled_once_including_last_schedule_row(self):
        values = self.values(area=1, F39=.1)
        values['SCHEDULE'].update(self.values(area=1, row=1009)['SCHEDULE'])
        engine = self.engine(values)
        self.assertEqual(engine.value('SCHEDULE', 'U10'), 1)
        self.assertEqual(engine.value('SCHEDULE', 'U1009'), 1)
        self.assertAlmostEqual(engine.product_totals(FX5090)['net'], 2/.69/18)
        self.assertEqual(engine.product_totals(FX5090)['whole'], 1)
        self.assertEqual(engine.product_totals(PRODUCT)['whole'], 0)

    def test_selecting_product_and_default_does_not_mutate_saved_settings(self):
        values = self.values(D6=FX5090, D9=2, D15=2.4, F37=20)
        del values['SCHEDULE']['B10']
        before = deepcopy(values)
        engine = self.engine(values)
        self.assertAlmostEqual(engine.value('SCHEDULE', 'T10'), 10/.69/20)
        self.assertEqual(engine.value('SETTINGS', 'D9'), 2)
        self.assertEqual(engine.value('SETTINGS', 'D15'), 2.4)
        self.assertEqual(values, before)
        self.assertEqual(model()['source']['sha256'], REPORT_SHA256)

    def test_explicit_cases_dispatch_without_changing_legacy_sc902_case(self):
        expected = {'I/H beam - 3 sides':'R3', 'I/H column - 4 sides':'R4',
                    'I/H column - 3 sides':'R4',
                    'I/H beam - 4 sides':'R4B', 'RHS/SHS column - 4 sides':'H4',
                    'RHS/SHS beam - 3 sides':'H3', 'RHS/SHS beam - 4 sides':'H4B',
                    'CHS column - 4 sides':'CHS'}
        for label, case in expected.items():
            with self.subTest(label=label):
                values = self.values()
                values['SCHEDULE']['C10'] = label
                self.assertEqual(self.engine(values).value('SCHEDULE', 'O10'), 1000)
                self.assertEqual(self.fx_lookup.call_args.kwargs['case'], case)
        values['SCHEDULE']['C10'] = 'I/H beam/column - 4 sides'
        blocked = self.engine(values)
        self.assertEqual(blocked.value('SCHEDULE', 'O10'), '')
        self.assertEqual(blocked.value('SCHEDULE', 'X10'), '')
        self.assertIn('separate', blocked.value('SCHEDULE', 'Y10'))
        for label in ('I/H beam/column - 4 sides','I/H column - 4 sides','I/H beam - 4 sides','I/H column - 3 sides'):
            values['SCHEDULE'].update(B10=PRODUCT, C10=label)
            self.assertEqual(self.engine(values).value('SCHEDULE', 'O10'), 1000)
            self.assertEqual(self.sc_lookup.call_args.kwargs['case'], 'R4')

    def test_new_case_choices_omit_only_combined_alias_and_temperature_control_is_closed(self):
        self.assertNotIn('I/H beam/column - 4 sides', CASE_CHOICES)
        self.assertIn('I/H beam/column - 4 sides', CASES)
        self.assertIn('I/H column - 4 sides', CASE_CHOICES)
        self.assertIn('I/H beam - 4 sides', CASE_CHOICES)
        self.assertEqual(set(CASE_CHOICES), set(CASES)-{'I/H beam/column - 4 sides'})
        source = model()
        columns = {field['column']:field for field in source['schedule']['columns']}
        self.assertEqual(columns['D']['type'], 'select')
        self.assertEqual(columns['D']['validation']['errorStyle'], 'stop')
        engine = self.engine()
        reference = engine.evaluate(('name','CChoices'),'SCHEDULE',10,3)
        self.assertEqual(tuple(reference.values()), CASE_CHOICES)

    def test_temperature_choices_match_each_product_and_native_case_without_changing_inputs(self):
        sc_open = (350,400,450,500,550,580,600,620,650,700,750)
        fx5090 = (350,400,450,500,550,600,650,700,750)
        fx6010_r4 = (350,400,450,500,550,580,600,650,700,750)
        fx6010_h3 = (350,400,450,500,550,600,620,650,700,750)
        fx6010_h4 = (350,400,450,500,520,550,600,650,700,750)
        expected = {
            PRODUCT: {'I/H beam - 3 sides':sc_open, 'I/H beam/column - 4 sides':sc_open,
                      'I/H column - 3 sides':sc_open,
                      'I/H column - 4 sides':sc_open, 'I/H beam - 4 sides':sc_open,
                      'RHS/SHS column - 4 sides':TEMPERATURES, 'CHS column - 4 sides':TEMPERATURES,
                      'RHS/SHS beam - 3 sides':(), 'RHS/SHS beam - 4 sides':()},
            FX5090: {label:fx5090 for label in CASES},
            FX6010: {'I/H beam - 3 sides':sc_open, 'I/H beam/column - 4 sides':fx6010_r4,
                     'I/H column - 3 sides':fx6010_r4,
                     'I/H column - 4 sides':fx6010_r4, 'I/H beam - 4 sides':fx6010_r4,
                     'RHS/SHS beam - 3 sides':fx6010_h3,
                     'RHS/SHS column - 4 sides':fx6010_h4,
                     'RHS/SHS beam - 4 sides':fx6010_h4, 'CHS column - 4 sides':fx6010_h4}}
        for product_name, cases in expected.items():
            for label, choices in cases.items():
                with self.subTest(product=product_name, case=label):
                    values = self.values(product_name)
                    values['SCHEDULE'].update(C10=label,D10=612.1234567890123)
                    before = deepcopy(values)
                    engine = self.engine(values)
                    self.assertEqual(engine.temperature_options(10), choices)
                    self.assertEqual(values, before)
                    self.assertEqual(engine.inputs['SCHEDULE']['D10'], 612.1234567890123)
            values = self.values(product_name, D6=product_name)
            del values['SCHEDULE']['B10']
            values['SCHEDULE']['C10'] = None
            engine = self.engine(values)
            union = tuple(sorted({temperature for axis in cases.values() for temperature in axis}))
            self.assertEqual(engine.temperature_options(10), union)

    def test_temperature_options_keep_row_product_case_and_saved_temperature_independent(self):
        values = self.values(FX6010)
        values['SCHEDULE'].update(C10='I/H column - 4 sides', D10=620)
        values['SCHEDULE'].update(self.values(FX6010,row=11)['SCHEDULE'])
        values['SCHEDULE'].update(C11='RHS/SHS beam - 3 sides',D11=580)
        values['SCHEDULE'].update(self.values(PRODUCT,row=12)['SCHEDULE'])
        values['SCHEDULE'].update(C12='I/H beam/column - 4 sides',D12=520)
        before = deepcopy(values)
        engine = self.engine(values)
        self.assertNotIn(620, engine.temperature_options(10))
        self.assertIn(580, engine.temperature_options(10))
        self.assertIn(620, engine.temperature_options(11))
        self.assertNotIn(580, engine.temperature_options(11))
        self.assertNotIn(520, engine.temperature_options(12))
        self.assertEqual(engine.value('SCHEDULE','O12'),1000)
        self.assertEqual(self.sc_lookup.call_args.args[2],520)
        self.assertEqual(self.sc_lookup.call_args.kwargs['case'],'R4')
        self.assertEqual(values,before)

    def test_legacy_sc902_and_separate_four_side_cases_retain_identical_named_geometry(self):
        results = []
        for label in ('I/H beam/column - 4 sides','I/H column - 4 sides','I/H beam - 4 sides'):
            values = self.values(PRODUCT, area=None)
            values['SCHEDULE'].update(C10=label,E10='Section',F10='310UB40.4',
                                      I10=2,J10=3)
            before = deepcopy(values)
            engine = self.engine(values)
            results.append(tuple(engine.value('SCHEDULE',f'{column}10') for column in
                                 ('M','N','O','Q','R','S','T','V','X')))
            self.assertEqual(self.sc_lookup.call_args.kwargs['case'],'R4')
            self.assertEqual(values,before)
        self.assertEqual(results[0],results[1])
        self.assertEqual(results[1],results[2])
        self.assertGreater(results[0][0],0)
        self.assertGreater(results[0][4],0)

    def test_named_firetex_open_members_restrict_geometry_to_ih_families(self):
        engine = self.engine()
        from estimator.excel_engine import WorkbookEngine
        original = WorkbookEngine.value
        for family in ('PFC','EA','UA','T','BT','CT'):
            def changed(instance, sheet, address):
                return family if sheet == 'SECTIONS' and address.startswith('B') and address[1:].isdigit() else original(instance,sheet,address)
            with self.subTest(family=family), patch.object(WorkbookEngine,'value',changed):
                factor,_,note = engine._geometry('310UB40.4','R3',FX5090)
                self.assertIsNone(factor)
                self.assertIn('verified A/V',note)
                self.assertIsNotNone(engine._geometry('310UB40.4','R3',PRODUCT)[0])
        self.assertIsNotNone(engine._geometry('310UB40.4','R3',FX5090)[0])
        self.assertIsNone(engine._geometry('100X100X4SHS','R4B',FX5090)[0])

    def test_esa_uses_selected_product_steel_density(self):
        values = self.values(F40=8000, D17=7000)
        values['SCHEDULE'].update(E10='ESA/M',G10=12.5)
        self.assertEqual(self.engine(values).value('SCHEDULE','M10'),100)
        values['SCHEDULE']['E10'] = 'Hp/A'
        self.assertEqual(self.engine(values).value('SCHEDULE','M10'),12.5)

    def test_unknown_product_counts_once_and_withholds_unassigned_orders(self):
        values = self.values('Unknown product')
        values['SCHEDULE'].update(self.values(row=11)['SCHEDULE'])
        engine = self.engine(values)
        self.assertEqual(engine.value('SCHEDULE','S5'),1)
        for pool in engine.product_totals().values():
            self.assertEqual(pool['whole'],'')
        self.assertEqual(engine.value('SCHEDULE','O10'),'')

    def test_missing_firetex_data_never_falls_back_to_sc902(self):
        engine = self.engine(self.values(),installed=False)
        self.assertEqual(engine.value('SCHEDULE','O10'),'')
        self.assertEqual(engine.value('SCHEDULE','V10'),'ASSESSMENT DATA UNAVAILABLE')
        self.assertEqual(engine.value('SCHEDULE','S10'),'')
        self.fx_lookup.assert_not_called()
        self.sc_lookup.assert_not_called()

    def test_invalid_or_unrepresentable_firetex_quantities_are_withheld(self):
        for settings in ({'F37':0},{'F38':0},{'F36':-1},{'F39':1.1},{'F37':5e-324},
                         {'F38':5e-324},{'F36':5e-324,'F38':1}):
            with self.subTest(settings=settings):
                values = self.values(area=5e-324 if settings.get('F36')==5e-324 else 10,**settings)
                engine = self.engine(values)
                self.assertEqual(engine.value('SCHEDULE','O10'),1000)
                self.assertEqual(engine.value('SCHEDULE','S10'),'')
                self.assertEqual(engine.product_totals(FX5090)['whole'],'')

    def test_schedule_aggregate_overflow_is_withheld(self):
        engine = self.engine({})
        pools = {name:dict(area=1e308,litres=1e308,pending=0) for name in PRODUCTS}
        with patch.object(engine,'product_totals',return_value=pools):
            self.assertEqual(engine.value('SCHEDULE','A5'),'')
            self.assertEqual(engine.value('SCHEDULE','G5'),'')

    def test_component_mass_overflow_does_not_escape_into_derived_settings(self):
        values = self.values(FX6010,H66=1,H68=.8,H71='Component mass (kg)',
                             H72=1e308,H73=1e308,H74=0)
        engine = self.engine(values)
        self.assertEqual(engine.value('SETTINGS','H80'),'')
        self.assertEqual(engine.value('SETTINGS','H81'),'')
        self.assertEqual(engine.value('SCHEDULE','O10'),1000)
        self.assertEqual(engine.value('SCHEDULE','S10'),'')
        self.assertEqual(engine.product_totals(FX6010)['pack_kg'],'')

    def test_new_settings_normalization_and_template_preserve_all_products(self):
        values = self.values(FX6010,H67=8,H68=.8,H71='Pack volume (L)',F37=20,D16=.15)
        normalized = normalize_calculator_inputs('steel_intumescent',values)
        self.assertEqual(normalized['SETTINGS'],values['SETTINGS'])
        for address in ('F42','H80'):
            with self.subTest(address=address),self.assertRaises(ValidationError):
                normalize_calculator_inputs('steel_intumescent',{'SETTINGS':{address:10}})
        payload = export_schedule_template('steel_intumescent')
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook['SCHEDULE']
        headings = [cell.value for cell in sheet[1]]
        sheet.cell(2,headings.index('Product')+1,FX6010)
        sheet.cell(2,headings.index('Exposure / member case')+1,'RHS/SHS beam - 3 sides')
        sheet.cell(2,headings.index('Fire period (min)')+1,75)
        stream = BytesIO()
        workbook.save(stream)
        result = import_schedule_workbook('steel_intumescent',stream.getvalue(),'mixed.xlsx',values)
        self.assertEqual(result['inputs']['SCHEDULE']['B10'],FX6010)
        self.assertEqual(result['inputs']['SCHEDULE']['C10'],'RHS/SHS beam - 3 sides')
        self.assertEqual(result['inputs']['SCHEDULE']['H10'],75)
        self.assertEqual(result['inputs']['SETTINGS'],values['SETTINGS'])


if __name__ == '__main__':
    unittest.main()
