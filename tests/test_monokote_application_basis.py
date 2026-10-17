"""Application method changes estimating yield, never the assessed thickness."""

from copy import deepcopy
from io import BytesIO
import base64
import http.client
import json
import math
from pathlib import Path
import tempfile
import threading
import unittest

from openpyxl import load_workbook
from pypdf import PdfReader

from estimator.calculator_defaults import default_calculator_inputs
from estimator.calculator_register import build_calculator_register
from estimator.calculator_report import build_calculator_report, build_calculator_summary_report, project_calculator_report
from estimator.catalog import ValidationError
from estimator.project_file import export_project, load_project_bytes
from estimator.schedule_rows import blank_schedule_defaults
from estimator.schedule_workbook import export_schedule_template, import_schedule_workbook
from estimator.storage import Store
from estimator.server import create_server
from estimator.workbook_calculators import calculator_definition, calculator_session, normalize_calculator_inputs, source_model
from estimator.workbook_catalog import load_workbook_catalog
from estimator.workbook_runtime import application_editable_cells


IDENTITY = 'steel_vermiculite'
PRODUCTS = (('MONOKOTE MK-6 HY', 'D569', 'D235', 'D238', 21.8 / 344, 24),
            ('MONOKOTE Z106', 'D570', 'D562', 'D565', 22.2 / 325, 25))


def inputs_for(product, *, row=10, mode='Uninjected', direct=None):
    inputs = default_calculator_inputs(IDENTITY)
    basis, direct_cell = next((basis, direct) for name, basis, direct, *_ in PRODUCTS if name == product)
    inputs['SETTINGS'].update({basis: mode, direct_cell: direct})
    inputs['SCHEDULE'] = {f'{column}{row}': value for column, value in {
        'A': 'Synthetic member', 'B': product, 'C': 'Re-entrant - 3 sides', 'D': 620,
        'E': 'Section', 'F': '410UB54', 'H': 120, 'I': 165, 'J': 2}.items()}
    inputs['CALCULATOR'] = {'D6': product}
    return inputs


def evaluated(inputs, row=10):
    _, engine, lock = calculator_session(IDENTITY, inputs)
    engine_row = row - 7
    with lock:
        return {'technical': tuple(engine.value('ENGINE', f'{column}{engine_row}') for column in ('L', 'AD', 'AE', 'AF')),
                'quick_yield': engine.value('ENGINE', 'AH2'),
                'yield': engine.value('ENGINE', f'AH{engine_row}'),
                'volume': engine.value('SCHEDULE', f'S{row}'),
                'bags': engine.value('SCHEDULE', f'T{row}'),
                'line_order': engine.value('SCHEDULE', f'U{row}'),
                'status': engine.value('SCHEDULE', f'W{row}'),
                'orders': tuple(engine.value('BAGS', f'G{r}') for r in (24, 25))}


class MonokoteApplicationBasisTests(unittest.TestCase):
    def test_derived_selectors_are_owned_by_their_panels_immediately_after_waste(self):
        raw = load_workbook_catalog(IDENTITY)
        original = deepcopy(raw)
        model = source_model(IDENTITY)
        settings = next(sheet for sheet in model['sheets'] if sheet['name'] == 'SETTINGS')
        definition = calculator_definition(IDENTITY)
        metadata = next(sheet for sheet in definition['sheets'] if sheet['name'] == 'SETTINGS')
        for cell, owner, waste_row in (('D569', 'A229', 237), ('D570', 'A559', 564)):
            with self.subTest(cell=cell):
                self.assertIn(cell, application_editable_cells(IDENTITY, 'SETTINGS'))
                field = next(field for field in model['fields']['SETTINGS'] if field['cell'] == cell)
                self.assertEqual(field['type'], 'select')
                self.assertEqual(field['validation']['formula1'], '"Injected,Uninjected"')
                self.assertEqual(settings['cells'][cell]['value'], 'Uninjected')
                explanation = settings['cells'][f'G{cell[1:]}']['value']
                self.assertIn('does not increase yield' if cell == 'D570' else '40/27', explanation)
                self.assertEqual(definition['defaults']['SETTINGS'][cell], 'Uninjected')
                section = next(section for section in metadata['settings_sections'] if section['id'] == owner)
                self.assertIn(f'A{cell[1:]}:N{cell[1:]}', section['ranges'])
                self.assertEqual(section['move_after'], [{'row': waste_row, 'rows': [int(cell[1:])]}])
                self.assertEqual(metadata['display_cells'][cell]['control'], 'select')
                self.assertNotIn(cell, next(sheet for sheet in raw['sheets'] if sheet['name'] == 'SETTINGS')['cells'])
        self.assertEqual(raw, original)
        self.assertEqual(load_workbook_catalog(IDENTITY), original)
        self.assertEqual(model['source'], raw['source'])

    def test_missing_and_explicit_uninjected_preserve_legacy_quantities_and_inputs(self):
        for product, basis, _, _, _, _ in PRODUCTS:
            for direct in (None, .123456789012345):
                with self.subTest(product=product, direct=direct):
                    legacy = inputs_for(product, direct=direct)
                    for cell in ('D569', 'D570'):
                        legacy['SETTINGS'].pop(cell, None)
                    original = deepcopy(legacy)
                    explicit = deepcopy(legacy)
                    explicit['SETTINGS'][basis] = 'Uninjected'
                    self.assertEqual(evaluated(legacy), evaluated(explicit))
                    self.assertEqual(normalize_calculator_inputs(IDENTITY, legacy), original)
                    self.assertEqual(legacy, original)

    def test_mk6_injection_applies_after_direct_priority_or_density_fallback(self):
        for direct, base_yield in ((None, 21.8 / 344), ('', 21.8 / 344), (.123456789012345, .123456789012345)):
            with self.subTest(direct=direct):
                plain = inputs_for('MONOKOTE MK-6 HY', direct=direct)
                plain['SETTINGS']['D237'] = .125
                injected = deepcopy(plain)
                injected['SETTINGS']['D569'] = 'Injected'
                a, b = evaluated(plain), evaluated(injected)
                self.assertEqual(a['yield'], base_yield)
                self.assertEqual(b['yield'], base_yield * (40 / 27))
                self.assertEqual(a['technical'], b['technical'])
                self.assertEqual(a['volume'], b['volume'])
                self.assertAlmostEqual(b['bags'], a['bags'] * 27 / 40, places=12)
                self.assertEqual(b['quick_yield'], b['yield'])
                self.assertEqual(b['line_order'], math.ceil(b['bags'] * 1.125))
                self.assertEqual(injected['SETTINGS']['D237'], .125)

    def test_z106_records_method_without_a_yield_gain_for_either_base(self):
        for direct in (None, '', .123456789012345):
            with self.subTest(direct=direct):
                plain = inputs_for('MONOKOTE Z106', direct=direct)
                plain['SETTINGS']['D564'] = .075
                injected = deepcopy(plain)
                injected['SETTINGS']['D570'] = 'Injected'
                self.assertEqual(evaluated(plain), evaluated(injected))
                self.assertEqual(normalize_calculator_inputs(IDENTITY, injected)['SETTINGS']['D570'], 'Injected')

    def test_invalid_choices_are_rejected_and_invalid_direct_yields_stay_withheld(self):
        for product, basis, direct_cell, used, _, _ in PRODUCTS:
            for value in (None, '', 'injected', 'Injected ', 'Other', 1, True, [], {}):
                with self.subTest(cell=basis, invalid=value):
                    with self.assertRaisesRegex(ValidationError, 'Injected or Uninjected'):
                        normalize_calculator_inputs(IDENTITY, {'SETTINGS': {basis: value}})
            for mode in ('Injected', 'Uninjected'):
                for value in (0, -1):
                    with self.subTest(product=product, mode=mode, direct=value):
                        inputs = inputs_for(product, mode=mode, direct=value)
                        _, engine, lock = calculator_session(IDENTITY, inputs)
                        with lock:
                            self.assertEqual(engine.value('SETTINGS', used), '')
                        result = evaluated(inputs)
                        self.assertEqual(result['bags'], '')
                        self.assertNotEqual(result['status'], 'QUANTIFIED - ESTIMATE')
            with self.assertRaises(ValidationError):
                normalize_calculator_inputs(IDENTITY, {'SETTINGS': {direct_cell: 'invalid'}})

    def test_products_stay_independent_and_last_row_uses_pooled_waste_rounding(self):
        inputs = inputs_for('MONOKOTE MK-6 HY', mode='Injected')
        second = {address.replace('10', '11'): value for address, value in inputs['SCHEDULE'].items()}
        second.update({'A11': 'Small second member', 'I11': 1, 'J11': .000000123456789})
        inputs['SCHEDULE'].update(second)
        final = inputs_for('MONOKOTE Z106', row=1009, mode='Injected')
        inputs['SCHEDULE'].update(final['SCHEDULE'])
        inputs['SETTINGS'].update({'D570': 'Injected', 'D237': .15, 'D564': .075})
        _, engine, lock = calculator_session(IDENTITY, inputs)
        with lock:
            mk6_bags = sum(engine.value('SCHEDULE', f'T{row}') for row in (10, 11))
            z106_bags = engine.value('SCHEDULE', 'T1009')
            self.assertEqual(engine.value('ENGINE', 'AH3'), (21.8 / 344) * (40 / 27))
            self.assertEqual(engine.value('ENGINE', 'AH1002'), 22.2 / 325)
            self.assertEqual(engine.value('BAGS', 'G24'), math.ceil(mk6_bags * 1.15))
            self.assertNotEqual(engine.value('BAGS', 'G24'),
                                sum(engine.value('SCHEDULE', f'U{row}') for row in (10, 11)))
            self.assertEqual(engine.value('BAGS', 'G25'), math.ceil(z106_bags * 1.075))
        changed = deepcopy(inputs)
        changed['SETTINGS'].update({'D569': 'Uninjected', 'D570': 'Uninjected'})
        a, b = evaluated(changed, 1009), evaluated(inputs, 1009)
        self.assertEqual({key: a[key] for key in ('technical', 'yield', 'volume', 'bags', 'line_order', 'status')},
                         {key: b[key] for key in ('technical', 'yield', 'volume', 'bags', 'line_order', 'status')})
        self.assertEqual(a['orders'][1], b['orders'][1])
        self.assertNotEqual(a['orders'][0], b['orders'][0])
        self.assertNotEqual(a['quick_yield'], b['quick_yield'])

    def test_save_project_and_schedule_import_preserve_basis_without_implicit_saves(self):
        inputs = inputs_for('MONOKOTE MK-6 HY', mode='Injected')
        inputs['SETTINGS']['D570'] = 'Injected'
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory) / 'basis.sqlite3')
            legacy = deepcopy(inputs)
            legacy['SETTINGS'].pop('D569')
            legacy['SETTINGS'].pop('D570')
            store.save_calculator_state(IDENTITY, legacy)
            self.assertEqual(store.calculator_state(IDENTITY)['inputs'], blank_schedule_defaults(IDENTITY, legacy))
            saved = store.save_calculator_state(IDENTITY, inputs)
            self.assertEqual(store.calculator_state(IDENTITY), saved)
            with store.connect() as db:
                before = db.execute('SELECT * FROM calculator_states ORDER BY id').fetchall()
            payload = export_project(store, {'estimate': {}, 'calculators': {IDENTITY: {'inputs': inputs}}})
            loaded = load_project_bytes(store, payload)
            self.assertEqual(loaded['calculators'][IDENTITY]['inputs']['SETTINGS'], inputs['SETTINGS'])
            imported = import_schedule_workbook(IDENTITY, export_schedule_template(IDENTITY), 'schedule.xlsx', inputs)
            self.assertEqual(imported['inputs']['SETTINGS'], inputs['SETTINGS'])
            with store.connect() as db:
                self.assertEqual(db.execute('SELECT * FROM calculator_states ORDER BY id').fetchall(), before)
            for cell in ('D569', 'D570'):
                invalid = deepcopy(inputs)
                invalid['SETTINGS'][cell] = 'Other'
                with self.assertRaises(ValidationError):
                    store.save_calculator_state(IDENTITY, invalid)
                with self.assertRaises(ValidationError):
                    export_project(store, {'estimate': {}, 'calculators': {IDENTITY: {'inputs': invalid}}})
                altered = json.loads(payload)
                altered['calculators'][IDENTITY]['inputs']['SETTINGS'][cell] = 'Other'
                with self.assertRaises(ValidationError):
                    load_project_bytes(store, json.dumps(altered).encode())
                with self.assertRaises(ValidationError):
                    # The importer preserves settings; the API validates the
                    # proposed inputs before they can be applied or calculated.
                    normalize_calculator_inputs(IDENTITY, import_schedule_workbook(
                        IDENTITY, export_schedule_template(IDENTITY), 'schedule.xlsx', invalid)['inputs'])
            with store.connect() as db:
                self.assertEqual(db.execute('SELECT * FROM calculator_states ORDER BY id').fetchall(), before)

    def test_current_draft_pdf_and_values_only_xlsx_use_the_same_method_quantities(self):
        plain = inputs_for('MONOKOTE MK-6 HY')
        draft = inputs_for('MONOKOTE MK-6 HY', mode='Injected')
        original = deepcopy(draft)
        projection = project_calculator_report(IDENTITY, draft)
        self.assertEqual(projection['rows'][0]['values']['T'], evaluated(draft)['bags'])
        self.assertNotEqual(projection['rows'][0]['values']['T'], evaluated(plain)['bags'])
        workbook = load_workbook(BytesIO(build_calculator_register(IDENTITY, draft)), data_only=False)
        try:
            self.assertFalse(any(cell.data_type == 'f' for sheet in workbook for row in sheet for cell in row))
            self.assertEqual(workbook['Summary']['E8'].value, evaluated(draft)['orders'][0])
        finally:
            workbook.close()
        for builder in (build_calculator_report, build_calculator_summary_report):
            pdf = PdfReader(BytesIO(builder(IDENTITY, draft)))
            text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
            self.assertIn('MONOKOTE MK-6 HY', text)
            if builder is build_calculator_summary_report:
                self.assertRegex(text, rf'\b{int(evaluated(draft)["orders"][0])}(?:\.00)?\b')
        self.assertEqual(draft, original)


class MonokoteApplicationBasisHttpTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.store = Store(Path(cls.temp.name) / 'basis-http.sqlite3')
        cls.server = create_server(0, cls.store.path, library_directory=Path(cls.temp.name) / 'library')
        assert cls.server.server_port != 8765
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.template = base64.b64encode(export_schedule_template(IDENTITY)).decode()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()
        cls.temp.cleanup()

    def setUp(self):
        with self.store.connect() as db:
            db.execute('DELETE FROM calculator_states')

    def request(self, method, action, payload, expected=200):
        connection = http.client.HTTPConnection('127.0.0.1', self.server.server_port, timeout=120)
        try:
            connection.request(method, f'/api/calculators/{IDENTITY}/{action}', json.dumps(payload),
                               {'Content-Type': 'application/json'})
            result = connection.getresponse()
            content = result.read()
            self.assertEqual(result.status, expected, content[:300])
            return content
        finally:
            connection.close()

    def saved_rows(self):
        with self.store.connect() as db:
            return db.execute('SELECT * FROM calculator_states ORDER BY id').fetchall()

    def test_invalid_basis_is_rejected_at_all_http_boundaries_without_writes(self):
        for cell in ('D569', 'D570'):
            for value in (None, 'Other', [], {}):
                inputs = {'SETTINGS': {cell: value}}
                for action in ('calculate', 'worksheet', 'report.pdf', 'summary.pdf', 'register.xlsx', 'state', 'import'):
                    with self.subTest(cell=cell, value=value, action=action):
                        payload = {'inputs': inputs}
                        if action == 'import':
                            payload.update(filename='schedule.xlsx', content_base64=self.template)
                        error = json.loads(self.request('PUT' if action == 'state' else 'POST', action, payload, 400))
                        self.assertIn('Injected or Uninjected', error['error'])
                        self.assertEqual(self.saved_rows(), [])

    def test_draft_calculation_and_both_exports_do_not_save_and_explicit_save_roundtrips(self):
        saved = inputs_for('MONOKOTE MK-6 HY')
        self.store.save_calculator_state(IDENTITY, saved)
        before = self.saved_rows()
        draft = inputs_for('MONOKOTE MK-6 HY', mode='Injected')
        draft['SETTINGS']['D570'] = 'Injected'
        result = json.loads(self.request('POST', 'calculate',
                            {'inputs': draft, 'sheet': 'SCHEDULE', 'start_row': 10, 'row_count': 1}))
        self.assertEqual(result['inputs']['SETTINGS']['D569'], 'Injected')
        self.assertEqual(next(cell['value'] for cell in result['rows'][0]['cells'] if cell['address'] == 'T10'),
                         evaluated(draft)['bags'])
        for action in ('report.pdf', 'summary.pdf', 'register.xlsx'):
            content = self.request('POST', action, {'inputs': draft})
            self.assertTrue(content.startswith(b'%PDF') if action.endswith('.pdf') else content.startswith(b'PK'))
            self.assertEqual(self.saved_rows(), before)
        receipt = json.loads(self.request('PUT', 'state', {'inputs': draft}))
        self.assertEqual(receipt['inputs']['SETTINGS']['D569'], 'Injected')
        self.assertEqual(receipt['inputs']['SETTINGS']['D570'], 'Injected')
        self.assertEqual(self.store.calculator_state(IDENTITY), receipt)


if __name__ == '__main__':
    unittest.main()
