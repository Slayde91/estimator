"""Private assessment policy with synthetic tables, never publisher table data."""

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from pypdf import PdfReader

from estimator import monokote_hollow as policy
from estimator.calculator_defaults import default_calculator_inputs
from estimator.calculator_register import build_calculator_register
from estimator.calculator_report import build_calculator_report, build_calculator_summary_report, project_calculator_report
from estimator.excel_engine import WorkbookEngine
from estimator.workbook_calculators import calculator_session, calculate_worksheet, calculator_definition, source_model


def synthetic_dataset():
    """Invented values exercise addressing without redistributing evidence."""
    return {
        'schema_version': 1, 'id': policy.DATASET_ID,
        'source': {'sha256': '1' * 64, 'report': 'FAR4856 Issue 2', 'date': '2020-11-09'},
        'scope': {'exposure': 'H4', 'member': 'hollow_column', 'families': ['SHS', 'RHS', 'CHS']},
        'temperatures': list(policy.TEMPERATURES), 'factors': list(policy.FACTORS),
        'tables': [{'period': period, 'page': index + 9, 'table': index + 1,
                    'values': [[round((index + 1) * 10 + factor_index / 10 + temp_index / 100, 2)
                                for temp_index in range(len(policy.TEMPERATURES))]
                               for factor_index in range(len(policy.FACTORS))]}
                   for index, period in enumerate(policy.PERIODS)],
        'blocked_cells': [{'period': 240, 'factor': 250, 'temperature': 650,
                           'reason': 'Synthetic publisher anomaly awaiting verification.'}],
    }


def encoded(data):
    return json.dumps(data, separators=(',', ':'), allow_nan=False).encode()


def row_inputs(product='MONOKOTE MK-6 HY', section='100X100X9SHS', *, row=10, **changes):
    inputs = default_calculator_inputs('steel_vermiculite')
    values = {'A': 'Hollow diagnostic', 'B': product, 'C': 'Hollow - 4 sides',
              'D': 550, 'E': 'Section', 'F': section, 'H': 120, 'I': 2, 'J': 3}
    values.update(changes)
    inputs['SCHEDULE'] = {f'{c}{row}': value for c, value in values.items()}
    inputs['CALCULATOR'] = {f'D{target}': values[col] for target, col in (
        (6, 'B'), (7, 'C'), (8, 'D'), (9, 'E'), (10, 'F'), (12, 'H'), (14, 'I'), (15, 'J'))}
    if 'G' in values:
        inputs['CALCULATOR']['D11'] = values['G']
    return inputs


class MonokoteHollowTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.data = synthetic_dataset()
        # A source dash is represented by null, never zero or a guessed value.
        self.data['tables'][5]['values'][-1][0] = None
        self.payload = encoded(self.data)
        self.path = Path(self.folder.name) / policy.DATASET_FILENAME
        self.path.write_bytes(self.payload)
        self.pin = patch.object(policy, 'DATASET_SHA256', sha256(self.payload).hexdigest())
        self.pin.start()
        self.addCleanup(self.pin.stop)
        self.environment = patch.dict('os.environ', {'CEASEFIRE_CALCULATOR_EVIDENCE_DIRECTORY': self.folder.name})
        self.environment.start()
        self.addCleanup(self.environment.stop)

    def engine(self, inputs):
        return calculator_session('steel_vermiculite', inputs)[1]

    def test_loader_accepts_bound_snapshot_and_rejects_missing_tampered_and_special_files(self):
        assessment = policy.load_assessment()
        self.assertFalse(assessment.error)
        self.assertEqual(assessment.digest, sha256(self.payload).hexdigest())
        self.assertEqual(len(assessment.tables), 6)
        self.path.write_bytes(self.payload + b' ')
        self.assertTrue(policy.load_assessment().error)
        self.path.unlink()
        self.assertTrue(policy.load_assessment().error)
        self.path.mkdir()
        self.assertTrue(policy.load_assessment().error)

    def test_matching_hash_does_not_bypass_schema_scope_or_required_hold(self):
        changes = [lambda d: d.update(temperatures=[350]),
                   lambda d: d['scope'].update(member='beam'),
                   lambda d: d['tables'].pop(),
                   lambda d: d['tables'][0]['values'][0].__setitem__(0, -1),
                   lambda d: d.update(blocked_cells=[])]
        for change in changes:
            data = deepcopy(self.data)
            change(data)
            payload = encoded(data)
            with self.subTest(change=change), patch.object(policy, 'DATASET_SHA256', sha256(payload).hexdigest()):
                with self.assertRaises(ValueError):
                    policy.validate_assessment_payload(payload)

    def test_exact_named_user_case_both_products_quantities_and_original_excel_defect(self):
        for product, mass, density in [('MONOKOTE MK-6 HY', 21.8, 344), ('MONOKOTE Z106', 22.2, 325)]:
            inputs = row_inputs(product)
            original = WorkbookEngine(source_model('steel_vermiculite'), inputs)
            engine = self.engine(inputs)
            with self.subTest(product=product):
                self.assertEqual(original.value('SCHEDULE', 'V10'), 'CALCULATION ERROR')
                self.assertEqual(original.value('SCHEDULE', 'O10'), 25)
                self.assertEqual(engine.value('SCHEDULE', 'N10'), 125)
                self.assertEqual(engine.value('SCHEDULE', 'P10'), 41.94)
                self.assertAlmostEqual(engine.value('SCHEDULE', 'S10'), 2 * 3 * .4 * 41.94 / 1000)
                self.assertAlmostEqual(engine.value('SCHEDULE', 'T10'), (2 * 3 * .4 * 41.94 / 1000) / (mass / density))
                self.assertEqual(engine.value('SCHEDULE', 'W10'), 'QUANTIFIED - ESTIMATE')
                self.assertIn('Retained printed section factor 122', engine.value('SCHEDULE', 'Y10'))

    def test_named_factor_is_temperature_independent_and_family_scope_is_explicit(self):
        for section, upper in [('100X100X9SHS', 125), ('150X150X6SHS', 175),
                               ('150X100X6RHS', 175), ('168.3X6.4CHS', 165)]:
            for temperature in (350, 550, 620, 750):
                with self.subTest(section=section, temperature=temperature):
                    engine = self.engine(row_inputs(section=section, D=temperature))
                    self.assertEqual(engine.value('SCHEDULE', 'N10'), upper)
                    self.assertIsInstance(engine.value('SCHEDULE', 'P10'), float)
                    self.assertIn('columns only', engine.value('SCHEDULE', 'Y10'))
        engine = self.engine(row_inputs(section='410UB54'))
        self.assertEqual(engine.value('SCHEDULE', 'V10'), 'UNSUPPORTED HOLLOW SECTION')
        self.assertEqual(engine.value('SCHEDULE', 'T10'), '')

    def test_all_period_comparison_uses_new_evidence_and_visible_basis(self):
        inputs = row_inputs()
        engine = self.engine(inputs)
        for column, expected in [('C', 11.94), ('E', 21.94), ('F', 31.94),
                                 ('G', 41.94), ('H', 51.94), ('I', 61.94)]:
            self.assertEqual(engine.value('CALCULATOR', column + '29'), expected)
            self.assertEqual(engine.value('CALCULATOR', column + '30'), 'UPPER ROW')
        for column in ('B', 'D'):
            self.assertEqual(engine.value('CALCULATOR', column + '29'), '')
            self.assertEqual(engine.value('CALCULATOR', column + '30'), 'FRL NOT TABULATED')
        result = calculate_worksheet('steel_vermiculite', inputs, 'CALCULATOR')
        self.assertNotIn('H23:N24', result['omitted_ranges'])
        self.assertIn('FAR4856', engine.value('CALCULATOR', 'H23'))
        self.assertIn('columns only', engine.value('CALCULATOR', 'H23'))

    def test_shared_factor_and_entered_esa_follow_retained_input_basis(self):
        engine = self.engine(row_inputs(section='100X100X10SHS'))
        self.assertIsInstance(engine.value('SCHEDULE', 'P10'), float)
        self.assertIn('Retained shared section factor', engine.value('SCHEDULE', 'Y10'))
        engine = self.engine(row_inputs(E='ESA/M', G=15.4))
        self.assertEqual(engine.value('SCHEDULE', 'N10'), 125)
        self.assertAlmostEqual(engine.value('ENGINE', 'R3'), 15.4 * 7850 / 1000)
        self.assertIn('User-entered ESA/M', engine.value('SCHEDULE', 'Y10'))

    def test_start_guidance_describes_assessment_without_changing_workbook_source(self):
        definition = calculator_definition('steel_vermiculite')
        settings = next(sheet for sheet in definition['sheets'] if sheet['name'] == 'SETTINGS')
        self.assertIn('Hp/A and ESA/M', settings['display_text']['A285'])
        self.assertIn('FAR4856 Issue 2', settings['display_text']['A285'])
        self.assertIn('columns only', settings['display_text']['A289'])
        original = WorkbookEngine(source_model('steel_vermiculite'))
        self.assertIn('hollow cases need named sections', original.value('SETTINGS', 'A289'))

    def test_lookup_policies_bounds_unpublished_and_explicit_publisher_hold(self):
        for factor, temperature, period, expected in [
                (0, 550, 120, 'ENTER POSITIVE FACTOR'), (-1, 550, 120, 'ENTER POSITIVE FACTOR'),
                (29.99, 550, 120, 'FACTOR OUT OF RANGE'), (365.01, 550, 120, 'FACTOR OUT OF RANGE'),
                (125, 575, 120, 'UNSUPPORTED CASE / TEMP'), (125, 550, 45, 'FRL NOT TABULATED'),
                (365, 350, 240, 'NP - NOT PUBLISHED'), (250, 650, 240, 'BLOCKED - SOURCE REVIEW')]:
            with self.subTest(factor=factor, temperature=temperature, period=period):
                engine = self.engine(row_inputs(E='Hp/A', G=factor, D=temperature, H=period))
                self.assertEqual(engine.value('SCHEDULE', 'V10'), expected)
                self.assertEqual(engine.value('SCHEDULE', 'T10'), '')
                self.assertIn(expected, engine.value('SCHEDULE', 'Y10'))
        for factor, expected in [(125, 'PUBLISHED - ASSESSMENT TABLE'), (122, 'NO EXACT FACTOR ROW')]:
            inputs = row_inputs(E='Hp/A', G=factor)
            inputs['SETTINGS']['D14'] = 'Exact only'
            self.assertEqual(self.engine(inputs).value('SCHEDULE', 'V10'), expected)

    def test_missing_and_tampered_evidence_invalidate_identical_cached_inputs(self):
        inputs = row_inputs()
        first = self.engine(inputs)
        self.assertEqual(first.value('SCHEDULE', 'P10'), 41.94)
        self.path.write_bytes(self.payload + b' ')
        changed = self.engine(inputs)
        self.assertIsNot(changed, first)
        self.assertEqual(changed.value('SCHEDULE', 'V10'), 'ASSESSMENT DATA UNAVAILABLE')
        self.assertEqual(changed.value('SCHEDULE', 'T10'), '')
        self.assertEqual(changed.value('CALCULATOR', 'G29'), '')
        self.path.unlink()
        missing = self.engine(inputs)
        self.assertIsNot(missing, changed)
        self.assertEqual(missing.value('SCHEDULE', 'T10'), '')
        self.path.write_bytes(self.payload)
        self.assertEqual(self.engine(inputs).value('SCHEDULE', 'P10'), 41.94)

    def test_h3_other_products_and_solid_cases_still_match_original(self):
        controls = [row_inputs(C='Hollow - 3 sides', D=620),
                    row_inputs(section='410UB54', C='Re-entrant - 3 sides', D=620),
                    row_inputs('CAFCO 300')]
        self.path.unlink()
        for inputs in controls:
            original = WorkbookEngine(source_model('steel_vermiculite'), inputs)
            engine = self.engine(inputs)
            for column in ('M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U', 'V', 'W', 'X', 'Y'):
                self.assertEqual(engine.value('SCHEDULE', column + '10'), original.value('SCHEDULE', column + '10'))

    def test_last_row_blank_zero_quantity_and_invalid_yield_still_withhold_bags(self):
        engine = self.engine(row_inputs(row=1009))
        self.assertEqual(engine.value('SCHEDULE', 'P1009'), 41.94)
        self.assertEqual(engine.value('SCHEDULE', 'W1009'), 'QUANTIFIED - ESTIMATE')
        for quantity in (None, 0, -1, 1.5):
            engine = self.engine(row_inputs(I=quantity))
            self.assertEqual(engine.value('SCHEDULE', 'T10'), '')
            self.assertNotEqual(engine.value('SCHEDULE', 'W10'), 'QUANTIFIED - ESTIMATE')
        inputs = row_inputs()
        inputs['SETTINGS']['D235'] = 0
        engine = self.engine(inputs)
        self.assertEqual(engine.value('SCHEDULE', 'T10'), '')
        self.assertEqual(engine.value('SCHEDULE', 'W10'), 'ENTER VERIFIED YIELD')

    def test_report_and_values_only_xlsx_share_corrected_snapshot_and_basis(self):
        inputs = row_inputs()
        data = project_calculator_report('steel_vermiculite', inputs)
        self.assertEqual(data['rows'][0]['values']['P'], 41.94)
        self.assertIn('FAR4856 Issue2 p12 Table 4', data['rows'][0]['values']['Y'])
        self.assertTrue(data['rows'][0]['complete'])
        self.assertEqual(data['assessment_evidence']['dataset_sha256'], sha256(self.payload).hexdigest())
        workbook = load_workbook(BytesIO(build_calculator_register('steel_vermiculite', inputs)))
        cells = [cell for sheet in workbook for row in sheet for cell in row]
        self.assertTrue(any(cell.value == 41.94 for cell in cells))
        self.assertTrue(any('FAR4856 Issue2 p12 Table 4' in str(cell.value) for cell in cells))
        self.assertTrue(any('Retained printed section factor 122' in str(cell.value) for cell in cells))
        self.assertTrue(any('columns only' in str(cell.value) for cell in cells))
        self.assertTrue(any(sha256(self.payload).hexdigest() in str(cell.value) for cell in cells))
        self.assertFalse(any(cell.data_type == 'f' for cell in cells))
        pdf = PdfReader(BytesIO(build_calculator_report('steel_vermiculite', inputs)))
        text = '\n'.join(page.extract_text() for page in pdf.pages)
        self.assertGreaterEqual(len(pdf.pages), 2)
        self.assertNotIn('Full schedule', pdf.pages[0].extract_text())
        self.assertIn('Full schedule', pdf.pages[1].extract_text())
        self.assertIn('100X100X9SHS', pdf.pages[1].extract_text())
        self.assertIn('41.94', text)
        self.assertIn('FAR4856', text)
        self.assertIn('columns only', text)
        self.assertIn('Retained printed section factor 122', ' '.join(text.split()))
        self.assertIn(sha256(self.payload).hexdigest(), text.replace('\n', ''))
        summary = PdfReader(BytesIO(build_calculator_summary_report('steel_vermiculite', inputs)))
        summary_text = '\n'.join(page.extract_text() for page in summary.pages)
        self.assertGreaterEqual(len(summary.pages), 2)
        self.assertNotIn('Net bags', summary.pages[0].extract_text())
        self.assertIn('Net bags', summary.pages[1].extract_text())
        self.assertIn('MONOKOTE MK-6 HY', summary.pages[1].extract_text())
        self.assertIn('columns only', summary_text)
        self.assertIn('Retained printed section factor 122', ' '.join(summary_text.split()))
        self.assertIn(sha256(self.payload).hexdigest(), summary_text.replace('\n', ''))
        self.path.unlink()
        data = project_calculator_report('steel_vermiculite', inputs)
        self.assertFalse(data['rows'][0]['complete'])
        self.assertEqual(data['rows'][0]['values']['P'], '')


if __name__ == '__main__':
    unittest.main()
