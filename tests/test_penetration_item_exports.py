"""Current-item exports stay independent of schedules and retain source identity."""
import base64
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import http.client
import json
from pathlib import Path
import tempfile
import threading
import unittest

from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from estimator.catalog import ValidationError
from estimator.penetration_calculator import calculate, definition
from estimator.penetration_item_exports import prepare_item_snapshot, build_item_xlsx, render_item_pdf
from estimator.reference_library import ReferenceNotFound
from estimator.server import create_server


def png():
    stream = BytesIO()
    image = Image.new('RGB', (320, 160), 'white')
    # Visible, asymmetric diagram prevents an empty/image-count-only fixture.
    for x in range(40, 150):
        for y in range(25, 110):
            image.putpixel((x, y), (180, 25, 20))
    image.save(stream, format='PNG')
    return stream.getvalue()


class DiagramLibrary:
    def __init__(self, payload=None):
        self.payload = payload
        self.calls = []

    def diagram_asset(self, item_id, thumbnail=False):
        self.calls.append((item_id, thumbnail))
        if self.payload is None:
            raise ReferenceNotFound()
        return self.payload, 'image/png', 'Reviewed source drawing.png'


def body():
    return {'item': {'id': 'unsaved-current-item', 'inputs': {
        'T': 'Current item only', 'U': 'Unsaved system details', 'O': 3,
        'K': 'Plastic Pipes', 'AW': 123.456789012345, 'AX': 234.567890123456,
        'AI': 7.12345678901234, 'AJ': 4.23456789012345}},
        'globals': {}, 'configuration': {}}


class CurrentItemExportTests(unittest.TestCase):
    def snapshot(self, value, library=None):
        return prepare_item_snapshot(value, library=library or DiagramLibrary(), default_configuration={})

    def test_exact_single_unsaved_item_and_engine_prices_without_mutating_capture(self):
        value = body()
        before = deepcopy(value)
        snapshot = self.snapshot(value)
        expected = calculate({'globals': value['globals'], 'rows': [value['item']]}, value['configuration'])
        self.assertEqual(snapshot['result'], expected)
        self.assertEqual(value, before)
        self.assertEqual(len(snapshot['result']['rows']), 1)
        self.assertEqual(snapshot['result']['rows'][0]['id'], 'unsaved-current-item')

    def test_structural_schedule_paths_urls_unknown_and_calculated_fields_rejected(self):
        invalid = [dict(body(), rows=[body()['item']]), dict(body(), draft={'rows': [body()['item']]}),
                   dict(body(), item=[body()['item']]), dict(body(), attachments=[]),
                   dict(body(), item={**body()['item'], 'outputs': {'H': 0}}),
                   dict(body(), item={**body()['item'], 'path': 'C:/private.pdf'})]
        for value in invalid:
            with self.subTest(value=list(value)):
                with self.assertRaises(ValidationError):
                    self.snapshot(value)
        for filename in ('C:/private.png', '../private.png', 'https://host/image.png', 'item.pdf'):
            with self.subTest(filename=filename):
                with self.assertRaises(ValidationError):
                    self.snapshot(dict(body(), diagram={'filename': filename, 'content_base64': base64.b64encode(png()).decode()}))

    def test_pending_image_overrides_library_with_no_managed_lookup_and_keeps_identity(self):
        value = body()
        value['item']['library_item_id'] = 'pfp-saved'
        payload = png()
        value['diagram'] = {'filename': 'Current source.png', 'content_base64': base64.b64encode(payload).decode()}
        library = DiagramLibrary()
        snapshot = self.snapshot(value, library)
        self.assertEqual(library.calls, [])
        self.assertEqual(snapshot['diagram_identity']['kind'], 'current unsaved image')
        self.assertEqual(snapshot['diagram_identity']['source_sha256'], sha256(payload).hexdigest())
        self.assertEqual(snapshot['diagram']['width'], 320)
        self.assertEqual(snapshot['diagram']['height'], 160)

    def test_managed_diagram_full_not_thumbnail_missing_and_unbounded_fail_clearly(self):
        value = body()
        value['item']['library_item_id'] = 'pfp-managed'
        library = DiagramLibrary(png())
        snapshot = self.snapshot(value, library)
        self.assertEqual(library.calls, [('pfp-managed', False)])
        self.assertEqual(snapshot['diagram_identity']['source_sha256'], sha256(png()).hexdigest())
        with self.assertRaisesRegex(ValidationError, 'unavailable'):
            self.snapshot(value)
        with self.assertRaises(ValidationError):
            self.snapshot(value, DiagramLibrary(b'x' * (15 * 1048576 + 1)))
        with self.assertRaises(ValidationError):
            self.snapshot(value, DiagramLibrary(b'%PDF-1.7 invalid image'))

    def test_xlsx_all_raw_and_effective_fields_exact_no_formula_or_other_item(self):
        value = body()
        value['item']['inputs']['T'] = '=HYPERLINK("https://example.invalid","not executable")'
        snapshot = self.snapshot(value)
        before = deepcopy(snapshot)
        workbook = load_workbook(BytesIO(build_item_xlsx(snapshot)), data_only=False)
        self.assertEqual(snapshot, before)
        self.assertNotIn('Schedule', workbook.sheetnames)
        rows = {row[2].value: row for row in workbook['Inputs'].iter_rows(min_row=5)}
        self.assertEqual(rows['AW'][3].value, value['item']['inputs']['AW'])
        self.assertEqual(rows['AX'][3].value, value['item']['inputs']['AX'])
        self.assertEqual(rows['T'][3].value, value['item']['inputs']['T'])
        self.assertEqual(rows['T'][3].data_type, 's')
        self.assertEqual(len(rows), len(definition({})['row_fields']))
        self.assertTrue(all(cell.data_type != 'f' and cell.hyperlink is None for sheet in workbook for row in sheet for cell in row))
        self.assertEqual(workbook['Item']['B5'].value, 'unsaved-current-item')

    def test_pdf_relevant_unsaved_fields_and_one_real_diagram_no_schedule(self):
        value = body()
        payload = png()
        value['diagram'] = {'filename': 'Current source.png', 'content_base64': base64.b64encode(payload).decode()}
        snapshot = self.snapshot(value)
        before = deepcopy(snapshot)
        reader = PdfReader(BytesIO(render_item_pdf(snapshot)))
        text = '\n'.join(page.extract_text() for page in reader.pages)
        self.assertIn('Current Firestopping item', text)
        self.assertIn('Current item only', text)
        self.assertIn('Unsaved system details', text)
        self.assertIn(sha256(payload).hexdigest(), text)
        self.assertNotIn('Schedule', text)
        # First-page source image plus company logo; the source is 320x160.
        dimensions = [(image.image.width, image.image.height) for page in reader.pages for image in page.images]
        self.assertIn((320, 160), dimensions)
        self.assertEqual(dimensions.count((320, 160)), 1)
        self.assertEqual(snapshot, before)

    def test_absent_diagram_is_explicit_not_silently_omitted(self):
        reader = PdfReader(BytesIO(render_item_pdf(self.snapshot(body()))))
        self.assertIn('No source diagram is attached to this current item.', '\n'.join(p.extract_text() for p in reader.pages))

    def test_linked_missing_diagram_xlsx_details_available_pdf_stays_explicit(self):
        value = body()
        value['item']['library_item_id'] = 'pfp-no-image'
        snapshot = prepare_item_snapshot(value, library=DiagramLibrary(), default_configuration={}, require_diagram=False)
        workbook = load_workbook(BytesIO(build_item_xlsx(snapshot)), data_only=False)
        identity = dict(workbook['Source'].iter_rows(min_row=5, values_only=True))
        self.assertEqual(identity['library_item_id'], 'pfp-no-image')
        self.assertEqual(identity['status'], 'unavailable')
        self.assertEqual(workbook['Item']['B5'].value, value['item']['id'])
        with self.assertRaisesRegex(ValidationError, 'unavailable'):
            render_item_pdf(snapshot)


class CurrentItemHTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.server = create_server(0, Path(cls.temp.name) / 'test.sqlite3', library_directory=Path(cls.temp.name) / 'library')
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()
        cls.temp.cleanup()

    def request(self, route, payload):
        connection = http.client.HTTPConnection('127.0.0.1', self.server.server_port, timeout=30)
        connection.request('POST', route, json.dumps(payload), {'Content-Type': 'application/json'})
        response = connection.getresponse()
        result = response.status, response.getheader('Content-Type'), response.getheader('Content-Disposition'), response.read()
        connection.close()
        return result

    def test_both_routes_export_one_item_and_reject_schedule_body(self):
        for extension, magic in (('pdf', b'%PDF-'), ('xlsx', b'PK')):
            with self.subTest(extension=extension):
                status, mime, disposition, content = self.request('/api/penetration/item.' + extension, body())
                self.assertEqual(status, 200)
                self.assertTrue(content.startswith(magic))
                self.assertIn('Firestopping-Item.' + extension, disposition)
                status, _, _, _ = self.request('/api/penetration/item.' + extension, {'draft': {'rows': [body()['item']]}})
                self.assertEqual(status, 400)


if __name__ == '__main__':
    unittest.main()
