"""Resolved application labour hours survive exports without filling raw drafts."""

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from pypdf import PdfReader

from estimator.catalog import ROOT
from estimator.firestopping_library import FirestoppingLibrary
from estimator.penetration_calculator import calculate, definition, source_model
from estimator.penetration_report import build_penetration_register, render_penetration_pdf
from estimator.project_file import export_project, load_project_bytes
from estimator.report import render_quote_pdf
from estimator.storage import Store


def draft(globals=None, **overrides):
    choices = {field['column']: field.get('options', []) for field in definition({})['row_fields']}
    return {'globals': globals or {}, 'rows': [{'id': 'service', 'inputs': {
        'K': 'Synthetic service', 'O': 3, 'Y': choices['Y'][0], 'W': choices['W'][0],
        'AL': 40, 'AN': 2, 'AH': .5, 'AI': 5, 'AJ': 7, **overrides}}]}


def input_records(payload):
    workbook = load_workbook(BytesIO(payload))
    return {row[1].value: (row[2].value, row[4].value)
            for row in workbook['Inputs'].iter_rows(min_row=5)}


def setting_records(payload):
    workbook = load_workbook(BytesIO(payload))
    return {row[0].value: (row[1].value, row[2].value)
            for row in workbook['Settings'].iter_rows(min_row=5)}


def pdf_text(payload):
    return '\n'.join(page.extract_text() for page in PdfReader(BytesIO(payload)).pages)


class PenetrationLabourExportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.store = Store(self.root / 'state.sqlite3')

    def test_automatic_exports_show_resolved_hours_and_keep_raw_values_absent(self):
        source = draft()
        result = calculate(source)
        before = deepcopy(result)
        self.assertEqual(result['rows'][0]['input_defaults'], {'register_allowance_hours': .25, 'pipe_labour_hours': .25})
        records = input_records(build_penetration_register(result, result['definition'], {}))
        self.assertEqual(records['Pipe Labour'], (.25, 'Automatic'))
        self.assertEqual(setting_records(build_penetration_register(result, result['definition'], {}))['Register Allowance'], (.25, 'hrs'))
        text = pdf_text(render_penetration_pdf(result, result['definition'], {}))
        self.assertIn('Schedule', text)
        for removed in ('Settings', 'Register Allowance', 'Pipe Labour', 'Inputs', 'Calculated detail'):
            self.assertNotIn(removed, text)
        self.assertEqual(result, before)
        self.assertNotIn('register_allowance_hours', source['rows'][0]['inputs'])
        self.assertNotIn('pipe_labour_hours', result['draft']['rows'][0]['inputs'])

    def test_explicit_manual_zero_exports_as_zero_and_never_becomes_automatic(self):
        result = calculate(draft(globals={'register_allowance_hours': 0}, pipe_labour_hours=0))
        self.assertEqual(result['rows'][0]['outputs']['DK'], 1.5)  # AH .5 × Item QTY3 only.
        records = input_records(build_penetration_register(result, result['definition'], {}))
        self.assertEqual(records['Pipe Labour'], (0, 'Manual'))
        self.assertEqual(setting_records(build_penetration_register(result, result['definition'], {}))['Register Allowance'], (0, 'hrs'))
        text = pdf_text(render_penetration_pdf(result, result['definition'], {}))
        self.assertNotIn('0.00 (Manual)', text)
        self.assertEqual(result['draft']['rows'][0]['inputs']['pipe_labour_hours'], 0)

    def test_register_exports_editable_service_routes_and_task_hour_bands(self):
        result = calculate(draft(globals={
            'service_routes': {'Lagged Pipes': ['Lagged Pipes', 'Insulated Pipes']},
            'labour_bands': {
                'pipe': [{'maximum': 80, 'hours': .65}],
                'board': [{'maximum': .2, 'hours': .55}],
            },
        }))
        workbook = load_workbook(BytesIO(build_penetration_register(
            result, result['definition'], {})))
        rows = [tuple(cell.value for cell in row)
                for row in workbook['Settings'].iter_rows()]
        self.assertIn(('Lagged Pipes', 'Lagged Pipes; Insulated Pipes', None, None), rows)
        self.assertIn(('Pipe Labour', 80, 'mm', .65), rows)
        self.assertIn(('Board Task Hours', .2, 'm²', .55), rows)

    def test_paired_dimensions_export_once_without_changing_source_values(self):
        source = draft(AQ=450, AR=100, AW=1000, AX=500)
        result = calculate(source)
        records = input_records(build_penetration_register(result, result['definition'], {}))
        self.assertEqual(records['Cabletray W x D'], ('450.00 x 100.00', None))
        self.assertEqual(records['Board/Batt W x L'], ('1,000.00 x 500.00', None))
        self.assertNotIn('Cabletray Depth', records)
        self.assertNotIn('Board Length', records)
        text = pdf_text(render_penetration_pdf(result, result['definition'], {}))
        self.assertNotIn('Cabletray W x D', text)
        self.assertNotIn('450.00 x 100.00', text)
        self.assertEqual(result['draft']['rows'][0]['inputs']['AQ'], 450)
        self.assertEqual(result['draft']['rows'][0]['inputs']['AR'], 100)

    def test_auto_tracks_diameter_manual_hours_and_source_identity_survive_project_roundtrip(self):
        model = deepcopy(source_model())
        source_hash = sha256((ROOT / 'data/penetration.json.gz').read_bytes()).hexdigest()
        auto = draft(register_allowance_hours=None, pipe_labour_hours=None, AL=50)
        first = calculate(auto)
        auto['rows'][0]['inputs']['AL'] = 100
        second = calculate(auto)
        self.assertEqual(first['rows'][0]['input_defaults']['pipe_labour_hours'], .25)
        self.assertEqual(second['rows'][0]['input_defaults']['pipe_labour_hours'], .30)
        self.assertEqual(input_records(build_penetration_register(second, second['definition'], {}))['Pipe Labour'], (.30, 'Automatic'))
        manual = draft(globals={'register_allowance_hours': 0}, pipe_labour_hours=.1234567890123, AL=301)
        payload = export_project(self.store, {'estimate': {'title': 'Manual labour'},
            'penetration': {'draft': auto, 'composer': manual, 'library_tracking_version': 1}})
        loaded = load_project_bytes(self.store, payload)
        stored = json.loads(payload)['penetration']
        self.assertIsNone(stored['draft']['rows'][0]['inputs']['pipe_labour_hours'])
        self.assertEqual(stored['composer']['rows'][0]['inputs']['pipe_labour_hours'], .1234567890123)
        self.assertEqual(stored['composer']['globals']['register_allowance_hours'], 0)
        self.assertEqual(loaded['penetration'], stored)
        self.assertEqual(stored['source_sha256'], model['source']['sha256'])
        self.assertEqual(source_model(), model)
        self.assertEqual(sha256((ROOT / 'data/penetration.json.gz').read_bytes()).hexdigest(), source_hash)
        result = calculate(manual)
        self.assertEqual(input_records(build_penetration_register(result, result['definition'], {}))['Pipe Labour'], (.1234567890123, 'Manual'))
        self.assertEqual(self.store.list_quotes(), [])

    def test_unavailable_auto_and_unselected_pipe_are_distinguished_in_exports(self):
        result = calculate(draft(AL=301))
        self.assertEqual(result['summary']['grand_total'], '#VALUE!')
        records = input_records(build_penetration_register(result, result['definition'], {}))
        self.assertEqual(records['Pipe Labour'], ('Unavailable', 'Manual value required'))
        text = pdf_text(render_penetration_pdf(result, result['definition'], {}))
        self.assertNotIn('Unavailable (Manual value required)', text)
        self.assertIn('Enter Pipe Labour hours', text)
        result = calculate(draft(Y=None, AL=None))
        records = input_records(build_penetration_register(result, result['definition'], {}))
        self.assertEqual(records['Pipe Labour'], ('Not applicable', 'Automatic; no collar selected'))
        self.assertEqual(result['errors'], [])
        result = calculate(draft(Y=None, AL=None, pipe_labour_hours=0))
        self.assertNotIn('pipe_labour_hours', result['rows'][0]['inputs'])
        self.assertEqual(input_records(build_penetration_register(result, result['definition'], {}))['Pipe Labour'],
                         ('Not applicable', 'Automatic; no collar selected'))

    def test_main_quote_costs_hours_days_and_material_projection_reconcile_once(self):
        request = {'title': 'Combined labour', 'inputs': {'B8': 1}, 'penetration': {'draft': draft()}}
        quote = self.store.save_quote(request)
        result = quote['result']
        fire = result['firestopping']['result']
        row = fire['rows'][0]
        self.assertEqual(row['outputs']['DF'], .5)  # .25 per pipe × AN2.
        self.assertEqual(fire['summary']['labour_hours'], 3.75)  # (.25 register + .5 pipes + .5 additional) × O3.
        hourly = row['outputs']['CW']
        self.assertEqual(fire['summary']['labour'], 3.75 * hourly + 7)
        tasks = {item['label']: item for item in fire['schedule_breakdown']['rows']}
        self.assertEqual(len(tasks), 8)
        self.assertEqual(tasks['Additional Labour']['task_hours'], 1.5)
        self.assertEqual(tasks['Additional Labour']['labour_costs'], 1.5 * hourly + 7)
        self.assertEqual(tasks['Register allowance']['task_hours'], .75)
        self.assertEqual(tasks['Other']['task_hours'], 0)
        self.assertEqual(tasks['Other']['labour_costs'], 0)
        self.assertAlmostEqual(sum(item['task_hours'] for item in tasks.values()), fire['summary']['labour_hours'])
        self.assertAlmostEqual(sum(item['labour_costs'] for item in tasks.values()), fire['summary']['labour'])
        collars = next(item for item in fire['material_breakdown'] if item.get('quantity_column') == 'AN')
        self.assertEqual(collars['quantity'], 6)
        self.assertEqual(collars['days'], 1.5 / 8)
        self.assertEqual(result['summary']['total'], result['base_summary']['total'] + fire['summary']['grand_total'])
        self.assertEqual(result['summary']['days'], result['base_summary']['days'] + 3.75 / 8)
        text = pdf_text(render_quote_pdf(quote))
        for label in ('Additional Labour', 'Register allowance', f"${result['summary']['total']:,.2f}"):
            self.assertIn(label, text)

        # The hourly price already contains the selected worker count. Doubling
        # the crew changes cost, never register/pipe task hours or labour days.
        teams = next(field['options'] for field in definition({})['row_fields'] if field['column'] == 'W')
        two_workers = calculate(draft(W=teams[3]))
        self.assertEqual(two_workers['summary']['labour_hours'], 3.75)
        self.assertEqual(two_workers['summary']['total_days'], 3.75 / 8)
        self.assertEqual(two_workers['rows'][0]['outputs']['CW'], hourly * 2)
        self.assertEqual(two_workers['summary']['labour'], 3.75 * hourly * 2 + 7)

    def test_saved_quote_report_remains_frozen_until_explicit_recalculation(self):
        quote = self.store.save_quote({'title': 'Saved automatically resolved labour', 'inputs': {'B8': 1}, 'penetration': {'draft': draft()}})
        with patch('estimator.penetration_calculator.resolve_labour', side_effect=AssertionError('Snapshot recalculated')):
            stored = Store(self.store.path).quote(quote['id'])
            render_quote_pdf(stored)
        self.assertEqual(stored, quote)
        # An older snapshot can retain its original seven task labels. Export
        # copy must describe those stored labels without resolving new policy.
        historical = deepcopy(stored)
        tasks = historical['result']['labour']['firestopping_tasks']
        additional = next(task for task in tasks if task['name'].endswith('Additional Labour'))
        register = next(task for task in tasks if task['name'].endswith('Register allowance'))
        other = next(task for task in tasks if task['name'].endswith('Other'))
        register['name'] = 'Firestopping · Labour'
        for key in ('task_hours', 'days', 'total'):
            other[key] += additional[key]
        tasks.remove(additional)
        with patch('estimator.penetration_calculator.resolve_labour', side_effect=AssertionError('Historical snapshot recalculated')):
            text = pdf_text(render_quote_pdf(historical))
        self.assertIn('Firestopping Labour includes source setup time', text)
        self.assertNotIn('Register allowance shows', text)
        later = deepcopy(quote['penetration'])
        later['draft']['rows'][0]['inputs']['AL'] = 300
        preview = self.store.prepare_quote({'penetration': later}, quote['id'])
        self.assertNotEqual(preview['result']['summary']['total'], quote['result']['summary']['total'])
        self.assertEqual(self.store.quote(quote['id']), quote)
        saved = self.store.save_quote({'penetration': later}, quote['id'])
        self.assertEqual(saved['result'], preview['result'])
        self.assertNotIn('pipe_labour_hours', saved['penetration']['draft']['rows'][0]['inputs'])

    def test_existing_library_resolves_defaults_without_rewriting_source_or_saved_edits(self):
        from test_firestopping_library import editable_library
        library_path = self.root / 'library'
        editable_library(library_path)
        source_bytes = (library_path / 'library.json').read_bytes()
        library = FirestoppingLibrary(library_path, self.store)
        with self.store.connect() as db:
            schema = db.execute('SELECT name,sql FROM sqlite_master WHERE type="table" ORDER BY name').fetchall()
            version = db.execute('PRAGMA user_version').fetchone()[0]
        edit = library.edit('pkb-001')
        self.assertEqual(edit['result']['rows'][0]['input_defaults']['register_allowance_hours'], .25)
        self.assertNotIn('register_allowance_hours', edit['draft']['rows'][0]['inputs'])
        self.assertNotIn('pipe_labour_hours', edit['draft']['rows'][0]['inputs'])
        self.assertEqual(library.edits.stamp(), (0, 0))
        self.assertEqual(edit['source_price']['amount'], 150)
        self.assertEqual((library_path / 'library.json').read_bytes(), source_bytes)
        with self.store.connect() as db:
            self.assertEqual(db.execute('SELECT name,sql FROM sqlite_master WHERE type="table" ORDER BY name').fetchall(), schema)
            self.assertEqual(db.execute('PRAGMA user_version').fetchone()[0], version)


if __name__ == '__main__':
    unittest.main()
