"""Separate composer persistence and source-based whole-schedule presentation."""

from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from pypdf import PdfReader

from estimator.catalog import ValidationError
from estimator.excel_engine import FormulaError, column_number
from estimator.penetration_breakdown import schedule_breakdown
from estimator.penetration_calculator import (
    PenetrationEngine, SUMMARY_COLUMNS, calculate, definition, engine_for_draft,
    normalize_composer, normalize_draft, source_model,
)
from estimator.penetration_report import build_penetration_register, render_penetration_pdf
from estimator.project_file import export_project, load_project_bytes
from estimator.storage import Store


class PenetrationScheduleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary = tempfile.TemporaryDirectory()
        cls.store = Store(Path(cls.temporary.name) / 'projects.sqlite3')
        cls.spec = definition({})

    @classmethod
    def tearDownClass(cls):
        cls.temporary.cleanup()

    def options(self, column):
        return next(field['options'] for field in self.spec['row_fields'] if field['column'] == column)

    def mixed_schedule(self):
        return {'globals': {'J': 'Yes', 'K': 2.1234567890123, 'L': .15, 'M': .2}, 'rows': [
            {'id': f'item-{index}', 'inputs': {
                'O': index + 2, 'T': f'Service Ø65 – {index}', 'W': self.options('W')[index],
                'X': self.options('X')[index], 'AA': self.options('AA')[index],
                'AW': 600, 'AX': 450, 'AY': 1, 'AZ': .1,
                'AL': 65, 'AM': 300, 'AN': 1, 'AO': .15,
                'P': 'Concrete wall', 'Q': 'Standard', 'R': 'Standard',
            }} for index in range(2)]}

    def test_empty_schedule_keeps_globals_but_has_no_template_or_fixed_charges(self):
        draft = {'globals': {'J': 'Yes', 'K': 7.25, 'L': .35, 'M': .45}, 'rows': []}
        before_source = deepcopy(source_model())
        result = calculate(draft)
        self.assertEqual(result['draft'], draft)
        self.assertEqual(result['rows'], [])
        self.assertEqual(result['errors'], [])
        self.assertTrue(all(value == 0 for value in result['summary'].values()))
        self.assertEqual(result['summary_cells'], dict.fromkeys(SUMMARY_COLUMNS, 0))
        self.assertTrue(all(value == 0 for value in result['breakdown_cells'].values()))
        table = result['schedule_breakdown']
        self.assertEqual(table['totals'], {'material_costs': 0, 'labour_costs': 0, 'task_hours': 0})
        for row in table['rows']:
            self.assertEqual(row['unit_prices'], [])
            self.assertEqual(row['material_quantities'], [])
            self.assertEqual([row[key] for key in ('material_costs', 'labour_costs', 'task_hours')], [0, 0, 0])
        self.assertTrue(all(not group['rows'] for group in table['source_groups']))
        engine, normalized = engine_for_draft(draft)
        self.assertEqual(normalized['rows'], [])
        self.assertEqual(engine.value('CALC', 'D2'), 0)
        self.assertIsNone(engine.value('CALC', 'T4'))
        self.assertEqual(source_model(), before_source)

    def test_empty_exports_have_zero_summary_and_no_invented_line(self):
        result = calculate({'globals': {'J': 'Yes', 'K': 10}, 'rows': []})
        workbook = load_workbook(BytesIO(build_penetration_register(result, result['definition'], {})))
        self.assertFalse(any(row[0].value for row in workbook['Schedule'].iter_rows(min_row=5)))
        summary_values = {row[0].value: row[1].value for row in workbook['Summary'].iter_rows()}
        for name in ('Labour', 'Materials', 'Grand total'):
            self.assertEqual(summary_values[name], 0)
        pdf = PdfReader(BytesIO(render_penetration_pdf(result, result['definition'], {})))
        text = '\n'.join(page.extract_text() for page in pdf.pages)
        self.assertIn('Schedule', text)
        self.assertNotIn('Line 1', text)
        self.assertNotIn('Pair Coil Bundle', text)

    def test_composer_and_schedule_defaults_are_independent(self):
        spec = definition({})
        self.assertEqual(spec['schedule_defaults']['rows'], [])
        self.assertEqual(len(spec['defaults']['rows']), 1)
        spec['schedule_defaults']['globals']['K'] = 3
        self.assertIsNone(spec['defaults']['globals']['K'])
        self.assertEqual(len(normalize_draft(None)['rows']), 1)  # historical API default
        for invalid in (None, {'rows': []}, {'rows': [{'id': 'a'}, {'id': 'b'}]}):
            with self.subTest(invalid=invalid), self.assertRaises(ValidationError):
                normalize_composer(invalid)

    def test_project_round_trip_preserves_independent_drafts_and_frozen_configuration(self):
        composer = deepcopy(self.spec['defaults'])
        composer['globals'].update(J='Yes', K=9.876543210987, L=.333, M=.777)
        composer['rows'][0]['inputs'].update(T='Ø65 未保存', O=0, AL=None, AI=12.345678901234,
                                           W='Historical labour selection')
        for draft in (self.spec['schedule_defaults'], self.mixed_schedule()):
            with self.subTest(rows=len(draft['rows'])):
                request = {'estimate': {'configuration': {'inventory': {'0': {'sales_price': 123.456789}}}},
                           'penetration': {'draft': deepcopy(draft), 'composer': deepcopy(composer)}}
                original = deepcopy(request)
                payload = export_project(self.store, request)
                loaded = load_project_bytes(self.store, payload)
                self.assertEqual(loaded['penetration'], {'source_sha256': self.spec['source_sha256'],
                                 'draft': draft, 'composer': composer})
                self.assertEqual(request, original)
                before = calculate(draft, loaded['estimate']['configuration'])
                self.store.save_configuration({'inventory': {'0': {'sales_price': 876.54321}}})
                again = load_project_bytes(self.store, payload)
                self.assertEqual(again['estimate']['configuration'], loaded['estimate']['configuration'])
                self.assertEqual(calculate(again['penetration']['draft'], again['estimate']['configuration']), before)

    def test_legacy_snapshot_keeps_schedule_without_making_it_a_composer(self):
        draft = self.mixed_schedule()
        payload = export_project(self.store, {'estimate': {}, 'penetration': {'draft': draft}})
        self.assertEqual(load_project_bytes(self.store, payload)['penetration'],
                         {'source_sha256': self.spec['source_sha256'], 'draft': draft})

    def test_project_rejects_invalid_composer_and_source_version(self):
        valid = json.loads(export_project(self.store, {'estimate': {}, 'penetration': {
            'draft': self.spec['schedule_defaults'], 'composer': self.spec['defaults']}}))
        for replacement in (None, {'rows': []}, {'rows': [{'id': 'a'}, {'id': 'b'}]},
                            {'rows': [{'id': 'a', 'inputs': {'H': 5}}]},
                            {'rows': [{'id': 'a', 'inputs': {'AL': '65'}}]}):
            invalid = deepcopy(valid)
            invalid['penetration']['composer'] = replacement
            with self.subTest(composer=replacement), self.assertRaises(ValidationError):
                load_project_bytes(self.store, json.dumps(invalid).encode())
        invalid = deepcopy(valid)
        invalid['penetration']['source_sha256'] = '0' * 64
        with self.assertRaisesRegex(ValidationError, 'different Firestopping'):
            load_project_bytes(self.store, json.dumps(invalid).encode())

    def test_mixed_products_keep_distinct_prices_quantities_and_per_line_multipliers(self):
        draft = self.mixed_schedule()
        before = deepcopy(draft)
        result = calculate(draft)
        self.assertEqual(result['errors'], [])
        table = result['schedule_breakdown']
        self.assertEqual(table['totals'], {'material_costs': result['summary']['materials'],
                         'labour_costs': result['summary']['labour'], 'task_hours': result['summary']['labour_hours']})
        by_task = {row['label']: row for row in table['rows']}
        board = by_task['Board']
        self.assertEqual([item['product'] for item in board['unit_prices']], self.options('X')[:2])
        self.assertEqual([item['value'] for item in board['unit_prices']],
                         [row['outputs']['CX'] for row in result['rows']])
        self.assertEqual([item['row_ids'] for item in board['unit_prices']], [['item-0'], ['item-1']])
        for item in board['material_quantities']:
            self.assertEqual(len(item['row_ids']), 1)
            row = next(row for row in result['rows'] if row['id'] == item['row_ids'][0])
            self.assertEqual(item['value'], row['outputs'][item['column']] * row['inputs']['O'])
            self.assertEqual(item['product'], row['inputs']['X'])
            self.assertIn(item['product'], item['label'])
        for key in ('material_costs', 'labour_costs', 'task_hours'):
            self.assertAlmostEqual(sum(row[key] for row in table['rows']), table['totals'][key], places=9)
        for group in table['source_groups']:
            self.assertEqual([row['label'] for row in group['rows']], ['Line 1', 'Line 2'])
            for line in group['rows']:
                row = next(row for row in result['rows'] if row['id'] == line['row_id'])
                self.assertTrue(all(value['value'] == row['outputs'][value['column']] for value in line['values']))
        no_travel = deepcopy(draft)
        no_travel['globals']['K'] = None
        without = calculate(no_travel)
        source_lists = next(sheet for sheet in source_model()['sheets'] if sheet['name'] == 'LISTS')['cells']
        expected_travel = source_lists['BZ2']['value'] * draft['globals']['K']
        self.assertEqual(result['summary'], without['summary'])
        self.assertIn(result['summary']['travel_lafha'], ('', 0))
        self.assertEqual(draft, before)

    def test_matching_products_deduplicate_prices_and_total_mastic_and_other_quantities(self):
        draft = self.mixed_schedule()
        draft['rows'][0]['inputs'].update(AB=self.options('AB')[0], AC=.5, O=3)
        draft['rows'][1]['inputs'] = deepcopy(draft['rows'][0]['inputs'])
        draft['rows'][1]['inputs'].update(AC=.75, O=4)
        original = deepcopy(draft)
        result = calculate(draft)
        self.assertEqual(result['errors'], [])
        tasks = {row['label']: row for row in result['schedule_breakdown']['rows']}
        for task in ('Labour', 'Board', 'Wrap', 'Mastic'):
            self.assertEqual(len(tasks[task]['unit_prices']), 1)
            self.assertEqual(tasks[task]['unit_prices'][0]['row_ids'], ['item-0', 'item-1'])
            self.assertNotIn('Line ', tasks[task]['unit_prices'][0]['label'])
        mastic = tasks['Mastic']['material_quantities']
        self.assertEqual(len(mastic), 1)
        self.assertEqual(mastic[0]['column'], 'AC')
        self.assertEqual(mastic[0]['value'], .5 * 3 + .75 * 4)
        self.assertEqual(mastic[0]['product'], self.options('AB')[0])
        self.assertEqual(mastic[0]['row_ids'], ['item-0', 'item-1'])
        for task in ('Board', 'Wrap'):
            for entry in tasks[task]['material_quantities']:
                expected = sum(row['outputs'][entry['column']] * row['inputs']['O'] for row in result['rows'])
                self.assertEqual(entry['value'], expected)
                self.assertEqual(entry['row_ids'], ['item-0', 'item-1'])
        for row, expected in zip(result['rows'], (1.5, 3)):
            item = next(task for task in row['breakdown']['rows'] if task['label'] == 'Mastic')
            self.assertEqual(item['material_quantities'][0]['value'], expected)
        self.assertEqual(draft, original)

    def test_grouping_keeps_distinct_rates_units_contexts_and_unselected_products(self):
        result = calculate(self.mixed_schedule())
        first, second = result['rows']
        second['inputs']['X'] = first['inputs']['X']
        board1 = next(task for task in first['breakdown']['rows'] if task['label'] == 'Board')
        board2 = next(task for task in second['breakdown']['rows'] if task['label'] == 'Board')
        board1['unit_prices'][0]['value'] = 1.23456789012345
        board2['unit_prices'][0]['value'] = 1.23456789012346
        board1['material_quantities'] = [{'column': 'CJ', 'label': 'Substrate', 'value': .123456789012345, 'units': 'sheets', 'format': 'number'},
                                          {'column': 'CQ', 'label': 'Bulkhead', 'value': 2, 'units': 'sheets', 'format': 'number'}]
        board2['material_quantities'] = [{'column': 'CJ', 'label': 'Substrate', 'value': .234567890123456, 'units': 'sheets', 'format': 'number'},
                                          {'column': 'CJ', 'label': 'Substrate', 'value': 4, 'units': 'm²', 'format': 'number'}]
        original = deepcopy(result)
        board = next(task for task in schedule_breakdown(result)['rows'] if task['label'] == 'Board')
        self.assertEqual([item['value'] for item in board['unit_prices']], [1.23456789012345, 1.23456789012346])
        self.assertEqual([item['value'] for item in board['material_quantities']],
                         [.123456789012345 + .234567890123456, 2, 4])
        self.assertEqual(board['material_quantities'][0]['row_ids'], ['item-0', 'item-1'])
        self.assertEqual(result, original)
        # Without a product identity, coincidentally equal values cannot be combined.
        for row in result['rows']:
            row['inputs']['X'] = None
        board = next(task for task in schedule_breakdown(result)['rows'] if task['label'] == 'Board')
        self.assertEqual(len(board['material_quantities']), 4)
        self.assertTrue(all(len(item['row_ids']) == 1 for item in board['material_quantities']))

    def test_grouped_mastic_keeps_blank_zero_negative_and_error_distinctions(self):
        result = calculate(self.mixed_schedule())
        for row in result['rows']:
            row['inputs']['AB'] = self.options('AB')[0]
            mastic = next(task for task in row['breakdown']['rows'] if task['label'] == 'Mastic')
            mastic['unit_prices'][0]['value'] = 0
        for values, expected in (((None, ''), None), ((0, None), 0), ((2, -2), 0),
                                 ((2, '#DIV/0!'), '#DIV/0!'), ((float('inf'), 2), '#NUM!')):
            with self.subTest(values=values):
                for row, value in zip(result['rows'], values):
                    next(task for task in row['breakdown']['rows'] if task['label'] == 'Mastic')['material_quantities'][0]['value'] = value
                before = deepcopy(result)
                mastic = next(task for task in schedule_breakdown(result)['rows'] if task['label'] == 'Mastic')
                self.assertEqual(len(mastic['unit_prices']), 1)
                self.assertEqual(mastic['unit_prices'][0]['value'], 0)
                if expected is None:
                    self.assertEqual(mastic['material_quantities'], [])
                else:
                    self.assertEqual(mastic['material_quantities'][0]['value'], expected)
                self.assertEqual(result, before)

    def test_schedule_projection_retains_zero_and_errors_without_mutating_calculation(self):
        result = calculate(self.mixed_schedule())
        row = result['rows'][0]
        board = next(task for task in row['breakdown']['rows'] if task['label'] == 'Board')
        board['unit_prices'][0]['value'] = 0
        board['material_quantities'][0]['value'] = '#DIV/0!'
        board['material_quantities'][1]['value'] = ''
        row['outputs']['BI'] = 0
        before = deepcopy(result)
        table = schedule_breakdown(result)
        self.assertEqual(next(task for task in table['rows'] if task['label'] == 'Board')['unit_prices'][0]['value'], 0)
        quantities = next(task for task in table['rows'] if task['label'] == 'Board')['material_quantities']
        self.assertEqual(quantities[0]['value'], '#DIV/0!')
        self.assertFalse(any(value['value'] == '' for value in quantities))
        self.assertEqual([group['label'] for group in table['source_groups']], ['Summary'])
        self.assertEqual(result, before)

    def test_calculation_error_remains_visible_in_aggregate_and_canonical_footer(self):
        original = PenetrationEngine.cell
        def failed_material(engine, sheet, row, column):
            if sheet == 'CALC' and row == 5 and column == column_number('DM'):
                raise FormulaError('#DIV/0!')
            return original(engine, sheet, row, column)
        with patch.object(PenetrationEngine, 'cell', failed_material):
            result = calculate(self.mixed_schedule())
        table = result['schedule_breakdown']
        self.assertEqual(next(row for row in table['rows'] if row['label'] == 'Board')['material_costs'], '#DIV/0!')
        self.assertEqual(table['totals']['material_costs'], '#DIV/0!')
        self.assertEqual(result['summary']['grand_total'], '#DIV/0!')
        self.assertTrue(any(error['row_id'] == 'item-1' and error['cell'] == 'DM5' for error in result['errors']))


if __name__ == '__main__':
    unittest.main()
