"""Manufacturer metadata is portable and cannot alter pricing or quantities."""
from copy import deepcopy
import ast
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from estimator.calculator import calculate, fields
from estimator.catalog import (baseline, catalog_signature, effective_catalog,
    validate_catalog, validate_configuration, ValidationError)
from estimator.pricing_manufacturers import with_manufacturer_defaults, load_manufacturer_seed
from estimator.pricing_workbook import (export_pricing_workbook, import_pricing_workbook,
    COMBINED_SHEET, PRE_MANUFACTURER_PRODUCT_SERVICE_HEADERS, _serialize_exact)
from estimator.project_file import export_project, load_project_bytes
from estimator.storage import Store


def remove_manufacturer_metadata(value):
    if isinstance(value, dict):
        return {key: remove_manufacturer_metadata(item) for key, item in value.items()
                if key not in {'manufacturer', 'manufacturer_assignment'}}
    if isinstance(value, list):
        return [remove_manufacturer_metadata(item) for item in value]
    return value


class PricingManufacturerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = baseline()
        cls.product = next(item for item in cls.data['inventory'] if item['id'] == '204')
        cls.linked = next(rate for rows in cls.data['rate_groups'].values() for rate in rows if rate['inventory_id'] == '204')

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.seed = self.root / 'reviewed.json'
        self.write_seed()

    def write_seed(self, **changes):
        assignment = {'kind': 'inventory', 'id': self.product['id'], 'name': self.product['name'],
                      'manufacturer': 'Reviewed manufacturer',
                      'evidence': [{'url': 'https://example.com/source-review', 'title': 'Synthetic test evidence',
                                    'note': 'Fixture only; this is not a real product assignment.'}]}
        assignment.update(changes)
        self.seed.write_text(json.dumps({'schema_version': 1, 'reviewed_at': '2026-10-30', 'assignments': [assignment]}), encoding='utf-8')

    def enriched(self, configuration=None):
        return with_manufacturer_defaults({'inventory': {}, 'rates': {}} if configuration is None else configuration, seed_path=self.seed)

    def test_seed_exact_binding_metadata_traceability_and_complete_numeric_invariance(self):
        original = {'inventory': {}, 'rates': {}}
        before = deepcopy(original)
        enriched = self.enriched(original)
        self.assertEqual(original, before)
        effective = effective_catalog(enriched)
        product = next(item for item in effective['inventory'] if item['id'] == '204')
        self.assertEqual(product['manufacturer'], 'Reviewed manufacturer')
        seed, digest = load_manufacturer_seed(self.seed)
        proof = product['source']['manufacturer_assignment']
        self.assertEqual(proof['seed_sha256'], digest)
        self.assertEqual(proof['value'], product['manufacturer'])
        self.assertEqual(proof['binding']['name'], self.product['name'])
        self.assertEqual(proof['evidence'], seed['assignments'][0]['evidence'])
        # Embedded catalogs use the validator's existing explicit rule flags;
        # compare that canonical structure as well as the legacy implicit
        # baseline's independently calculated cells and quantities below.
        canonical_before = effective_catalog({'catalog': self.data, **original})
        self.assertEqual(remove_manufacturer_metadata(effective), canonical_before)
        self.assertEqual(catalog_signature(effective), catalog_signature(effective_catalog(original)))
        signed = {**original, 'catalog_signature': catalog_signature(self.data)}
        self.assertEqual(self.enriched(signed)['catalog_signature'], signed['catalog_signature'])
        teams = {field['cell']: next(option for option in field['options'] if option and option != 'N/A')
                 for field in fields(self.data) if field['cell'] in {'D2','D3','D4','D5','D6','D7','D8','D9','D10'}}
        for inputs in ({}, {**teams, 'D15': 'Promat Cafco 300', 'B15': 234.567890123, 'B16': 12.345678901},
                       {**teams, 'D15': 'Promat Fendolite MII', 'B15': 999.123456789, 'B23': 23.456789012}):
            self.assertEqual(calculate(inputs, enriched), calculate(inputs, original))
        self.assertEqual(self.enriched(enriched), enriched)

    def test_manual_values_explicit_clears_and_renamed_custom_records_win(self):
        for value in ('Manual manufacturer', ''):
            configuration = {'inventory': {'204': {'manufacturer': value}}, 'rates': {}}
            self.assertEqual(self.enriched(configuration), configuration)
            catalog = deepcopy(self.data)
            next(item for item in catalog['inventory'] if item['id'] == '204')['manufacturer'] = value
            configuration = {'catalog': catalog, 'inventory': {}, 'rates': {}}
            self.assertEqual(self.enriched(configuration), configuration)
        for field, value in (('name', 'Custom replacement product'), ('product_service', 'Custom display product')):
            configuration = {'inventory': {'204': {field: value}}, 'rates': {}}
            self.assertEqual(self.enriched(configuration), configuration)
        self.write_seed(name='Different source name')
        self.assertEqual(self.enriched(), {'inventory': {}, 'rates': {}})
        self.write_seed(id='custom-item-not-present')
        self.assertEqual(self.enriched(), {'inventory': {}, 'rates': {}})

    def test_manufacturer_schema_is_bounded_optional_text_and_linked_uses_follow_owner(self):
        for value in ('', 'Promat 日本', 'x' * 200):
            configuration = validate_configuration({'inventory': {'204': {'manufacturer': value}}})
            effective = effective_catalog(configuration)
            rate = next(item for rows in effective['rate_groups'].values() for item in rows if item['id'] == self.linked['id'])
            self.assertEqual(rate['manufacturer'], value)
        for value in (None, True, [], {}, 'x' * 201, 'bad\x00text', '\ud800'):
            with self.subTest(value=repr(value)), self.assertRaises(ValidationError):
                validate_configuration({'inventory': {'204': {'manufacturer': value}}})
        catalog = deepcopy(self.data)
        rate = next(item for rows in catalog['rate_groups'].values() for item in rows if item['id'] == self.linked['id'])
        rate['manufacturer'] = 'Earlier inherited workbook value'
        configuration = {'catalog': catalog, 'inventory': {'204': {'manufacturer': 'New owner value'}}, 'rates': {}}
        self.assertEqual(next(item for rows in effective_catalog(configuration)['rate_groups'].values() for item in rows if item['id'] == self.linked['id'])['manufacturer'], 'New owner value')

    def test_shared_defaults_do_not_write_settings_or_backfill_frozen_quote_and_project(self):
        store = Store(self.root / 'state.sqlite3')
        raw = store.save_configuration({'inventory': {}, 'rates': {}})
        old = store.save_quote({'title': 'Frozen older project', 'inputs': {}, 'configuration': raw})
        with store.connect() as db:
            stored_before = db.execute('SELECT data FROM quotes WHERE id=?', (old['id'],)).fetchone()[0]
            settings_before = db.execute('SELECT data FROM settings WHERE id=1').fetchone()[0]
        with patch('estimator.pricing_manufacturers.SEED_PATH', self.seed):
            shared = store.shared_configuration()
            self.assertEqual(store.configuration(), raw)
            new = store.prepare_quote({'title': 'New project', 'inputs': {}})
            self.assertIn('manufacturer', next(item for item in effective_catalog(new['configuration'])['inventory'] if item['id'] == '204'))
            self.assertNotIn('manufacturer', next(item for item in effective_catalog(store.quote(old['id'])['configuration'])['inventory'] if item['id'] == '204'))
            store.save_configuration(shared)
            restarted = Store(store.path)
            self.assertEqual(restarted.shared_configuration(), shared)
            cleared = deepcopy(shared)
            cleared['inventory']['204'] = {'manufacturer': ''}
            restarted.save_configuration(cleared)
            self.assertEqual(restarted.shared_configuration(), cleared)
        with store.connect() as db:
            self.assertEqual(db.execute('SELECT data FROM quotes WHERE id=?', (old['id'],)).fetchone()[0], stored_before)
        # Explicit frozen configurations must not even consult new seed data.
        with patch('estimator.pricing_manufacturers.with_manufacturer_defaults', side_effect=AssertionError('Frozen data consulted shared defaults')):
            reopened = store.prepare_quote({'configuration': old['configuration'], 'inputs': {}}, old['id'])
            self.assertEqual(reopened['configuration'], old['configuration'])
            payload = export_project(store, {'estimate': {'title': 'Portable frozen', 'configuration': old['configuration']}})
            restored = load_project_bytes(store, payload)
            self.assertEqual(restored['estimate']['configuration'], old['configuration'])
        self.assertTrue(settings_before)

    def test_standalone_manufacturer_and_blank_survive_quote_freezing(self):
        catalog = deepcopy(self.data)
        standalone = deepcopy(catalog['rate_groups']['labour_rates'][0])
        standalone.update(id='standalone-manufacturer-test', name='Standalone metadata test', inventory_id=None, price_mode='override')
        catalog['rate_groups']['labour_rates'].append(standalone)
        store = Store(self.root / 'standalone.sqlite3')
        for value in ('Entered manufacturer', ''):
            configuration = {'catalog': catalog, 'inventory': {}, 'rates': {standalone['id']: {'manufacturer': value}}}
            saved = store.save_quote({'title': 'Standalone', 'inputs': {}, 'configuration': configuration})
            self.assertEqual(saved['configuration']['rates'][standalone['id']]['manufacturer'], value)
            self.assertEqual(store.quote(saved['id'])['configuration'], saved['configuration'])

    def test_existing_custom_linked_rate_and_its_yield_survive_seed_and_workbook_round_trip(self):
        catalog = deepcopy(self.data)
        custom = deepcopy(catalog['rate_groups']['boards'][0])
        identity = 'rate_6052d7e6-d7a2-4e60-9db2-1e54a7cf2ba4'
        custom.update({'id': identity, 'name': 'Promat SYSTEMPANELTM', 'inventory_id': '260',
                       'price_mode': 'inventory', 'yield': 1.62, 'source': {'origin': 'user'}})
        catalog['rate_groups']['boards'].append(custom)
        configuration = {'catalog': catalog, 'inventory': {}, 'rates': {identity: {'yield': 1.62}}}
        enriched = with_manufacturer_defaults(configuration)
        before, after = effective_catalog(configuration), effective_catalog(enriched)
        self.assertEqual(remove_manufacturer_metadata(after), before)
        result = import_pricing_workbook(export_pricing_workbook(enriched), 'custom-rate.xlsx', enriched)
        rates = [item for rows in effective_catalog(result['configuration'])['rate_groups'].values() for item in rows]
        self.assertEqual(len(rates), 167)
        retained = next(item for item in rates if item['id'] == identity)
        self.assertEqual(retained['yield'], 1.62)
        self.assertEqual(retained['inventory_id'], '260')
        self.assertEqual(retained['name'], custom['name'])
        self.assertEqual(result['summary']['rates']['updated'], 0)

    def test_excel_manufacturer_edits_clears_and_previous_layout_are_lossless(self):
        configuration = {'inventory': {'204': {'manufacturer': 'Reviewed manufacturer'}}, 'rates': {}}
        payload = export_pricing_workbook(configuration)
        unchanged = import_pricing_workbook(payload, 'unchanged.xlsx', configuration)
        self.assertEqual(unchanged['summary']['inventory']['updated'], 0)
        self.assertEqual(unchanged['summary']['rates']['updated'], 0)
        book = load_workbook(BytesIO(payload))
        try:
            sheet = book[COMBINED_SHEET]
            columns = {cell.value: cell.column for cell in sheet[1]}
            number = next(row[0].row for row in sheet.iter_rows(min_row=2) if sheet.cell(row[0].row, columns['Inventory ID']).value == '204')
            sheet.cell(number, columns['Manufacturer']).value = None
            cleared = import_pricing_workbook(_serialize_exact(book, escape_text=True), 'cleared.xlsx', configuration)['configuration']
            self.assertEqual(next(item for item in effective_catalog(cleared)['inventory'] if item['id'] == '204')['manufacturer'], '')
            self.assertEqual(self.enriched(cleared), cleared)
            sheet.delete_cols(columns['Manufacturer'])
            self.assertEqual(tuple(cell.value for cell in sheet[1]), PRE_MANUFACTURER_PRODUCT_SERVICE_HEADERS)
            legacy = import_pricing_workbook(_serialize_exact(book, escape_text=True), 'older-layout.xlsx', configuration)
            self.assertEqual(legacy['summary']['inventory']['updated'], 0)
            self.assertEqual(legacy['summary']['rates']['updated'], 0)
        finally:
            book.close()
        self.assertEqual(calculate({}, configuration), calculate({}, unchanged['configuration']))

    def test_desktop_resource_manifest_keeps_reviewed_seed_and_missing_resource_fails_clearly(self):
        source = Path(__file__).resolve().parents[1] / 'scripts' / 'build_windows.py'
        tree = ast.parse(source.read_text(encoding='utf-8'))
        data_files = next(ast.literal_eval(node.value) for node in tree.body if isinstance(node, ast.Assign)
                          and any(isinstance(target, ast.Name) and target.id == 'DATA_FILES' for target in node.targets))
        self.assertEqual(data_files.count('pricing_manufacturers.json'), 1)
        # Full and standard builds use the same explicit DATA_FILES list; the
        # frozen resource loader must read the packaged bytes, not an optional
        # external manufacturer's directory or a silent empty fallback.
        frozen = self.root / '_internal' / 'data'
        frozen.mkdir(parents=True)
        (frozen / 'pricing_manufacturers.json').write_bytes(self.seed.read_bytes())
        reviewed, digest = load_manufacturer_seed(frozen / 'pricing_manufacturers.json')
        self.assertEqual(reviewed['assignments'][0]['id'], '204')
        self.assertEqual(len(digest), 64)
        with self.assertRaisesRegex(ValidationError, 'manufacturer defaults could not be read'):
            load_manufacturer_seed(frozen / 'missing.json')

    def test_truly_older_templates_preserve_newer_manual_manufacturer_but_new_blank_column_clears(self):
        # The historical hidden snapshot has no manufacturer field at all.
        # It must not erase metadata the user entered after this file existed.
        older = export_pricing_workbook({})
        for value in ('Newer manual manufacturer', ''):
            current = {'inventory': {'204': {'manufacturer': value}}, 'rates': {}}
            for previous_columns in (8, 9):
                with self.subTest(value=value, columns=previous_columns):
                    book = load_workbook(BytesIO(older))
                    try:
                        sheet = book[COMBINED_SHEET]
                        sheet.delete_cols(10)
                        if previous_columns == 8:
                            sheet.delete_cols(7)
                        historical = _serialize_exact(book, escape_text=True)
                    finally:
                        book.close()
                    result = import_pricing_workbook(historical, 'pre-manufacturer.xlsx', current)
                    retained = next(item for item in effective_catalog(result['configuration'])['inventory'] if item['id'] == '204')
                    self.assertEqual(retained['manufacturer'], value)
                    self.assertEqual(result['summary']['inventory']['updated'], 0)
            # The new format does have a visible Manufacturer field. Its blank
            # clears newer metadata even if its hidden snapshot was older.
            new_format = import_pricing_workbook(older, 'new-visible-column.xlsx', current)
            retained = next(item for item in effective_catalog(new_format['configuration'])['inventory'] if item['id'] == '204')
            self.assertEqual(retained['manufacturer'], '')

    def test_old_standalone_service_template_preserves_current_manufacturer_and_new_column_can_clear(self):
        catalog = deepcopy(self.data)
        standalone = deepcopy(catalog['rate_groups']['labour_rates'][0])
        identity = 'old-template-standalone'
        standalone.update(id=identity, name='Independent legacy service', inventory_id=None, price_mode='override')
        catalog['rate_groups']['labour_rates'].append(standalone)
        historical = {'catalog': catalog, 'inventory': {}, 'rates': {}}
        current = {**historical, 'rates': {identity: {'manufacturer': 'Current service metadata'}}}
        exported = export_pricing_workbook(historical)
        book = load_workbook(BytesIO(exported))
        try:
            book[COMBINED_SHEET].delete_cols(10)
            older = _serialize_exact(book, escape_text=True)
        finally:
            book.close()
        restored = import_pricing_workbook(older, 'old-standalone.xlsx', current)
        retained = next(item for rows in effective_catalog(restored['configuration'])['rate_groups'].values() for item in rows if item['id'] == identity)
        self.assertEqual(retained['manufacturer'], 'Current service metadata')
        self.assertEqual(restored['summary']['rates']['updated'], 0)
        cleared = import_pricing_workbook(exported, 'new-column-standalone.xlsx', current)
        retained = next(item for rows in effective_catalog(cleared['configuration'])['rate_groups'].values() for item in rows if item['id'] == identity)
        self.assertEqual(retained['manufacturer'], '')

    def test_packaged_review_seed_is_nonempty_and_proof_is_portable_metadata_only(self):
        reviewed, _ = load_manufacturer_seed()
        self.assertGreater(len(reviewed['assignments']), 0)
        original = {item['id']: item for item in self.data['inventory']}
        for assignment in reviewed['assignments']:
            if assignment['kind'] == 'inventory':
                self.assertEqual(assignment['name'], original[assignment['id']]['name'])
            for proof in assignment['evidence']:
                encoded = json.dumps(proof, ensure_ascii=False)
                self.assertNotIn('file://', encoded.lower())
                self.assertNotIn('c:\\\\users\\\\', encoded.lower())
                self.assertNotRegex(encoded, r'\$\s*\d')


if __name__ == '__main__':
    unittest.main()
