"""Derived yield units preserve coverage arithmetic and legacy file exchanges."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from estimator.calculator import calculate
from estimator.catalog import baseline, effective_catalog, validate_catalog, ValidationError, yield_unit
from estimator.pricing_workbook import export_pricing_workbook, import_pricing_workbook, _serialize_exact
from estimator.storage import Store


class PricingUnitTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.payload = export_pricing_workbook({})

    def test_every_export_cell_is_centred_and_units_are_not_repeated(self):
        book = load_workbook(BytesIO(self.payload))
        sheet = book['Inventory & Rates']
        headers = {cell.value: cell.column for cell in sheet[1]}
        for page in book:
            for row in page:
                for cell in row:
                    self.assertEqual((cell.alignment.horizontal, cell.alignment.vertical), ('center', 'center'))
        product = next(row[0].row for row in sheet if sheet.cell(row[0].row, headers['Inventory ID']).value == '204')
        self.assertEqual(sheet.cell(product, headers['Yield unit']).value, 'm² / unit')
        self.assertEqual(headers['Yield unit'], headers['Yield'] + 1)
        result = import_pricing_workbook(self.payload, 'pricing.xlsx', {})
        self.assertEqual(result['summary']['rates']['updated'], 0)
        book.close()

    def test_legacy_unit_metadata_is_not_displayed_or_converted(self):
        catalog = baseline()
        catalog['rate_groups']['primers'][0]['yield_unit'] = 'm² / drum; coverage'
        catalog['rate_groups']['topcoats'][0]['yield_unit'] = 'm² / 20 kg'
        configuration = {'catalog': catalog, 'rates': {'primers:1': {'yield': 71, 'price': 401.25}}}
        inputs = {'D20': '20kg SBR Latex - Promat', 'B20': 142, 'D21': '20kg SBR Latex - Promat', 'B21': 142}
        before = calculate(inputs, configuration)
        exported = export_pricing_workbook(configuration)
        book = load_workbook(BytesIO(exported))
        sheet = book['Inventory & Rates']
        headers = {cell.value: cell.column for cell in sheet[1]}
        product = next(row[0].row for row in sheet if sheet.cell(row[0].row, headers['Inventory ID']).value == '204')
        self.assertEqual(sheet.cell(product, headers['Yield unit']).value, 'm² / unit')
        book.close()
        after = import_pricing_workbook(exported, 'pricing.xlsx', configuration)
        self.assertEqual(after['summary']['rates']['updated'], 0)
        rates = effective_catalog(after['configuration'])['rate_groups']
        self.assertEqual(rates['primers'][0]['yield_unit'], 'm² / drum; coverage')
        self.assertEqual(rates['topcoats'][0]['yield_unit'], 'm² / 20 kg')
        self.assertEqual(calculate(inputs, after['configuration'])['cells'], before['cells'])

    def test_units_follow_existing_calculator_dimensions_for_all_categories(self):
        data = baseline()
        area_groups = {'mesh', 'pins', 'boards', 'primers', 'topcoats'}
        for group, rates in data['rate_groups'].items():
            expected = 'm² / unit' if group in area_groups else 'm / unit' if group == 'mastic' else ''
            for rate in rates:
                with self.subTest(group=group, rate=rate['id']):
                    self.assertEqual(yield_unit(group, rate), expected)
                    if expected:
                        self.assertEqual(yield_unit(group, {**rate, 'yield_unit': 'incorrect label'}), expected)

    def test_excel_units_are_locked_but_editable_input_cells_and_new_columns_are_unlocked(self):
        book = load_workbook(BytesIO(self.payload))
        sheet = book['Inventory & Rates']
        self.assertTrue(sheet.protection.sheet)
        self.assertFalse(sheet.protection.autoFilter)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                self.assertEqual(cell.protection.locked, cell.column == 12, cell.coordinate)
        self.assertTrue(sheet.column_dimensions['L'].protection.locked)
        self.assertFalse(sheet.column_dimensions['K'].protection.locked)
        self.assertFalse(sheet.column_dimensions['R'].protection.locked)
        self.assertTrue(all(cell.protection.locked for cell in sheet[1]))
        book.close()

    def test_xlsx_unit_edits_cannot_change_the_catalogue(self):
        book = load_workbook(BytesIO(self.payload))
        sheet = book['Inventory & Rates']
        headers = {cell.value: cell.column for cell in sheet[1]}
        product = next(row[0].row for row in sheet if sheet.cell(row[0].row, headers['Inventory ID']).value == '204')
        # A file edited by a tool bypassing sheet protection still cannot change
        # units. Old unit-list lengths also cannot prevent ordinary use edits.
        sheet.cell(product, headers['Yield unit']).value = 'm² / drum;not a unit;extra'
        result = import_pricing_workbook(_serialize_exact(book), 'changed-labels.xlsx', {})
        self.assertEqual(result['summary']['rates']['updated'], 0)
        self.assertNotIn('yield_unit', result['configuration']['catalog']['rate_groups']['primers'][0])
        self.assertEqual(yield_unit('primers', result['configuration']['catalog']['rate_groups']['primers'][0]), 'm² / unit')
        book.close()

    def test_adding_use_does_not_require_editing_the_locked_unit_column(self):
        book = load_workbook(BytesIO(self.payload))
        sheet = book['Inventory & Rates']
        headers = {cell.value: cell.column for cell in sheet[1]}
        product = next(row[0].row for row in sheet if sheet.cell(row[0].row, headers['Inventory ID']).value == '204')
        values = {'Group': 'boards', 'Selection name': 'Extra area coverage use', 'Price source': 'Inventory',
                  'Sell rate': '161', 'Yield type': 'Number', 'Yield': '7.125', 'Rate ID': '', 'Use order': ''}
        for name, entry in values.items():
            cell = sheet.cell(product, headers[name])
            cell.value = str(cell.value) + ';' + entry
        result = import_pricing_workbook(_serialize_exact(book), 'extra-use.xlsx', {})
        self.assertEqual(result['summary']['rates']['added'], 1)
        added = next(rate for rate in result['configuration']['catalog']['rate_groups']['boards'] if rate['name'] == 'Extra area coverage use')
        self.assertEqual(added['yield'], 7.125)
        self.assertEqual(yield_unit('boards', added), 'm² / unit')
        book.close()

    def test_non_yield_rows_export_blank_units_instead_of_empty_separators(self):
        book = load_workbook(BytesIO(self.payload))
        sheet = book['Inventory & Rates']
        headers = {cell.value: cell.column for cell in sheet[1]}
        for row in sheet.iter_rows(min_row=2):
            groups = str(row[headers['Group'] - 1].value or '').split(';')
            if not any(group in {'mesh', 'pins', 'boards', 'primers', 'topcoats', 'mastic'} for group in groups):
                self.assertIsNone(row[headers['Yield unit'] - 1].value)
        book.close()

    def test_distinct_units_remain_visible_when_one_product_has_area_and_linear_uses(self):
        catalog = baseline()
        catalog['rate_groups']['mastic'][0]['inventory_id'] = '204'
        book = load_workbook(BytesIO(export_pricing_workbook({'catalog': catalog})))
        sheet = book['Inventory & Rates']
        product = next(row[0].row for row in sheet if sheet.cell(row[0].row, 15).value == '204')
        self.assertEqual(sheet.cell(product, 12).value, 'm / unit; m² / unit')
        book.close()

    def test_clearing_editable_cells_removes_a_product_without_touching_locked_units(self):
        book = load_workbook(BytesIO(self.payload))
        sheet = book['Inventory & Rates']
        product = next(row[0].row for row in sheet if sheet.cell(row[0].row, 15).value == '204')
        for cell in sheet[product]:
            if not cell.protection.locked:
                cell.value = None
        self.assertEqual(sheet.cell(product, 12).value, 'm² / unit')
        result = import_pricing_workbook(_serialize_exact(book), 'removed.xlsx', {})
        self.assertEqual(result['summary']['inventory']['removed'], 1)
        self.assertEqual(result['summary']['rates']['removed'], 2)
        book.close()

    def test_earlier_compact_workbook_without_unit_column_still_imports(self):
        book = load_workbook(BytesIO(self.payload))
        book['Inventory & Rates'].delete_cols(12)
        book['Inventory & Rates'].auto_filter.ref = 'A1:AA418'
        result = import_pricing_workbook(_serialize_exact(book), 'older-pricing.xlsx', {})
        self.assertEqual(result['summary']['rates']['updated'], 0)
        self.assertEqual(result['summary']['inventory']['updated'], 0)
        book.close()

    def test_invalid_unit_metadata_is_rejected_before_persistence(self):
        for group, value in [('primers', 42), ('primers', 'x' * 101), ('labour_rates', 'm² / unit')]:
            catalog = baseline()
            catalog['rate_groups'][group][0]['yield_unit'] = value
            with self.subTest(group=group, value=value), self.assertRaises(ValidationError):
                validate_catalog(catalog)
        self.assertEqual(yield_unit('mastic', baseline()['rate_groups']['mastic'][0]), 'm / unit')

    def test_saved_library_and_reset_persist_across_restart_without_mutating_quotes(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'library.sqlite3'
            store = Store(path)
            custom = {'inventory': {'204': {'supplier_price': 300}}, 'rates': {'primers:1': {'yield': 71}}}
            store.save_configuration(custom)
            quote = store.save_quote({'title': 'Frozen library', 'inputs': {'D20': '20kg SBR Latex - Promat', 'B20': 142}})
            self.assertEqual(Store(path).configuration(), custom)
            store.save_configuration({'inventory': {}, 'rates': {}})
            self.assertEqual(Store(path).configuration(), {'inventory': {}, 'rates': {}})
            self.assertEqual(Store(path).quote(quote['id']), quote)


if __name__ == '__main__':
    unittest.main()
