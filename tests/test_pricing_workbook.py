"""Pricing XLSX round trips, list replacement, and untrusted-file rejection."""

from copy import deepcopy
import csv
from io import BytesIO, StringIO
import json
import hashlib
import math
import os
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
import unittest
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook

from estimator.calculator import calculate
from estimator.catalog import ValidationError, baseline, effective_catalog
from estimator.pricing_workbook import (
    INVENTORY_HEADERS, RATE_HEADERS, PROPERTY_HEADERS, COMBINED_HEADERS, COMPACT_HEADERS, COMBINED_SHEET,
    _legacy_pricing_workbook,
    import_pricing_workbook, _format_sheet, _serialize_exact,
)


def export_compact_pricing_workbook(configuration):
    """Retained earlier compact-file fixture; new layout has dedicated tests."""
    workbook = _legacy_pricing_workbook(configuration)
    try:
        return _serialize_exact(workbook, escape_text=True)
    finally:
        workbook.close()


def export_pricing_workbook(configuration):
    """Independent fixture for the earlier two-list template import contract."""
    data = effective_catalog(configuration)
    workbook = Workbook()
    inventory = workbook.active
    inventory.title = "Inventory"
    inventory.append(INVENTORY_HEADERS)
    for item in data["inventory"]:
        inventory.append([item["id"], item.get("item_code", ""), item["name"], item["sales_description"],
                          "Supplier markup" if item["pricing_mode"] == "supplier_markup" else "Manual",
                          item.get("supplier_price"), item["markup"], item["sales_price"],
                          *(item.get("properties", {}).get(key) for key in PROPERTY_HEADERS.values())])
    rates = workbook.create_sheet("Rates")
    rates.append(RATE_HEADERS)
    for group, records in data["rate_groups"].items():
        for rate in records:
            value = rate.get("yield")
            mode = "override" if "price" in configuration.get("rates", {}).get(rate["id"], {}) else rate.get("price_mode", "inventory" if rate.get("inventory_id") else "override")
            yield_type = "Not used" if not data["rate_group_rules"][group]["yield_column"] else "Blank" if value is None else "Empty text" if value == "" else "Number"
            rates.append([rate["id"], group, rate.get("inventory_id"), rate["name"], mode.title(), rate["price"], yield_type,
                          value if isinstance(value, (int, float)) else None])
    workbook.create_sheet("Instructions")
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    return _serialize_exact(workbook)


def modify(payload, action, *, normal_excel_precision=False):
    workbook = load_workbook(BytesIO(payload))
    action(workbook)
    if normal_excel_precision:
        stream = BytesIO()
        workbook.save(stream)
        result = stream.getvalue()
    else:
        result = _serialize_exact(workbook)
    workbook.close()
    return result


def export_combined_pricing_workbook(configuration):
    """Independent fixture of the former Inventory/Use-row template."""
    original = load_workbook(BytesIO(export_pricing_workbook(configuration)))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = COMBINED_SHEET
    sheet.append(COMBINED_HEADERS)
    inventory = [dict(zip(INVENTORY_HEADERS, row)) for row in original["Inventory"].iter_rows(min_row=2, values_only=True)]
    rates = [dict(zip(RATE_HEADERS, row)) for row in original["Rates"].iter_rows(min_row=2, values_only=True)]
    groups = {}
    for rate in rates:
        groups[rate["Group"]] = groups.get(rate["Group"], 0) + 1
        rate["Use order"] = groups[rate["Group"]]
    def append(kind, record):
        values = {**record, "Row type": kind}
        values["Name"] = record.get("Product name", record.get("Rate name"))
        values["Sell price / rate"] = record.get("Sell price", record.get("Unit sell rate"))
        sheet.append([values.get(header) for header in COMBINED_HEADERS])
    for item in inventory:
        append("Inventory", item)
        for rate in rates:
            if rate["Inventory ID"] == item["Inventory ID"]:
                append("Use", rate)
    for rate in rates:
        if rate["Inventory ID"] is None:
            append("Use", rate)
    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Legacy template", "Retained import fixture"])
    _format_sheet(sheet, COMBINED_HEADERS, {}, percent_columns=(6,), freeze_panes="C2")
    _format_sheet(instructions, (), {}, freeze_panes="A2")
    for page in workbook:
        for row in page:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    original.close()
    return _serialize_exact(workbook)


def find(sheet, identity):
    return next(row[0].row for row in sheet.iter_rows(min_row=2) if row[0].value == identity)


def replace_part(payload, name, content):
    output = BytesIO()
    with ZipFile(BytesIO(payload)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item.filename, content if item.filename == name else source.read(item.filename))
        if name not in source.namelist():
            target.writestr(name, content)
    return output.getvalue()


class LegacyPricingWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.exported = export_pricing_workbook({})

    def imported(self, payload=None, current=None):
        return import_pricing_workbook(self.exported if payload is None else payload,
                                       "ceasefire-pricing.xlsx", {} if current is None else current)

    def test_legacy_template_contains_complete_typed_lists(self):
        workbook = load_workbook(BytesIO(self.exported))
        self.assertEqual(workbook.sheetnames, ["Inventory", "Rates", "Instructions"])
        inventory, rates = workbook["Inventory"], workbook["Rates"]
        self.assertEqual(tuple(cell.value for cell in inventory[1]), INVENTORY_HEADERS)
        self.assertEqual(tuple(cell.value for cell in rates[1]), RATE_HEADERS)
        self.assertEqual((inventory.max_row, rates.max_row), (418, 167))
        self.assertEqual(inventory["G3"].value, 0.3)
        self.assertFalse(any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row))
        self.assertEqual(inventory.cell(find(inventory, "239"), 8).value,
                         next(item for item in baseline()["inventory"] if item["id"] == "239")["sales_price"])
        workbook.close()

    def test_unchanged_import_preserves_every_calculator_cell_and_reports_no_changes(self):
        result = self.imported()
        empty = {"added": 0, "removed": 0, "updated": 0}
        self.assertEqual(result["summary"], {"inventory": empty, "rates": empty})
        catalog = effective_catalog(result["configuration"])
        for old, new in zip(baseline()["inventory"], catalog["inventory"]):
            for field in ("name", "sales_description", "sales_price", "supplier_price", "markup", "properties"):
                self.assertEqual(old[field], new[field], (old["id"], field))
        fixture_path = Path(__file__).parent / "fixtures" / "excel-calculator-oracle.json"
        fixture = json.loads(fixture_path.read_text(encoding="utf-8-sig"))
        self.assertEqual(len(fixture["scenarios"]), 216)
        source_index = {}
        for rows in baseline()["rate_groups"].values():
            for rate in rows:
                for field in ("price", "yield"):
                    if source := rate["source"].get(field):
                        source_index[source.replace("Lists!", "")] = (rate["id"], field)
        for case in fixture["scenarios"]:
            config = deepcopy(result["configuration"])
            for source, expected in case.get("lookupOverrides", {}).items():
                rate_id, field = source_index[source.replace("Lists!", "")]
                config["rates"].setdefault(rate_id, {})[field] = expected
            actual = calculate(case["inputs"], config)["cells"]
            for cell, expected in case["expected"].items():
                with self.subTest(scenario=case["id"], cell=cell):
                    if isinstance(expected, (int, float)):
                        self.assertTrue(math.isclose(actual[cell], expected, rel_tol=1e-12, abs_tol=1e-8))
                    else:
                        self.assertEqual(actual[cell], expected)
        self.assertEqual(calculate()["cells"], calculate(configuration=result["configuration"])["cells"])
        self.assertEqual(result["configuration"]["catalog"]["sources"]["pricing_import"],
                         {"filename": "ceasefire-pricing.xlsx", "sha256": hashlib.sha256(self.exported).hexdigest()})

    def test_excel_save_precision_preserves_stored_prices_and_trailing_choice_names(self):
        saved_by_excel = modify(self.exported, lambda workbook: None, normal_excel_precision=True)
        result = self.imported(saved_by_excel)
        self.assertEqual(result["summary"]["inventory"]["updated"], 0)
        self.assertEqual(result["summary"]["rates"]["updated"], 0)
        self.assertEqual(calculate({"D6": "1 Team (sheeting) - 1x "})["cells"],
                         calculate({"D6": "1 Team (sheeting) - 1x "}, result["configuration"])["cells"])

    def test_missing_or_stale_dimensions_do_not_omit_inventory_or_rates(self):
        for dimension in (None, "A1:A1"):
            payload = self.exported
            for part in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
                with ZipFile(BytesIO(payload)) as archive:
                    root = ET.fromstring(archive.read(part))
                element = root.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}dimension")
                if dimension is None:
                    root.remove(element)
                else:
                    element.set("ref", dimension)
                payload = replace_part(payload, part, ET.tostring(root, encoding="utf-8"))
            with self.subTest(dimension=dimension):
                result = self.imported(payload)
                catalog = effective_catalog(result["configuration"])
                self.assertEqual(len(catalog["inventory"]), 417)
                self.assertEqual(sum(map(len, catalog["rate_groups"].values())), 166)
                self.assertEqual(result["summary"], {"inventory": {"added": 0, "removed": 0, "updated": 0},
                                                     "rates": {"added": 0, "removed": 0, "updated": 0}})
        with ZipFile(BytesIO(self.exported)) as archive:
            root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        root.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheetData").clear()
        payload = replace_part(self.exported, "xl/worksheets/sheet1.xml", ET.tostring(root, encoding="utf-8"))
        with self.assertRaisesRegex(ValidationError, "headers"):
            self.imported(payload)

    def test_supplier_markup_changes_propagate_without_freezing_exported_rates(self):
        def edit(workbook):
            row = find(workbook["Inventory"], "204")
            workbook["Inventory"].cell(row, 6, 300)
            workbook["Inventory"].cell(row, 7, 0.2)
        result = self.imported(modify(self.exported, edit))
        catalog = effective_catalog(result["configuration"])
        self.assertEqual(result["summary"]["inventory"]["updated"], 1)
        self.assertEqual(result["summary"]["rates"]["updated"], 2)
        for group in ("primers", "topcoats"):
            self.assertEqual(catalog["rate_groups"][group][0]["price"], 360)
        second = self.imported(export_pricing_workbook(result["configuration"]), result["configuration"])
        self.assertEqual(second["summary"]["rates"]["updated"], 0)

    def test_rate_price_changes_become_persistent_overrides_and_can_restore_link(self):
        def edit(workbook):
            workbook["Rates"].cell(find(workbook["Rates"], "primers:1"), 6, 401.25)
        first = self.imported(modify(self.exported, edit))["configuration"]
        catalog = effective_catalog(first)
        self.assertEqual(catalog["rate_groups"]["primers"][0]["price_mode"], "override")
        self.assertEqual(catalog["rate_groups"]["primers"][0]["price"], 401.25)
        changed = deepcopy(first)
        changed["inventory"]["204"] = {"supplier_price": 300}
        catalog = effective_catalog(changed)
        self.assertEqual(catalog["rate_groups"]["primers"][0]["price"], 401.25)
        self.assertEqual(catalog["rate_groups"]["topcoats"][0]["price"], 390)
        def restore(workbook):
            workbook["Rates"].cell(find(workbook["Rates"], "primers:1"), 5, "Inventory")
        restored = self.imported(modify(export_pricing_workbook(changed), restore), changed)
        self.assertEqual(effective_catalog(restored["configuration"])["rate_groups"]["primers"][0]["price"], 390)

    def test_exported_existing_overrides_round_trip_without_loss(self):
        configuration = {"inventory": {"200": {"supplier_price": 99}, "915": {"sales_price": 2222}},
                         "rates": {"boards:1": {"price": 123.45, "yield": 1.2345}}}
        before = deepcopy(configuration)
        result = self.imported(export_pricing_workbook(configuration), configuration)
        self.assertEqual(configuration, before)
        self.assertEqual(result["summary"]["inventory"]["updated"], 0)
        self.assertEqual(result["summary"]["rates"]["updated"], 0)
        inputs = {"B15": 18.75, "B19": 50.5, "D19": baseline()["rate_groups"]["boards"][0]["name"]}
        self.assertEqual(calculate(inputs, configuration)["cells"], calculate(inputs, result["configuration"])["cells"])

    def test_add_remove_products_choices_and_move_group_preserves_explicit_ids(self):
        def edit(workbook):
            inventory, rates = workbook["Inventory"], workbook["Rates"]
            inventory.delete_rows(find(inventory, "200"))
            rates.delete_rows(find(rates, "sprays:2"))
            inventory.append(["new-product", "C-1", "New coating", "New coating", "Supplier markup", 25, 0.2, None] + [None] * 11)
            rates.append([None, "sprays", "new-product", "New coating", "Inventory", None, "Not used", None])
            rates.cell(find(rates, "sprays:3"), 2, "access_panels")
        result = self.imported(modify(self.exported, edit))
        self.assertEqual(result["summary"]["inventory"], {"added": 1, "removed": 1, "updated": 0})
        self.assertEqual(result["summary"]["rates"], {"added": 1, "removed": 1, "updated": 1})
        catalog = effective_catalog(result["configuration"])
        new_rate = next(item for item in catalog["rate_groups"]["sprays"] if item["name"] == "New coating")
        self.assertTrue(new_rate["id"].startswith("rate_"))
        self.assertEqual(new_rate["price"], 30)
        self.assertEqual(calculate({"D15": "New coating", "B15": 2}, result["configuration"])["cells"]["A63"], 30)

    def test_blank_inventory_id_allocates_new_identity_and_independent_rate_is_supported(self):
        def edit(workbook):
            workbook["Inventory"].append([None, "X", "New product", "", "Manual", None, 0, 42] + [None] * 11)
            workbook["Rates"].append([None, "labour_rates", None, "New team", "Override", 1234.56, "Not used", None])
        result = self.imported(modify(self.exported, edit))
        catalog = effective_catalog(result["configuration"])
        self.assertTrue(catalog["inventory"][-1]["id"].startswith("inv_"))
        self.assertEqual(catalog["rate_groups"]["labour_rates"][-1]["price"], 1234.56)

    def test_blank_empty_text_and_numeric_yields_remain_distinct(self):
        for yield_type, value, expected in (("Blank", None, "#DIV/0!"), ("Empty text", None, "#VALUE!"), ("Number", 2.5, 4)):
            def edit(workbook):
                row = find(workbook["Rates"], "mesh:2")
                workbook["Rates"].cell(row, 7, yield_type)
                workbook["Rates"].cell(row, 8).value = value
            result = self.imported(modify(self.exported, edit))
            self.assertEqual(calculate({"D16": "Promat Promamesh", "B16": 10}, result["configuration"])["cells"]["B68"], expected)

    def test_invalid_ids_links_names_headers_numbers_and_yields_are_rejected(self):
        def change(sheet, identity, column, value):
            def edit(workbook):
                workbook[sheet].cell(find(workbook[sheet], identity), column).value = value
            return edit
        invalid = [
            change("Inventory", "200", 1, "204"),
            change("Rates", "sprays:2", 1, "sprays:1"),
            change("Rates", "sprays:2", 4, "N/A"),
            change("Rates", "sprays:2", 3, "missing"),
            change("Rates", "sprays:2", 2, "invented_group"),
            change("Rates", "sprays:2", 7, "Number"),
            change("Rates", "mesh:2", 7, "Blank"),
            change("Inventory", "200", 6, -1),
            change("Inventory", "200", 6, "55"),
            change("Inventory", "200", 6, True),
            change("Inventory", "200", 7, -1.1),
            change("Inventory", "200", 6, 1e13),
            change("Inventory", "200", 8, 123.45),
            lambda workbook: setattr(workbook["Inventory"]["A1"], "value", "Wrong header"),
            lambda workbook: workbook["Inventory"].delete_rows(find(workbook["Inventory"], "200")),
            lambda workbook: workbook["Inventory"].cell(5002, 1, "oversized"),
        ]
        for action in invalid:
            with self.subTest(action=action), self.assertRaises(ValidationError):
                self.imported(modify(self.exported, action))

    def test_formulas_macros_external_links_xml_entities_and_oversized_archives_rejected(self):
        with self.assertRaisesRegex(ValidationError, "formulas"):
            self.imported(modify(self.exported, lambda workbook: setattr(workbook["Inventory"]["F3"], "value", "=1+2")))
        with self.assertRaisesRegex(ValidationError, "Macros"):
            self.imported(replace_part(self.exported, "xl/vbaProject.bin", b"anything"))
        with self.assertRaisesRegex(ValidationError, "External links"):
            self.imported(replace_part(self.exported, "xl/_rels/extra.xml.rels", b'<Relationships><Relationship TargetMode="External" Target="https://invalid.test"/></Relationships>'))
        with self.assertRaisesRegex(ValidationError, "entities"):
            self.imported(replace_part(self.exported, "xl/worksheets/sheet1.xml", b'<!DOCTYPE a [<!ENTITY b "x">]><a>&b;</a>'))
        with self.assertRaisesRegex(ValidationError, "20 MB"):
            self.imported(replace_part(self.exported, "padding.bin", b"0" * (21 * 1024 * 1024)))
        with ZipFile(BytesIO(self.exported)) as archive:
            xml = archive.read("xl/worksheets/sheet1.xml").replace(b">112.5<", b">not-a-number<", 1)
        with self.assertRaisesRegex(ValidationError, "malformed cells"):
            self.imported(replace_part(self.exported, "xl/worksheets/sheet1.xml", xml))
        for content, filename in ((b"garbage", "list.xlsx"), (self.exported, "list.xlsm"), (b"", "list.xlsx")):
            with self.subTest(filename=filename), self.assertRaises(ValidationError):
                import_pricing_workbook(content, filename, {})

    def test_labels_starting_equals_are_exported_as_literal_text(self):
        configuration = {"inventory": {"200": {"name": "=A1+A2"}}}
        payload = export_pricing_workbook(configuration)
        workbook = load_workbook(BytesIO(payload))
        cell = workbook["Inventory"].cell(find(workbook["Inventory"], "200"), 3)
        self.assertEqual((cell.value, cell.data_type), ("=A1+A2", "s"))
        workbook.close()
        self.imported(payload, configuration)


def combined_row(sheet, kind, identity):
    key = COMBINED_HEADERS.index("Inventory ID" if kind == "Inventory" else "Rate ID")
    return next(row[0].row for row in sheet.iter_rows(min_row=2) if row[0].value == kind and row[key].value == identity)


def combined_cell(sheet, kind, identity, field):
    return sheet.cell(combined_row(sheet, kind, identity), COMBINED_HEADERS.index(field) + 1)


def append_combined(sheet, **values):
    sheet.append([values.get(header) for header in COMBINED_HEADERS])


class CombinedPricingWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.exported = export_combined_pricing_workbook({})

    def imported(self, payload=None, current=None):
        return import_pricing_workbook(self.exported if payload is None else payload, "ceasefire-pricing.xlsx", {} if current is None else current)

    # Exercise the same independent 216-scenario Excel oracle through both file
    # layouts; the combined format must not weaken the legacy parity coverage.
    test_combined_import_matches_216_native_excel_scenarios = LegacyPricingWorkbookTests.test_unchanged_import_preserves_every_calculator_cell_and_reports_no_changes

    def test_export_has_one_owner_per_product_and_visible_uses_with_exact_values(self):
        workbook = load_workbook(BytesIO(self.exported))
        self.assertEqual(workbook.sheetnames, [COMBINED_SHEET, "Instructions"])
        sheet = workbook[COMBINED_SHEET]
        self.assertEqual(tuple(cell.value for cell in sheet[1]), COMBINED_HEADERS)
        self.assertEqual((sheet.max_row, sheet.max_column), (584, 26))
        self.assertEqual(sheet.freeze_panes, "C2")
        self.assertEqual(sheet.auto_filter.ref, "A1:Z584")
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(sum(row[0] == "Inventory" for row in rows), 417)
        self.assertEqual(sum(row[0] == "Use" for row in rows), 166)
        self.assertEqual(combined_cell(sheet, "Inventory", "204", "Supplier price").value, 296.1)
        self.assertEqual(combined_cell(sheet, "Inventory", "204", "Markup").value, .3)
        self.assertIn("%", combined_cell(sheet, "Inventory", "204", "Markup").number_format)
        self.assertEqual(combined_cell(sheet, "Inventory", "204", "Sell price / rate").value, 384.93)
        parent = combined_row(sheet, "Inventory", "204")
        self.assertFalse(sheet.row_dimensions[parent].collapsed)
        for rate_id, offset in (("primers:1", 1), ("topcoats:1", 2)):
            child = combined_row(sheet, "Use", rate_id)
            self.assertEqual(child, parent + offset)
            self.assertEqual(sheet.row_dimensions[child].outlineLevel, 0)
            self.assertFalse(sheet.row_dimensions[child].hidden)
            self.assertEqual(combined_cell(sheet, "Use", rate_id, "Inventory ID").value, "204")
            self.assertEqual(combined_cell(sheet, "Use", rate_id, "Yield").value, 142)
            self.assertIsNone(combined_cell(sheet, "Use", rate_id, "Supplier price").value)
        for row in range(2, sheet.max_row + 1):
            self.assertFalse(sheet.row_dimensions[row].hidden, row)
            self.assertFalse(sheet.row_dimensions[row].collapsed, row)
            self.assertEqual(sheet.row_dimensions[row].outlineLevel, 0, row)
        use_headers = ("Group", "Price source", "Yield type", "Yield", "Rate ID", "Use order")
        for header in use_headers:
            self.assertIsNone(combined_cell(sheet, "Inventory", "204", header).value, header)
        self.assertEqual([combined_cell(sheet, "Use", "primers:1", header).value for header in use_headers],
                         ["primers", "Inventory", "Number", 142, "primers:1", 1])
        self.assertFalse(any(cell.data_type == "f" for page in workbook for row in page for cell in row))
        workbook.close()

    def test_exported_views_have_only_valid_unique_panes_before_and_after_openpyxl_save(self):
        namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        expected = {
            "xl/worksheets/sheet1.xml": (
                {"xSplit": "2", "ySplit": "1", "topLeftCell": "C2", "activePane": "bottomRight", "state": "frozen"},
                [("topRight", "C1"), ("bottomLeft", "A2"), ("bottomRight", "C2")]),
            "xl/worksheets/sheet2.xml": (
                {"ySplit": "1", "topLeftCell": "A2", "activePane": "bottomLeft", "state": "frozen"},
                [("bottomLeft", "A2")]),
        }
        for payload in (self.exported, modify(self.exported, lambda workbook: None, normal_excel_precision=True)):
            with ZipFile(BytesIO(payload)) as archive:
                for path, (pane_attributes, selections) in expected.items():
                    root = ET.fromstring(archive.read(path))
                    views = root.findall("s:sheetViews/s:sheetView", namespace)
                    self.assertEqual(len(views), 1, path)
                    view = views[0]
                    panes = view.findall("s:pane", namespace)
                    self.assertEqual(len(panes), 1, path)
                    self.assertEqual(panes[0].attrib, pane_attributes, path)
                    selected = view.findall("s:selection", namespace)
                    self.assertEqual([(node.get("pane"), node.get("activeCell")) for node in selected], selections, path)
                    self.assertEqual(len({node.get("pane") for node in selected}), len(selected), path)
                    self.assertTrue(all(node.get("sqref") == node.get("activeCell") for node in selected))
                    self.assertIn(panes[0].get("activePane"), {node.get("pane") for node in selected})

    def test_formatting_reinitializes_selections_when_freeze_changes(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Product", "Rate"])
        sheet.append(["Test product", 12.3456789012345])
        original = list(sheet.values)
        for freeze, expected in (("E2", ["topRight", "bottomLeft", "bottomRight"]),
                                 ("A2", ["bottomLeft"]), ("C2", ["topRight", "bottomLeft", "bottomRight"]),
                                 ("C2", ["topRight", "bottomLeft", "bottomRight"]),
                                 ("C1", ["topRight"]), ("A1", [None])):
            _format_sheet(sheet, (), {}, freeze_panes=freeze)
            self.assertEqual([selection.pane for selection in sheet.sheet_view.selection], expected)
            self.assertEqual(len(sheet.views.sheetView), 1)
            self.assertEqual(list(sheet.values), original)
        workbook.close()

    def test_roundtrip_preserves_every_product_rate_dropdown_order_and_calculation(self):
        for normal_precision in (False, True):
            payload = modify(self.exported, lambda workbook: None, normal_excel_precision=normal_precision)
            result = self.imported(payload)
            empty = {"added": 0, "removed": 0, "updated": 0}
            self.assertEqual(result["summary"], {"inventory": empty, "rates": empty})
            catalog = effective_catalog(result["configuration"])
            original = baseline()
            self.assertEqual(catalog["inventory"], original["inventory"])
            for group, rates in original["rate_groups"].items():
                self.assertEqual([r["id"] for r in catalog["rate_groups"][group]], [r["id"] for r in rates])
                for before, after in zip(rates, catalog["rate_groups"][group]):
                    for key in ("id", "name", "inventory_id", "price", "yield"):
                        self.assertEqual(after[key], before[key], (group, key))
            for inputs in ({}, {"B15": 37.625, "B16": 60, "B9": .17, "B27": .125}, {"D6": "1 Team (sheeting) - 1x "}, {"C15": 0}):
                self.assertEqual(calculate(inputs)["cells"], calculate(inputs, result["configuration"])["cells"])

    def test_links_and_custom_dropdown_order_survive_sorting_collapse_and_filters(self):
        catalog = baseline()
        for rates in catalog["rate_groups"].values():
            rates.reverse()
        current = {"catalog": catalog, "inventory": {}, "rates": {"primers:1": {"price": 413.1234567890123, "yield": 137.25}}}
        def sort(workbook):
            sheet = workbook[COMBINED_SHEET]
            data = list(sheet.iter_rows(min_row=2, values_only=True))
            sheet.delete_rows(2, sheet.max_row)
            for row in reversed(data):
                sheet.append(row)
            for row in range(2, sheet.max_row + 1):
                sheet.row_dimensions[row].hidden = True
        result = self.imported(modify(export_combined_pricing_workbook(current), sort), current)
        before, after = effective_catalog(current), effective_catalog(result["configuration"])
        self.assertEqual({r["id"]: r for r in before["inventory"]}, {r["id"]: r for r in after["inventory"]})
        for group, rates in before["rate_groups"].items():
            self.assertEqual([r["id"] for r in after["rate_groups"][group]], [r["id"] for r in rates])
        self.assertEqual(calculate({"D20": "20kg SBR Latex - Promat", "B20": 280}, current)["cells"],
                         calculate({"D20": "20kg SBR Latex - Promat", "B20": 280}, result["configuration"])["cells"])

    def test_inventory_and_independent_use_edits_keep_price_and_yield_ownership(self):
        def edit(workbook):
            sheet = workbook[COMBINED_SHEET]
            for kind, identity, key, value in (("Inventory", "204", "Supplier price", 300), ("Inventory", "204", "Markup", .2),
                                             ("Inventory", "915", "Sell price / rate", 2222.125),
                                             ("Use", "primers:1", "Sell price / rate", 401.25), ("Use", "primers:1", "Yield", 71)):
                combined_cell(sheet, kind, identity, key).value = value
        result = self.imported(modify(self.exported, edit))
        catalog = effective_catalog(result["configuration"])
        primer, topcoat = catalog["rate_groups"]["primers"][0], catalog["rate_groups"]["topcoats"][0]
        self.assertEqual((primer["price"], primer["yield"], primer["price_mode"]), (401.25, 71, "override"))
        self.assertEqual((topcoat["price"], topcoat["yield"]), (360, 142))
        self.assertEqual(catalog["rate_groups"]["labour_rates"][3]["price"], 2222.125)
        self.assertEqual(self.imported(export_combined_pricing_workbook(result["configuration"]), result["configuration"])["summary"]["rates"]["updated"], 0)
        def restore(workbook):
            combined_cell(workbook[COMBINED_SHEET], "Use", "primers:1", "Price source").value = "Inventory"
        restored = self.imported(modify(export_combined_pricing_workbook(result["configuration"]), restore), result["configuration"])
        self.assertEqual(effective_catalog(restored["configuration"])["rate_groups"]["primers"][0]["price"], 360)

    def test_add_remove_rename_move_groups_and_unlinked_uses_are_explicit(self):
        def edit(workbook):
            sheet = workbook[COMBINED_SHEET]
            sheet.delete_rows(combined_row(sheet, "Use", "sprays:2"))
            sheet.delete_rows(combined_row(sheet, "Inventory", "200"))
            # Child precedes parent deliberately; IDs, not adjacency, link them.
            append_combined(sheet, **{"Row type": "Use", "Name": "New coating", "Group": "sprays", "Inventory ID": "new-product", "Price source": "Inventory", "Yield type": "Not used"})
            append_combined(sheet, **{"Row type": "Inventory", "Name": "New coating", "Item code": "C-1", "Inventory ID": "new-product", "Pricing mode": "Supplier markup", "Supplier price": 25, "Markup": .2})
            append_combined(sheet, **{"Row type": "Use", "Name": "New team", "Group": "labour_rates", "Price source": "Override", "Sell price / rate": 1234.56, "Yield type": "Not used"})
            combined_cell(sheet, "Use", "sprays:3", "Group").value = "access_panels"
            combined_cell(sheet, "Use", "sprays:3", "Use order").value = None
            combined_cell(sheet, "Use", "primers:1", "Name").value = "Renamed primer"
            combined_cell(sheet, "Inventory", "204", "Name").value = "Inventory display only"
        result = self.imported(modify(self.exported, edit))
        catalog = effective_catalog(result["configuration"])
        self.assertEqual(result["summary"]["inventory"], {"added": 1, "removed": 1, "updated": 1})
        self.assertEqual(result["summary"]["rates"], {"added": 2, "removed": 1, "updated": 2})
        self.assertEqual(catalog["rate_groups"]["sprays"][-1]["price"], 30)
        self.assertEqual(catalog["rate_groups"]["access_panels"][-1]["id"], "sprays:3")
        self.assertIsNone(catalog["rate_groups"]["labour_rates"][-1]["inventory_id"])
        self.assertEqual(calculate({"D15": "New coating", "D20": "Renamed primer"}, result["configuration"])["cells"]["A63"], 30)

    def test_blank_yield_variants_and_unused_inventory_are_retained(self):
        for kind, value, expected in (("Blank", None, "#DIV/0!"), ("Empty text", None, "#VALUE!"), ("Number", 2.5, 4)):
            def edit(workbook):
                sheet = workbook[COMBINED_SHEET]
                combined_cell(sheet, "Use", "mesh:2", "Yield type").value = kind
                combined_cell(sheet, "Use", "mesh:2", "Yield").value = value
                append_combined(sheet, **{"Row type": "Inventory", "Name": "Unused item", "Pricing mode": "Manual", "Sell price / rate": 42, "Markup": 0})
            result = self.imported(modify(self.exported, edit))
            self.assertEqual(calculate({"D16": "Promat Promamesh", "B16": 10}, result["configuration"])["cells"]["B68"], expected)
            self.assertTrue(result["configuration"]["catalog"]["inventory"][-1]["id"].startswith("inv_"))
            again = self.imported(export_combined_pricing_workbook(result["configuration"]), result["configuration"])
            self.assertEqual(again["summary"]["inventory"]["added"], 0)

    def test_invalid_row_ownership_links_names_orders_and_mixed_formats_fail(self):
        invalid = [("Use", "primers:1", "Supplier price", 5), ("Inventory", "204", "Yield", 5),
                   ("Inventory", "204", "Rate ID", "unexpected"), ("Use", "sprays:2", "Inventory ID", "missing"),
                   ("Use", "sprays:2", "Rate ID", "sprays:1"), ("Inventory", "200", "Inventory ID", "204"),
                   ("Use", "sprays:2", "Name", "N/A"), ("Use", "sprays:2", "Group", "invented"),
                   ("Use", "sprays:2", "Use order", 1), ("Use", "sprays:2", "Use order", -1),
                   ("Use", "sprays:2", "Use order", 0), ("Use", "sprays:2", "Use order", 1.5),
                   ("Use", "sprays:2", "Use order", 5001), ("Use", "sprays:2", "Use order", "2"),
                   ("Use", "sprays:2", "Use order", True), ("Use", "sprays:2", "Row type", "Product")]
        for kind, identity, field, value in invalid:
            def edit(workbook):
                combined_cell(workbook[COMBINED_SHEET], kind, identity, field).value = value
            with self.subTest(field=field, value=value), self.assertRaises(ValidationError):
                self.imported(modify(self.exported, edit))
        for action in (lambda w: w.create_sheet("Rates"), lambda w: w[COMBINED_SHEET].cell(10002, 1, "Inventory"),
                       lambda w: w[COMBINED_SHEET].cell(2, 27, "extra")):
            with self.assertRaises(ValidationError):
                self.imported(modify(self.exported, action))

    def test_per_type_limit_and_security_are_enforced_before_import(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = COMBINED_SHEET
        sheet.append(COMBINED_HEADERS)
        for index in range(3):
            append_combined(sheet, **{"Row type": "Inventory", "Inventory ID": f"new{index}", "Name": "Unused", "Pricing mode": "Manual", "Markup": 0, "Sell price / rate": 2})
        with patch("estimator.pricing_workbook.MAX_ROWS", 2), self.assertRaisesRegex(ValidationError, "Inventory rows"):
            self.imported(_serialize_exact(workbook))
        with self.assertRaisesRegex(ValidationError, "formulas"):
            self.imported(modify(self.exported, lambda w: setattr(w[COMBINED_SHEET]["E3"], "value", "=1+2")))
        with self.assertRaisesRegex(ValidationError, "Macros"):
            self.imported(replace_part(self.exported, "xl/vbaProject.bin", b"anything"))
        with self.assertRaisesRegex(ValidationError, "External links"):
            self.imported(replace_part(self.exported, "xl/_rels/extra.xml.rels", b'<Relationships><Relationship TargetMode="External" Target="https://invalid.test"/></Relationships>'))
        with self.assertRaisesRegex(ValidationError, "entities"):
            self.imported(replace_part(self.exported, "xl/workbook.xml", b'<!DOCTYPE a [<!ENTITY b "x">]><a>&b;</a>'))


VECTOR_HEADERS = ("Group", "Selection name", "Price source", "Sell rate", "Yield type", "Yield", "Rate ID", "Use order", "Yield unit")


def vector(values):
    """Independent fixture codec, including deliberate empty CSV slots."""
    if len(values) == 1 and not isinstance(values[0], str):
        return values[0]
    output = StringIO(newline="")
    csv.writer(output, delimiter=";", lineterminator="\r\n").writerow(values)
    return output.getvalue()[:-2]


def vector_values(value):
    return next(csv.reader(StringIO(value, newline=""), delimiter=";")) if isinstance(value, str) and value else [value]


def compact_cell(sheet, identity, field):
    row = next(row[0].row for row in sheet.iter_rows(min_row=2)
               if row[COMPACT_HEADERS.index("Inventory ID")].value == identity)
    return sheet.cell(row, COMPACT_HEADERS.index(field) + 1)


def append_compact(sheet, **values):
    sheet.append([values.get(header) for header in COMPACT_HEADERS])


def raw_pricing_text(payload, replacements, *, shared=False):
    """Independent Excel-style XML fixture; values already contain OOXML escapes."""
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(BytesIO(payload)) as archive:
        sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        types = ET.fromstring(archive.read("[Content_Types].xml"))
    strings = ET.Element(ns + "sst", count=str(len(replacements)), uniqueCount=str(len(replacements)))
    for address, parts in replacements.items():
        cell = sheet.find(f".//{ns}c[@r='{address}']")
        cell.clear()
        cell.set("r", address)
        cell.set("t", "s" if shared else "inlineStr")
        if shared:
            ET.SubElement(cell, ns + "v").text = str(len(strings))
            container = ET.SubElement(strings, ns + "si")
        else:
            container = ET.SubElement(cell, ns + "is")
        for part in parts:
            holder = ET.SubElement(container, ns + "r") if len(parts) > 1 else container
            node = ET.SubElement(holder, ns + "t", {"{http://www.w3.org/XML/1998/namespace}space": "preserve"})
            node.text = part
        if len(parts) > 1:
            phonetic = ET.SubElement(container, ns + "rPh", sb="0", eb="1")
            ET.SubElement(phonetic, ns + "t").text = "not displayed"
    payload = replace_part(payload, "xl/worksheets/sheet1.xml", ET.tostring(sheet, encoding="utf-8"))
    if shared:
        ET.SubElement(relations, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                      Id="pricingText", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
                      Target="sharedStrings.xml")
        ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                      PartName="/xl/sharedStrings.xml", ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml")
        payload = replace_part(payload, "xl/sharedStrings.xml", ET.tostring(strings, encoding="utf-8"))
        payload = replace_part(payload, "xl/_rels/workbook.xml.rels", ET.tostring(relations, encoding="utf-8"))
        payload = replace_part(payload, "[Content_Types].xml", ET.tostring(types, encoding="utf-8"))
    return payload


class CompactPricingWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.exported = export_compact_pricing_workbook({})

    def imported(self, payload=None, current=None):
        return import_pricing_workbook(self.exported if payload is None else payload, "ceasefire-pricing.xlsx", {} if current is None else current)

    test_compact_import_matches_216_native_excel_scenarios = LegacyPricingWorkbookTests.test_unchanged_import_preserves_every_calculator_cell_and_reports_no_changes
    test_roundtrip_preserves_all_values_and_calculations = CombinedPricingWorkbookTests.test_roundtrip_preserves_every_product_rate_dropdown_order_and_calculation
    test_exported_views_remain_valid_after_save = CombinedPricingWorkbookTests.test_exported_views_have_only_valid_unique_panes_before_and_after_openpyxl_save

    def test_one_row_per_product_has_complete_aligned_uses_and_readable_height(self):
        workbook = load_workbook(BytesIO(self.exported))
        sheet = workbook[COMBINED_SHEET]
        self.assertEqual(workbook.sheetnames, [COMBINED_SHEET, "Instructions"])
        self.assertEqual(tuple(cell.value for cell in sheet[1]), COMPACT_HEADERS)
        self.assertEqual((sheet.max_row, sheet.max_column), (418, 28))
        self.assertEqual(sheet.auto_filter.ref, "A1:AB418")
        self.assertTrue(sheet.column_dimensions["R"].hidden)
        self.assertEqual(compact_cell(sheet, "204", "Group").value, "primers;topcoats")
        self.assertEqual(compact_cell(sheet, "204", "Rate ID").value, "primers:1;topcoats:1")
        self.assertEqual(compact_cell(sheet, "204", "Sell rate").value, "384.93;384.93")
        self.assertEqual(compact_cell(sheet, "204", "Yield").value, "142;142")
        self.assertIsInstance(compact_cell(sheet, "204", "Sell price").value, float)
        self.assertIn("%", compact_cell(sheet, "204", "Markup").number_format)
        uses = 0
        for row in sheet.iter_rows(min_row=2):
            self.assertFalse(sheet.row_dimensions[row[0].row].hidden)
            self.assertEqual(sheet.row_dimensions[row[0].row].outlineLevel, 0)
            data = dict(zip(COMPACT_HEADERS, (cell.value for cell in row)))
            if data["Group"] is None:
                self.assertTrue(all(data[header] is None for header in VECTOR_HEADERS))
                continue
            count = len(vector_values(data["Group"]))
            self.assertTrue(all(len(vector_values(data[header])) == count for header in VECTOR_HEADERS if header != "Yield unit"))
            if data["Yield unit"] is not None:
                units = data["Yield unit"].split("; ")
                self.assertEqual(len(units), len(set(units)))
                self.assertTrue(set(units).issubset({"m² / unit", "m / unit"}))
            if count == 1:
                self.assertIsInstance(data["Sell rate"], (int, float))
            uses += count
        self.assertEqual(uses, 166)
        self.assertGreater(sheet.row_dimensions[compact_cell(sheet, "0", "Group").row].height, 31)
        self.assertFalse(any(cell.data_type == "f" for page in workbook for row in page for cell in row))
        workbook.close()

    def test_edits_preserve_independent_prices_yields_and_link_restoration(self):
        def edit(workbook):
            sheet = workbook[COMBINED_SHEET]
            for header, value in (("Supplier price", 300), ("Markup", .2), ("Sell rate", "401.1234567890123;384.93"), ("Yield", "71.12345678901234;142")):
                compact_cell(sheet, "204", header).value = value
        changed = self.imported(modify(self.exported, edit))
        catalog = effective_catalog(changed["configuration"])
        primer, topcoat = catalog["rate_groups"]["primers"][0], catalog["rate_groups"]["topcoats"][0]
        self.assertEqual((primer["price"], primer["yield"]), (401.1234567890123, 71.12345678901234))
        self.assertEqual((topcoat["price"], topcoat["yield"]), (360, 142))
        self.assertEqual(primer["price_mode"], "override")
        current = changed["configuration"]
        def restore(workbook):
            compact_cell(workbook[COMBINED_SHEET], "204", "Price source").value = "Inventory;Inventory"
        restored = self.imported(modify(export_compact_pricing_workbook(current), restore), current)
        rate = effective_catalog(restored["configuration"])["rate_groups"]["primers"][0]
        self.assertEqual((rate["price"], rate["yield"]), (360, 71.12345678901234))

    def test_csv_names_precision_ids_and_group_order_survive_sorting(self):
        def reorder(workbook):
            sheet = workbook[COMBINED_SHEET]
            data = list(sheet.iter_rows(min_row=2, values_only=True))
            sheet.delete_rows(2, sheet.max_row)
            for row in reversed(data):
                values = dict(zip(COMPACT_HEADERS, row))
                if values["Group"] is not None:
                    for header in VECTOR_HEADERS:
                        values[header] = vector(list(reversed(vector_values(values[header]))))
                append_compact(sheet, **values)
        # The stdlib openpyxl writer itself loses raw CR in ordinary saves.
        # Verify exact CRLF through our exporter, and ordinary-save precision
        # separately with LF, which that external writer represents faithfully.
        for ordinary_save in (False, True):
            with self.subTest(ordinary_save=ordinary_save):
                catalog = baseline()
                special = '  Name; with "quotes"' + ('\n' if ordinary_save else '\r\n') + 'next line  '
                catalog["rate_groups"]["primers"][0]["name"] = special
                current = {"catalog": catalog, "rates": {"primers:1": {"price": 413.1234567890123, "yield": 137.1234567890123}}, "inventory": {}}
                original = effective_catalog(current)
                payload = modify(export_compact_pricing_workbook(current), reorder, normal_excel_precision=ordinary_save)
                after = effective_catalog(self.imported(payload, current)["configuration"])
                self.assertEqual({r["id"]: r for r in after["inventory"]}, {r["id"]: r for r in original["inventory"]})
                for group, rates in original["rate_groups"].items():
                    self.assertEqual([r["id"] for r in after["rate_groups"][group]], [r["id"] for r in rates])
                    for old, new in zip(rates, after["rate_groups"][group]):
                        for key in ("name", "inventory_id", "price", "yield"):
                            self.assertEqual(new[key], old[key], (group, key))
                self.assertEqual(after["rate_groups"]["primers"][0]["name"], special)

    def test_stdlib_serializer_preserves_line_breaks_on_text_only_sheets(self):
        script = r'''
from io import BytesIO
from zipfile import ZipFile
import openpyxl
from estimator.pricing_workbook import _serialize_exact
assert not openpyxl.LXML
workbook = openpyxl.Workbook()
text = workbook.active
values = ["one\rtwo", "one\r\ntwo", "one\ntwo", "literal &#13;", '  name;"quoted"\t\u03bb  ']
for value in values:
    text.append([value])
numbers = workbook.create_sheet("Numbers")
numbers.append([0.12345678901234568, "numeric neighbour\r\ntext"])
for iteration in range(2):
    payload = _serialize_exact(workbook)
    with ZipFile(BytesIO(payload)) as archive:
        assert b"&#13;" in archive.read("xl/worksheets/sheet1.xml")
    workbook.close()
    workbook = openpyxl.load_workbook(BytesIO(payload))
    assert [row[0].value for row in workbook.worksheets[0]] == values
    assert workbook["Numbers"]["A1"].value == 0.12345678901234568
    assert workbook["Numbers"]["B1"].value == "numeric neighbour\r\ntext"
workbook.close()
'''
        process = subprocess.run([sys.executable, "-c", script], cwd=Path(__file__).resolve().parents[1],
                                 env={**os.environ, "OPENPYXL_LXML": "False"}, capture_output=True, text=True, timeout=30)
        self.assertEqual(process.returncode, 0, process.stdout + process.stderr)

    def test_excel_shared_and_inline_text_preserve_real_and_literal_escapes(self):
        workbook = load_workbook(BytesIO(self.exported))
        row = compact_cell(workbook[COMBINED_SHEET], "204", "Selection name").row
        workbook.close()
        expected = ['  Name; "quoted"\r\nnext line  ', 'literal _x000D_ _x005F_ bare x005F_ \U0001f600']
        encoded = vector(['  Name; "quoted"_x000D_\nnext line  ', 'literal _x005F_x000D_ _x005F_x005F_ bare x005F_ _xD83D__xDE00_'])
        for shared in (False, True):
            with self.subTest(shared=shared):
                payload = raw_pricing_text(self.exported, {
                    f"G{row}": [encoded], f"B{row}": ["split _x00", "0D_ product"], "F1": ["Gr_x006F_up"]}, shared=shared)
                data = effective_catalog(self.imported(payload)["configuration"])
                self.assertEqual(data["rate_groups"]["primers"][0]["name"], expected[0])
                self.assertEqual(data["rate_groups"]["topcoats"][0]["name"], expected[1])
                self.assertEqual(next(item["name"] for item in data["inventory"] if item["id"] == "204"), "split _x000D_ product")
                for group, identity in (("primers", "primers:1"), ("topcoats", "topcoats:1")):
                    rate = data["rate_groups"][group][0]
                    self.assertEqual((rate["id"], rate["inventory_id"], rate["price"], rate["yield"]), (identity, "204", 384.93, 142))

    def test_pricing_writer_protects_literal_escape_tokens_and_precise_use_values(self):
        catalog = baseline()
        names = ['  _x000D_; "literal" _x005F_ x005F_ _X000D_\r\nreal  ', 'Other _x0041_ \n line']
        for group, name in zip(("primers", "topcoats"), names):
            catalog["rate_groups"][group][0]["name"] = name
        config = {"catalog": catalog, "rates": {"primers:1": {"price": 413.1234567890123, "yield": 137.1234567890123},
                                                    "topcoats:1": {"price": 207.23456789012346, "yield": 62.34567890123456}}}
        payload = export_compact_pricing_workbook(config)
        with ZipFile(BytesIO(payload)) as archive:
            xml = archive.read("xl/worksheets/sheet1.xml")
        self.assertIn(b"_x005F_x000D_", xml)
        self.assertIn(b"_x005F_x005F_", xml)
        self.assertIn(b"&#13;", xml)
        result = self.imported(payload, config)
        data = effective_catalog(result["configuration"])
        for group, name in zip(("primers", "topcoats"), names):
            rate = data["rate_groups"][group][0]
            self.assertEqual(rate["name"], name)
            self.assertEqual((rate["price"], rate["yield"]), tuple(config["rates"][rate["id"]][field] for field in ("price", "yield")))
        self.assertEqual(result["summary"]["rates"]["updated"], 0)
        # The serializer's shared default keeps schedule export semantics.
        workbook = Workbook()
        workbook.active["A1"] = "_x000D_"
        with ZipFile(BytesIO(_serialize_exact(workbook))) as archive:
            self.assertNotIn(b"_x005F_x000D_", archive.read("xl/worksheets/sheet1.xml"))
        workbook.close()

    def test_decoded_controls_invalid_unicode_and_shared_references_are_rejected(self):
        workbook = load_workbook(BytesIO(self.exported))
        address = compact_cell(workbook[COMBINED_SHEET], "204", "Selection name").coordinate
        workbook.close()
        for shared in (False, True):
            for token in ("_x0000_", "_x0001_", "_x000B_", "_xFFFF_", "_xFFFE_", "_xD800_"):
                with self.subTest(shared=shared, token=token), self.assertRaisesRegex(ValidationError, "control characters|Unicode escape"):
                    self.imported(raw_pricing_text(self.exported, {address: [f"Bad{token};Other"]}, shared=shared))
        payload = raw_pricing_text(self.exported, {address: ["One;Two"]}, shared=True)
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        with ZipFile(BytesIO(payload)) as archive:
            xml = archive.read("xl/worksheets/sheet1.xml")
        for index in ("-1", "1", "1.5"):
            root = ET.fromstring(xml)
            root.find(f".//{ns}c[@r='{address}']/{ns}v").text = index
            with self.subTest(index=index), self.assertRaisesRegex(ValidationError, "shared text reference"):
                self.imported(replace_part(payload, "xl/worksheets/sheet1.xml", ET.tostring(root)))
        root = ET.fromstring(xml)
        row = root.find(f".//{ns}c[@r='{address}']/..")
        duplicate = ET.SubElement(row, ns + "c", r=address, t="n")
        ET.SubElement(duplicate, ns + "v").text = "1"
        with self.assertRaisesRegex(ValidationError, "duplicate cell addresses"):
            self.imported(replace_part(payload, "xl/worksheets/sheet1.xml", ET.tostring(root)))

    def test_decoded_text_does_not_change_formula_or_nonfinite_number_validation(self):
        workbook = load_workbook(BytesIO(self.exported))
        address = compact_cell(workbook[COMBINED_SHEET], "204", "Selection name").coordinate
        workbook.close()
        payload = raw_pricing_text(self.exported, {address: ["_x003D_1+1;Other"]}, shared=True)
        data = effective_catalog(self.imported(payload)["configuration"])
        self.assertEqual(data["rate_groups"]["primers"][0]["name"], "=1+1")
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        with ZipFile(BytesIO(payload)) as archive:
            xml = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        ET.SubElement(xml.find(f".//{ns}c[@r='{address}']"), ns + "f").text = "1+1"
        with self.assertRaisesRegex(ValidationError, "formulas"):
            self.imported(replace_part(payload, "xl/worksheets/sheet1.xml", ET.tostring(xml)))
        for field in ("Sell rate", "Yield", "Use order"):
            with self.subTest(field=field), self.assertRaisesRegex(ValidationError, "finite"):
                self.imported(modify(self.exported, lambda book: setattr(compact_cell(book[COMBINED_SHEET], "204", field), "value", "1e999;1")))

    def test_new_products_unused_products_and_standalone_rates_keep_ownership(self):
        def add(workbook):
            sheet = workbook[COMBINED_SHEET]
            append_compact(sheet, **{"Product name": "New linked product", "Pricing mode": "Manual", "Markup": 0, "Sell price": 12,
                "Group": "primers;topcoats", "Selection name": "New primer;New topcoat", "Price source": "Inventory;Inventory",
                "Sell rate": "12;12", "Yield type": "Number;Number", "Yield": "10;20", "Rate ID": ";", "Use order": ";"})
            append_compact(sheet, **{"Product name": "Unused product", "Pricing mode": "Manual", "Markup": 0, "Sell price": 3})
            append_compact(sheet, **{"Group": "primers", "Selection name": "Standalone primer", "Price source": "Override", "Sell rate": 7,
                                    "Yield type": "Number", "Yield": 0, "Rate ID": "new_standalone"})
        result = self.imported(modify(self.exported, add))
        catalog = effective_catalog(result["configuration"])
        self.assertEqual(len(catalog["inventory"]), 419)
        identity = next(item["id"] for item in catalog["inventory"] if item["name"] == "New linked product")
        new = [rate for rows in catalog["rate_groups"].values() for rate in rows if rate["name"] in ("New primer", "New topcoat")]
        self.assertEqual({rate["inventory_id"] for rate in new}, {identity})
        self.assertEqual(len({rate["id"] for rate in new}), 2)
        standalone = next(rate for rate in catalog["rate_groups"]["primers"] if rate["id"] == "new_standalone")
        self.assertIsNone(standalone["inventory_id"])
        self.assertEqual((standalone["price"], standalone["yield"]), (7, 0))
        again = self.imported(export_compact_pricing_workbook(result["configuration"]), result["configuration"])
        self.assertEqual(again["summary"], {"inventory": {"added": 0, "removed": 0, "updated": 0}, "rates": {"added": 0, "removed": 0, "updated": 0}})

    def test_blank_empty_text_zero_and_separator_padding_are_preserved(self):
        def edit(workbook):
            sheet = workbook[COMBINED_SHEET]
            for header, value in (("Group", "primers; topcoats"), ("Price source", "Inventory; Inventory"),
                                  ("Rate ID", "primers:1; topcoats:1"), ("Yield type", "Blank; Empty text"), ("Yield", " ; "), ("Use order", "1; 1")):
                compact_cell(sheet, "204", header).value = value
        catalog = effective_catalog(self.imported(modify(self.exported, edit))["configuration"])
        self.assertIsNone(catalog["rate_groups"]["primers"][0]["yield"])
        self.assertEqual(catalog["rate_groups"]["topcoats"][0]["yield"], "")
        def zero(workbook):
            compact_cell(workbook[COMBINED_SHEET], "204", "Yield").value = "0;142"
        catalog = effective_catalog(self.imported(modify(self.exported, zero))["configuration"])
        self.assertEqual(catalog["rate_groups"]["primers"][0]["yield"], 0)

    def test_mismatched_lists_bad_quotes_numbers_and_duplicate_orders_are_rejected(self):
        for header, value, message in (("Yield", "142", "same number"), ("Rate ID", "primers:1", "same number"),
                                       ("Selection name", 'Bad"name;Other', "CSV"), ("Selection name", '"Unclosed;Other', "CSV"),
                                       ("Sell rate", "NaN;360", "number"), ("Use order", "1.5;1", "whole number"),
                                       ("Group", "primers;primers", "Duplicate Use order")):
            with self.subTest(header=header, value=value), self.assertRaisesRegex(ValidationError, message):
                self.imported(modify(self.exported, lambda w: setattr(compact_cell(w[COMBINED_SHEET], "204", header), "value", value)))

    def test_expanded_use_limits_cell_size_and_security_remain_enforced(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = COMBINED_SHEET
        sheet.append(COMPACT_HEADERS)
        append_compact(sheet, **{"Product name": "Limit", "Inventory ID": "limit", "Pricing mode": "Manual", "Markup": 0, "Sell price": 1,
            "Group": "primers;primers;primers", "Selection name": "One;Two;Three", "Price source": "Inventory;Inventory;Inventory",
            "Sell rate": "1;1;1", "Yield type": "Number;Number;Number", "Yield": "1;1;1", "Rate ID": "one;two;three", "Use order": "1;2;3"})
        with patch("estimator.pricing_workbook.MAX_ROWS", 2), self.assertRaisesRegex(ValidationError, "after expansion"):
            self.imported(_serialize_exact(workbook))
        for header in ("Selection name", "Sell rate"):
            with self.assertRaisesRegex(ValidationError, "formulas"):
                self.imported(modify(self.exported, lambda w: setattr(compact_cell(w[COMBINED_SHEET], "204", header), "value", "=1+1")))
        with self.assertRaisesRegex(ValidationError, "provided columns|unexpected columns"):
            self.imported(modify(self.exported, lambda w: setattr(w[COMBINED_SHEET]["AC2"], "value", "extra")))
        catalog = baseline()
        seed = catalog["rate_groups"]["primers"][0]
        catalog["rate_groups"]["primers"] = [{**deepcopy(seed), "id": f"long:{i}", "name": f"{i}" + "x" * 990} for i in range(40)]
        with self.assertRaisesRegex(ValidationError, "32,767"):
            export_compact_pricing_workbook({"catalog": catalog})


if __name__ == "__main__":
    unittest.main()
