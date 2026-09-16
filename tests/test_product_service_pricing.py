"""Single-column descriptions and yields preserve estimator business inputs."""

from copy import deepcopy
from io import BytesIO
import json
import unittest

from openpyxl import load_workbook

from estimator.calculator import calculate
from estimator.catalog import (baseline, effective_catalog, catalog_signature,
                               product_service_name, validate_catalog,
                               validate_configuration, ValidationError)
from estimator.pricing_workbook import (export_pricing_workbook, import_pricing_workbook,
    PRODUCT_SERVICE_VISIBLE_HEADERS, PRODUCT_SERVICE_HEADERS, COMBINED_SHEET, _serialize_exact)
import test_pricing_workbook as pricing_fixtures


def edit(payload, changes, *, ordinary_save=False):
    book = load_workbook(BytesIO(payload))
    try:
        sheet = book[COMBINED_SHEET]
        headers = {cell.value: cell.column for cell in sheet[1]}
        for identity, fields in changes:
            number = next(row[0].row for row in sheet.iter_rows(min_row=2)
                if sheet.cell(row[0].row, headers['Inventory ID']).value == identity
                or sheet.cell(row[0].row, headers['Rate ID']).value == identity)
            for field, value in fields.items():
                cell = sheet.cell(number, headers[field])
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = 's'
        if ordinary_save:
            stream = BytesIO()
            book.save(stream)
            return stream.getvalue()
        return _serialize_exact(book, escape_text=True)
    finally:
        book.close()


class ProductServicePricingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.exported = export_pricing_workbook({})

    def imported(self, payload=None, current=None):
        return import_pricing_workbook(self.exported if payload is None else payload,
            'ceasefire-pricing.xlsx', {} if current is None else current)

    test_simplified_import_matches_216_native_excel_scenarios = pricing_fixtures.LegacyPricingWorkbookTests.test_unchanged_import_preserves_every_calculator_cell_and_reports_no_changes

    def test_export_has_only_eight_visible_columns_and_one_numeric_yield(self):
        book = load_workbook(BytesIO(self.exported))
        try:
            sheet = book[COMBINED_SHEET]
            headers = {cell.value: cell.column for cell in sheet[1]}
            visible = [cell.value for cell in sheet[1] if not sheet.column_dimensions[cell.column_letter].hidden]
            self.assertEqual(tuple(visible), PRODUCT_SERVICE_VISIBLE_HEADERS)
            self.assertEqual(tuple(cell.value for cell in sheet[1]), PRODUCT_SERVICE_HEADERS)
            self.assertEqual(sheet.auto_filter.ref, 'A1:H418')
            number = next(row[0].row for row in sheet if sheet.cell(row[0].row, headers['Inventory ID']).value == '204')
            self.assertEqual(sheet.cell(number, headers['Yield']).value, 142)
            self.assertEqual(sheet.cell(number, headers['Use yields']).value, '142;142')
            self.assertEqual(sheet.cell(number, headers['Product/Service']).value,
                product_service_name(next(item for item in baseline()['inventory'] if item['id'] == '204')))
        finally:
            book.close()

    def test_descriptive_choice_is_one_existing_label_and_has_stable_ties(self):
        product = {'name': 'Short', 'sales_description': 'A longer sales description'}
        rate = {'name': 'Other selection', 'display_name': 'Another'}
        self.assertEqual(product_service_name(product, [rate]), product['sales_description'])
        self.assertEqual(product_service_name({'name': 'Name', 'sales_description': 'Desc'}), 'Desc')
        self.assertEqual(product_service_name(product, [{**rate, 'product_service': 'Explicit rate'}]), 'Explicit rate')
        self.assertEqual(product_service_name({**product, 'product_service': 'Explicit product'}, [rate]), 'Explicit product')
        self.assertEqual(product_service_name(None, [{'name': '  standalone name  '}]), '  standalone name  ')

    def test_metadata_edit_does_not_change_identity_prices_yields_or_calculation(self):
        inputs = {'B20': 284, 'D20': '20kg SBR Latex - Promat', 'B21': 142, 'D21': '20kg SBR Latex - Promat'}
        before = effective_catalog({})
        renamed = {'inventory': {'204': {'product_service': 'SBR latex primer and topcoat, 20 kg'}}}
        after = effective_catalog(renamed)
        self.assertEqual(catalog_signature(before), catalog_signature(after))
        self.assertEqual(calculate(inputs, renamed)['cells'], calculate(inputs)['cells'])
        for group in ('primers', 'topcoats'):
            old, new = before['rate_groups'][group][0], after['rate_groups'][group][0]
            self.assertEqual(new['display_name'], renamed['inventory']['204']['product_service'])
            for key in ('id', 'name', 'price', 'yield', 'inventory_id'):
                self.assertEqual(new[key], old[key])

    def test_legacy_inventory_label_overrides_survive_noop_export_import(self):
        configuration = {'inventory': {'204': {'name': 'Legacy edited product',
                                                'sales_description': 'Legacy edited sales description'}}}
        before = effective_catalog(configuration)
        after = effective_catalog(self.imported(export_pricing_workbook(configuration), configuration)['configuration'])
        for group in ('primers', 'topcoats'):
            self.assertEqual(after['rate_groups'][group][0]['display_name'], before['rate_groups'][group][0]['display_name'])
            self.assertEqual(after['rate_groups'][group][0]['name'], before['rate_groups'][group][0]['name'])
        self.assertEqual(calculate(configuration=configuration)['cells'], calculate(configuration={'catalog': after})['cells'])

    def test_invalid_metadata_is_rejected_for_catalog_and_configuration(self):
        for value in ('', ' ', 'x' * 1001, 12, None, 'Bad\0label'):
            for category, identity in (('inventory', '204'), ('rates', 'primers:1')):
                with self.subTest(category=category, value=value), self.assertRaises(ValidationError):
                    validate_configuration({category: {identity: {'product_service': value}}})
        for kind in ('product', 'rate'):
            catalog = baseline()
            record = catalog['inventory'][0] if kind == 'product' else catalog['rate_groups']['primers'][0]
            record['product_service'] = 'Bad\0label'
            with self.assertRaises(ValidationError):
                validate_catalog(catalog)

    def test_excel_label_renames_all_linked_display_names_without_renaming_keys(self):
        label = '=Descriptive product; with "quotes" and _x000D_'
        result = self.imported(edit(self.exported, [('204', {'Product/Service': label})]))
        data = effective_catalog(result['configuration'])
        self.assertEqual(catalog_signature(data), catalog_signature())
        self.assertEqual(next(item['product_service'] for item in data['inventory'] if item['id'] == '204'), label)
        for group in ('primers', 'topcoats'):
            self.assertEqual(data['rate_groups'][group][0]['display_name'], label)
            self.assertEqual(data['rate_groups'][group][0]['name'], baseline()['rate_groups'][group][0]['name'])
        again = self.imported(export_pricing_workbook(result['configuration']), result['configuration'])
        self.assertEqual(again['summary']['inventory']['updated'], 0)
        self.assertEqual(again['summary']['rates']['updated'], 0)

    def test_standalone_label_and_sell_price_are_editable_and_remain_standalone(self):
        catalog = baseline()
        rate = catalog['rate_groups']['primers'][0]
        rate.update({'inventory_id': None, 'price_mode': 'override', 'product_service': 'Original independent coating'})
        original = {'catalog': catalog}
        payload = edit(export_pricing_workbook(original), [('primers:1', {
            'Product/Service': 'Independent primer service', 'Sell price': 444.125, 'Yield': 7.25})])
        result = self.imported(payload, original)['configuration']
        saved = effective_catalog(result)['rate_groups']['primers'][0]
        self.assertEqual((saved['inventory_id'], saved['name'], saved['display_name'], saved['price'], saved['yield']),
            (None, rate['name'], 'Independent primer service', 444.125, 7.25))
        # A different current library cannot replace the exported stable key or label.
        other = {'catalog': baseline(), 'rates': {'primers:1': {'price': 5, 'product_service': 'Different'}}}
        again = self.imported(export_pricing_workbook(result), other)['configuration']
        self.assertEqual(effective_catalog(again)['rate_groups']['primers'][0], saved)

    def test_differing_hidden_overrides_and_yields_survive_noop_and_inventory_edit(self):
        config = {'rates': {'primers:1': {'price': 401.1234567890123, 'yield': 71.12345678901234},
                            'topcoats:1': {'yield': 142}}, 'inventory': {'204': {'product_service': 'Shared coating'}}}
        payload = export_pricing_workbook(config)
        book = load_workbook(BytesIO(payload))
        sheet = book[COMBINED_SHEET]
        headers = {cell.value: cell.column for cell in sheet[1]}
        row = next(row[0].row for row in sheet if sheet.cell(row[0].row, headers['Inventory ID']).value == '204')
        self.assertEqual(sheet.cell(row, headers['Yield']).value, 'Mixed')
        book.close()
        copied = self.imported(payload, {'rates': {'primers:1': {'price': 900, 'yield': 1}}})['configuration']
        rates = effective_catalog(copied)['rate_groups']
        self.assertEqual((rates['primers'][0]['price'], rates['primers'][0]['yield']), (401.1234567890123, 71.12345678901234))
        self.assertEqual((rates['topcoats'][0]['price'], rates['topcoats'][0]['yield']), (384.93, 142))
        edited = self.imported(edit(payload, [('204', {'Supplier price': 300, 'Markup': .2})]), config)['configuration']
        rates = effective_catalog(edited)['rate_groups']
        self.assertEqual(rates['primers'][0]['price'], 401.1234567890123)
        self.assertEqual(rates['topcoats'][0]['price'], 360)

    def test_scalar_yield_edit_updates_only_yield_bearing_uses_and_retains_precision(self):
        catalog = baseline()
        catalog['rate_groups']['labour_rates'][0]['inventory_id'] = '204'
        config = {'catalog': catalog}
        before = effective_catalog(config)
        value = 71.12345678901234
        result = self.imported(edit(export_pricing_workbook(config), [('204', {'Yield': value})]), config)['configuration']
        after = effective_catalog(result)
        self.assertEqual(after['rate_groups']['primers'][0]['yield'], value)
        self.assertEqual(after['rate_groups']['topcoats'][0]['yield'], value)
        self.assertEqual(after['rate_groups']['labour_rates'][0]['yield'], before['rate_groups']['labour_rates'][0]['yield'])
        saved = self.imported(edit(export_pricing_workbook(result), [], ordinary_save=True), result)['configuration']
        self.assertEqual(effective_catalog(saved)['rate_groups']['primers'][0]['yield'], value)

    def test_mixed_units_reject_shared_edits_without_rejecting_unchanged_file(self):
        catalog = baseline()
        catalog['rate_groups']['mastic'][0]['inventory_id'] = '204'
        config = {'catalog': catalog}
        payload = export_pricing_workbook(config)
        result = self.imported(payload, config)
        self.assertEqual(result['summary']['rates']['updated'], 0)
        with self.assertRaisesRegex(ValidationError, 'incompatible yield units'):
            self.imported(edit(payload, [('204', {'Yield': 7})]), config)
        added = self.imported(edit(self.exported, [('204', {'Group': 'primers;topcoats;mastic'})]))['configuration']
        rates = effective_catalog(added)['rate_groups']
        self.assertEqual(rates['primers'][0]['yield'], 142)
        self.assertIsNone(next(rate['yield'] for rate in rates['mastic'] if rate['inventory_id'] == '204'))

    def test_blank_empty_text_zero_and_differing_states_remain_exact(self):
        for first, second in ((None, None), ('', ''), (None, ''), (0, 0), (None, 0)):
            config = {'rates': {'primers:1': {'yield': first}, 'topcoats:1': {'yield': second}}}
            with self.subTest(values=(first, second)):
                result = self.imported(export_pricing_workbook(config), config)['configuration']
                rates = effective_catalog(result)['rate_groups']
                self.assertEqual((rates['primers'][0]['yield'], rates['topcoats'][0]['yield']), (first, second))
        zero = self.imported(edit(self.exported, [('204', {'Yield': 0})]))['configuration']
        self.assertEqual(effective_catalog(zero)['rate_groups']['primers'][0]['yield'], 0)

    def test_visible_group_changes_preserve_existing_keys_and_create_new_use(self):
        payload = edit(self.exported, [('204', {'Group': 'primers;boards', 'Yield': 142})])
        result = self.imported(payload)['configuration']
        rates = effective_catalog(result)['rate_groups']
        self.assertEqual(rates['primers'][0]['id'], 'primers:1')
        self.assertFalse(any(rate['id'] == 'topcoats:1' for rate in rates['topcoats']))
        new = next(rate for rate in rates['boards'] if rate['inventory_id'] == '204')
        self.assertEqual((new['price'], new['yield']), (384.93, 142))

    def test_malformed_snapshot_or_scalar_yield_is_rejected(self):
        for field, value in [('Saved values', '{bad'), ('Saved values', '[]'),
                             ('Saved values', json.dumps({'product': []})), ('Yield', '142;142'), ('Yield', -1)]:
            with self.subTest(field=field, value=value), self.assertRaises(ValidationError):
                self.imported(edit(self.exported, [('204', {field: value})]))
        book = load_workbook(BytesIO(self.exported))
        sheet = book[COMBINED_SHEET]
        headers = {cell.value: cell.column for cell in sheet[1]}
        number = next(row[0].row for row in sheet if sheet.cell(row[0].row, headers['Inventory ID']).value == '204')
        original = json.loads(sheet.cell(number, headers['Saved values']).value)
        book.close()
        for malformed in ('properties-list', 'properties-null', 'price-mode-number', 'group-object', 'bad-order', 'deep-source'):
            saved = deepcopy(original)
            if malformed.startswith('properties'):
                saved['product']['properties'] = [] if malformed.endswith('list') else None
            elif malformed == 'price-mode-number':
                saved['uses'][0]['rate']['price_mode'] = 42
            elif malformed == 'group-object':
                saved['uses'][0]['group'] = {}
            elif malformed == 'bad-order':
                saved['uses'][0]['order'] = 'first'
            else:
                nested = []
                for _ in range(400):
                    nested = [nested]
                saved['product']['source'] = {'nested': nested}
            with self.subTest(snapshot=malformed), self.assertRaises(ValidationError):
                self.imported(edit(self.exported, [('204', {'Saved values': json.dumps(saved), 'Group': 'topcoats;primers'})]))


if __name__ == '__main__':
    unittest.main()
