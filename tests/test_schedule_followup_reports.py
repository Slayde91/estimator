"""Quantity exports and the revised board evidence stay reviewable."""
from io import BytesIO
import json
from pathlib import Path
import unittest

from openpyxl import load_workbook
from pypdf import PdfReader

from estimator.calculator_register import build_calculator_register
from estimator.calculator_report import build_calculator_report, build_calculator_summary_report, project_calculator_report
from estimator.schedule_rows import empty_schedule_inputs
from tests.test_fyrewrap_rules import schedule_input


class ScheduleFollowupReportTests(unittest.TestCase):
    def test_mixed_case_board_product_keeps_exact_input_and_assessment_trace(self):
        from tests.test_board_product_totals import example_inputs
        inputs = example_inputs(9, length=2)
        inputs['CALCULATOR'].update(C9='promatect 100', D9='100PFC', G9=4,
                                     H9=120, I9='Column', J9=550, CJ9=2)
        data = project_calculator_report('steel_board', inputs)
        self.assertEqual(data['inputs']['CALCULATOR']['C9'], 'promatect 100')
        self.assertEqual(data['rows'][0]['values']['C'], 'promatect 100')
        self.assertIn('FC15555-01-03', data['assessment_receipt'])
        self.assertIn('Table ', data['assessment_notes'][0])
        self.assertIn('m⁻¹', data['assessment_notes'][0])
        workbook = load_workbook(BytesIO(build_calculator_register('steel_board', inputs)))
        sheet = workbook['Schedule']
        headers = {cell.value: cell.column for cell in sheet[5]}
        self.assertEqual(headers['Quantity'], headers['Lineal metres'] + 1)
        self.assertEqual(sheet.cell(6, headers['Quantity']).value, 2)
        self.assertIn('FC15555-01-03', sheet.cell(6, headers['Status']).value)
        self.assertIn('m⁻¹', sheet.cell(6, headers['Status']).value)
        for builder in (build_calculator_report, build_calculator_summary_report):
            text = '\n'.join(page.extract_text() or '' for page in PdfReader(BytesIO(builder('steel_board', inputs))).pages)
            self.assertIn('1/m', text)
            self.assertNotIn('m⁻¹', text)

    def test_implicit_legacy_quantity_does_not_create_blank_report_rows(self):
        for identity in ('ductwork', 'steel_board'):
            with self.subTest(identity=identity):
                data = project_calculator_report(identity, empty_schedule_inputs(identity))
                self.assertEqual(data['rows'], [])
                self.assertEqual(data['incomplete_rows'], 0)

    def test_quantity_repeats_complete_duct_takeoff_in_pdf_and_excel(self):
        inputs = schedule_input({11: {'F': 1, 'G': 1, 'CM': 3}})
        data = project_calculator_report('ductwork', inputs)
        self.assertEqual(len(data['rows']), 1)
        values = data['rows'][0]['values']
        self.assertEqual((values['D'], values['CM']), (10, 3))
        # Independent published-detail hand takeoff for one run is27.8096m².
        self.assertAlmostEqual(values['N'], 27.8096 * 3, places=8)
        workbook = load_workbook(BytesIO(build_calculator_register('ductwork', inputs)))
        sheet = workbook['Schedule']
        headers = {cell.value: cell.column for cell in sheet[5]}
        self.assertEqual(headers['Quantity'], headers['Length (m)'] + 1)
        self.assertEqual(sheet.cell(6, headers['Quantity']).value, 3)
        self.assertEqual(sheet.cell(6, headers['Length (m)']).value, 10)
        self.assertAlmostEqual(sheet.cell(6, headers['Wrap material (m²)']).value, values['N'])
        text = '\n'.join(page.extract_text() or '' for page in PdfReader(BytesIO(build_calculator_report('ductwork', inputs))).pages)
        self.assertIn('qty', text)
        self.assertIn('10', text)
        self.assertIn('3', text)

    def test_board_document_uses_actual_downloaded_assessment_identity(self):
        path = Path(__file__).resolve().parents[1] / 'data' / 'calculator_documents.json'
        documents = json.loads(path.read_text(encoding='utf-8'))['sections']['steel_board']
        relevant = [item for item in documents if set(item.get('products', [])) & {'PROMATECT 100', 'PROMATECT 250'}]
        self.assertFalse(any('handbook' in json.dumps(item).casefold() for item in relevant))
        authorities = [item for item in relevant if item['kind'] == 'Fire assessment — thickness authority']
        self.assertEqual(len(authorities), 1)
        authority = authorities[0]
        self.assertIn('FC15555-01-03', authority['title'])
        self.assertIn('01-02.pdf', authority['url'])
        self.assertEqual(authority['source_sha256'], 'e2b47d71893d5837625a4280aca6ee7939c6c23bf53018e954e54ba32894128e')


if __name__ == '__main__':
    unittest.main()
