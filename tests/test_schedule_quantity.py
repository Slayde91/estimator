"""Repeated members/runs preserve raw lengths, legacy inputs and per-run overlap."""

from copy import deepcopy
from io import BytesIO
import re
import unittest

from openpyxl import Workbook, load_workbook

from estimator.catalog import ValidationError
from estimator.schedule_quantity import quantity, total_length
from estimator.schedule_workbook import _schedule, export_schedule_template, import_schedule_workbook
from estimator.workbook_calculators import (
    _board_product_totals, _missing_schedule_fields, approved_formula_overrides,
    calculate_page, calculate_worksheet, calculator_definition, calculator_engine, normalize_calculator_inputs, source_model,
)
from tests.test_board_product_totals import example_inputs


OMITTED = object()


def duct_inputs(product='CAFCO 300', count=OMITTED, *, length=1.13, wall=1, floor=0):
    values = {'B11': '300 x 200', 'C11': product, 'D11': length,
              'E11': '120/120/120', 'F11': wall, 'G11': floor,
              'H11': 'Internal', 'I11': 'Vertical' if floor else 'Horizontal'}
    if count is not OMITTED:
        values['CM11'] = count
    return {'CALCULATOR': values}


def engine(identity, inputs):
    return calculator_engine(identity, normalize_calculator_inputs(identity, inputs))


class ScheduleQuantityTests(unittest.TestCase):
    def test_private_quantity_helpers_never_become_advanced_display_columns(self):
        for identity, original_last, quantity_column in (('steel_board', 35, 88), ('ductwork', 90, 91)):
            result = calculate_worksheet(identity, {}, 'CALCULATOR', include_advanced=True,
                                         schedule_view={'rows': [source_model(identity)['schedule']['first_row']],
                                                        'offset': 0, 'limit': 1})
            expected = [*range(1, original_last + 1), quantity_column]
            self.assertEqual(result['visible_columns'], expected)
            self.assertEqual(result['application_display_columns'], expected)
            self.assertEqual(result['max_column'], quantity_column)
            self.assertTrue(all(cell['column'] in expected for row in result['rows'] for cell in row['cells']))

    def test_board_section_picker_offers_clean_ids_and_exact_chs_aliases(self):
        from estimator.promat_board_assessment import chs_section_aliases, steel_id_choices
        model = source_model('steel_board')
        cells = next(sheet['cells'] for sheet in model['sheets'] if sheet['name'] == 'STEEL LIBRARY')
        historical = next(cell['value'] for cell in cells.values()
                          if isinstance(cell.get('value'), str) and '[P100]' in cell['value'])
        inputs = {'CALCULATOR': {'C9': 'PROMATECT 100', 'D9': historical}}
        before = deepcopy(inputs)
        page = calculate_page('steel_board', inputs, 'CALCULATOR', 9, 1)
        picker = next(cell for cell in page['rows'][0]['cells'] if cell['address'] == 'D9')
        self.assertEqual(picker['value'], historical)
        self.assertEqual(picker['options'], steel_id_choices(model))
        self.assertFalse(any('[P100]' in option for option in picker['options']))
        aliases = chs_section_aliases(model)
        self.assertTrue(aliases)
        self.assertTrue(set(aliases).issubset(picker['options']))
        self.assertEqual(inputs, before)

    def test_new_fields_are_appended_without_changing_source_length_addresses(self):
        for identity, count, length, total in (('steel_board', 'CJ', 'F', 'CK'), ('ductwork', 'CM', 'D', 'CN')):
            model = source_model(identity)
            self.assertEqual(model['schedule']['quantity_column'], count)
            self.assertEqual(model['schedule']['quantity_length_column'], length)
            self.assertEqual(model['schedule']['total_length_column'], total)
            self.assertIn(count + str(model['schedule']['first_row']) + ':' + count + str(model['schedule']['last_row']), model['input_ranges'][model['schedule']['sheet']])

    def test_spray_quantities_scale_once_and_dft_and_raw_inputs_stay_exact(self):
        for product in ('CAFCO 300', 'MONOKOTE'):
            one_inputs = duct_inputs(product, 1)
            three_inputs = duct_inputs(product, 3)
            before = deepcopy(three_inputs)
            one, three = engine('ductwork', one_inputs), engine('ductwork', three_inputs)
            for column in ('K', 'M', 'AN', 'AE', 'AJ'):
                self.assertAlmostEqual(three.value('CALCULATOR', column + '11'), 3 * one.value('CALCULATOR', column + '11'), places=12, msg=product + column)
            self.assertEqual(three.value('CALCULATOR', 'L11'), one.value('CALCULATOR', 'L11'))
            self.assertEqual(three.value('CALCULATOR', 'D11'), 1.13)
            self.assertEqual(three.value('CALCULATOR', 'F11'), 1)
            self.assertEqual(three_inputs, before)

    def test_wrap_preserves_each_run_overlap_and_repeats_penetration_materials(self):
        for wall, floor in ((1, 0), (0, 1)):
            one = engine('ductwork', duct_inputs('FyreWrap', 1, wall=wall, floor=floor))
            three = engine('ductwork', duct_inputs('FyreWrap', 3, wall=wall, floor=floor))
            for column in ('K', 'V', 'W', 'X', 'N', 'O', 'P', 'Z', 'AA', 'AE', 'AG', 'AI', 'AJ', 'CH', 'CI', 'CJ'):
                expected, actual = one.value('CALCULATOR', column + '11'), three.value('CALCULATOR', column + '11')
                if type(expected) in (int, float):
                    self.assertAlmostEqual(actual, 3 * expected, places=12, msg=column)
                else:
                    self.assertEqual(actual, expected, column)
            self.assertEqual(three.value('CALCULATOR', 'BP11'), one.value('CALCULATOR', 'BP11'))
            self.assertEqual(three.value('CALCULATOR', 'BQ11'), one.value('CALCULATOR', 'BQ11'))

    def test_legacy_omitted_quantity_is_one_without_rewriting_maps(self):
        inputs = duct_inputs('FyreWrap')
        before = deepcopy(inputs)
        legacy, explicit = engine('ductwork', inputs), engine('ductwork', duct_inputs('FyreWrap', 1))
        for column in ('K', 'V', 'W', 'N', 'P', 'AE', 'AJ'):
            self.assertEqual(legacy.value('CALCULATOR', column + '11'), explicit.value('CALCULATOR', column + '11'))
        self.assertEqual(quantity(legacy, 11), 1)
        self.assertEqual(inputs, before)
        self.assertNotIn('CM11', normalize_calculator_inputs('ductwork', inputs)['CALCULATOR'])

    def test_explicit_blank_quantity_withholds_new_duct_quantities(self):
        value = engine('ductwork', duct_inputs('CAFCO 300', None))
        for column in ('K', 'M', 'AN', 'AE', 'AJ'):
            self.assertIn(value.value('CALCULATOR', column + '11'), (None, ''))
        self.assertIn('Quantity', value.value('CALCULATOR', 'J11'))

    def test_board_quantity_uses_raw_length_once_at_full_precision(self):
        inputs = example_inputs(9, length=0.123456789012345)
        inputs['CALCULATOR'].update({'C9': 'TRAFALGAR COREX', 'CJ9': 1})
        one = engine('steel_board', inputs)
        inputs['CALCULATOR']['CJ9'] = 3
        three = engine('steel_board', inputs)
        self.assertEqual(three.value('CALCULATOR', 'F9'), 0.123456789012345)
        self.assertEqual(total_length(three, 9), 0.123456789012345 * 3)
        for column in ('AD', 'AE', 'AF'):
            self.assertAlmostEqual(three.value('CALCULATOR', column + '9'), 3 * one.value('CALCULATOR', column + '9'), places=13)

    def test_new_board_missing_critical_temperature_cannot_order(self):
        for product in ('TRAFALGAR COREX', 'PROMATECT-XS'):
            inputs = example_inputs(9)
            inputs['CALCULATOR'].update({'C9': product, 'J9': None, 'CJ9': 1})
            value = engine('steel_board', inputs)
            self.assertNotEqual(value.value('CALCULATOR', 'AR9'), 'CLADDING ESTIMATE')
            for column in ('AE', 'AF', 'AG'):
                self.assertIn(value.value('CALCULATOR', column + '9'), (None, ''))
            totals = next(row for row in _board_product_totals(value) if row['product'] == product)
            self.assertEqual(totals['incomplete_rows'], 1)
            inputs['CALCULATOR'].pop('CJ9')
            legacy = engine('steel_board', inputs)
            self.assertEqual(legacy.required_missing_by_row, {})
            self.assertGreater(legacy.value('CALCULATOR', 'AE9'), 0)
            self.assertIn('Critical temp (C)', legacy.all_required_missing_by_row[9])

    def test_invalid_explicit_quantity_is_rejected_at_backend_boundary(self):
        for identity, address in (('ductwork', 'CM11'), ('steel_board', 'CJ9')):
            for value in (0, -1, 1.5, True, float('inf'), float('nan'), '2'):
                with self.subTest(identity=identity, value=value), self.assertRaises(ValidationError):
                    normalize_calculator_inputs(identity, {'CALCULATOR': {address: value}})

    def test_blank_padding_stays_blank_and_area_override_exempts_member_quantity(self):
        for identity in ('steel_board', 'ductwork'):
            definition = calculator_definition(identity)
            schedule = definition['schedule']
            result = calculate_page(identity, definition['inputs'], schedule['sheet'], schedule['first_row'], 1)
            cells = {cell['address']: cell for cell in result['rows'][0]['cells']}
            self.assertIsNone(cells[schedule['quantity_column'] + str(schedule['first_row'])]['value'])
            self.assertEqual(result['required_review_rows'], [])
            self.assertFalse(any(cell.get('required_missing') for cell in cells.values()))
        for identity in ('steel_vermiculite', 'steel_intumescent'):
            inputs = {'SCHEDULE': {'B10': 'Nullifire SC902' if identity == 'steel_intumescent' else 'CAFCO 300',
                                   'C10': 'I/H beam - 3 sides', 'D10': 550, 'E10': 'Section', 'H10': 120, 'L10': 3.5}}
            missing = _missing_schedule_fields(source_model(identity), inputs)
            self.assertEqual(missing, {})

    def test_effective_length_ranges_keep_both_endpoints_and_source_strings(self):
        overrides = approved_formula_overrides('steel_board')
        for cells in overrides.values():
            for formula in cells.values():
                self.assertNotRegex(formula, r'CK\$?\d+:\$?F\$?\d+')

    def test_templates_put_quantity_beside_length_and_retain_legacy_imports(self):
        for identity in ('steel_board', 'ductwork'):
            catalog, schedule, fields = _schedule(identity)
            positions = {field['column']: index for index, field in enumerate(fields)}
            self.assertEqual(positions[schedule['quantity_column']], positions[schedule['quantity_length_column']] + 1)
            exported = load_workbook(BytesIO(export_schedule_template(identity)))
            self.assertIn('Quantity', [cell.value for cell in exported[schedule['sheet']][1]])
            exported.close()
            workbook = Workbook(); sheet = workbook.active; sheet.title = schedule['sheet']
            old = [field for field in fields if not field.get('quantity')]
            sheet.append([field['label'] for field in old])
            values = {schedule['quantity_length_column']: 2.5}
            sheet.append([values.get(field['column']) for field in old])
            output = BytesIO(); workbook.save(output)
            result = import_schedule_workbook(identity, output.getvalue(), 'legacy.xlsx')
            self.assertNotIn(schedule['quantity_column'] + str(schedule['first_row']), result['inputs'][schedule['sheet']])
            self.assertEqual(result['inputs'][schedule['sheet']][schedule['quantity_length_column'] + str(schedule['first_row'])], 2.5)


if __name__ == '__main__':
    unittest.main()
