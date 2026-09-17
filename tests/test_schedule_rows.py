"""Dynamic display rows preserve exact data and the complete calculation model."""

from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from estimator.catalog import ValidationError
from estimator.project_file import export_project, load_project_bytes
from estimator.schedule_rows import empty_schedule_inputs, normalize_schedule_rows, populated_schedule_rows, schedule_window
from estimator.schedule_workbook import import_schedule_workbook
from estimator.storage import Store
from estimator.workbook_calculators import calculate_worksheet, calculator_session, source_model


IDS = ('steel_vermiculite', 'steel_board', 'ductwork')


class ScheduleRowsTests(unittest.TestCase):
    def test_new_schedules_are_blank_but_legacy_inputs_keep_examples_and_saved_state(self):
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory) / 'state.sqlite3')
            for identity in IDS:
                schedule = source_model(identity)['schedule']
                state = store.calculator_state(identity)
                self.assertEqual(state['schedule_rows'], [schedule['first_row']])
                self.assertEqual(populated_schedule_rows(identity, state['inputs']), set())
                self.assertTrue(populated_schedule_rows(identity, {}), identity)
                legacy = {'inputs': {}, 'source_sha256': source_model(identity)['source']['sha256']}
                with store.connect() as db:
                    db.execute('INSERT INTO calculator_states VALUES(?,?,?)', (identity, json.dumps(legacy), 'earlier'))
                reopened = store.calculator_state(identity)
                self.assertEqual(reopened['inputs'], {})
                self.assertEqual(reopened['schedule_rows'], normalize_schedule_rows(identity, {}))
                with store.connect() as db:
                    self.assertEqual(json.loads(db.execute('SELECT data FROM calculator_states WHERE id=?', (identity,)).fetchone()[0]), legacy)

    def test_partial_late_rows_zero_and_advanced_inputs_cannot_be_hidden(self):
        for identity in IDS:
            schedule = source_model(identity)['schedule']
            first, last = schedule['first_row'], schedule['last_row']
            inputs = empty_schedule_inputs(identity)
            field = next(field for field in reversed(schedule['columns']) if field.get('editable') and not field.get('generated'))
            inputs[schedule['sheet']][f"{field['column']}{last}"] = 0 if field['type'] == 'number' else 'Late partial row'
            self.assertEqual(normalize_schedule_rows(identity, inputs), list(range(first, last + 1)))
            self.assertEqual(normalize_schedule_rows(identity, inputs, [first]), [first, last])

    def test_invalid_row_metadata_and_windows_are_rejected(self):
        identity = 'ductwork'
        inputs = empty_schedule_inputs(identity)
        first = source_model(identity)['schedule']['first_row']
        for rows in ([], True, 1, {}, [True], [first, first], [first + 1, first], [first - 1], [1011], [float(first)]):
            with self.subTest(rows=rows), self.assertRaises(ValidationError):
                normalize_schedule_rows(identity, inputs, rows)
        for view in ([], True, {'extra': 1}, {'offset': True}, {'offset': -1}, {'offset': 1}, {'limit': 0}, {'limit': 101}, {'limit': 1.5}):
            with self.subTest(view=view), self.assertRaises(ValidationError):
                schedule_window(identity, inputs, view)

    def test_removed_rows_do_not_shift_data_and_blank_added_rows_survive_project_and_local_save(self):
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory) / 'state.sqlite3')
            drafts = {}
            for identity in IDS:
                schedule = source_model(identity)['schedule']
                first, third, blank = schedule['first_row'], schedule['first_row'] + 2, schedule['first_row'] + 3
                inputs = empty_schedule_inputs(identity)
                field = next(field for field in schedule['columns'] if field.get('editable') and field['type'] == 'number')
                inputs[schedule['sheet']][f"{field['column']}{third}"] = 7.123456789012345
                inputs[schedule['sheet']][f"{field['column']}{first + 1}"] = None
                rows = [first, third, blank]
                saved = store.save_calculator_state(identity, inputs, rows)
                self.assertEqual(saved['inputs'], inputs)
                self.assertEqual(store.calculator_state(identity)['schedule_rows'], rows)
                drafts[identity] = {'inputs': inputs, 'schedule_rows': rows}
            loaded = load_project_bytes(store, export_project(store, {'estimate': {}, 'calculators': drafts}))
            for identity, draft in drafts.items():
                self.assertEqual(loaded['calculators'][identity]['inputs'], draft['inputs'])
                self.assertEqual(loaded['calculators'][identity]['schedule_rows'], draft['schedule_rows'])
            legacy = json.loads(export_project(store, {'estimate': {}, 'calculators': drafts}))
            for entry in legacy['calculators'].values():
                entry.pop('schedule_rows')
                entry['inputs'] = {}
            reopened = load_project_bytes(store, json.dumps(legacy).encode())
            for identity in IDS:
                # A legacy empty overlay retains source examples. The ductwork
                # example's Mixed orientation is now the canonical Both alias;
                # it is the only required overlay and no source row moves.
                expected_inputs = {'CALCULATOR': {'I13': 'Both'}} if identity == 'ductwork' else {}
                self.assertEqual(reopened['calculators'][identity]['inputs'], expected_inputs)
                self.assertEqual(legacy['calculators'][identity]['inputs'], {})
                self.assertEqual(reopened['calculators'][identity]['schedule_rows'], normalize_schedule_rows(identity, {}))

    def test_import_extent_ignores_generated_lines_and_preserves_gaps_partial_data(self):
        for identity in IDS:
            schedule = source_model(identity)['schedule']
            fields = [field for field in schedule['columns'] if field.get('editable')]
            if identity == 'steel_vermiculite':
                fields.sort(key=lambda field: field['column'] != 'AA')
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = schedule['sheet']
            sheet.append(['Line', *[field['label'] for field in fields]])
            for row in range(2, 1002):
                sheet.cell(row, 1, row - 1)
            number_column = next(index for index, field in enumerate(fields, 2) if field['type'] == 'number')
            sheet.cell(8, number_column, 7.123456789012345)
            output = BytesIO()
            workbook.save(output)
            result = import_schedule_workbook(identity, output.getvalue(), 'schedule.xlsx')
            self.assertEqual(result['imported_rows'], 1)
            self.assertEqual(result['schedule_rows'], list(range(schedule['first_row'], schedule['first_row'] + 7)))
            sheet.cell(8, number_column).value = None
            output = BytesIO()
            workbook.save(output)
            result = import_schedule_workbook(identity, output.getvalue(), 'schedule.xlsx')
            self.assertEqual(result['schedule_rows'], [schedule['first_row']])
            self.assertEqual(populated_schedule_rows(identity, result['inputs']), set())
            workbook.close()

    def test_bounded_windows_match_full_formula_results_and_keep_offscreen_totals(self):
        for identity in IDS:
            schedule = source_model(identity)['schedule']
            first, last = schedule['first_row'], schedule['last_row']
            # Preserve source examples and add a precise late input. Windows may
            # never constrain the inputs available to source formulas/totals.
            number = next(field for field in schedule['columns'] if field.get('editable') and field['type'] == 'number')
            inputs = {schedule['sheet']: {f"{number['column']}{last}": 7.123456789012345}}
            original = deepcopy(inputs)
            full = calculate_worksheet(identity, inputs, schedule['sheet'])
            view = calculate_worksheet(identity, inputs, schedule['sheet'], schedule_view={'offset': 999, 'limit': 1})
            self.assertEqual(inputs, original)
            self.assertEqual(view['schedule_view']['total_rows'], 1000)
            visible = {row['row']: row for row in view['rows']}
            expected = {row['row']: row for row in full['rows']}
            self.assertNotIn(first, visible)
            self.assertIn(last, visible)
            self.assertLess(len(view['rows']), 30)
            for row, data in visible.items():
                # Shared dropdown IDs are response-local, so compare resolved values.
                for actual, baseline in zip(data['cells'], expected[row]['cells']):
                    self.assertEqual(actual['address'], baseline['address'])
                    self.assertEqual(actual['value'], baseline['value'])
                    if 'options_ref' in actual:
                        self.assertEqual(view['option_sets'][actual['options_ref']], full['option_sets'][baseline['options_ref']])
            for key in ('product_totals', 'board_product_totals'):
                if key in full:
                    self.assertEqual(view[key], full[key])
            _, engine, lock = calculator_session(identity, inputs)
            with lock:
                self.assertEqual(engine.value(schedule['sheet'], f"{number['column']}{last}"), 7.123456789012345)


if __name__ == '__main__':
    unittest.main()
