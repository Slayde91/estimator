"""Schedule templates preserve source input shape and isolate uploaded data."""

from copy import deepcopy
from datetime import datetime
import hashlib
from io import BytesIO
import unittest
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook

from estimator.catalog import ValidationError
from estimator.pricing_workbook import _serialize_exact, export_pricing_workbook
from estimator.schedule_workbook import export_schedule_template, import_schedule_workbook
from estimator.workbook_catalog import load_workbook_catalog
from estimator.workbook_runtime import load_application_catalog


IDS = ("steel_vermiculite", "ductwork", "steel_board")
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def edit(payload, action):
    workbook = load_workbook(BytesIO(payload))
    action(workbook)
    result = _serialize_exact(workbook)
    workbook.close()
    return result


def patch_xml(payload, part, transform):
    output = BytesIO()
    with ZipFile(BytesIO(payload)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            value = source.read(entry.filename)
            if entry.filename == part:
                value = transform(value)
            target.writestr(entry.filename, value)
    return output.getvalue()


class ScheduleWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.templates = {identity: export_schedule_template(identity) for identity in IDS}
        cls.catalogs = {identity: load_application_catalog(identity) for identity in IDS}

    def imported(self, identity="ductwork", payload=None, current=None):
        return import_schedule_workbook(identity, payload or self.templates[identity], "schedule.xlsx", current)

    def test_templates_omit_lines_and_keep_1000_centered_blank_input_rows(self):
        for identity in IDS:
            with self.subTest(identity=identity):
                catalog = self.catalogs[identity]
                schedule = catalog["schedule"]
                columns = [field for field in schedule["columns"]
                           if field["editable"] and not field.get("advanced")]
                if identity == "steel_vermiculite":
                    columns.sort(key=lambda field: field["column"] != "AA")
                workbook = load_workbook(BytesIO(self.templates[identity]))
                self.assertEqual(workbook.sheetnames, [schedule["sheet"], "Instructions"])
                sheet = workbook[schedule["sheet"]]
                self.assertEqual([cell.value for cell in sheet[1]], [field["label"] for field in columns])
                self.assertFalse(any(field.get("advanced") for field in columns))
                self.assertEqual(sheet.max_row, 1001)
                self.assertTrue(all(cell.value is None for row in sheet.iter_rows(min_row=2) for cell in row))
                self.assertTrue(all(cell.alignment.horizontal == "center" and cell.alignment.vertical == "center"
                                    for page in workbook for row in page for cell in row))
                self.assertFalse(any(cell.data_type == "f" for page in workbook for row in page for cell in row))
                self.assertEqual(sheet.freeze_panes, "B2")
                self.assertTrue(sheet.auto_filter.ref)
                self.assertTrue(sheet.data_validations.dataValidation)
                workbook.close()

    def test_blank_import_clears_all_source_examples_and_retains_settings(self):
        for identity in IDS:
            catalog = self.catalogs[identity]
            schedule = catalog["schedule"]
            current = {name: {fields[0]["cell"]: 0.125} for name, fields in catalog["fields"].items() if fields}
            before = deepcopy(current)
            result = self.imported(identity, current=current)
            self.assertEqual(current, before)
            self.assertEqual(result["imported_rows"], 0)
            for name, values in before.items():
                self.assertEqual(result["inputs"][name], values)
            expected = (schedule["last_row"] - schedule["first_row"] + 1) * sum(field["editable"] for field in schedule["columns"])
            self.assertEqual(len(result["inputs"][schedule["sheet"]]), expected)
            self.assertTrue(all(value is None for value in result["inputs"][schedule["sheet"]].values()))

    def test_actual_duct_row_preserves_zero_precision_and_source_mapping(self):
        precise = 0.12345678901234566
        def fill(workbook):
            sheet = workbook["CALCULATOR"]
            for index, value in enumerate(["250x250", "FyreWrap", precise, "120/120/120", 0, 1, "Internal", "Mixed"], 1):
                sheet.cell(2, index, value)
        payload = edit(self.templates["ductwork"], fill)
        result = self.imported(payload=payload, current={"PRODUCT SETTINGS": {"B97": 1.22}})
        self.assertEqual(result["inputs"]["CALCULATOR"]["D11"], precise)
        self.assertEqual(result["inputs"]["CALCULATOR"]["F11"], 0)
        self.assertEqual(result["inputs"]["CALCULATOR"]["G11"], 1)
        self.assertEqual(result["inputs"]["CALCULATOR"]["B11"], "250x250")
        self.assertEqual(result["inputs"]["CALCULATOR"]["I11"], "Both")
        self.assertIsNone(result["inputs"]["CALCULATOR"]["B12"])
        self.assertEqual(result["inputs"]["PRODUCT SETTINGS"], {"B97": 1.22})
        self.assertEqual(result["source_sha256"], hashlib.sha256(payload).hexdigest())
        self.assertEqual(result["imported_rows"], 1)

    def test_rows_keep_positions_and_hidden_final_row_is_imported(self):
        def fill(workbook):
            sheet = workbook["CALCULATOR"]
            sheet["A1001"] = "1500x1000"
            sheet["C1001"] = 0
            sheet.row_dimensions[1001].hidden = True
        result = self.imported(payload=edit(self.templates["ductwork"], fill))
        self.assertEqual(result["imported_rows"], 1)
        self.assertEqual(result["inputs"]["CALCULATOR"]["B1010"], "1500x1000")
        self.assertEqual(result["inputs"]["CALCULATOR"]["D1010"], 0)
        self.assertIsNone(result["inputs"]["CALCULATOR"]["B1009"])

    def test_legacy_templates_keep_source_coordinates_and_clear_new_tail(self):
        for identity in IDS:
            with self.subTest(identity=identity):
                original = load_workbook_catalog(identity)["schedule"]
                fields = [field for field in original["columns"] if field["editable"]]
                book = Workbook()
                sheet = book.active
                sheet.title = original["sheet"]
                sheet.append([field["label"] for field in fields])
                sheet.cell(2, 1, "legacy-first")
                old_capacity = original["last_row"] - original["first_row"] + 1
                sheet.cell(old_capacity + 1, 1, "legacy-last")
                payload = _serialize_exact(book)
                book.close()
                current = {original["sheet"]: {fields[0]["column"] + str(original["first_row"] + 999): "old-tail"}}
                result = self.imported(identity, payload, current)["inputs"][original["sheet"]]
                self.assertEqual(result[fields[0]["column"] + str(original["first_row"])], "legacy-first")
                self.assertEqual(result[fields[0]["column"] + str(original["last_row"])], "legacy-last")
                if old_capacity < 1000:
                    self.assertIsNone(result[fields[0]["column"] + str(original["first_row"] + 999)])
                if identity == "steel_vermiculite":
                    self.assertIsNone(result["AA10"])

    def test_new_steel_location_and_final_row_retain_identity_and_precision(self):
        precise = 0.12345678901234566
        for identity, first, mark_column, location_column, quantity_column, source_quantity in (
                ("steel_vermiculite", 10, 2, 1, 10, "I"),
                ("steel_board", 9, 1, 2, 6, "F")):
            def fill(book):
                sheet = book[self.catalogs[identity]["schedule"]["sheet"]]
                for row in (2, 1001):
                    sheet.cell(row, mark_column, f"MARK-{row}")
                    sheet.cell(row, location_column, f"North / level {row}\nZone B")
                    sheet.cell(row, quantity_column, precise)
            result = self.imported(identity, edit(self.templates[identity], fill))
            self.assertEqual(result["imported_rows"], 2)
            cells = result["inputs"][self.catalogs[identity]["schedule"]["sheet"]]
            source_location = "AA" if identity == "steel_vermiculite" else "B"
            for target, template_row in ((first, 2), (first + 999, 1001)):
                self.assertEqual(cells[f"A{target}"], f"MARK-{template_row}")
                self.assertEqual(cells[f"{source_location}{target}"], f"North / level {template_row}\nZone B")
                self.assertEqual(cells[f"{source_quantity}{target}"], precise)
            self.assertNotIn(f"Z{first}", cells)

    def test_previous_numbered_templates_keep_physical_positions_and_locations(self):
        for identity in IDS:
            def fill(book):
                sheet = book[self.catalogs[identity]["schedule"]["sheet"]]
                sheet.insert_cols(1)
                sheet["A1"] = "Line"
                sheet["A2"] = 1000
                sheet["B2"] = "First location" if identity == "steel_vermiculite" else "First item"
                sheet["A1001"] = 1
                sheet["B1001"] = "Last location" if identity == "steel_vermiculite" else "Last item"
            with self.subTest(identity=identity):
                result = self.imported(identity, edit(self.templates[identity], fill))
                schedule = self.catalogs[identity]["schedule"]
                column = "AA" if identity == "steel_vermiculite" else "B" if identity == "ductwork" else "A"
                cells = result["inputs"][schedule["sheet"]]
                self.assertTrue(cells[column + str(schedule["first_row"])].startswith("First"))
                self.assertTrue(cells[column + str(schedule["last_row"])].startswith("Last"))
                self.assertEqual(result["imported_rows"], 2)

    def test_line_is_informational_but_rejects_malformed_values(self):
        def legacy(book):
            sheet = book["CALCULATOR"]
            sheet.insert_cols(1)
            sheet["A1"] = "Line"
        template = edit(self.templates["ductwork"], legacy)
        for value in (1.5, 0, 1001, "12", "=1+1", True):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                self.imported(payload=edit(template, lambda book: setattr(book["CALCULATOR"]["A2"], "value", value)))
        result = self.imported(payload=template)
        self.assertEqual(result["imported_rows"], 0)
        self.assertTrue(all(value is None for value in result["inputs"]["CALCULATOR"].values()))

    def test_board_advanced_columns_and_warning_values_are_preserved(self):
        # Older exports included all optional M:X fields. They remain accepted
        # even though new templates contain only the normal A:L schedule input.
        fields = [field for field in self.catalogs["steel_board"]["schedule"]["columns"] if field["editable"]]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "CALCULATOR"
        sheet.append([field["label"] for field in fields])
        values = {"A": "M-001", "H": 95, "J": 537.5, "L": 0,
                  "M": "Custom girth", "V": 1.125, "W": 0, "X": "DESIGN-01.2345"}
        for index, field in enumerate(fields, 1):
            if field["column"] in values:
                sheet.cell(2, index, values[field["column"]])
        payload = _serialize_exact(workbook)
        workbook.close()
        result = self.imported("steel_board", payload)
        cells = result["inputs"]["CALCULATOR"]
        self.assertEqual(cells["H9"], 95)
        self.assertEqual(cells["J9"], 537.5)
        self.assertEqual(cells["X9"], "DESIGN-01.2345")
        self.assertEqual(cells["L9"], 0)

    def test_board_template_contains_only_normal_schedule_inputs(self):
        workbook = load_workbook(BytesIO(self.templates["steel_board"]))
        sheet = workbook["CALCULATOR"]
        self.assertEqual([cell.value for cell in sheet[1]], [
            "Member mark", "Location", "Product", "Steel ID", "ESA/M input (m2/t)",
            "Lineal metres", "Sides", "FRL required (min)", "Beam or Column",
            "Critical temp (C)", "Family (ESA input)", "Waste (%)",
        ])
        self.assertEqual(sheet.max_column, 12)
        self.assertNotIn("Exposure layout", [cell.value for cell in sheet[1]])
        self.assertNotIn("Design reference", [cell.value for cell in sheet[1]])
        workbook.close()

    def test_dropdowns_reference_local_lists_and_preserve_dependent_choices(self):
        board = load_workbook(BytesIO(self.templates["steel_board"]))
        rules = {str(rule.sqref).split(":")[0]: rule for rule in board["CALCULATOR"].data_validations.dataValidation}
        self.assertIn('$C2="TRAFALGAR COREX"', rules["H2"].formula1)
        self.assertIn('$I2="Column"', rules["J2"].formula1)
        self.assertNotIn("SETTINGS!", rules["H2"].formula1)
        self.assertEqual(rules["H2"].errorStyle, "warning")
        self.assertGreater(len(board.defined_names), 5)
        self.assertEqual(board["CALCULATOR"]["L2"].number_format, "0.00%")
        self.assertIn("H1001", str(rules["H2"].sqref))
        board.close()
        duct = load_workbook(BytesIO(self.templates["ductwork"]))
        rules = {str(rule.sqref).split(":")[0]: rule for rule in duct["CALCULATOR"].data_validations.dataValidation}
        # The ductwork lists are now closed application choices; unlike the
        # dependent board lists above, they need no copied reference lookup.
        expected = {
            "B2": '"CAFCO 300,MONOKOTE,FyreWrap"',
            "D2": '"60/60/60,90/90/90,120/120/120,180/180/180,240/240/180"',
            "G2": '"Internal,External,Both,Stair pressurisation,Other pressurisation"',
            "H2": '"Horizontal,Vertical,Both"',
        }
        for address, choices in expected.items():
            self.assertEqual(rules[address].formula1, choices)
            self.assertEqual(rules[address].errorStyle, "stop")
            self.assertTrue(rules[address].showErrorMessage)
            self.assertTrue(rules[address].allowBlank)
            self.assertIn(f"{address[0]}1001", str(rules[address].sqref))
        duct.close()

    def test_missing_or_stale_dimensions_do_not_truncate_schedule(self):
        payload = edit(self.templates["ductwork"], lambda book: setattr(book["CALCULATOR"]["A1001"], "value", "250x250"))
        for mode in ("missing", "stale"):
            def change(content):
                root = ET.fromstring(content)
                dimension = root.find(f"{{{NS}}}dimension")
                if mode == "missing":
                    root.remove(dimension)
                else:
                    dimension.set("ref", "A1:H2")
                return ET.tostring(root)
            result = self.imported(payload=patch_xml(payload, "xl/worksheets/sheet1.xml", change))
            self.assertEqual(result["inputs"]["CALCULATOR"]["B1010"], "250x250")

    def test_rejects_headers_extra_columns_rows_and_foreign_sheets(self):
        actions = [lambda book: setattr(book["CALCULATOR"]["A1"], "value", "Wrong header"),
                   lambda book: setattr(book["CALCULATOR"]["J2"], "value", 123),
                   lambda book: setattr(book["CALCULATOR"]["B1002"], "value", "250x250"),
                   lambda book: book.create_sheet("DATABASE")]
        for action in actions:
            with self.assertRaises(ValidationError):
                self.imported(payload=edit(self.templates["ductwork"], action))
        with self.assertRaises(ValidationError):
            self.imported("ductwork", self.templates["steel_board"])

    def test_rejects_formulas_external_links_dates_boolean_errors_and_numeric_text(self):
        for address, value in [("A2", "=1+1"), ("C2", "12.5"), ("C2", True),
                               ("C2", datetime(2026, 9, 13)), ("C2", "#DIV/0!"), ("C2", 1e13)]:
            with self.subTest(value=str(value)), self.assertRaises(ValidationError):
                self.imported(payload=edit(self.templates["ductwork"], lambda book: setattr(book["CALCULATOR"][address], "value", value)))
        with self.assertRaisesRegex(ValidationError, "External"):
            self.imported(payload=edit(self.templates["ductwork"], lambda book: setattr(book["CALCULATOR"]["B2"], "hyperlink", "https://example.com")))

    def test_pricing_workbook_is_rejected_as_wrong_schedule_template(self):
        payload = export_pricing_workbook({})
        for identity in IDS:
            with self.subTest(identity=identity), self.assertRaisesRegex(ValidationError, "schedule template"):
                self.imported(identity, payload)

    def test_rejects_malformed_numeric_xml_and_invalid_file_types(self):
        payload = edit(self.templates["ductwork"], lambda book: setattr(book["CALCULATOR"]["D2"], "value", 12.5))
        invalid = patch_xml(payload, "xl/worksheets/sheet1.xml", lambda value: value.replace(b">12.5<", b">not-a-number<"))
        with self.assertRaisesRegex(ValidationError, "malformed"):
            self.imported(payload=invalid)
        for filename, payload in [("schedule.xlsm", self.templates["ductwork"]), ("schedule.xlsx", b"no zip"),
                                   ("schedule.xlsx", b"x" * (5 * 1024 * 1024 + 1))]:
            with self.assertRaises(ValidationError):
                import_schedule_workbook("ductwork", payload, filename)

    def test_preserves_literal_equals_text_and_ignores_reference_list_changes(self):
        def fill(workbook):
            sheet = workbook["CALCULATOR"]
            sheet["A2"] = "=literal mark"
            sheet["A2"].data_type = "s"
            workbook["Instructions"]["A10"] = "invented reference choice"
        result = self.imported("steel_board", edit(self.templates["steel_board"], fill))
        self.assertEqual(result["inputs"]["CALCULATOR"]["A9"], "=literal mark")
        self.assertNotIn("Instructions", result["inputs"])
        with self.assertRaises(ValidationError):
            self.imported(current={"PRODUCT SETTINGS": {"J137": "invented"}})

    def test_duplicate_rows_or_cells_are_rejected_without_silent_data_loss(self):
        for kind in ("row", "cell"):
            def duplicate(content):
                root = ET.fromstring(content)
                data = root.find(f"{{{NS}}}sheetData")
                row = data.find(f"{{{NS}}}row")
                if kind == "row":
                    data.insert(1, deepcopy(row))
                else:
                    row.insert(1, deepcopy(row[0]))
                return ET.tostring(root)
            with self.subTest(kind=kind), self.assertRaisesRegex(ValidationError, "duplicate"):
                self.imported(payload=patch_xml(self.templates["ductwork"], "xl/worksheets/sheet1.xml", duplicate))


if __name__ == "__main__":
    unittest.main()
