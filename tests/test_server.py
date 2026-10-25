import base64
import csv
import http.client
from io import BytesIO, StringIO
import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from estimator.catalog import ROOT, baseline
from estimator.calculator import labour_breakdown
from estimator.server import create_server
from estimator.report import render_quote_pdf
from estimator.storage import Store
from openpyxl import load_workbook
from test_report import pdf_text


class ServerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.server = create_server(0, Path(cls.temp.name) / "test.sqlite3")
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()
        cls.temp.cleanup()

    def request(self, method, path, body=None, headers=None):
        # Artifact serialization can exceed five seconds on a busy CI worker.
        connection = http.client.HTTPConnection("127.0.0.1", self.server.server_port, timeout=30)
        try:
            supplied = {"Content-Type": "application/json", **(headers or {})}
            connection.request(method, path, body=json.dumps(body) if isinstance(body, dict) else body, headers=supplied)
            response = connection.getresponse()
            return response.status, dict(response.getheaders()), response.read()
        finally:
            connection.close()

    def test_bootstrap_calculation_save_and_reopen_round_trip(self):
        status, _, payload = self.request("GET", "/api/bootstrap")
        self.assertEqual(status, 200)
        bootstrap = json.loads(payload)
        self.assertEqual(len(bootstrap["fields"]), 64)
        self.assertEqual(bootstrap["pricing_usage"]["firestopping"]["label"], "Firestopping Estimator")
        self.assertEqual(bootstrap["pricing_usage"]["firestopping"]["source_field"], "sales_description")
        self.assertIn("collar", bootstrap["pricing_usage"]["firestopping"]["keywords"])
        collar_group = next(group for group in bootstrap["pricing_usage"]["firestopping"]["groups"] if group["key"] == "collars")
        self.assertEqual((collar_group["label"], collar_group["list_column"]), ("Collar Type", "B"))
        self.assertIn("collar", collar_group["keywords"])
        status, _, payload = self.request("POST", "/api/calculate", {"inputs": {"B15": 12.25}})
        self.assertEqual(status, 200)
        total = json.loads(payload)["summary"]["total"]
        status, _, payload = self.request("POST", "/api/quotes", {"title": "Integration quote", "inputs": {"B15": 12.25}})
        self.assertEqual(status, 201)
        quote = json.loads(payload)
        status, _, payload = self.request("GET", "/api/quotes/" + quote["id"])
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(payload)["result"]["summary"]["total"], total)

    def test_cross_site_and_dns_rebinding_requests_are_rejected(self):
        for headers in ({"Host": "attacker.example"}, {"Origin": "https://attacker.example"}, {"Origin": "null"}):
            with self.subTest(headers=headers):
                self.assertEqual(self.request("POST", "/api/configuration", {}, headers)[0], 403)
        self.assertEqual(self.request("PUT", "/api/configuration", "inventory=x", {"Content-Type": "text/plain"})[0], 400)

    def test_download_progress_tracks_real_file_save_and_failure(self):
        from estimator.download_files import DownloadDestination
        destination = DownloadDestination(Path(self.temp.name), 'downloads')
        status, _, body = self.request('POST', '/api/download-progress', {})
        self.assertEqual(status, 200)
        token = json.loads(body)['token']
        header = {'X-Estimator-Download-Progress': token}
        with patch('estimator.project_library.ProjectLibrary.capture_download', return_value=destination):
            status, _, body = self.request('POST', '/api/calculators/steel_board/template',
                                           {'download': {'project_token': None}}, header)
        self.assertEqual(status, 200)
        receipt = json.loads(body)
        self.assertTrue(receipt['saved'])
        self.assertTrue(Path(receipt['path']).is_file())
        status, _, body = self.request('GET', '/api/download-progress/' + token)
        self.assertEqual(json.loads(body), {'percent': 100, 'state': 'saved'})
        self.assertEqual(self.request('POST', '/api/quote-report', {}, header)[0], 400)
        status, _, body = self.request('POST', '/api/download-progress', {})
        failed_token = json.loads(body)['token']
        self.assertEqual(self.request('POST', '/api/calculators/steel_board/template',
                         {'extra': True}, {'X-Estimator-Download-Progress': failed_token})[0], 400)
        status, _, body = self.request('GET', '/api/download-progress/' + failed_token)
        self.assertEqual(json.loads(body)['state'], 'failed')
        self.assertLess(json.loads(body)['percent'], 100)

    def test_a_second_estimator_cannot_share_the_live_port(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(OSError):
                create_server(self.server.server_port, Path(directory) / "second.sqlite3")

    def test_legacy_labour_response_projects_saved_cells_without_rewriting_quote(self):
        store = Store(Path(self.temp.name) / "test.sqlite3")
        original = store.save_quote({"title": "Legacy labour days", "inputs": {"B15": 27.125, "B27": .15, "F26": 1.5, "F27": 2}})
        original["result"].pop("labour")
        before = json.dumps(original, allow_nan=False)
        with store.connect() as db:
            db.execute("UPDATE quotes SET data=? WHERE id=?", (before, original["id"]))
        with patch("estimator.server.calculate", side_effect=AssertionError("Legacy result recalculated")), patch("estimator.storage.calculate", side_effect=AssertionError("Legacy result recalculated")):
            status, _, body = self.request("GET", f'/api/quotes/{original["id"]}')
        self.assertEqual(status, 200)
        response = json.loads(body)
        self.assertEqual(response["result"]["labour"], labour_breakdown(original["result"]))
        self.assertEqual({key: value for key, value in response["result"].items() if key != "labour"}, original["result"])
        self.assertEqual(response["configuration"], original["configuration"])
        with store.connect() as db:
            self.assertEqual(db.execute("SELECT data FROM quotes WHERE id=?", (original["id"],)).fetchone()[0], before)

    def test_bad_json_and_unknown_cells_are_rejected(self):
        for body in ("[1]", '{"inputs":{"B15":NaN}}', "{", {"inputs": {"F7": 0}}, {"inputs": {"B15": 10**400}}):
            self.assertEqual(self.request("POST", "/api/calculate", body)[0], 400)

    def test_only_allowlisted_static_files_are_served(self):
        for path in ("/.runtime/estimator.sqlite3", "/../README.md", "/data/baseline.json", "/%2e%2e/README.md"):
            self.assertEqual(self.request("GET", path)[0], 404)
        status, headers, body = self.request("GET", "/")
        self.assertEqual(status, 200)
        self.assertIn(b"Ceasefire", body)
        self.assertIn("frame-ancestors 'none'", headers["Content-Security-Policy"])
        status, headers, body = self.request("GET", "/fonts/Montserrat-Variable.ttf")
        self.assertEqual(status, 200)
        self.assertEqual(headers["Content-Type"], "font/ttf")
        self.assertEqual(body, (ROOT / "static/fonts/Montserrat-Variable.ttf").read_bytes())
        self.assertEqual(self.request("GET", "/fonts/OFL.txt")[0], 404)
        status, headers, body = self.request("GET", "/icons/document.png")
        self.assertEqual(status, 200)
        self.assertEqual(headers["Content-Type"], "image/png")
        self.assertEqual(body, (ROOT / "static/icons/document.png").read_bytes())
        status, headers, body = self.request("GET", "/library-detail-text.js")
        self.assertEqual(status, 200)
        self.assertEqual(body, (ROOT / "static/library-detail-text.js").read_bytes())

    def test_help_and_signature_assets_have_exact_bytes_and_content_types(self):
        for path, content_type in (
            ('/help-glossary.js', 'text/javascript; charset=utf-8'),
            ('/help-glossary.css', 'text/css; charset=utf-8'),
            ('/takeoff-signatures.js', 'text/javascript; charset=utf-8'),
            ('/takeoff-graphics.js', 'text/javascript; charset=utf-8'),
            ('/takeoff-graphics-ui.js', 'text/javascript; charset=utf-8'),
            ('/takeoff-pdf-tools.js', 'text/javascript; charset=utf-8'),
            ('/takeoff-pdf-tools.css', 'text/css; charset=utf-8'),
            ('/icons/default-memory.svg', 'image/svg+xml'),
            ('/icons/signature.svg', 'image/svg+xml'),
        ):
            with self.subTest(path=path):
                status, headers, body = self.request('GET', path)
                self.assertEqual(status, 200)
                self.assertEqual(headers['Content-Type'], content_type)
                self.assertEqual(body, (ROOT / 'static' / path.lstrip('/')).read_bytes())

    def test_custom_cursor_assets_are_served_as_exact_png_bytes(self):
        for path in sorted((ROOT / 'static/cursors').glob('*.png')):
            with self.subTest(asset=path.name):
                status, headers, body = self.request('GET', '/cursors/' + path.name)
                self.assertEqual(status, 200)
                self.assertEqual(headers['Content-Type'], 'image/png')
                self.assertEqual(body, path.read_bytes())

    def test_header_menu_sounds_have_exact_bytes_and_audio_mime_without_wider_file_access(self):
        import hashlib
        expected = {
            'mgs1-menu-move.mp3': (8313, 'ac2d686c78fa5dabc2415a0f3b39f4ced124f5fc1613ba3918e06daa7b47f8d0'),
            'mgs-confirm2.mp3': (32391, 'e9e7e156e24a4fde706f922ee135e88df51dc5ba94eed6344e66bbd6298e3e86'),
            'new-explosion.mp3': (73398, '1ed5c489c6e7ab08caacf1cba198e7aae07cb7a904be51a4800c56eb32512c21'),
            'hover.mp3': (5015, '052d0d7f3fb41ddb67766dedd8f4c237b46c8b18745dfad0a73a14c8afb6fad5'),
            'select.mp3': (11520, '60e089eeac138ce5301fb59dc9d649c0404a556499c8f78c15c81a7cdfde1969'),
            'select-2.mp3': (192768, '37e5c8645aa76e11dc8039545940436ea9201e533b1c946fbc9dfc3845efeb7b'),
            'mixkit-vintage-dot-printing-1371.wav': (249332, '5c39082dfe0aa5c24cae5faa0ba616bfb1354d8b7686b0b0784fa32689a9decf'),
        }
        for name, (size, digest) in expected.items():
            with self.subTest(name=name):
                status, headers, body = self.request('GET', '/sounds/' + name)
                self.assertEqual(status, 200)
                self.assertEqual(headers['Content-Type'], 'audio/wav' if name.endswith('.wav') else 'audio/mpeg')
                self.assertEqual(headers['X-Content-Type-Options'], 'nosniff')
                self.assertEqual(body, (ROOT / 'static/sounds' / name).read_bytes())
                self.assertEqual(len(body), size)
                self.assertEqual(hashlib.sha256(body).hexdigest(), digest)
                self.assertNotIn('unsafe-inline', headers['Content-Security-Policy'])
        status, headers, body = self.request('GET', '/header-menu-sounds.js')
        self.assertEqual(status, 200)
        self.assertEqual(headers['Content-Type'], 'text/javascript; charset=utf-8')
        self.assertEqual(body, (ROOT / 'static/header-menu-sounds.js').read_bytes())
        for route in ('/sounds/other.mp3', '/sounds/../app.js', '/sounds/%2e%2e/app.js', '/sounds/'):
            with self.subTest(route=route):
                self.assertEqual(self.request('GET', route)[0], 404)

    def test_official_logo_is_served_unchanged_with_image_content_type(self):
        status, headers, body = self.request("GET", "/ceasefire-logo.png")
        self.assertEqual(status, 200)
        self.assertEqual(headers["Content-Type"], "image/png")
        self.assertEqual(body, (ROOT / "static/ceasefire-logo.png").read_bytes())

    def test_parent_control_icons_are_served_as_exact_png_files(self):
        for name in ('takeoff-add-service.png', 'takeoff-add-barrier.png', 'takeoff-transfer.png'):
            with self.subTest(name=name):
                status, headers, body = self.request('GET', '/icons/' + name)
                self.assertEqual(status, 200)
                self.assertEqual(headers['Content-Type'], 'image/png')
                self.assertEqual(body, (ROOT / 'static/icons' / name).read_bytes())
                self.assertEqual(int(headers['Content-Length']), len(body))
                self.assertEqual(self.request('GET', '/icons/' + name + '.bak')[0], 404)

    def test_current_pdf_is_downloadable_and_does_not_save_or_change_prices(self):
        quotes_before = self.request("GET", "/api/quotes")[2]
        pricing_before = self.request("GET", "/api/configuration")[2]
        status, headers, payload = self.request("POST", "/api/quote-report", {"title": 'Duct "/..\\\r\nreport', "inputs": {"B15": 12.25}})
        self.assertEqual(status, 200)
        self.assertEqual(headers["Content-Type"], "application/pdf")
        self.assertEqual(headers["Content-Disposition"], 'attachment; filename="CEASEFIRE-Estimate.pdf"')
        self.assertEqual(int(headers["Content-Length"]), len(payload))
        self.assertEqual(headers["Cache-Control"], "no-store")
        text = pdf_text(payload)
        self.assertIn("Current unsaved estimate", text)
        self.assertIn("Current estimate", text)
        self.assertEqual(self.request("GET", "/api/quotes")[2], quotes_before)
        self.assertEqual(self.request("GET", "/api/configuration")[2], pricing_before)

    def test_saved_pdf_uses_its_snapshot_even_when_catalogue_changes(self):
        _, _, body = self.request("POST", "/api/quotes", {"title": "Historical PDF", "inputs": {"B15": 19.75}})
        quote = json.loads(body)
        quotes_before = self.request("GET", "/api/quotes")[2]
        pricing_before = self.request("GET", "/api/configuration")[2]
        with patch("estimator.catalog.baseline", side_effect=AssertionError("Saved report read today's catalogue")), patch("estimator.storage.calculate", side_effect=AssertionError("Saved report recalculated")):
            status, headers, payload = self.request("GET", f'/api/quotes/{quote["id"]}/report.pdf')
        self.assertEqual(status, 200)
        self.assertEqual(headers["Content-Disposition"], 'attachment; filename="CEASEFIRE-Estimate.pdf"')
        self.assertEqual(headers["Content-Type"], "application/pdf")
        self.assertEqual(int(headers["Content-Length"]), len(payload))
        self.assertEqual(headers["Cache-Control"], "no-store")
        text = pdf_text(payload)
        self.assertIn("Saved quote", text)
        self.assertIn("Historical PDF", text)
        self.assertIn(quote["id"], text)
        self.assertIn(f'${quote["result"]["summary"]["total"]:,.2f}', text)
        self.assertEqual(self.request("GET", f'/api/quotes/{quote["id"]}')[2], body)
        self.assertEqual(self.request("GET", "/api/quotes")[2], quotes_before)
        self.assertEqual(self.request("GET", "/api/configuration")[2], pricing_before)

    def test_pdf_filename_is_independent_of_unicode_and_header_like_titles(self):
        for title in ("\u4e2d\u6587 \U0001f9ef", '../../"\r\nX-Injected: yes'):
            with self.subTest(title=title):
                status, _, saved = self.request("POST", "/api/quotes", {"title": title})
                self.assertEqual(status, 201)
                quote = json.loads(saved)
                self.assertEqual(quote["title"], title)
                quotes_before = self.request("GET", "/api/quotes")[2]
                pricing_before = self.request("GET", "/api/configuration")[2]
                for method, path, body in (("POST", "/api/quote-report", {"title": title}),
                                           ("GET", f'/api/quotes/{quote["id"]}/report.pdf', None)):
                    with self.subTest(method=method), patch("estimator.report.render_quote_pdf", return_value=b"%PDF-test") as renderer:
                        status, headers, payload = self.request(method, path, body)
                    self.assertEqual(status, 200)
                    self.assertEqual(headers["Content-Disposition"], 'attachment; filename="CEASEFIRE-Estimate.pdf"')
                    self.assertNotIn("X-Injected", headers)
                    self.assertEqual(renderer.call_args.args[0]["title"], title)
                    self.assertEqual(payload, b"%PDF-test")
                self.assertEqual(self.request("GET", f'/api/quotes/{quote["id"]}')[2], saved)
                self.assertEqual(self.request("GET", "/api/quotes")[2], quotes_before)
                self.assertEqual(self.request("GET", "/api/configuration")[2], pricing_before)

    def test_edited_saved_preview_preserves_source_lineage_without_saving(self):
        _, _, body = self.request("POST", "/api/quotes", {"title": "Original PDF pricing", "inputs": {"B15": 10}})
        quote = json.loads(body)
        changed = baseline()
        changed["sources"]["quote"]["sha256"] = "later-source-hash"
        with patch("estimator.catalog.baseline", return_value=changed), patch("estimator.report.render_quote_pdf", wraps=render_quote_pdf) as renderer:
            status, headers, payload = self.request("POST", "/api/quote-report", {"source_quote_id": quote["id"], "title": "Unsaved changes", "inputs": {"B15": 12}, "configuration": quote["configuration"]})
        self.assertEqual(status, 200)
        self.assertEqual(headers["Content-Disposition"], 'attachment; filename="CEASEFIRE-Estimate.pdf"')
        self.assertEqual(renderer.call_args.args[0]["source_hashes"], quote["source_hashes"])
        self.assertNotIn("later-source-hash", pdf_text(payload))
        self.assertEqual(self.request("GET", f'/api/quotes/{quote["id"]}')[2], body)

    def test_excel_import_is_a_draft_until_save_and_replaces_choices(self):
        previous = json.loads(self.request("GET", "/api/configuration")[2])
        try:
            self.assertEqual(self.request("PUT", "/api/configuration", {"inventory": {}, "rates": {}})[0], 200)
            _, _, original = self.request("POST", "/api/quotes", {"title": "Before library replacement", "inputs": {"B15": 12.25}})
            quote = json.loads(original)
            status, headers, exported = self.request("POST", "/api/pricing/export", {})
            self.assertEqual(status, 200)
            self.assertIn("spreadsheetml.sheet", headers["Content-Type"])
            workbook = load_workbook(BytesIO(exported))
            sheet = workbook["Inventory & Rates"]
            columns = {cell.value: cell.column for cell in sheet[1]}
            records = list(sheet.iter_rows(min_row=2, values_only=True))
            self.assertNotIn("Row type", columns)
            self.assertEqual(len(records), 417)
            self.assertEqual(len({row[columns["Inventory ID"] - 1] for row in records}), 417)
            use_columns = ("Group", "Selection name", "Price source", "Sell rate",
                           "Yield type", "Use yields", "Rate ID", "Use order")

            def read_uses(record):
                values = {name: record[columns[name] - 1] for name in use_columns}
                if all(value in (None, "") for value in values.values()):
                    return {name: [] for name in use_columns}
                result = {name: next(csv.reader(StringIO(value, newline=""), delimiter=";", strict=True))
                          if isinstance(value, str) and value else [value]
                          for name, value in values.items()}
                self.assertEqual(len({len(entries) for entries in result.values()}), 1)
                return result

            self.assertEqual(sum(len(read_uses(record)["Rate ID"]) for record in records), 166)
            # Remove one choice from the eight editable lists; units are derived.
            removed_choice = None
            for number, record in enumerate(records, 2):
                uses = read_uses(record)
                for index, (group, name) in enumerate(zip(uses["Group"], uses["Selection name"])):
                    if group != "sprays" or name != "Promat Cafco 300":
                        continue
                    self.assertIsNone(removed_choice)
                    removed_choice = (record[columns["Inventory ID"] - 1], uses["Rate ID"][index])
                    for heading, entries in uses.items():
                        entries.pop(index)
                        value = None
                        if len(entries) == 1 and not isinstance(entries[0], str):
                            value = entries[0]
                        elif entries:
                            encoded = StringIO(newline="")
                            csv.writer(encoded, delimiter=";").writerow(entries)
                            value = encoded.getvalue()[:-2]
                        sheet.cell(number, columns[heading]).value = value
                    break
            self.assertIsNotNone(removed_choice)
            record = {"Inventory ID": "qa-new-product", "Item code": "QA-001",
                      "Product/Service": "Replacement spray", "Product name": "Replacement spray", "Sales description": "Replacement spray",
                      "Pricing mode": "Supplier markup", "Supplier price": 100, "Markup": .25,
                      "Sell price": 125, "Rate ID": "qa-new-rate", "Group": "sprays",
                      "Selection name": "Replacement spray", "Price source": "Inventory", "Sell rate": 125,
                      "Yield type": "Not used", "Use order": 999}
            sheet.append([record.get(header) for header in columns])
            stream = BytesIO()
            workbook.save(stream)
            workbook.close()
            pricing_before = self.request("GET", "/api/configuration")[2]
            status, _, body = self.request("POST", "/api/pricing/import", {"filename": "replacement.xlsx", "content_base64": base64.b64encode(stream.getvalue()).decode()})
            self.assertEqual(status, 200, body)
            proposed = json.loads(body)
            self.assertEqual(proposed["summary"]["inventory"]["added"], 1)
            self.assertEqual(proposed["summary"]["inventory"]["removed"], 0)
            self.assertEqual(proposed["summary"]["rates"]["added"], 1)
            self.assertEqual(proposed["summary"]["rates"]["removed"], 1)
            catalog = proposed["configuration"]["catalog"]
            self.assertIn(removed_choice[0], {item["id"] for item in catalog["inventory"]})
            self.assertNotIn(removed_choice[1], {rate["id"] for rates in catalog["rate_groups"].values() for rate in rates})
            self.assertEqual(self.request("GET", "/api/configuration")[2], pricing_before)
            self.assertEqual(self.request("PUT", "/api/configuration", proposed["configuration"])[0], 200)
            self.assertEqual(Store(Path(self.temp.name) / "test.sqlite3").configuration(), proposed["configuration"])
            _, _, body = self.request("GET", "/api/bootstrap")
            field = next(field for field in json.loads(body)["fields"] if field["cell"] == "D15")
            self.assertIn("Replacement spray", field["options"])
            self.assertNotIn("Promat Cafco 300", field["options"])
            status, _, body = self.request("POST", "/api/calculate", {"inputs": {"D15": "Replacement spray", "B15": 10}})
            self.assertEqual(status, 200)
            result = json.loads(body)
            self.assertEqual(result["cells"]["A63"], 125)
            self.assertIsInstance(result["summary"]["total"], (int, float))
            # Existing quote owns the removed product, prices and dropdowns.
            self.assertEqual(self.request("GET", f'/api/quotes/{quote["id"]}')[2], original)
            status, _, body = self.request("POST", "/api/calculate", {"inputs": quote["inputs"], "configuration": quote["configuration"]})
            self.assertEqual(status, 200)
            self.assertEqual(json.loads(body)["cells"], quote["result"]["cells"])
            self.assertEqual(self.request("GET", f'/api/quotes/{quote["id"]}/report.pdf')[0], 200)
        finally:
            self.assertEqual(self.request("PUT", "/api/configuration", previous)[0], 200)

    def test_bad_pricing_uploads_do_not_change_configuration(self):
        before = self.request("GET", "/api/configuration")[2]
        for body in ({}, {"filename": "rates.xlsm", "content_base64": "QQ=="},
                     {"filename": "rates.xlsx", "content_base64": "broken!"},
                     {"filename": "rates.xlsx", "content_base64": ""},
                     {"filename": "rates.xlsx", "content_base64": "QQ=="}):
            with self.subTest(body=body):
                self.assertEqual(self.request("POST", "/api/pricing/import", body)[0], 400)
                self.assertEqual(self.request("GET", "/api/configuration")[2], before)
        self.assertEqual(self.request("POST", "/api/pricing/export", {}, {"Origin": "https://attacker.example"})[0], 403)

    def test_estimate_details_name_summary_and_pdf_round_trip(self):
        data = {"title": "Name is generated", "project_no": " CF-123 ", "client": " Example Client ",
                "site_address": " 42 Test Street ", "workflow": "Fire wrap to ductwork",
                "inputs": {"D15": "Trafalgar FyreWRAP 610", "B15": 12.34567}}
        status, _, payload = self.request("POST", "/api/quotes", data)
        self.assertEqual(status, 201, payload)
        quote = json.loads(payload)
        self.assertEqual(quote["title"], "CF-123- Example Client- 42 Test Street")
        self.assertEqual(quote["client"], "Example Client")
        self.assertEqual(quote["inputs"]["B15"], 12.34567)
        self.assertIn("Fire wrap to ductwork", quote["work_summary"])
        self.assertIn("Trafalgar FyreWRAP 610", quote["work_summary"])
        self.assertNotIn("12.34567", quote["work_summary"])
        self.assertEqual(self.request("GET", f'/api/quotes/{quote["id"]}')[2], payload)
        status, _, pdf = self.request("GET", f'/api/quotes/{quote["id"]}/report.pdf')
        self.assertEqual(status, 200)
        content = pdf_text(pdf)
        for text in ("CF-123", "Example Client", "42 Test Street", "Trafalgar FyreWRAP 610"):
            self.assertIn(text, content)
        self.assertNotIn("Work summary", content)
        status, _, edited = self.request("PUT", f'/api/quotes/{quote["id"]}', {
            "client": "Updated Client", "inputs": quote["inputs"], "workflow": quote["workflow"],
        })
        self.assertEqual(status, 200, edited)
        self.assertEqual(json.loads(edited)["title"], "CF-123- Updated Client- 42 Test Street")

    def test_work_summary_reacts_to_workflow_and_products_without_changing_calculation(self):
        inputs = {"D15": "Promat Cafco 300", "B15": 2.34567}
        status, _, first = self.request("POST", "/api/calculate", {"inputs": inputs, "workflow": "Intumescent spray to slabs"})
        self.assertEqual(status, 200)
        status, _, second = self.request("POST", "/api/calculate", {"inputs": inputs, "workflow": "Intumescent spray to walls"})
        self.assertEqual(status, 200)
        first, second = json.loads(first), json.loads(second)
        self.assertEqual(first["cells"], second["cells"])
        self.assertIn("Intumescent spray to slabs", first["work_summary"])
        self.assertIn("Intumescent spray to walls", second["work_summary"])
        for invalid in ({"workflow": {}}, {"workflow": "x" * 201}, {"work_summary": "Invented work"}):
            self.assertEqual(self.request("POST", "/api/calculate", invalid)[0], 400)

    def test_pdf_requests_validate_inputs_and_respect_origin_boundary(self):
        for data in ({"title": ""}, {"title": "Invalid", "inputs": {"F7": 99}}, {"title": "Invalid", "source_quote_id": 123}):
            self.assertEqual(self.request("POST", "/api/quote-report", data)[0], 400)
        self.assertEqual(self.request("POST", "/api/quote-report", {"title": "Blocked"}, {"Origin": "https://attacker.example"})[0], 403)
        self.assertEqual(self.request("GET", "/api/quotes/missing/report.pdf")[0], 404)


if __name__ == "__main__":
    unittest.main()
