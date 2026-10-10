"""Surface topology, scale-squared measurements and area authority boundaries."""

from copy import deepcopy
import csv
from io import BytesIO, StringIO
import json
import math
import unittest
from uuid import uuid4

from openpyxl import load_workbook

from estimator.catalog import ValidationError
from estimator.takeoff_area import MAX_AREA_EXCLUSIONS, MAX_AREA_VERTICES, ring_area, validate_polygon
from estimator.takeoff_model import item_digest, item_result, validate_snapshot
from estimator.takeoff_transfer import mapped_values
from tests import test_takeoff_workspace as fixtures


class TakeoffAreaTests(unittest.TestCase):
    def setUp(self):
        self.case = fixtures.TakeoffWorkspaceTests()
        self.case.setUp()
        self.addCleanup(self.case.doCleanups)

    def proposal(self, mode='wall'):
        case = self.case
        return {'mode': mode, 'quantity': 1,
                'geometry': {'kind': 'polygon', 'document_id': case.doc['id'], 'page': 1,
                             'points': [[20, 30], [120, 30], [120, 80], [20, 80]],
                             'exclusions': [{'id': str(uuid4()), 'points': [[40, 40], [60, 40], [60, 50], [40, 50]], 'note': 'Door opening'}]},
                'measurement': {'method': 'calibrated', 'calibration_id': case.calibration_id},
                'fields': {'mark': 'W-01' if mode == 'wall' else 'S-01', 'treatment': 'Nominated board treatment',
                           'substrate': 'Concrete', 'frl': '120/120/120',
                           'surface_basis': 'wall-face' if mode == 'wall' else 'slab-soffit',
                           'surface_citation': 'Elevation A shows the true treatment face'},
                'evidence': [{'document_id': case.doc['id'], 'page': 1, 'fields': ['surface_basis', 'net_area_m2'], 'note': 'Elevation A'}]}

    def create(self, proposal=None):
        self.case.command('create_item', item=proposal or self.proposal())
        return self.case.state['snapshot']['items'][-1]['id']

    def current(self, identity):
        return next(item for item in self.case.state['snapshot']['items'] if item['id'] == identity)

    def test_calibrated_net_area_uses_scale_squared_and_ignores_render_transforms(self):
        for mode in ('wall', 'slab'):
            identity = self.create(self.proposal(mode))
            result = next(value for value in self.case.state['item_results'] if value['id'] == identity)
            self.assertEqual(result, {'id': identity, 'gross_area_m2': 50, 'excluded_area_m2': 2, 'net_area_m2': 48,
                                      'layers': 1, 'total_area_m2': 48, 'issues': []})
            original = deepcopy(self.case.state['snapshot'])
            for rotation in (0, 90, 180, 270):
                for user_unit in (1, 2, 5):
                    state = deepcopy(original)
                    state['documents'][0]['pages'][0].update(rotation=rotation, user_unit=user_unit)
                    self.assertEqual(item_result(self.current(identity), state)['net_area_m2'], 48)

    def test_area_is_invariant_under_winding_and_translation(self):
        concave = [[0, 0], [10, 0], [10, 2], [2, 2], [2, 10], [0, 10]]
        self.assertEqual(ring_area(concave), 36)
        self.assertEqual(ring_area(list(reversed(concave))), 36)
        self.assertEqual(ring_area([[x+1e10, y-1e10] for x, y in concave]), 36)
        proposed = self.proposal(); proposed['geometry']['points'].reverse()
        proposed['geometry']['exclusions'][0]['points'].reverse()
        identity = self.create(proposed)
        self.assertEqual(item_result(self.current(identity), self.case.state['snapshot'])['net_area_m2'], 48)

    def test_multiple_scales_are_independent_and_revision_invalidates_dependents(self):
        first = self.create()
        calibration = deepcopy(self.case.state['snapshot']['calibrations'][0])
        calibration.update(id=str(uuid4()), distance_m=5, name='Inset scale')
        self.case.command('add_calibration', calibration=calibration)
        proposed = self.proposal('slab'); proposed['measurement']['calibration_id'] = calibration['id']
        second = self.create(proposed)
        self.assertEqual(item_result(self.current(first), self.case.state['snapshot'])['net_area_m2'], 48)
        self.assertEqual(item_result(self.current(second), self.case.state['snapshot'])['net_area_m2'], 12)
        self.case.confirm(first); self.case.confirm(second)
        self.case.command('update_calibration', calibration_id=self.case.calibration_id, changes={'distance_m': 20})
        self.assertEqual(self.current(first)['state'], 'draft')
        self.assertEqual(self.current(second)['state'], 'confirmed')
        self.assertEqual(item_result(self.current(first), self.case.state['snapshot'])['net_area_m2'], 192)

    def test_invalid_polygon_topologies_fail_atomically(self):
        variants = [
            [[20, 30], [120, 80], [20, 80], [120, 30]],
            [[20, 30], [120, 30], [120, 80], [20, 30]],
            [[20, 30], [120, 30], [70, 30], [70, 80], [20, 80]],
            [[20, 30], [70, 30], [120, 30]],
            [[20, 30], [120, 30], [120, 80], [70, 30], [20, 80]],
        ]
        before = deepcopy(self.case.state['snapshot'])
        for outer in variants:
            proposed = self.proposal(); proposed['geometry'].update(points=outer, exclusions=[])
            with self.subTest(outer=outer), self.assertRaises(ValidationError):
                self.create(proposed)
            self.assertEqual(self.case.service.get(self.case.sid)['snapshot'], before)

    def test_exclusions_cannot_escape_touch_overlap_or_nest(self):
        variants = [
            [[10, 30], [30, 30], [30, 40], [10, 40]],
            [[20, 40], [30, 40], [30, 50], [20, 50]],
            [[110, 40], [130, 40], [130, 50], [110, 50]],
            [[45, 42], [50, 42], [50, 48], [45, 48]],
            [[55, 42], [70, 42], [70, 48], [55, 48]],
            [[60, 42], [70, 42], [70, 48], [60, 48]],
        ]
        for ring in variants:
            proposed = self.proposal()
            proposed['geometry']['exclusions'].append({'id': str(uuid4()), 'points': ring, 'note': 'Invalid second exclusion'})
            with self.subTest(ring=ring), self.assertRaises(ValidationError):
                self.create(proposed)

    def test_concave_outer_rejects_hole_crossing_boundary_with_inside_first_vertex(self):
        proposed = self.proposal()
        proposed['geometry']['points'] = [[20, 30], [120, 30], [120, 50], [40, 50], [40, 100], [20, 100]]
        proposed['geometry']['exclusions'][0]['points'] = [[25, 35], [80, 35], [80, 75], [25, 75]]
        with self.assertRaisesRegex(ValidationError, 'touch or cross'):
            self.create(proposed)

    def test_vertex_and_exclusion_budgets_precede_quadratic_topology_work(self):
        proposed = self.proposal(); geometry = proposed['geometry']
        geometry['points'] = [[70+30*math.cos(2*math.pi*i/MAX_AREA_VERTICES), 60+20*math.sin(2*math.pi*i/MAX_AREA_VERTICES)] for i in range(MAX_AREA_VERTICES)]
        with self.assertRaisesRegex(ValidationError, 'vertices in total'):
            self.create(proposed)
        proposed = self.proposal()
        proposed['geometry']['exclusions'] = [deepcopy(proposed['geometry']['exclusions'][0]) for _ in range(MAX_AREA_EXCLUSIONS+1)]
        with self.assertRaisesRegex(ValidationError, 'explicit exclusions'):
            self.create(proposed)

    def test_area_quantity_and_measurement_cannot_manufacture_surface_multipliers(self):
        for quantity in (None, 0, 2, 1.0, True):
            proposed = self.proposal(); proposed['quantity'] = quantity
            with self.subTest(quantity=quantity), self.assertRaises(ValidationError):
                self.create(proposed)
        proposed = self.proposal(); proposed['measurement'] = {'method': 'cited', 'length_m': 10, 'citation': 'Unrelated length'}
        with self.assertRaisesRegex(ValidationError, 'cited length'):
            self.create(proposed)
        proposed = self.proposal(); proposed['geometry']['points'][0] = [0, 0]
        with self.assertRaisesRegex(ValidationError, 'outside'):
            self.create(proposed)

    def test_confirmation_requires_actual_surface_basis_and_complete_human_fields(self):
        proposed = self.proposal()
        proposed['fields'].update(surface_basis='wall-footprint', substrate='', treatment=' ', surface_citation='', frl='120')
        identity = self.create(proposed)
        issues = item_result(self.current(identity), self.case.state['snapshot'])['issues']
        self.assertIn('INVALID_SURFACE_BASIS', {issue['code'] for issue in issues})
        self.assertIn('INVALID_FRL', {issue['code'] for issue in issues})
        self.case.command('review_items', item_ids=[identity])
        with self.assertRaises(ValidationError):
            self.case.command('confirm_items', item_ids=[identity])
        self.case.command('update_item', item_id=identity, changes={'fields': self.proposal()['fields']})
        self.case.confirm(identity)
        checks = self.current(identity)['confirmation']['checks']
        self.assertEqual(checks, {'engine': 'takeoffs-area-v1', 'quantity': 1, 'gross_area_m2': 50,
                                'excluded_area_m2': 2, 'net_area_m2': 48, 'evidence_verified': True, 'issues': []})

    def test_confirmation_binds_exclusions_and_registry_rejects_forged_area_receipts(self):
        identity = self.create(); self.case.confirm(identity)
        original = deepcopy(self.case.state['snapshot'])
        self.assertEqual(self.case.service.open(original, source_path='verified.json')['snapshot']['items'][0]['state'], 'confirmed')
        forged = deepcopy(original); forged['items'][0]['confirmation']['checks'].update(net_area_m2=49, gross_area_m2=51)
        self.assertEqual(self.case.service.open(forged, source_path='verified.json')['snapshot']['items'][0]['state'], 'reviewed')
        inconsistent = deepcopy(original); inconsistent['items'][0]['confirmation']['checks']['net_area_m2'] = 49
        with self.assertRaisesRegex(ValidationError, 'exactly equal'):
            validate_snapshot(inconsistent)
        changed = deepcopy(self.current(identity)['geometry'])
        excluded_id = changed['exclusions'][0]['id']; changed['exclusions'][0]['points'][1][0] = 70
        self.case.command('update_item', item_id=identity, changes={'geometry': changed})
        self.assertEqual(self.current(identity)['geometry']['exclusions'][0]['id'], excluded_id)
        self.assertEqual(self.current(identity)['state'], 'draft')
        self.assertNotEqual(item_digest(self.current(identity), self.case.state['snapshot']), original['items'][0]['confirmation']['digest'])
        with self.assertRaises(ValidationError):
            self.case.service.export(self.case.sid, 'csv', [identity])

    def test_empty_exclusions_and_missing_geometry_have_distinct_results(self):
        proposed = self.proposal(); proposed['geometry']['exclusions'] = []
        identity = self.create(proposed); self.case.confirm(identity)
        self.assertEqual(self.current(identity)['confirmation']['checks']['excluded_area_m2'], 0)
        proposed = self.proposal(); proposed.update(geometry=None, measurement=None)
        identity = self.create(proposed)
        self.assertIsNone(item_result(self.current(identity), self.case.state['snapshot'])['net_area_m2'])
        with self.assertRaisesRegex(ValidationError, 'Render'):
            self.case.command('review_items', item_ids=[identity])

    def test_collinear_forward_boundary_vertices_are_valid_and_duplicate_exclusion_ids_are_not(self):
        proposed = self.proposal()
        proposed['geometry']['points'].insert(1, [70, 30])
        identity = self.create(proposed)
        self.assertEqual(item_result(self.current(identity), self.case.state['snapshot'])['net_area_m2'], 48)
        proposed['geometry']['exclusions'].append(deepcopy(proposed['geometry']['exclusions'][0]))
        with self.assertRaisesRegex(ValidationError, 'IDs must be distinct'):
            self.create(proposed)

    def test_near_crossing_boundaries_are_not_accepted_through_topology_tolerance(self):
        proposed = self.proposal()
        proposed['geometry'].update(points=[[20, 30], [120, 80], [20, 80], [120, 30+1e-12]], exclusions=[])
        with self.assertRaisesRegex(ValidationError, 'cross or touch'):
            self.create(proposed)

    def test_area_exports_retain_exclusion_identity_and_do_not_emit_linear_quantities(self):
        identity = self.create(); self.case.confirm(identity)
        csv_bytes, _, _ = self.case.service.export(self.case.sid, 'csv', [identity])
        row = next(csv.DictReader(StringIO(csv_bytes.decode('utf-8-sig'))))
        self.assertEqual(row['Length per item m'], ''); self.assertEqual(row['Total length m'], '')
        self.assertEqual(float(row['Net area m2']), 48)
        self.assertEqual(json.loads(row['Area exclusions']), self.current(identity)['geometry']['exclusions'])
        payload, _, _ = self.case.service.export(self.case.sid, 'xlsx', [identity])
        workbook = load_workbook(BytesIO(payload)); sheet = workbook['Confirmed Takeoffs']
        values = {cell.value: sheet.cell(2, cell.column) for cell in sheet[1]}
        self.assertIsNone(values['Total length m'].value)
        self.assertEqual(values['Net area m2'].value, 48)
        self.assertEqual(values['Treatment'].data_type, 's')
        self.assertEqual(json.loads(values['Measurement basis'].value)['calibration']['distance_m'], 10)
        workbook.close()

    def test_area_transfer_and_split_merge_are_explicitly_unsupported_and_atomic(self):
        first = self.create(); second = self.create(self.proposal('slab'))
        self.case.confirm(first); self.case.confirm(second)
        before = deepcopy(self.case.state['snapshot'])
        for target in ('steel_vermiculite', 'steel_board', 'ductwork'):
            with self.subTest(target=target), self.assertRaisesRegex(ValidationError, 'no approved calculator transfer mapping'):
                mapped_values(self.current(first), before, target, 10)
        with self.assertRaisesRegex(ValidationError, 'Surface split is unavailable'):
            self.case.command('split_item', item_id=first, parts=[self.proposal(), self.proposal()])
        with self.assertRaisesRegex(ValidationError, 'Surface merge is unavailable'):
            self.case.command('merge_items', item_ids=[first, second], item=self.proposal())
        self.assertEqual(self.case.service.get(self.case.sid)['snapshot'], before)

    def test_existing_linear_receipt_and_result_shapes_are_unchanged_in_mixed_projects(self):
        linear = self.case.create(); self.case.confirm(linear)
        previous = deepcopy(self.current(linear)['confirmation'])
        self.create()
        self.assertEqual(self.current(linear)['confirmation'], previous)
        self.assertEqual(set(item_result(self.current(linear), self.case.state['snapshot'])), {'id', 'length_m', 'total_length_m', 'issues'})
        self.assertEqual(set(previous['checks']), {'engine', 'quantity', 'length_m', 'total_length_m', 'evidence_verified', 'issues'})
        validate_snapshot(self.case.state['snapshot'])

    def test_combined_register_exports_native_types_status_and_exact_source_area_without_changing_legacy_schema(self):
        wall = self.create(); proposed = self.proposal('slab')
        proposed['geometry']['points'][1][0] = 120.1234567890123
        proposed['geometry']['points'][2][0] = 120.1234567890123
        floor = self.create(proposed); self.case.confirm(wall)
        before = deepcopy(self.case.state['snapshot'])
        def export(mode, confirmation='all'):
            payload, _, name = self.case.service.export_workspace(self.case.sid, 'schedule-xlsx', {'expected_revision': before['revision'], 'mode': mode, 'confirmation': confirmation})
            workbook = load_workbook(BytesIO(payload)); self.addCleanup(workbook.close)
            values = list(workbook['Current Takeoffs'].values)
            return [dict(zip(values[0], row)) for row in values[1:]], values[0], name
        rows, headers, name = export('walls_floors')
        self.assertEqual(name, 'CEASEFIRE-Walls-Floors-Takeoffs.xlsx')
        self.assertEqual([row['Item ID'] for row in rows], [wall, floor])
        self.assertEqual([row['Mode'] for row in rows], ['wall', 'slab'])
        self.assertEqual([row['Confirmation'] for row in rows], ['Confirmed', 'Unconfirmed'])
        for row in rows:
            item = self.current(row['Item ID']); result = item_result(item, before)
            for column, key in [('Gross area m2', 'gross_area_m2'), ('Excluded area m2', 'excluded_area_m2'), ('Net area m2', 'net_area_m2')]:
                self.assertEqual(row[column], result[key])
            self.assertEqual(json.loads(row['Geometry']), item['geometry'])
            self.assertEqual(json.loads(row['All item properties']), item['fields'])
            self.assertEqual(row['Source SHA-256'], self.case.doc['sha256'])
        self.assertEqual([row['Item ID'] for row in export('walls_floors', 'confirmed')[0]], [wall])
        self.assertEqual([row['Item ID'] for row in export('walls_floors', 'unconfirmed')[0]], [floor])
        for mode, identity in [('wall', wall), ('slab', floor)]:
            legacy, legacy_headers, legacy_name = export(mode)
            self.assertEqual([row['Item ID'] for row in legacy], [identity])
            self.assertNotIn('Gross area m2', legacy_headers); self.assertNotIn('Excluded area m2', legacy_headers)
            self.assertEqual(legacy_name, f'CEASEFIRE-{mode.title()}-Takeoffs.xlsx')
        self.assertEqual(self.case.service.get(self.case.sid)['snapshot'], before)


if __name__ == '__main__':
    unittest.main()
