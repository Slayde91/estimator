"""Manual Steel counts preserve physical identities and native calculator inputs."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
import unittest
from unittest.mock import patch
from uuid import uuid4

from openpyxl import load_workbook
from pypdf import PdfReader

from estimator.catalog import ValidationError
from estimator.takeoff_model import item_digest, markup_appearance, validate_snapshot
from tests import test_takeoff_workspace as fixtures
from tests import test_takeoff_project as project_fixtures
from tests.test_takeoff_marked_exports import drawing_fixture


FIELDS = {'mark': 'C1', 'section': '100UC15', 'member_type': 'Beam',
          'exposure': 'Re-entrant - 3 sides', 'fire_period_min': 120,
          'critical_temperature': 550, 'product': 'CAFCO 300', 'sides': 3}


class TakeoffCountTests(unittest.TestCase):
    def setUp(self):
        self.case = fixtures.TakeoffWorkspaceTests(); self.case.setUp()
        self.addCleanup(self.case.doCleanups)
        # A separate source-only session proves the feature has no scale gate.
        self.case.state = self.case.service.open(); self.case.sid = self.case.state['session_id']
        self.case.state = self.case.service.add_document(self.case.sid, self.case.doc, self.case.state['revision'])
        self.case.command('record_render', document_id=self.case.doc['id'], page=1, success=True, warnings=[])

    def add(self, lengths=(3.5, 3.5, 4.25), **changes):
        request = {'document_id': self.case.doc['id'], 'page': 1,
                   'markers': [{'point': [20+index*15, 40+index*2], 'length_m': length}
                               for index, length in enumerate(lengths)],
                   'fields': deepcopy(FIELDS), 'appearance': {}, **changes}
        return self.case.command('add_count_items', **request)['created_item_ids']

    def item(self, identifier):
        return next(item for item in self.case.state['snapshot']['items'] if item['id'] == identifier)

    def assert_rejected(self, op, **values):
        before = deepcopy(self.case.state['snapshot']); blobs = deepcopy(self.case.documents.blobs)
        with self.assertRaises(ValidationError): self.case.command(op, **values)
        self.assertEqual(self.case.service.get(self.case.sid)['snapshot'], before)
        self.assertEqual(self.case.documents.blobs, blobs, 'Rejected commands must not publish an audit event.')

    def test_batch_groups_exact_lengths_only_and_retries_keep_generated_identities(self):
        old = deepcopy(self.case.state['snapshot']); old_blobs = deepcopy(self.case.documents.blobs)
        ids = self.add((3.5, 4, 3.5, 3.50000000001))
        first, second, third = [self.item(value) for value in ids]
        self.assertEqual(len(ids), 3); self.assertEqual(first['quantity'], 2)
        self.assertEqual(first['geometry']['points'], [[20, 40], [50, 44]])
        self.assertEqual(first['measurement'], {'method': 'manual', 'length_m': 3.5})
        self.assertEqual(first['count_id'], second['count_id']); self.assertEqual(second['count_id'], third['count_id'])
        members = [member for item in (first, second, third) for member in item['member_ids']]
        self.assertEqual(len(set(members)), 4)
        self.assertEqual(self.case.state['snapshot']['calibrations'], [])
        self.assertEqual(self.case.state['item_results'][0]['length_m'], 3.5)
        self.assertEqual(self.case.state['item_results'][0]['total_length_m'], 7)
        self.assertEqual(self.case.state['item_results'][0]['issues'], [])
        for key, value in old_blobs.items(): self.assertEqual(self.case.documents.blobs[key], value)
        self.assertEqual(validate_snapshot(old), old)
        self.assertEqual(self.case.state['snapshot']['version'], old['version'])
        other = self.add((3.5,), fields={**FIELDS, 'section': '410UB54'})
        self.assertNotEqual(self.item(other[0])['count_id'], first['count_id'])
        self.assertEqual(self.item(ids[0])['fields'], FIELDS)
        request = {'op': 'add_count_items', 'request_id': str(uuid4()), 'expected_revision': self.case.state['revision'],
                   'document_id': self.case.doc['id'], 'page': 1,
                   'markers': [{'point': [100.123456789, 100.987654321], 'length_m': 5.123456789}],
                   'fields': {}, 'appearance': {}}
        once = self.case.service.command(self.case.sid, request)
        self.assertEqual(once, self.case.service.command(self.case.sid, request))
        self.assertEqual(once['snapshot']['items'][-1]['geometry']['points'], [[100.123456789, 100.987654321]])

    def test_invalid_lengths_coordinates_styles_and_unknown_details_are_atomic(self):
        valid = {'document_id': self.case.doc['id'], 'page': 1,
                 'markers': [{'point': [20, 40], 'length_m': 3.5}], 'fields': {}, 'appearance': {}}
        for length in (None, 0, -1, True, float('nan'), float('inf'), 1e13):
            with self.subTest(length=length):
                self.assert_rejected('add_count_items', **{**valid, 'markers': [valid['markers'][0], {'point': [30, 40], 'length_m': length}]})
        for changes in ({'markers': []}, {'markers': [{'point': [0, 0], 'length_m': 2}]},
                        {'page': 2}, {'fields': {'invented': 'value'}}, {'appearance': {'marker_shape': 'star'}},
                        {'appearance': {'marker_size': 0}}, {'appearance': {'marker_size': 73}},
                        {'appearance': {'marker_size': True}}, {'appearance': {'opacity': 1.1}}):
            with self.subTest(changes=changes): self.assert_rejected('add_count_items', **{**valid, **changes})

    def test_quantity_identity_and_geometry_shortcuts_are_rejected_but_single_marker_can_move(self):
        identifier = self.add((3.5,))[0]; old = deepcopy(self.item(identifier))
        for changes in ({'quantity': 2}, {'member_ids': [str(uuid4())]}, {'geometry': None},
                        {'geometry': {**old['geometry'], 'points': [[20, 40], [30, 40]]}},
                        {'geometry': {key: value for key, value in old['geometry'].items() if key != 'kind'}},
                        {'measurement': {'method': 'cited', 'length_m': 3.5, 'citation': 'fake source length'}}):
            with self.subTest(changes=changes): self.assert_rejected('update_item', item_id=identifier, changes=changes)
        self.case.command('confirm_items', item_ids=[identifier]); confirmed = deepcopy(self.item(identifier))
        self.case.command('update_item', item_id=identifier, changes={
            'quantity': 1, 'member_ids': old['member_ids'], 'geometry': {**old['geometry'], 'points': [[40.125, 70.875]]}})
        current = self.item(identifier)
        self.assertEqual(current['member_ids'], old['member_ids']); self.assertEqual(current['measurement'], old['measurement'])
        self.assertEqual(self.case.state['item_results'][0]['length_m'], 3.5)
        self.assertIsNone(current['confirmation']); self.assertNotEqual(item_digest(current, self.case.state['snapshot']), confirmed['confirmation']['digest'])
        self.assert_rejected('split_steel_group', item_id=identifier, quantities=[1, 1])
        other = self.add((3.5,))[0]
        self.assert_rejected('merge_steel_groups', item_ids=[identifier, other])
        self.assert_rejected('merge_items', item_ids=[identifier, other], item={})

    def test_snapshot_rejects_forged_counts_and_manual_measurements_on_traces(self):
        identifier = self.add((3.5,))[0]; original = deepcopy(self.case.state['snapshot'])
        for changes in ({'quantity': 2}, {'count_id': None}, {'measurement': None}, {'mode': 'duct'},
                        {'geometry': {key: value for key, value in self.item(identifier)['geometry'].items() if key != 'kind'}}):
            value = deepcopy(original); value['items'][0].update(changes)
            with self.subTest(changes=changes), self.assertRaises(ValidationError): validate_snapshot(value)
        value = deepcopy(original); value['items'][0].pop('count_id')
        with self.assertRaises(ValidationError): validate_snapshot(value)
        self.assert_rejected('create_item', item={key: deepcopy(value) for key, value in self.item(identifier).items()
                             if key in ('mode', 'geometry', 'measurement', 'quantity', 'fields', 'count_id')})

    def test_shared_details_require_whole_count_and_appearance_preserves_confirmation(self):
        ids = self.add(); self.case.command('confirm_items', item_ids=ids)
        before = deepcopy(self.case.state['snapshot']); receipts = {item['id']: deepcopy(item['confirmation']) for item in before['items']}
        self.assert_rejected('update_item', item_id=ids[0], changes={'fields': {'section': '410UB54'}})
        self.assert_rejected('bulk_update', item_ids=[ids[0]], changes={'appearance': {'marker_shape': 'diamond'}})
        addition = {'id': str(uuid4()), 'kind': 'riser', 'length_mm': 500, 'note': 'Explicit elevation rise',
                    'document_id': self.case.doc['id'], 'page': 1}
        self.assert_rejected('update_item', item_id=ids[0], changes={'length_additions': [addition]})
        self.case.command('bulk_update', item_ids=ids, changes={'appearance': {'marker_shape': 'diamond', 'marker_size': 18,
                           'fill_enabled': True, 'fill_color': '#123456', 'opacity': .6}})
        for identifier in ids:
            self.assertEqual(self.item(identifier)['confirmation'], receipts[identifier])
            self.assertEqual(markup_appearance(self.item(identifier))['marker_size'], 18)
        self.case.command('bulk_update', item_ids=ids, changes={'fields': {'fire_period_min': 90}})
        for identifier in ids:
            self.assertEqual(self.item(identifier)['fields']['fire_period_min'], 90)
            self.assertIsNone(self.item(identifier)['confirmation'])
        forged = deepcopy(self.case.state['snapshot']); forged['items'][0]['fields']['section'] = 'OTHER'
        with self.assertRaisesRegex(ValidationError, 'same source, details'): validate_snapshot(forged)
        forged = deepcopy(self.case.state['snapshot']); forged['items'][0]['length_additions'] = [addition]
        with self.assertRaisesRegex(ValidationError, 'same source, details'): validate_snapshot(forged)

    def test_manual_length_collision_merges_only_same_count_and_undo_restores_identities(self):
        ids = self.add(); other = self.add((4.25,))[0]
        self.case.command('confirm_items', item_ids=ids)
        first, second = [deepcopy(self.item(identifier)) for identifier in ids]
        self.case.apply(self.case.preview(ids[1])); calculator = deepcopy(self.case.state['calculator'])
        self.case.command('update_item', item_id=ids[0], changes={'measurement': {'method': 'manual', 'length_m': 4.25}})
        self.assertEqual(self.case.state['regrouped_item_ids'], [ids[0]])
        merged = self.item(ids[0]); self.assertEqual(merged['quantity'], 3)
        self.assertEqual(merged['member_ids'], first['member_ids'] + second['member_ids'])
        self.assertEqual(merged['geometry']['points'], first['geometry']['points'] + second['geometry']['points'])
        self.assertIn(second['id'], merged['predecessor_ids']); self.assertEqual(self.item(other)['quantity'], 1)
        self.assertEqual(self.case.state['snapshot']['transfers'][0]['status'], 'deleted')
        self.assertNotIn('calculator', self.case.state)
        self.assertEqual(calculator['inputs']['SCHEDULE']['I10'], 1)
        self.case.command('undo')
        for previous in (first, second):
            restored = self.item(previous['id'])
            for name in ('count_id', 'member_ids', 'geometry', 'measurement', 'quantity'):
                self.assertEqual(restored[name], previous[name])
            self.assertIsNone(restored['confirmation'])

    def test_member_length_edits_survive_regrouping_and_keep_unedited_members_unchanged(self):
        ids = self.add((3, 4, 4)); other = self.add((3,))[0]
        originals = [deepcopy(self.item(identifier)) for identifier in ids]
        members = [item['member_ids'] for item in originals]
        self.case.command('confirm_items', item_ids=ids)
        self.case.apply(self.case.preview(ids[0])); calculator = deepcopy(self.case.state['calculator'])
        self.case.command('update_count_lengths', groups=[{'member_ids': members[0], 'length_m': 4}])
        merged = deepcopy(self.item(ids[0])); self.assertEqual(merged['quantity'], 3)
        self.assertEqual(self.case.state['snapshot']['transfers'][0]['status'], 'stale')
        # A second edit captured before the merge still targets its own members.
        request = {'op': 'update_count_lengths', 'request_id': str(uuid4()), 'expected_revision': self.case.state['revision'],
                   'groups': [{'member_ids': members[1], 'length_m': 6.123456789}]}
        self.case.state = self.case.service.command(self.case.sid, request)
        self.assertEqual(self.case.state, self.case.service.command(self.case.sid, request))
        rows = [item for item in self.case.state['snapshot']['items'] if item['count_id'] == merged['count_id']]
        self.assertEqual(len(rows), 2)
        by_member = {member: (item['measurement']['length_m'], point) for item in rows
                     for member, point in zip(item['member_ids'], item['geometry']['points'])}
        for index, original in enumerate(originals):
            for member, point in zip(original['member_ids'], original['geometry']['points']):
                self.assertEqual(by_member[member], (4 if index == 0 else 6.123456789, point))
        for item in rows:
            self.assertEqual(item['fields'], originals[0]['fields'])
            self.assertEqual(item['evidence'], originals[0]['evidence'])
            self.assertIsNone(item['confirmation'])
        self.assertEqual(self.item(other)['measurement']['length_m'], 3)
        self.assertNotIn('calculator', self.case.state)
        self.assertEqual(calculator['inputs']['SCHEDULE']['I10'], 1)
        self.case.command('undo')
        restored = self.item(merged['id'])
        for key in ('geometry', 'measurement', 'member_ids', 'evidence', 'quantity', 'count_id'):
            self.assertEqual(restored[key], merged[key])

    def test_member_length_swaps_are_atomic_and_invalid_or_duplicate_members_are_rejected(self):
        ids = self.add((3, 4)); before = [deepcopy(self.item(identifier)) for identifier in ids]
        groups = [{'member_ids': before[0]['member_ids'], 'length_m': 4},
                  {'member_ids': before[1]['member_ids'], 'length_m': 3}]
        self.case.command('update_count_lengths', groups=groups)
        for old, length in zip(before, (4, 3)):
            current = self.item(old['id'])
            self.assertEqual(current['measurement']['length_m'], length)
            self.assertEqual(current['member_ids'], old['member_ids'])
        for invalid in ([], [{'member_ids': [], 'length_m': 3}], groups + [groups[0]],
                        [{'member_ids': [str(uuid4())], 'length_m': 3}]):
            self.assert_rejected('update_count_lengths', groups=invalid)
        for length in (None, 0, -1, True, float('nan'), float('inf'), 1e13):
            self.assert_rejected('update_count_lengths', groups=[{'member_ids': before[0]['member_ids'], 'length_m': length}])

    def test_bulk_manual_length_groups_each_count_separately_and_rejects_mixed_traces(self):
        ids = self.add(); other = self.add((6.25,), fields={**FIELDS, 'mark': 'Different count'})[0]
        self.case.command('bulk_update', item_ids=ids+[other], changes={'measurement': {'method': 'manual', 'length_m': 8}})
        self.assertEqual(self.case.state['regrouped_item_ids'], [ids[0], other])
        self.assertEqual(self.item(ids[0])['quantity'], 3); self.assertEqual(self.item(other)['quantity'], 1)
        self.assertNotEqual(self.item(ids[0])['count_id'], self.item(other)['count_id'])
        trace = self.case.create(measurement={'method': 'cited', 'length_m': 4, 'citation': 'Source dimension'})
        self.assert_rejected('bulk_update', item_ids=[ids[0], trace], changes={'measurement': {'method': 'manual', 'length_m': 10}})

    def test_delete_specific_marker_and_last_group_preserve_audit_and_calculator(self):
        identifier = self.add((3.5, 3.5, 3.5))[0]
        self.case.confirm(identifier); self.case.apply(self.case.preview(identifier)); original = deepcopy(self.item(identifier))
        calculator = deepcopy(self.case.state['calculator']); member = original['member_ids'][1]
        self.case.command('delete_count_marker', item_id=identifier, member_id=member)
        current = self.item(identifier)
        self.assertEqual(current['quantity'], 2)
        self.assertEqual(current['member_ids'], [original['member_ids'][0], original['member_ids'][2]])
        self.assertEqual(current['geometry']['points'], [original['geometry']['points'][0], original['geometry']['points'][2]])
        self.assertEqual(self.case.state['snapshot']['transfers'][0]['status'], 'stale')
        self.assertEqual(calculator['inputs']['SCHEDULE']['I10'], 3)
        self.assert_rejected('delete_count_marker', item_id=identifier, member_id=member)
        self.case.command('undo'); self.assertEqual(self.item(identifier)['member_ids'], original['member_ids'])
        for member in original['member_ids']:
            self.case.command('delete_count_marker', item_id=identifier, member_id=member)
        self.assertEqual(self.case.state['snapshot']['items'], [])
        self.assertEqual(self.case.state['snapshot']['transfers'][0]['status'], 'deleted')
        self.case.command('undo')
        self.assertEqual(self.item(identifier)['member_ids'], [original['member_ids'][-1]])
        self.assertIsNone(self.item(identifier)['confirmation'])

    def test_continue_groups_lengths_preserves_rows_and_retries_without_duplicate_members(self):
        ids = self.add(appearance={'marker_shape': 'diamond', 'marker_size': 18})
        for identifier in ids:
            self.case.command('update_item', item_id=identifier, changes={'evidence': [
                {'document_id': self.case.doc['id'], 'page': 1, 'note': identifier}]})
        self.case.command('confirm_items', item_ids=ids)
        self.case.apply(self.case.preview(ids[0])); calculator = deepcopy(self.case.state['calculator'])
        before = deepcopy(self.case.state['snapshot']); original = {item['id']: item for item in before['items']}
        request = {'op': 'continue_count', 'request_id': str(uuid4()), 'expected_revision': before['revision'],
                   'item_id': ids[1], 'markers': [{'point': [70, 60], 'length_m': 3.5},
                                               {'point': [80, 65], 'length_m': 8},
                                               {'point': [90, 70], 'length_m': 8}]}
        self.case.state = self.case.service.command(self.case.sid, request)
        self.assertEqual(self.case.state, self.case.service.command(self.case.sid, request))
        new_id, = self.case.state['created_item_ids']
        self.assertEqual(self.case.state['regrouped_item_ids'], ids + [new_id])
        appended = self.item(ids[0]); untouched = self.item(ids[1]); new = self.item(new_id)
        self.assertEqual(appended['geometry']['points'], original[ids[0]]['geometry']['points'] + [[70, 60]])
        self.assertEqual(appended['member_ids'][:2], original[ids[0]]['member_ids'])
        self.assertEqual(appended['quantity'], 3); self.assertIsNone(appended['confirmation'])
        self.assertEqual(appended['evidence'], original[ids[0]]['evidence'])
        self.assertEqual(untouched, original[ids[1]], 'A row with no new marker retains its confirmation.')
        for key in ('count_id', 'fields', 'appearance', 'evidence'):
            self.assertEqual(new[key], original[ids[1]][key])
        self.assertEqual(new['measurement'], {'method': 'manual', 'length_m': 8})
        self.assertEqual(new['quantity'], 2)
        members = [member for item in self.case.state['snapshot']['items'] for member in item['member_ids']]
        self.assertEqual(len(set(members)), 6)
        for key, value in before['transfers'][0].items():
            self.assertEqual(self.case.state['snapshot']['transfers'][0][key], 'stale' if key == 'status' else value)
        self.assertNotIn('calculator', self.case.state)
        self.assertEqual(calculator['inputs']['SCHEDULE']['I10'], 2)
        self.case.command('undo')
        self.assertEqual([item['id'] for item in self.case.state['snapshot']['items']], ids)
        for identifier in ids:
            for key in ('geometry', 'measurement', 'member_ids', 'evidence', 'quantity', 'count_id', 'appearance'):
                self.assertEqual(self.item(identifier)[key], original[identifier][key])

    def test_continue_new_length_retains_additions_and_does_not_invalidate_existing_links(self):
        identifier = self.add((3.5,))[0]
        addition = {'id': str(uuid4()), 'kind': 'riser', 'length_mm': 500, 'note': 'Manually entered rise',
                    'document_id': self.case.doc['id'], 'page': 1}
        self.case.command('update_item', item_id=identifier, changes={'length_additions': [addition]})
        self.case.confirm(identifier); self.case.apply(self.case.preview(identifier))
        previous = deepcopy(self.case.state['snapshot'])
        self.case.command('continue_count', item_id=identifier, markers=[{'point': [60, 50], 'length_m': 6.25}])
        new_id, = self.case.state['created_item_ids']
        self.assertEqual(self.item(identifier), previous['items'][0])
        self.assertEqual(self.case.state['snapshot']['transfers'], previous['transfers'])
        self.assertEqual(self.item(new_id)['length_additions'], [addition])
        result = next(result for result in self.case.state['item_results'] if result['id'] == new_id)
        self.assertEqual(result['length_m'], 6.75)
        self.assertEqual(result['total_length_m'], 6.75)

    def test_continue_rejects_invalid_batches_and_unauthorised_detail_changes_atomically(self):
        identifier = self.add((3.5,))[0]
        valid = {'point': [50, 50], 'length_m': 4}
        for marker in ({'point': [0, 0], 'length_m': 5}, {'point': [60, 60], 'length_m': 0},
                       {'point': [60, 60], 'length_m': True}, {'point': [60, 60], 'length_m': float('inf')},
                       {'point': [60, 60], 'length_m': 4, 'member_id': str(uuid4())}):
            with self.subTest(marker=marker):
                self.assert_rejected('continue_count', item_id=identifier, markers=[valid, marker])
        for details in ({'fields': {}}, {'appearance': {}}, {'document_id': self.case.doc['id']}, {'page': 1}):
            self.assert_rejected('continue_count', item_id=identifier, markers=[valid], **details)
        for target in (str(uuid4()), None, [], self.case.create(measurement={'method': 'cited', 'length_m': 4, 'citation': 'Plan'})):
            self.assert_rejected('continue_count', item_id=target, markers=[valid])
        self.assert_rejected('continue_count', item_id=identifier, markers=[])
        with patch('estimator.takeoff_workspace.MAX_ITEMS', 4):
            self.assert_rejected('continue_count', item_id=identifier, markers=[valid, valid])

    def test_selected_marker_move_is_atomic_and_keeps_manual_lengths_and_other_markers(self):
        ids = self.add((3.5, 3.5, 3.5, 4.25)); other = self.add((7,))[0]
        self.case.command('confirm_items', item_ids=ids + [other])
        self.case.apply(self.case.preview(ids[0])); calculator = deepcopy(self.case.state['calculator'])
        before = deepcopy(self.case.state['snapshot']); original = {item['id']: item for item in before['items']}
        targets = [{'item_id': identifier, 'member_id': original[identifier]['member_ids'][index]}
                   for identifier, index in ((ids[0], 0), (ids[0], 2), (ids[1], 0))]
        request = {'op': 'move_count_markers', 'request_id': str(uuid4()), 'expected_revision': before['revision'],
                   'markers': targets, 'delta_pdf': [12.125, 8.875]}
        self.case.state = self.case.service.command(self.case.sid, request)
        self.assertEqual(self.case.state, self.case.service.command(self.case.sid, request))
        self.assertEqual(self.case.state['revision'], before['revision'] + 1)
        for identifier in ids:
            current = self.item(identifier); old = original[identifier]
            self.assertEqual(current['version'], old['version'] + 1)
            for key in ('quantity', 'measurement', 'member_ids', 'fields', 'appearance', 'count_id'):
                self.assertEqual(current[key], old[key])
            self.assertIsNone(current['confirmation'])
            for index, (x, y) in enumerate(old['geometry']['points']):
                expected = [x + 12.125, y + 8.875] if identifier != ids[0] or index != 1 else [x, y]
                self.assertEqual(current['geometry']['points'][index], expected)
        self.assertEqual(self.item(other), original[other])
        self.assertEqual(self.case.state['snapshot']['transfers'][0]['status'], 'stale')
        self.assertNotIn('calculator', self.case.state)
        self.assertEqual(calculator['inputs']['SCHEDULE']['I10'], 3)
        self.assertEqual(calculator['inputs']['SCHEDULE']['J10'], 3.5)
        self.case.command('undo')
        for identifier in ids:
            self.assertEqual(self.item(identifier)['geometry'], original[identifier]['geometry'])
            self.assertEqual(self.item(identifier)['member_ids'], original[identifier]['member_ids'])

    def test_marker_move_rejects_bad_targets_offsets_and_mixed_sources_without_partial_write(self):
        identifier = self.add((3.5, 3.5))[0]
        valid = {'item_id': identifier, 'member_id': self.item(identifier)['member_ids'][0]}
        for target in (valid, {'item_id': identifier, 'member_id': str(uuid4())},
                       {'item_id': str(uuid4()), 'member_id': valid['member_id']},
                       {'item_id': identifier, 'member_id': []}, {**valid, 'point': [40, 40]}):
            self.assert_rejected('move_count_markers', markers=[valid, target], delta_pdf=[2, 3])
        for delta in ([], [0, 0], [1000, 0], [True, 0], [float('nan'), 0], [1], '1,2'):
            self.assert_rejected('move_count_markers', markers=[valid], delta_pdf=delta)
        self.assert_rejected('move_count_markers', markers=[], delta_pdf=[2, 3])
        document = {**fixtures.document(), 'sha256': 'b'*64}
        self.case.state = self.case.service.add_document(self.case.sid, document, self.case.state['revision'])
        other = self.add((4,), document_id=document['id'])[0]
        target = {'item_id': other, 'member_id': self.item(other)['member_ids'][0]}
        self.assert_rejected('move_count_markers', markers=[valid, target], delta_pdf=[2, 3])
        with patch('estimator.takeoff_workspace.MAX_ITEMS', 1):
            self.assert_rejected('move_count_markers', markers=[valid, target], delta_pdf=[2, 3])

    def test_continue_and_move_reject_stale_revisions(self):
        identifier = self.add((3.5,))[0]
        revision = self.case.state['revision']; member = self.item(identifier)['member_ids'][0]
        self.case.command('continue_count', item_id=identifier, markers=[{'point': [40, 50], 'length_m': 3.5}])
        self.assert_rejected('continue_count', expected_revision=revision, item_id=identifier,
                             markers=[{'point': [60, 60], 'length_m': 8}])
        self.assert_rejected('move_count_markers', expected_revision=revision,
                             markers=[{'item_id': identifier, 'member_id': member}], delta_pdf=[2, 3])

    def test_both_steel_transfer_mappings_keep_existing_quantity_and_length_contract(self):
        spray = self.add((2.75, 2.75))[0]
        self.case.confirm(spray); preview = self.case.preview(spray); self.case.apply(preview)
        values = self.case.state['calculator']['inputs']['SCHEDULE']
        self.assertEqual(values['I10'], 2); self.assertEqual(values['J10'], 2.75)
        board = self.add((3.125, 3.125), fields={**FIELDS, 'product': 'TRAFALGAR COREX', 'critical_temperature': 620})[0]
        self.case.confirm(board); self.case.apply(self.case.preview(board, 'steel_board'))
        self.assertEqual(self.case.state['calculator']['inputs']['CALCULATOR']['F9'], 6.25)
        self.assertEqual(self.item(spray)['quantity'], 2)

    def test_current_and_confirmed_exports_retain_count_source_and_ordered_member_mapping(self):
        identifier = self.add((3.5, 3.5))[0]; self.case.confirm(identifier); item = deepcopy(self.item(identifier))
        for format, sheet in (('schedule-xlsx', 'Current Takeoffs'), ('xlsx', 'Confirmed Takeoffs')):
            if format == 'schedule-xlsx':
                payload = self.case.service.export_workspace(self.case.sid, format,
                    {'mode': 'steel', 'expected_revision': self.case.state['revision']})[0]
            else: payload = self.case.service.export(self.case.sid, format, [identifier])[0]
            book = load_workbook(BytesIO(payload)); self.addCleanup(book.close)
            rows = list(book[sheet].values); row = dict(zip(rows[0], rows[1]))
            self.assertEqual(row['Quantity'], 2); self.assertEqual(row['Count ID'], item['count_id'])
            self.assertEqual(json.loads(row['Physical member IDs']), item['member_ids'])
            self.assertEqual(json.loads(row['Measurement basis']), {'method': 'manual', 'length_m': 3.5})
            geometry = json.loads(row['Geometry'] if format == 'schedule-xlsx' else row['Geometry PDF points'])
            self.assertEqual(geometry['points'] if isinstance(geometry, dict) else geometry, item['geometry']['points'])
        self.assertEqual(self.item(identifier), item)


class TakeoffCountEvidenceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): project_fixtures.TakeoffProjectTests.setUpClass()

    def setUp(self):
        self.case = project_fixtures.TakeoffProjectTests(); self.case.pdf = drawing_fixture(); self.case.setUp()
        self.addCleanup(self.case.doCleanups)
        self.document = self.case.session['snapshot']['documents'][0]

    def command(self, op, **values):
        case = self.case
        case.session = case.service.command(case.session['session_id'], {'op': op, 'request_id': str(uuid4()),
            'expected_revision': case.session['revision'], **values})
        return case.session

    def add(self, page=1, shape='circle'):
        self.command('record_render', document_id=self.document['id'], page=page, success=True, warnings=[])
        return self.command('add_count_items', document_id=self.document['id'], page=page,
            markers=[{'point': [30, 40], 'length_m': 3.5}, {'point': [160, 80], 'length_m': 3.5}],
            fields={**FIELDS, 'mark': 'COUNT-'+shape}, appearance={'marker_shape': shape, 'marker_size': 12,
            'stroke_color': '#F12175', 'fill_color': '#F12175', 'fill_enabled': True})['created_item_ids'][0]

    def test_real_save_reopen_preserves_counts_originals_and_existing_calculator_project_values(self):
        case = self.case; initial = deepcopy(case.session['snapshot'])
        old_audit = {path.name: path.read_bytes() for path in (case.documents.root/'audit').glob('*')}
        identifier = self.add(); self.command('confirm_items', item_ids=[identifier])
        snapshot = deepcopy(case.session['snapshot'])
        case.request.update(takeoffs=snapshot, takeoffs_session_id=case.session['session_id'])
        case.library.save_as(case.request); saved = json.loads(case.target.read_bytes())
        self.assertEqual(saved['takeoffs']['items'], snapshot['items'])
        for name in ('estimate', 'calculators'): self.assertEqual(saved[name], json.loads(case.legacy)[name])
        case.dialogs.opened = str(case.target); reopened = case.library.open_file()
        self.assertEqual(reopened['takeoffs']['items'], snapshot['items']); self.assertEqual(reopened['takeoffs_issues'], [])
        self.assertEqual(validate_snapshot(initial), initial)
        for name, content in old_audit.items(): self.assertEqual((case.documents.root/'audit'/name).read_bytes(), content)
        self.assertEqual(case.documents.document_path(self.document).read_bytes(), case.pdf)

    def test_continued_and_moved_markers_save_reopen_with_original_member_mapping(self):
        identifier = self.add(); case = self.case
        original = deepcopy(case.session['snapshot']['items'][0])
        self.command('continue_count', item_id=identifier, markers=[{'point': [80, 60], 'length_m': 3.5},
                                                                 {'point': [120, 70], 'length_m': 5.25}])
        new_id, = case.session['created_item_ids']
        new = next(item for item in case.session['snapshot']['items'] if item['id'] == new_id)
        self.command('move_count_markers', markers=[{'item_id': identifier, 'member_id': original['member_ids'][0]},
                                                  {'item_id': new_id, 'member_id': new['member_ids'][0]}],
                     delta_pdf=[10, 15])
        snapshot = deepcopy(case.session['snapshot'])
        case.request.update(takeoffs=snapshot, takeoffs_session_id=case.session['session_id'])
        case.library.save_as(case.request)
        case.dialogs.opened = str(case.target); reopened = case.library.open_file()
        self.assertEqual(reopened['takeoffs']['items'], snapshot['items'])
        self.assertEqual(reopened['takeoffs_issues'], [])
        first = next(item for item in reopened['takeoffs']['items'] if item['id'] == identifier)
        self.assertEqual(first['member_ids'][:2], original['member_ids'])
        self.assertEqual(first['geometry']['points'], [[40, 55], [160, 80], [80, 60]])
        for name in ('estimate', 'calculators'):
            self.assertEqual(json.loads(case.target.read_bytes())[name], json.loads(case.legacy)[name])
        self.assertEqual(case.documents.document_path(self.document).read_bytes(), case.pdf)

    def test_actual_pdf_renders_separate_shapes_and_retains_rotation_crop_and_manual_legend(self):
        shapes = ('circle', 'square', 'triangle', 'diamond')
        identifiers = [self.add(page, shape) for page, shape in enumerate(shapes, 1)]
        case = self.case; before = deepcopy(case.session['snapshot'])
        original = case.documents.document_path(self.document); original_hash = sha256(original.read_bytes()).hexdigest()
        payload = case.service.export_workspace(case.session['session_id'], 'marked-pdf', {
            'expected_revision': case.session['revision'], 'mode': 'steel', 'document_id': self.document['id'],
            'item_ids': identifiers})[0]
        reader = PdfReader(BytesIO(payload)); self.assertEqual(len(reader.pages), 4)
        for index, (page, shape) in enumerate(zip(reader.pages, shapes)):
            text = page.extract_text(); self.assertIn('COUNT-'+shape, text)
            self.assertIn('2 markers x 3.50 m manual length each; 7.00 m total', text)
            self.assertIn(f'ORIGINAL DRAWING ROTATION {index*90}', text)
            operations = page.get_contents().operations
            # Examine actual exported vector paths. Each painted symbol must
            # terminate before the next marker; no connecting segment exists.
            paths = []; current = []
            for values, operator in operations:
                if operator == b'n': current = []
                elif operator in (b'm', b'l', b'c', b're', b'h'): current.append((values, operator))
                elif operator in (b'B*', b'B'):
                    paths.append(current); current = []
            if shape == 'circle': matching = [path for path in paths if sum(op == b'c' for _, op in path) == 4]
            elif shape == 'square': matching = [path for path in paths if any(op == b're' and list(map(float, vals[2:])) == [12, 12] for vals, op in path)]
            else:
                lines = 2 if shape == 'triangle' else 3
                matching = [path for path in paths if sum(op == b'l' for _, op in path) == lines and sum(op == b'm' for _, op in path) == 1]
            self.assertEqual(len(matching), 2, shape)
        self.assertEqual(sha256(original.read_bytes()).hexdigest(), original_hash)
        self.assertEqual(case.service.get(case.session['session_id'])['snapshot'], before)

    def test_pdf_and_register_distinguish_manual_base_from_cited_riser_additions(self):
        identifier = self.add(); case = self.case
        self.command('update_item', item_id=identifier, changes={'length_additions': [
            {'id': str(uuid4()), 'kind': 'riser', 'length_mm': 500, 'note': 'Section A explicit height',
             'document_id': self.document['id'], 'page': 1}]})
        self.command('confirm_items', item_ids=[identifier])
        result = case.session['item_results'][0]
        self.assertEqual((result['base_length_m'], result['additions_length_m'], result['length_m'], result['total_length_m']),
                         (3.5, .5, 4, 8))
        payload = case.service.export_workspace(case.session['session_id'], 'marked-pdf', {
            'expected_revision': case.session['revision'], 'mode': 'steel', 'document_id': self.document['id'],
            'item_ids': [identifier]})[0]
        text = ' '.join(PdfReader(BytesIO(payload)).pages[0].extract_text().split())
        self.assertIn('2 markers x 3.50 m manual base + 0.50 m explicit additions each; 8.00 m total', text)
        self.assertNotIn('4.00 m manual length', text)
        payload = case.service.export(case.session['session_id'], 'xlsx', [identifier])[0]
        book = load_workbook(BytesIO(payload)); self.addCleanup(book.close)
        data = list(book['Confirmed Takeoffs'].values); row = dict(zip(data[0], data[1]))
        self.assertEqual(row['Base length per item m'], 3.5)
        self.assertEqual(row['Riser/drop additions per item m'], .5)
        self.assertEqual(row['Length per item m'], 4); self.assertEqual(row['Total length m'], 8)
        provenance = dict(book['Provenance'].values)
        self.assertIn('explicitly entered manual length', provenance['Coordinates'])
        self.assertIn('derived from its retained member markers', provenance['Quantity'])


if __name__ == '__main__': unittest.main()
