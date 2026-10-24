"""Confirmed, source-bound intumescent transfers and native DFT unit projection.

The fixed assessed result isolates bridge arithmetic from assessment-table tests.
No result is copied into Takeoff properties, calculator inputs or saved receipts.
"""
from copy import deepcopy
from io import BytesIO
import unittest
from unittest.mock import patch
from uuid import uuid4

from openpyxl import load_workbook

from estimator.catalog import ValidationError
from estimator.intumescent import PRODUCT, FX5090, FX6010
from estimator.schedule_rows import empty_schedule_inputs
from estimator.takeoff_exports import linked_register_results, linked_result_text
from estimator.takeoff_model import validate_snapshot
from estimator.takeoff_presentation import apply_presentation, effective_appearance, legend_rows
from estimator.workbook_calculators import calculator_session, _session
from tests import test_takeoff_workspace as fixtures


DFT = 1234.56789012345
ASSESSED = {'dft': DFT, 'factor': 210, 'status': 'ESTIMATE - CONSERVATIVE ASSESSMENT',
            'source': 'FC18580-01-01 p12 Table 6; 550C; A/V 210; 60 min', 'note': 'Test assessment result'}


class IntumescentTakeoffTests(unittest.TestCase):
    def setUp(self):
        self.case = fixtures.TakeoffWorkspaceTests()
        self.case.setUp()
        self.addCleanup(self.case.doCleanups)
        _session.cache_clear()
        self.addCleanup(_session.cache_clear)
        self.lookup = patch('estimator.sc902_assessment.lookup', return_value=ASSESSED).start()
        self.firetex = patch('estimator.firetex_assessment.lookup', return_value={
            **ASSESSED, 'status': 'PUBLISHED - CERTIFICATE TABLE', 'source': 'FIRETEX certificate Table 1'}).start()
        self.addCleanup(patch.stopall)

    def create(self, product=PRODUCT, **changes):
        identity = self.case.create(quantity=3, measurement={
            'method': 'cited', 'length_m': 7.123456789012345, 'citation': 'Explicit plan dimension'})
        self.case.command('update_item', item_id=identity, changes={'fields': {
            'product': product, 'section': '360UB45', 'exposure': 'I/H beam - 3 sides',
            'critical_temperature': 550, 'fire_period_min': 60, 'level': 'L2', 'zone': 'East', **changes}})
        self.case.confirm(identity)
        return identity

    def linked(self, product=PRODUCT, inputs=None, **changes):
        identity = self.create(product, **changes)
        preview = self.case.preview(identity, 'steel_intumescent', inputs=inputs)
        self.case.apply(preview, inputs=inputs)
        return identity, {'steel_intumescent': {key: deepcopy(preview[key]) for key in ('inputs', 'schedule_rows')}}

    def read(self, drafts):
        return self.case.service.linked_register_results(self.case.sid, {
            'expected_revision': self.case.state['revision'], 'calculator_drafts': drafts})['linked_results']

    def test_transfer_matches_native_manual_inputs_and_all_row_outputs_without_precision_loss(self):
        length = 7.123456789012345
        for product in (PRODUCT, FX5090, FX6010):
            with self.subTest(product=product):
                identity = self.create(product, girth_override_m=.912345678901234,
                                       area_override_m2=19.87654321098765)
                original = empty_schedule_inputs('steel_intumescent')
                original['SCHEDULE']['A11'] = 'Unrelated manual member'
                original['SETTINGS']['H70'] = .123456789012345
                before = deepcopy(original)
                preview = self.case.preview(identity, 'steel_intumescent', inputs=original, rows=[10, 11])
                self.assertEqual(original, before)
                row = preview['changes'][0]['row']
                manual = deepcopy(original)
                manual['SCHEDULE'].update({column+str(row): value for column, value in {
                    'A': 'B17', 'AA': 'L2 / East', 'B': product, 'C': 'I/H beam - 3 sides',
                    'D': 550, 'E': 'Section', 'F': '360UB45', 'G': None,
                    'H': 60, 'I': 3, 'J': length, 'K': .912345678901234, 'L': 19.87654321098765}.items()})
                self.assertEqual(preview['inputs'], manual)
                _, actual, lock = calculator_session('steel_intumescent', preview['inputs'])
                _, expected, expected_lock = calculator_session('steel_intumescent', manual)
                with lock, expected_lock:
                    for column in 'MNOPQRSTUVWXY':
                        self.assertEqual(actual.value('SCHEDULE', column+str(row)), expected.value('SCHEDULE', column+str(row)))
                self.case.apply(preview, inputs=original, rows=[10, 11])
                validate_snapshot(self.case.state['snapshot'])

    def test_all_products_return_exact_dry_film_mm_and_report_provenance_read_only(self):
        for product in (PRODUCT, FX5090, FX6010):
            with self.subTest(product=product):
                identity, drafts = self.linked(product)
                original = deepcopy(drafts)
                before = deepcopy(self.case.state['snapshot'])
                result = self.read(drafts)[identity][0]
                binding = next(row for row in before['transfers'] if row['item_id'] == identity)
                _, engine, lock = calculator_session('steel_intumescent', drafts['steel_intumescent']['inputs'])
                with lock:
                    native = engine.value('SCHEDULE', f"O{binding['row']}")
                    wet = engine.value('SCHEDULE', f"P{binding['row']}")
                    self.assertEqual(result['intumescent_thickness_source'], engine.value('SCHEDULE', f"X{binding['row']}"))
                self.assertEqual(result['status'], 'Current')
                self.assertEqual(result['intumescent_dft_um'], native)
                self.assertEqual(result['intumescent_thickness_mm'], native / 1000)
                self.assertNotEqual(result['intumescent_thickness_mm'], wet / 1000)
                self.assertEqual(result['calculator_source_sha256'], binding['source_sha256'])
                self.assertRegex(result['calculator_draft_sha256'], r'^[0-9a-f]{64}$')
                self.assertIn('mm DFT', linked_result_text([result]))
                self.assertEqual(self.case.service.get(self.case.sid)['snapshot'], before)
                self.assertEqual(drafts, original)

    def test_quantity_hold_does_not_discard_valid_native_dft(self):
        inputs = empty_schedule_inputs('steel_intumescent')
        inputs['SETTINGS']['H74'] = 5  # retained legacy Part C
        identity, drafts = self.linked(FX6010, inputs)
        _, engine, lock = calculator_session('steel_intumescent', drafts['steel_intumescent']['inputs'])
        with lock:
            self.assertIn('legacy Part C', engine.value('SCHEDULE', 'W10'))
        result = self.read(drafts)[identity][0]
        self.assertEqual(result['status'], 'Current')
        self.assertEqual(result['intumescent_thickness_mm'], DFT / 1000)

    def test_blocked_thickness_never_becomes_zero_or_usable_result(self):
        for status in ('NP - NOT PUBLISHED', 'BEAM THICKNESS LIMIT', 'UNSUPPORTED CASE', 'ASSESSMENT DATA UNAVAILABLE'):
            with self.subTest(status=status):
                self.lookup.return_value = {**ASSESSED, 'status': status, 'dft': None}
                _session.cache_clear()
                identity, drafts = self.linked()
                result = self.read(drafts)[identity][0]
                self.assertEqual(result['status'], 'Unavailable')
                self.assertIsNone(result['intumescent_thickness_mm'])
                self.assertEqual(result['reason'], status)

    def test_edited_omitted_stale_and_forged_links_withhold_results(self):
        identity, drafts = self.linked()
        binding = self.case.state['snapshot']['transfers'][0]
        for fault in ('edit', 'remove', 'omit'):
            supplied = deepcopy(drafts)
            target = supplied['steel_intumescent']
            if fault == 'edit': target['inputs']['SCHEDULE']['K10'] = .123456789
            elif fault == 'remove': target['schedule_rows'] = [11]
            else: target['inputs']['SCHEDULE'].pop('J10')
            with self.subTest(fault=fault):
                self.assertEqual(self.read(supplied)[identity][0]['status'], 'Unavailable')
        snapshot = deepcopy(self.case.state['snapshot'])
        for changes in ({'id': str(uuid4())}, {'source_sha256': 'b'*64}, {'item_version': 999}, {'status': 'conflict'}):
            forged = deepcopy(snapshot)
            forged['transfers'][0].update(changes)
            result = linked_register_results(forged, {'expected_revision': snapshot['revision'],
                'calculator_drafts': drafts}, store=self.case.store)['linked_results'][identity][0]
            self.assertEqual(result['status'], 'Unavailable')
        self.case.command('update_item', item_id=identity, changes={'quantity': 4})
        self.case.confirm(identity)
        self.assertEqual(self.read(drafts)[identity][0]['status'], 'Unavailable')
        self.case.apply(self.case.preview(identity, 'steel_intumescent', inputs=drafts['steel_intumescent']['inputs'],
            rows=[10], update=True), inputs=drafts['steel_intumescent']['inputs'], rows=[10])
        current = self.case.state['calculator']
        self.assertEqual(self.read({'steel_intumescent': {key: current[key] for key in ('inputs', 'schedule_rows')}})[identity][0]['status'], 'Current')

    def test_repeat_transfer_skips_and_dirty_linked_delete_preserves_then_undo_restores(self):
        identity, drafts = self.linked()
        target = drafts['steel_intumescent']
        skipped = self.case.preview(identity, 'steel_intumescent', inputs=target['inputs'], rows=[10])
        self.assertEqual(skipped['inputs'], target['inputs'])
        self.assertFalse(skipped['changes'])
        self.assertEqual(len(skipped['skipped']), 1)
        dirty = deepcopy(drafts)
        dirty['steel_intumescent']['inputs']['SCHEDULE']['K10'] = .123456789
        before = deepcopy(self.case.state['snapshot'])
        request = lambda values: {'expected_revision': self.case.state['revision'], 'request_id': str(uuid4()),
                                 'item_ids': [identity], 'calculators': values}
        with self.assertRaisesRegex(ValidationError, 'row was edited'):
            self.case.service.delete_linked_items(self.case.sid, request(dirty))
        self.assertEqual(self.case.service.get(self.case.sid)['snapshot'], before)
        self.case.state = self.case.service.delete_linked_items(self.case.sid, request(drafts))
        self.case.state = self.case.service.undo_linked_delete(self.case.sid, {
            'expected_revision': self.case.state['revision'], 'request_id': str(uuid4()),
            'calculators': self.case.state['calculators']})
        self.assertEqual(self.case.state['calculators'], drafts)
        self.assertEqual(self.case.state['snapshot']['items'][0]['id'], identity)
        self.assertEqual(self.case.state['snapshot']['transfers'][0]['status'], 'stale')

    def test_register_and_drawing_use_same_native_mm_without_changing_source(self):
        identity, drafts = self.linked()
        linked = self.read(drafts)
        before = deepcopy(self.case.state['snapshot'])
        payload = self.case.service.export_workspace(self.case.sid, 'register-xlsx', {
            'expected_revision': before['revision'], 'mode': 'steel', 'target': 'steel_intumescent',
            'calculator_drafts': drafts})[0]
        workbook = load_workbook(BytesIO(payload))
        try:
            rows = list(workbook.active.values)
            self.assertEqual(rows[1][rows[0].index('Thickness (mm)')], DFT / 1000)
            self.assertFalse(any(cell.data_type == 'f' for row in workbook.active for cell in row))
        finally:
            workbook.close()
        drawing = deepcopy(before)
        apply_presentation(drawing, {'op': 'toggle_thickness_colours', 'document_id': self.case.doc['id'],
                                    'calculator_id': 'steel_intumescent'}, linked)
        validate_snapshot(drawing)
        original = drawing['drawing_presentation']['colour_modes'][0]['originals'][0]
        self.assertEqual(original['thickness_mm'], DFT / 1000)
        self.assertEqual(original['calculator_id'], 'steel_intumescent')
        self.assertEqual(effective_appearance(drawing, drawing['items'][0], linked)['stroke_color'], original['colour'])
        legend = {'document_id': self.case.doc['id'], 'page': 1, 'mode': 'steel'}
        self.assertIn(f'{DFT/1000:g} mm', legend_rows(drawing, legend, {}, linked)[0]['text'])
        self.assertEqual(drawing['items'], before['items'])
        self.assertEqual(drawing['transfers'], before['transfers'])
        self.assertEqual(self.case.service.get(self.case.sid)['snapshot'], before)

    def test_detach_keeps_manual_inputs_and_removes_only_link(self):
        identity, drafts = self.linked()
        original = deepcopy(drafts)
        self.case.command('detach_transfers', item_ids=[identity], calculator_id='steel_intumescent')
        self.assertEqual(self.read(drafts)[identity], [])
        self.assertEqual(drafts, original)

    def test_intumescent_register_format_retains_submillimetre_dft_display(self):
        self.lookup.return_value = {**ASSESSED, 'dft': 548}
        identity, drafts = self.linked()
        payload = self.case.service.export_workspace(self.case.sid, 'register-xlsx', {
            'expected_revision': self.case.state['revision'], 'mode': 'steel',
            'target': 'steel_intumescent', 'calculator_drafts': drafts})[0]
        workbook = load_workbook(BytesIO(payload))
        try:
            sheet = workbook.active
            columns = {cell.value: cell.column for cell in sheet[1]}
            thickness = sheet.cell(2, columns['Thickness (mm)'])
            self.assertEqual(thickness.value, .548)
            self.assertEqual(thickness.data_type, 'n')
            self.assertEqual(thickness.number_format, '0.000############')
            self.assertEqual(sheet.cell(2, columns['Length each (m)']).number_format, '#,##0.00')
            self.assertEqual(sheet.cell(2, columns['Qty']).number_format, '#,##0')
        finally:
            workbook.close()

    def test_no_silent_spray_exposure_alias_is_created(self):
        identity = self.create(exposure='Re-entrant - 3 sides')
        with self.assertRaisesRegex(ValidationError, 'exact supported'):
            self.case.preview(identity, 'steel_intumescent')


if __name__ == '__main__':
    unittest.main()
