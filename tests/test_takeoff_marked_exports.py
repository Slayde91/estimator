"""Current-register exports and static marked PDFs preserve evidence authority."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import subprocess
import unittest
from unittest.mock import patch
from uuid import uuid4

from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, DecodedStreamObject, DictionaryObject, FloatObject, NameObject, NumberObject, RectangleObject, TextStringObject
from reportlab.pdfgen import canvas

from estimator.catalog import ValidationError
from estimator.takeoff_exports import export_workspace
from estimator.takeoff_markup_pdf_worker import page_transform, transform
from tests import test_takeoff_workspace as fixtures
from tests import test_takeoff_project as project_fixtures


class TakeoffCurrentExportTests(unittest.TestCase):
    def setUp(self):
        self.case = fixtures.TakeoffWorkspaceTests(); self.case.setUp(); self.addCleanup(self.case.doCleanups)

    def export(self, **request):
        snapshot = self.case.state['snapshot']
        return self.case.service.export_workspace(self.case.sid, 'schedule-xlsx',
            {'expected_revision': snapshot['revision'], 'mode': 'steel', **request})

    def rows(self, payload):
        workbook = load_workbook(BytesIO(payload)); self.addCleanup(workbook.close)
        rows = list(workbook['Current Takeoffs'].values)
        return [dict(zip(rows[0], row)) for row in rows[1:]]

    def test_all_current_items_include_drafts_literal_text_unknowns_and_complete_properties(self):
        first = self.case.create(); second = self.case.create(quantity=None, measurement=None)
        self.case.command('update_item', item_id=first, changes={'fields': {'mark': '=2+2', 'notes': 'Retained notes', 'zone': 'Z9'}})
        self.case.command('confirm_items', item_ids=[first]); before = deepcopy(self.case.state['snapshot'])
        payload, kind, filename = self.export()
        rows = self.rows(payload)
        self.assertIn('spreadsheetml', kind); self.assertEqual(filename, 'CEASEFIRE-Steel-Takeoffs.xlsx')
        self.assertEqual(len(rows), 2); self.assertEqual(rows[0]['Mark / Run ID'], '=2+2')
        self.assertEqual(rows[0]['Confirmation'], 'Confirmed'); self.assertEqual(rows[1]['Confirmation'], 'Unconfirmed')
        self.assertEqual(rows[0]['Exposed sides'], 3); self.assertEqual(rows[0]['Critical temperature C'], 550)
        self.assertEqual(rows[0]['Notes'], 'Retained notes'); self.assertEqual(rows[0]['Zone'], 'Z9')
        self.assertIsNone(rows[1]['Quantity']); self.assertIsNone(rows[1]['Total length m'])
        self.assertTrue(rows[1]['Issues']); self.assertIsNone(rows[1]['Confirmation digest'])
        self.assertEqual(json.loads(rows[0]['All item properties'])['critical_temperature'], 550)
        self.assertEqual(self.case.service.get(self.case.sid)['snapshot'], before)
        with self.assertRaisesRegex(ValidationError, 'every item'):
            self.export(item_ids=[first])

    def test_linked_spray_values_come_from_current_engine_and_stale_or_forged_links_are_unavailable(self):
        identifier = self.case.create()
        self.case.command('update_item', item_id=identifier, changes={'fields': {
            'section': '410UB54', 'product': 'MONOKOTE MK-6 HY', 'critical_temperature': 620}})
        self.case.command('confirm_items', item_ids=[identifier])
        self.case.apply(self.case.preview(identifier))
        draft = self.case.state['calculator']; supplied = {draft['id']: {key: draft[key] for key in ('inputs', 'schedule_rows')}}
        row = self.rows(self.export(calculator_drafts=supplied)[0])[0]
        data = json.loads(row['Calculator result provenance'])[0]
        self.assertEqual(data['status'], 'Current'); self.assertGreater(data['estimating_thickness_mm'], 0)
        self.assertIn('mm estimating', row['Linked calculator result'])
        from estimator.workbook_calculators import calculator_session
        _, engine, lock = calculator_session(draft['id'], draft['inputs'])
        with lock: self.assertEqual(data['estimating_thickness_mm'], engine.value('SCHEDULE', 'P10'))
        self.assertIn('Unavailable', self.rows(self.export()[0])[0]['Linked calculator result'])
        edited = deepcopy(supplied); edited[draft['id']]['inputs']['SCHEDULE']['J10'] += 1
        self.assertIn('Unavailable', self.rows(self.export(calculator_drafts=edited)[0])[0]['Linked calculator result'])
        # Receipt-shaped imported text does not regain local binding authority.
        forged = deepcopy(self.case.state['snapshot']); forged['transfers'][0]['id'] = str(uuid4())
        output = export_workspace(forged, {'expected_revision': forged['revision'], 'mode': 'steel', 'calculator_drafts': supplied},
                                  'schedule-xlsx', documents=self.case.documents, store=self.case.store)
        self.assertIn('Unavailable', self.rows(output[0])[0]['Linked calculator result'])

    def test_board_layers_and_thickness_are_native_cells_and_unconfirmed_edits_remove_result(self):
        identifier = self.case.create()
        self.case.command('update_item', item_id=identifier, changes={'fields': {'product': 'TRAFALGAR COREX', 'critical_temperature': 620}})
        self.case.command('confirm_items', item_ids=[identifier]); self.case.apply(self.case.preview(identifier, 'steel_board'))
        draft = self.case.state['calculator']; supplied = {draft['id']: {key: draft[key] for key in ('inputs', 'schedule_rows')}}
        record = json.loads(self.rows(self.export(calculator_drafts=supplied)[0])[0]['Calculator result provenance'])[0]
        from estimator.workbook_calculators import calculator_session
        _, engine, lock = calculator_session('steel_board', draft['inputs'])
        with lock:
            self.assertEqual(record['status'], 'Current')
            for name, column in (('board_layers', 'AA'), ('board_stack_mm', 'Z'), ('total_thickness_mm', 'AB')):
                self.assertEqual(record[name], engine.value('CALCULATOR', column+'9'))
        self.case.command('update_item', item_id=identifier, changes={'quantity': 3})
        row = self.rows(self.export(calculator_drafts=supplied)[0])[0]
        self.assertEqual(row['Confirmation'], 'Unconfirmed'); self.assertIn('Unavailable', row['Linked calculator result'])

    def test_revision_ids_modes_and_calculator_output_injection_are_rejected(self):
        self.case.create()
        for changes in ({'mode': 'physical'}, {'expected_revision': True}, {'expected_revision': -1},
                        {'item_ids': [str(uuid4())]}, {'unexpected': 'data'},
                        {'calculator_drafts': {'steel_board': {'inputs': {}, 'schedule_rows': None}}},
                        {'calculator_drafts': {'steel_board': {'inputs': [], 'schedule_rows': [9]}}},
                        {'calculator_drafts': {'steel_board': {'inputs': {'CALCULATOR': []}, 'schedule_rows': [9]}}},
                        {'calculator_drafts': {'steel_board': {'inputs': {'CALCULATOR': {'AB9': 12}}, 'schedule_rows': [9]}}}):
            with self.subTest(changes=changes), self.assertRaises(ValidationError): self.export(**changes)
        self.case.documents.blocked = True
        with self.assertRaisesRegex(ValidationError, 'Changed'): self.export()

    def test_current_export_http_routes_return_attachments_and_reject_query_options(self):
        from estimator.catalog import ROOT
        from estimator.takeoff_http import TakeoffHTTP
        self.case.create()
        adapter = TakeoffHTTP(self.case.service, ROOT, "default-src 'self'")
        request = {'expected_revision': self.case.state['revision'], 'mode': 'steel'}
        class Handler:
            command = 'POST'
            path = '/api/takeoffs/sessions/' + self.case.sid + '/export/schedule-xlsx'
            def read_json(self): return request
            def send_download(self, *values): self.download = values
        handler = Handler()
        self.assertTrue(adapter.dispatch(handler, handler.path))
        self.assertTrue(handler.download[0].startswith(b'PK'))
        self.assertEqual(handler.download[2], 'CEASEFIRE-Steel-Takeoffs.xlsx')
        route = handler.path; handler.path += '?extra=1'
        with self.assertRaisesRegex(ValidationError, 'structured request'): adapter.dispatch(handler, route)

    def test_multi_document_sources_and_calibration_remain_hash_bound_in_current_xlsx(self):
        identifier = self.case.create(); other = fixtures.document(); other.update(name='Elevation.pdf', sha256='b'*64)
        self.case.state = self.case.service.add_document(self.case.sid, other, self.case.state['revision'])
        self.case.command('record_render', document_id=other['id'], page=1, success=True, warnings=[])
        self.case.command('update_item', item_id=identifier, changes={
            'evidence': [{'document_id': other['id'], 'page': 1, 'note': 'Elevation source'}],
            'length_additions': [{'id': str(uuid4()), 'kind': 'riser', 'length_mm': 2345.678901234567,
                                 'note': 'Cited height', 'document_id': other['id'], 'page': 1}]})
        payload = self.export()[0]; row = self.rows(payload)[0]
        for key in ('Supporting evidence', 'Riser/drop additions'):
            source = json.loads(row[key])[0]
            self.assertEqual(source['document_sha256'], other['sha256']); self.assertEqual(source['document_name'], other['name'])
        basis = json.loads(row['Measurement basis'])['calibration']
        self.assertEqual(basis['document_sha256'], self.case.doc['sha256']); self.assertEqual(basis['id'], self.case.calibration_id)
        self.assertEqual(json.loads(row['Physical member IDs']), self.case.state['snapshot']['items'][0]['member_ids'])
        workbook = load_workbook(BytesIO(payload)); self.addCleanup(workbook.close)
        cells = dict(zip([cell.value for cell in workbook['Current Takeoffs'][1]], workbook['Current Takeoffs'][2]))
        self.assertEqual(cells['Quantity'].number_format, '0'); self.assertEqual(cells['Source page'].number_format, '0')
        self.assertEqual(cells['Total length m'].number_format, '0.00')
        self.assertEqual(cells['Total length m'].value, 2*(10+2345.678901234567/1000))


def drawing_fixture(annotations=False, layers=False):
    buffer = BytesIO(); source = canvas.Canvas(buffer, pagesize=(220, 140), invariant=True)
    for rotation in (0, 90, 180, 270):
        source.setStrokeColorRGB(.7, .7, .7); source.rect(20, 30, 100, 60)
        source.setFont('Helvetica', 7); source.drawString(20, 65, f'ORIGINAL DRAWING ROTATION {rotation}')
        source.line(20, 40, 100, 40); source.showPage()
    source.save(); reader = PdfReader(BytesIO(buffer.getvalue())); writer = PdfWriter()
    for index, page in enumerate(reader.pages):
        page.cropbox = RectangleObject([10, 20, 210, 120]); page[NameObject('/Rotate')] = NumberObject(index*90)
        page.trimbox = RectangleObject([80, 50, 120, 70])
        page[NameObject('/UserUnit')] = FloatObject(2); writer.add_page(page)
    # A source document action is neither imported nor executed by the export.
    writer.add_js('app.alert("must not survive in exported PDF")')
    if annotations:
        link = annotations == 'link'
        border = writer._add_object(ArrayObject([NumberObject(0), NumberObject(0), NumberObject(0)]))
        border_style = writer._add_object(DictionaryObject({NameObject('/W'): NumberObject(0)}))
        annotation = writer._add_object(DictionaryObject({NameObject('/Type'): NameObject('/Annot'),
            NameObject('/Subtype'): NameObject('/Link' if link else '/Text'), NameObject('/Rect'): RectangleObject([20, 30, 40, 50]),
            NameObject('/Border'): border, NameObject('/BS'): border_style,
            NameObject('/Contents'): TextStringObject('Unflattened source evidence')}))
        writer.pages[0][NameObject('/Annots')] = writer._add_object(ArrayObject([annotation]))
    if layers:
        names = ('Same layer name', 'Same layer name') if layers == 'same-names' else ('Visible green geometry', 'Hidden red geometry')
        groups = [writer._add_object(DictionaryObject({NameObject('/Type'): NameObject('/OCG'), NameObject('/Name'): TextStringObject(name)}))
                  for name in names]
        writer._root_object[NameObject('/OCProperties')] = DictionaryObject({NameObject('/OCGs'): ArrayObject(groups),
            NameObject('/D'): DictionaryObject({NameObject('/BaseState'): NameObject('/ON'), NameObject('/OFF'): ArrayObject([groups[1]])})})
        writer.pages[0]['/Resources'][NameObject('/Properties')] = DictionaryObject({NameObject('/Visible'): groups[0], NameObject('/Hidden'): groups[1]})
        stream = DecodedStreamObject()
        stream.set_data(b'/OC /Visible BDC q 0 0.8 0 rg 140 30 20 20 re f Q EMC\n/OC /Hidden BDC q 1 0 0 rg 165 30 20 20 re f Q EMC\n')
        original = writer.pages[0].raw_get('/Contents')
        writer.pages[0][NameObject('/Contents')] = ArrayObject([original, writer._add_object(stream)])
    output = BytesIO(); writer.write(output); return output.getvalue()


class TakeoffMarkedPDFTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): project_fixtures.TakeoffProjectTests.setUpClass()

    def setUp(self):
        self.case = project_fixtures.TakeoffProjectTests(); self.case.pdf = drawing_fixture(); self.case.setUp()
        self.addCleanup(self.case.doCleanups); self.document = self.case.session['snapshot']['documents'][0]

    def command(self, op, **values):
        case = self.case
        case.session = case.service.command(case.session['session_id'], {'op': op, 'request_id': str(uuid4()),
                                  'expected_revision': case.session['revision'], **values})
        return case.session

    def create_marks(self):
        identifiers = []
        for page in range(1, 5):
            self.command('record_render', document_id=self.document['id'], page=page, success=True, warnings=[])
            calibration = str(uuid4())
            self.command('add_calibration', calibration={'id': calibration, 'document_id': self.document['id'], 'page': page,
                'points': [[20, 40], [100, 40]], 'distance_m': 8, 'uniform_scale': True})
            self.command('create_item', item={'mode': 'steel', 'quantity': 2,
                'geometry': {'document_id': self.document['id'], 'page': page, 'points': [[20, 40], [100, 40]]},
                'measurement': {'method': 'calibrated', 'calibration_id': calibration},
                'fields': {'mark': f'B-ROT-{(page-1)*90}', 'section': '100UC15', 'member_type': 'Beam',
                           'exposure': 'Re-entrant - 3 sides', 'fire_period_min': 120}, 'evidence': []})
            identifier = self.case.session['snapshot']['items'][-1]['id']; identifiers.append(identifier)
            self.command('update_item', item_id=identifier, changes={'appearance': {
                'stroke_color': '#ED1476', 'fill_color': '#16699B', 'fill_enabled': False, 'stroke_width': 4, 'opacity': .8}})
        return identifiers

    def export(self, ids, **changes):
        case = self.case
        return case.service.export_workspace(case.session['session_id'], 'marked-pdf', {
            'expected_revision': case.session['revision'], 'mode': 'steel', 'document_id': self.document['id'],
            'item_ids': ids, **changes})

    def test_source_drawings_rotation_crop_userunit_styles_and_readable_legend_survive(self):
        ids = self.create_marks(); before = deepcopy(self.case.session['snapshot'])
        source = self.case.documents.document_path(self.document); original_hash = sha256(source.read_bytes()).hexdigest()
        payload, kind, filename = self.export(ids)
        self.assertEqual((kind, filename), ('application/pdf', 'CEASEFIRE-Marked-Drawing.pdf'))
        reader = PdfReader(BytesIO(payload), strict=True); self.assertEqual(len(reader.pages), 4)
        for index, page in enumerate(reader.pages):
            text = page.extract_text()
            self.assertIn(f'ORIGINAL DRAWING ROTATION {index*90}', text)
            self.assertIn(f'B-ROT-{index*90}', text); self.assertIn('100UC15', text)
            self.assertIn('16.00 m total', text); self.assertIn('Unconfirmed', text); self.assertIn('Unavailable', text)
            self.assertIn('TAKEOFF LEGEND', text); self.assertIn(ids[index], text)
            self.assertNotIn('/Annots', page); self.assertEqual(page.get('/Rotate', 0), 0); self.assertEqual(page.get('/UserUnit', 1), 1)
            self.assertEqual(float(page.mediabox.width), 595)
            self.assertGreater(float(page.mediabox.height), 200 if index%2 == 0 else 400)
            operators = page.get_contents().operations
            self.assertTrue(any(op == b'w' and float(values[0]) == 4 for values, op in operators))
            self.assertTrue(any(op == b'RG' and abs(float(values[0])-237/255) < 1e-5 for values, op in operators))
        self.assertNotIn('/Names', reader.trailer['/Root']); self.assertNotIn('/OpenAction', reader.trailer['/Root'])
        self.assertEqual(sha256(source.read_bytes()).hexdigest(), original_hash)
        self.assertEqual(self.case.service.get(self.case.session['session_id'])['snapshot'], before)

    def test_selection_scope_is_explicit_and_other_pages_remain_unmarked(self):
        ids = self.create_marks(); payload = self.export([ids[1]])[0]; reader = PdfReader(BytesIO(payload))
        self.assertNotIn(ids[0], reader.pages[0].extract_text()); self.assertIn('No visible markups', reader.pages[0].extract_text())
        self.assertIn(ids[1], reader.pages[1].extract_text())
        with self.assertRaises(ValidationError): self.export(ids, mode='duct')
        with self.assertRaises(ValidationError): self.export(ids, document_id=str(uuid4()))
        with self.assertRaises(ValidationError): self.export(ids, item_ids=ids+[ids[0]])

    def test_cited_rectangles_fill_without_converting_source_dimensions(self):
        ids = self.create_marks()
        self.command('update_item', item_id=ids[0], changes={
            'geometry': {'document_id': self.document['id'], 'page': 1, 'points': [[20, 40], [100, 80]]},
            'measurement': {'method': 'cited', 'length_m': 7.123456789, 'citation': 'Cited elevation height'},
            'appearance': {'fill_enabled': True, 'fill_color': '#09B644'}})
        reader = PdfReader(BytesIO(self.export([ids[0]])[0])); page = reader.pages[0]
        self.assertIn('14.25 m total', page.extract_text())
        operations = page.get_contents().operations
        self.assertTrue(any(operator == b'B*' for _, operator in operations), 'Cited region is a closed filled outline.')
        self.assertTrue(any(operator == b'rg' and abs(float(values[1])-182/255) < 1e-5 for values, operator in operations))
        item = self.case.service.get(self.case.session['session_id'])['snapshot']['items'][0]
        self.assertEqual(item['measurement']['length_m'], 7.123456789)

    def test_complete_long_identifiers_wrap_and_paginate_without_truncation(self):
        ids = self.create_marks()
        mark = 'MARK_' + 'LONGIDENTIFIER_'*90 + 'LAST_MARK_TOKEN'
        section = 'SECTION_' + 'COMPLETE_SECTION_'*90 + 'LAST_SECTION_TOKEN'
        self.command('update_item', item_id=ids[0], changes={'fields': {'mark': mark, 'section': section}})
        reader = PdfReader(BytesIO(self.export([ids[0]])[0])); text = ''.join(page.extract_text() for page in reader.pages)
        compact = ''.join(text.split())
        self.assertIn(mark, compact); self.assertIn(section, compact)
        self.assertIn('TAKEOFF LEGEND - CONTINUED', text)
        self.assertIn('Markup details continue on the following legend page.', reader.pages[0].extract_text())
        self.assertNotIn('No visible markups', reader.pages[0].extract_text())
        self.assertNotIn('...', text)

    def test_area_fill_retains_even_odd_exclusion_holes(self):
        self.create_marks(); calibration = self.case.session['snapshot']['calibrations'][0]['id']
        self.command('create_item', item={'mode': 'wall', 'quantity': 1,
            'geometry': {'kind': 'polygon', 'document_id': self.document['id'], 'page': 1,
                         'points': [[20, 30], [120, 30], [120, 80], [20, 80]],
                         'exclusions': [{'id': str(uuid4()), 'points': [[40, 45], [60, 45], [60, 55], [40, 55]], 'note': 'Opening'}]},
            'measurement': {'method': 'calibrated', 'calibration_id': calibration},
            'fields': {'mark': 'W-EXCLUSION'}, 'evidence': []})
        identifier = self.case.session['snapshot']['items'][-1]['id']
        reader = PdfReader(BytesIO(self.export([identifier], mode='wall')[0]))
        self.assertIn('48.00 m2 net', reader.pages[0].extract_text())
        operations = reader.pages[0].get_contents().operations
        painted = [index for index, (_, operator) in enumerate(operations) if operator == b'B*']
        self.assertTrue(painted)
        self.assertEqual(sum(1 for _, operator in operations[painted[0]-13:painted[0]] if operator == b'h'), 2)

    def test_rendered_source_and_markup_use_identical_transform_and_verified_crop(self):
        ids = self.create_marks(); reader = PdfReader(BytesIO(self.export(ids)[0]))
        for index, page in enumerate(reader.pages):
            metadata = self.document['pages'][index]
            _, width, height = page_transform(metadata)
            matrix, _, _ = page_transform(metadata, float(page.mediabox.height)-height, (float(page.mediabox.width)-width)/2)
            operations = page.get_contents().operations
            # The actual source drawing stream is prefixed with the same CTM
            # whose coordinates were used to paint the markup, with the verified
            # crop rather than the deliberately smaller print TrimBox.
            self.assertTrue(any(op == b'cm' and tuple(float(v) for v in values) == matrix for values, op in operations))
            self.assertTrue(any(op == b're' and list(map(float, values)) == [10, 20, 200, 100] for values, op in operations))
            for point in ([20, 40], [100, 40]):
                expected = transform(point, matrix)
                self.assertTrue(any(op in (b'm', b'l') and len(values) == 2 and
                                    all(abs(float(a)-b) < 1e-5 for a, b in zip(values, expected))
                                    for values, op in operations))

    def test_geometry_transform_maps_unrotated_crop_corners_once(self):
        metadata = deepcopy(self.document['pages'][0])
        for rotation, expected in ((0, ((5, 9), (405, 209))), (90, ((5, 409), (205, 9))),
                                   (180, ((405, 209), (5, 9))), (270, ((205, 9), (5, 409)))):
            metadata['rotation'] = rotation; matrix, width, height = page_transform(metadata, 9, 5)
            self.assertEqual(transform([10, 20], matrix), expected[0]); self.assertEqual(transform([210, 120], matrix), expected[1])
            self.assertEqual((width, height), (200, 400) if rotation in (90, 270) else (400, 200))

    def test_annotation_fidelity_tamper_timeout_and_disk_limit_fail_explicitly(self):
        ids = self.create_marks()
        from estimator import takeoff_markup_pdf as module
        with patch.object(module.subprocess, 'run', side_effect=subprocess.TimeoutExpired('worker', 65)):
            with self.assertRaisesRegex(ValidationError, 'time limit'): self.export(ids)
        with patch.object(module.shutil, 'disk_usage', return_value=type('Space', (), {'free': 1})()):
            with self.assertRaisesRegex(ValidationError, 'disk space'): self.export(ids)
        source = self.case.documents.document_path(self.document)
        source.write_bytes(source.read_bytes()+b'\nchanged\n')
        with self.assertRaisesRegex(ValidationError, 'changed'): self.export(ids)
        other = project_fixtures.TakeoffProjectTests(); other.pdf = drawing_fixture(annotations=True); other.setUp(); self.addCleanup(other.doCleanups)
        doc = other.session['snapshot']['documents'][0]
        with self.assertRaisesRegex(ValidationError, 'annotations or form widgets'):
            other.service.export_workspace(other.session['session_id'], 'marked-pdf', {
                'expected_revision': other.session['revision'], 'mode': 'steel', 'document_id': doc['id'], 'item_ids': []})

    def test_layer_default_visibility_keeps_resource_identity_and_invisible_links_are_removed(self):
        other = project_fixtures.TakeoffProjectTests(); other.pdf = drawing_fixture(annotations='link', layers='same-names')
        other.setUp(); self.addCleanup(other.doCleanups); doc = other.session['snapshot']['documents'][0]
        payload = other.service.export_workspace(other.session['session_id'], 'marked-pdf', {
            'expected_revision': other.session['revision'], 'mode': 'steel', 'document_id': doc['id'], 'item_ids': []})[0]
        reader = PdfReader(BytesIO(payload)); root = reader.trailer['/Root']; properties = reader.pages[0]['/Resources']['/Properties']
        self.assertEqual(root['/OCProperties']['/D']['/OFF'][0], properties.raw_get('/Hidden'))
        self.assertNotEqual(properties.raw_get('/Visible'), properties.raw_get('/Hidden'))
        self.assertEqual(properties['/Visible']['/Name'], properties['/Hidden']['/Name'])
        self.assertEqual(root['/OCProperties']['/D']['/BaseState'], '/ON')
        self.assertEqual(set(root['/OCProperties']['/OCGs']), {properties.raw_get('/Visible'), properties.raw_get('/Hidden')})
        self.assertNotIn('/Names', root); self.assertNotIn('/OpenAction', root); self.assertNotIn('/Annots', reader.pages[0])


if __name__ == '__main__': unittest.main()
