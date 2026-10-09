"""Parent review owns child eligibility while historical field bytes stay intact."""
from copy import deepcopy
import json
import unittest
from uuid import uuid4

from estimator.catalog import ValidationError
from estimator.takeoff_library_links import validate_history
from estimator.takeoff_physical import effective_confirmation
from estimator.takeoff_physical_exports import matrix_rows
from estimator.takeoff_physical_markers import barrier_summary
from tests import test_takeoff_library_links as fixtures
from tests import test_takeoff_library_barrier_import as defect_import_fixtures
from tests import test_takeoff_service_plan_library_import as plan_import_fixtures
from tests.test_takeoff_physical_v2 import create
from tests.test_takeoff_service_plans import create as create_plan


class ParentConfirmationImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixtures.TakeoffLibraryLinkTests.setUpClass()

    def fixture(self, kind):
        case = kind()
        case.setUp()
        self.addCleanup(case.doCleanups)
        return case

    def test_defect_library_addition_invalidates_only_review_metadata_not_location_or_source(self):
        case = self.fixture(defect_import_fixtures.LibraryBarrierImportTests)
        owner_id = case.defect['entity']['id']
        case.fixture.review_physical([owner_id])
        before = deepcopy(case.state())
        result, _ = case.apply(case.proposal(item_quantity=3))
        after = result['snapshot']
        old, current = before['physical']['defects'][0], after['physical']['defects'][0]
        self.assertEqual(current, {**old, 'confirmation': 'unconfirmed', 'revision': old['revision'] + 1})
        self.assertEqual(effective_confirmation(after['physical'], after['physical']['services'][0]['id']), 'unconfirmed')
        self.assertEqual(after['physical']['services'][0]['quantity'], 3)
        self.assertIsNone(after['library_assignments']['records'][0]['schedule_binding'])
        event = case.fixture.case.documents.get_blob(after['audit_head'])
        validate_history(event)
        forged = deepcopy(event)
        forged['after']['physical']['defects'][0]['fields']['location'] = 'Forged location'
        with self.assertRaisesRegex(ValidationError, 'cannot change retained'):
            validate_history(forged)
        case.fixture.case.documents.validate_audit(after, owner=case.sid)

    def test_existing_plan_library_addition_invalidates_review_and_preserves_prior_service_facts(self):
        case = self.fixture(plan_import_fixtures.ServicePlanLibraryImportTests)
        owner_id = case.root['entity']['id']
        case.fixture.review_physical([owner_id], 'service_plans')
        before = deepcopy(case.state())
        proposal = case.proposal(existing=True)
        after = case.apply(proposal)[0]['snapshot']
        old, current = before['service_plans']['barriers'][0], after['service_plans']['barriers'][0]
        self.assertEqual(current, {**old, 'confirmation': 'unconfirmed', 'revision': old['revision'] + 1})
        self.assertEqual(after['service_plans']['services'][0], before['service_plans']['services'][0])
        added = after['service_plans']['services'][-1]
        self.assertEqual(effective_confirmation(after['service_plans'], added['id']), 'unconfirmed')
        selection = after['library_assignments']['records'][-1]['barrier_selection']
        self.assertEqual(selection['context_barrier_revision'], old['revision'])
        self.assertEqual(selection['barrier_revision'], current['revision'])
        event = case.fixture.case.documents.get_blob(after['audit_head'])
        validate_history(event)
        for key, value in (('confirmation', 'confirmed'), ('revision', current['revision'] + 1)):
            forged = deepcopy(event)
            forged['after']['service_plans']['barriers'][0][key] = value
            with self.subTest(key=key), self.assertRaises(ValidationError):
                validate_history(forged)
        case.fixture.case.documents.validate_audit(after, owner=case.sid)

    def test_new_plan_root_does_not_invalidate_unrelated_existing_review(self):
        case = self.fixture(plan_import_fixtures.ServicePlanLibraryImportTests)
        case.fixture.review_physical([case.root['entity']['id']], 'service_plans')
        before = deepcopy(case.state())
        after = case.apply(case.proposal(existing=False))[0]['snapshot']
        self.assertEqual(after['service_plans']['barriers'][0], before['service_plans']['barriers'][0])
        self.assertEqual(after['service_plans']['services'][0], before['service_plans']['services'][0])
        self.assertEqual(effective_confirmation(after['service_plans'], after['service_plans']['services'][0]['id']), 'confirmed')
        self.assertEqual(effective_confirmation(after['service_plans'], after['service_plans']['services'][-1]['id']), 'unconfirmed')


class ParentConfirmationProjectionAndProjectTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixtures.TakeoffLibraryLinkTests.setUpClass()

    def setUp(self):
        self.links = fixtures.TakeoffLibraryLinkTests()
        self.links.setUp()
        self.addCleanup(self.links.doCleanups)

    def test_old_child_review_and_location_survive_real_project_save_open_without_overriding_owner(self):
        for scope in ('defect_reports', 'service_plans'):
            with self.subTest(scope=scope):
                # Each native Save As begins at a fresh disposable target.
                self.links = fixtures.TakeoffLibraryLinkTests()
                self.links.setUp()
                self.addCleanup(self.links.doCleanups)
                owner = self.links.member(scope)
                commands = []
                barrier_id = owner
                if scope == 'defect_reports':
                    barrier = create('barrier', owner, location='Legacy barrier location', substrate='Concrete')
                    barrier_id = barrier['entity']['id']
                    commands.append(barrier)
                service = (create if scope == 'defect_reports' else create_plan)('service', barrier_id, service_type='Single Cables')
                commands.append(service)
                self.links.physical(commands, scope)
                self.links.review_physical([service['entity']['id']], scope)
                before = deepcopy(self.links.state()['snapshot'])
                after = deepcopy(before)
                key = 'physical' if scope == 'defect_reports' else 'service_plans'
                graph = after[key]
                # Reconstruct the old independent-child review metadata as a
                # genuine retained draft event, never a technical approval.
                graph['services'][0]['confirmation'] = 'unconfirmed'
                graph['services'][0]['revision'] += 1
                graph['revision'] += 1
                legacy = self.links.service._commit(self.links.sid, {
                    'op': 'apply_physical', 'expected_revision': before['revision'],
                    'request_id': str(uuid4())}, before, after)['snapshot']
                expected = deepcopy(legacy[key])
                case = self.links.case
                case.library.save_as({**deepcopy(case.base), 'takeoffs': legacy,
                    'takeoffs_session_id': self.links.sid, 'penetration': {'draft': deepcopy(self.links.draft)}})
                saved_bytes = case.target.read_bytes()
                saved = json.loads(saved_bytes)
                self.assertEqual(saved['takeoffs'][key], expected)
                case.dialogs.opened = str(case.target)
                reopened = case.library.open_file()
                self.assertEqual(reopened['takeoffs_issues'], [])
                self.assertEqual(reopened['takeoffs'][key], expected)
                self.assertEqual(effective_confirmation(expected, service['entity']['id']), 'confirmed')
                self.assertEqual(case.target.read_bytes(), saved_bytes)

    def test_defect_location_projects_only_owner_value_and_keeps_raw_legacy_barrier_field(self):
        owner = create('defect', location='Defect ground floor')
        barrier = create('barrier', owner['entity']['id'], location='Legacy incorrect upper floor', substrate='Concrete')
        service = create('service', barrier['entity']['id'], service_type='Single Cables')
        self.links.physical([owner, barrier, service])
        graph = self.links.state()['snapshot']['physical']
        before = deepcopy(graph)
        self.assertEqual(matrix_rows(graph)[0][3], 'Defect ground floor')
        label = '\n'.join(barrier_summary(graph, graph['barriers'][0]))
        self.assertIn('Defect ground floor', label)
        self.assertNotIn('Legacy incorrect upper floor', label)
        from estimator.takeoff_physical_exports import export_physical_graph
        import csv
        from io import StringIO
        rows = list(csv.DictReader(StringIO(export_physical_graph(graph, 'csv')[0].decode('utf-8-sig'))))
        self.assertTrue(all(row['location'] == 'Defect ground floor' for row in rows))
        row = next(row for row in rows if row['entity_type'] == 'barrier')
        self.assertEqual(json.loads(row['fields_json'])['location'], 'Legacy incorrect upper floor')
        from io import BytesIO
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(export_physical_graph(graph, 'xlsx')[0]), read_only=True)
        try:
            provenance = dict(workbook['Provenance'].iter_rows(values_only=True))
            self.assertIn('Location columns display the owning Defect', provenance['Facts'])
            self.assertIn('legacy child Location', provenance['Facts'])
            self.assertIn('fields_json', provenance['Facts'])
        finally:
            workbook.close()
        self.assertEqual(graph, before)
        empty = deepcopy(graph)
        empty['defects'][0]['fields'].pop('location')
        self.assertEqual(matrix_rows(empty)[0][3], '')
        self.assertNotIn('Legacy incorrect upper floor', '\n'.join(barrier_summary(empty, empty['barriers'][0])))


if __name__ == '__main__':
    unittest.main()
