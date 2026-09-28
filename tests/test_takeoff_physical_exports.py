"""Lossless draft physical exports, independent of workspace or live sources."""

from copy import deepcopy
import csv
from hashlib import sha256
from io import BytesIO, StringIO
import json
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook

from estimator.catalog import ValidationError
from estimator import takeoff_physical as physical
from estimator import takeoff_physical_exports as exports
from tests.test_takeoff_physical import evidence, proposal, uid


def records(sheet):
    values = list(sheet.iter_rows(values_only=True))
    return [dict(zip(values[0], row)) for row in values[1:]]


class PhysicalExportTests(unittest.TestCase):
    def setUp(self):
        self.graph = physical.new_graph(uid(1000), uid(2000))
        self.add('barrier', 1, fields={'label': 'B1', 'substrate': 'Concrete',
                 'orientation': 'Horizontal', 'thickness_mm': 123.12345678901235})
        self.add('defect', 2, 1, fields={'label': 'D1', 'location': 'Level 1 - Grid A', 'frl': '120/120/120'})
        self.add('opening', 3, 2, fields={'label': 'O1', 'opening_type': 'Corehole', 'size': '100 mm'},
                 evidence=[evidence(occurrence=101, page=1)])
        self.add('service', 4, 3, fields={'label': 'P1', 'service': 'Pipe', 'service_type': 'Copper',
                 'size': '25 mm', 'diameter_mm': 25.123456789012345}, quantity=2,
                 evidence=[evidence(occurrence=103, page=2)])
        self.add('service', 5, 3, fields={'label': 'C1', 'service': 'Cable', 'service_type': 'Cable bundle',
                 'size': '40 mm bundle'}, quantity=1)
        self.add('opening', 6, 2, fields={'opening_type': 'Oversized Opening', 'size': 'Unknown'})

    def change(self, command):
        preview = physical.preview_change(self.graph, command)
        self.graph = physical.apply_change(self.graph, command, expected_revision=self.graph['revision'],
                                           preview_digest=preview['preview_digest'])['graph']

    def add(self, kind, identifier, parent=None, **kwargs):
        self.change(proposal(kind, identifier, parent, **kwargs))

    def csv(self, names=None):
        payload, mime, filename = exports.export_physical_graph(self.graph, 'csv', names)
        self.assertEqual(mime, 'text/csv; charset=utf-8')
        self.assertIn('UNAPPROVED-DRAFT', filename)
        self.assertTrue(payload.startswith(b'\xef\xbb\xbf'))
        return list(csv.DictReader(StringIO(payload.decode('utf-8-sig'))))

    def xlsx(self, names=None):
        payload, mime, filename = exports.export_physical_graph(self.graph, 'xlsx', names)
        self.assertIn('spreadsheetml', mime)
        self.assertIn('UNAPPROVED-DRAFT', filename)
        workbook = load_workbook(BytesIO(payload))
        self.addCleanup(workbook.close)
        return workbook

    def test_csv_keeps_all_stable_hierarchy_links_and_typed_own_facts(self):
        rows = self.csv()
        self.assertEqual(len(rows), 6)
        by_id = {row['entity_id']: row for row in rows}
        service = by_id[uid(4)]
        self.assertEqual((service['barrier_id'], service['defect_id'], service['opening_id'],
                          service['service_id'], service['parent_type'], service['parent_id']),
                         (uid(1), uid(2), uid(3), uid(4), 'opening', uid(3)))
        self.assertEqual((service['label'], service['service'], service['service_type'],
                          service['service_size'], service['quantity']), ('P1', 'Pipe', 'Copper', '25 mm', '2'))
        self.assertEqual((service['substrate'], service['frl'], service['location']), ('', '', ''))
        self.assertEqual(by_id[uid(1)]['substrate'], 'Concrete')
        self.assertEqual(by_id[uid(1)]['orientation'], 'Horizontal')
        self.assertEqual(by_id[uid(2)]['location'], 'Level 1 - Grid A')
        self.assertEqual(by_id[uid(2)]['frl'], '120/120/120')
        self.assertEqual(by_id[uid(3)]['opening_size'], '100 mm')
        self.assertEqual(by_id[uid(3)]['opening_type'], 'Corehole')
        self.assertEqual(json.loads(service['fields_json']), self.graph['services'][0]['fields'])
        self.assertEqual(service['graph_sha256'], physical.graph_digest(self.graph))
        self.assertTrue(all(row[exports.STATUS_COLUMN] == exports.STATUS for row in rows))

    def test_empty_opening_keeps_its_row_without_placeholder_service_or_quantity(self):
        rows = self.csv()
        empty = next(row for row in rows if row['entity_id'] == uid(6))
        self.assertEqual(empty['entity_type'], 'opening')
        self.assertEqual((empty['service_id'], empty['quantity']), ('', ''))
        self.assertEqual(len([row for row in rows if row['entity_type'] == 'service']), 2)
        self.assertTrue(all(row['quantity'] == '' for row in rows if row['entity_type'] != 'service'))
        workbook = self.xlsx()
        self.assertEqual(len(records(workbook['Openings'])), 2)
        self.assertEqual(len(records(workbook['Services'])), 2)
        self.assertNotIn('quantity', [cell.value for cell in workbook['Openings'][1]])
        self.assertNotIn('quantity', [cell.value for cell in workbook['Barriers'][1]])

    def test_xlsx_has_explicit_draft_provenance_and_all_typed_sheets(self):
        workbook = self.xlsx()
        self.assertEqual(workbook.sheetnames, ['Barriers', 'Defects', 'Openings', 'Services',
                                             'Evidence', 'Historical Entities', 'Provenance'])
        info = {row[0].value: row[1].value for row in workbook['Provenance']}
        self.assertEqual(info['Export status'], 'UNAPPROVED DRAFT')
        self.assertIn('unverified assertions', info['Source status'])
        self.assertIn('does not read or verify source bytes', info['Source status'])
        self.assertIn('No approval', info['Authority'])
        self.assertEqual(info['State'], 'draft')
        self.assertEqual(info['Physical graph SHA-256'], physical.graph_digest(self.graph))
        self.assertFalse(any(cell.data_type == 'f' for sheet in workbook for row in sheet for cell in row))
        for sheet in ('Barriers', 'Defects', 'Openings', 'Services', 'Evidence', 'Historical Entities'):
            self.assertIn('UNAPPROVED DRAFT', workbook[sheet]['A1'].value)

    def test_tombstones_stay_in_csv_history_and_evidence_with_exact_parent_links(self):
        self.change({'op': 'delete', 'entity_id': uid(3), 'cascade': True})
        rows = self.csv()
        self.assertEqual(len(rows), 6)
        historical = [row for row in rows if row['record_scope'] == 'historical']
        self.assertEqual({row['entity_id'] for row in historical}, {uid(3), uid(4), uid(5)})
        self.assertTrue(all(row['deleted'] == 'True' for row in historical))
        self.assertTrue(all(row['deleted_at_revision'] == str(self.graph['revision']) for row in historical))
        workbook = self.xlsx()
        self.assertEqual([row['entity_id'] for row in records(workbook['Openings'])], [uid(6)])
        self.assertEqual(records(workbook['Services']), [])
        history = records(workbook['Historical Entities'])
        self.assertEqual({row['entity_id'] for row in history}, {uid(3), uid(4), uid(5)})
        service = next(row for row in history if row['entity_id'] == uid(4))
        self.assertEqual((service['barrier_id'], service['defect_id'], service['opening_id'], service['quantity']),
                         (uid(1), uid(2), uid(3), 2))
        associations = records(workbook['Evidence'])
        self.assertEqual(len(associations), 2)
        self.assertTrue(all(row['record_scope'] == 'historical' for row in associations))

    def test_evidence_keeps_image_hash_occurrence_region_fields_and_optional_names(self):
        ref = evidence(occurrence=101, region=[[0, 0], [20, 0], [20, 10]])
        ref['fields'] = ['size', 'defect_id']
        self.change({'op': 'update', 'entity_id': uid(3), 'changes': {'evidence': [ref]}})
        names = {uid(100): 'Original report.pdf', uid(999): 'Unused name'}
        rows = self.csv(names)
        opening = next(row for row in rows if row['entity_id'] == uid(3))
        self.assertEqual(json.loads(opening['evidence_json']), [ref])
        self.assertEqual(json.loads(opening['source_names_json']), {uid(100): 'Original report.pdf'})
        workbook = self.xlsx(names)
        association = next(row for row in records(workbook['Evidence']) if row['entity_id'] == uid(3))
        self.assertEqual((association['document_id'], association['document_name'], association['document_sha256'],
                          association['page'], association['image_id'], association['image_sha256'], association['occurrence_id']),
                         (uid(100), 'Original report.pdf', 'a'*64, 1, uid(102), 'b'*64, uid(101)))
        self.assertEqual(json.loads(association['region_json']), ref['region'])
        self.assertEqual(json.loads(association['supported_fields_json']), ref['fields'])
        self.assertEqual(json.loads(association['association_json']), ref)
        self.assertEqual(association['association_key'], f'{uid(3)}/1')
        self.assertEqual(association[exports.SOURCE_STATUS_COLUMN], 'UNVERIFIED ASSERTIONS')

    def test_csv_formula_prefixes_are_escaped_and_xlsx_text_remains_literal(self):
        for text in ('=HYPERLINK("bad")', '+SUM(A1)', '-1+2', '@SUM(A1)', '  =A1', '\ttext', '\rtext'):
            with self.subTest(text=text):
                self.change({'op': 'update', 'entity_id': uid(4), 'changes': {'fields': {'label': text}}})
                csv_row = next(row for row in self.csv() if row['entity_id'] == uid(4))
                self.assertEqual(csv_row['label'], "'" + text)
                self.assertEqual(json.loads(csv_row['fields_json']), {'label': text})
                workbook = self.xlsx({uid(100): text})
                sheet = workbook['Services']; headers = {cell.value: cell.column for cell in sheet[1]}
                self.assertEqual(sheet.cell(2, headers['label']).value, text)
                self.assertEqual(sheet.cell(2, headers['label']).data_type, 's')
                self.assertEqual(records(workbook['Evidence'])[0]['document_name'], text)
                self.assertFalse(any(cell.data_type == 'f' for page in workbook for row in page for cell in row))

    def test_shared_occurrence_preserves_distinct_entity_annotation_regions(self):
        opening_ref = evidence(region=[[0, 0], [20, 0], [20, 10]])
        service_ref = evidence(region=[[1, 1], [2, 1], [2, 2]])
        self.change({'op': 'update', 'entity_id': uid(3), 'changes': {'evidence': [opening_ref]}})
        self.change({'op': 'update', 'entity_id': uid(4), 'changes': {'evidence': [service_ref]}})
        associations = records(self.xlsx()['Evidence'])
        self.assertEqual({row['occurrence_id'] for row in associations}, {uid(101)})
        by_entity = {row['entity_id']: row for row in associations}
        self.assertEqual(json.loads(by_entity[uid(3)]['region_json']), opening_ref['region'])
        self.assertEqual(json.loads(by_entity[uid(4)]['region_json']), service_ref['region'])

    def test_exact_numeric_values_and_line_breaks_round_trip_without_formulas(self):
        content = 'First\r\nSecond\rThird\nFourth; "quoted", Ω 😀'
        self.change({'op': 'update', 'entity_id': uid(4), 'changes': {'fields': {
            'label': content, 'diameter_mm': 25.123456789012345, 'notes': ''}}})
        workbook = self.xlsx()
        barrier = records(workbook['Barriers'])[0]
        service = records(workbook['Services'])[0]
        self.assertEqual(barrier['thickness_mm'], 123.12345678901235)
        self.assertEqual(service['diameter_mm'], 25.123456789012345)
        self.assertEqual(service['label'], content)
        self.assertEqual(json.loads(service['fields_json'])['notes'], '')
        row = next(row for row in self.csv() if row['entity_id'] == uid(4))
        self.assertEqual(row['label'], content)
        self.assertEqual(float(row['diameter_mm']), 25.123456789012345)

    def test_long_unicode_json_is_split_losslessly_into_ordered_provenance_chunks(self):
        refs = []
        for index in range(32):
            ref = evidence(occurrence=500+index, page=index+1)
            ref['note'] = '😀' * 1990 + f' {index}'
            refs.append(ref)
        self.change({'op': 'update', 'entity_id': uid(3), 'changes': {'evidence': refs}})
        original = exports._json(refs)
        workbook = self.xlsx()
        opening = records(workbook['Openings'])[0]
        self.assertIn('See Provenance Detail:', opening['evidence_json'])
        parts = [row for row in records(workbook['Provenance Detail'])
                 if row['sheet'] == 'Openings' and row['record_id'] == uid(3) and row['column'] == 'evidence_json']
        self.assertGreater(len(parts), 1)
        self.assertEqual([part['part'] for part in parts], list(range(1, len(parts)+1)))
        self.assertTrue(all(part['total_parts'] == len(parts) for part in parts))
        restored = ''.join(part[exports.CHUNK_HEADERS[-1]] for part in parts)
        self.assertEqual(restored, original)
        self.assertEqual(json.loads(restored), refs)
        self.assertTrue(all(part['sha256'] == sha256(original.encode()).hexdigest() for part in parts))
        self.assertTrue(all(len(part[exports.CHUNK_HEADERS[-1]].encode('utf-16-le')) <= 60000 for part in parts))
        self.assertEqual(len(records(workbook['Evidence'])), 33)
        csv_row = next(row for row in self.csv() if row['entity_id'] == uid(3))
        self.assertEqual(csv_row['evidence_json'], original)

    def test_empty_graph_stays_draft_without_inventing_entity_rows(self):
        self.graph = physical.new_graph(uid(1000), uid(2000))
        payload, _, _ = exports.export_physical_graph(self.graph, 'csv')
        self.assertIn(b'UNAPPROVED DRAFT', payload)
        self.assertIn(b'UNVERIFIED ASSERTIONS', payload)
        self.assertEqual(self.csv(), [])
        workbook = self.xlsx()
        for sheet in ('Barriers', 'Defects', 'Openings', 'Services', 'Evidence', 'Historical Entities'):
            self.assertEqual(records(workbook[sheet]), [])

    def test_export_never_mutates_graph_or_name_lookup_and_labels_do_not_change_digest(self):
        original = deepcopy(self.graph)
        names = {uid(100): '=Untrusted display name'}
        original_names = dict(names)
        first = self.csv(names)
        names[uid(100)] = 'Changed display name'
        second = self.csv(names)
        self.assertEqual([row['graph_sha256'] for row in first], [row['graph_sha256'] for row in second])
        self.assertNotEqual(first[2]['source_names_json'], second[2]['source_names_json'])
        self.xlsx(original_names)
        self.assertEqual(self.graph, original)
        self.assertEqual(original_names, {uid(100): '=Untrusted display name'})

    def test_invalid_or_authoritative_graph_and_invalid_format_are_rejected(self):
        for format in ('pdf', 'CSV', None, [], {'approved': True}):
            with self.subTest(format=format), self.assertRaises(ValidationError):
                exports.export_physical_graph(self.graph, format)
        variants = []
        for key, value in (('state', 'approved'), ('approval', 'draft'), ('physical_model_lock', {})):
            graph = deepcopy(self.graph); graph[key] = value; variants.append(graph)
        for key, value in (('quantity', 0), ('quantity', True), ('quantity', float('nan')), ('opening_id', uid(999))):
            graph = deepcopy(self.graph); graph['services'][0][key] = value; variants.append(graph)
        for graph in variants:
            for format in ('csv', 'xlsx'):
                with self.subTest(format=format, graph=graph), self.assertRaises(ValidationError):
                    exports.export_physical_graph(graph, format)
        with self.assertRaises(TypeError):
            exports.export_physical_graph(self.graph, 'csv', approved=True)

    def test_source_name_lookup_is_bounded_plain_text_only(self):
        for names in ([], {1: 'name'}, {'bad': 'name'}, {uid(100): {'path': 'not read'}},
                      {uid(100): '\x00'}, {uid(100): '\ud800'},
                      {uid(100): 'x'*(exports.MAX_SOURCE_NAME_LENGTH+1)}):
            with self.subTest(names=names), self.assertRaises(ValidationError):
                exports.export_physical_graph(self.graph, 'csv', names)
        with patch.object(exports, 'MAX_SOURCE_NAMES', 1), self.assertRaises(ValidationError):
            exports.export_physical_graph(self.graph, 'csv', {uid(100): 'One', uid(101): 'Two'})
        with patch.object(exports, 'MAX_SOURCE_NAME_TEXT', 3), self.assertRaises(ValidationError):
            exports.export_physical_graph(self.graph, 'csv', {uid(100): 'More'})

    def test_exports_are_byte_deterministic_and_zip_metadata_time_is_not_an_event(self):
        for format in ('csv', 'xlsx'):
            first = exports.export_physical_graph(self.graph, format)[0]
            second = exports.export_physical_graph(deepcopy(self.graph), format)[0]
            self.assertEqual(first, second)
        with ZipFile(BytesIO(first)) as archive:
            self.assertTrue(all(entry.date_time == exports._EXCEL_DATE for entry in archive.infolist()))
            self.assertIn(b'2000-01-01T00:00:00Z', archive.read('docProps/core.xml'))
            self.assertFalse(any('/externalLinks/' in entry.filename or 'vbaProject' in entry.filename for entry in archive.infolist()))


if __name__ == '__main__':
    unittest.main()
