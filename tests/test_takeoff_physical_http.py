"""Real local HTTP integration with disposable synthetic PDF image evidence."""

import hashlib
import http.client
from io import BytesIO
import json
from pathlib import Path
import tempfile
from threading import Thread
import unittest
from uuid import uuid4

from openpyxl import load_workbook

from estimator.server import create_server
from tests.test_takeoff_images_worker import make_pdf


class PhysicalHTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.server = create_server(0, Path(cls.tmp.name)/'physical-qa.sqlite3')
        cls.thread = Thread(target=cls.server.serve_forever, daemon=True); cls.thread.start()
        cls.pdf = make_pdf()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown(); cls.server.server_close(); cls.thread.join(); cls.tmp.cleanup()

    def request(self, method, route, value=None, headers=None):
        connection = http.client.HTTPConnection('127.0.0.1', self.server.server_port, timeout=90)
        try:
            body = json.dumps(value).encode() if isinstance(value, dict) else value
            connection.request(method, route, body, {'Content-Type': 'application/json', **(headers or {})})
            response = connection.getresponse()
            return response.status, dict(response.getheaders()), response.read()
        finally:
            connection.close()

    def json_request(self, method, route, body=None):
        status, _, payload = self.request(method, route, body)
        self.assertEqual(status, 200, payload)
        return json.loads(payload)

    def session(self, *, upload=True, payload=None):
        state = self.json_request('POST', '/api/takeoffs/sessions', {})
        base = '/api/takeoffs/sessions/'+state['session_id']
        if upload:
            payload = self.pdf if payload is None else payload
            upload = self.json_request('POST', base+'/uploads', {'filename': 'synthetic-image.pdf',
                'size': len(payload), 'sha256': hashlib.sha256(payload).hexdigest()})
            endpoint = base+'/uploads/'+upload['upload_id']
            self.assertEqual(self.request('PUT', endpoint+'?offset=0', payload,
                                          {'Content-Type': 'application/octet-stream'})[0], 200)
            state = self.json_request('POST', endpoint+'/complete', {'expected_revision': 0})
        return base, state

    def test_real_image_extract_associate_export_and_same_origin_file(self):
        base, state = self.session()
        document = state['snapshot']['documents'][0]
        request = {'expected_revision': state['revision'], 'request_id': str(uuid4()),
                   'document_id': document['id'], 'first_page': 1, 'page_count': 1}
        state = self.json_request('POST', base+'/images/extract', request)
        self.assertEqual(state['snapshot']['version'], 2)
        self.assertIsNone(state['snapshot']['physical'])
        self.assertEqual(self.json_request('POST', base+'/images/extract', request), state)
        inventory = self.json_request('GET', base+'/images')
        self.assertEqual(inventory['total'], 1)
        self.assertFalse(inventory['has_more'])
        self.assertEqual(self.json_request('GET', base+'/images?extraction_id='+state['extraction_id']+'&offset=0&limit=1')['items'], inventory['items'])
        self.assertEqual(inventory['extractions'][0]['coverage']['complete_pages'], [1])
        self.assertEqual(inventory['extractions'][0]['page_results'][0]['status'], 'complete')
        image = inventory['items'][0]
        route = base+f'/images/{image["extraction_id"]}/{image["asset_id"]}/file'
        status, headers, payload = self.request('GET', route)
        self.assertEqual(status, 200, payload)
        self.assertEqual(headers['Content-Type'], 'image/png')
        self.assertTrue(payload.startswith(b'\x89PNG'))
        self.assertEqual(hashlib.sha256(payload).hexdigest(), image['image_sha256'])
        self.assertEqual(headers['X-Content-Type-Options'], 'nosniff')
        self.assertEqual(headers['Cache-Control'], 'no-store')
        self.assertIn('Content-Security-Policy', headers)
        reference = {key: image[key] for key in ('document_id', 'document_sha256', 'page', 'image_id', 'image_sha256', 'occurrence_id')}
        reference['region'] = [[11, 21], [12, 21], [12, 22]]
        ids = [str(uuid4()) for _ in range(3)]
        commands = []
        for kind, identifier, parent_key, parent in (('defect', ids[0], None, None),
                ('barrier', ids[1], 'defect_id', ids[0]), ('service', ids[2], 'barrier_id', ids[1])):
            entity = {'id': identifier, 'fields': {'label': kind}, 'evidence': [reference],
                      'uncertainty': {'state': 'unresolved', 'note': 'Synthetic draft'}}
            if parent_key: entity[parent_key] = parent
            if kind == 'service': entity['quantity'] = 2
            commands.append({'op': 'create', 'kind': kind, 'entity': entity})
        preview = self.json_request('POST', base+'/physical/preview', {'expected_revision': state['revision'], 'commands': commands})
        self.assertEqual(set(preview['changed_ids']), set(ids))
        state = self.json_request('POST', base+'/physical/apply', {'expected_revision': state['revision'],
            'request_id': str(uuid4()), 'preview_id': preview['preview_id']})
        self.assertEqual(state['snapshot']['physical']['version'], 2)
        self.assertNotIn('openings', state['snapshot']['physical'])
        self.assertEqual(state['snapshot']['physical']['services'][0]['display_id'], 'S-0001')
        self.assertEqual(state['snapshot']['physical']['services'][0]['quantity'], 2)
        self.assertEqual(state['snapshot']['physical']['state'], 'draft')
        for format in ('csv', 'xlsx'):
            status, headers, payload = self.request('POST', base+'/physical/export/'+format, {})
            self.assertEqual(status, 200, payload)
            self.assertIn('UNAPPROVED-DRAFT', headers['Content-Disposition'])
            if format == 'csv': self.assertIn(b'UNAPPROVED DRAFT', payload)
            else:
                with BytesIO(payload) as stream:
                    workbook = load_workbook(stream)
                    self.assertEqual(workbook['Services'].max_row, 2)
                    self.assertNotIn('Openings', workbook.sheetnames)
                    workbook.close()
        other, _ = self.session(upload=False)
        other_route = other+f'/images/{image["extraction_id"]}/{image["asset_id"]}/file'
        self.assertEqual(self.request('GET', other_route)[0], 400)
        self.assertEqual(self.request('POST', base+'/physical/export/csv', {'approved': True})[0], 400)
        self.assertEqual(self.request('POST', base+'/physical/approve', {})[0], 404)

    def test_cropped_image_marker_preserves_original_placement_and_can_be_associated(self):
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import RectangleObject
        writer = PdfWriter(); writer.add_page(PdfReader(BytesIO(self.pdf)).pages[0])
        writer.pages[0].cropbox = RectangleObject([15, 25, 40, 35])
        stream = BytesIO(); writer.write(stream)
        base, state = self.session(payload=stream.getvalue())
        document = state['snapshot']['documents'][0]
        state = self.json_request('POST', base+'/images/extract', {'expected_revision': state['revision'],
            'request_id': str(uuid4()), 'document_id': document['id'], 'first_page': 1, 'page_count': 1})
        image = self.json_request('GET', base+'/images')['items'][0]
        self.assertEqual(image['quad_pdf'], [[10, 20], [50, 20], [50, 40], [10, 40]])
        self.assertTrue(image['region_clipped'])
        self.assertIn('SOURCE_MARKER_CLIPPED', {issue['code'] for issue in image['issues']})
        self.assertTrue(all(15 <= point[0] <= 40 and 25 <= point[1] <= 35 for point in image['region']))
        reference = {key: image[key] for key in ('document_id', 'document_sha256', 'page', 'image_id', 'image_sha256', 'occurrence_id', 'region')}
        command = {'op': 'create', 'kind': 'defect', 'entity': {'id': str(uuid4()), 'fields': {},
            'evidence': [reference], 'uncertainty': {'state': 'unresolved', 'note': 'Cropped source view'}}}
        preview = self.json_request('POST', base+'/physical/preview', {'expected_revision': state['revision'], 'commands': [command]})
        result = self.json_request('POST', base+'/physical/apply', {'expected_revision': state['revision'],
            'request_id': str(uuid4()), 'preview_id': preview['preview_id']})
        self.assertEqual(result['snapshot']['physical']['defects'][0]['evidence'], [reference])

    def test_invalid_ranges_queries_and_stale_previews_fail_without_mutation(self):
        base, state = self.session()
        document = state['snapshot']['documents'][0]
        for first, count in ((0, 1), (1, 2), (1, True), (True, 1)):
            status, _, _ = self.request('POST', base+'/images/extract', {'expected_revision': state['revision'],
                'request_id': str(uuid4()), 'document_id': document['id'], 'first_page': first, 'page_count': count})
            self.assertEqual(status, 400)
        self.assertEqual(self.json_request('GET', base), state)

    def test_matrix_pdf_requires_current_revision_and_preserves_both_scopes(self):
        from pypdf import PdfReader
        from tests.test_takeoff_physical_v2 import create
        base, state = self.session(upload=False)
        preview = self.json_request('POST', base+'/physical/preview', {
            'expected_revision':state['revision'], 'commands':[create('defect',label='Matrix only',frl='-/60/60')]})
        state = self.json_request('POST',base+'/physical/apply',{'expected_revision':state['revision'],
            'request_id':str(uuid4()),'preview_id':preview['preview_id']})
        for revision in (None, True, state['revision']-1,state['revision']+1):
            self.assertEqual(self.request('POST',base+'/physical/export/pdf',
                {'scope':'defect_reports','expected_revision':revision})[0],400)
        status,headers,payload=self.request('POST',base+'/physical/export/pdf',
            {'scope':'defect_reports','expected_revision':state['revision']})
        self.assertEqual(status,200,payload);self.assertEqual(headers['Content-Type'],'application/pdf')
        self.assertIn('Passive_Fire_Matrix.pdf',headers['Content-Disposition'])
        reader=PdfReader(BytesIO(payload));self.assertEqual(reader.metadata.title,'Passive_Fire_Matrix')
        self.assertIn('D-0001',reader.pages[0].extract_text())
        self.assertEqual(self.json_request('GET',base),state)
        self.assertEqual(self.request('GET', base+'/images?approved=true')[0], 400)
        self.assertEqual(self.request('GET', base+'/images?limit=101')[0], 400)
        self.assertEqual(self.request('GET', base+'/images?offset=bad')[0], 400)
        self.assertEqual(self.request('POST', base+'/physical/preview', {'expected_revision': state['revision'], 'commands': [{'op': 'approve'}]})[0], 400)
        self.assertEqual(self.request('POST', base+'/physical/apply', {'expected_revision': state['revision'],
            'request_id': str(uuid4()), 'preview_id': []})[0], 400)
        self.assertEqual(self.json_request('GET', base), state)
