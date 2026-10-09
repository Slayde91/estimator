"""Numbered three-level exports and outer snapshot/audit integration."""

from copy import deepcopy
import csv
from io import BytesIO, StringIO
import json
import unittest

from openpyxl import load_workbook

from estimator import takeoff_physical as physical
from estimator.catalog import ValidationError
from estimator.takeoff_model import audit_affected, new_snapshot, upgrade_snapshot, validate_snapshot
from estimator.takeoff_physical_exports import export_physical_graph
from tests.test_takeoff_physical import evidence, uid
from tests.test_takeoff_physical_exports import records


class PhysicalV2ExportTests(unittest.TestCase):
    def setUp(self):
        self.graph = physical.new_graph(uid(1000), uid(2000), version=2)
        self.add('defect', 1, fields={'label': 'Location issue', 'frl': '120/120/120'})
        self.add('barrier', 2, 1, fields={'substrate': 'Concrete', 'thickness_mm': 120})
        self.add('service', 3, 2, fields={'service': 'Pipe', 'size': '25 mm'}, quantity=2)
        self.add('barrier', 4, 1, fields={'substrate': 'Masonry'})

    def change(self, command):
        preview = physical.preview_change(self.graph, command)
        self.graph = physical.apply_change(self.graph, command,
            expected_revision=self.graph['revision'], preview_digest=preview['preview_digest'])['graph']

    def add(self, kind, number, parent=None, *, fields=None, quantity=None):
        entity = {'id': uid(number), 'fields': fields or {}, 'evidence': [],
                  'uncertainty': {'state': 'not_assessed', 'note': ''}}
        if parent is not None:
            entity[physical.graph_parents(self.graph)[kind][1]] = uid(parent)
        if kind == 'service':
            entity['quantity'] = quantity
        self.change({'op': 'create', 'kind': kind, 'entity': entity})

    def csv_rows(self):
        payload, _, _ = export_physical_graph(self.graph, 'csv')
        reader = csv.DictReader(StringIO(payload.decode('utf-8-sig')))
        return list(reader), reader.fieldnames

    def workbook(self):
        payload, _, _ = export_physical_graph(self.graph, 'xlsx')
        workbook = load_workbook(BytesIO(payload))
        self.addCleanup(workbook.close)
        return workbook

    def test_csv_numbered_ids_and_uuid_links_follow_defect_barrier_service(self):
        rows, headers = self.csv_rows()
        self.assertEqual([row['entity_type'] for row in rows], ['defect', 'barrier', 'barrier', 'service'])
        self.assertEqual(headers[headers.index('defect_id'):headers.index('service_id') + 1],
                         ['defect_id', 'barrier_id', 'service_id'])
        self.assertFalse(any('opening' in header for header in headers))
        self.assertNotIn('shape', headers)
        self.assertNotIn('depth_mm', headers)
        service = next(row for row in rows if row['entity_type'] == 'service')
        self.assertEqual((service['defect_id'], service['barrier_id'], service['service_id']),
                         ('D-0001', 'B-0001', 'S-0001'))
        self.assertEqual((service['defect_uuid'], service['barrier_uuid'], service['service_uuid'],
                          service['parent_id'], service['parent_type']),
                         (uid(1), uid(2), uid(3), uid(2), 'barrier'))
        self.assertEqual(service['quantity'], '2')
        self.assertEqual(service['substrate'], '')
        self.assertEqual(service['frl'], '')
        empty = next(row for row in rows if row['entity_id'] == uid(4))
        self.assertEqual((empty['barrier_id'], empty['service_id'], empty['quantity']), ('B-0002', '', ''))

    def test_xlsx_has_three_typed_sheets_and_tombstones_keep_serials(self):
        self.change({'op': 'delete', 'entity_id': uid(2), 'cascade': True})
        self.add('barrier', 5, 1)
        workbook = self.workbook()
        self.assertEqual(workbook.sheetnames, ['Defects', 'Barriers', 'Services',
                                              'Evidence', 'Historical Entities', 'Provenance'])
        self.assertEqual([row['barrier_id'] for row in records(workbook['Barriers'])], ['B-0002', 'B-0003'])
        self.assertEqual(records(workbook['Services']), [])
        historic = {row['entity_id']: row for row in records(workbook['Historical Entities'])}
        self.assertEqual(historic[uid(2)]['display_id'], 'B-0001')
        self.assertEqual(historic[uid(3)]['service_id'], 'S-0001')
        self.assertEqual(historic[uid(3)]['parent_id'], uid(2))
        for sheet in workbook:
            self.assertFalse(any('opening' in str(cell.value).lower() for cell in sheet[1]))
            self.assertFalse(any(cell.data_type == 'f' for row in sheet for cell in row))

    def test_new_exports_are_deterministic_and_preserve_input(self):
        original = deepcopy(self.graph)
        for format in ('csv', 'xlsx'):
            self.assertEqual(export_physical_graph(self.graph, format)[0],
                             export_physical_graph(deepcopy(self.graph), format)[0])
        self.assertEqual(self.graph, original)

    def test_selected_xlsx_uses_parent_confirmation_and_keeps_child_legacy_fields_in_both_workspaces(self):
        self.add('service', 5, 2, fields={'service': 'Cable', 'notes': '=literal source text'}, quantity=3)
        self.add('service', 6, 2, quantity=4)
        self.change({'op': 'delete', 'entity_id': uid(6), 'cascade': False})
        for version in (2, 3):
            graph = deepcopy(self.graph)
            if version == 3:
                graph['version'] = 3
                del graph['defects']
                for barrier in graph['barriers']:
                    del barrier['defect_id']
                    barrier['fields']['frl'] = '-/120/120'
            graph['services'][0]['confirmation'] = 'unconfirmed'
            graph['services'][0]['evidence'] = [evidence(occurrence=101, page=1)]
            graph['services'][1]['confirmation'] = 'confirmed'
            graph['services'][1]['evidence'] = [evidence(occurrence=102, page=2)]
            graph['services'][2]['confirmation'] = 'confirmed'
            graph['defects' if version == 2 else 'barriers'][0]['confirmation'] = 'confirmed'
            original = deepcopy(graph)
            active = {entry['id'] for collection in physical.graph_collections(graph).values()
                      for entry in graph[collection] if not entry['deleted']}
            reviewed = active if version == 2 else {uid(2), uid(3), uid(5)}
            for selection, expected in (('confirmed', reviewed), ('unconfirmed', active - reviewed), ('all', active)):
                with self.subTest(version=version, selection=selection):
                    payload, mime, filename = export_physical_graph(graph, 'xlsx', confirmation=selection)
                    self.assertIn('spreadsheetml', mime)
                    self.assertIn(selection.title() + '-Items', filename)
                    self.assertEqual(payload, export_physical_graph(deepcopy(graph), 'xlsx', confirmation=selection)[0])
                    workbook = load_workbook(BytesIO(payload))
                    try:
                        rows = [row for collection in physical.graph_collections(graph).values()
                                for row in records(workbook[collection.title()])]
                        self.assertEqual({row['entity_id'] for row in rows}, expected)
                        self.assertTrue(all(row['record_scope'] == 'active' for row in rows))
                        if selection != 'all':
                            self.assertTrue(all(row['confirmation'] == selection for row in rows))
                        self.assertTrue(all(row['graph_sha256'] == physical.graph_digest(graph) for row in rows))
                        self.assertEqual(records(workbook['Historical Entities']), [])
                        self.assertEqual({row['entity_id'] for row in records(workbook['Evidence'])}, expected & {uid(3), uid(5)})
                        service = next((row for row in rows if row['entity_id'] == uid(3)), None)
                        if service:
                            self.assertEqual((service['barrier_id'], service['parent_id'], service['quantity']), ('B-0001', uid(2), 2))
                        self.assertFalse(any(cell.data_type == 'f' for sheet in workbook for row in sheet for cell in row))
                        self.assertIn('Deleted records are excluded', dict(workbook['Provenance'].values)['History'])
                    finally:
                        workbook.close()
            self.assertEqual(graph, original)

    def test_confirmation_selection_is_validated_and_is_xlsx_only(self):
        for selection in ('approved', '', True, [], {}):
            with self.subTest(selection=selection), self.assertRaises(ValidationError):
                export_physical_graph(self.graph, 'xlsx', confirmation=selection)
        for format in ('csv', 'pdf'):
            with self.assertRaises(ValidationError):
                export_physical_graph(self.graph, format, confirmation='all')
        graph = physical.new_graph(uid(1000), uid(2000), version=2)
        workbook = load_workbook(BytesIO(export_physical_graph(graph, 'xlsx', confirmation='confirmed')[0]))
        try:
            self.assertEqual(records(workbook['Services']), [])
            self.assertIn('confirmation', [cell.value for cell in workbook['Services'][1]])
        finally:
            workbook.close()

    def test_outer_snapshot_and_audit_keep_all_three_entity_kinds(self):
        before = upgrade_snapshot(new_snapshot())
        before['project_id'] = self.graph['project_id']
        after = deepcopy(before)
        after['physical'] = self.graph
        self.assertEqual(validate_snapshot(after), after)
        self.assertEqual(audit_affected(before, after)['physical'], sorted(uid(number) for number in (1, 2, 3, 4)))
        self.assertEqual(validate_snapshot(json.loads(json.dumps(after)))['physical'], self.graph)
        changed = deepcopy(after)
        changed['physical']['services'][0]['fields']['notes'] = 'Updated note'
        self.assertEqual(audit_affected(after, changed)['physical'], [uid(3)])


if __name__ == '__main__':
    unittest.main()
