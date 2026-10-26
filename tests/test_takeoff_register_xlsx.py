"""Public register copies match screen data without changing source authority."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import unittest

from openpyxl import load_workbook

from estimator.catalog import ValidationError
from estimator.takeoff_register_xlsx import (AREA, DUCT, DUCT_LENGTH, LENGTH, STEEL, SURFACE,
    export_current_register, export_physical_register, register_workbook)
from tests import test_takeoff_workspace as workspace_fixtures
from tests import test_takeoff_physical_v2_exports as physical_fixtures
from tests.test_takeoff_physical import uid


def read(payload):
    workbook = load_workbook(BytesIO(payload))
    sheet = workbook.active
    return workbook, sheet, list(sheet.values)


def presentation_fixture(mode, target=None, *, recorded_fields=None):
    area = mode == 'walls_floors'
    fields = {'mark': '=literal identifier', 'level': 'Level 1', 'member_type': 'Beam', 'section': '410UB54',
              'fire_period_min': 120, 'sides': 3, 'critical_temperature': 620.123456789,
              'exposure': 'Both', 'width_mm': 600, 'height_mm': 400, 'shape': 'rectangular',
              'frl': '-/120/120', 'orientation': 'Horizontal', 'surface_basis': 'wall-face',
              'substrate': 'Concrete wall', 'treatment': 'Spray', 'system': 'System A', 'product': 'CAFCO 300', 'layers': 2}
    fields.update(recorded_fields or {})
    item = {'id': 'one', 'mode': 'wall' if area else mode, 'fields': fields, 'quantity': 2,
            'geometry': {'document_id': 'drawing', 'page': 1}}
    snapshot = {'documents': [{'id': 'drawing', 'name': 'Drawing 001.pdf'}]}
    results = {'one': {'length_m': 7.123456789012345, 'total_length_m': 14.24691357802469,
                       'gross_area_m2': 32.54, 'excluded_area_m2': 2.21, 'net_area_m2': 30.33,
                       'total_area_m2': 60.66, 'issues': []}}
    linked = {'one': [] if area else [{'calculator_id': target, 'status': 'Current',
        'estimating_thickness_mm': 28.123456789, 'total_thickness_mm': 50, 'duct_thickness_mm': 38,
        'duct_surface_status': 'Current', 'duct_surface_m2': 31.123456789012345,
        'report_references': [{'label': '=Literal report', 'url': 'https://tfire.com.au/documents/report.pdf'}, {'label': 'Private report', 'url': None}]}]}
    return export_current_register(snapshot, [item], results, {'one': False}, linked,
        mode=mode, target=target, confirmation='all')


class RegisterCopyTests(unittest.TestCase):
    def test_each_register_has_exact_data_headers_and_one_sheet(self):
        steel = ('Confirmation', 'Member mark', 'Level', 'Member type', 'Steel section', 'Product',
                 'Fire period (min)', 'Crit. Temp (°C)', 'Exposure')
        length = ('Qty', 'Length each (m)', 'Total (m)', 'Thickness (mm)', 'Report', 'Evidence / issues')
        duct = ('Confirmation', 'Item', 'Level', 'WxH (mm)', 'FRL', 'Orientation', 'Product',
                'Qty', 'Length each (m)', 'Total (m)', 'Duct surface (m²)', 'Thickness (mm)', 'Report', 'Evidence / issues')
        cases = [('steel', 'steel_vermiculite', steel + length),
                 ('steel', 'steel_board', steel + ('Exposed sides',) + length),
                 ('duct', 'ductwork', duct), ('walls_floors', None, SURFACE + AREA)]
        for mode, target, expected in cases:
            with self.subTest(mode=mode, target=target):
                workbook, sheet, rows = read(presentation_fixture(mode, target)[0])
                try:
                    self.assertEqual(workbook.sheetnames, ['Current Takeoffs'])
                    self.assertEqual(tuple(rows[0]), expected)
                    self.assertEqual(len(rows), 2)
                    self.assertEqual(rows[1][1], '=literal identifier')
                    self.assertEqual(sheet.cell(2, 2).data_type, 's')
                    self.assertFalse(any(cell.data_type == 'f' for row in sheet for cell in row))
                    self.assertEqual(sheet['A1'].fill.fgColor.rgb, '00B90A15')
                    self.assertEqual(sheet['A1'].font.color.rgb, '00FFFFFF')
                    self.assertFalse(sheet.sheet_view.showGridLines)
                    self.assertEqual(sheet.freeze_panes, 'C2')
                    self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
                    self.assertEqual(sheet.print_title_rows, '$1:$1')
                    self.assertEqual(sheet.page_setup.orientation, 'landscape')
                    self.assertGreaterEqual(sheet.row_dimensions[1].height, 32)
                    if not mode == 'walls_floors':
                        report = sheet.cell(2, rows[0].index('Report') + 1)
                        self.assertEqual(report.value, '=Literal report — https://tfire.com.au/documents/report.pdf | Private report')
                        self.assertEqual(report.data_type, 's'); self.assertIsNone(report.hyperlink)
                        index = rows[0].index('Length each (m)')
                        self.assertEqual(rows[1][index], 7.123456789012345)
                        self.assertEqual(sheet.cell(2, index+1).number_format, '#,##0.00')
                        self.assertEqual(rows[1][rows[0].index('Product')], 'CAFCO 300')
                        if mode == 'steel':
                            self.assertEqual(rows[1][rows[0].index('Crit. Temp (°C)')], 620.123456789)
                            self.assertEqual(rows[1][rows[0].index('Exposure')], 'Both')
                        else:
                            index = rows[0].index('Duct surface (m²)')
                            self.assertEqual(rows[1][index], 31.123456789012345)
                            self.assertEqual(sheet.cell(2, index+1).number_format, '#,##0.00')
                finally:
                    workbook.close()

    def test_manual_and_legacy_count_rows_use_main_columns_without_invented_length(self):
        for mode, target in [('steel', 'steel_vermiculite'), ('duct', 'ductwork')]:
            quantity = 3 if mode == 'steel' else 1
            total = 3.703703673702 if mode == 'steel' else 1.234567891234
            manual = {'id': 'manual', 'mode': mode, 'count_id': 'retained-batch',
                      'fields': {'mark': 'Manual Count', 'product': 'MONOKOTE'}, 'quantity': quantity,
                      'geometry': {'kind': 'count', 'document_id': 'drawing', 'page': 1}}
            legacy = {'id': 'legacy', 'mode': mode, 'purpose': 'count-only',
                      'count_id': 'old-batch', 'fields': {'mark': 'Legacy Count'}, 'quantity': 4,
                      'geometry': {'kind': 'count', 'document_id': 'drawing', 'page': 2}}
            selected = [manual, legacy]
            snapshot = {'documents': [{'id': 'drawing', 'name': 'Original.pdf'}]}
            results = {'manual': {'length_m': 1.234567891234, 'total_length_m': total, 'issues': []},
                       'legacy': {'length_m': None, 'total_length_m': None, 'issues': []}}
            before = deepcopy((snapshot, selected, results))
            payload = export_current_register(snapshot, selected, results,
                {'manual': False, 'legacy': False}, {'manual': [], 'legacy': []},
                mode=mode, target=target, confirmation='all')[0]
            workbook, sheet, rows = read(payload)
            try:
                self.assertEqual(workbook.sheetnames, ['Current Takeoffs'])
                self.assertEqual(len(rows), 3, 'Count rows share the single main header; no Counts section')
                self.assertEqual([row[1] for row in rows[1:]], ['Manual Count', 'Legacy Count'])
                self.assertEqual(rows[1][rows[0].index('Qty')], quantity)
                self.assertEqual(rows[2][rows[0].index('Qty')], 4)
                self.assertEqual(rows[1][rows[0].index('Length each (m)')], 1.234567891234)
                self.assertEqual(rows[1][rows[0].index('Total (m)')], total)
                for label in ['Length each (m)', 'Total (m)', 'Thickness (mm)']:
                    self.assertIsNone(rows[2][rows[0].index(label)])
                self.assertIn('Original.pdf · p2', rows[2][rows[0].index('Evidence / issues')])
                self.assertNotIn('non-rectangular duct', rows[2][rows[0].index('Evidence / issues')])
                self.assertFalse(any(cell.data_type == 'f' for row in sheet for cell in row))
            finally:
                workbook.close()
            self.assertEqual((snapshot, selected, results), before)

    def test_current_duct_surface_does_not_require_an_available_thickness(self):
        item = {'id': 'one', 'mode': 'duct', 'fields': {'mark': 'D1'}, 'quantity': 1,
                'geometry': {'document_id': 'drawing', 'page': 1}}
        snapshot = {'documents': [{'id': 'drawing', 'name': 'Original.pdf'}]}
        results = {'one': {'length_m': 4.567891234567, 'total_length_m': 4.567891234567, 'issues': []}}
        for status, expected in [('Current', 9.135782469134), ('Unavailable', None)]:
            linked = {'one': [{'calculator_id': 'ductwork', 'status': 'Unavailable',
                'duct_surface_status': status, 'duct_surface_m2': 9.135782469134, 'duct_thickness_mm': None}]}
            workbook, sheet, rows = read(export_current_register(snapshot, [item], results,
                {'one': True}, linked, mode='duct', target='ductwork', confirmation='all')[0])
            try:
                self.assertEqual(rows[1][rows[0].index('Duct surface (m²)')], expected)
                self.assertIsNone(rows[1][rows[0].index('Thickness (mm)')])
            finally:
                workbook.close()

    def test_real_service_receipts_filter_and_source_revision_gates(self):
        case = workspace_fixtures.TakeoffWorkspaceTests(); case.setUp(); self.addCleanup(case.doCleanups)
        confirmed = case.create(); case.command('update_item', item_id=confirmed, changes={'fields': {'mark': 'B-confirmed'}}); case.confirm(confirmed)
        draft = case.create(quantity=None, measurement=None, fields={'mark': 'B-draft'})
        original = deepcopy(case.state['snapshot'])
        def export(**changes):
            return case.service.export_workspace(case.sid, 'register-xlsx', {
                'expected_revision': original['revision'], 'mode': 'steel', 'target': 'steel_vermiculite', **changes})
        for category, marks in [('all', ['B-confirmed', 'B-draft']), ('confirmed', ['B-confirmed']), ('unconfirmed', ['B-draft'])]:
            workbook, sheet, rows = read(export(confirmation=category)[0])
            try:
                self.assertEqual([row[1] for row in rows[1:]], marks)
                if category == 'unconfirmed':
                    self.assertIsNone(rows[1][rows[0].index('Qty')])
                    self.assertIsNone(rows[1][rows[0].index('Thickness (mm)')])
            finally: workbook.close()
        for target in ('arbitrary', True, [], {}):
            with self.subTest(target=target), self.assertRaises(ValidationError): export(target=target)
        with self.assertRaises(ValidationError): export(expected_revision=original['revision']-1)
        with self.assertRaises(ValidationError): export(item_ids=[draft])
        case.documents.blocked = True
        with self.assertRaises(ValidationError): export()
        self.assertEqual(case.state['snapshot'], original)

    def test_retained_fractional_fire_period_and_exposed_sides_display_without_rounding(self):
        for target in ('steel_vermiculite', 'steel_board'):
            workbook, sheet, rows = read(presentation_fixture('steel', target,
                recorded_fields={'fire_period_min': 45.5, 'sides': 2.5})[0])
            try:
                fields = [('Fire period (min)', 45.5)]
                if target == 'steel_board': fields.append(('Exposed sides', 2.5))
                for label, value in fields:
                    cell = sheet.cell(2, rows[0].index(label) + 1)
                    self.assertEqual(cell.value, value)
                    self.assertEqual(cell.number_format, 'General')
                self.assertEqual(sheet.cell(2, rows[0].index('Qty') + 1).number_format, '#,##0')
            finally: workbook.close()

    def test_physical_hierarchy_columns_order_ancestry_and_inherited_review(self):
        case = physical_fixtures.PhysicalV2ExportTests(); case.setUp()
        case.graph['defects'][0]['confirmation'] = 'confirmed'
        original = deepcopy(case.graph)
        workbook, sheet, rows = read(export_physical_register(case.graph, confirmation='confirmed')[0])
        try:
            self.assertEqual(workbook.sheetnames, ['Defect Reports'])
            self.assertEqual(rows[0], ('Defect ID', 'Barrier ID', 'Service ID', 'Confirmation', 'Location', 'FRL',
                'Substrate', 'Orientation', 'Category', 'Service type', 'Service quantity', 'Service Size (mm)', 'Report', 'Source evidence'))
            self.assertEqual([row[:3] for row in rows[1:]], [('D-0001', None, None), ('D-0001', 'B-0001', None),
                ('D-0001', 'B-0001', 'S-0001'), ('D-0001', 'B-0002', None)])
            self.assertTrue(all(row[3] == 'Confirmed' for row in rows[1:]))
            self.assertEqual(rows[3][10], 2)
            self.assertTrue(all(row[5] == '120/120/120' for row in rows[1:]))
        finally: workbook.close()
        self.assertEqual(original, case.graph)
        case.change({'op': 'delete', 'entity_id': uid(3), 'cascade': False})
        workbook, _, rows = read(export_physical_register(case.graph)[0])
        try: self.assertFalse(any(row[2] == 'S-0001' for row in rows[1:]))
        finally: workbook.close()

    def test_empty_filter_stays_a_formatted_header_only_register(self):
        case = physical_fixtures.PhysicalV2ExportTests(); case.setUp()
        workbook, sheet, rows = read(export_physical_register(case.graph, confirmation='confirmed')[0])
        try:
            self.assertEqual(len(rows), 1)
            self.assertEqual(sheet.auto_filter.ref, 'A1:N1')
        finally: workbook.close()

    def test_physical_reference_cells_are_literal_complete_and_do_not_change_graph(self):
        case = physical_fixtures.PhysicalV2ExportTests(); case.setUp(); before = deepcopy(case.graph)
        references = {uid(3): {'reports': [{'label': '=Exact Rev A', 'url': 'https://tfire.com.au/documents/report.pdf'},
                                           {'label': 'Private report', 'url': None}],
                               'systems': [{'id': 'test-1', 'manufacturer': 'Alpha', 'number': 1},
                                           {'id': 'test-2', 'manufacturer': 'Alpha', 'number': 2}]}}
        workbook, sheet, rows = read(export_physical_register(case.graph, references=references)[0])
        try:
            row = next(row for row in rows[1:] if row[2] == 'S-0001')
            self.assertEqual(row[rows[0].index('Report')], '=Exact Rev A — https://tfire.com.au/documents/report.pdf | Private report')
            self.assertNotIn('System', rows[0])
            self.assertFalse(any(cell.data_type == 'f' or cell.hyperlink for cells in sheet for cell in cells))
        finally: workbook.close()
        self.assertEqual(case.graph, before)

    def test_service_size_does_not_infer_a_blank_display_field_from_dimensions(self):
        case = physical_fixtures.PhysicalV2ExportTests(); case.setUp()
        case.change({'op': 'update', 'entity_id': uid(3), 'changes': {'fields': {
            'service': 'Pipe', 'width_mm': 25.123456789, 'height_mm': 40, 'diameter_mm': 20}}})
        original = deepcopy(case.graph)
        self.assertNotIn('size', original['services'][0]['fields'])
        workbook, sheet, rows = read(export_physical_register(case.graph)[0])
        try:
            service = next(row for row in rows[1:] if row[2] == 'S-0001')
            self.assertIsNone(service[rows[0].index('Service Size (mm)')])
            self.assertEqual(service[rows[0].index('Service quantity')], 2)
        finally: workbook.close()
        self.assertEqual(case.graph, original)

    def test_single_sheet_count_table_long_wrapping_and_cell_limit(self):
        payload = register_workbook([(STEEL + LENGTH, []), (('Item', 'Notes'), [['C1', 'A long recorded note. ' * 30]])])
        workbook, sheet, rows = read(payload)
        try:
            self.assertEqual(len(workbook.sheetnames), 1)
            self.assertEqual(rows[2][:2], ('Item', 'Notes'))
            self.assertGreater(sheet.row_dimensions[4].height, 26)
        finally: workbook.close()
        with self.assertRaisesRegex(ValidationError, 'cell limit'):
            register_workbook([(('Item',), [['a' * 32768]])])

    def test_public_browser_routes_use_register_projection(self):
        source = (Path(__file__).parents[1] / 'static' / 'takeoffs.js').read_text(encoding='utf-8')
        self.assertIn('physical/export/${format === "xlsx" ? "register-xlsx" : format}', source)
        self.assertIn('export/${pdf ? format : "register-xlsx"}', source)
        self.assertIn('target: isArea(mode) ? null : state.ui.target.value', source)


if __name__ == '__main__':
    unittest.main()
