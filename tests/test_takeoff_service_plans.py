"""Independent service plans, reviewed source markers and portable draft state."""
from copy import deepcopy
import csv
from hashlib import sha256
from io import BytesIO, StringIO
import json
import unittest
from uuid import uuid4

from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import RectangleObject

from estimator.catalog import ValidationError
from estimator.takeoff_model import audit_affected, new_snapshot, upgrade_snapshot, validate_snapshot
from estimator.takeoff_physical import graph_parents, preview_change
from estimator.takeoff_physical_exports import export_physical_graph
from estimator.takeoff_physical_markers import barrier_summary
from estimator.takeoff_physical_operations import current_graph, prepare_changes
from tests import test_takeoff_physical_v2 as v2
from tests import test_takeoff_workspace as fixtures
from tests import test_takeoff_project as project_fixtures
from tests.test_takeoff_marked_exports import drawing_fixture


def create(kind='barrier', parent=None, *, marker=None, quantity=1, **fields):
    entity = {'id': str(uuid4()), 'fields': fields, 'evidence': [],
              'uncertainty': {'state': 'not_assessed', 'note': ''}}
    if kind == 'service':
        entity.update(barrier_id=parent, quantity=quantity)
    if marker is not None:
        entity['marker'] = marker
    return {'op': 'create', 'kind': kind, 'entity': entity}


class ServicePlansModelTests(unittest.TestCase):
    def setUp(self):
        self.snapshot = new_snapshot()
        self.document = fixtures.document()
        self.snapshot['documents'] = [self.document]
        self.marker = {'document_id': self.document['id'], 'document_sha256': self.document['sha256'],
                       'page': 1, 'point': [50, 50]}
        self.barrier = create(marker=self.marker, substrate='Concrete', frl='-/120/120')
        self.service = create('service', self.barrier['entity']['id'], service='Pipe', quantity=4)

    def graph(self):
        return prepare_changes(self.snapshot, [self.barrier, self.service], scope='service_plans')['graph']

    def test_separate_stable_identity_and_barrier_root_without_synthetic_defect(self):
        graph = self.graph()
        self.assertEqual(graph['version'], 3)
        self.assertEqual(graph_parents(graph), {'service': ('barrier', 'barrier_id')})
        self.assertNotIn('defects', graph)
        self.assertNotIn('defect_id', graph['barriers'][0])
        self.assertEqual(graph['barriers'][0]['display_id'], 'B-0001')
        self.assertNotEqual(current_graph(self.snapshot)['id'], graph['id'])
        self.assertEqual(current_graph(self.snapshot, 'service_plans'), current_graph(self.snapshot, 'service_plans'))
        self.assertNotIn('service_plans', self.snapshot)

    def test_legacy_snapshots_remain_exact_and_service_plans_optional(self):
        legacy = deepcopy(self.snapshot)
        self.assertEqual(validate_snapshot(legacy), legacy)
        current = upgrade_snapshot(legacy)
        self.assertNotIn('service_plans', current)
        self.assertEqual(validate_snapshot(current), current)
        current.update(documents=[self.document], service_plans=self.graph())
        self.assertEqual(validate_snapshot(current), current)
        current['version'] = 1
        with self.assertRaises(ValidationError): validate_snapshot(current)

    def test_typed_scope_parent_and_frl_ownership_reject_injections(self):
        for command in (create('defect'), create('service', self.barrier['entity']['id'], frl='-/60/60'),
                        {**self.barrier, 'entity': {**self.barrier['entity'], 'defect_id': str(uuid4())}}):
            with self.subTest(command=command), self.assertRaises(ValidationError):
                prepare_changes(self.snapshot, [command], scope='service_plans')
        for scope in ('physical', '', None, [], True):
            with self.subTest(scope=scope), self.assertRaises(ValidationError): current_graph(self.snapshot, scope)
        with self.assertRaises(ValidationError): prepare_changes(self.snapshot, [self.barrier])
        wrong = {**upgrade_snapshot(self.snapshot), 'physical': self.graph()}
        with self.assertRaises(ValidationError): validate_snapshot(wrong)

    def test_marker_rejects_bounds_hash_missing_source_extra_fields_and_bad_numbers(self):
        changes = ({'point': [1e12, 0]}, {'point': [float('nan'), 0]}, {'point': [True, 1]},
                   {'point': [[1, 2]]}, {'point': [1, 2, 3]}, {'page': 999}, {'page': True},
                   {'document_sha256': 'f'*64}, {'document_id': str(uuid4())}, {'quantity': 9})
        for patch in changes:
            command = deepcopy(self.barrier); command['entity']['marker'].update(patch)
            with self.subTest(patch=patch), self.assertRaises(ValidationError):
                prepare_changes(self.snapshot, [command], scope='service_plans')
        service = deepcopy(self.service); service['entity']['marker'] = self.marker
        with self.assertRaises(ValidationError):
            prepare_changes(self.snapshot, [self.barrier, service], scope='service_plans')

    def test_marker_movement_removal_and_delete_restore_never_change_service_quantity(self):
        graph = self.graph(); identifier = self.barrier['entity']['id']
        for marker in ({**self.marker, 'point': [60, 70]}, None, self.marker):
            graph = v2.apply(graph, {'op': 'update', 'entity_id': identifier, 'changes': {'marker': marker}})
            self.assertEqual(graph['barriers'][0]['marker'], marker)
            self.assertEqual(graph['services'][0]['quantity'], 4)
        with self.assertRaises(ValidationError):
            preview_change(graph, {'op': 'delete', 'entity_id': identifier, 'cascade': False})
        graph = v2.apply(graph, {'op': 'delete', 'entity_id': identifier, 'cascade': True})
        self.assertEqual(graph['barriers'][0]['marker'], self.marker)
        self.assertTrue(graph['services'][0]['deleted'])
        graph = v2.apply(graph, {'op': 'restore', 'entity_id': identifier, 'mode': 'same_deletion'})
        self.assertEqual(graph['barriers'][0]['marker'], self.marker)
        self.assertEqual(graph['services'][0]['quantity'], 4)

    def test_callout_layout_is_optional_source_space_and_never_changes_service_facts(self):
        graph = self.graph(); identifier = self.barrier['entity']['id']
        layout = {'offset': [-30.123456789, 12.987654321], 'width': 140.123456789, 'height': 50.987654321}
        updated = v2.apply(graph, {'op': 'update', 'entity_id': identifier,
            'changes': {'marker': {**self.marker, 'callout': layout}}})
        self.assertEqual(updated['barriers'][0]['marker']['callout'], layout)
        self.assertEqual(updated['barriers'][0]['marker']['point'], self.marker['point'])
        self.assertEqual(updated['services'], graph['services'])
        self.assertEqual(barrier_summary(updated, updated['barriers'][0]), barrier_summary(graph, graph['barriers'][0]))
        for patch in ({'offset': [0]}, {'offset': [True, 2]}, {'offset': [float('inf'), 2]},
                      {'offset': [-10001, 2]}, {'offset': [2, 10001]}, {'width': 0}, {'height': 10001},
                      {'width': True}, {'height': float('nan')}, {'text': 'Approved'}):
            with self.subTest(patch=patch), self.assertRaises(ValidationError):
                v2.apply(graph, {'op': 'update', 'entity_id': identifier,
                    'changes': {'marker': {**self.marker, 'callout': {**layout, **patch}}}})

    def test_numbered_defect_report_barrier_marker_is_additive(self):
        defect = v2.create('defect', frl='-/90/90')
        barrier = v2.create('barrier', defect['entity']['id'], substrate='Concrete')
        barrier['entity']['marker'] = self.marker
        graph = prepare_changes(self.snapshot, [defect, barrier])['graph']
        self.assertEqual(graph['barriers'][0]['marker'], self.marker)
        self.assertIn('FRL -/90/90', barrier_summary(graph, graph['barriers'][0])[0])
        invalid = deepcopy(defect); invalid['entity']['marker'] = self.marker
        with self.assertRaises(ValidationError): prepare_changes(self.snapshot, [invalid])

    def test_frl_is_derived_from_current_parent_in_exports_and_summary(self):
        graph = self.graph()
        other = create(frl='-/60/60', substrate='Masonry')
        graph = v2.apply(graph, other)
        graph = v2.apply(graph, {'op': 'reparent', 'entity_id': self.service['entity']['id'], 'parent_id': other['entity']['id']})
        csv_bytes, _, filename = export_physical_graph(graph, 'csv')
        rows = list(csv.DictReader(StringIO(csv_bytes.decode('utf-8-sig'))))
        self.assertIn('Service-Plans', filename)
        self.assertNotIn('defect_id', rows[0]); self.assertNotIn('defect_uuid', rows[0])
        service = next(row for row in rows if row['entity_type'] == 'service')
        self.assertEqual(service['frl'], "'-/60/60")
        self.assertNotIn('frl', json.loads(service['fields_json']))
        self.assertEqual(barrier_summary(graph, graph['barriers'][0])[-1], '0 services')
        self.assertIn('4 ×', barrier_summary(graph, graph['barriers'][1])[-1])
        self.assertEqual(json.loads(rows[0]['marker_json']), self.marker)
        workbook = load_workbook(BytesIO(export_physical_graph(graph, 'xlsx')[0])); self.addCleanup(workbook.close)
        self.assertNotIn('Defects', workbook.sheetnames)
        self.assertIn('frl', [cell.value for cell in workbook['Barriers'][1]])

    def test_snapshot_rejects_cross_scope_identity_reuse(self):
        snapshot = upgrade_snapshot(self.snapshot)
        defect = v2.create('defect')
        defect['entity']['id'] = self.barrier['entity']['id']
        snapshot['physical'] = prepare_changes(self.snapshot, [defect])['graph']
        snapshot['service_plans'] = self.graph()
        with self.assertRaisesRegex(ValidationError, 'distinct across'): validate_snapshot(snapshot)


class ServicePlansWorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.case = fixtures.TakeoffWorkspaceTests(); self.case.setUp(); self.addCleanup(self.case.doCleanups)
        self.service, self.sid = self.case.service, self.case.sid
        self.marker = {'document_id': self.case.doc['id'], 'document_sha256': self.case.doc['sha256'],
                       'page': 1, 'point': [50, 50]}
        self.barrier = create(marker=self.marker, frl='-/120/120')

    def state(self): return self.service.get(self.sid, verify_evidence=False)

    def apply(self, commands, scope='service_plans'):
        preview = self.service.preview_physical(self.sid, {'expected_revision': self.state()['revision'], 'commands': commands, 'scope': scope})
        request = {'expected_revision': self.state()['revision'], 'request_id': str(uuid4()), 'preview_id': preview['preview_id'], 'scope': scope}
        return self.service.apply_physical(self.sid, request), request

    def test_scope_preview_apply_digest_and_idempotent_replay_are_bound(self):
        before = self.state()
        preview = self.service.preview_physical(self.sid, {'expected_revision': before['revision'], 'commands': [self.barrier], 'scope': 'service_plans'})
        request = {'expected_revision': before['revision'], 'request_id': str(uuid4()), 'preview_id': preview['preview_id'], 'scope': 'defect_reports'}
        with self.assertRaisesRegex(ValidationError, 'scope'): self.service.apply_physical(self.sid, request)
        self.assertEqual(self.state(), before)
        request['scope'] = 'service_plans'
        result = self.service.apply_physical(self.sid, request)
        self.assertIsNone(result['snapshot']['physical'])
        self.assertEqual(result['snapshot']['service_plans']['barriers'][0]['marker'], self.marker)
        self.assertEqual(self.service.apply_physical(self.sid, request), result)
        with self.assertRaises(ValidationError): self.service.apply_physical(self.sid, {**request, 'scope': 'defect_reports'})
        event = self.case.documents.get_blob(result['snapshot']['audit_head'])
        self.assertEqual(event['affected_ids']['service_plans'], [self.barrier['entity']['id']])
        self.assertEqual(event['affected_ids']['physical'], [])

    def test_independent_graphs_and_undo_preserve_other_scope_and_tombstones(self):
        self.apply([v2.create('defect', label='Existing report')], 'defect_reports')
        reports = deepcopy(self.state()['snapshot']['physical'])
        self.apply([self.barrier])
        self.apply([{'op': 'update', 'entity_id': self.barrier['entity']['id'], 'changes': {'marker': {**self.marker, 'point': [70, 80]}}}])
        result = self.service.command(self.sid, {'op': 'undo', 'expected_revision': self.state()['revision'], 'request_id': str(uuid4())})
        self.assertEqual(result['snapshot']['service_plans']['barriers'][0]['marker'], self.marker)
        self.assertEqual(result['snapshot']['physical'], reports)
        self.apply([{'op': 'delete', 'entity_id': self.barrier['entity']['id'], 'cascade': False}])
        with self.assertRaisesRegex(ValidationError, 'tombstones'):
            self.service.command(self.sid, {'op': 'delete_document', 'document_id': self.case.doc['id'], 'expected_revision': self.state()['revision'], 'request_id': str(uuid4())})
        restored = self.service.command(self.sid, {'op': 'undo', 'expected_revision': self.state()['revision'], 'request_id': str(uuid4())})
        self.assertFalse(restored['snapshot']['service_plans']['barriers'][0]['deleted'])
        self.assertEqual(restored['snapshot']['service_plans']['barriers'][0]['marker'], self.marker)

    def test_callout_layout_applies_with_stale_guards_and_undo_preserves_marker_and_other_scope(self):
        self.apply([v2.create('defect', label='Report kept')], 'defect_reports')
        self.apply([self.barrier]); before = self.state()['snapshot']
        marker = {**self.marker, 'callout': {'offset': [23.125, -11.875], 'width': 130, 'height': 45}}
        result, request = self.apply([{'op': 'update', 'entity_id': self.barrier['entity']['id'], 'changes': {'marker': marker}}])
        self.assertEqual(result['snapshot']['service_plans']['barriers'][0]['marker'], marker)
        self.assertEqual(result, self.service.apply_physical(self.sid, request))
        self.assertEqual(result['snapshot']['physical'], before['physical'])
        undo = self.service.command(self.sid, {'op': 'undo', 'expected_revision': result['revision'], 'request_id': str(uuid4())})
        self.assertEqual(undo['snapshot']['service_plans']['barriers'][0]['marker'], self.marker)
        self.assertEqual(undo['snapshot']['physical'], before['physical'])

    def test_wrong_scope_entity_ids_and_cross_scope_duplicates_are_rejected_in_preview(self):
        self.apply([self.barrier])
        before = self.state()
        for command in ({'op': 'update', 'entity_id': self.barrier['entity']['id'], 'changes': {'fields': {}}},
                        {**v2.create('defect'), 'entity': {**v2.create('defect')['entity'], 'id': self.barrier['entity']['id']}}):
            with self.assertRaises(ValidationError): self.apply([command], 'defect_reports')
            self.assertEqual(before, self.state())

    def test_failed_marker_source_reverification_preserves_snapshot(self):
        preview = self.service.preview_physical(self.sid, {'expected_revision': self.state()['revision'], 'commands': [self.barrier], 'scope': 'service_plans'})
        before = self.state(); self.case.documents.blocked = True
        with self.assertRaises(ValidationError):
            self.service.apply_physical(self.sid, {'expected_revision': before['revision'], 'request_id': str(uuid4()), 'preview_id': preview['preview_id'], 'scope': 'service_plans'})
        self.assertEqual(before, self.state())

    def test_other_scope_edit_invalidates_marker_preview_and_new_graph_undo_retains_ids(self):
        preview = self.service.preview_physical(self.sid, {'expected_revision': self.state()['revision'], 'commands': [self.barrier], 'scope': 'service_plans'})
        self.apply([v2.create('defect')], 'defect_reports')
        with self.assertRaises(ValidationError):
            self.service.apply_physical(self.sid, {'expected_revision': self.state()['revision'], 'request_id': str(uuid4()), 'preview_id': preview['preview_id'], 'scope': 'service_plans'})
        self.apply([self.barrier])
        result = self.service.command(self.sid, {'op': 'undo', 'expected_revision': self.state()['revision'], 'request_id': str(uuid4())})
        retained = result['snapshot']['service_plans']['barriers'][0]
        self.assertTrue(retained['deleted'])
        self.assertEqual(retained['display_id'], 'B-0001')
        self.assertEqual(retained['marker'], self.marker)
        self.apply([create()])
        self.assertEqual(self.state()['snapshot']['service_plans']['barriers'][-1]['display_id'], 'B-0002')


class ServicePlansProjectTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): project_fixtures.TakeoffProjectTests.setUpClass()

    def setUp(self):
        self.case = project_fixtures.TakeoffProjectTests(); self.case.setUp(); self.addCleanup(self.case.doCleanups)
        self.doc = self.case.session['snapshot']['documents'][0]
        self.marker = {'document_id': self.doc['id'], 'document_sha256': self.doc['sha256'], 'page': 1, 'point': [120, 300]}
        self.barrier = create(marker=self.marker, label='Plan barrier', substrate='Concrete', frl='-/120/120')
        self.pipe = create('service', self.barrier['entity']['id'], service='Pipe', diameter_mm=50, quantity=3)
        self.apply([self.barrier, self.pipe])

    def apply(self, commands, scope='service_plans'):
        case = self.case
        preview = case.service.preview_physical(case.session['session_id'], {'expected_revision': case.session['revision'], 'scope': scope, 'commands': commands})
        case.session = case.service.apply_physical(case.session['session_id'], {'expected_revision': case.session['revision'], 'scope': scope,
            'request_id': str(uuid4()), 'preview_id': preview['preview_id']})

    def export(self, **changes):
        case = self.case
        return case.service.export_workspace(case.session['session_id'], 'marked-pdf', {'expected_revision': case.session['revision'],
            'mode': 'penetrations', 'physical_scope': 'service_plans', 'document_id': self.doc['id'],
            'item_ids': [self.barrier['entity']['id']], **changes})

    def test_save_reopen_preserves_both_scopes_marker_bytes_and_calculator_state(self):
        self.apply([v2.create('defect', label='Report kept')], 'defect_reports')
        case = self.case; before = deepcopy(case.session['snapshot'])
        request = {**deepcopy(case.base), 'takeoffs': before, 'takeoffs_session_id': case.session['session_id']}
        case.library.save_as(request)
        saved = json.loads(case.target.read_bytes())
        case.dialogs.opened = str(case.target); reopened = case.library.open_file()
        self.assertEqual(reopened['takeoffs_issues'], [])
        for key in ('physical', 'service_plans'):
            self.assertEqual(before[key], saved['takeoffs'][key]); self.assertEqual(before[key], reopened['takeoffs'][key])
        for key in ('estimate', 'calculators'):
            self.assertEqual(saved[key], json.loads(case.legacy)[key])
        case.documents.validate_audit(reopened['takeoffs'])

    def test_marked_pdf_contains_automatic_summary_and_does_not_mutate_source(self):
        case = self.case; source = case.documents.document_path(self.doc); checksum = sha256(source.read_bytes()).hexdigest()
        before = deepcopy(case.session['snapshot'])
        payload, mime, _ = self.export(); self.assertEqual(mime, 'application/pdf')
        text = ' '.join(' '.join(page.extract_text().split()) for page in PdfReader(BytesIO(payload)).pages)
        for expected in ('B-0001', 'Plan barrier', 'FRL -/120/120', 'Concrete', 'S-0001', 'Pipe', 'Ø 50 mm', '3 ×'):
            self.assertIn(expected, text)
        self.assertNotIn('TAKEOFF LEGEND', text)
        self.assertNotIn('see legend', text)
        self.assertNotIn(self.barrier['entity']['id'], text)
        self.assertNotIn('3 markers', text)
        self.assertEqual(sha256(source.read_bytes()).hexdigest(), checksum)
        self.assertEqual(case.service.get(case.session['session_id'])['snapshot'], before)
        self.apply([{'op': 'update', 'entity_id': self.pipe['entity']['id'], 'changes': {'fields': {'service': 'Cable'}}}])
        text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(self.export()[0])).pages)
        self.assertIn('Cable', text); self.assertNotIn('Ø 50 mm', text)

    def test_resized_positioned_callout_survives_pdf_export_save_and_reopen(self):
        case = self.case
        marker = deepcopy(case.session['snapshot']['service_plans']['barriers'][0]['marker'])
        marker['callout'] = {'offset': [60.125, -35.875], 'width': 180.25, 'height': 65.75}
        self.apply([{'op': 'update', 'entity_id': self.barrier['entity']['id'], 'changes': {'marker': marker}}])
        before = deepcopy(case.session['snapshot']); source = case.documents.document_path(self.doc).read_bytes()
        payload = self.export()[0]; reader = PdfReader(BytesIO(payload))
        text = '\n'.join(page.extract_text() for page in reader.pages)
        for expected in ('B-0001', 'Plan barrier', 'Concrete', '3 ×', 'Ø 50 mm'): self.assertIn(expected, text)
        self.assertEqual(case.service.get(case.session['session_id'])['snapshot'], before)
        self.assertEqual(case.documents.document_path(self.doc).read_bytes(), source)
        request = {**deepcopy(case.base), 'takeoffs': before, 'takeoffs_session_id': case.session['session_id']}
        case.library.save_as(request); case.dialogs.opened = str(case.target)
        reopened = case.library.open_file()
        self.assertEqual(reopened['takeoffs']['service_plans']['barriers'][0]['marker'], marker)
        self.assertEqual(reopened['takeoffs_issues'], [])
        saved = json.loads(case.target.read_bytes())
        for key in ('estimate', 'calculators'): self.assertEqual(saved[key], json.loads(case.legacy)[key])

    def test_upright_callout_dimensions_export_consistently_at_every_rotation_and_user_unit(self):
        # Enough visible space for this box at all rotations avoids confusing
        # dimension preservation with the intentional page-edge clamping.
        writer = PdfWriter()
        for page in PdfReader(BytesIO(drawing_fixture())).pages:
            page.mediabox = RectangleObject([0, 0, 600, 400])
            page.cropbox = RectangleObject([10, 20, 590, 380])
            writer.add_page(page)
        source_pdf = BytesIO(); writer.write(source_pdf)
        case = project_fixtures.TakeoffProjectTests(); case.pdf = source_pdf.getvalue(); case.setUp()
        self.addCleanup(case.doCleanups)
        document = case.session['snapshot']['documents'][0]
        layout = {'offset': [15.125, -20.875], 'width': 180.25, 'height': 65.75}
        barriers = [create(marker={'document_id': document['id'], 'document_sha256': document['sha256'],
            'page': page, 'point': [120, 180], 'callout': layout}, label=f'Rotation {90*(page-1)}')
            for page in range(1, 5)]
        preview = case.service.preview_physical(case.session['session_id'], {
            'expected_revision': case.session['revision'], 'scope': 'service_plans', 'commands': barriers})
        case.session = case.service.apply_physical(case.session['session_id'], {
            'expected_revision': case.session['revision'], 'scope': 'service_plans',
            'request_id': str(uuid4()), 'preview_id': preview['preview_id']})
        before = deepcopy(case.session['snapshot'])
        payload = case.service.export_workspace(case.session['session_id'], 'marked-pdf', {
            'expected_revision': case.session['revision'], 'mode': 'penetrations', 'physical_scope': 'service_plans',
            'document_id': document['id'], 'item_ids': [barrier['entity']['id'] for barrier in barriers]})[0]
        pages = PdfReader(BytesIO(payload)).pages
        self.assertEqual(len(pages), 4)
        for index, page in enumerate(pages):
            # Inspect the actual painted rounded box in the exported PDF, not
            # an intermediate spec that could still be drawn with swapped axes.
            painted_boxes, path = [], []
            for values, operator in page.get_contents().operations:
                if operator == b'n': path = []
                elif operator in (b'm', b'l', b'c'):
                    path.extend(zip(map(float, values[::2]), map(float, values[1::2])))
                elif operator in (b'B', b'B*', b'b', b'b*'):
                    if path:
                        xs, ys = zip(*path)
                        painted_boxes.append((max(xs)-min(xs), max(ys)-min(ys)))
                    path = []
                elif operator in (b'S', b's', b'f', b'F', b'f*'): path = []
            with self.subTest(rotation=90*index):
                self.assertTrue(any(abs(width-360.5) < 1e-5 and abs(height-131.5) < 1e-5
                    for width, height in painted_boxes), painted_boxes)
                self.assertIn(f'Rotation {90*index}', page.extract_text())
        self.assertEqual(case.service.get(case.session['session_id'])['snapshot'], before)
        self.assertEqual(case.documents.document_path(document).read_bytes(), source_pdf.getvalue())

    def test_marked_pdf_rejects_scope_other_entity_and_client_summary_injection(self):
        for changes in ({'physical_scope': 'defect_reports'}, {'item_ids': [self.pipe['entity']['id']]},
                        {'physical_scope': None}, {'summary': 'Approved'}, {'document_id': str(uuid4())},
                        {'item_ids': [self.barrier['entity']['id']]*2}):
            with self.subTest(changes=changes), self.assertRaises(ValidationError): self.export(**changes)

    def test_blank_source_pages_survive_with_browser_zoom_and_rotation_without_legends(self):
        case=project_fixtures.TakeoffProjectTests();case.pdf=drawing_fixture();case.setUp();self.addCleanup(case.doCleanups)
        document=case.session['snapshot']['documents'][0]
        command=create(marker={'document_id':document['id'],'document_sha256':document['sha256'],
            'page':1,'point':[60,60]},label='One marked page')
        preview=case.service.preview_physical(case.session['session_id'],{'expected_revision':case.session['revision'],
            'scope':'service_plans','commands':[command]})
        case.session=case.service.apply_physical(case.session['session_id'],{'expected_revision':case.session['revision'],
            'scope':'service_plans','request_id':str(uuid4()),'preview_id':preview['preview_id']})
        before=deepcopy(case.session['snapshot'])
        payload=case.service.export_workspace(case.session['session_id'],'marked-pdf',{'expected_revision':case.session['revision'],
            'mode':'penetrations','physical_scope':'service_plans','document_id':document['id'],
            'item_ids':[command['entity']['id']],'rendering':{'zoom':.32,'rotations':{'1':90,'3':270}}})[0]
        pages=PdfReader(BytesIO(payload)).pages;self.assertEqual(len(pages),len(document['pages']))
        from estimator.takeoff_markup_pdf_worker import page_transform
        for index,page in enumerate(pages):
            metadata=dict(document['pages'][index]);metadata['rotation']=(metadata['rotation']+{0:90,2:270}.get(index,0))%360
            _,width,height=page_transform(metadata)
            self.assertAlmostEqual(float(page.mediabox.width),width);self.assertAlmostEqual(float(page.mediabox.height),height)
            text=page.extract_text();self.assertNotIn('TAKEOFF LEGEND',text);self.assertNotIn('see legend',text)
        self.assertEqual(case.service.get(case.session['session_id'])['snapshot'],before)
        for rendering in ({'zoom':True,'rotations':{}},{'zoom':.32,'rotations':{'0':90}},
                          {'zoom':.32,'rotations':{'1':45}},{'zoom':float('inf'),'rotations':{}}):
            with self.subTest(rendering=rendering),self.assertRaises(ValidationError): self.export(rendering=rendering)

    def test_saved_audit_rejects_removed_graph_or_rewritten_numbered_identity(self):
        case = self.case; before = case.session['snapshot']
        for mutation in ('remove_graph', 'remove_barrier', 'renumber', 'uuid'):
            after = deepcopy(before); after['revision'] += 1
            if mutation == 'remove_graph': after.pop('service_plans')
            elif mutation == 'remove_barrier':
                after['service_plans']['barriers'] = []; after['service_plans']['services'] = []
            elif mutation == 'renumber': after['service_plans']['barriers'][0]['display_id'] = 'B-0002'
            else: after['service_plans']['services'][0]['id'] = str(uuid4())
            strip = lambda value: {key: entry for key, entry in value.items() if key != 'audit_head'}
            event = {'version': 2, 'project_id': before['project_id'], 'revision': after['revision'],
                'previous': before['audit_head'], 'request_id': str(uuid4()), 'op': 'apply_physical',
                'at': '2026-10-03T00:00:00Z', 'actor': {'kind': 'local-session', 'session_id': case.session['session_id']},
                'before': strip(before), 'after': strip(after), 'affected_ids': audit_affected(before, after)}
            after['audit_head'] = case.documents.put_blob(event)
            with self.subTest(mutation=mutation), self.assertRaises(ValidationError): case.documents.validate_audit(after)


if __name__ == '__main__': unittest.main()
