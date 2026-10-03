"""Source-only counts and lengths never acquire calculator transfer meaning."""
from copy import deepcopy
from io import BytesIO
import json
import unittest
from uuid import uuid4

from openpyxl import load_workbook
from pypdf import PdfReader

from estimator.catalog import ValidationError
from estimator.takeoff_exports import HEADERS, register_rows
from estimator.takeoff_model import digest, item_result, validate_snapshot
from estimator.takeoff_transfer import mapped_values
from tests import test_takeoff_workspace as fixtures
from tests import test_takeoff_project as project_fixtures
from tests.test_takeoff_marked_exports import drawing_fixture


FIELDS = {'mark': 'Ceiling grilles', 'level': 'Level 1', 'width_mm': 600, 'height_mm': 400,
          'frl': '-/60/60', 'orientation': 'Horizontal'}


class StandaloneTests(unittest.TestCase):
    def setUp(self):
        self.case = fixtures.TakeoffWorkspaceTests(); self.case.setUp()
        self.addCleanup(self.case.doCleanups)

    def count(self, mode='duct'):
        return self.case.command('add_standalone_count', mode=mode, document_id=self.case.doc['id'], page=1,
            markers=[{'point': [30, 40]}, {'point': [50, 60]}], fields=deepcopy(FIELDS), appearance={})['created_item_ids'][0]

    def length(self, mode='wall'):
        self.case.command('create_item', item={'mode': mode, 'purpose': 'length-only',
            'geometry': {'document_id': self.case.doc['id'], 'page': 1, 'points': [[20, 40], [70, 40], [70, 60]]},
            'measurement': {'method': 'calibrated', 'calibration_id': self.case.calibration_id},
            'quantity': 1, 'fields': {'mark': 'Opening perimeter'}})
        return self.case.state['snapshot']['items'][-1]['id']

    def item(self, identifier):
        return next(value for value in self.case.state['snapshot']['items'] if value['id'] == identifier)

    def rejected(self, op, **values):
        before = deepcopy(self.case.state['snapshot']); audit = deepcopy(self.case.documents.blobs)
        with self.assertRaises(ValidationError): self.case.command(op, **values)
        self.assertEqual(self.case.service.get(self.case.sid)['snapshot'], before)
        self.assertEqual(self.case.documents.blobs, audit)

    def test_counts_have_exact_markers_without_scale_length_or_calculator_requirements(self):
        self.case.state = self.case.service.open(); self.case.sid = self.case.state['session_id']
        self.case.state = self.case.service.add_document(self.case.sid, self.case.doc, 0)
        self.case.command('record_render', document_id=self.case.doc['id'], page=1, success=True, warnings=[])
        for mode in ('steel', 'duct'):
            identifier = self.count(mode); item = self.item(identifier)
            result = item_result(item, self.case.state['snapshot'])
            self.assertEqual((item['purpose'], item['measurement'], item['quantity']), ('count-only', None, 2))
            self.assertEqual((result['total_count'], result['length_m'], result['total_length_m'], result['issues']), (2, None, None, []))
            self.case.command('confirm_items', item_ids=[identifier])
            self.assertEqual(self.item(identifier)['confirmation']['checks']['engine'], 'takeoffs-count-v1')
            self.assertEqual(self.item(identifier)['confirmation']['checks']['total_count'], 2)
        self.assertEqual(validate_snapshot(self.case.state['snapshot']), self.case.state['snapshot'])

    def test_count_continuation_move_delete_and_undo_preserve_independent_identities(self):
        identifier = self.count(); original = deepcopy(self.item(identifier))
        request = {'op': 'continue_standalone_count', 'request_id': str(uuid4()), 'expected_revision': self.case.state['revision'],
                   'item_id': identifier, 'markers': [{'point': [80, 50]}]}
        self.case.state = self.case.service.command(self.case.sid, request)
        self.assertEqual(self.case.state, self.case.service.command(self.case.sid, request))
        current = self.item(identifier)
        self.assertEqual(current['member_ids'][:2], original['member_ids']); self.assertEqual(current['quantity'], 3)
        member = current['member_ids'][-1]
        self.case.command('move_count_markers', markers=[{'item_id': identifier, 'member_id': member}], delta_pdf=[10, 5])
        self.assertEqual(self.item(identifier)['geometry']['points'], [[30, 40], [50, 60], [90, 55]])
        self.case.command('update_item', item_id=identifier, changes={'fields': {'mark': 'Updated grille count', 'level': 'Level 2'}})
        self.assertEqual(self.item(identifier)['fields']['mark'], 'Updated grille count')
        self.case.command('delete_count_marker', item_id=identifier, member_id=member)
        self.assertEqual((self.item(identifier)['quantity'], self.item(identifier)['member_ids']), (2, original['member_ids']))
        self.case.command('undo')
        self.assertEqual(self.item(identifier)['geometry']['points'][-1], [90, 55])
        self.case.command('delete_items', item_ids=[identifier]); self.assertEqual(self.case.state['snapshot']['items'], [])

    def test_invalid_count_edits_are_atomic_and_legacy_length_ops_cannot_convert_counts(self):
        identifier = self.count(); member = self.item(identifier)['member_ids'][0]
        for markers in ([], [{'point': [0, 0]}], [{'point': [30, 40], 'length_m': 3}], [{'point': [True, 40]}]):
            self.rejected('continue_standalone_count', item_id=identifier, markers=markers)
        for changes in ({'quantity': 3}, {'measurement': {'method': 'manual', 'length_m': 5}},
                        {'fields': {'product': 'FyreWrap'}}, {'length_additions': [{'id': str(uuid4())}]},
                        {'geometry': {**self.item(identifier)['geometry'], 'kind': 'count'}}):
            self.rejected('update_item', item_id=identifier, changes=changes)
        self.rejected('continue_count', item_id=identifier, markers=[{'point': [80, 60], 'length_m': 3}])
        self.rejected('update_count_lengths', groups=[{'member_ids': [member], 'length_m': 5}])
        self.rejected('split_steel_group', item_id=identifier, quantities=[1, 1])

    def test_wall_and_slab_lengths_are_open_calibrated_lines_with_exact_precision(self):
        for mode in ('wall', 'slab'):
            identifier = self.length(mode)
            result = item_result(self.item(identifier), self.case.state['snapshot'])
            self.assertEqual((result['length_m'], result['total_length_m'], result['issues']), (7, 7, []))
            self.assertNotIn('net_area_m2', result)
            self.case.command('confirm_items', item_ids=[identifier])
            self.assertEqual(self.item(identifier)['confirmation']['checks']['engine'], 'takeoffs-length-v1')
            self.case.command('update_item', item_id=identifier, changes={'geometry': {
                **self.item(identifier)['geometry'], 'points': [[20, 40], [73.123456789, 40]]}})
            self.assertEqual(item_result(self.item(identifier), self.case.state['snapshot'])['length_m'], 5.3123456789)
            self.assertEqual(self.item(identifier)['state'], 'draft')

    def test_calibration_changes_recompute_lengths_without_creating_surface_areas(self):
        identifier = self.length(); self.case.command('confirm_items', item_ids=[identifier])
        self.case.command('update_calibration', calibration_id=self.case.calibration_id, changes={'distance_m': 20})
        self.assertEqual(item_result(self.item(identifier), self.case.state['snapshot'])['length_m'], 14)
        self.assertEqual(self.item(identifier)['state'], 'draft')
        self.rejected('update_item', item_id=identifier, changes={'quantity': 2})
        self.rejected('update_item', item_id=identifier, changes={'geometry': {
            **self.item(identifier)['geometry'], 'points': [[20, 40], [20, 40]]}})
        self.rejected('update_item', item_id=identifier, changes={'measurement': {'method': 'cited', 'length_m': 8, 'citation': 'Plan'}})

    def test_all_calculator_targets_reject_standalone_even_after_confirmation_and_import_forgery(self):
        for identifier in (self.count('steel'), self.count('duct'), self.length('wall'), self.length('slab')):
            self.case.command('confirm_items', item_ids=[identifier])
            for calculator_id, row in (('steel_vermiculite', 10), ('steel_board', 9), ('ductwork', 11)):
                with self.subTest(identifier=identifier, calculator=calculator_id):
                    before = deepcopy(self.case.state['snapshot'])
                    with self.assertRaisesRegex(ValidationError, 'cannot be transferred'):
                        mapped_values(self.item(identifier), before, calculator_id, row)
                    with self.assertRaisesRegex(ValidationError, 'cannot be transferred'):
                        self.case.preview(identifier, calculator_id)
                    self.assertEqual(self.case.service.get(self.case.sid)['snapshot'], before)
        forged = deepcopy(self.case.state['snapshot']); item = forged['items'][0]
        forged['transfers'] = [{'id': str(uuid4()), 'item_id': item['id'], 'item_version': item['version'],
            'item_digest': 'a'*64, 'calculator_id': 'steel_vermiculite', 'sheet': 'SCHEDULE', 'row': 10,
            'source_sha256': 'b'*64, 'input_hash': digest({}), 'values': {}, 'status': 'current'}]
        with self.assertRaisesRegex(ValidationError, 'cannot retain'): validate_snapshot(forged)

    def test_register_exports_show_counts_without_false_lengths_and_lengths_without_false_area(self):
        identifiers = [self.count(), self.length()]
        for identifier in identifiers: self.case.command('confirm_items', item_ids=[identifier])
        rows = {value['Purpose']: value for row in register_rows(self.case.state['snapshot'], self.case.state['snapshot']['items'])
                for value in [dict(zip(HEADERS, row))]}
        self.assertEqual((rows['count-only']['Quantity'], rows['count-only']['Length per item m']), (2, None))
        self.assertEqual(json.loads(rows['count-only']['Measurement basis']), {'method': 'count', 'quantity': 2})
        self.assertEqual((rows['length-only']['Length per item m'], rows['length-only']['Net area m2']), (7, None))
        for mode, identifier in zip(('duct', 'wall'), identifiers):
            payload = self.case.service.export_workspace(self.case.sid, 'schedule-xlsx', {
                'expected_revision': self.case.state['revision'], 'mode': mode, 'item_ids': [identifier]})[0]
            workbook = load_workbook(BytesIO(payload)); sheet = workbook['Current Takeoffs']
            values = dict(zip([cell.value for cell in sheet[1]], [cell.value for cell in sheet[2]]))
            self.assertEqual(values['Purpose'], self.item(identifier)['purpose'])

    def test_count_exports_stay_after_regular_items_when_created_first(self):
        count = self.count('steel'); regular = self.case.create()
        for identifier in (count, regular): self.case.command('confirm_items', item_ids=[identifier])
        rows = register_rows(self.case.state['snapshot'], self.case.state['snapshot']['items'])
        self.assertEqual([row[0] for row in rows], [regular, count])
        payload = self.case.service.export_workspace(self.case.sid, 'schedule-xlsx', {
            'expected_revision': self.case.state['revision'], 'mode': 'steel'})[0]
        workbook = load_workbook(BytesIO(payload)); self.addCleanup(workbook.close)
        self.assertEqual([row[0] for row in list(workbook['Current Takeoffs'].values)[1:]], [regular, count])


class StandaloneProjectTests(unittest.TestCase):
    def test_real_save_reopen_preserves_counts_lengths_source_bytes_and_calculator_inputs(self):
        case = project_fixtures.TakeoffProjectTests(); case.setUpClass(); case.setUp()
        self.addCleanup(case.doCleanups)
        def command(op, **values):
            case.session = case.service.command(case.session['session_id'], {'op': op, 'request_id': str(uuid4()),
                'expected_revision': case.session['revision'], **values})
            return case.session
        doc = case.session['snapshot']['documents'][0]
        command('record_render', document_id=doc['id'], page=1, success=True, warnings=[])
        calibration = str(uuid4())
        command('add_calibration', calibration={'id': calibration, 'document_id': doc['id'], 'page': 1,
            'points': [[50, 50], [150, 50]], 'distance_m': 10, 'uniform_scale': True})
        command('add_standalone_count', mode='duct', document_id=doc['id'], page=1,
            markers=[{'point': [80, 80]}, {'point': [120, 100]}], fields=FIELDS, appearance={})
        command('create_item', item={'mode': 'slab', 'purpose': 'length-only', 'quantity': 1, 'fields': {'mark': 'Slab edge'},
            'geometry': {'document_id': doc['id'], 'page': 1, 'points': [[50, 200], [150, 200]]},
            'measurement': {'method': 'calibrated', 'calibration_id': calibration}})
        snapshot = deepcopy(case.session['snapshot'])
        case.request.update(takeoffs=snapshot, takeoffs_session_id=case.session['session_id'])
        case.library.save_as(case.request); case.dialogs.opened = str(case.target)
        reopened = case.library.open_file()
        self.assertEqual(reopened['takeoffs']['items'], snapshot['items']); self.assertEqual(reopened['takeoffs_issues'], [])
        self.assertEqual(case.documents.document_path(doc).read_bytes(), case.pdf)
        saved = json.loads(case.target.read_bytes())
        for key in ('estimate', 'calculators'): self.assertEqual(saved[key], json.loads(case.legacy)[key])

    def test_count_and_length_marked_pdfs_retain_all_rotations_crops_units_and_source_bytes(self):
        case = project_fixtures.TakeoffProjectTests(); case.setUpClass(); case.pdf = drawing_fixture(); case.setUp()
        self.addCleanup(case.doCleanups)
        def command(op, **values):
            case.session = case.service.command(case.session['session_id'], {'op': op, 'request_id': str(uuid4()),
                'expected_revision': case.session['revision'], **values})
            return case.session
        doc = case.session['snapshot']['documents'][0]; counts, lengths = [], []
        for page in range(1, 5):
            command('record_render', document_id=doc['id'], page=page, success=True, warnings=[])
            calibration = str(uuid4())
            command('add_calibration', calibration={'id': calibration, 'document_id': doc['id'], 'page': page,
                'points': [[20, 30], [120, 30]], 'distance_m': 10, 'uniform_scale': True})
            counts.extend(command('add_standalone_count', mode='duct', document_id=doc['id'], page=page,
                markers=[{'point': [30, 40]}, {'point': [160, 80]}], fields=FIELDS,
                appearance={'marker_shape': 'diamond'})['created_item_ids'])
            command('create_item', item={'mode': 'slab', 'purpose': 'length-only', 'quantity': 1, 'fields': {'mark': 'Edge'},
                'geometry': {'document_id': doc['id'], 'page': page, 'points': [[20, 60], [120, 60]]},
                'measurement': {'method': 'calibrated', 'calibration_id': calibration}})
            lengths.append(case.session['snapshot']['items'][-1]['id'])
        before = deepcopy(case.session['snapshot']); source = case.documents.document_path(doc).read_bytes()
        for mode, identifiers in (('duct', counts), ('slab', lengths)):
            output = case.service.export_workspace(case.session['session_id'], 'marked-pdf', {
                'expected_revision': case.session['revision'], 'mode': mode, 'document_id': doc['id'], 'item_ids': identifiers})[0]
            pages = PdfReader(BytesIO(output)).pages
            self.assertEqual(len(pages), 4)
            for index, page in enumerate(pages):
                content = page.extract_text(); self.assertIn(f'ORIGINAL DRAWING ROTATION {index*90}', content)
                self.assertIn('2 markers counted' if mode == 'duct' else '10.00 m total', content)
                self.assertIn('no calculator transfer', content)
                self.assertNotIn('Net area unavailable', content); self.assertNotIn('manual length each', content)
                self.assertFalse(page.get('/Rotate')); self.assertFalse(page.get('/UserUnit'))
        self.assertEqual(case.service.get(case.session['session_id'])['snapshot'], before)
        self.assertEqual(case.documents.document_path(doc).read_bytes(), source)


if __name__ == '__main__': unittest.main()
