"""Printed scale, viewport boundaries and cited vertical length authority."""
from copy import deepcopy
import csv
from io import BytesIO, StringIO
import json
import math
import unittest
from uuid import uuid4

from openpyxl import load_workbook

from estimator.catalog import ValidationError
from estimator.native_dialogs import SaveSelection
from estimator.takeoff_model import (SCALE_DENOMINATORS, digest, item_digest,
    item_result, preset_distance, validate_snapshot)
from estimator.takeoff_transfer import mapped_values
from tests import test_takeoff_workspace as fixtures
from tests import test_takeoff_project as project_fixtures


class TakeoffViewportTests(unittest.TestCase):
    def setUp(self):
        self.case = fixtures.TakeoffWorkspaceTests(); self.case.setUp()
        self.addCleanup(self.case.doCleanups)

    def calibration(self, region=None, denominator=100, **extra):
        x, y, width, _ = region or [10, 20, 200, 100]
        value = {'id': str(uuid4()), 'document_id': self.case.doc['id'], 'page': 1,
                 'name': 'Printed scale', 'points': [[x, y], [x+width, y]],
                 'scale_denominator': denominator, 'uniform_scale': True, **extra}
        if region:
            value['region'] = region
        self.case.command('add_calibration', calibration=value)
        return self.case.state['snapshot']['calibrations'][-1]

    def addition(self, **changes):
        return {'id': str(uuid4()), 'kind': 'riser', 'length_mm': 1500,
                'note': 'Section A states a 1500 mm rise', 'document_id': self.case.doc['id'], 'page': 1, **changes}

    def item(self, identifier):
        return next(item for item in self.case.state['snapshot']['items'] if item['id'] == identifier)

    def test_all_printed_presets_use_user_unit_once_and_ignore_rotation_crop_offsets(self):
        for denominator in SCALE_DENOMINATORS:
            calibration = self.calibration(denominator=denominator)
            self.assertAlmostEqual(calibration['distance_m'], 200*2*0.0254/72*denominator)
            snapshot = deepcopy(self.case.state['snapshot'])
            for rotation in (0, 90, 180, 270):
                snapshot['documents'][0]['pages'][0]['rotation'] = rotation
                self.assertEqual(preset_distance(calibration, snapshot), calibration['distance_m'])
            snapshot['documents'][0]['pages'][0]['user_unit'] = 1
            self.assertEqual(preset_distance(calibration, snapshot), calibration['distance_m']/2)
        identifier = self.case.create()
        self.assertEqual(item_result(self.item(identifier), self.case.state['snapshot'])['length_m'], 10)

    def test_presets_reject_unsupported_or_forged_scale_values(self):
        for value in (1, 3, 0, -50, 100.0, True, '100', math.inf, math.nan):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                self.calibration(denominator=value)
        with self.assertRaisesRegex(ValidationError, 'preset distance'):
            self.calibration(distance_m=123)
        with self.assertRaisesRegex(ValidationError, 'preset distance'):
            self.calibration(points=[[10, 20], [10.00000000000001, 20]], distance_m=1e-13)
        with self.assertRaises(ValidationError):
            self.case.command('add_calibration', calibration={'scale_denominator': 100})

    def test_preset_commands_store_exact_scale_without_rewriting_loaded_values(self):
        expected = 200*2*(0.0254/72)*100
        calibration = self.calibration(distance_m=expected*(1+5e-13))
        self.assertEqual(calibration['distance_m'], expected)
        self.case.command('update_calibration', calibration_id=calibration['id'],
                          changes={'scale_denominator': 50, 'distance_m': (expected/2)*(1+5e-13)})
        self.assertEqual(self.case.state['snapshot']['calibrations'][-1]['distance_m'], expected/2)
        loaded = deepcopy(self.case.state['snapshot'])
        loaded['calibrations'][-1]['distance_m'] *= 1+5e-13
        before = digest(loaded)
        self.assertEqual(digest(validate_snapshot(loaded)), before)
        self.assertEqual(digest(loaded), before)

    def test_viewports_are_independent_and_full_geometry_must_stay_in_one(self):
        left = self.calibration([10, 20, 90, 100], 100)
        right = self.calibration([110, 20, 90, 100], 50)
        lengths = []
        for calibration, x in ((left, 20), (right, 120)):
            identifier = self.case.create(measurement={'method': 'calibrated', 'calibration_id': calibration['id']},
                geometry={'document_id': self.case.doc['id'], 'page': 1, 'points': [[x, 30], [x+50, 30]]})
            lengths.append(item_result(self.item(identifier), self.case.state['snapshot'])['length_m'])
        self.assertEqual(lengths[0], lengths[1]*2)
        for calibration_id in (left['id'], self.case.calibration_id):
            with self.assertRaisesRegex(ValidationError, 'viewport'):
                self.case.create(measurement={'method': 'calibrated', 'calibration_id': calibration_id})
        with self.assertRaisesRegex(ValidationError, 'overlap'):
            self.calibration([80, 25, 40, 40])
        with self.assertRaisesRegex(ValidationError, 'baseline'):
            self.calibration([20, 25, 20, 20], points=[[10, 25], [40, 25]])

    def test_new_viewport_invalidates_only_affected_items_and_cannot_be_bypassed(self):
        affected = self.case.create()
        outside = self.case.create(geometry={'document_id': self.case.doc['id'], 'page': 1, 'points': [[150, 30], [200, 30]]})
        self.case.command('confirm_items', item_ids=[affected, outside])
        unaffected_digest = self.item(outside)['confirmation']['digest']
        calibration = self.calibration([10, 20, 110, 100], 100)
        self.assertEqual(self.item(affected)['state'], 'draft')
        self.assertEqual(self.item(outside)['confirmation']['digest'], unaffected_digest)
        with self.assertRaisesRegex(ValidationError, 'viewport'):
            self.case.command('confirm_items', item_ids=[affected])
        self.case.command('update_item', item_id=affected, changes={'measurement': {'method': 'calibrated', 'calibration_id': calibration['id']}})
        self.case.command('confirm_items', item_ids=[affected])
        self.assertEqual(self.item(affected)['state'], 'confirmed')

    def test_viewport_intersection_uses_route_not_merely_its_bounding_box(self):
        self.calibration([50, 50, 30, 30])
        identifier = self.case.create(geometry={'document_id': self.case.doc['id'], 'page': 1,
                                               'points': [[20, 90], [100, 90], [100, 30]]})
        self.assertEqual(item_result(self.item(identifier), self.case.state['snapshot'])['length_m'], 14)
        with self.assertRaisesRegex(ValidationError, 'viewport'):
            self.case.create(geometry={'document_id': self.case.doc['id'], 'page': 1, 'points': [[20, 60], [100, 60]]})

    def test_revision_is_immutable_invalidates_confirmation_and_supports_manual_transition(self):
        calibration = self.calibration([10, 20, 110, 100])
        identifier = self.case.create(measurement={'method': 'calibrated', 'calibration_id': calibration['id']})
        self.case.command('confirm_items', item_ids=[identifier])
        self.case.command('update_calibration', calibration_id=calibration['id'], changes={'scale_denominator': 50})
        revised = self.case.state['snapshot']['calibrations'][-1]
        self.assertEqual(revised['supersedes_id'], calibration['id'])
        self.assertEqual(revised['region'], calibration['region'])
        self.assertEqual(revised['distance_m'], calibration['distance_m']/2)
        self.assertEqual(self.case.state['snapshot']['calibrations'][-2], calibration)
        self.assertEqual(self.item(identifier)['state'], 'draft')
        with self.assertRaisesRegex(ValidationError, 'superseded'):
            self.case.command('update_calibration', calibration_id=calibration['id'], changes={'name': 'Old'})
        with self.assertRaisesRegex(ValidationError, 'superseded'):
            self.case.create(measurement={'method': 'calibrated', 'calibration_id': calibration['id']})
        self.case.command('update_calibration', calibration_id=revised['id'], changes={'scale_denominator': None, 'distance_m': 11})
        manual = self.case.state['snapshot']['calibrations'][-1]
        self.assertNotIn('scale_denominator', manual)
        self.assertEqual(item_result(self.item(identifier), self.case.state['snapshot'])['length_m'], 10)
        with self.assertRaisesRegex(ValidationError, 'viewport'):
            self.case.command('update_calibration', calibration_id=manual['id'], changes={'region': [10, 20, 50, 100], 'points': [[10, 20], [60, 20]]})

    def test_riser_drop_per_member_lengths_exports_and_exact_transfer_mapping(self):
        additions = [self.addition(), self.addition(kind='drop', length_mm=250, note='Detail B drop 250 mm')]
        identifier = self.case.create(length_additions=additions, quantity=3)
        result = item_result(self.item(identifier), self.case.state['snapshot'])
        self.assertEqual((result['base_length_m'], result['additions_length_m'], result['length_m'], result['total_length_m']), (10, 1.75, 11.75, 35.25))
        self.case.command('confirm_items', item_ids=[identifier])
        preview = self.case.preview(identifier)
        self.assertEqual(preview['inputs']['SCHEDULE']['I10'], 3)
        self.assertEqual(preview['inputs']['SCHEDULE']['J10'], 11.75)
        board_item = deepcopy(self.item(identifier))
        board_item['fields'].update(product='TRAFALGAR COREX', critical_temperature=620)
        board, _ = mapped_values(board_item, self.case.state['snapshot'], 'steel_board', 9)
        self.assertEqual(board['F9'], 35.25)
        duct = self.case.create('duct', quantity=1, length_additions=additions)
        self.case.command('confirm_items', item_ids=[duct])
        self.assertEqual(self.case.preview(duct, 'ductwork')['inputs']['CALCULATOR']['D11'], 11.75)
        payload, _, _ = self.case.service.export(self.case.sid, 'csv', [identifier])
        row = next(csv.DictReader(StringIO(payload.decode('utf-8-sig'))))
        self.assertEqual(float(row['Base length per item m']), 10)
        self.assertEqual(float(row['Riser/drop additions per item m']), 1.75)
        evidence = json.loads(row['Riser/drop source dimensions'])
        self.assertEqual(evidence[0]['document_sha256'], self.case.doc['sha256'])
        payload, _, _ = self.case.service.export(self.case.sid, 'xlsx', [identifier])
        workbook = load_workbook(BytesIO(payload)); self.addCleanup(workbook.close)
        sheet = workbook['Confirmed Takeoffs']; headers = {cell.value: cell.column for cell in sheet[1]}
        self.assertEqual(sheet.cell(2, headers['Total length m']).value, 35.25)
        self.case.command('update_item', item_id=duct, changes={'fields': {'shape': 'circular', 'diameter_mm': 600}})
        self.case.command('confirm_items', item_ids=[duct])
        with self.assertRaisesRegex(ValidationError, 'Circular'):
            self.case.preview(duct, 'ductwork')

    def test_additions_require_explicit_positive_finite_dimensions_and_retained_source(self):
        for value in (0, -1, True, '1000', math.nan, math.inf):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                self.case.create(length_additions=[self.addition(length_mm=value)])
        for changes in ({'note': None}, {'kind': 'guess'}, {'page': 2}, {'document_id': str(uuid4())}):
            with self.subTest(changes=changes), self.assertRaises(ValidationError):
                self.case.create(length_additions=[self.addition(**changes)])
        addition = self.addition()
        with self.assertRaises(ValidationError):
            self.case.create(length_additions=[addition, addition])
        with self.assertRaises(ValidationError):
            self.case.create(length_additions=[self.addition() for _ in range(101)])

    def test_addition_edits_evidence_and_render_changes_invalidate_approval(self):
        other = fixtures.document(); other['sha256'] = 'b'*64
        self.case.state = self.case.service.add_document(self.case.sid, other, self.case.state['revision'])
        identifier = self.case.create(length_additions=[self.addition(document_id=other['id'])])
        with self.assertRaisesRegex(ValidationError, 'Render'):
            self.case.command('confirm_items', item_ids=[identifier])
        self.case.command('record_render', document_id=other['id'], page=1, success=True, warnings=[])
        self.case.command('confirm_items', item_ids=[identifier])
        original = deepcopy(self.case.state['snapshot'])
        tampered = deepcopy(original); tampered['documents'][1]['sha256'] = 'c'*64
        self.assertNotEqual(item_digest(tampered['items'][0], tampered), original['items'][0]['confirmation']['digest'])
        with self.assertRaisesRegex(ValidationError, 'linked'):
            self.case.command('delete_document', document_id=other['id'])
        self.case.command('record_render', document_id=other['id'], page=1, success=False, warnings=[])
        self.assertEqual(self.item(identifier)['state'], 'draft')
        self.case.command('update_item', item_id=identifier, changes={'length_additions': [self.addition(length_mm=2000)]})
        self.assertEqual(self.item(identifier)['state'], 'draft')
        self.assertEqual(item_result(self.item(identifier), self.case.state['snapshot'])['length_m'], 12)

    def test_duct_split_cannot_drop_additions_steel_groups_preserve_them(self):
        identifier = self.case.create('duct', quantity=1, length_additions=[self.addition()])
        original = self.item(identifier)
        parts = [{key: deepcopy(original[key]) for key in ('mode', 'quantity', 'fields', 'measurement', 'geometry', 'evidence')} for _ in range(2)]
        parts[0]['geometry']['points'] = [[10, 30], [60, 30]]
        parts[1]['geometry']['points'] = [[60, 30], [110, 30]]
        before = deepcopy(self.case.state['snapshot'])
        with self.assertRaisesRegex(ValidationError, 'riser/drop'):
            self.case.command('split_item', item_id=identifier, parts=parts)
        self.assertEqual(self.case.service.get(self.case.sid)['snapshot'], before)
        adjacent = self.case.create('duct', quantity=1, geometry={'document_id': self.case.doc['id'], 'page': 1,
                                                                'points': [[110, 30], [210, 30]]})
        merged_proposal = deepcopy(parts[0]); merged_proposal['geometry']['points'] = [[10, 30], [210, 30]]
        before = deepcopy(self.case.state['snapshot'])
        with self.assertRaisesRegex(ValidationError, 'riser/drop'):
            self.case.command('merge_items', item_ids=[identifier, adjacent], item=merged_proposal)
        self.assertEqual(self.case.service.get(self.case.sid)['snapshot'], before)
        steel = self.case.create(quantity=3, length_additions=[self.addition()])
        members = set(self.item(steel)['member_ids'])
        self.case.command('split_steel_group', item_id=steel, quantities=[1, 2])
        groups = [item for item in self.case.state['snapshot']['items'] if item['mode'] == 'steel']
        self.assertEqual(sum(item_result(item, self.case.state['snapshot'])['total_length_m'] for item in groups), 34.5)
        self.case.command('merge_steel_groups', item_ids=[item['id'] for item in groups])
        merged = self.case.state['snapshot']['items'][-1]
        self.assertEqual(set(merged['member_ids']), members)
        self.assertEqual(merged['length_additions'], groups[0]['length_additions'])

    def test_legacy_item_digest_and_snapshot_shape_are_unchanged(self):
        identifier = self.case.create()
        item, snapshot = self.item(identifier), self.case.state['snapshot']
        expected = {key: value for key, value in item.items() if key not in ('state', 'review', 'confirmation')}
        expected['source_documents'] = [{key: self.case.doc[key] for key in ('id', 'sha256', 'size', 'pages')}]
        expected['calibration'] = snapshot['calibrations'][0]
        self.assertEqual(item_digest(item, snapshot), digest(expected))
        self.assertNotIn('length_additions', item)
        self.assertEqual(validate_snapshot(snapshot), snapshot)

    def test_area_containing_or_crossing_viewport_needs_explicit_uniform_scope(self):
        from tests.test_takeoff_area import TakeoffAreaTests
        area_case = TakeoffAreaTests(); area_case.case = self.case
        proposal = area_case.proposal()
        calibration = self.calibration([10, 20, 130, 90], 50)
        with self.assertRaisesRegex(ValidationError, 'viewport'):
            self.case.command('create_item', item=proposal)
        proposal['measurement']['calibration_id'] = calibration['id']
        self.case.command('create_item', item=proposal)
        result = self.case.state['item_results'][-1]
        expected_scale = 2*0.0254/72*50
        self.assertAlmostEqual(result['net_area_m2'], 4800*expected_scale**2)
        with self.assertRaisesRegex(ValidationError, 'linear'):
            self.case.command('update_item', item_id=self.case.state['snapshot']['items'][-1]['id'],
                              changes={'length_additions': [self.addition()]})

    def test_direct_confirm_batch_is_atomic_and_retry_does_not_duplicate_receipts(self):
        valid = self.case.create(); invalid = self.case.create(quantity=None)
        with self.assertRaisesRegex(ValidationError, 'quantity'):
            self.case.command('confirm_items', item_ids=[valid, invalid])
        with self.case.store.connect() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM takeoff_approvals').fetchone()[0], 0)
        request = {'op': 'confirm_items', 'expected_revision': self.case.state['revision'],
                   'request_id': str(uuid4()), 'item_ids': [valid]}
        first = self.case.service.command(self.case.sid, request)
        self.assertEqual(self.case.service.command(self.case.sid, request), first)
        with self.case.store.connect() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM takeoff_approvals').fetchone()[0], 2)
        self.case.documents.blocked = True
        with self.assertRaisesRegex(ValidationError, 'Changed'):
            self.case.service.command(self.case.sid, request)


class TakeoffViewportProjectTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        project_fixtures.TakeoffProjectTests.setUpClass()

    def test_real_save_as_reopen_preserves_scales_additions_hashes_and_legacy_values(self):
        case = project_fixtures.TakeoffProjectTests(); case.setUp(); self.addCleanup(case.doCleanups)
        def command(op, **values):
            case.session = case.service.command(case.session['session_id'], {'op': op, 'request_id': str(uuid4()),
                                      'expected_revision': case.session['revision'], **values})
        document = case.session['snapshot']['documents'][0]
        command('record_render', document_id=document['id'], page=1, success=True, warnings=[])
        calibration, item_id = str(uuid4()), str(uuid4())
        command('add_calibration', calibration={'id': calibration, 'document_id': document['id'], 'page': 1,
                'name': 'Detail viewport', 'points': [[50, 50], [150, 50]], 'region': [50, 50, 200, 200],
                'scale_denominator': 100, 'uniform_scale': True})
        command('create_item', item={'id': item_id, 'mode': 'duct', 'quantity': 1,
                'geometry': {'document_id': document['id'], 'page': 1, 'points': [[50, 60], [150, 60]]},
                'measurement': {'method': 'calibrated', 'calibration_id': calibration},
                'fields': {'shape': 'rectangular', 'width_mm': 600, 'height_mm': 400, 'frl': '120/120/120', 'orientation': 'Horizontal'},
                'evidence': [], 'length_additions': [{'id': str(uuid4()), 'kind': 'riser', 'length_mm': 2100,
                    'note': 'Section riser dimension 2100 mm', 'document_id': document['id'], 'page': 1}]})
        command('confirm_items', item_ids=[item_id])
        snapshot = deepcopy(case.session['snapshot'])
        case.library.save_as({**deepcopy(case.base), 'takeoffs': snapshot, 'takeoffs_session_id': case.session['session_id']})
        saved = json.loads(case.target.read_bytes())
        for key in ('estimate', 'calculators'):
            self.assertEqual(saved[key], json.loads(case.legacy)[key])
        case.dialogs.opened = str(case.target); opened = case.library.open_file()
        self.assertEqual(opened['takeoffs_issues'], [])
        self.assertEqual(opened['takeoffs']['items'], snapshot['items'])
        self.assertEqual(opened['takeoffs']['calibrations'], snapshot['calibrations'])
        self.assertEqual(opened['takeoffs']['audit_head'], snapshot['audit_head'])
        second = case.root / 'copied' / 'copy.json'; second.parent.mkdir()
        case.dialogs.selection = SaveSelection(str(second), None)
        case.library.save_as({**deepcopy(case.base), 'takeoffs': opened['takeoffs'], 'takeoffs_session_id': opened['takeoffs_session_id']})
        case.dialogs.opened = str(second); final = case.library.open_file()
        self.assertEqual(final['takeoffs']['items'], snapshot['items'])
        self.assertEqual(final['takeoffs']['audit_head'], snapshot['audit_head'])
        source = next((second.parent / final['takeoffs']['companion_folder']).rglob('*.pdf'))
        source.write_bytes(source.read_bytes()+b'\nmodified evidence\n')
        broken = case.library.open_file()
        self.assertTrue(broken['takeoffs_issues'])
        self.assertEqual(broken['takeoffs']['items'][0]['state'], 'draft')


if __name__ == '__main__':
    unittest.main()
