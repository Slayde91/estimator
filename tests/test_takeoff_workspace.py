"""Takeoff review gates, measurement integrity, audit and guarded calculator drafts."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from uuid import uuid4

from openpyxl import load_workbook

from estimator.catalog import ValidationError
from estimator.schedule_rows import empty_schedule_inputs
from estimator.storage import Store
from estimator.takeoff_model import digest, item_digest, item_result, validate_snapshot
from estimator.takeoff_workspace import TakeoffService


class Documents:
    def __init__(self):
        self.blobs = {}; self.blocked = False; self.owners = []

    def put_blob(self, value, kind='audit'):
        key = digest(value); self.blobs[key] = deepcopy(value); return key

    def get_blob(self, key, kind='audit'):
        if key not in self.blobs:
            raise ValidationError('Missing retained audit evidence.')
        return deepcopy(self.blobs[key])

    def validate_project_documents(self, documents, owner=None):
        self.owners.append(owner)
        return [{'document_id': d['id'], 'code': 'CHANGED', 'message': 'Changed source'} for d in documents] if self.blocked else []

    def assert_documents(self, documents, owner=None):
        if self.validate_project_documents(documents, owner):
            raise ValidationError('Changed source evidence.')

    def restore(self, snapshot, path, owner=None):
        return self.validate_project_documents(snapshot['documents'], owner)

    def bind_source(self, owner, snapshot, path):
        self.owners.append(owner)

    def close_owner(self, owner):
        self.owners.append(owner)

    def validate_audit(self, snapshot, owner=None):
        self.owners.append(owner)

    def assert_add_capacity(self, snapshot, document):
        pass


def document():
    return {'id': str(uuid4()), 'name': 'Drawing.pdf', 'sha256': 'a'*64, 'size': 100,
            'pages': [{'page': 1, 'width': 200, 'height': 100, 'view': [10, 20, 210, 120],
                       'media_box': [0, 0, 220, 140], 'crop_box': [10, 20, 210, 120], 'rotation': 90, 'user_unit': 2}]}


class TakeoffWorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.store = Store(Path(self.directory.name) / 'takeoffs.sqlite3')
        self.documents = Documents()
        self.service = TakeoffService(self.store, self.documents)
        self.state = self.service.open(); self.sid = self.state['session_id']
        self.doc = document()
        self.state = self.service.add_document(self.sid, self.doc, self.state['revision'])
        self.command('record_render', document_id=self.doc['id'], page=1, success=True, warnings=[])
        self.calibration_id = str(uuid4())
        self.command('add_calibration', calibration={'id': self.calibration_id, 'document_id': self.doc['id'], 'page': 1,
            'name': 'Plan scale', 'points': [[10, 20], [110, 20]], 'distance_m': 10, 'uniform_scale': True})

    def command(self, op, **values):
        request = {'op': op, 'request_id': str(uuid4()), 'expected_revision': self.state['revision'], **values}
        self.state = self.service.command(self.sid, request)
        return self.state

    def create(self, mode='steel', **changes):
        fields = {'mark': 'B17', 'section': '100UC15', 'member_type': 'Beam', 'exposure': 'Re-entrant - 3 sides',
                  'fire_period_min': 120, 'critical_temperature': 550, 'product': 'CAFCO 300', 'sides': 3}
        if mode == 'duct':
            fields = {'mark': 'D1', 'shape': 'rectangular', 'width_mm': 600, 'height_mm': 400,
                      'frl': '120/120/120', 'orientation': 'Horizontal', 'exposure': 'Internal',
                      'product': 'FyreWrap', 'wall_penetrations': 0, 'floor_penetrations': 0}
        item = {'mode': mode, 'geometry': {'document_id': self.doc['id'], 'page': 1, 'points': [[10, 30], [110, 30]]},
                'measurement': {'method': 'calibrated', 'calibration_id': self.calibration_id},
                'quantity': 2, 'fields': fields, 'evidence': [{'document_id': self.doc['id'], 'page': 1, 'note': 'Plan member'}]}
        item.update(changes)
        self.command('create_item', item=item)
        return self.state['snapshot']['items'][-1]['id']

    def confirm(self, identity):
        self.command('review_items', item_ids=[identity]); self.command('confirm_items', item_ids=[identity])

    def preview(self, item_id, calculator_id='steel_vermiculite', inputs=None, rows=None, update=False):
        first = {'steel_vermiculite': 10, 'steel_board': 9, 'ductwork': 11}[calculator_id]
        return self.service.preview_transfer(self.sid, {'expected_revision': self.state['revision'],
            'calculator_id': calculator_id, 'inputs': inputs if inputs is not None else empty_schedule_inputs(calculator_id),
            'schedule_rows': rows or [first], 'item_ids': [item_id], 'update_linked': update})

    def apply(self, preview, inputs=None, rows=None):
        calculator_id = preview['calculator_id']; first = {'steel_vermiculite': 10, 'steel_board': 9, 'ductwork': 11}[calculator_id]
        request = {'expected_revision': self.state['revision'], 'request_id': str(uuid4()), 'preview_id': preview['preview_id'],
                   'inputs': inputs if inputs is not None else empty_schedule_inputs(calculator_id), 'schedule_rows': rows or [first]}
        self.state = self.service.apply_transfer(self.sid, request)
        return request

    def test_calibration_uses_unrotated_pdf_units_exactly_once(self):
        identity = self.create()
        result = self.state['item_results'][0]
        self.assertEqual(result['length_m'], 10)
        self.assertEqual(result['total_length_m'], 20)
        self.assertEqual(result['issues'], [])
        item = self.state['snapshot']['items'][0]
        self.assertEqual(item['id'], identity)
        self.assertEqual(item['geometry']['points'], [[10, 30], [110, 30]])
        with self.assertRaisesRegex(ValidationError, 'outside'):
            self.command('update_item', item_id=identity, changes={'geometry': {'document_id': self.doc['id'], 'page': 1, 'points': [[0, 0]]}})

    def test_direct_confirmation_records_local_hash_bound_review_and_confirmation(self):
        identity = self.create()
        self.command('confirm_items', item_ids=[identity])
        self.assertIsNotNone(self.state['snapshot']['items'][0]['review'])
        original = deepcopy(self.state['snapshot'])
        self.assertEqual(self.service.open(original, source_path='verified-project.json')['snapshot']['items'][0]['state'], 'confirmed')
        self.assertEqual(self.service.open(original)['snapshot']['items'][0]['state'], 'draft')
        forged = deepcopy(original); forged['items'][0]['fields']['section'] = 'Invented'
        self.assertEqual(self.service.open(forged, source_path='verified-project.json')['snapshot']['items'][0]['state'], 'draft')
        self.command('update_item', item_id=identity, changes={'quantity': 3})
        edited = self.state['snapshot']['items'][0]
        self.assertEqual(edited['state'], 'draft'); self.assertIsNone(edited['confirmation'])
        self.assertNotEqual(item_digest(edited, self.state['snapshot']), original['items'][0]['confirmation']['digest'])
        foreign = TakeoffService(Store(Path(self.directory.name)/'foreign.sqlite3'), self.documents)
        self.assertEqual(foreign.open(original)['snapshot']['items'][0]['state'], 'draft')

    def test_explicit_quantity_render_review_and_cited_dimensions_fail_closed(self):
        identity = self.create(quantity=None)
        self.command('review_items', item_ids=[identity])
        with self.assertRaisesRegex(ValidationError, 'quantity'):
            self.command('confirm_items', item_ids=[identity])
        self.command('update_item', item_id=identity, changes={'quantity': 1, 'measurement': {'method': 'cited', 'length_m': 7.82, 'citation': 'Section A dimension'}})
        self.confirm(identity)
        self.assertEqual(self.state['item_results'][0]['length_m'], 7.82)
        self.command('record_render', document_id=self.doc['id'], page=1, success=True, warnings=['Unsupported image'])
        self.assertEqual(self.state['snapshot']['items'][0]['state'], 'draft')
        with self.assertRaisesRegex(ValidationError, 'Render'):
            self.command('review_items', item_ids=[identity])
        for quantity in (0, -1, 1.5, True):
            with self.subTest(quantity=quantity), self.assertRaises(ValidationError):
                self.command('update_item', item_id=identity, changes={'quantity': quantity})

    def test_confirmation_receipt_binds_actor_and_successful_deterministic_checks(self):
        import json
        identity = self.create(quantity=None)
        self.command('review_items', item_ids=[identity])
        review = self.state['snapshot']['items'][0]['review']
        self.assertEqual(review['actor'], {'kind': 'local-session', 'session_id': self.sid})
        self.assertTrue(any(issue['code'] == 'MISSING_QUANTITY' for issue in review['checks']['issues']))
        self.command('update_item', item_id=identity, changes={'quantity': 3})
        self.confirm(identity)
        receipt = self.state['snapshot']['items'][0]['confirmation']
        self.assertEqual(receipt['actor'], {'kind': 'local-session', 'session_id': self.sid})
        self.assertEqual(receipt['checks'], {'engine': 'takeoffs-v1', 'quantity': 3, 'length_m': 10,
                         'total_length_m': 30, 'evidence_verified': True, 'issues': []})
        payload, _, _ = self.service.export(self.sid, 'xlsx', [identity])
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook['Confirmed Takeoffs']; headers = {c.value: c.column for c in sheet[1]}
        self.assertEqual(json.loads(sheet.cell(2, headers['Confirmation checks']).value), receipt['checks'])
        self.assertEqual(json.loads(sheet.cell(2, headers['Confirmed by']).value), receipt['actor'])
        workbook.close()
        snapshot = deepcopy(self.state['snapshot'])
        snapshot['items'][0]['confirmation']['checks']['length_m'] = 9
        reopened = self.service.open(snapshot, source_path='verified-project.json')
        self.assertEqual(reopened['snapshot']['items'][0]['state'], 'reviewed')
        invalid = deepcopy(self.state['snapshot'])
        invalid['items'][0]['confirmation']['checks']['issues'] = review['checks']['issues']
        with self.assertRaisesRegex(ValidationError, 'Confirmation requires'):
            validate_snapshot(invalid)

    def test_idempotent_commands_and_atomic_bulk_revision_guards(self):
        first = self.create(); second = self.create()
        revision = self.state['revision']
        request = {'request_id': str(uuid4()), 'expected_revision': revision, 'op': 'bulk_update', 'item_ids': [first, second], 'changes': {'fields': {'level': 'Level 2'}}}
        once = self.service.command(self.sid, request)
        self.assertEqual(once, self.service.command(self.sid, request))
        self.assertEqual(once['revision'], revision+1)
        with self.assertRaisesRegex(ValidationError, 'different operation'):
            self.service.command(self.sid, {**request, 'changes': {'quantity': 3}})
        with self.assertRaisesRegex(ValidationError, 'changed'):
            self.service.command(self.sid, {**request, 'request_id': str(uuid4())})
        self.state = once
        before = deepcopy(self.state)
        with self.assertRaises(ValidationError):
            self.command('bulk_update', item_ids=[first, second], changes={'quantity': 0})
        self.assertEqual(self.service.get(self.sid)['snapshot'], before['snapshot'])

    def test_idempotent_command_retry_returns_current_state_and_keeps_compact_receipts(self):
        import json
        identity = self.create()
        request = {'request_id': str(uuid4()), 'expected_revision': self.state['revision'],
                   'op': 'update_item', 'item_id': identity, 'changes': {'fields': {'level': 'L2'}}}
        self.state = self.service.command(self.sid, request)
        self.command('update_item', item_id=identity, changes={'fields': {'level': 'L3'}})
        current = deepcopy(self.state)
        self.assertEqual(self.service.command(self.sid, request), current)
        for _ in range(100):
            self.command('record_render', document_id=self.doc['id'], page=1, success=True, warnings=[])
        receipts = self.service._sessions[self.sid]['requests']
        self.assertEqual(len(receipts), 100)
        self.assertLess(len(json.dumps(receipts).encode()), 20000)
        self.assertEqual(self.state['snapshot'], current['snapshot'])

    def test_transfer_payload_cache_is_bounded_across_sessions_without_losing_receipt_hashes(self):
        from unittest.mock import patch
        with patch('estimator.takeoff_workspace.SESSION_CACHE_BYTES', 1024), patch('estimator.takeoff_workspace.PROCESS_CACHE_BYTES', 2048):
            request_id = str(uuid4())
            session = self.service._sessions[self.sid]
            session['requests'][request_id] = {'request_hash': 'a'*64, 'applied_transfer': True, 'applied_revision': self.state['revision']}
            self.service._cache_payload(self.sid, 'requests', request_id, {'text': 'a'*700})
            self.service._cache_payload(self.sid, 'previews', str(uuid4()), {'text': 'b'*700})
            self.assertEqual(session['requests'][request_id]['request_hash'], 'a'*64)
            self.assertNotIn('payload', session['requests'][request_id])
            with self.assertRaisesRegex(ValidationError, 'already applied'):
                self.service._request_response(self.sid, session['requests'][request_id])
            for _ in range(3):
                sid = self.service.open()['session_id']
                self.service._cache_payload(sid, 'previews', str(uuid4()), {'text': 'c'*700})
            retained = [value for current in self.service._sessions.values()
                        for bucket in ('previews', 'requests') for value in current[bucket].values() if 'payload' in value]
            self.assertLessEqual(sum(value['cache_bytes'] for value in retained), 2048)
            with self.assertRaisesRegex(ValidationError, 'memory limit'):
                self.service._cache_payload(self.sid, 'previews', str(uuid4()), {'text': 'd'*1024})

    def test_delete_undo_preserves_identity_and_requires_new_review(self):
        identity = self.create(); self.confirm(identity)
        with self.assertRaisesRegex(ValidationError, 'linked items'):
            self.command('delete_document', document_id=self.doc['id'])
        self.command('delete_items', item_ids=[identity]); self.assertFalse(self.state['snapshot']['items'])
        self.command('undo')
        self.assertEqual(self.state['snapshot']['items'][0]['id'], identity)
        self.assertEqual(self.state['snapshot']['items'][0]['state'], 'draft')
        history = self.service.history(self.sid, limit=2)
        self.assertEqual([e['op'] for e in history['items']], ['undo', 'delete_items'])
        self.assertTrue(history['has_more'])

    def test_spray_transfer_preserves_per_member_length_exact_settings_and_repeat_identity(self):
        identity = self.create(); self.confirm(identity)
        original = empty_schedule_inputs('steel_vermiculite')
        original['SCHEDULE']['A10'] = 'Existing manual row'
        original['SCHEDULE']['I10'] = 0
        original['SETTINGS']['D37'] = 0.123456789012345
        preview = self.preview(identity, inputs=original)
        self.assertEqual(preview['bindings'][0]['row'], 11)
        self.assertEqual(preview['inputs']['SCHEDULE']['I11'], 2)
        self.assertEqual(preview['inputs']['SCHEDULE']['J11'], 10)
        self.assertEqual(preview['inputs']['SCHEDULE']['A10'], 'Existing manual row')
        self.assertEqual(preview['inputs']['SETTINGS'], original['SETTINGS'])
        apply_request = self.apply(preview, inputs=original)
        self.assertEqual(self.service.apply_transfer(self.sid, apply_request), self.state)
        # JavaScript JSON.stringify serializes integral floats as integers.
        browser_inputs = deepcopy(preview['inputs']); browser_inputs['SCHEDULE']['J11'] = 10
        repeated = self.preview(identity, inputs=browser_inputs, rows=preview['schedule_rows'])
        self.assertEqual(repeated['changes'][0]['action'], 'unchanged')
        self.assertEqual(repeated['bindings'][0]['id'], preview['bindings'][0]['id'])
        with self.store.connect() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM calculator_states').fetchone()[0], 0)

    def test_board_and_duct_transfer_total_length_once_and_circular_remains_register_only(self):
        identity = self.create()
        self.command('update_item', item_id=identity, changes={'fields': {'product': 'TRAFALGAR COREX', 'critical_temperature': 620}})
        self.confirm(identity)
        board = self.preview(identity, 'steel_board')
        self.assertEqual(board['inputs']['CALCULATOR']['F9'], 20)
        duct_id = self.create('duct', quantity=1); self.confirm(duct_id)
        duct = self.preview(duct_id, 'ductwork')
        self.assertEqual(duct['inputs']['CALCULATOR']['D11'], 10)
        self.assertEqual(duct['inputs']['CALCULATOR']['B11'], '600x400')
        self.command('update_item', item_id=duct_id, changes={'fields': {'shape': 'circular', 'diameter_mm': 300}})
        self.confirm(duct_id)
        with self.assertRaisesRegex(ValidationError, 'Circular'):
            self.preview(duct_id, 'ductwork')
        data, _, _ = self.service.export(self.sid, 'csv', [duct_id])
        self.assertIn(b'circular', data)

    def test_explicit_linked_refresh_and_conflicting_manual_edit(self):
        identity = self.create(); self.confirm(identity)
        preview = self.preview(identity); self.apply(preview)
        self.command('update_item', item_id=identity, changes={'quantity': 3}); self.confirm(identity)
        with self.assertRaisesRegex(ValidationError, 'explicit update'):
            self.preview(identity, inputs=preview['inputs'], rows=preview['schedule_rows'])
        updated = self.preview(identity, inputs=preview['inputs'], rows=preview['schedule_rows'], update=True)
        self.assertEqual(updated['changes'][0]['action'], 'update')
        self.assertEqual(updated['inputs']['SCHEDULE']['I10'], 3)
        edited = deepcopy(preview['inputs']); edited['SCHEDULE']['K10'] = 5
        with self.assertRaisesRegex(ValidationError, 'linked calculator row'):
            self.preview(identity, inputs=edited, rows=preview['schedule_rows'], update=True)
        with self.assertRaisesRegex(ValidationError, 'changed during'):
            self.service.apply_transfer(self.sid, {'request_id': str(uuid4()), 'expected_revision': self.state['revision'], 'preview_id': updated['preview_id'], 'inputs': edited, 'schedule_rows': preview['schedule_rows']})

    def test_capacity_failure_never_changes_takeoff_or_calculator_and_unsupported_profiles_block(self):
        identity = self.create(); self.confirm(identity)
        inputs = empty_schedule_inputs('steel_vermiculite')
        inputs['SCHEDULE'].update({f'A{row}': f'Manual {row}' for row in range(10, 1010)})
        before = deepcopy(self.state['snapshot']); original = deepcopy(inputs)
        with self.assertRaisesRegex(ValidationError, 'enough empty rows'):
            self.preview(identity, inputs=inputs, rows=list(range(10, 1010)))
        self.assertEqual(inputs, original)
        self.assertEqual(self.service.get(self.sid)['snapshot'], before)
        self.command('update_item', item_id=identity, changes={'fields': {'section': '460UB fabricated'}})
        self.confirm(identity)
        with self.assertRaisesRegex(ValidationError, 'exact supported'):
            self.preview(identity)

    def test_duct_normalization_is_visible_dimensions_exact_and_grouped_runs_blocked(self):
        identity = self.create('duct', quantity=1)
        self.command('update_item', item_id=identity, changes={'fields': {'width_mm': 600.1234567890123, 'frl': '60/60/60'}})
        self.confirm(identity)
        preview = self.preview(identity, 'ductwork')
        self.assertEqual(preview['inputs']['CALCULATOR']['B11'], '600.1234567890123x400')
        self.assertEqual(preview['inputs']['CALCULATOR']['E11'], '120/120/120')
        self.assertEqual(preview['normalizations'][0]['before'], '60/60/60')
        self.command('update_item', item_id=identity, changes={'quantity': 2}); self.confirm(identity)
        with self.assertRaisesRegex(ValidationError, 'each physical duct run'):
            self.preview(identity, 'ductwork')

    def test_explicit_detach_preserves_calculator_and_identical_render_is_read_only(self):
        identity = self.create(); self.confirm(identity)
        preview = self.preview(identity); self.apply(preview)
        calculated = deepcopy(self.state['calculator'])
        revision, head = self.state['revision'], self.state['snapshot']['audit_head']
        self.command('record_render', document_id=self.doc['id'], page=1, success=True, warnings=[])
        self.assertEqual(self.state['revision'], revision)
        self.assertEqual(self.state['snapshot']['audit_head'], head)
        self.command('detach_transfers', item_ids=[identity], calculator_id='steel_vermiculite')
        self.assertEqual(self.state['snapshot']['transfers'], [])
        self.assertEqual(calculated['inputs'], preview['inputs'])
        new = self.preview(identity, inputs=calculated['inputs'], rows=calculated['schedule_rows'])
        self.assertEqual(new['bindings'][0]['row'], 11)

    def test_member_identities_survive_quantity_edits_and_cannot_be_double_counted(self):
        identity = self.create()
        original = self.state['snapshot']['items'][0]['member_ids']
        self.assertEqual(len(original), 2); self.assertNotEqual(original[0], original[1])
        self.command('update_item', item_id=identity, changes={'quantity': 3})
        updated = self.state['snapshot']['items'][0]['member_ids']
        self.assertEqual(updated[:2], original)
        self.assertEqual(len(updated), 3)
        with self.assertRaisesRegex(ValidationError, 'multiple takeoff items'):
            self.create(member_ids=original)
        self.command('update_item', item_id=identity, changes={'quantity': 1})
        self.assertEqual(self.state['snapshot']['items'][0]['member_ids'], original[:1])
        self.command('undo')
        self.assertEqual(self.state['snapshot']['items'][0]['member_ids'], updated)

    def test_calibration_revision_reassigns_dependents_and_invalidates_all_confirmations(self):
        first = self.create(); second = self.create(); self.confirm(first); self.confirm(second)
        old = deepcopy(self.state['snapshot']['calibrations'][0])
        self.command('update_calibration', calibration_id=self.calibration_id, changes={'distance_m': 12, 'name': 'Corrected scale', 'uniform_scale': True})
        revised = self.state['revised_calibration_id']
        self.assertNotEqual(revised, self.calibration_id)
        self.assertEqual(self.state['snapshot']['calibrations'][0], old)
        for item, result in zip(self.state['snapshot']['items'], self.state['item_results']):
            self.assertEqual(item['measurement']['calibration_id'], revised)
            self.assertEqual(item['state'], 'draft'); self.assertIsNone(item['confirmation'])
            self.assertEqual(result['length_m'], 12)
        self.command('undo')
        self.assertTrue(all(item['measurement']['calibration_id'] == self.calibration_id for item in self.state['snapshot']['items']))
        self.assertTrue(all(item['state'] == 'draft' for item in self.state['snapshot']['items']))

    def test_duct_transfer_matches_manual_calculator_geometry_and_quantities(self):
        from estimator.workbook_calculators import calculator_session
        identity = self.create('duct', quantity=1); self.confirm(identity)
        preview = self.preview(identity, 'ductwork')
        manual = empty_schedule_inputs('ductwork')
        manual['CALCULATOR'].update({'B11': '600x400', 'C11': 'FyreWrap', 'D11': 10,
                                    'E11': '120/120/120', 'F11': 0, 'G11': 0, 'H11': 'Internal', 'I11': 'Horizontal'})
        _, actual, actual_lock = calculator_session('ductwork', preview['inputs'])
        _, expected, expected_lock = calculator_session('ductwork', manual)
        with actual_lock, expected_lock:
            self.assertEqual(actual.value('CALCULATOR', 'K11'), 20)
            for cell in ('J11', 'K11', 'L11', 'M11', 'N11', 'O11', 'R11', 'V11', 'W11', 'X11'):
                self.assertEqual(actual.value('CALCULATOR', cell), expected.value('CALCULATOR', cell), cell)

    def test_spray_and_board_full_outputs_match_manual_nonintegral_grouped_members(self):
        from estimator.workbook_calculators import calculator_session, source_model
        length = 7.123456789012345
        identity = self.create(quantity=3, measurement={'method': 'cited', 'length_m': length,
                                                     'citation': 'Explicit precision dimension on plan'})
        self.command('update_item', item_id=identity, changes={'fields': {'level': 'L2', 'zone': 'East'}})
        for calculator_id, sheet, manual_row in (
            ('steel_vermiculite', 'SCHEDULE', {'A10': 'B17', 'AA10': 'L2 / East', 'B10': 'CAFCO 300',
             'C10': 'Re-entrant - 3 sides', 'D10': 550, 'E10': 'Section', 'F10': '100UC15',
             'H10': 120, 'I10': 3, 'J10': length}),
            ('steel_board', 'CALCULATOR', {'A9': 'B17', 'B9': 'L2 / East', 'C9': 'TRAFALGAR COREX',
             'D9': '100UC15', 'F9': length * 3, 'G9': 3, 'H9': 120, 'I9': 'Beam', 'J9': 620}),
        ):
            with self.subTest(calculator=calculator_id):
                if calculator_id == 'steel_board':
                    self.command('update_item', item_id=identity, changes={'fields': {'product': 'TRAFALGAR COREX', 'critical_temperature': 620}})
                self.confirm(identity)
                preview = self.preview(identity, calculator_id)
                manual = empty_schedule_inputs(calculator_id)
                manual[sheet].update(manual_row)
                _, actual, actual_lock = calculator_session(calculator_id, preview['inputs'])
                _, expected, expected_lock = calculator_session(calculator_id, manual)
                model = source_model(calculator_id)
                with actual_lock, expected_lock:
                    compared = 0
                    for page in model['sheets']:
                        if page['name'] not in model['pages']:
                            continue
                        for address, cell in page['cells'].items():
                            if 'formula' in cell:
                                self.assertEqual(actual.value(page['name'], address), expected.value(page['name'], address),
                                                 (calculator_id, page['name'], address))
                                compared += 1
                    self.assertGreater(compared, 10000)
                if calculator_id == 'steel_vermiculite':
                    self.assertEqual(preview['inputs'][sheet]['I10'], 3)
                    self.assertEqual(preview['inputs'][sheet]['J10'], length)
                else:
                    self.assertEqual(preview['inputs'][sheet]['F9'], length * 3)
                    self.assertGreater(actual.value(sheet, 'AE9'), 0)
                    self.assertGreater(actual.value(sheet, 'AG9'), 0)
                    self.assertTrue(any('supports' in warning['message'] for warning in preview['warnings']))

    def test_board_native_unsupported_design_invalidates_entire_transfer_batch_before_apply(self):
        def temperatures(member):
            return next(column['options'] for column in self.service.options('steel_board',
                {'product': 'TRAFALGAR COREX', 'member_type': member})['columns'] if column['column'] == 'J')
        self.assertEqual(temperatures('Beam'), [620])
        self.assertEqual(temperatures('Column'), [550])
        supported = self.create()
        self.command('update_item', item_id=supported, changes={'fields': {'mark': 'Supported beam',
            'product': 'TRAFALGAR COREX', 'critical_temperature': 620}})
        self.confirm(supported)
        unsupported = self.create()
        self.command('update_item', item_id=unsupported, changes={'fields': {'mark': 'Unsupported beam',
            'product': 'TRAFALGAR COREX', 'critical_temperature': 550}})
        self.confirm(unsupported)
        original = empty_schedule_inputs('steel_board'); before = deepcopy(self.state['snapshot'])
        request = {'expected_revision': self.state['revision'], 'calculator_id': 'steel_board',
                   'inputs': deepcopy(original), 'schedule_rows': [9], 'item_ids': [supported, unsupported]}
        with self.assertRaisesRegex(ValidationError, 'Board transfer contains invalid items') as caught:
            self.service.preview_transfer(self.sid, request)
        message = str(caught.exception)
        self.assertIn('Unsupported beam', message)
        self.assertIn('NO BOARD DESIGN', message)
        self.assertIn('No thickness table for TRAFALGAR COREX, beam, 120 min at 550 C.', message)
        self.assertIn('keep the project requirements unchanged', message)
        self.assertEqual(request['inputs'], original)
        self.assertEqual(self.service.get(self.sid)['snapshot'], before)
        self.assertFalse(self.service._sessions[self.sid]['previews'])
        self.assertFalse(before['transfers'])
        valid = self.preview(supported, 'steel_board')
        self.assertEqual(valid['inputs']['CALCULATOR']['J9'], 620)
        self.assertTrue(valid['warnings'])

    def test_applied_transfer_retry_rechecks_changed_source_evidence_before_returning_calculator(self):
        identity = self.create(); self.confirm(identity)
        preview = self.preview(identity); request = self.apply(preview)
        snapshot = deepcopy(self.state['snapshot'])
        self.documents.blocked = True
        with self.assertRaisesRegex(ValidationError, 'Changed source evidence'):
            self.service.apply_transfer(self.sid, request)
        self.assertEqual(self.service.get(self.sid, verify_evidence=False)['snapshot'], snapshot)
        self.documents.blocked = False
        self.assertEqual(self.service.apply_transfer(self.sid, request), self.state)

    def test_failed_saved_source_binding_blocks_later_capture(self):
        identity = self.create(); self.confirm(identity)
        def fail(*args):
            raise OSError('Source removed between atomic save and binding')
        self.documents.bind_source = fail
        with self.assertRaises(OSError):
            self.service.saved_source(self.sid, self.state['snapshot'], 'saved.json')
        self.assertEqual(self.service.get(self.sid)['issues'][0]['code'], 'SAVE_SOURCE_UNAVAILABLE')
        with self.assertRaisesRegex(ValidationError, 'Retained evidence'):
            self.service.capture(self.sid, self.state['snapshot'])

    def test_stale_audit_head_cannot_approve_a_changed_native_snapshot(self):
        identity = self.create(); self.confirm(identity)
        snapshot = deepcopy(self.state['snapshot'])
        snapshot['items'][0]['fields']['level'] = 'Unaudited alteration'
        opened = self.service.open(snapshot, source_path='original-project.json')
        self.assertTrue(any(issue['code'] == 'AUDIT_STATE_MISMATCH' for issue in opened['issues']))
        self.assertEqual(opened['snapshot']['items'][0]['state'], 'draft')
        with self.assertRaisesRegex(ValidationError, 'Retained evidence'):
            self.service.capture(opened['session_id'], opened['snapshot'])

    def test_evidence_changes_and_restore_failures_block_every_gate(self):
        identity = self.create(); self.confirm(identity)
        self.documents.blocked = True
        with self.assertRaisesRegex(ValidationError, 'Changed source'):
            self.preview(identity)
        with self.assertRaisesRegex(ValidationError, 'Changed source'):
            self.service.export(self.sid, 'xlsx', [identity])
        self.assertTrue(self.service.get(self.sid)['issues'])
        self.assertIn(self.sid, self.documents.owners)
        self.documents.blocked = False
        opened = self.service.open(self.state['snapshot'], evidence_issues=[{'code': 'MISSING_AUDIT', 'message': 'Missing audit'}])
        with self.assertRaisesRegex(ValidationError, 'Retained evidence'):
            self.service.capture(opened['session_id'], opened['snapshot'])

    def test_imported_transfer_bindings_cannot_claim_a_manual_row(self):
        identity = self.create(); self.confirm(identity)
        preview = self.preview(identity); self.apply(preview)
        snapshot = deepcopy(self.state['snapshot'])
        snapshot['transfers'][0]['id'] = str(uuid4())
        opened = self.service.open(snapshot)
        self.assertEqual(opened['snapshot']['transfers'][0]['status'], 'conflict')

    def test_export_keeps_ids_sources_precision_and_literal_formula_text(self):
        identity = self.create()
        self.command('update_item', item_id=identity, changes={'fields': {'mark': '=HYPERLINK("bad")'},
            'measurement': {'method': 'cited', 'length_m': 7.123456789012345, 'citation': 'Dimension 7.123456789012345'}})
        self.confirm(identity)
        payload, mime, name = self.service.export(self.sid, 'xlsx', [identity])
        self.assertIn('spreadsheetml', mime); self.assertTrue(name.endswith('.xlsx'))
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook['Confirmed Takeoffs']; headers = {c.value: c.column for c in sheet[1]}
        self.assertEqual(sheet.cell(2, headers['Item ID']).value, identity)
        self.assertEqual(sheet.cell(2, headers['Mark / run']).value, '=HYPERLINK("bad")')
        self.assertEqual(sheet.cell(2, headers['Mark / run']).data_type, 's')
        self.assertEqual(sheet.cell(2, headers['Source SHA-256']).value, 'a'*64)
        self.assertEqual(sheet.cell(2, headers['Length per item m']).value, 7.123456789012345)
        self.assertFalse(any(c.data_type == 'f' for page in workbook for row in page for c in row))
        workbook.close()
        csv, _, _ = self.service.export(self.sid, 'csv', [identity])
        self.assertIn(b"'=HYPERLINK", csv)
        self.create()
        with self.assertRaisesRegex(ValidationError, 'Review|confirmed'):
            self.service.export(self.sid, 'xlsx')

    def test_xlsx_never_truncates_large_member_identity_or_geometry_evidence(self):
        import json
        identity = self.create(quantity=1000); self.confirm(identity)
        payload, _, _ = self.service.export(self.sid, 'xlsx', [identity])
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook['Confirmed Takeoffs']; headers = {c.value: c.column for c in sheet[1]}
        self.assertIn('Provenance Detail', sheet.cell(2, headers['Physical member IDs']).value)
        parts = [row for row in workbook['Provenance Detail'].iter_rows(min_row=2, values_only=True)
                 if row[0] == identity and row[1] == 'Physical member IDs']
        self.assertEqual(json.loads(''.join(row[4] for row in sorted(parts, key=lambda row: row[2]))),
                         self.state['snapshot']['items'][0]['member_ids'])
        basis = json.loads(sheet.cell(2, headers['Measurement basis']).value)
        self.assertEqual(basis['calibration']['distance_m'], 10)
        workbook.close()

    def test_evidence_field_support_is_mode_specific_hash_bound_and_exported(self):
        import json
        identity = self.create()
        original = item_digest(self.state['snapshot']['items'][0], self.state['snapshot'])
        reference = {'document_id': self.doc['id'], 'page': 1, 'note': 'Schedule row B17',
                     'fields': ['section', 'quantity', 'length_m']}
        self.command('update_item', item_id=identity, changes={'evidence': [reference]})
        self.assertNotEqual(item_digest(self.state['snapshot']['items'][0], self.state['snapshot']), original)
        self.confirm(identity)
        payload, _, _ = self.service.export(self.sid, 'xlsx', [identity])
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook['Confirmed Takeoffs']; headers = {c.value: c.column for c in sheet[1]}
        exported = json.loads(sheet.cell(2, headers['Evidence references']).value)
        self.assertEqual(exported[0]['fields'], reference['fields'])
        workbook.close()
        for fields in (['frl'], ['section', 'section'], ['quantity'] * 33, 'section', [None]):
            with self.subTest(fields=fields), self.assertRaisesRegex(ValidationError, 'Evidence fields'):
                self.command('update_item', item_id=identity, changes={'evidence': [{**reference, 'fields': fields}]})
        self.command('update_item', item_id=identity, changes={'evidence': [{**reference, 'fields': []}]})
        self.assertEqual(self.state['snapshot']['items'][0]['state'], 'draft')
        duct = self.create('duct', quantity=1, evidence=[{**reference, 'fields': ['frl', 'diameter_mm']}])
        with self.assertRaisesRegex(ValidationError, 'Evidence fields'):
            self.command('update_item', item_id=duct, changes={'evidence': [{**reference, 'fields': ['section']}]})

    def test_split_and_merge_preserve_actual_geometry_and_predecessors(self):
        original_id = self.create('duct', quantity=1)
        original = self.state['snapshot']['items'][0]
        def part(points):
            result = {k: deepcopy(original[k]) for k in ('mode', 'quantity', 'fields', 'measurement', 'geometry', 'evidence')}
            result['geometry']['points'] = points
            return result
        self.command('split_item', item_id=original_id, parts=[part([[10, 30], [60, 30]]), part([[60, 30], [110, 30]])])
        items = self.state['snapshot']['items']
        self.assertEqual(len(items), 2)
        self.assertEqual(sum(r['length_m'] for r in self.state['item_results']), 10)
        self.assertTrue(all(i['predecessor_ids'] == [original_id] for i in items))
        self.command('merge_items', item_ids=[i['id'] for i in items], item=part([[10, 30], [60, 30], [110, 30]]))
        self.assertEqual(len(self.state['snapshot']['items']), 1)
        self.assertEqual(set(self.state['snapshot']['items'][0]['predecessor_ids']), {original_id, *(item['id'] for item in items)})

    def test_pure_snapshot_validation_rejects_orphans_unknown_values_and_invalid_calibration(self):
        self.create()
        valid = self.state['snapshot']; self.assertEqual(validate_snapshot(valid), valid)
        invalid = deepcopy(valid); invalid['items'][0]['geometry']['document_id'] = str(uuid4())
        with self.assertRaises(ValidationError): validate_snapshot(invalid)
        invalid = deepcopy(valid); invalid['items'][0]['fields']['invented_quantity'] = 4
        with self.assertRaises(ValidationError): validate_snapshot(invalid)
        invalid = deepcopy(valid); invalid['calibrations'][0]['points'][1] = invalid['calibrations'][0]['points'][0]
        with self.assertRaises(ValidationError): validate_snapshot(invalid)

    def test_audit_digest_and_validation_avoid_redundant_copies_without_changing_evidence(self):
        from unittest.mock import patch
        from estimator.takeoff_model import audit_state_digest
        identity = self.create(); self.confirm(identity)
        state = deepcopy(self.state['snapshot'])
        state['companion_folder'] = 'portable locator'
        state['documents'][0]['pages'][0]['width'] = 200.0
        original = deepcopy(state)
        legacy = {key: deepcopy(value) for key, value in state.items() if key not in ('audit_head', 'companion_folder')}
        legacy['items'] = [{key: value for key, value in item.items() if key not in ('state', 'review', 'confirmation')}
                           for item in legacy['items']]
        legacy['transfers'] = [{key: value for key, value in binding.items() if key != 'status'}
                               for binding in legacy['transfers']]
        expected = digest(legacy)
        with patch('estimator.takeoff_model.deepcopy', side_effect=AssertionError('Unexpected evidence-tree copy')):
            self.assertEqual(audit_state_digest(state), expected)
            self.assertIs(validate_snapshot(state, copy_result=False), state)
        self.assertEqual(state, original)
        copied = validate_snapshot(state)
        copied['documents'][0]['pages'][0]['view'][0] = 999
        self.assertEqual(state, original)

    def test_malformed_operations_previews_and_huge_numbers_return_validation_errors(self):
        from estimator.takeoff_model import number
        before = deepcopy(self.state['snapshot'])
        with self.assertRaises(ValidationError):
            self.command([])
        with self.assertRaises(ValidationError):
            self.service.apply_transfer(self.sid, {'expected_revision': self.state['revision'],
                'request_id': str(uuid4()), 'preview_id': [], 'inputs': {}, 'schedule_rows': []})
        for value in (10**400, -(10**400), float('inf'), float('nan')):
            with self.subTest(value=str(value)), self.assertRaises(ValidationError):
                number(value, 'test number')
        self.assertEqual(self.service.get(self.sid)['snapshot'], before)

    def test_create_item_detaches_caller_geometry_and_measurement_before_auditing(self):
        geometry = {'document_id': self.doc['id'], 'page': 1, 'points': [[10, 30], [110, 30]]}
        measurement = {'method': 'cited', 'length_m': 10, 'citation': 'Source dimension'}
        self.create(geometry=geometry, measurement=measurement)
        committed = deepcopy(self.state['snapshot'])
        head = committed['audit_head']
        audit = self.documents.get_blob(head)
        geometry['points'][0][0] = 99
        measurement['length_m'] = 77
        self.assertEqual(self.service.get(self.sid)['snapshot'], committed)
        self.assertEqual(self.documents.get_blob(head), audit)

    def test_new_audit_events_record_actor_and_exact_changed_ids_in_bounded_history(self):
        from estimator.takeoff_model import audit_affected
        first = self.create(); second = self.create()
        self.command('bulk_update', item_ids=[second, first], changes={'fields': {'level': 'L2'}})
        row = self.service.history(self.sid, limit=1)['items'][0]
        self.assertEqual(row['actor'], {'kind': 'local-session', 'session_id': self.sid})
        self.assertEqual(row['affected_ids'], {'items': sorted([first, second]), 'documents': [],
                                             'calibrations': [], 'transfers': []})
        self.assertNotIn('before', row); self.assertNotIn('after', row)
        self.command('record_render', document_id=self.doc['id'], page=1, success=False, warnings=['Raster failed'])
        event = self.documents.get_blob(self.state['snapshot']['audit_head'])
        self.assertEqual(event['affected_ids']['documents'], [self.doc['id']])
        self.assertEqual(event['affected_ids']['items'], sorted([first, second]))
        self.assertEqual(audit_affected(event['before'], event['after']), event['affected_ids'])
        for malformed in (None, {'items': {}}, {'items': [None]}, {'items': [{'id': []}]},
                          {'render_checks': [{'document_id': self.doc['id'], 'page': True}]}):
            with self.subTest(malformed=malformed), self.assertRaises(ValidationError):
                audit_affected(malformed, event['after'])


if __name__ == '__main__':
    unittest.main()
