"""Canonical evidence, both file formats, and read-only HTTP boundaries."""
from copy import deepcopy
from contextlib import closing
from hashlib import sha256
from io import BytesIO
import http.client
import gc
import json
from pathlib import Path
import sqlite3
import tempfile
import threading
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from pypdf import PdfReader

from estimator.catalog import ValidationError
from estimator.reference_library import ReferenceLibrary
from estimator.server import create_server
from estimator.technical_item_exports import prepare_technical_item, build_technical_item_xlsx, render_technical_item_pdf
from technical_item_fixture_data import technical_library, REPORT_URL


def inventory(root):
    return {str(path.relative_to(root)): sha256(path.read_bytes()).hexdigest() for path in root.rglob('*') if path.is_file()}


def logical_database(path):
    with closing(sqlite3.connect(path.resolve().as_uri() + '?mode=ro', uri=True)) as connection:
        connection.execute('PRAGMA query_only=ON')
        return {name: sorted(repr(row) for row in connection.execute('SELECT * FROM "' + name.replace('"', '""') + '"'))
                for (name,) in connection.execute("SELECT name FROM sqlite_schema WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")}


class TechnicalItemExportsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name) / 'library'
        technical_library(self.root)
        self.library = ReferenceLibrary(self.root)

    def capture(self, body=None):
        return prepare_technical_item(body or {'id': 'report-a-v1'}, library=self.library)

    def test_refetched_canonical_record_and_diagram_identity_without_mutation(self):
        before = inventory(self.root)
        snapshot = self.capture()
        self.assertEqual(snapshot['record']['id'], 'report-a-v1')
        self.assertEqual(snapshot['images'][0]['source_sha256'], sha256((self.root / 'images/diagram-a.png').read_bytes()).hexdigest())
        self.assertEqual(len(snapshot['images']), 1)
        self.assertNotIn('Lead Time', [field['label'] for field in snapshot['fields']])
        self.assertEqual(snapshot['tables'][0]['columns'], ['Service', 'FRL'])
        self.assertEqual(inventory(self.root), before)

    def test_payload_cannot_supply_record_content_paths_urls_or_schedule(self):
        for body in ({'id': '../private'}, {'id': 'report-a-v1', 'record': {}}, {'id': 'report-a-v1', 'images': []},
                     {'id': 'report-a-v1', 'url': REPORT_URL}, {'id': ['report-a-v1']}, {'items': ['report-a-v1']}):
            with self.subTest(body=body), self.assertRaises(ValidationError):
                self.capture(body)

    def test_both_formats_preserve_report_hyperlink_diagram_and_literal_cell(self):
        snapshot = self.capture()
        before = deepcopy(snapshot)
        workbook = load_workbook(BytesIO(build_technical_item_xlsx(snapshot)), data_only=False)
        cells = [cell for sheet in workbook for row in sheet for cell in row]
        self.assertTrue(any(cell.value == '=Literal source text' and cell.data_type == 's' for cell in cells))
        linked = [cell for cell in cells if cell.hyperlink]
        self.assertEqual([(cell.value, cell.hyperlink.target) for cell in linked], [('SYN-123', REPORT_URL)])
        self.assertEqual(len(workbook['Source Diagrams']._images), 1)
        source = workbook['Source Diagrams']._images[0]
        self.assertEqual((source.width, source.height), (640, 360))
        self.assertTrue(any(cell.value == 'Copper pipe 25 mm' for cell in cells))
        self.assertFalse(any(cell.data_type == 'f' for cell in cells))
        reader = PdfReader(BytesIO(render_technical_item_pdf(snapshot)))
        for page in reader.pages:
            self.assertIn('ABN: 50 612 231 562', page.extract_text())
            self.assertIn('TECHNICAL ITEM', page.extract_text())
        text = '\n'.join(page.extract_text() for page in reader.pages)
        for value in ('SYN-123', 'Unlinked source report', '=Literal source text', 'Copper pipe 25 mm', 'Source information', snapshot['images'][0]['source_sha256']):
            self.assertIn(value, text)
        uris = [annotation.get_object()['/A']['/URI'] for page in reader.pages for annotation in page.get('/Annots', []) if '/A' in annotation.get_object() and '/URI' in annotation.get_object()['/A']]
        self.assertEqual(uris, [REPORT_URL])
        source_images = [image.image for page in reader.pages for image in page.images if image.image.size == (640, 360)]
        self.assertEqual(len(source_images), 1)
        self.assertLess(source_images[0].convert('RGB').getpixel((90, 90))[1], 70)
        self.assertGreater(source_images[0].convert('RGB').getpixel((480, 160))[2], 100)
        self.assertEqual(snapshot, before)

    def test_wide_table_all_source_columns_and_rows_remain_readable_and_complete(self):
        snapshot = self.capture({'id': 'report-a-v2'})
        snapshot['tables'] = [{'label': 'Source alternatives', 'caption': 'Every row retains its own source conditions.',
                               'columns': [f'Column {index}' for index in range(1, 21)],
                               'rows': [[f'row{row}-column{column}' for column in range(1, 21)] for row in (1, 2)]}]
        reader = PdfReader(BytesIO(render_technical_item_pdf(snapshot)))
        text = '\n'.join(page.extract_text() for page in reader.pages)
        self.assertIn('source row numbers align these parts.', text)
        for row in (1, 2):
            for column in range(1, 21):
                self.assertIn(f'row{row}-column{column}', text)
        workbook = load_workbook(BytesIO(build_technical_item_xlsx(snapshot)))
        self.assertEqual(workbook['Source Table 1']['A3'].value, snapshot['tables'][0]['caption'])
        self.assertEqual(list(workbook['Source Table 1'].iter_rows(min_row=6, max_row=7, values_only=True)), [tuple(row) for row in snapshot['tables'][0]['rows']])

    def test_tampered_registered_image_fails_closed_before_export(self):
        (self.root / 'images/diagram-a.png').write_bytes(b'changed original')
        with self.assertRaises(ValidationError):
            self.capture()

    def test_invalid_report_link_and_finite_work_budgets_fail_closed(self):
        canonical = self.library.detail('technical', 'report-a-v1')
        field = next(field for field in canonical['fields'] if field['label'] == 'Report Number')
        field['report_links'][0]['url'] = 'https://untrusted.example/report.pdf'
        with patch.object(self.library, 'detail', return_value=canonical), self.assertRaisesRegex(ValidationError, 'report reference'):
            self.capture()
        for name, value, expected in (('MAX_CELLS', 1, 'download size'), ('MAX_IMAGE_BYTES', 1, 'bounded'), ('MAX_IMAGE_PIXELS', 1, 'pixel budget'), ('MAX_IMAGES', 0, 'too many')):
            with self.subTest(name=name), patch('estimator.technical_item_exports.' + name, value), patch('estimator.technical_item_exports._render_diagram') as decode:
                with self.assertRaisesRegex(ValidationError, expected):
                    self.capture()
                decode.assert_not_called()

    def test_record_without_diagram_explicit_in_both_formats(self):
        snapshot = self.capture({'id': 'report-a-v2'})
        text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(render_technical_item_pdf(snapshot))).pages)
        self.assertIn('No source diagram is recorded.', text)
        workbook = load_workbook(BytesIO(build_technical_item_xlsx(snapshot)))
        self.assertEqual(len(workbook['Source Diagrams']._images), 0)

    def test_oversized_canonical_xlsx_cells_fail_before_workbook_assignment(self):
        canonical = self.library.detail('technical', 'report-a-v2')
        canonical['fields'][0]['value'] = 'A' * 40000
        with patch.object(self.library, 'detail', return_value=canonical):
            snapshot = self.capture({'id': 'report-a-v2'})
        self.assertEqual(snapshot['fields'][0]['value'], 'A' * 40000)
        with patch('estimator.technical_item_exports.Workbook') as workbook:
            with self.assertRaisesRegex(ValidationError, '32,767.*PDF.*original source'):
                build_technical_item_xlsx(snapshot)
            workbook.assert_not_called()
        for slot in ('title', 'caption', 'column', 'table_cell', 'diagram_caption'):
            with self.subTest(slot=slot):
                value = self.capture()
                if slot == 'title': value['tables'][0]['label'] = 'A' * 40000
                elif slot == 'caption': value['tables'][0]['caption'] = 'A' * 40000
                elif slot == 'column': value['tables'][0]['columns'][0] = 'A' * 40000
                elif slot == 'table_cell': value['tables'][0]['rows'][0][0] = 'A' * 40000
                else: value['images'][0]['caption'] = 'A' * 40000
                with patch('estimator.technical_item_exports.Workbook') as workbook, self.assertRaises(ValidationError):
                    build_technical_item_xlsx(value)
                workbook.assert_not_called()

    def test_xlsx_cell_utf16_boundary_and_long_literal_text_round_trip_exactly(self):
        for text in ('=' + 'A' * 32766, '\U0001f4cb' * 16383 + 'x'):
            with self.subTest(utf16_units=len(text.encode('utf-16-le')) // 2):
                snapshot = self.capture({'id': 'report-a-v2'})
                snapshot['fields'][0]['value'] = text
                workbook = load_workbook(BytesIO(build_technical_item_xlsx(snapshot)))
                try:
                    cells = [cell for sheet in workbook for row in sheet for cell in row if cell.value == text]
                    self.assertEqual(len(cells), 1)
                    self.assertEqual(cells[0].data_type, 's')
                finally: workbook.close()
        snapshot['fields'][0]['value'] = '\U0001f4cb' * 16384
        with patch('estimator.technical_item_exports.Workbook') as workbook, self.assertRaises(ValidationError):
            build_technical_item_xlsx(snapshot)
        workbook.assert_not_called()


class TechnicalItemHTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.folder = Path(cls.temp.name)
        cls.library = cls.folder / 'library'
        technical_library(cls.library)
        cls.database = cls.folder / 'qa.sqlite3'
        cls.server = create_server(0, cls.database, library_directory=cls.library)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown(); cls.server.server_close(); cls.thread.join(); gc.collect(); cls.temp.cleanup()

    def request(self, route, body, method='POST', headers=None):
        connection = http.client.HTTPConnection('127.0.0.1', self.server.server_port, timeout=30)
        connection.request(method, route, json.dumps(body), {'Content-Type': 'application/json', **(headers or {})})
        response = connection.getresponse()
        result = response.status, response.getheader('Content-Type'), response.getheader('Content-Disposition'), response.read()
        connection.close()
        return result

    def test_routes_refetch_canonical_evidence_and_preserve_database_library(self):
        files, database = inventory(self.library), logical_database(self.database)
        for extension, magic in (('pdf', b'%PDF-'), ('xlsx', b'PK')):
            status, mime, disposition, output = self.request('/api/technical/item.' + extension, {'id': 'report-a-v1'})
            self.assertEqual(status, 200)
            self.assertTrue(output.startswith(magic))
            self.assertIn('Technical-Item.' + extension, disposition)
            self.assertIn('application/', mime)
        self.assertEqual(inventory(self.library), files)
        self.assertEqual(logical_database(self.database), database)

    def test_method_origin_query_unknown_field_and_missing_record_guards(self):
        for route, body, method, headers, expected in (
                ('/api/technical/item.pdf', {'id': 'report-a-v1'}, 'PUT', {}, 405),
                ('/api/technical/item.pdf', {'id': 'report-a-v1'}, 'POST', {'Origin': 'https://untrusted.example'}, 403),
                ('/api/technical/item.xlsx?path=private', {'id': 'report-a-v1'}, 'POST', {}, 400),
                ('/api/technical/item.pdf', {'id': 'report-a-v1', 'text': 'untrusted'}, 'POST', {}, 400),
                ('/api/technical/item.xlsx', {'id': 'missing-record'}, 'POST', {}, 404)):
            with self.subTest(route=route, method=method):
                self.assertEqual(self.request(route, body, method, headers)[0], expected)

    def test_measure_icon_route_returns_exact_retained_asset_with_avif_mime(self):
        status, mime, _, payload = self.request('/icons/takeoff-measure.avif', {}, method='GET')
        self.assertEqual(status, 200)
        self.assertEqual(mime, 'image/avif')
        self.assertEqual(payload, (Path(__file__).resolve().parents[1] / 'static/icons/takeoff-measure.avif').read_bytes())


if __name__ == '__main__':
    unittest.main()
