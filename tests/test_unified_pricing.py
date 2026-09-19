"""Compact pricing rows retain the calculator and saved-quote contracts."""

import base64
from copy import deepcopy
import csv
import hashlib
import http.client
from io import BytesIO, StringIO
import json
from pathlib import Path
import tempfile
import threading
import unittest

from openpyxl import load_workbook

from estimator.calculator import specification
from estimator.catalog import baseline, effective_catalog
from estimator.estimate_composition import calculate
from estimator.server import create_server
from estimator.storage import Store


SHARED_INPUTS = {
    'D3': '1 Team - 1x', 'D4': '1 Team - 1x', 'D5': '1 Team - 1x',
    'B16': 10, 'D16': 'Promat Promamesh', 'E16': 0,
    'B20': 142, 'D20': '20kg SBR Latex - Promat', 'E20': 0,
    'B21': 284, 'D21': '20kg SBR Latex - Promat', 'E21': 0,
    'B26': 0, 'B27': 0,
}

USE_COLUMNS = ('Group', 'Selection name', 'Price source', 'Sell rate',
               'Yield type', 'Use yields', 'Rate ID', 'Use order')


def read_use_values(value):
    """Read test fixtures independently of the application's vector parser."""
    if value is None:
        return ['']
    if not isinstance(value, str):
        return [value]
    records = list(csv.reader(StringIO(value), delimiter=';', strict=True))
    return records[0] if records else ['']


def write_use_values(values):
    if len(values) == 1 and not isinstance(values[0], str):
        return values[0]
    stream = StringIO()
    csv.writer(stream, delimiter=';').writerow(values)
    return stream.getvalue()[:-2]


class UnifiedPricingIntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.database = Path(self.temp.name) / 'unified-pricing.sqlite3'
        self.store = Store(self.database)
        self.start_server()

    def start_server(self):
        self.server = create_server(0, self.database)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    def stop_server(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()

    def tearDown(self):
        self.stop_server()
        self.temp.cleanup()

    def request(self, method, path, payload=None, *, expected_status=(200, 201)):
        connection = http.client.HTTPConnection('127.0.0.1', self.server.server_port, timeout=30)
        try:
            connection.request(method, path, None if payload is None else json.dumps(payload),
                               {'Content-Type': 'application/json'})
            response = connection.getresponse()
            content = response.read()
            self.assertIn(response.status, expected_status, content[:800])
            return json.loads(content) if response.getheader('Content-Type', '').startswith('application/json') else content
        finally:
            connection.close()

    def stored_rows(self):
        with self.store.connect() as database:
            return {table: database.execute(f'SELECT * FROM {table} ORDER BY 1').fetchall()
                    for table in ('settings', 'quotes')}

    def export(self, configuration=None):
        return self.request('POST', '/api/pricing/export',
                            {} if configuration is None else {'configuration': configuration})

    def edit(self, payload, edits):
        workbook = load_workbook(BytesIO(payload))
        try:
            self.assertEqual(workbook.sheetnames, ['Inventory & Rates', 'Instructions'])
            sheet = workbook['Inventory & Rates']
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertNotIn('Row type', headers)
            for row_type, identity, fields in edits:
                if row_type == 'Inventory':
                    matches = [row[0].row for row in sheet.iter_rows(min_row=2)
                               if sheet.cell(row[0].row, headers['Inventory ID']).value == identity]
                    self.assertEqual(len(matches), 1, (row_type, identity))
                    for heading, value in fields.items():
                        sheet.cell(matches[0], headers[heading]).value = value
                else:
                    self.assertEqual(row_type, 'Use')
                    matches = [(row[0].row, index)
                               for row in sheet.iter_rows(min_row=2)
                               for index, rate_id in enumerate(read_use_values(
                                   sheet.cell(row[0].row, headers['Rate ID']).value))
                               if rate_id == identity]
                    self.assertEqual(len(matches), 1, (row_type, identity))
                    number, index = matches[0]
                    count = len(read_use_values(sheet.cell(number, headers['Rate ID']).value))
                    for heading, value in fields.items():
                        self.assertIn(heading, USE_COLUMNS)
                        cell = sheet.cell(number, headers[heading])
                        values = read_use_values(cell.value)
                        self.assertEqual(len(values), count, (identity, heading))
                        values[index] = value
                        cell.value = write_use_values(values)
            # Exercise an ordinary spreadsheet save, including its numeric precision.
            stream = BytesIO()
            workbook.save(stream)
            return stream.getvalue()
        finally:
            workbook.close()

    def preview(self, payload, current=None):
        body = {'filename': 'compact-pricing.xlsx', 'content_base64': base64.b64encode(payload).decode('ascii')}
        if current is not None:
            body['configuration'] = current
        proposed = self.request('POST', '/api/pricing/import', body)
        self.assertEqual(proposed['configuration']['catalog']['sources']['pricing_import'],
                         {'filename': 'compact-pricing.xlsx', 'sha256': hashlib.sha256(payload).hexdigest()})
        return proposed

    def assert_calculation(self, actual, inputs, configuration):
        expected = calculate(inputs, configuration)
        formula_cells = set(specification()['formulas'])
        self.assertEqual(len(formula_cells), 151)
        self.assertTrue(formula_cells <= actual['cells'].keys())
        # Compare all stored cells as well as the 151 source formula outputs.
        self.assertEqual(actual['cells'], expected['cells'])
        self.assertEqual(actual['errors'], expected['errors'])
        self.assertEqual(actual['summary'], expected['summary'])
        return actual['cells']

    def test_draft_export_and_shared_inventory_edit_preview_without_saving(self):
        draft = {'inventory': {'915': {'sales_price': 2222}}, 'rates': {}}
        before = self.stored_rows()
        payload = self.edit(self.export(draft), [
            ('Inventory', '204', {'Supplier price': 300, 'Markup': 0.2}),
        ])
        proposed = self.preview(payload, draft)
        self.assertEqual(self.stored_rows(), before)
        expected = deepcopy(draft)
        expected['inventory']['204'] = {'supplier_price': 300, 'markup': 0.2}
        actual = self.request('POST', '/api/calculate',
                              {'inputs': SHARED_INPUTS, 'configuration': proposed['configuration']})
        cells = self.assert_calculation(actual, SHARED_INPUTS, expected)
        self.assertEqual({address: cells[address] for address in ('A87', 'A92', 'B87', 'B92', 'F87', 'F92')},
                         {'A87': 360, 'A92': 360, 'B87': 1, 'B92': 2, 'F87': 360, 'F92': 720})
        self.assertEqual(cells['A89'], 2222)
        self.assertEqual(self.stored_rows(), before)
        self.assert_calculation(self.request('POST', '/api/calculate', {'inputs': SHARED_INPUTS}),
                                SHARED_INPUTS, {})

    def test_separate_use_override_and_yield_survive_round_trip_and_later_price_change(self):
        original = self.export()
        first_payload = self.edit(original, [
            ('Use', 'primers:1', {'Sell rate': 401.25, 'Yield type': 'Number', 'Use yields': 71}),
            ('Use', 'mesh:2', {'Yield type': 'Number', 'Use yields': 2.5}),
        ])
        first = self.preview(first_payload)['configuration']
        expected = {'inventory': {}, 'rates': {'primers:1': {'price': 401.25, 'yield': 71},
                                               'mesh:2': {'yield': 2.5}}}
        second_payload = self.edit(self.export(first), [
            ('Inventory', '204', {'Supplier price': 300, 'Markup': 0.2}),
        ])
        second = self.preview(second_payload, first)['configuration']
        expected['inventory']['204'] = {'supplier_price': 300, 'markup': 0.2}
        result = self.request('POST', '/api/calculate', {'inputs': SHARED_INPUTS, 'configuration': second})
        cells = self.assert_calculation(result, SHARED_INPUTS, expected)
        self.assertEqual({address: cells[address] for address in ('A87', 'F20', 'B87', 'A92', 'F21', 'B68')},
                         {'A87': 401.25, 'F20': 71, 'B87': 2, 'A92': 360, 'F21': 142, 'B68': 4})
        catalog = effective_catalog(second)
        primer = next(rate for rate in catalog['rate_groups']['primers'] if rate['id'] == 'primers:1')
        topcoat = next(rate for rate in catalog['rate_groups']['topcoats'] if rate['id'] == 'topcoats:1')
        self.assertEqual((primer['inventory_id'], topcoat['inventory_id']), ('204', '204'))
        self.assertEqual((primer['price_mode'], topcoat['price_mode']), ('override', 'inventory'))
        self.assertEqual((primer['price'], primer['yield'], topcoat['price'], topcoat['yield']),
                         (401.25, 71, 360, 142))
        workbook = load_workbook(BytesIO(self.export(second)))
        try:
            sheet = workbook['Inventory & Rates']
            headers = {cell.value: cell.column for cell in sheet[1]}
            rows = [row[0].row for row in sheet.iter_rows(min_row=2)
                    if sheet.cell(row[0].row, headers['Inventory ID']).value == '204']
            self.assertEqual(len(rows), 1)
            vectors = {name: read_use_values(sheet.cell(rows[0], headers[name]).value)
                       for name in USE_COLUMNS}
            self.assertEqual({len(values) for values in vectors.values()}, {2})
            uses = {values['Rate ID']: values for values in
                    (dict(zip(USE_COLUMNS, values)) for values in zip(*vectors.values()))}
            self.assertEqual(set(uses), {'primers:1', 'topcoats:1'})
            for rate_id, group, price_source, price, yield_value in (
                ('primers:1', 'primers', 'Override', 401.25, 71),
                ('topcoats:1', 'topcoats', 'Inventory', 360, 142),
            ):
                self.assertEqual((uses[rate_id]['Group'], uses[rate_id]['Price source']),
                                 (group, price_source))
                self.assertEqual((float(uses[rate_id]['Sell rate']), float(uses[rate_id]['Use yields'])),
                                 (price, yield_value))
        finally:
            workbook.close()
        round_trip = self.preview(self.export(second), second)
        self.assertEqual(round_trip['summary'], {'inventory': {'added': 0, 'removed': 0, 'updated': 0},
                                                'rates': {'added': 0, 'removed': 0, 'updated': 0}})
        self.assert_calculation(self.request('POST', '/api/calculate', {
            'inputs': SHARED_INPUTS, 'configuration': round_trip['configuration']}), SHARED_INPUTS, expected)
        restored_payload = self.edit(self.export(second), [('Use', 'primers:1', {'Price source': 'Inventory'})])
        restored = self.preview(restored_payload, second)['configuration']
        expected['rates']['primers:1'].pop('price')
        restored_result = self.request('POST', '/api/calculate', {'inputs': SHARED_INPUTS, 'configuration': restored})
        self.assert_calculation(restored_result, SHARED_INPUTS, expected)
        self.assertEqual(restored_result['cells']['A87'], 360)

    def test_imported_yield_kinds_keep_existing_numeric_and_error_results(self):
        original = self.export()
        before = self.stored_rows()
        for kind, workbook_value, stored_value, output in (
            ('Number', 2.5, 2.5, 4), ('Number', 0, 0, '#DIV/0!'),
            ('Blank', None, None, '#DIV/0!'), ('Empty text', None, '', '#VALUE!'),
        ):
            with self.subTest(kind=kind):
                payload = self.edit(original, [('Use', 'mesh:2', {'Yield type': kind, 'Use yields': workbook_value})])
                proposed = self.preview(payload)['configuration']
                mesh = next(rate for rate in effective_catalog(proposed)['rate_groups']['mesh']
                            if rate['id'] == 'mesh:2')
                self.assertEqual(mesh['yield'], stored_value)
                self.assertEqual(type(mesh['yield']), type(stored_value))
                actual = self.request('POST', '/api/calculate', {'inputs': SHARED_INPUTS, 'configuration': proposed})
                self.assert_calculation(actual, SHARED_INPUTS, {'rates': {'mesh:2': {'yield': stored_value}}})
                self.assertEqual(actual['cells']['B68'], output)
                self.assertEqual(self.stored_rows(), before)

    def test_saved_quote_keeps_all_formula_results_and_its_library_after_save_and_restart(self):
        old_quote = self.request('POST', '/api/quotes', {'title': 'Before compact pricing', 'inputs': SHARED_INPUTS})
        before = self.stored_rows()
        payload = self.edit(self.export(), [
            ('Inventory', '204', {'Supplier price': 300, 'Markup': 0.2}),
            ('Inventory', '915', {'Sell price': 2222}),
            ('Use', 'mesh:2', {'Yield type': 'Number', 'Use yields': 2.5}),
        ])
        proposed = self.preview(payload)['configuration']
        self.assertEqual(self.stored_rows(), before)
        self.request('PUT', '/api/configuration', proposed)
        self.assertEqual(self.store.configuration(), proposed)
        expected = {'inventory': {'204': {'supplier_price': 300, 'markup': 0.2}, '915': {'sales_price': 2222}},
                    'rates': {'mesh:2': {'yield': 2.5}}}
        self.assert_calculation(self.request('POST', '/api/calculate', {'inputs': SHARED_INPUTS}), SHARED_INPUTS, expected)
        new_quote = self.request('POST', '/api/quotes', {'title': 'After compact pricing', 'inputs': SHARED_INPUTS})
        self.assert_calculation(new_quote['result'], SHARED_INPUTS, expected)
        self.assertNotEqual(old_quote['result']['summary']['total'], new_quote['result']['summary']['total'])
        self.assertEqual(self.request('GET', f'/api/quotes/{old_quote["id"]}'), old_quote)
        self.assertEqual(old_quote['source_hashes'], baseline()['sources'])
        self.assertEqual(new_quote['source_hashes']['pricing_import']['sha256'], hashlib.sha256(payload).hexdigest())
        self.stop_server()
        self.start_server()
        self.assertEqual(self.store.configuration(), proposed)
        self.assertEqual(self.request('GET', f'/api/quotes/{old_quote["id"]}'), old_quote)
        self.assertEqual(self.request('GET', f'/api/quotes/{new_quote["id"]}'), new_quote)
        self.assert_calculation(self.request('POST', '/api/calculate', {
            'inputs': old_quote['inputs'], 'configuration': old_quote['configuration']}), old_quote['inputs'], {})
        self.assert_calculation(self.request('POST', '/api/calculate', {'inputs': SHARED_INPUTS}), SHARED_INPUTS, expected)

    def test_mismatched_use_lists_are_rejected_without_saving_or_losing_the_draft(self):
        saved = {'inventory': {'915': {'sales_price': 1900}}, 'rates': {}}
        self.request('PUT', '/api/configuration', saved)
        old_quote = self.request('POST', '/api/quotes', {'title': 'Before rejected import', 'inputs': SHARED_INPUTS})
        before = self.stored_rows()
        draft = {'inventory': {'915': {'sales_price': 2222}},
                 'rates': {'primers:1': {'price': 401.25, 'yield': 71}}}
        original = self.export(draft)
        # Group is now the editable scalar-view membership list. Its edits are
        # reconciled by stable group identities; hidden vectors still align.
        for heading, shortened_value in (('Use yields', 71), ('Selection name', 'Only one'), ('Rate ID', 'primers:1')):
            with self.subTest(heading=heading):
                malformed = self.edit(original, [('Inventory', '204', {heading: shortened_value})])
                rejected = self.request('POST', '/api/pricing/import', {
                    'filename': 'mismatched-pricing.xlsx',
                    'content_base64': base64.b64encode(malformed).decode('ascii'),
                    'configuration': draft,
                }, expected_status=(400,))
                self.assertTrue(rejected.get('error'))
                self.assertNotIn('configuration', rejected)
                self.assertEqual(self.stored_rows(), before)
        # The same unsaved draft can still be exported, previewed and calculated.
        proposed = self.preview(self.export(draft), draft)
        self.assertEqual(proposed['summary'], {'inventory': {'added': 0, 'removed': 0, 'updated': 0},
                                             'rates': {'added': 0, 'removed': 0, 'updated': 0}})
        self.assert_calculation(self.request('POST', '/api/calculate', {
            'inputs': SHARED_INPUTS, 'configuration': proposed['configuration']}), SHARED_INPUTS, draft)
        self.assert_calculation(self.request('POST', '/api/calculate', {'inputs': SHARED_INPUTS}),
                                SHARED_INPUTS, saved)
        self.assertEqual(self.request('GET', f'/api/quotes/{old_quote["id"]}'), old_quote)
        self.assertEqual(self.stored_rows(), before)


if __name__ == '__main__':
    unittest.main()
