"""Real loopback HTTP coverage of calculator boundaries and saved input state.

Every run owns a temporary database and ephemeral port. No live estimator data,
authoritative workbooks, or the separate manual QA server are modified.
"""

import base64
import hashlib
import http.client
from io import BytesIO
import json
from pathlib import Path
import tempfile
import threading
import unittest

from openpyxl import load_workbook

from estimator.catalog import ROOT
from estimator.calculator_defaults import default_calculator_inputs
from estimator.pricing_workbook import _serialize_exact
from estimator.server import create_server, MAX_BODY
from estimator.storage import Store


PAGES = {
    "steel_vermiculite": ["CALCULATOR", "SCHEDULE", "BAGS", "SETTINGS"],
    "ductwork": ["CALCULATOR", "SUMMARY", "PRODUCT SETTINGS"],
    "steel_board": ["START", "CALCULATOR", "BOARD SUMMARY", "EXTRA BOARDS", "SETTINGS"],
}


class WorkbookCalculatorApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.database = Path(cls.temp.name) / "calculator-api.sqlite3"
        cls.server = create_server(0, cls.database)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.source_hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest()
                             for path in (ROOT / "data" / "calculators").glob("*.json.gz")}
        cls.template = None

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()
        cls.temp.cleanup()

    def setUp(self):
        self.store = Store(self.database)
        with self.store.connect() as db:
            db.execute("DELETE FROM calculator_states")

    def request(self, method, path, body=None, headers=None):
        connection = http.client.HTTPConnection("127.0.0.1", self.server.server_port, timeout=60)
        try:
            connection.request(method, path, body=json.dumps(body) if isinstance(body, dict) else body,
                               headers={"Content-Type": "application/json", **(headers or {})})
            response = connection.getresponse()
            return response.status, dict(response.getheaders()), response.read()
        finally:
            connection.close()

    def json_request(self, method, path, body=None, expected=200):
        status, _, payload = self.request(method, path, body)
        self.assertEqual(status, expected, payload.decode("utf-8", errors="replace")[:1000])
        return json.loads(payload)

    def definition(self, identity="ductwork"):
        return self.json_request("GET", f"/api/calculators/{identity}")

    def calculate(self, inputs=None, *, identity="ductwork", sheet="CALCULATOR", row=11, count=1):
        body = {"sheet": sheet, "start_row": row, "row_count": count}
        if inputs is not None:
            body["inputs"] = inputs
        return self.json_request("POST", f"/api/calculators/{identity}/calculate", body)

    def save(self, inputs, identity="ductwork"):
        return self.json_request("PUT", f"/api/calculators/{identity}/state", {"inputs": inputs})

    def exported(self):
        if self.__class__.template is None:
            status, _, payload = self.request("POST", "/api/calculators/ductwork/template", {})
            self.assertEqual(status, 200)
            self.__class__.template = payload
        return self.__class__.template

    def upload(self, payload=None, inputs=None):
        body = {"filename": "duct-schedule.xlsx", "content_base64": base64.b64encode(payload or self.exported()).decode()}
        if inputs is not None:
            body["inputs"] = inputs
        return body

    @staticmethod
    def cell(page, address):
        return next(cell for row in page["rows"] for cell in row["cells"] if cell["address"] == address)

    def stored_rows(self):
        with self.store.connect() as db:
            return db.execute("SELECT id,data,updated_at FROM calculator_states ORDER BY id").fetchall()

    def test_definitions_expose_exact_twelve_pages_and_official_documents(self):
        listing = self.json_request("GET", "/api/calculators")
        self.assertEqual([item["id"] for item in listing["calculators"]], list(PAGES))
        total_pages = 0
        for identity, pages in PAGES.items():
            with self.subTest(identity=identity):
                definition = self.definition(identity)
                self.assertEqual(definition["pages"], pages)
                self.assertEqual([sheet["name"] for sheet in definition["sheets"]], pages)
                self.assertEqual(definition["inputs"], default_calculator_inputs(identity))
                self.assertEqual(len(definition["source"]["sha256"]), 64)
                self.assertTrue(definition["documents"])
                self.assertTrue(all(item["url"].startswith("https://") for item in definition["documents"]))
                for name in pages:
                    page = self.calculate(identity=identity, sheet=name, row=1, count=1)
                    self.assertEqual(page["sheet"], name)
                    self.assertEqual((page["start_row"], page["end_row"]), (1, 1))
                    total_pages += 1
        self.assertEqual(total_pages, 12)
        self.assertEqual(self.stored_rows(), [])

    def test_zero_width_source_helpers_are_hidden_in_worksheet_metadata(self):
        definition = self.definition()
        metadata = next(sheet for sheet in definition["sheets"] if sheet["name"] == "CALCULATOR")
        # The original workbook hides AS:CL by width=0, while hidden="0".
        # These are helper formulas; the visible estimating columns end at AQ.
        self.assertEqual(metadata["hidden_columns"], list(range(45, 91)))
        self.assertTrue(all(metadata["column_widths"][str(column)] == 0
                            for column in range(45, 91)))
        self.assertFalse(set(range(1, 44)) & set(metadata["hidden_columns"]))
        self.assertEqual(metadata["hidden_rows"], [])
        page = self.calculate()
        self.assertEqual(page["hidden_columns"], metadata["hidden_columns"])
        self.assertEqual(page["hidden_rows"], metadata["hidden_rows"])
        self.assertEqual(self.stored_rows(), [])

    def test_draft_calculation_changes_expected_outputs_without_saving(self):
        initial = self.calculate()
        self.assertAlmostEqual(self.cell(initial, "K11")["value"], 10)
        changed = self.calculate({"CALCULATOR": {"D11": 12.345678901234567}})
        self.assertEqual(self.cell(changed, "D11")["value"], 12.345678901234567)
        self.assertAlmostEqual(self.cell(changed, "K11")["value"], 12.345678901234567)
        self.assertFalse(self.cell(changed, "K11")["editable"])
        self.assertTrue(self.cell(changed, "K11")["calculated"])
        self.assertTrue(self.cell(changed, "D11")["editable"])
        self.assertEqual(self.definition()["inputs"], {})
        self.assertEqual(self.stored_rows(), [])
        self.assertEqual(self.cell(self.calculate(), "K11")["value"], self.cell(initial, "K11")["value"])

    def test_settings_affect_yield_and_calculations_are_session_isolated(self):
        inputs = {"PRODUCT SETTINGS": {"B44": 10, "B45": 60, "B46": 20}}
        changed = self.calculate(inputs)
        self.assertAlmostEqual(self.cell(changed, "M11")["value"], 20)
        original = self.calculate({})
        self.assertAlmostEqual(self.cell(original, "M11")["value"], 11.7)
        self.assertEqual(self.definition()["inputs"], {})
        self.assertEqual(inputs["PRODUCT SETTINGS"]["B46"], 20)

    def test_save_reopen_and_fresh_store_preserve_precision_and_source_hash(self):
        inputs = {"CALCULATOR": {"D11": 2.1234567890123457, "F11": 0, "G11": None},
                  "PRODUCT SETTINGS": {"B97": 1.22}}
        saved = self.save(inputs)
        self.assertEqual(saved["inputs"], inputs)
        self.assertEqual(saved["source_sha256"], self.definition()["source"]["sha256"])
        self.assertTrue(saved["updated_at"])
        self.assertEqual(self.definition()["inputs"], inputs)
        self.assertEqual(Store(self.database).calculator_state("ductwork"), saved)
        page = self.calculate()
        self.assertEqual(self.cell(page, "D11")["value"], inputs["CALCULATOR"]["D11"])
        self.assertEqual(self.cell(page, "F11")["value"], 0)
        self.assertIsNone(self.cell(page, "G11")["value"])
        self.assertEqual(self.json_request("GET", "/api/quotes"), {"quotes": []})
        self.assertEqual(self.json_request("GET", "/api/configuration"), {"inventory": {}, "rates": {}})

    def test_saved_states_are_separate_for_all_calculators(self):
        supplied = {"ductwork": {"CALCULATOR": {"D11": 1.25}},
                    "steel_vermiculite": {"SCHEDULE": {"A10": "VERM-1"}},
                    "steel_board": {"CALCULATOR": {"A9": "BOARD-1"}, "EXTRA BOARDS": {"A6": "Extra"}}}
        for identity, inputs in supplied.items():
            self.save(inputs, identity)
        self.assertEqual(len(self.stored_rows()), 3)
        for identity, inputs in supplied.items():
            self.assertEqual(self.definition(identity)["inputs"], inputs)
        self.save({"CALCULATOR": {"D11": 7}})
        for identity in ("steel_vermiculite", "steel_board"):
            self.assertEqual(self.definition(identity)["inputs"], supplied[identity])

    def test_template_response_is_a_bounded_values_only_excel_download(self):
        status, headers, payload = self.request("POST", "/api/calculators/ductwork/template", {})
        self.assertEqual(status, 200)
        self.assertEqual(headers["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertEqual(headers["Content-Disposition"], 'attachment; filename="ceasefire-ductwork-schedule.xlsx"')
        self.assertEqual(headers["Cache-Control"], "no-store")
        self.assertEqual(int(headers["Content-Length"]), len(payload))
        workbook = load_workbook(BytesIO(payload))
        self.assertEqual(workbook.sheetnames, ["CALCULATOR", "Instructions"])
        self.assertEqual(workbook["CALCULATOR"].max_column, 8)
        self.assertEqual(workbook["CALCULATOR"].max_row, 301)
        self.assertFalse(any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row))
        workbook.close()
        self.assertEqual(self.stored_rows(), [])

    def test_import_is_a_draft_then_save_replaces_schedule_and_keeps_settings(self):
        old = {"CALCULATOR": {"B11": "400x200", "D11": 4}, "PRODUCT SETTINGS": {"B97": 1.22}}
        self.save(old)
        before = self.stored_rows()
        workbook = load_workbook(BytesIO(self.exported()))
        sheet = workbook["CALCULATOR"]
        for column, value in enumerate(["250x250", "CAFCO 300", 0.12345678901234566, "120/120/120", 0, 0, "External", "Horizontal"], 1):
            sheet.cell(2, column, value)
        sheet["A301"] = "500x250"
        sheet.row_dimensions[301].hidden = True
        payload = _serialize_exact(workbook)
        workbook.close()
        proposed = self.json_request("POST", "/api/calculators/ductwork/import", self.upload(payload))
        self.assertEqual(proposed["imported_rows"], 2)
        self.assertEqual(proposed["source_sha256"], hashlib.sha256(payload).hexdigest())
        self.assertEqual(proposed["inputs"]["CALCULATOR"]["D11"], 0.12345678901234566)
        self.assertEqual(proposed["inputs"]["CALCULATOR"]["F11"], 0)
        self.assertEqual(proposed["inputs"]["CALCULATOR"]["B310"], "500x250")
        self.assertIsNone(proposed["inputs"]["CALCULATOR"]["B12"])
        self.assertEqual(proposed["inputs"]["PRODUCT SETTINGS"], old["PRODUCT SETTINGS"])
        self.assertEqual(self.stored_rows(), before)
        self.assertEqual(self.definition()["inputs"], old)
        saved = self.save(proposed["inputs"])
        self.assertEqual(self.definition()["inputs"], proposed["inputs"])
        self.assertEqual(Store(self.database).calculator_state("ductwork"), saved)
        page = self.calculate()
        self.assertAlmostEqual(self.cell(page, "K11")["value"], 0.12345678901234566)
        self.assertEqual(self.definition("steel_board")["inputs"], {})

    def test_source_version_guard_preserves_existing_saved_inputs_on_every_write(self):
        stale = {"inputs": {"CALCULATOR": {"D11": 123}}, "source_sha256": "old-workbook-hash", "updated_at": "old-time"}
        with self.store.connect() as db:
            db.execute("INSERT INTO calculator_states VALUES(?,?,?)", ("ductwork", json.dumps(stale), "old-time"))
        before = self.stored_rows()
        for method, route, body in [
            ("GET", "/api/calculators/ductwork", None),
            ("POST", "/api/calculators/ductwork/calculate", {"inputs": {}, "row_count": 1}),
            ("POST", "/api/calculators/ductwork/import", self.upload()),
            ("PUT", "/api/calculators/ductwork/state", {"inputs": {"CALCULATOR": {"D11": 2}}}),
        ]:
            with self.subTest(route=route):
                response = self.json_request(method, route, body, expected=400)
                self.assertIn("version", response["error"])
                self.assertEqual(self.stored_rows(), before)
        self.assertEqual(self.definition("steel_board")["inputs"], {})

    def test_calculated_cells_reference_databases_and_unknown_cells_cannot_be_written(self):
        before = self.stored_rows()
        protected = [("ductwork", {"CALCULATOR": {"K11": 999}}),
                     ("ductwork", {"PRODUCT SETTINGS": {"J137": "Invented"}}),
                     ("ductwork", {"CALCULATOR": {"B311": "Outside capacity"}}),
                     ("steel_vermiculite", {"SECTIONS": {"A6": "Invented"}}),
                     ("steel_board", {"STEEL LIBRARY": {"A6": "Invented"}}),
                     ("steel_board", {"CALCULATOR": {"AI9": "Invented"}})]
        for identity, inputs in protected:
            for action, method in (("calculate", "POST"), ("state", "PUT")):
                with self.subTest(identity=identity, action=action, inputs=inputs):
                    self.json_request(method, f"/api/calculators/{identity}/{action}", {"inputs": inputs}, expected=400)
                    self.assertEqual(self.stored_rows(), before)
        for path, original in self.source_hashes.items():
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), original)

    def test_input_payload_shapes_types_and_oversized_numbers_fail_cleanly(self):
        self.save({"CALCULATOR": {"D11": 10}})
        before = self.stored_rows()
        invalid = [[], True, {"CALCULATOR": []}, {"no worksheet": {}},
                   {"CALCULATOR": {"D11": "12.5"}}, {"CALCULATOR": {"D11": True}},
                   {"CALCULATOR": {"D11": {"value": 12}}}, {"CALCULATOR": {"D11": 10**400}},
                   {"CALCULATOR": {"D11": 1e101}}, {"CALCULATOR": {"B11": "x" * 2001}},
                   {"CALCULATOR": {"B11": "control\u0000character"}}]
        for inputs in invalid:
            with self.subTest(inputs=str(inputs)[:100]):
                self.json_request("PUT", "/api/calculators/ductwork/state", {"inputs": inputs}, expected=400)
                self.assertEqual(self.stored_rows(), before)
        self.json_request("PUT", "/api/calculators/ductwork/state", {"inputs": None}, expected=400)
        self.assertEqual(self.stored_rows(), before)

    def test_page_bounds_and_hidden_database_pages_are_rejected(self):
        invalid = [{"start_row": 0}, {"start_row": 311}, {"start_row": True}, {"start_row": "11"},
                   {"row_count": 0}, {"row_count": 51}, {"row_count": 1.5}, {"sheet": []},
                   {"sheet": {}}, {"sheet": False}, {"sheet": 0}, {"sheet": "DATABASE"}]
        for body in invalid:
            request = {"sheet": "CALCULATOR", "start_row": 11, "row_count": 1, **body}
            with self.subTest(body=body):
                self.json_request("POST", "/api/calculators/ductwork/calculate", request, expected=400)
        self.json_request("POST", "/api/calculators/steel_board/calculate", {"sheet": "STEEL LIBRARY"}, expected=400)
        self.assertEqual(self.calculate(row=310, count=50)["end_row"], 310)

    def test_unknown_payload_fields_methods_identifiers_and_json_are_rejected(self):
        for action, method, body in [("calculate", "POST", {"formulas": {}}), ("state", "PUT", {}),
                                     ("state", "PUT", {"inputs": {}, "source_sha256": "fake"}),
                                     ("template", "POST", {"inputs": {}}), ("import", "POST", {"configuration": {}})]:
            self.json_request(method, f"/api/calculators/ductwork/{action}", body, expected=400)
        for action, method in (("calculate", "PUT"), ("state", "POST"), ("template", "PUT"), ("import", "PUT")):
            self.json_request(method, f"/api/calculators/ductwork/{action}", {}, expected=405)
        self.json_request("GET", "/api/calculators/unknown", expected=400)
        for body in ("[]", "{", '{"inputs":{"CALCULATOR":{"D11":NaN}}}', '{"inputs":{"CALCULATOR":{"D11":Infinity}}}'):
            self.assertEqual(self.request("POST", "/api/calculators/ductwork/calculate", body)[0], 400)
        self.assertEqual(self.stored_rows(), [])

    def test_bad_uploads_leave_saved_state_unchanged(self):
        self.save({"CALCULATOR": {"D11": 7}})
        before = self.stored_rows()
        invalid = [{}, {"filename": "schedule.xlsm", "content_base64": "QQ=="},
                   {"filename": "schedule.xlsx", "content_base64": "bad!"},
                   {"filename": "schedule.xlsx", "content_base64": ""},
                   {"filename": "schedule.xlsx", "content_base64": "QQ=="},
                   {"filename": "schedule.xlsx", "content_base64": "A" * (7 * 1024 * 1024)}]
        for body in invalid:
            self.json_request("POST", "/api/calculators/ductwork/import", body, expected=400)
            self.assertEqual(self.stored_rows(), before)
        workbook = load_workbook(BytesIO(self.exported()))
        workbook["CALCULATOR"]["C2"] = "=1+1"
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        self.json_request("POST", "/api/calculators/ductwork/import", self.upload(stream.getvalue()), expected=400)
        self.json_request("POST", "/api/calculators/ductwork/import", self.upload(inputs={"PRODUCT SETTINGS": {"J137": "Injected"}}), expected=400)
        self.assertEqual(self.stored_rows(), before)

    def test_origin_host_content_type_and_request_size_boundaries_apply(self):
        for headers in ({"Host": "attacker.example"}, {"Origin": "https://attacker.example"}, {"Origin": "null"}):
            self.assertEqual(self.request("PUT", "/api/calculators/ductwork/state", {"inputs": {}}, headers)[0], 403)
        self.assertEqual(self.request("POST", "/api/calculators/ductwork/calculate", "inputs=123", {"Content-Type": "text/plain"})[0], 400)
        self.assertEqual(self.request("POST", "/api/calculators/ductwork/calculate", "{}", {"Content-Length": str(MAX_BODY + 1)})[0], 400)
        for path in ("/data/calculators/ductwork.json.gz", "/.runtime/calculator-api.sqlite3", "/estimator/workbook_catalog.py"):
            self.assertEqual(self.request("GET", path)[0], 404)
        self.assertEqual(self.stored_rows(), [])


if __name__ == "__main__":
    unittest.main()
