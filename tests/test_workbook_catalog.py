"""Source preservation and editable-boundary checks for packaged calculators."""

from copy import deepcopy
import gzip
import hashlib
import json
import os
from pathlib import Path
import tempfile
import unittest
import xml.etree.ElementTree as ET

from estimator.catalog import ValidationError
from estimator.workbook_catalog import (
    CATALOG_SPECS, DATA_DIRECTORY, column_name, column_number, editable_cells,
    list_workbook_catalogs, load_workbook_catalog, range_addresses, setting_cells,
)
from scripts.import_calculators import DEFAULT_SOURCE_DIRECTORY, XML_SPACE, extract_calculator, normalize_literal_text, write_catalog
from scripts.import_workbooks import SourceWorkbook, NS
from scripts.check_calculator_integrity import compare_catalogs, format_catalog_difference


SOURCE_HASHES = {
    "steel_vermiculite": "1ea62906d13f2f5d34bae9c6f2391e084f26da5d597c9d69b154ad99579028ad",
    "ductwork": "9b2e5388a0118c4b3f66ea57f582d487d156585b45d9f34ebc8178f270ff7462",
    "steel_board": "934e951e4255976a7d3fea64f77b0b0a7f1852174fecb126e76a640f0a38546d",
}


class WorkbookCatalogTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.books = {identifier: load_workbook_catalog(identifier) for identifier in CATALOG_SPECS}

    def test_exact_modules_pages_and_source_inventory(self):
        index = list_workbook_catalogs()
        self.assertEqual([entry["id"] for entry in index], list(CATALOG_SPECS))
        self.assertEqual([entry["title"] for entry in index],
                         ["Structural Steel (vermiculite)", "Ductwork", "Structural Steel (board)"])
        expected = {"steel_vermiculite": (8, 120973), "ductwork": (3, 24059), "steel_board": (11, 16534)}
        for identifier, book in self.books.items():
            with self.subTest(identifier=identifier):
                self.assertEqual(book["source"]["sha256"], SOURCE_HASHES[identifier])
                self.assertEqual((len(book["sheets"]), book["counts"]["formulas"]), expected[identifier])
                self.assertFalse(book["source_features"]["vba"])
                self.assertEqual(book["source_features"]["external_links"], [])
                self.assertEqual(book["source_features"]["connections"], [])
                self.assertTrue(all(page in {sheet["name"] for sheet in book["sheets"]} for page in book["pages"]))
        self.assertEqual(self.books["steel_vermiculite"]["pages"], ["CALCULATOR", "SCHEDULE", "BAGS", "SETTINGS"])
        self.assertEqual(self.books["steel_board"]["pages"], ["START", "CALCULATOR", "BOARD SUMMARY", "EXTRA BOARDS", "SETTINGS"])

    def test_formula_caches_never_replace_executable_formula_or_literal(self):
        for book in self.books.values():
            for sheet in book["sheets"]:
                for address, cell in sheet["cells"].items():
                    if "formula" in cell:
                        self.assertIsInstance(cell["formula"], str, (sheet["name"], address))
                        self.assertTrue(cell["formula"], (sheet["name"], address))
                        self.assertFalse(cell["formula"].startswith("="))
                        self.assertNotIn("value", cell)
                    else:
                        self.assertNotIn("cached_value", cell)
        sheet = next(sheet for sheet in self.books["steel_vermiculite"]["sheets"] if sheet["name"] == "SCHEDULE")
        self.assertEqual(sheet["cells"]["M11"]["cached_value"], "")
        self.assertEqual(sheet["cells"]["M11"]["data_type"], "str")
        self.assertNotIn("value", sheet["cells"]["B11"])

    def test_existing_source_errors_and_known_duct_formula_are_preserved(self):
        errors = self.books["steel_vermiculite"]["cached_errors"]
        self.assertEqual(len(errors), 31)
        self.assertTrue({"sheet": "AUDIT", "cell": "E371", "error": "#VALUE!"} in errors)
        self.assertFalse(any(error["sheet"] in self.books["steel_vermiculite"]["pages"] for error in errors))
        sheet = next(sheet for sheet in self.books["ductwork"]["sheets"] if sheet["name"] == "CALCULATOR")
        self.assertIn("M6", sheet["cells"]["AL11"]["formula"])
        self.assertIn("M7", sheet["cells"]["AL12"]["formula"])

    def test_literal_text_whitespace_matches_independent_native_excel_probe(self):
        fixture = json.loads((Path(__file__).parent / "fixtures/calculators/literal-whitespace-probe.json").read_text(encoding="utf-8"))
        self.assertEqual(fixture["oracle"]["application"], "Microsoft Excel")
        for item in fixture["cases"]:
            label, kind, value, space = item["input"]
            if kind not in ("str", "formula"):
                continue
            with self.subTest(case=label):
                namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                node = ET.Element(namespace + "c", {"t": "str"})
                raw = ET.SubElement(node, namespace + "v")
                raw.text = value
                if space in ("value", "default"):
                    raw.attrib[XML_SPACE] = "preserve" if space == "value" else "default"
                elif space == "cell":
                    node.attrib[XML_SPACE] = "preserve"
                if kind == "formula":
                    ET.SubElement(node, namespace + "f").text = '"' + value + '"'
                actual = normalize_literal_text(node, value, "preserve" if space == "row" else "default")
                self.assertEqual(actual, item["expected_cell"])
        # The raw source is retained alongside Excel's interpreted literal.
        lookup = next(sheet for sheet in self.books["steel_board"]["sheets"] if sheet["name"] == "LOOKUP CACHE")
        self.assertEqual(len(lookup["cells"]["B314"]["value"]), 289)
        self.assertEqual(len(lookup["cells"]["B314"]["source_value"]), 302)

    def test_editable_boundaries_include_optional_workflows_and_fixed_line_ids(self):
        self.assertEqual(len(editable_cells("steel_vermiculite", "SCHEDULE")), 12000)
        self.assertIn("L1009", editable_cells("steel_vermiculite", "SCHEDULE"))
        self.assertNotIn("M10", editable_cells("steel_vermiculite", "SCHEDULE"))
        self.assertNotIn("Z10", editable_cells("steel_vermiculite", "SCHEDULE"))
        self.assertEqual(len(editable_cells("ductwork", "CALCULATOR")), 2400)
        self.assertNotIn("A11", editable_cells("ductwork", "CALCULATOR"))
        self.assertEqual(len(editable_cells("steel_board", "CALCULATOR")), 4800)
        self.assertIn("X208", editable_cells("steel_board", "CALCULATOR"))
        self.assertNotIn("Y9", editable_cells("steel_board", "CALCULATOR"))
        self.assertEqual(len(editable_cells("steel_board", "EXTRA BOARDS")), 400)
        self.assertIn("N45", editable_cells("steel_board", "EXTRA BOARDS"))
        self.assertNotIn("J6", editable_cells("steel_board", "EXTRA BOARDS"))
        self.assertFalse(editable_cells("steel_vermiculite", "THICKNESS DATA"))

    def test_settings_defaults_preserve_small_values_and_formula_backed_overrides(self):
        settings = next(sheet for sheet in self.books["steel_vermiculite"]["sheets"] if sheet["name"] == "SETTINGS")
        allowed = setting_cells("steel_vermiculite", "SETTINGS")
        self.assertEqual(len(allowed), 39)
        for address in ("D70", "D102", "D179", "D235"):
            self.assertIn(address, allowed)
            self.assertIn("formula", settings["cells"][address])
        self.assertNotIn("D40", allowed)
        self.assertEqual(settings["cells"]["D15"]["value"], 1e-8)
        self.assertNotIn("D15", allowed)
        self.assertEqual(len(setting_cells("steel_board", "SETTINGS")), 28)
        self.assertIn("B21", setting_cells("steel_board", "SETTINGS"))
        self.assertNotIn("B22", setting_cells("steel_board", "SETTINGS"))
        self.assertEqual(len(setting_cells("ductwork", "PRODUCT SETTINGS")), 11)
        self.assertNotIn("B8", setting_cells("ductwork", "PRODUCT SETTINGS"))

    def test_presentation_metadata_preserves_validation_and_source_visibility(self):
        book = self.books["steel_vermiculite"]
        self.assertEqual(book["defined_names"]["SectionList"], "SETTINGS!$BM$6:$BM$558")
        schedule = next(sheet for sheet in book["sheets"] if sheet["name"] == "SCHEDULE")
        self.assertTrue(any(column.get("hidden") == "1" and column["min"] == "26" for column in schedule["columns"]))
        self.assertTrue(any(item["formula1"] == "ProductList" and item["sqref"] == "B10:B1009" for item in schedule["validations"]))
        self.assertEqual(book["schedule"]["columns"][1]["type"], "select")
        self.assertTrue(book["schedule"]["columns"][1]["editable"])
        self.assertFalse(book["schedule"]["columns"][-1]["editable"])
        self.assertTrue(book["styles"]["cell_styles"])
        self.assertTrue(book["comments"])
        for identifier in ("ductwork", "steel_board"):
            for source_sheet in self.books[identifier]["sheets"]:
                self.assertNotEqual(source_sheet["dimension"], "A1")
        duct = next(sheet for sheet in self.books["ductwork"]["sheets"] if sheet["name"] == "CALCULATOR")
        self.assertEqual(duct["dimension"], "A1:CL310")

    def test_structured_table_references_retain_their_exact_data_ranges(self):
        board = self.books["steel_board"]
        table = board["tables"]["tSchedule"]
        self.assertEqual(table["sheet"], "CALCULATOR")
        self.assertEqual(table["ref"], "A8:CI208")
        self.assertEqual(table["header_row_count"], 1)
        self.assertEqual(table["totals_row_count"], 0)
        self.assertIn("Active", [column["name"] for column in table["columns"]])
        self.assertIn("Quantity status", [column["name"] for column in table["columns"]])
        self.assertEqual(len(table["columns"]), 87)

    def test_loaded_models_and_index_are_independent_from_cached_defaults(self):
        book = load_workbook_catalog("ductwork")
        book["sheets"][0]["cells"]["A1"]["value"] = "Changed by one estimate"
        book["pages"].clear()
        fresh = load_workbook_catalog("ductwork")
        self.assertNotEqual(fresh["sheets"][0]["cells"]["A1"]["value"], "Changed by one estimate")
        self.assertTrue(fresh["pages"])
        index = list_workbook_catalogs()
        index.clear()
        self.assertEqual(len(list_workbook_catalogs()), 3)

    def test_identifiers_and_ranges_are_bounded(self):
        for invalid in ("../ductwork", "unknown", None):
            with self.assertRaises(ValidationError):
                load_workbook_catalog(invalid)
        self.assertEqual(list(range_addresses("$Z$2:$AA$3")), ["Z2", "AA2", "Z3", "AA3"])
        self.assertEqual(column_name(16384), "XFD")
        self.assertEqual(column_number("XFD"), 16384)
        for invalid in ("A0", "A2:A1", "B1:A1", "A1:XFD1048576", "XFE1", "A1;B1", "A1048577"):
            with self.subTest(invalid=invalid), self.assertRaises(ValidationError):
                list(range_addresses(invalid))

    def test_gzip_packaging_is_deterministic_and_round_trips_exact_values(self):
        data = {"id": "test", "value": 0.05128205128205128, "empty": "", "blank": None,
                "formula": 'IF(A1="", "", A1*1.15)', "text": "m²"}
        with tempfile.TemporaryDirectory() as directory:
            path = write_catalog(data, directory)
            first = path.read_bytes()
            write_catalog(deepcopy(data), directory)
            self.assertEqual(path.read_bytes(), first)
            self.assertEqual(json.loads(gzip.decompress(first)), data)


class WorkbookSourceRegressionTests(unittest.TestCase):
    def test_full_source_extraction_matches_every_packaged_cell_formula_and_metadata(self):
        directory = Path(os.environ.get("ESTIMATOR_CALCULATOR_SOURCE_DIR", str(DEFAULT_SOURCE_DIRECTORY)))
        if not all((directory / spec["filename"]).is_file() for spec in CATALOG_SPECS.values()):
            self.skipTest("Original three calculator workbooks are unavailable on this host.")
        for identifier, spec in CATALOG_SPECS.items():
            with self.subTest(identifier=identifier):
                path = directory / spec["filename"]
                before = hashlib.sha256(path.read_bytes()).hexdigest()
                actual = extract_calculator(path, identifier)
                with gzip.open(DATA_DIRECTORY / f"{identifier}.json.gz", "rt", encoding="utf-8") as source:
                    expected = json.load(source)
                self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)
                comparison = compare_catalogs(expected, actual)
                self.assertTrue(comparison["matches"], format_catalog_difference(comparison))

    def test_shared_formula_boundaries_match_independent_openpyxl_translation(self):
        from openpyxl.formula.translate import Translator
        directory = Path(os.environ.get("ESTIMATOR_CALCULATOR_SOURCE_DIR", str(DEFAULT_SOURCE_DIRECTORY)))
        if not all((directory / spec["filename"]).is_file() for spec in CATALOG_SPECS.values()):
            self.skipTest("Original three calculator workbooks are unavailable on this host.")
        compared = 0
        for identifier, specification in CATALOG_SPECS.items():
            source = SourceWorkbook(directory / specification["filename"])
            try:
                for name, path in source.sheet_paths.items():
                    root = ET.fromstring(source.archive.read(path))
                    masters = {}
                    followers = {}
                    for node in root.findall("x:sheetData/x:row/x:c", NS):
                        formula = node.find("x:f", NS)
                        if formula is None or formula.attrib.get("t") != "shared":
                            continue
                        key = formula.attrib["si"]
                        if formula.text:
                            masters[key] = (node.attrib["r"], formula.text)
                        else:
                            followers.setdefault(key, []).append(node.attrib["r"])
                    if not followers:
                        continue
                    expanded = source.sheet(name)
                    for key, addresses in followers.items():
                        origin, formula = masters[key]
                        for target in {addresses[0], addresses[len(addresses)//2], addresses[-1]}:
                            expected = Translator("=" + formula, origin=origin).translate_formula(target)[1:]
                            self.assertEqual(expanded[target]["formula"], expected, (identifier, name, target))
                            compared += 1
            finally:
                source.archive.close()
        self.assertGreater(compared, 50)


if __name__ == "__main__":
    unittest.main()
